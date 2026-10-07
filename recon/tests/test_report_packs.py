"""Reports grouped by reader + the in-app viewer (UX phase 3, track B).

The #reports hub and the #report/<key> viewer read ONE registry (RPT_PACKS /
RPT_GROUPS / RPTS in recon/frontend/js/45-report-viewer.js). A registry entry
pointing at a URL the backend does not serve shows a blank iframe with every
other test green — these pin it:
  * every /api path in the registry is a GET route of the app;
  * every pack / group holds at least one entry; every screen is a route() target;
  * (node) the param builders put the month / branch / fixed extras in the query;
  * the two server reports this track added (branch-errors.print,
    accepted-variance-year.print/.xlsx) render and agree with their JSON siblings.

Run:  py -3.12 -m pytest recon/tests/test_report_packs.py
"""
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import recon.main as main                     # noqa: E402
import recon.core as core                     # noqa: E402
from test_recon import _make_db, DEMO_POS     # noqa: E402
from test_recon_routes import _routes         # noqa: E402

FE = Path(main.__file__).resolve().parent / "frontend"
VIEWER = FE / "js" / "45-report-viewer.js"


def _registry_src():
    src = VIEWER.read_text(encoding="utf-8")
    m = re.search(r"RPT-REGISTRY:BEGIN(.*?)RPT-REGISTRY:END", src, flags=re.S)
    assert m, "registry markers missing in 45-report-viewer.js"
    return m.group(1)


def _get_route_patterns():
    pats = []
    for path, methods in _routes():
        if "GET" in methods:
            pats.append(re.compile("^" + re.sub(r"\{[^}]+\}", "[^/]+", re.escape(path).replace(r"\{", "{").replace(r"\}", "}")) + "$"))
    return pats


# ── static: registry ↔ backend route table ───────────────────────────────────
def test_every_registry_path_is_a_backend_get_route():
    reg = _registry_src()
    paths = set()
    for raw in re.findall(r"['`](/api/[^'`?]+)", reg):
        paths.add(re.sub(r"\$\{[^}]*\}", "X", raw))     # `/api/sessions/${…}/export.xlsx` → a path param
    assert len(paths) >= 20, sorted(paths)
    pats = _get_route_patterns()
    missing = sorted(p for p in paths if not any(rx.match(p) for rx in pats))
    assert not missing, f"registry links with no GET route: {missing}"


def test_every_pack_and_group_has_an_entry_and_every_entry_a_known_pack():
    reg = _registry_src()
    packs = re.findall(r"\{key:'(\w+)', icon:'[\w-]+', title:'[^']+', reader:", reg)
    groups = re.findall(r"\{key:'(\w+)', icon:'[\w-]+', title:'[^']+', note:", reg)
    assert packs == ["branch", "manager", "acct", "auditor"], packs
    assert groups == ["tools", "diag"], groups
    used = re.findall(r"pack:'(\w+)'", reg)
    for k in packs + groups:
        assert used.count(k) >= 1, f"pack/group {k} has no entry"
    assert set(used) <= set(packs + groups), set(used) - set(packs + groups)


def test_every_registry_screen_is_routed_and_diagnostics_left_the_nav():
    reg = _registry_src()
    core_js = (FE / "js" / "00-core.js").read_text(encoding="utf-8")
    for scr in re.findall(r"screen:'#(\w+)'", reg):
        assert f"h==='{scr}'" in core_js, f"#{scr} is not handled by route()"
    assert "h.startsWith('report/')" in core_js and "loadReportViewer(" in core_js
    nav = re.search(r"const _NAV_TABS=\[(.*?)\];", core_js, flags=re.S).group(1)
    assert "'#board'" not in nav and "วินิจฉัย" not in nav        # ตรวจระบบ: from รายงาน + Admin now
    assert "rptDiagCardHtml()" in (FE / "js" / "80-setup.js").read_text(encoding="utf-8")
    idx = (FE / "index.html").read_text(encoding="utf-8")
    assert idx.index('js/40-reports.js') < idx.index('js/45-report-viewer.js') < idx.index('js/50-daybook')


# ── node harness: the param builders ─────────────────────────────────────────
_HARNESS = r"""
const fs=require('fs'), vm=require('vm');
const ctx=vm.createContext({console, window:{}, encodeURIComponent});
vm.runInContext(fs.readFileSync(process.argv[1],'utf8'), ctx);
const out=vm.runInContext(`(()=>{const r={};for(const k of Object.keys(RPTS)){
  r[k]={pack:RPTS[k].pack, branch:RPTS[k].branch||'', screen:RPTS[k].screen||'', session:!!RPTS[k].session,
        all:rptLinks(k,'2026-06','',''), one:rptLinks(k,'2026-06','4',''), ses:rptLinks(k,'2026-06','4','17')};}
  return JSON.stringify(r);})()`, ctx);
process.stdout.write(out);
"""


@pytest.fixture(scope="module")
def built():
    node = shutil.which("node")
    if not node:
        pytest.skip("node not installed")
    r = subprocess.run([node, "-e", _HARNESS, str(VIEWER)], capture_output=True, text=True,
                       encoding="utf-8", timeout=60)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout)


def _q(url):
    from urllib.parse import urlsplit, parse_qs
    return {k: v[0] for k, v in parse_qs(urlsplit(url).query).items()}


def test_builders_put_month_and_branch_in_the_query(built):
    for k, e in built.items():
        for which in ("all", "one"):
            for kind in ("url", "xlsx", "zip"):
                u = e[which][kind]
                if not u or k == "codes":
                    continue
                if e["session"]:
                    continue
                q = _q(u)
                assert q.get("period_month") == "2026-06", (k, which, u)
                want_branch = which == "one" and e["branch"] in ("opt", "req")
                assert ("branch_id" in q) == want_branch, (k, which, u)
                if want_branch:
                    assert q["branch_id"] == "4", (k, u)


def test_builders_fixed_extras_and_special_entries(built):
    assert _q(built["closereg"]["all"]["url"])["months"] == "6"
    assert _q(built["aging"]["one"]["xlsx"])["lookback"] == "2"
    assert _q(built["accepted"]["all"]["url"])["threshold"] == "1000"
    assert _q(built["acceptedyear"]["one"]["xlsx"])["threshold"] == "1000"
    assert _q(built["matchdetail"]["one"]["url"])["detail"] == "1"
    assert built["codes"]["one"]["url"] == "/api/statement-codes.html"          # not month/branch scoped
    assert built["codes"]["one"]["xlsx"] == "/api/statement-codes.xlsx"
    assert built["auditorpack"]["one"]["zip"].startswith("/api/reports/auditor-pack.zip?")
    # the session download only exists once the session is known
    assert built["closexlsx"]["one"]["xlsx"] == ""
    assert built["closexlsx"]["ses"]["xlsx"] == "/api/sessions/17/export.xlsx"
    # every entry is a document (url/xlsx/zip) or an in-app screen
    for k, e in built.items():
        assert e["one"]["url"] or e["ses"]["xlsx"] or e["one"]["zip"] or e["screen"], k


# ── backend: the two reports added for the packs ─────────────────────────────
PM = "2026-06"


@pytest.fixture
def client(tmp_path):
    core._DB_PATH = main._DB_PATH = _make_db(tmp_path, DEMO_POS)
    os.environ.pop("OSHINEI_SEED_TOKEN", None)
    with TestClient(main.app) as cl:
        yield cl


def _sql(sql, *args):
    c = sqlite3.connect(main._DB_PATH)
    c.execute(sql, args)
    c.commit()
    c.close()


def test_branch_errors_print_month_branch_and_escaped(client):
    _sql("INSERT INTO recon_lane_adj(branch_id,period_month,category,bank,note,adj_type,created_by) "
         "VALUES(4,?,'QR','ALL','<script>x</script> พนักงานกดผิด','accept','mgr')", PM)
    _sql("INSERT INTO recon_lane_adj(branch_id,period_month,category,bank,note,adj_type,created_by) "
         "VALUES(4,'2026-05','CASH','ALL','เดือนก่อน','accept','mgr')")
    r = client.get("/api/reports/branch-errors.print", params={"period_month": PM})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
    assert "พนักงานกดผิด" in r.text and "<script>x</script>" not in r.text
    assert "เดือนก่อน" not in r.text                                    # month-scoped
    assert "class='toolbar'" in r.text                                   # the viewer hides it in-app
    one = client.get("/api/reports/branch-errors.print", params={"period_month": PM, "branch_id": 4})
    assert one.status_code == 200 and "พนักงานกดผิด" in one.text
    # the JSON / xlsx take the same branch filter; omitted = unchanged (all branches)
    assert client.get("/api/reports/branch-errors", params={"branch_id": 999}).json()["items"] == []
    assert len(client.get("/api/reports/branch-errors").json()["items"]) == 2
    assert client.get("/api/reports/branch-errors.print").status_code == 422   # month required


def test_accepted_variance_year_rolls_up_months(client):
    for pm in ("2026-05", PM):
        _sql("INSERT INTO recon_day_ack(branch_id,period_month,category,day,variance,note,created_by) "
             "VALUES(4,?,'QR',?,-250,'ตรวจแล้ว','t')", pm, pm + "-10")
    month = client.get("/api/reports/accepted-variance", params={"period_month": PM}).json()
    r = client.get("/api/reports/accepted-variance-year.print", params={"period_month": PM})
    assert r.status_code == 200 and "2026-01 ถึง 2026-06" in r.text
    x = client.get("/api/reports/accepted-variance-year.xlsx", params={"period_month": PM, "branch_id": 4})
    assert x.status_code == 200 and x.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    months = [row for row in wb["รายเดือน"].iter_rows(min_row=2, values_only=True)]
    assert [m[0] for m in months] == [f"2026-{i:02d}" for i in range(1, 7)]
    june = next(m for m in months if m[0] == PM)
    assert june[1] == month["count"] and round(june[2]) == round(month["grand_total_abs"])
    may = next(m for m in months if m[0] == "2026-05")
    assert may[1] == 1                                                  # the earlier month counts too
    assert len(list(wb["รายการทั้งหมด"].iter_rows(min_row=2))) == month["count"] + 1
    bad = client.get("/api/reports/accepted-variance-year.print", params={"period_month": "2026-6"})
    assert bad.status_code == 400
