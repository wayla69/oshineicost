"""Golden-file regression for the recon app — every GET that answers on the
fixture must return identical output (Phase 7, 2026-09).

Same idea as backend/test_golden.py: recon/main.py (7,270 lines, 116
endpoints) is about to be split into a core module and routers by family, and
this is what makes "nothing changed" checkable. The fixture is the one the
other recon tests use (test_recon._make_db with DEMO_POS); the parameter
defaults are what that fixture holds (day 2026-05-01, branch 1).

What is pinned depends on the body:
  JSON      canonical dump (floats to 6 dp, write clocks masked, record lists as sets)
  HTML      the text, with dates/times masked (the print views stamp "printed at")
  xlsx      sheet names and dimensions (the bytes carry a creation timestamp)
  zip       the member names
  other     content-type and length

  regenerate:  GOLDEN_UPDATE=1 python -m pytest recon/tests/test_recon_golden.py
"""
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

# same bootstrap as test_recon.py: repo root for `recon.*`, this folder for the fixture helpers
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import recon.main as main                      # noqa: E402
import recon.core as core   # get_db() reads core._DB_PATH; main._DB_PATH is a copy
from test_recon import _make_db, DEMO_POS      # noqa: E402

GOLDEN = Path(__file__).parent / "golden"
UPDATE = os.environ.get("GOLDEN_UPDATE") == "1"
_TS_KEYS = {"updated_at", "created_at", "issued_at", "generated_at", "captured_at", "uploaded_at",
            "closed_at", "signed_at", "acked_at", "ts", "server_time", "now", "printed_at"}
_TS_RE = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?|\d{1,2}/\d{1,2}/\d{4} \d{2}:\d{2}")

# (path, params) — every parameter-free or default-parameterised GET that
# answers 200 on the fixture. Generated 2026-09-04; extend when a route is added.
SMOKE = [
    ('/api/_diag/auth-wiring', {}),
    ('/api/account-map', {}),
    ('/api/acquirer-mids', {}),
    ('/api/attachments', {'branch_id': 1, 'day': '2026-05-01'}),
    ('/api/audit', {}),
    ('/api/bank-accounts', {}),
    ('/api/branches', {}),
    ('/api/channel-move', {'branch_id': 1, 'period_month': '2026-05'}),
    ('/api/daily-accounts', {'period_month': '2026-05'}),
    ('/api/daily-accounts.print', {'period_month': '2026-05'}),
    ('/api/daily-board/dates', {}),
    ('/api/daily-dashboard', {'day': '2026-05-01'}),
    ('/api/day-ack', {'branch_id': 1, 'period_month': '2026-05'}),
    ('/api/deposit-slips', {'branch_id': 1, 'period_month': '2026-05'}),
    ('/api/diagnostics/customer-types', {'period_month': '2026-05'}),
    ('/api/diagnostics/payment-codes', {'period_month': '2026-05'}),
    ('/api/diagnostics/reclass-pattern', {'month': '2026-05'}),
    ('/api/lanes/adj', {'branch_id': 1, 'period_month': '2026-05'}),
    ('/api/map-suggestions', {'branch_id': 1, 'period_month': '2026-05'}),
    ('/api/method-config', {}),
    ('/api/monthly-report', {'period_month': '2026-05', 'day': '2026-05-01'}),
    ('/api/overview', {'period_month': '2026-05'}),
    ('/api/payment-map', {}),
    ('/api/reports/accepted-variance', {'period_month': '2026-05'}),
    ('/api/reports/accepted-variance.print', {'period_month': '2026-05'}),
    ('/api/reports/accepted-variance.xlsx', {'period_month': '2026-05'}),
    ('/api/reports/audit-trail', {'period_month': '2026-05'}),
    ('/api/reports/audit-trail.print', {'period_month': '2026-05'}),
    ('/api/reports/audit-trail.xlsx', {'period_month': '2026-05'}),
    ('/api/reports/auditor-pack.zip', {'period_month': '2026-05'}),
    ('/api/reports/branch-errors', {}),
    ('/api/reports/branch-errors.xlsx', {}),
    ('/api/reports/close-register', {'period_month': '2026-05'}),
    ('/api/reports/close-register.print', {'period_month': '2026-05'}),
    ('/api/reports/close-register.xlsx', {'period_month': '2026-05'}),
    ('/api/reports/open-items-aging', {'period_month': '2026-05'}),
    ('/api/reports/open-items-aging.print', {'period_month': '2026-05'}),
    ('/api/reports/open-items-aging.xlsx', {'period_month': '2026-05'}),
    ('/api/reports/outstanding-docs', {'period_month': '2026-05'}),
    ('/api/reports/outstanding-docs.print', {'period_month': '2026-05'}),
    ('/api/reports/override-register', {'period_month': '2026-05'}),
    ('/api/reports/override-register.print', {'period_month': '2026-05'}),
    ('/api/reports/override-register.xlsx', {'period_month': '2026-05'}),
    ('/api/reports/recon-certificate.print', {'period_month': '2026-05'}),
    ('/api/reports/wrong-channel', {}),
    ('/api/reports/wrong-channel.print', {}),
    ('/api/reports/wrong-channel.xlsx', {}),
    ('/api/sessions', {}),
    ('/api/setup/master', {'period_month': '2026-05'}),
    ('/api/statement-codes', {}),
    ('/api/statement-codes.html', {}),
    ('/api/statement-codes.xlsx', {}),
    ('/api/uat-signoff', {}),
    ('/api/uploads', {}),
    ('/api/v2/overview', {'month': '2026-05'}),
    # before the first sync these three answer 'not available yet' (Phase 7a) — pinned as such
    ('/api/daily-board', {'day': '2026-05-01'}),
    ('/api/daily-report', {'day': '2026-05-01'}),
    ('/api/diagnostics/excel-payment-cols', {'period_month': '2026-05'}),
]


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    core._DB_PATH = main._DB_PATH = _make_db(tmp_path_factory.mktemp("golden"), DEMO_POS)
    # The open-items aging report buckets by age from TODAY, so its goldens
    # turned red at midnight on 2026-09-05: pin the day they were generated.
    import datetime
    import recon.routers.reports as reports
    mp = pytest.MonkeyPatch()
    mp.setattr(reports, "_today", lambda: datetime.date(2026, 9, 4))
    try:
        with TestClient(main.app) as cl:
            yield cl
    finally:
        mp.undo()


def _slug(path, params):
    base = re.sub(r"[^a-z0-9]+", "_", path.replace("/api/", "")).strip("_")
    if params:
        base += "__" + re.sub(r"[^A-Za-z0-9._-]+", "_", "_".join(f"{k}-{v}" for k, v in sorted(params.items())))
    return base[:150] + ".txt"


def _canon(obj):
    if isinstance(obj, float):
        return round(obj, 6)
    if isinstance(obj, dict):
        return {k: ("<ts>" if k in _TS_KEYS and isinstance(v, (str, int, float)) else _canon(v))
                for k, v in obj.items()}
    if isinstance(obj, list):
        items = [_canon(v) for v in obj]
        if items and all(isinstance(v, dict) for v in items):
            items.sort(key=lambda v: json.dumps(v, sort_keys=True, ensure_ascii=False))
        return items
    return obj


def _pin(r) -> str:
    ct = (r.headers.get("content-type") or "").lower()
    if "json" in ct:
        return json.dumps(_canon(r.json()), ensure_ascii=False, sort_keys=True, indent=1) + "\n"
    if "html" in ct or "text/" in ct:
        return _TS_RE.sub("<ts>", r.text)
    if "spreadsheet" in ct or "openxml" in ct:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        return "\n".join(f"{ws.title}: {ws.max_row}x{ws.max_column}" for ws in wb.worksheets) + "\n"
    if "zip" in ct:
        return "\n".join(sorted(zipfile.ZipFile(io.BytesIO(r.content)).namelist())) + "\n"
    return f"{ct} {len(r.content)} bytes\n"


@pytest.mark.parametrize("path,params", SMOKE, ids=[_slug(p, q)[:-4] for p, q in SMOKE])
def test_golden(client, path, params):
    r = client.get(path, params=params)
    assert r.status_code == 200, r.text[:300]
    got = _pin(r)
    f = GOLDEN / _slug(path, params)
    if UPDATE or not f.exists():
        GOLDEN.mkdir(exist_ok=True)
        f.write_text(got, encoding="utf-8")
        if not UPDATE:
            pytest.skip(f"golden file created: {f.name}")
        return
    want = f.read_text(encoding="utf-8")
    if got != want:
        gl, wl = got.splitlines(), want.splitlines()
        for i, (a, b) in enumerate(zip(gl, wl)):
            if a != b:
                pytest.fail(f"{f.name} differs at line {i + 1}:\n  golden: {b[:160]}\n  now:    {a[:160]}")
        pytest.fail(f"{f.name}: length differs ({len(wl)} -> {len(gl)} lines)")
