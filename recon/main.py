"""
Revenue Reconciliation API
Second service in the same Fly app as the main dashboard.
Shares /data/oshinei.db via the shared Fly volume.

Run locally:
    uvicorn recon.main:app --port 8001 --reload

Environment:
    OSHINEI_DB   shared database path  (default: /data/oshinei.db)
"""

from __future__ import annotations
import hashlib
import hmac
import json
import os
import re
import sqlite3
import uuid
from contextlib import contextmanager
from pathlib import Path
from datetime import datetime as _dt, timedelta as _td
from typing import Literal

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from .matcher import (
    TOL_ABS, TOL_PCT, TOL_ABS_DAY, TOL_PCT_DAY,
    _eff_fee_vat,
    bank_statement_coverage,
    cash_three_way,
    category_reconciliation,
    coupon_reconciliation,
    daily_board,
    daily_reconciliation,
    lane_ext_lines,
    lane_lines,
    lane_pos_days,
    pos_lane_lines,
    map_bank_suggestions,
    upload_match_stats,
    payout_matching,
    payout_reconciliation,
    point_pos_total_daily,
    coupon_discount_total_daily,
    pos_codes_scope,
    run_auto_match,
    suggest_reclassifications,
    _BANK_LABEL,
)
from .reconcile_v2 import reconcile_month, reconcile_overview, _CH_LABEL
from .seed_payment_map import seed_payment_map
from .parsers import (
    parse_bank_statement,
    parse_cc_settlement,
    parse_grab_report,
    parse_lineman_report,
)
from .parsers.bank import STATEMENT_CODE_MAP, _normalise_channel
from .parsers.loader import (to_parseable_xlsx, is_encrypted_upload as _is_encrypted_upload,
                             unlock_pdf as _unlock_pdf)
from .parsers.pdf_formats import (
    detect_and_parse as _detect_pdf_format,
    fee_invoice_by_name as _fee_invoice_by_name,
    fee_invoice_by_text as _fee_invoice_by_text,
    unsupported_bank_pdf_message as _unsupported_bank_pdf_message,
)

# ── Setup ─────────────────────────────────────────────────────────────────────

_DB_PATH    = Path(os.getenv("OSHINEI_DB", "/data/oshinei.db"))
_SCHEMA     = Path(__file__).parent / "schema.sql"
_ATTACH_DIR = Path(os.getenv("OSHINEI_ATTACH_DIR", str(_DB_PATH.parent / "attachments")))
_SNAPSHOT_DIR = Path(os.getenv("OSHINEI_SNAPSHOT_DIR", str(_DB_PATH.parent / "snapshots")))

app = FastAPI(title="Oshinei Revenue Reconciliation", version="1.0.0")
# The frontend is served same-origin (/app on this service), so it needs no CORS
# grant; restrict cross-origin reads to the known hosts instead of a blanket '*'.
# Override with OSHINEI_CORS_ORIGINS (comma-separated) if another origin must call in.
_CORS_ORIGINS = [o.strip() for o in os.getenv(
    "OSHINEI_CORS_ORIGINS",
    "https://oshinei-dashboard.fly.dev,https://oshinei-dashboard.fly.dev:8443,"
    "http://localhost:8001,http://127.0.0.1:8001,http://localhost:8013,http://127.0.0.1:8013"
).split(",") if o.strip()]
app.add_middleware(CORSMiddleware, allow_origins=_CORS_ORIGINS, allow_methods=["*"], allow_headers=["*"])


@app.middleware("http")
async def _frontend_no_cache(request, call_next):
    """The frontend is a single inline-HTML file served by StaticFiles, which
    sends only last-modified/etag (no Cache-Control). Browsers then apply
    heuristic caching and may serve a stale build even after a refresh. Force
    revalidation on /app so a fresh deploy is visible immediately (the etag still
    yields a cheap 304 when nothing changed)."""
    resp = await call_next(request)
    if request.url.path.startswith("/app"):
        resp.headers["Cache-Control"] = "no-cache, must-revalidate"
    return resp


@contextmanager
def get_db():
    conn = sqlite3.connect(_DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        yield conn
    finally:
        conn.close()


# Idempotent column adds for DBs created before a schema column existed
# (CREATE TABLE IF NOT EXISTS never alters an existing table).
_MIGRATIONS = {
    "recon_bank_lines":    [("account_no", "TEXT")],
    "recon_cc_lines":      [("acquirer", "TEXT"), ("fee_manual", "REAL"), ("vat_manual", "REAL"),
                            ("fee_vat_by", "TEXT"), ("fee_vat_at", "TEXT"), ("vat_amount", "REAL")],
    "recon_grab_lines":    [("fee_manual", "REAL"), ("vat_manual", "REAL"),
                            ("fee_vat_by", "TEXT"), ("fee_vat_at", "TEXT"), ("vat_amount", "REAL")],
    "recon_lineman_lines": [("fee_manual", "REAL"), ("vat_manual", "REAL"),
                            ("fee_vat_by", "TEXT"), ("fee_vat_at", "TEXT")],
    "recon_discrepancies": [("category", "TEXT"), ("bank", "TEXT"), ("assigned_to", "TEXT")],
    "recon_uploads":       [("file_hash", "TEXT")],
    # Ties the frozen working paper to the exact source documents (name + sha256) that
    # produced it, so a later upload delete/replace can't silently break the chain.
    "recon_close_snapshot": [("source_manifest", "TEXT")],
}


def _migrate(conn):
    for table, cols in _MIGRATIONS.items():
        try:
            existing = {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
        except sqlite3.OperationalError:
            continue
        if not existing:
            continue
        for name, decl in cols:
            if name not in existing:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {name} {decl}")


def _migrate_payment_bank(conn):
    """Widen recon_payment_map.bank to allow BAY/SCB/BBL + 'ANY' (multi-bank
    codes like the generic 'Bank Transfer' — the CHECK constraint can't be
    ALTERed in SQLite → recreate the table once, preserving rows). Then
    future-proof the franchise-only 'Credit BAY' tender (seed_payment_map is
    master-scoped so it never picks it up)."""
    try:
        row = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='recon_payment_map'").fetchone()
        if row and any(f"'{b}'" not in (row[0] or "") for b in ("BAY", "ANY", "TTB")):
            # Explicit transaction (NOT executescript): a CHECK failure mid-script
            # must roll the whole recreate back, never leave the table half-gone.
            # (Live data taught us this — it contained bank='TTB'.)
            conn.commit()
            conn.execute("BEGIN IMMEDIATE")
            try:
                conn.execute("""
                    CREATE TABLE recon_payment_map__new (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        payment_code_norm TEXT NOT NULL UNIQUE,
                        category TEXT NOT NULL CHECK(category IN
                            ('CASH','QR','TRANSFER','CREDIT','GRAB','LINEMAN','COUPON','ENTERTAIN','OTHER')),
                        bank TEXT NOT NULL DEFAULT 'NONE' CHECK(bank IN
                            ('KTB','KBANK','KTC','BAY','SCB','BBL','TTB','PROMPTPAY','ANY','NONE','UNKNOWN')),
                        reconcile INTEGER NOT NULL DEFAULT 1,
                        notes TEXT, updated_at TEXT DEFAULT (datetime('now')), updated_by TEXT
                    )""")
                conn.execute("""
                    INSERT INTO recon_payment_map__new
                        (id,payment_code_norm,category,bank,reconcile,notes,updated_at,updated_by)
                        SELECT id,payment_code_norm,category,bank,reconcile,notes,updated_at,updated_by
                        FROM recon_payment_map""")
                conn.execute("DROP TABLE recon_payment_map")
                conn.execute("ALTER TABLE recon_payment_map__new RENAME TO recon_payment_map")
                conn.execute("COMMIT")
            except Exception:
                conn.execute("ROLLBACK")
                raise
        conn.execute("""INSERT INTO recon_payment_map (payment_code_norm,category,bank,reconcile,updated_by)
            VALUES ('Credit BAY','CREDIT','BAY',1,'seed-franchise')
            ON CONFLICT(payment_code_norm) DO NOTHING""")
        # widen recon_acquirer_mid.acquirer to allow BAY/SCB/BBL + point acquirers
        arow = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='recon_acquirer_mid'").fetchone()
        if arow and "KTCPOINT" not in (arow[0] or ""):
            conn.executescript("""
                CREATE TABLE recon_acquirer_mid__new (
                    mid_id INTEGER PRIMARY KEY AUTOINCREMENT,
                    branch_id INTEGER NOT NULL REFERENCES branches(branch_id),
                    acquirer TEXT NOT NULL CHECK(acquirer IN
                        ('KTB','KBANK','KTC','BAY','SCB','BBL','KTCPOINT','BAYPOINT')),
                    merchant_id TEXT, terminal_id TEXT, active INTEGER NOT NULL DEFAULT 1
                );
                INSERT INTO recon_acquirer_mid__new (mid_id,branch_id,acquirer,merchant_id,terminal_id,active)
                    SELECT mid_id,branch_id,acquirer,merchant_id,terminal_id,active FROM recon_acquirer_mid;
                DROP TABLE recon_acquirer_mid;
                ALTER TABLE recon_acquirer_mid__new RENAME TO recon_acquirer_mid;
            """)
    except sqlite3.OperationalError:
        pass


def _migrate_lane_adj_accept(conn):
    """Widen recon_lane_adj.adj_type to allow 'accept' (v2 acknowledged
    variance). SQLite can't ALTER a CHECK → recreate once, preserving rows."""
    try:
        row = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='recon_lane_adj'").fetchone()
        if row and "'accept'" not in (row[0] or ""):
            conn.executescript("""
                CREATE TABLE recon_lane_adj__new (
                    id           INTEGER PRIMARY KEY AUTOINCREMENT,
                    branch_id    INTEGER NOT NULL REFERENCES branches(branch_id),
                    period_month TEXT    NOT NULL,
                    category     TEXT    NOT NULL,
                    bank         TEXT    NOT NULL DEFAULT 'NONE',
                    adj_pos      REAL    NOT NULL DEFAULT 0,
                    adj_ext      REAL    NOT NULL DEFAULT 0,
                    note         TEXT,
                    adj_type     TEXT    NOT NULL DEFAULT 'other'
                                     CHECK(adj_type IN ('off_pos','deposit_note','other','accept')),
                    created_by   TEXT,
                    created_at   TEXT    DEFAULT (datetime('now')),
                    UNIQUE(branch_id, period_month, category, bank)
                );
                INSERT INTO recon_lane_adj__new
                    (id,branch_id,period_month,category,bank,adj_pos,adj_ext,note,adj_type,created_by,created_at)
                    SELECT id,branch_id,period_month,category,bank,adj_pos,adj_ext,note,adj_type,created_by,created_at
                    FROM recon_lane_adj;
                DROP TABLE recon_lane_adj;
                ALTER TABLE recon_lane_adj__new RENAME TO recon_lane_adj;
                CREATE INDEX IF NOT EXISTS idx_rla_scope ON recon_lane_adj(branch_id, period_month);
            """)
    except sqlite3.OperationalError:
        pass


def _migrate_attachment_to_fs(conn):
    """Replace recon_attachment.data (BLOB) with file_path (TEXT).
    Files uploaded before this migration are lost (acceptable — feature was unused).
    Idempotent: skips if file_path already present."""
    try:
        existing = {r[1] for r in conn.execute("PRAGMA table_info(recon_attachment)")}
    except sqlite3.OperationalError:
        return
    if "file_path" in existing:
        return  # already migrated
    conn.executescript("""
        CREATE TABLE recon_attachment__new (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            branch_id   INTEGER NOT NULL,
            day         TEXT NOT NULL,
            category    TEXT,
            file_name   TEXT NOT NULL,
            mime        TEXT,
            size_bytes  INTEGER,
            file_path   TEXT,
            note        TEXT,
            uploaded_by TEXT,
            uploaded_at TEXT
        );
        INSERT INTO recon_attachment__new
            (id,branch_id,day,category,file_name,mime,size_bytes,note,uploaded_by,uploaded_at)
            SELECT id,branch_id,day,category,file_name,mime,size_bytes,note,uploaded_by,uploaded_at
            FROM recon_attachment;
        DROP TABLE recon_attachment;
        ALTER TABLE recon_attachment__new RENAME TO recon_attachment;
        CREATE INDEX IF NOT EXISTS idx_recon_attachment ON recon_attachment(branch_id, day);
    """)


# ── Requirement-table seeds (owner's "ตารางจับคู่รหัสรับเงิน", 2026-06) ─────────
# Both seeds are idempotent and NEVER overwrite a staff edit (guarded on
# updated_by/ON CONFLICT DO NOTHING). Re-runnable from /api/setup/seed-requirement.

# Lane defaults encode the owner's table. Parameter values = the engine's current
# constants (so seeding alone changes no live result) EXCEPT cash tol_abs_day=10
# (owner decision 2026-06-12: deposit slips may differ by a few baht — pass ≤฿10
# but keep the variance visible).
_METHOD_DEFAULTS = [
    # (category, bank, doc1, doc2, method, tol_abs, tol_abs_day, lag, merge, fee_vat, notes)
    ("CASH", "*", "deposit_slip", "bank_stmt_*", "cash_3way", None, 10.0, 3, 4, 0,
     "เงินสดของวัน D ฝากวันถัดไปก่อน 13:00 — เทียบ POS ↔ สลิปนำฝาก ↔ สเตทเมนต์ (ต่างได้หลักบาท)"),
    # itemized lanes: lag_days/merge_days seeded for display in Setup UI only —
    # the itemized matcher uses daily exact-match, not lag/merge grouping.
    ("QR", "KBANK", None, "bank_stmt_KBANK", "itemized", None, None, 1, 1, 0,
     "QR พร้อมเพย์ (K-SHOP/KNNQR/QR code K-bank) — เงินเข้าวันเดียวกันหรือวันถัดไป แยกรายการตาม POS"),
    ("TRANSFER", "KTB", None, "bank_stmt_KTB", "itemized", None, None, 1, 1, 0,
     "ลูกค้าโอนเข้า KTB — เข้าวันเดียวกันหรือวันถัดไป แยกรายการตาม POS"),
    ("TRANSFER", "KBANK", None, "bank_stmt_KBANK", "itemized", None, None, 1, 1, 0,
     "ลูกค้าโอนเข้า KBANK — เข้าวันเดียวกันหรือวันถัดไป แยกรายการตาม POS"),
    ("TRANSFER", "ANY", None, "bank_stmt_*", "itemized", None, None, 1, 1, 0,
     "Bank Transfer ทั่วไป — เข้าได้ทั้ง KTB และ KBANK ตรวจกับทั้งสองสเตทเมนต์"),
    ("CREDIT", "KBANK", "cc_summary", "bank_stmt_KBANK", "two_stage", None, None, 7, 4, 1,
     "ขั้น 1: สรุปยอด KBANK ↔ สเตทเมนต์ KBANK (อาจรวมหลายวัน) · ขั้น 2: gross ก่อนค่าธรรมเนียม+VAT ↔ POS"),
    ("CREDIT", "KTB", "cc_summary", "bank_stmt_KTB", "two_stage", None, None, 7, 4, 1,
     "ขั้น 1: สรุปยอด KTC ↔ สเตทเมนต์ KTB (อาจรวมหลายวัน) · ขั้น 2: gross ก่อนค่าธรรมเนียม+VAT ↔ POS"),
    ("CREDIT", "*", "cc_summary", "bank_stmt_*", "two_stage", None, None, 7, 4, 1, None),
    ("GRAB", "*", "grab_report", "bank_stmt_KBANK", "two_stage", None, None, 7, 4, 1,
     "Grab โอนเข้า KBANK วันทำการถัดไป (อาจรวม 2 วันเป็นยอดเดียว) — สรุปมีค่าคอม+VAT"),
    ("LINEMAN", "*", "lineman_report", "bank_stmt_KBANK", "two_stage", None, None, 7, 4, 1,
     "Lineman โอนเข้า KBANK วันทำการถัดไป (อาจรวม 2 วันเป็นยอดเดียว) — สรุปมีค่าคอม+VAT"),
    ("COUPON", "KTC", "points_summary", "bank_stmt_KTB", "issuer_points", None, None, None, None, 0,
     "แต้ม KTC — ใบสรุปคะแนนจาก KTC ↔ สเตทเมนต์ KTB (อยู่ในยอดฝากรวมของบัตร KTC)"),
    ("COUPON", "BAY", "points_summary", "bank_stmt_KBANK", "issuer_points", None, None, None, None, 0,
     "แต้ม BAY (กรุงศรี) — ใบสรุปคะแนนจาก BAY ↔ สเตทเมนต์ KBANK"),
]


def _seed_method_defaults(conn) -> dict:
    inserted = 0
    for row in _METHOD_DEFAULTS:
        cur = conn.execute("""
            INSERT INTO recon_method_config
                (category, bank, doc1, doc2, method, tol_abs, tol_abs_day,
                 lag_days, merge_days, fee_vat_expected, notes, updated_by)
            VALUES (?,?,?,?,?,?,?,?,?,?,?, 'seed')
            ON CONFLICT(category, bank) DO NOTHING
        """, row)
        inserted += cur.rowcount
    return {"inserted": inserted, "skipped": len(_METHOD_DEFAULTS) - inserted}


# The 12 codes of the owner's table with their now-known routing. Inserted when
# missing; existing SEED rows with an unresolved bank are upgraded (staff edits
# are never touched — guarded on updated_by).
_REQUIREMENT_CODES = [
    ("CASH",            "CASH",     "NONE"),
    ("Cash",            "CASH",     "NONE"),
    ("Coupon",          "COUPON",   "NONE"),
    ("Credit กสิกรไทย",  "CREDIT",   "KBANK"),
    ("Credit กรุงไทย",   "CREDIT",   "KTB"),
    ("GB",              "GRAB",     "NONE"),
    ("LM",              "LINEMAN",  "NONE"),
    ("K-SHOP",          "QR",       "KBANK"),
    ("KNNQR",           "QR",       "KBANK"),
    ("QR code K-bank",  "QR",       "KBANK"),
    ("Bank Transfer",   "TRANSFER", "ANY"),
    ("โอน KTB",         "TRANSFER", "KTB"),
]
_SEED_AUTHORS = ("seed", "seed-req")


def _seed_requirement_codes(conn) -> dict:
    """Apply the owner's routing answers to un-edited seed rows + insert any of
    the 12 canonical codes still missing. Every change is audited."""
    upgraded = inserted = 0
    for code, category, bank in _REQUIREMENT_CODES:
        old = conn.execute(
            "SELECT id, category, bank, updated_by FROM recon_payment_map "
            "WHERE payment_code_norm=?", (code,)).fetchone()
        if old is None:
            cur = conn.execute("""
                INSERT INTO recon_payment_map (payment_code_norm, category, bank, reconcile, updated_by)
                VALUES (?,?,?,1,'seed-req') ON CONFLICT(payment_code_norm) DO NOTHING
            """, (code, category, bank))
            inserted += cur.rowcount
            continue
        # Upgrade seed-authored rows (any bank) whose mapping differs from the canonical value.
        if (old["updated_by"] in _SEED_AUTHORS
                and (old["category"], old["bank"]) != (category, bank)):
            conn.execute(
                "UPDATE recon_payment_map SET category=?, bank=?, "
                "updated_by='seed-req', updated_at=datetime('now') WHERE id=?",
                (category, bank, old["id"]))
            diffs = [{"field": "bank", "old": old["bank"], "new": bank}]
            if old["category"] != category:
                diffs.append({"field": "category", "old": old["category"], "new": category})
            _audit(conn, "payment_map", old["id"], "update", "seed-req", diffs)
            upgraded += 1
    return {"upgraded": upgraded, "inserted": inserted}


@app.on_event("startup")
def startup():
    _ATTACH_DIR.mkdir(parents=True, exist_ok=True)
    with get_db() as conn:
        conn.executescript(_SCHEMA.read_text(encoding="utf-8"))
        _migrate(conn)
        _migrate_payment_bank(conn)   # widen bank CHECK + future-proof Credit BAY
        _migrate_lane_adj_accept(conn)  # widen adj_type CHECK for v2 'accept'
        _migrate_attachment_to_fs(conn)
        # recon_lane_adj: created by schema.sql executescript above (IF NOT EXISTS)
        seed_payment_map(conn)   # idempotent; never overwrites staff edits
        _seed_method_defaults(conn)     # owner's per-lane method table
        _seed_requirement_codes(conn)   # owner's 12-code routing answers
        conn.commit()


# ── Static frontend ───────────────────────────────────────────────────────────

_FRONTEND = Path(__file__).parent / "frontend"
if _FRONTEND.exists():
    app.mount("/app", StaticFiles(directory=str(_FRONTEND), html=True), name="frontend")


# ── Branches (read from shared branches table) ────────────────────────────────
# Reconciliation is scoped to MASTER (company-owned) branches only; franchises
# are out of scope. COALESCE keeps legacy NULL rows (default 'master') in scope.

_MASTER_ONLY = "COALESCE(branch_type,'master')='master' AND active=1"

# Reconciliation only considers data from this month onward — the app reconciles
# the live business, not the full historical DB (which the dashboard shares).
# Floor only; nothing is deleted. Override with env RECON_START_MONTH (YYYY-MM).
RECON_START_MONTH = os.getenv("RECON_START_MONTH", "2026-05")
RECON_START_DAY   = RECON_START_MONTH + "-01"


# ── Write auth + upload validation ────────────────────────────────────────────
# Mirrors the dashboard's gate: when OSHINEI_SEED_TOKEN is set (hosted deploys) a
# matching token is required via the `X-Upload-Token` header or `?token=`. Unset
# locally → writes are open, so the gate is invisible in dev.

MAX_UPLOAD_BYTES = 15 * 1024 * 1024          # 15 MB
_ALLOWED_EXT     = (".xlsx", ".xls", ".csv", ".pdf")
_PERIOD_RE       = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
_DAY_RE          = re.compile(r"^\d{4}-(0[1-9]|1[0-2])-(0[1-9]|[12]\d|3[01])$")


def _accepted_tokens() -> dict:
    """Parse OSHINEI_SEED_TOKEN into {token: label}. Comma/newline list so several
    people can each hold a token; entries may be 'label:token'. Single token still
    works. Shared shape with the dashboard gate."""
    raw = os.getenv("OSHINEI_SEED_TOKEN", "") or ""
    out: dict = {}
    for part in raw.replace("\n", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if ":" in part:
            label, tok = part.split(":", 1)
            tok = tok.strip()
            if tok:
                out[tok] = label.strip()
        else:
            out[part] = ""
    return out

def require_write_token(token: str = Query(""), x_upload_token: str = Header(None)):
    toks = _accepted_tokens()
    if toks:
        provided = (x_upload_token or token or "").strip()
        if not any(hmac.compare_digest(provided, t) for t in toks):
            raise HTTPException(403, "A valid write token is required.")


def _require_manager(x_manager_token: str | None):
    """Two-tier close: approving a month needs the manager code when
    OSHINEI_MANAGER_TOKEN is set (open locally when unset, like the write gate)."""
    expected = (os.environ.get("OSHINEI_MANAGER_TOKEN") or "").strip()
    if expected and not hmac.compare_digest((x_manager_token or "").strip(), expected):
        raise HTTPException(403, "ต้องใช้รหัสผู้จัดการเพื่อปิดเดือน (X-Manager-Token)")


def require_manager_token(x_manager_token: str = Header(None)):
    """Dependency for advanced-setup mutations (method parameters, tolerances):
    a bad value can mass-flip lane statuses, so these need the MANAGER code —
    ordinary code↔bank mapping stays on the staff write token (owner decision
    2026-06-12). Open locally when OSHINEI_MANAGER_TOKEN is unset."""
    expected = (os.environ.get("OSHINEI_MANAGER_TOKEN") or "").strip()
    if expected and not hmac.compare_digest((x_manager_token or "").strip(), expected):
        raise HTTPException(403, "ต้องใช้รหัสผู้จัดการ (X-Manager-Token) สำหรับการตั้งค่าขั้นสูง")


# ── Trusted actor identity (control gap fix 2026-06-19) ───────────────────────
# Recon shares the machine + DB + signing secret with the dashboard and answers on
# the SAME domain (different port → cookies are not port-scoped), so the dashboard's
# signed session cookie reaches us too. Derive the acting username from it SERVER-SIDE
# instead of trusting the client-supplied created_by/matched_by/approved_by — that
# makes the audit "by" tamper-resistant when login is in use. Falls back to the
# client value (then "staff") when there is no valid session, so open/local mode and
# the test suite are unchanged.
try:
    from backend.auth import parse_token as _parse_session, SESSION_COOKIE as _SESSION_COOKIE
except Exception:                                    # recon can also run standalone
    _parse_session, _SESSION_COOKIE = None, "oshinei_session"


def _actor(request, fallback=None, default="staff"):
    """Authenticated username from the signed session cookie, else the client-supplied
    fallback, else `default`. Pass default='' for identity fields that must stay blank
    when the actor is unknown (e.g. SoD compares created_by vs approved_by — two blanks
    must NOT read as the same person). Never raises."""
    if _parse_session is not None and request is not None:
        try:
            tok = request.cookies.get(_SESSION_COOKIE)
            if tok:
                payload = _parse_session(tok)
                if payload and payload.get("u"):
                    return str(payload["u"])
        except Exception:
            pass
    return (str(fallback).strip() if fallback else "") or default


async def _read_upload(file: UploadFile) -> bytes:
    if not (file.filename or "").lower().endswith(_ALLOWED_EXT):
        raise HTTPException(415, "Only .xlsx/.xls/.csv/.pdf files are accepted.")
    raw = await file.read()
    if not raw:
        raise HTTPException(422, "Empty file.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"File exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.")
    return raw


def _validate_period(period_month: str):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")


def _validate_master_branch(conn, branch_id: int):
    if not conn.execute(
        f"SELECT 1 FROM branches WHERE branch_id=? AND {_MASTER_ONLY}", (branch_id,)
    ).fetchone():
        raise HTTPException(400, f"branch_id {branch_id} is not an active master branch.")


def _assert_period_open(conn, branch_id: int, period_month: str):
    """Reject mutations to an approved/closed period. Reopen by setting the
    session status back to In-Review (or Draft). 423 = Locked."""
    row = conn.execute(
        "SELECT status FROM recon_sessions WHERE branch_id=? AND period_month=?",
        (branch_id, period_month)).fetchone()
    if row and row["status"] in ("approved", "closed"):
        raise HTTPException(
            423, f"{period_month} is {row['status']} for this branch — set the "
                 "session back to In-Review to make changes.")


@app.get("/api/branches")
def list_branches():
    with get_db() as conn:
        rows = conn.execute(
            f"SELECT branch_id, name, branch_type FROM branches "
            f"WHERE {_MASTER_ONLY} ORDER BY name"
        ).fetchall()
        return [dict(r) for r in rows]


# ── Config: payment-code map ──────────────────────────────────────────────────

_CATEGORIES = {"CASH", "QR", "TRANSFER", "CREDIT", "GRAB", "LINEMAN", "COUPON", "ENTERTAIN", "OTHER"}
_BANKS      = {"KTB", "KBANK", "KTC", "BAY", "SCB", "BBL", "TTB", "PROMPTPAY", "ANY", "NONE", "UNKNOWN"}


class PaymentMapUpsert(BaseModel):
    payment_code_norm: str | None = None   # required on create, ignored on update
    category:          str
    bank:              str = "NONE"
    reconcile:         int = 1
    notes:             str = ""
    updated_by:        str = ""


def _validate_map(body: PaymentMapUpsert):
    if body.category not in _CATEGORIES:
        raise HTTPException(400, f"category must be one of {sorted(_CATEGORIES)}")
    if body.bank not in _BANKS:
        raise HTTPException(400, f"bank must be one of {sorted(_BANKS)}")


@app.get("/api/payment-map")
def list_payment_map(needs_attention: int = 0):
    from .seed_payment_map import normalize
    with get_db() as conn:
        sql = "SELECT * FROM recon_payment_map"
        if needs_attention:
            # Codes the daily engine can't lane: bank not set, or PROMPTPAY in a
            # category that reconciles per settling bank (matcher → unmapped_bank).
            sql += (" WHERE bank='UNKNOWN'"
                    " OR (bank='PROMPTPAY' AND category IN ('QR','TRANSFER','CREDIT'))")
        sql += " ORDER BY (bank='UNKNOWN') DESC, category, bank, payment_code_norm"
        rows = [dict(r) for r in conn.execute(sql)]
        # Attach POS volume per normalized code (master branches) so staff can see
        # which codes actually matter — e.g. KNNQR is the single biggest tender.
        amounts: dict = {}
        try:
            for r in conn.execute("""
                SELECT p.payment_code, SUM(p.amount) amt, SUM(p.txn_count) n
                FROM pos_payment_month p JOIN branches b ON b.branch_id = p.branch_id
                WHERE COALESCE(b.branch_type,'master')='master' AND b.active=1
                  AND p.period_month >= ?
                GROUP BY p.payment_code
            """, (RECON_START_MONTH,)):
                a = amounts.setdefault(normalize(r["payment_code"]), {"amount": 0.0, "txn": 0})
                a["amount"] += float(r["amt"] or 0)
                a["txn"]    += int(r["n"] or 0)
        except sqlite3.OperationalError:
            pass
        for row in rows:
            a = amounts.get(row["payment_code_norm"], {"amount": 0.0, "txn": 0})
            row["pos_amount"] = round(a["amount"], 2)
            row["pos_txn"]    = a["txn"]
        if needs_attention:
            # Only codes that actually carry money block reconciliation. Dormant
            # seed codes (฿0, 0 txns) are mappable but not urgent — drop them here
            # so the attention badge counts real work, not dead POS codes.
            rows = [r for r in rows if r["pos_txn"] > 0 or r["pos_amount"] > 0]
        return rows


@app.post("/api/payment-map", dependencies=[Depends(require_write_token)])
def create_payment_map(body: PaymentMapUpsert):
    _validate_map(body)
    code = (body.payment_code_norm or "").strip()
    if not code:
        raise HTTPException(400, "payment_code_norm is required")
    with get_db() as conn:
        try:
            cur = conn.execute("""
                INSERT INTO recon_payment_map (payment_code_norm, category, bank, reconcile, notes, updated_by)
                VALUES (?,?,?,?,?,?)
            """, (code, body.category, body.bank, body.reconcile, body.notes, body.updated_by or "manual"))
            conn.commit()
            return {"id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"payment_code_norm {code!r} already mapped")


@app.patch("/api/payment-map/{map_id}", dependencies=[Depends(require_write_token)])
def update_payment_map(map_id: int, body: PaymentMapUpsert):
    _validate_map(body)
    with get_db() as conn:
        old = conn.execute(
            "SELECT category, bank, reconcile FROM recon_payment_map WHERE id=?",
            (map_id,)).fetchone()
        conn.execute("""
            UPDATE recon_payment_map
            SET category=?, bank=?, reconcile=?, notes=?, updated_by=?, updated_at=datetime('now')
            WHERE id=?
        """, (body.category, body.bank, body.reconcile, body.notes, body.updated_by or "manual", map_id))
        if old:
            diffs = []
            if old["category"] != body.category:
                diffs.append({"field": "category", "old": old["category"], "new": body.category})
            if old["bank"] != body.bank:
                diffs.append({"field": "bank", "old": old["bank"], "new": body.bank})
            if old["reconcile"] != body.reconcile:
                diffs.append({"field": "reconcile", "old": old["reconcile"], "new": body.reconcile})
            if diffs:
                _audit(conn, "payment_map", map_id, "update", body.updated_by or "staff", diffs)
        conn.commit()
    return {"id": map_id}


@app.delete("/api/payment-map/{map_id}", dependencies=[Depends(require_write_token)])
def delete_payment_map(map_id: int):
    with get_db() as conn:
        conn.execute("DELETE FROM recon_payment_map WHERE id=?", (map_id,))
        conn.commit()
    return {"deleted": map_id}


@app.post("/api/payment-map/seed", dependencies=[Depends(require_write_token)])
def reseed_payment_map():
    with get_db() as conn:
        result = seed_payment_map(conn)
    return result


# ── Config: per-lane method parameters (advanced setup — manager token) ───────

_METHODS = {"direct", "two_stage", "cash_3way", "itemized", "issuer_points"}
_DOC1S   = {"deposit_slip", "cc_summary", "grab_report", "lineman_report", "points_summary"}
_DOC2S   = {"bank_stmt_KTB", "bank_stmt_KBANK", "bank_stmt_*"}


class MethodConfig(BaseModel):
    category:         str
    bank:             str = "*"
    doc1:             str | None = None
    doc2:             str | None = None
    method:           str = "direct"
    tol_abs:          float | None = None
    tol_abs_day:      float | None = None
    lag_days:         int | None = None
    merge_days:       int | None = None
    fee_vat_expected: int = 0
    notes:            str = ""
    updated_by:       str = ""


def _validate_method_config(body: MethodConfig):
    if body.category not in _CATEGORIES:
        raise HTTPException(400, f"category must be one of {sorted(_CATEGORIES)}")
    if body.bank != "*" and body.bank not in _BANKS:
        raise HTTPException(400, f"bank must be '*' or one of {sorted(_BANKS)}")
    if body.method not in _METHODS:
        raise HTTPException(400, f"method must be one of {sorted(_METHODS)}")
    if body.doc1 and body.doc1 not in _DOC1S:
        raise HTTPException(400, f"doc1 must be one of {sorted(_DOC1S)}")
    if body.doc2 and body.doc2 not in _DOC2S:
        raise HTTPException(400, f"doc2 must be one of {sorted(_DOC2S)}")
    # Clamp parameters server-side — a typo must not detonate the matcher.
    if body.tol_abs     is not None: body.tol_abs     = max(0.0, min(100_000.0, body.tol_abs))
    if body.tol_abs_day is not None: body.tol_abs_day = max(0.0, min(100_000.0, body.tol_abs_day))
    if body.lag_days    is not None: body.lag_days    = max(0,   min(14,        body.lag_days))
    if body.merge_days  is not None: body.merge_days  = max(1,   min(7,         body.merge_days))


@app.get("/api/method-config")
def list_method_config():
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM recon_method_config ORDER BY category, bank")]


@app.post("/api/method-config", dependencies=[Depends(require_manager_token)])
def create_method_config(body: MethodConfig):
    _validate_method_config(body)
    with get_db() as conn:
        try:
            cur = conn.execute("""
                INSERT INTO recon_method_config
                    (category, bank, doc1, doc2, method, tol_abs, tol_abs_day,
                     lag_days, merge_days, fee_vat_expected, notes, updated_by)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
            """, (body.category, body.bank, body.doc1, body.doc2, body.method,
                  body.tol_abs, body.tol_abs_day, body.lag_days, body.merge_days,
                  body.fee_vat_expected, body.notes or None, body.updated_by or "manager"))
            _audit(conn, "method_config", cur.lastrowid, "create", body.updated_by or "manager")
            conn.commit()
            return {"id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"lane {body.category}/{body.bank} already configured")


@app.patch("/api/method-config/{cfg_id}", dependencies=[Depends(require_manager_token)])
def update_method_config(cfg_id: int, body: MethodConfig):
    _validate_method_config(body)
    with get_db() as conn:
        old = conn.execute("SELECT * FROM recon_method_config WHERE id=?", (cfg_id,)).fetchone()
        if not old:
            raise HTTPException(404, f"method config {cfg_id} not found")
        conn.execute("""
            UPDATE recon_method_config
            SET category=?, bank=?, doc1=?, doc2=?, method=?, tol_abs=?, tol_abs_day=?,
                lag_days=?, merge_days=?, fee_vat_expected=?, notes=?,
                updated_by=?, updated_at=datetime('now')
            WHERE id=?
        """, (body.category, body.bank, body.doc1, body.doc2, body.method,
              body.tol_abs, body.tol_abs_day, body.lag_days, body.merge_days,
              body.fee_vat_expected, body.notes or None,
              body.updated_by or "manager", cfg_id))
        diffs = [{"field": f, "old": old[f], "new": getattr(body, f)}
                 for f in ("doc1", "doc2", "method", "tol_abs", "tol_abs_day",
                           "lag_days", "merge_days", "fee_vat_expected")
                 if old[f] != getattr(body, f)]
        if diffs:
            _audit(conn, "method_config", cfg_id, "update", body.updated_by or "manager", diffs)
        conn.commit()
    return {"id": cfg_id}


@app.delete("/api/method-config/{cfg_id}", dependencies=[Depends(require_manager_token)])
def delete_method_config(cfg_id: int):
    with get_db() as conn:
        cur = conn.execute("DELETE FROM recon_method_config WHERE id=?", (cfg_id,))
        if cur.rowcount == 0:
            raise HTTPException(404, f"method config {cfg_id} not found")
        _audit(conn, "method_config", cfg_id, "delete", "manager")
        conn.commit()
    return {"deleted": cfg_id}


# ── Advanced setup: composed master view + requirement reseed ─────────────────

# Thai label per method for the master table (the spreadsheet's กระทบยอด column)
_METHOD_TH = {
    "direct":        "เทียบยอดตรง POS ↔ เอกสาร",
    "two_stage":     "2 ขั้น: สรุปยอด↔สเตทเมนต์ แล้ว gross↔POS",
    "cash_3way":     "3 ทาง: POS ↔ สลิปนำฝาก ↔ สเตทเมนต์",
    "itemized":      "รายตามรายการในสเตทเมนต์ (วันเดียวกัน/วันถัดไป)",
    "issuer_points": "แต้มบัตร: ใบสรุปคะแนน ↔ สเตทเมนต์",
}


def _doc_presence(conn, period_month: str) -> dict:
    """Company-wide 'is this document type in for the month?' map for the master
    table's doc chips: {doc_key: True/False}."""
    have: dict = {}
    try:
        for r in conn.execute(
            "SELECT DISTINCT source_type, provider FROM recon_uploads "
            "WHERE period_month=? AND status='confirmed'", (period_month,)):
            st, prov = r["source_type"], (r["provider"] or "")
            if st == "bank":
                from .matcher import _provider_bank
                have[f"bank_stmt_{_provider_bank(prov)}"] = True
                have["bank_stmt_*"] = True
            elif st == "credit_card":
                have["cc_summary"] = True
                if "POINT" in prov.upper():
                    have["points_summary"] = True
            elif st == "grab":
                have["grab_report"] = True
            elif st == "lineman":
                have["lineman_report"] = True
        if conn.execute("SELECT 1 FROM recon_deposit_slip WHERE substr(sale_day,1,7)=? LIMIT 1",
                        (period_month,)).fetchone():
            have["deposit_slip"] = True
    except sqlite3.OperationalError:
        pass
    return have


@app.get("/api/setup/master")
def setup_master(period_month: str):
    """Everything the advanced Setup page renders in one call: the master
    payment-code table (the owner's 'ทุกรหัสในตารางเดียว'), joined with each
    lane's method config + POS volume + document presence, plus the raw
    method_configs / bank_accounts / acquirer_mids and the global tolerances."""
    _validate_period(period_month)
    rows = list_payment_map()          # includes pos_amount / pos_txn
    with get_db() as conn:
        cfgs = {(r["category"], r["bank"]): dict(r) for r in conn.execute(
            "SELECT * FROM recon_method_config")}
        have_docs = _doc_presence(conn, period_month)
        accounts = [dict(r) for r in conn.execute(
            """SELECT a.*, b.name AS branch_name FROM recon_bank_account a
               JOIN branches b ON b.branch_id=a.branch_id ORDER BY b.name, a.bank""")]
        mids = [dict(r) for r in conn.execute(
            """SELECT m.*, b.name AS branch_name FROM recon_acquirer_mid m
               JOIN branches b ON b.branch_id=m.branch_id ORDER BY b.name, m.acquirer""")]

    out_rows = []
    for r in rows:
        mc = cfgs.get((r["category"], r["bank"])) or cfgs.get((r["category"], "*")) or {}
        needs_bank = (r["bank"] == "UNKNOWN" or
                      (r["bank"] == "PROMPTPAY" and r["category"] in ("QR", "TRANSFER", "CREDIT")))
        status = ("off" if not r["reconcile"]
                  else "needs_bank" if needs_bank
                  else "dormant" if (r["pos_amount"] == 0 and r["pos_txn"] == 0)
                  else "ok")
        out_rows.append({
            **r,
            "method": {
                "id":               mc.get("id"),
                "doc1":             mc.get("doc1"),
                "doc2":             mc.get("doc2"),
                "method":           mc.get("method"),
                "method_label":     _METHOD_TH.get(mc.get("method"), mc.get("method")),
                "tol_abs":          mc.get("tol_abs"),
                "tol_abs_day":      mc.get("tol_abs_day"),
                "lag_days":         mc.get("lag_days"),
                "merge_days":       mc.get("merge_days"),
                "fee_vat_expected": mc.get("fee_vat_expected", 0),
                "lane_bank":        mc.get("bank"),   # '*' = applies to the whole category
            },
            "doc_status": {
                "doc1": (None if not mc.get("doc1")
                         else "uploaded" if have_docs.get(mc["doc1"]) else "missing"),
                "doc2": (None if not mc.get("doc2")
                         else "uploaded" if have_docs.get(mc["doc2"]) else "missing"),
            },
            "status": status,
        })
    return {
        "period_month":   period_month,
        "rows":           out_rows,
        "method_configs": [dict(v) for v in sorted(cfgs.values(), key=lambda x: (x["category"], x["bank"]))],
        "bank_accounts":  accounts,
        "acquirer_mids":  mids,
        "globals": {"TOL_ABS": TOL_ABS, "TOL_PCT": TOL_PCT,
                    "TOL_ABS_DAY": TOL_ABS_DAY, "TOL_PCT_DAY": TOL_PCT_DAY},
        "manager_required": bool((os.environ.get("OSHINEI_MANAGER_TOKEN") or "").strip()),
    }


@app.post("/api/setup/seed-requirement", dependencies=[Depends(require_manager_token)])
def seed_requirement():
    """Re-apply the owner's requirement table (codes + lane methods) on demand.
    Idempotent; never overwrites staff/manager edits."""
    with get_db() as conn:
        codes = _seed_requirement_codes(conn)
        methods = _seed_method_defaults(conn)
        conn.commit()
    return {"codes": codes, "methods": methods}


# ── Config: bank accounts ─────────────────────────────────────────────────────

class BankAccount(BaseModel):
    branch_id:    int
    bank:         str
    account_no:   str
    account_name: str = ""
    active:       int = 1


@app.get("/api/bank-accounts")
def list_bank_accounts(branch_id: int | None = None):
    with get_db() as conn:
        sql = """SELECT a.*, b.name AS branch_name FROM recon_bank_account a
                 JOIN branches b ON b.branch_id = a.branch_id"""
        params: list = []
        if branch_id:
            sql += " WHERE a.branch_id=?"; params.append(branch_id)
        sql += " ORDER BY b.name, a.bank"
        return [dict(r) for r in conn.execute(sql, params)]


@app.post("/api/bank-accounts", dependencies=[Depends(require_write_token)])
def create_bank_account(body: BankAccount):
    if body.bank not in ("KTB", "KBANK"):
        raise HTTPException(400, "bank must be KTB or KBANK")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        try:
            cur = conn.execute("""
                INSERT INTO recon_bank_account (branch_id, bank, account_no, account_name, active)
                VALUES (?,?,?,?,?)
            """, (body.branch_id, body.bank, body.account_no, body.account_name, body.active))
            conn.commit()
            return {"account_id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"account {body.bank}/{body.account_no} already exists")


@app.patch("/api/bank-accounts/{account_id}", dependencies=[Depends(require_write_token)])
def update_bank_account(account_id: int, body: BankAccount):
    if body.bank not in ("KTB", "KBANK"):
        raise HTTPException(400, "bank must be KTB or KBANK")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        cur = conn.execute("""
            UPDATE recon_bank_account SET branch_id=?, bank=?, account_no=?, account_name=?, active=?
            WHERE account_id=?
        """, (body.branch_id, body.bank, body.account_no, body.account_name, body.active, account_id))
        if cur.rowcount == 0:
            raise HTTPException(404, f"bank account {account_id} not found")
        conn.commit()
    return {"account_id": account_id}


@app.delete("/api/bank-accounts/{account_id}", dependencies=[Depends(require_write_token)])
def delete_bank_account(account_id: int):
    with get_db() as conn:
        cur = conn.execute("DELETE FROM recon_bank_account WHERE account_id=?", (account_id,))
        if cur.rowcount == 0:
            raise HTTPException(404, f"bank account {account_id} not found")
        conn.commit()
    return {"deleted": account_id}


# ── Config: credit-card acquirer MIDs ─────────────────────────────────────────

class AcquirerMid(BaseModel):
    branch_id:   int
    acquirer:    str
    merchant_id: str = ""
    terminal_id: str = ""
    active:      int = 1


@app.get("/api/acquirer-mids")
def list_acquirer_mids(branch_id: int | None = None):
    with get_db() as conn:
        sql = """SELECT m.*, b.name AS branch_name FROM recon_acquirer_mid m
                 JOIN branches b ON b.branch_id = m.branch_id"""
        params: list = []
        if branch_id:
            sql += " WHERE m.branch_id=?"; params.append(branch_id)
        sql += " ORDER BY b.name, m.acquirer"
        return [dict(r) for r in conn.execute(sql, params)]


@app.post("/api/acquirer-mids", dependencies=[Depends(require_write_token)])
def create_acquirer_mid(body: AcquirerMid):
    if body.acquirer not in ("KTB", "KBANK", "KTC", "BAY", "SCB", "BBL", "KTCPOINT", "BAYPOINT"):
        raise HTTPException(400, "acquirer must be KTB, KBANK, KTC, BAY, SCB, BBL, KTCPOINT or BAYPOINT")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        cur = conn.execute("""
            INSERT INTO recon_acquirer_mid (branch_id, acquirer, merchant_id, terminal_id, active)
            VALUES (?,?,?,?,?)
        """, (body.branch_id, body.acquirer, body.merchant_id, body.terminal_id, body.active))
        conn.commit()
        return {"mid_id": cur.lastrowid}


@app.patch("/api/acquirer-mids/{mid_id}", dependencies=[Depends(require_write_token)])
def update_acquirer_mid(mid_id: int, body: AcquirerMid):
    if body.acquirer not in ("KTB", "KBANK", "KTC", "BAY", "SCB", "BBL", "KTCPOINT", "BAYPOINT"):
        raise HTTPException(400, "acquirer must be KTB, KBANK, KTC, BAY, SCB, BBL, KTCPOINT or BAYPOINT")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        cur = conn.execute("""
            UPDATE recon_acquirer_mid SET branch_id=?, acquirer=?, merchant_id=?, terminal_id=?, active=?
            WHERE mid_id=?
        """, (body.branch_id, body.acquirer, body.merchant_id, body.terminal_id, body.active, mid_id))
        if cur.rowcount == 0:
            raise HTTPException(404, f"acquirer MID {mid_id} not found")
        conn.commit()
    return {"mid_id": mid_id}


@app.delete("/api/acquirer-mids/{mid_id}", dependencies=[Depends(require_write_token)])
def delete_acquirer_mid(mid_id: int):
    with get_db() as conn:
        cur = conn.execute("DELETE FROM recon_acquirer_mid WHERE mid_id=?", (mid_id,))
        if cur.rowcount == 0:
            raise HTTPException(404, f"acquirer MID {mid_id} not found")
        conn.commit()
    return {"deleted": mid_id}


# ── Uploads ───────────────────────────────────────────────────────────────────

SourceType = Literal["bank", "credit_card", "grab", "lineman"]

_LINE_TABLE = {
    "bank":        "recon_bank_lines",
    "credit_card": "recon_cc_lines",
    "grab":        "recon_grab_lines",
    "lineman":     "recon_lineman_lines",
}


def _parse(source_type: SourceType, raw: bytes, provider: str, filename: str = ""):
    """Returns (provider, rows, skipped_rows). Recognised Thai bank/acquirer/Grab
    PDFs are handled by format-specific extractors; otherwise CSV/PDF inputs are
    normalised to an in-memory .xlsx so the domain parsers run unchanged."""
    if (filename or "").lower().endswith(".pdf"):
        special = _detect_pdf_format(source_type, raw)
        if special is not None:
            return special
    raw = to_parseable_xlsx(raw, filename)
    if source_type == "bank":
        prov, rows, skipped = parse_bank_statement(raw, provider or None)
    elif source_type == "credit_card":
        prov, rows, skipped = parse_cc_settlement(raw, provider or None)
    elif source_type == "grab":
        rows, skipped = parse_grab_report(raw)
        prov = "Grab"
    else:
        rows, skipped = parse_lineman_report(raw)
        prov = "Lineman"
    return prov, rows, skipped


def _specific_failure_message(source_type: SourceType, raw: bytes, filename: str) -> str | None:
    """The most specific actionable message for an upload that failed/parsed empty:
    a renamed fee invoice (A2), or an unsupported bank's statement PDF (C6)."""
    msg = _fee_invoice_by_text(raw, filename)
    if not msg and source_type == "bank":
        msg = _unsupported_bank_pdf_message(raw, filename)
    return msg


def _parse_or_422(source_type: SourceType, raw: bytes, provider: str, filename: str = ""):
    # C3: a password-locked file can't be type-detected or parsed — say so first.
    locked = _is_encrypted_upload(raw, filename)
    if locked:
        # KTB ships each branch a statement PDF locked with that branch's account
        # number. Try the configured passwords to unlock, then parse the decrypted
        # copy (file_hash above stays on the original bytes). Reject only if none fit.
        unlocked = _unlock_pdf(raw)
        if unlocked is not None:
            raw = unlocked
        else:
            raise HTTPException(422, locked)
    # A2: reject a known WRONG document type up front (cheap filename check) with
    # specific guidance on the correct file to upload — before any parse attempt.
    byname = _fee_invoice_by_name(source_type, filename)
    if byname:
        raise HTTPException(422, byname)
    try:
        prov, rows, skipped = _parse(source_type, raw, provider, filename)
    except HTTPException:
        raise
    except Exception as exc:
        # A statement with zero movements (e.g. KBank CA 'No Data' month) has no
        # table for any parser to find — that's a valid empty document, not an error.
        if _is_empty_statement(raw, filename):
            return (provider or "", [], 0)
        raise HTTPException(422, _specific_failure_message(source_type, raw, filename)
                            or f"Parse error: {exc}")
    # Parsed but found nothing (and not a legit empty statement): give the most
    # specific message we can; otherwise A1's generic 0-row guard handles it.
    if not rows and not _is_empty_statement(raw, filename):
        msg = _specific_failure_message(source_type, raw, filename)
        if msg:
            raise HTTPException(422, msg)
    return prov, rows, skipped


@app.post("/api/uploads/{source_type}/preview")
async def preview_upload(
    source_type: SourceType,
    file: UploadFile = File(...),
    branch_id: int   = Form(...),
    period_month: str = Form(...),
    provider: str    = Form(""),
):
    _validate_period(period_month)
    raw = await _read_upload(file)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
    prov, rows, skipped = _parse_or_422(source_type, raw, provider, file.filename)
    detected = _detect_branch_from_rows(source_type, rows)   # by bank account # / acquirer MID
    return {
        "provider":     prov,
        "row_count":    len(rows),
        "skipped_rows": skipped,   # data rows dropped for an unparseable date
        "preview":      rows[:50],
        "columns":      list(rows[0].keys()) if rows else [],
        "empty_statement": _is_empty_statement(raw, file.filename) if not rows else False,
        # C7: warn when no format-specific extractor matched and the generic grid
        # was used (its debit/credit split can be wrong) — operator should eyeball.
        "generic_fallback": str(prov).strip().lower() in ("", "generic", "none"),
        # OCR fallback used (no-text KTB PDF): amounts passed the balance-chain
        # self-check, but the operator should still eyeball the total before close.
        "ocr_used": "ocr" in str(prov).strip().lower(),
        **detected,                # detected_branch_id / detected_branch_name / detected_key
    }


def _is_empty_statement(raw: bytes, filename: str | None) -> bool:
    """Valid statement with zero movements (e.g. KBank CA 'No Data' month)."""
    if not (filename or "").lower().endswith(".pdf"):
        return False
    from .parsers.pdf_formats import is_empty_statement
    return is_empty_statement(raw)


def _detect_branch_from_rows(source_type: str, rows: list) -> dict:
    """Auto-attribute an upload to a branch from the identifiers inside it:
    a bank statement's account number (recon_bank_account) or an acquirer
    settlement's merchant MID (recon_acquirer_mid). Returns {} if no unambiguous
    match (the operator's chosen branch then stands)."""
    if not rows:
        return {}
    key, sql = None, None
    if source_type == "bank":
        key = next((r.get("account_no") for r in rows if r.get("account_no")), None)
        sql = "SELECT branch_id FROM recon_bank_account WHERE account_no=? AND COALESCE(active,1)=1"
    elif source_type == "credit_card":
        key = next((r.get("merchant_id") for r in rows if r.get("merchant_id")), None)
        sql = "SELECT branch_id FROM recon_acquirer_mid WHERE merchant_id=? AND COALESCE(active,1)=1"
    if not key:
        return {}
    key = str(key).strip()
    try:
        with get_db() as conn:
            hit = conn.execute(sql, (key,)).fetchone()
            if not hit:
                return {}
            name = conn.execute("SELECT name FROM branches WHERE branch_id=?", (hit["branch_id"],)).fetchone()
            return {"detected_branch_id": hit["branch_id"],
                    "detected_branch_name": name["name"] if name else None,
                    "detected_key": key}
    except sqlite3.OperationalError:
        return {}


def _ingest_one(conn, source_type: str, raw: bytes, filename: str,
                branch_id: int, period_month: str, provider: str = "") -> dict:
    """Parse + validate + persist ONE document into recon_uploads + its line table.
    Runs every guard (dedup, C3 lock, A2/C6 wrong-type, wrong-branch, A1 zero-row)
    and raises HTTPException on any. Does NOT commit — the caller owns the
    transaction (single commit in confirm; one commit per file in bulk)."""
    file_hash = hashlib.sha256(raw).hexdigest()
    dup = conn.execute("""
        SELECT file_name, uploaded_at FROM recon_uploads
        WHERE branch_id=? AND period_month=? AND source_type=?
          AND file_hash=? AND status='confirmed'
    """, (branch_id, period_month, source_type, file_hash)).fetchone()
    if dup:
        raise HTTPException(
            409, f"ไฟล์นี้อัปโหลดไปแล้วสำหรับสาขา/เดือนนี้ "
                 f"('{dup['file_name']}' เมื่อ {dup['uploaded_at']})")

    prov, rows, skipped = _parse_or_422(source_type, raw, provider, filename)

    # Wrong-branch guard: a document that identifies its own branch (account # /
    # acquirer MID) and isn't the chosen one is refused (how a UBN stmt lands on UDN).
    detected = _detect_branch_from_rows(source_type, rows)
    det_bid = detected.get("detected_branch_id")
    if det_bid and det_bid != branch_id:
        raise HTTPException(
            409, f"เอกสารนี้เป็นของสาขา '{detected.get('detected_branch_name')}' "
                 f"(ตรวจจาก {detected.get('detected_key')}) ไม่ใช่สาขาที่เลือก")

    # Zero-row guard (A1): a non-empty file that parses to nothing is the wrong
    # type — never save it as a silent 'confirmed, 0 rows'.
    if not rows and not _is_empty_statement(raw, filename):
        raise HTTPException(
            422, "อ่านได้ 0 รายการ — น่าจะเป็นเอกสารผิดประเภทหรือรูปแบบที่ยังไม่รองรับ "
                 "(เช่น ใบกำกับภาษีค่าธรรมเนียม แทนรายงาน settlement)")

    note = "ไม่มีรายการในเอกสาร (No Data)" if (
        not rows and _is_empty_statement(raw, filename)) else None
    cur = conn.execute("""
        INSERT INTO recon_uploads
            (source_type, branch_id, provider, file_name, period_month, row_count, file_hash, status, error_msg)
        VALUES (?,?,?,?,?,?,?,'confirmed',?)
    """, (source_type, branch_id, prov, filename, period_month, len(rows), file_hash, note))
    upload_id = cur.lastrowid
    table = _LINE_TABLE[source_type]
    overrides = _statement_code_overrides(conn) if source_type == "bank" else {}
    for row in rows:
        row["upload_id"] = upload_id
        row["branch_id"] = branch_id
        if overrides and "channel" in row:       # manager remaps win over the parser
            row["channel"] = _override_channel(row.get("description"), row.get("channel"), overrides)
        cols         = ", ".join(row.keys())
        placeholders = ", ".join("?" * len(row))
        conn.execute(f"INSERT INTO {table} ({cols}) VALUES ({placeholders})",
                     list(row.values()))
    _audit(conn, "upload", upload_id, "create", "staff",
           [{"field": "file_name", "old": None, "new": filename}])
    return {"upload_id": upload_id, "rows_saved": len(rows),
            "skipped_rows": skipped, "provider": prov}


@app.post("/api/uploads/{source_type}/confirm", dependencies=[Depends(require_write_token)])
async def confirm_upload(
    source_type: SourceType,
    file: UploadFile = File(...),
    branch_id: int   = Form(...),
    period_month: str = Form(...),
    provider: str    = Form(""),
):
    _validate_period(period_month)
    raw = await _read_upload(file)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        _assert_period_open(conn, branch_id, period_month)
        result = _ingest_one(conn, source_type, raw, file.filename, branch_id, period_month, provider)
        conn.commit()
    return result


# ── Bulk upload (B1): many files / whole folder / zips in one pass ─────────────

_MAX_BULK_FILES   = 2000
_MAX_MEMBER_BYTES = 25 * 1024 * 1024


def _infer_source_type(filename: str) -> str | None:
    """Best-effort source_type from filename conventions (for 'auto' bulk). Returns
    None when unsure → the file is reported in the manifest for manual routing,
    never guessed into the wrong table."""
    n = (filename or "").lower()
    if any(k in n for k in ("rdb003", "tiv_mdr", "rtxinv", "e-tax_invoice",
                            "_rmer", "_cc_", "kbank_cc", "settle")):
        return "credit_card"
    if n.startswith("3-cz") or re.match(r"^im\d{6,}", n):
        return "grab"
    if any(k in n for k in ("lineman", "line man", "line_man", "lmin", "lmrn",
                            "balance-report", "report_may", "report_jun")):
        return "lineman"
    if "statement" in n or "stmt" in n or "_kb_" in n or "_ktb" in n:
        return "bank"
    return None


def _expand_zip(zip_name: str, raw: bytes, manifest: list) -> list:
    """Yield (member_name, bytes) for a .zip; password-locked / oversized / bad
    members are recorded in the manifest and skipped."""
    import io as _io
    import zipfile as _zip
    out = []
    try:
        with _zip.ZipFile(_io.BytesIO(raw)) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                label = f"{zip_name}:{Path(info.filename).name}"
                if info.file_size > _MAX_MEMBER_BYTES:
                    manifest.append({"file": label, "status": "skipped",
                                     "rows_saved": 0, "reason": "ไฟล์ใหญ่เกินไป"})
                    continue
                try:
                    out.append((Path(info.filename).name, zf.read(info)))
                except RuntimeError:
                    manifest.append({"file": label, "status": "skipped", "rows_saved": 0,
                                     "reason": "ไฟล์ใน zip ถูกล็อกด้วยรหัสผ่าน — ถอดรหัสก่อน"})
    except Exception:
        manifest.append({"file": zip_name, "status": "error", "rows_saved": 0,
                         "reason": "เปิดไฟล์ zip ไม่ได้"})
    return out


@app.post("/api/uploads/bulk", dependencies=[Depends(require_write_token)])
async def bulk_upload(
    files: list[UploadFile] = File(...),
    branch_id: int   = Form(...),
    period_month: str = Form(...),
    source_type: str = Form("auto"),   # a real type, or 'auto' (infer per file)
    provider: str    = Form(""),
):
    """Upload many files (and/or .zip archives) in one pass. Each file runs the
    same validated pipeline as a single confirm; the response is a per-file
    manifest so wrong-type / locked / duplicate files surface together instead of
    one-click-at-a-time. One commit per file → a bad file never rolls back good ones."""
    _validate_period(period_month)
    if source_type not in ("bank", "credit_card", "grab", "lineman", "auto"):
        raise HTTPException(400, "source_type ต้องเป็น bank|credit_card|grab|lineman|auto")
    if len(files) > _MAX_BULK_FILES:
        raise HTTPException(413, f"ไฟล์มากเกินไป (สูงสุด {_MAX_BULK_FILES} ไฟล์ต่อครั้ง)")

    manifest: list = []
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        _assert_period_open(conn, branch_id, period_month)

        # Flatten: each upload is either a file or a .zip to expand.
        items: list = []
        for f in files:
            raw = await f.read()
            name = f.filename or "(unnamed)"
            if name.lower().endswith(".zip"):
                items.extend(_expand_zip(name, raw, manifest))
            else:
                items.append((name, raw))

        for name, raw in items:
            ext = Path(name).suffix.lower()
            if ext and ext not in _ALLOWED_EXT:
                manifest.append({"file": name, "status": "skipped", "rows_saved": 0,
                                 "reason": f"ไม่รองรับนามสกุล {ext}"}); continue
            if not raw:
                manifest.append({"file": name, "status": "skipped", "rows_saved": 0,
                                 "reason": "ไฟล์ว่าง"}); continue
            if len(raw) > _MAX_MEMBER_BYTES:
                manifest.append({"file": name, "status": "skipped", "rows_saved": 0,
                                 "reason": "ไฟล์ใหญ่เกินไป"}); continue
            st = source_type if source_type != "auto" else _infer_source_type(name)
            if not st:
                manifest.append({"file": name, "status": "skipped", "rows_saved": 0,
                                 "reason": "ไม่ทราบประเภทไฟล์ — เลือกช่องเองแล้วอัปโหลด"}); continue
            try:
                res = _ingest_one(conn, st, raw, name, branch_id, period_month, provider)
                conn.commit()
                manifest.append({"file": name, "source_type": st, "status": "ok",
                                 "rows_saved": res["rows_saved"], "reason": ""})
            except HTTPException as e:
                conn.rollback()
                status = "skipped" if e.status_code == 409 else "error"
                manifest.append({"file": name, "source_type": st, "status": status,
                                 "rows_saved": 0, "reason": str(e.detail)})
            except Exception as e:                       # pragma: no cover
                conn.rollback()
                manifest.append({"file": name, "source_type": st, "status": "error",
                                 "rows_saved": 0, "reason": f"error: {e}"})

    ok  = sum(1 for m in manifest if m["status"] == "ok")
    skip = sum(1 for m in manifest if m["status"] == "skipped")
    err = sum(1 for m in manifest if m["status"] == "error")
    return {"summary": {"total": len(manifest), "saved": ok, "skipped": skip,
                        "errored": err, "rows": sum(m["rows_saved"] for m in manifest)},
            "files": manifest}


@app.get("/api/uploads")
def list_uploads(branch_id: int | None = None, period_month: str | None = None):
    with get_db() as conn:
        filters, params = [], []
        if branch_id:
            filters.append("u.branch_id=?");    params.append(branch_id)
        if period_month:
            filters.append("u.period_month=?"); params.append(period_month)
        where = ("WHERE " + " AND ".join(filters)) if filters else ""
        rows = conn.execute(f"""
            SELECT u.*, b.name AS branch_name
            FROM recon_uploads u
            LEFT JOIN branches b ON b.branch_id = u.branch_id
            {where}
            ORDER BY u.uploaded_at DESC LIMIT 200
        """, params).fetchall()
        out = [dict(r) for r in rows]

        # Per-upload match status, computed LIVE by lane (no Run Match needed).
        # Recompute once per (branch, month) present in the result, then attach.
        scopes = {(u["branch_id"], u["period_month"]) for u in out
                  if u.get("branch_id") and u.get("period_month")}
        stats: dict = {}
        for bid, pm in scopes:
            try:
                stats.update(upload_match_stats(conn, bid, pm))
            except sqlite3.OperationalError:
                pass
        for u in out:
            s = stats.get(u["upload_id"], {"total": 0, "matched": 0, "excluded": 0})
            u["match_total"]    = s["total"]
            u["match_done"]     = s["matched"]
            u["match_excluded"] = s["excluded"]
        return out


@app.get("/api/uploads/{upload_id}/lines")
def upload_lines(upload_id: int, limit: int = 300):
    """The parsed lines of one upload WITH their match status — lets staff open a
    document and see exactly which lines reconciled and which still hang."""
    from .matcher import (_CHANNEL_CATEGORY, _provider_bank, _provider_acquirer,
                          _line_status, _lane_status_map, cash_deposit_matching)
    with get_db() as conn:
        up = conn.execute(
            "SELECT source_type, branch_id, period_month, provider FROM recon_uploads WHERE upload_id=?",
            (upload_id,)).fetchone()
        if not up:
            raise HTTPException(404, "upload not found")
        st_, bid, pm, prov = up["source_type"], up["branch_id"], up["period_month"], up["provider"]
        is_cc = st_ == "credit_card"
        spec = {
            "bank":        ("recon_bank_lines",    "txn_date",        "description", "credit"),
            "credit_card": ("recon_cc_lines",      "settlement_date", "batch_no",    "gross_amount"),
            "grab":        ("recon_grab_lines",    "order_date",      "order_id",    "food_amount"),
            "lineman":     ("recon_lineman_lines", "order_date",      "order_id",    "food_amount"),
        }[st_]
        table, date_col, ref_col, amt_col = spec
        extra = (", channel" if table == "recon_bank_lines"
                 else ", acquirer" if is_cc else "")
        rows = conn.execute(f"""
            SELECT id, {date_col} AS date, {ref_col} AS ref, {amt_col} AS amount,
                   match_status{extra}
            FROM {table} WHERE upload_id=? ORDER BY {date_col}, id LIMIT ?
        """, (upload_id, max(1, min(limit, 1000)))).fetchall()
        out = [dict(r) for r in rows]

        # Derive each line's LIVE reconciliation status by lane so the drill
        # agrees with the panel without a Run Match. One upload = one bank, so the
        # provider's bank is resolved once.
        lane_status = _lane_status_map(conn, bid, pm)
        cash_ids = {g["line_id"] for g in cash_deposit_matching(conn, bid, pm)["groups"]}
        prov_bank = _provider_bank(prov)
        prov_acq  = _provider_acquirer(prov)
        for r in out:
            if st_ == "bank":
                cat = _CHANNEL_CATEGORY.get(r.get("channel"))
                bank = "NONE" if cat == "CASH" else prov_bank
                # only revenue-channel CREDITS reconcile; debits / zero / FEE / payout
                # legs are informational (mirrors upload_match_stats' credit>0 filter)
                if cat is None or (r.get("amount") or 0) <= 0:
                    r["status"] = "informational"; continue
            elif is_cc:
                cat, bank = "CREDIT", (r.get("acquirer") or prov_acq)
            else:
                cat, bank = ("GRAB" if st_ == "grab" else "LINEMAN"), "NONE"
            r["status"] = _line_status(r["match_status"], cat, bank, lane_status, cash_ids, r["id"])
        done = sum(1 for r in out if r["status"] == "matched")
        excl = sum(1 for r in out if r["status"] == "excluded")
        # Denominator = reconcilable revenue lines (exclude informational + excluded),
        # so "X of Y" here equals the panel's match bar.
        reconcilable = sum(1 for r in out if r["status"] in ("matched", "unmatched", "manual"))
        return {"source_type": st_, "count": reconcilable, "total_lines": len(out),
                "matched": done, "excluded": excl, "lines": out}


@app.delete("/api/uploads/{upload_id}", dependencies=[Depends(require_write_token)])
def delete_upload(upload_id: int, request: Request):
    with get_db() as conn:
        up = conn.execute(
            "SELECT branch_id, period_month, file_name, file_hash FROM recon_uploads "
            "WHERE upload_id=?", (upload_id,)).fetchone()
        if up:
            _assert_period_open(conn, up["branch_id"], up["period_month"])
        conn.execute("DELETE FROM recon_uploads WHERE upload_id=?", (upload_id,))
        # Record the file hash too, so a deleted source document stays traceable in the
        # audit log (it may be referenced by an earlier close snapshot's manifest).
        _audit(conn, "upload", upload_id, "delete", _actor(request, None),
               [{"field": "file_name", "old": up["file_name"] if up else None, "new": None},
                {"field": "file_hash", "old": (up["file_hash"] if up else None), "new": None}])
        conn.commit()
    return {"deleted": upload_id}


class _UploadBranchBody(BaseModel):
    branch_id: int


@app.patch("/api/uploads/{upload_id}/branch", dependencies=[Depends(require_write_token)])
def reassign_upload_branch(upload_id: int, body: _UploadBranchBody):
    """Move a confirmed upload (and all its parsed lines) to another branch —
    the 2-click fix for a statement confirmed under the wrong branch."""
    with get_db() as conn:
        up = conn.execute(
            "SELECT source_type, branch_id, period_month, file_name FROM recon_uploads WHERE upload_id=?",
            (upload_id,)).fetchone()
        if not up:
            raise HTTPException(404, "upload not found")
        if up["branch_id"] == body.branch_id:
            return {"upload_id": upload_id, "branch_id": body.branch_id, "moved_lines": 0}
        _validate_master_branch(conn, body.branch_id)
        _assert_period_open(conn, up["branch_id"], up["period_month"])
        _assert_period_open(conn, body.branch_id, up["period_month"])

        table = _LINE_TABLE[up["source_type"]]
        moved = conn.execute(
            f"UPDATE {table} SET branch_id=? WHERE upload_id=?",
            (body.branch_id, upload_id)).rowcount
        conn.execute("UPDATE recon_uploads SET branch_id=? WHERE upload_id=?",
                     (body.branch_id, upload_id))
        _audit(conn, "upload", upload_id, "update", "staff",
               [{"field": "branch_id", "old": up["branch_id"], "new": body.branch_id}])
        conn.commit()
    return {"upload_id": upload_id, "branch_id": body.branch_id, "moved_lines": moved}


class _IgnoreBody(BaseModel):
    ignored: bool
    note:    str = ""
    by:      str = ""


@app.patch("/api/bank-lines/{line_id}/ignore", dependencies=[Depends(require_write_token)])
def set_bank_line_ignore(line_id: int, body: _IgnoreBody):
    """Flag a bank credit as an internal transfer (not revenue) so it stops
    polluting the lanes — e.g. a ฿2M intercompany sweep showing as TRANSFER
    ext_only. Reversible; survives Run Match; audited with the note."""
    with get_db() as conn:
        line = conn.execute(
            "SELECT branch_id, txn_date, description, credit, match_status "
            "FROM recon_bank_lines WHERE id=?", (line_id,)).fetchone()
        if not line:
            raise HTTPException(404, "bank line not found")
        _assert_period_open(conn, line["branch_id"], str(line["txn_date"])[:7])
        if body.ignored and line["match_status"] == "manual":
            raise HTTPException(409, "line is manually matched — undo the manual match first")
        if not body.ignored and line["match_status"] != "excluded":
            raise HTTPException(409, "line is not ignored")
        new_status = "excluded" if body.ignored else "unmatched"
        conn.execute("UPDATE recon_bank_lines SET match_status=? WHERE id=?",
                     (new_status, line_id))
        _audit(conn, "bank_line", line_id, "update", body.by or "staff",
               [{"field": "match_status", "old": line["match_status"], "new": new_status},
                {"field": "note", "old": None,
                 "new": body.note or f"{line['txn_date']} {line['description']} {line['credit']}"}])
        conn.commit()
    return {"line_id": line_id, "match_status": new_status}


# ── Evidence attachments (photos / PDFs per branch+day) ───────────────────────
_ATTACH_MAX = 10 * 1024 * 1024
_ATTACH_EXT = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".pdf", ".heic")

@app.get("/api/attachments")
def list_attachments(branch_id: int, day: str):
    if not _DAY_RE.match(day or ""):
        raise HTTPException(400, "day must be in 'YYYY-MM-DD' format.")
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT id, branch_id, day, category, file_name, mime, size_bytes, note, "
            "uploaded_by, uploaded_at FROM recon_attachment "
            "WHERE branch_id=? AND day=? ORDER BY uploaded_at DESC", (branch_id, day))]


@app.post("/api/attachments", dependencies=[Depends(require_write_token)])
async def create_attachment(branch_id: int = Form(...), day: str = Form(...),
                            category: str = Form(""), note: str = Form(""),
                            uploaded_by: str = Form(""), file: UploadFile = File(...)):
    if not _DAY_RE.match(day or ""):
        raise HTTPException(400, "day must be in 'YYYY-MM-DD' format.")
    fn = (file.filename or "").strip()
    if not fn.lower().endswith(_ATTACH_EXT):
        raise HTTPException(415, "Only images (jpg/png/webp/gif/heic) or PDF are accepted.")
    raw = await file.read()
    if not raw:
        raise HTTPException(422, "Empty file.")
    if len(raw) > _ATTACH_MAX:
        raise HTTPException(413, f"File exceeds the {_ATTACH_MAX // (1024*1024)} MB limit.")

    # Store file on volume, not in DB
    safe_ext = Path(fn).suffix.lower()
    rel_path = f"{day[:7]}/{branch_id}_{day}_{uuid.uuid4().hex[:8]}{safe_ext}"
    dest = _ATTACH_DIR / rel_path
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(raw)

    try:
        with get_db() as conn:
            _validate_master_branch(conn, branch_id)
            cur = conn.execute(
                "INSERT INTO recon_attachment"
                "(branch_id,day,category,file_name,mime,size_bytes,file_path,note,uploaded_by) "
                "VALUES(?,?,?,?,?,?,?,?,?)",
                (branch_id, day, (category or None), fn, file.content_type,
                 len(raw), rel_path, (note or None), (uploaded_by or None)))
            conn.commit()
    except Exception:
        dest.unlink(missing_ok=True)
        raise
    return {"id": cur.lastrowid, "file_name": fn, "size_bytes": len(raw)}


@app.get("/api/attachments/{att_id}/file")
def get_attachment_file(att_id: int):
    with get_db() as conn:
        r = conn.execute(
            "SELECT file_name, mime, file_path FROM recon_attachment WHERE id=?",
            (att_id,)).fetchone()
    if not r:
        raise HTTPException(404, "not found")
    fp = _ATTACH_DIR / r["file_path"] if r["file_path"] else None
    if not fp or not fp.exists():
        raise HTTPException(404, "file missing from storage")
    return Response(content=fp.read_bytes(),
                    media_type=r["mime"] or "application/octet-stream",
                    headers={"Content-Disposition": f'inline; filename="{r["file_name"]}"'})


@app.delete("/api/attachments/{att_id}", dependencies=[Depends(require_write_token)])
def delete_attachment(att_id: int):
    with get_db() as conn:
        r = conn.execute("SELECT file_path FROM recon_attachment WHERE id=?",
                         (att_id,)).fetchone()
        if r and r["file_path"]:
            fp = _ATTACH_DIR / r["file_path"]
            fp.unlink(missing_ok=True)
        conn.execute("DELETE FROM recon_attachment WHERE id=?", (att_id,))
        conn.commit()
    return {"deleted": att_id}


# ── Deposit slips (สลิปนำฝาก — CASH stage-1 evidence, typed by staff) ──────────

class DepositSlip(BaseModel):
    branch_id:     int
    sale_day:      str                  # POS day D the cash belongs to
    amount:        float
    deposit_day:   str | None = None    # normally D+1 (before 13:00)
    attachment_id: int | None = None    # optional slip photo (recon_attachment)
    note:          str = ""
    created_by:    str = ""


@app.get("/api/deposit-slips")
def list_deposit_slips(branch_id: int, period_month: str):
    _validate_period(period_month)
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM recon_deposit_slip WHERE branch_id=? AND substr(sale_day,1,7)=? "
            "ORDER BY sale_day", (branch_id, period_month))]


@app.post("/api/deposit-slips", dependencies=[Depends(require_write_token)])
def upsert_deposit_slip(body: DepositSlip):
    """One slip per (branch, sale_day) — POSTing again replaces the amount."""
    if not _DAY_RE.match(body.sale_day or ""):
        raise HTTPException(400, "sale_day must be in 'YYYY-MM-DD' format.")
    if body.deposit_day and not _DAY_RE.match(body.deposit_day):
        raise HTTPException(400, "deposit_day must be in 'YYYY-MM-DD' format.")
    if body.amount <= 0:
        raise HTTPException(400, "amount must be positive.")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        _assert_period_open(conn, body.branch_id, body.sale_day[:7])
        old = conn.execute(
            "SELECT id, amount FROM recon_deposit_slip WHERE branch_id=? AND sale_day=?",
            (body.branch_id, body.sale_day)).fetchone()
        if old:
            conn.execute("""
                UPDATE recon_deposit_slip SET amount=?, deposit_day=?, attachment_id=?,
                    note=?, created_by=?, created_at=datetime('now') WHERE id=?
            """, (body.amount, body.deposit_day, body.attachment_id,
                  body.note or None, body.created_by or "staff", old["id"]))
            slip_id = old["id"]
            if float(old["amount"]) != body.amount:
                _audit(conn, "deposit_slip", slip_id, "update", body.created_by or "staff",
                       [{"field": "amount", "old": old["amount"], "new": body.amount}])
        else:
            cur = conn.execute("""
                INSERT INTO recon_deposit_slip
                    (branch_id, sale_day, deposit_day, amount, attachment_id, note, created_by)
                VALUES (?,?,?,?,?,?,?)
            """, (body.branch_id, body.sale_day, body.deposit_day, body.amount,
                  body.attachment_id, body.note or None, body.created_by or "staff"))
            slip_id = cur.lastrowid
            _audit(conn, "deposit_slip", slip_id, "create", body.created_by or "staff",
                   [{"field": "amount", "old": None, "new": body.amount}])
        conn.commit()
    return {"id": slip_id, "sale_day": body.sale_day, "amount": body.amount}


@app.delete("/api/deposit-slips/{slip_id}", dependencies=[Depends(require_write_token)])
def delete_deposit_slip(slip_id: int):
    with get_db() as conn:
        row = conn.execute(
            "SELECT branch_id, sale_day FROM recon_deposit_slip WHERE id=?", (slip_id,)).fetchone()
        if not row:
            raise HTTPException(404, "deposit slip not found")
        _assert_period_open(conn, row["branch_id"], row["sale_day"][:7])
        conn.execute("DELETE FROM recon_deposit_slip WHERE id=?", (slip_id,))
        _audit(conn, "deposit_slip", slip_id, "delete", "staff")
        conn.commit()
    return {"deleted": slip_id}


# ── Sessions ──────────────────────────────────────────────────────────────────

class SessionCreate(BaseModel):
    branch_id:    int
    period_month: str
    notes:        str = ""
    created_by:   str = ""


@app.post("/api/sessions", dependencies=[Depends(require_write_token)])
def create_session(body: SessionCreate, request: Request):
    _validate_period(body.period_month)
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        try:
            cur = conn.execute("""
                INSERT INTO recon_sessions (branch_id, period_month, notes, created_by)
                VALUES (?,?,?,?)
            """, (body.branch_id, body.period_month, body.notes, _actor(request, body.created_by, default="")))
            conn.commit()
            return {"session_id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"Session already exists for branch {body.branch_id} / {body.period_month}")


@app.get("/api/sessions")
def list_sessions(branch_id: int | None = None, period_month: str | None = None):
    with get_db() as conn:
        filters, params = [], []
        if branch_id:
            filters.append("s.branch_id=?");    params.append(branch_id)
        if period_month:
            filters.append("s.period_month=?"); params.append(period_month)
        where = ("WHERE " + " AND ".join(filters)) if filters else ""
        rows = conn.execute(f"""
            SELECT s.*, b.name AS branch_name
            FROM recon_sessions s
            JOIN branches b ON b.branch_id = s.branch_id
            {where}
            ORDER BY s.period_month DESC, b.name
        """, params).fetchall()
        return [dict(r) for r in rows]


@app.get("/api/sessions/{session_id}")
def get_session(session_id: int):
    with get_db() as conn:
        row = conn.execute("""
            SELECT s.*, b.name AS branch_name
            FROM recon_sessions s
            JOIN branches b ON b.branch_id = s.branch_id
            WHERE s.session_id=?
        """, (session_id,)).fetchone()
        if not row:
            raise HTTPException(404)
        return dict(row)


@app.post("/api/sessions/{session_id}/run-match", dependencies=[Depends(require_write_token)])
def run_match(session_id: int):
    with get_db() as conn:
        session = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)
        ).fetchone()
        if not session:
            raise HTTPException(404)
        _assert_period_open(conn, session["branch_id"], session["period_month"])
        result = run_auto_match(conn, session_id, session["branch_id"], session["period_month"])
    return result


@app.patch("/api/sessions/{session_id}/status", dependencies=[Depends(require_write_token)])
def update_session_status(
    session_id:  int,
    request:     Request,
    status:      str = Form(...),
    approved_by: str = Form(""),
    reason:      str = Form(""),
    x_manager_token: str = Header(None),
):
    valid = {"draft", "in_review", "approved", "closed"}
    if status not in valid:
        raise HTTPException(400, f"status must be one of {valid}")
    if status in ("approved", "closed"):
        _require_manager(x_manager_token)
    actor = _actor(request, approved_by, default="")        # blank when unknown → SoD-safe
    with get_db() as conn:
        old_row = conn.execute(
            "SELECT status FROM recon_sessions WHERE session_id=?", (session_id,)).fetchone()
        if not old_row:
            raise HTTPException(404, f"session {session_id} not found")
        old_status = old_row["status"]
        # Reopen gate (control fix 2026-06-19): unlocking an already-signed-off month
        # (approved/closed → draft/in_review) is sensitive — require the MANAGER code +
        # a logged reason, just like approval. The next re-approval re-freezes a fresh
        # snapshot, so the working paper always matches the final figures.
        reopening = old_status in ("approved", "closed") and status not in ("approved", "closed")
        if reopening:
            _require_manager(x_manager_token)
            if len((reason or "").strip()) < 10:
                raise HTTPException(
                    400, "การเปิดงวดที่ปิดแล้วต้องระบุเหตุผล (อย่างน้อย 10 ตัวอักษร)")
        # Approval gate: don't let a month be signed off while issues are open.
        if status in ("approved", "closed"):
            pending = conn.execute(
                "SELECT COUNT(*) FROM recon_discrepancies "
                "WHERE session_id=? AND resolution='pending'",
                (session_id,)).fetchone()[0]
            if pending:
                raise HTTPException(
                    409, f"Cannot {status} — {pending} open discrepancy(ies) remain. "
                         "Resolve or accept them first.")
        # Only stamp approval metadata when the session is actually approved;
        # moving back to draft/in_review must not record a (false) approval.
        if status == "approved":
            # Segregation of duties (AUD-4): the person who prepared the month
            # cannot also approve it. A classic audit finding; enforced when both
            # identities are known.
            prep = (conn.execute("SELECT created_by FROM recon_sessions WHERE session_id=?",
                                  (session_id,)).fetchone()["created_by"] or "").strip().lower()
            appr = (actor or "").strip().lower()
            if prep and appr and prep == appr:
                raise HTTPException(
                    409, "แบ่งแยกหน้าที่ (SoD): ผู้จัดทำไม่สามารถอนุมัติงานของตัวเองได้ — "
                         "ต้องให้ผู้อื่นเป็นผู้อนุมัติ")
            conn.execute("""
                UPDATE recon_sessions
                SET status=?, approved_by=?, approved_at=datetime('now')
                WHERE session_id=?
            """, (status, actor, session_id))
        else:
            conn.execute(
                "UPDATE recon_sessions SET status=? WHERE session_id=?",
                (status, session_id),
            )
        if old_status != status:
            diffs = [{"field": "status", "old": old_status, "new": status}]
            if reopening:
                diffs.append({"field": "reopen_reason", "old": None, "new": reason.strip()})
            _audit(conn, "session", session_id, "update", actor or "staff", diffs)
        # AUD-5: freeze the working paper at each sign-off transition. Best-effort —
        # the export can be regenerated, so a snapshot hiccup must not block close.
        snap = None
        if status in ("approved", "closed") and old_status != status:
            try:
                full = conn.execute(
                    "SELECT s.*, b.name AS branch_name FROM recon_sessions s "
                    "JOIN branches b ON b.branch_id=s.branch_id WHERE s.session_id=?",
                    (session_id,)).fetchone()
                snap = _write_close_snapshot(conn, full, status, actor)
            except Exception:
                snap = None
        conn.commit()
    return {"session_id": session_id, "status": status, "snapshot": snap}


# ── Session summary + category reconciliation ─────────────────────────────────

def _norm_ts(s):
    """Normalise a stored timestamp for lexical comparison. computed_at /
    uploaded_at use 'YYYY-MM-DD HH:MM:SS' while audit_log.ts uses a 'T' separator
    — line them up so a plain string `>` is correct."""
    return s.replace("T", " ")[:19] if s else None


def _match_freshness(conn, session_id, branch_id, period_month):
    """When did Run Match last persist this session's category snapshot, vs. the
    latest data change (upload, manual match, lane adjustment, internal-transfer
    ignore)? The live `/summary` cards are always recomputed, but the *persisted*
    `recon_category_result` (Excel export, `/categories`, approval snapshot) only
    refreshes on Run Match — `stale=True` means it lags and staff should re-run."""
    computed_at = _norm_ts(conn.execute(
        "SELECT MAX(computed_at) FROM recon_category_result WHERE session_id=?",
        (session_id,)).fetchone()[0])
    changes = [
        conn.execute(
            "SELECT MAX(COALESCE(confirmed_at, uploaded_at)) FROM recon_uploads "
            "WHERE branch_id=? AND period_month=? AND status!='error'",
            (branch_id, period_month)).fetchone()[0],
        conn.execute(
            "SELECT MAX(matched_at) FROM recon_manual_match WHERE branch_id=? AND period_month=?",
            (branch_id, period_month)).fetchone()[0],
        conn.execute(
            "SELECT MAX(created_at) FROM recon_lane_adj WHERE branch_id=? AND period_month=?",
            (branch_id, period_month)).fetchone()[0],
        conn.execute(
            "SELECT MAX(a.ts) FROM recon_audit_log a "
            "JOIN recon_bank_lines bl ON bl.id = a.entity_id "
            "WHERE a.entity='bank_line' AND bl.branch_id=? "
            "AND strftime('%Y-%m', bl.txn_date)=?",
            (branch_id, period_month)).fetchone()[0],
    ]
    norm = [_norm_ts(x) for x in changes if x]
    last_change_at = max(norm) if norm else None
    # Stale when results were never computed, or a change post-dates the last run.
    stale = computed_at is None or (last_change_at is not None and last_change_at > computed_at)
    return {"computed_at": computed_at, "last_change_at": last_change_at, "stale": stale}


@app.get("/api/sessions/{session_id}/summary")
def session_summary(session_id: int):
    """Live per-category reconciliation (POS tender vs the right external source),
    computed fresh so it reflects current uploads/fee-VAT edits."""
    with get_db() as conn:
        session = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)
        ).fetchone()
        if not session:
            raise HTTPException(404)
        bid, pm = session["branch_id"], session["period_month"]

        rows, unmapped = category_reconciliation(conn, bid, pm)
        # per-lane manual-match tally (so the cards can flag manually-cleared lanes)
        manual: dict = {}
        for mr in conn.execute(
            "SELECT category, bank, COUNT(*) n, COALESCE(SUM(pos_amount),0) amt "
            "FROM recon_manual_match WHERE branch_id=? AND period_month=? GROUP BY category, bank",
            (bid, pm)):
            manual[f"{mr['category']}|{mr['bank']}"] = {"count": mr["n"], "amount": round(mr["amt"], 2)}
        categories = []
        for r in rows:
            c = {k: r[k] for k in
                 ("category", "bank", "pos_amount", "ext_amount", "ext_net", "variance", "status",
                  "adj_pos", "adj_ext", "adj_note", "adj_type", "adj_id")}
            c["manual"] = manual.get(f"{r['category']}|{r['bank']}")
            categories.append(c)
        payouts = payout_reconciliation(conn, bid, pm)

        pending = conn.execute(
            "SELECT COUNT(*) FROM recon_discrepancies WHERE session_id=? AND resolution='pending'",
            (session_id,)
        ).fetchone()[0]

        pos_total = sum(r["pos_amount"] for r in rows)
        ext_total = sum(r["ext_amount"] for r in rows)
        return {
            "categories":            categories,
            "payouts":               payouts,
            "unmapped":              [{"code": k, "amount": round(v, 2)} for k, v in sorted(unmapped.items())],
            "pos_total":             round(pos_total, 2),
            "ext_total":             round(ext_total, 2),
            "variance":              round(ext_total - pos_total, 2),
            "matched":               sum(1 for r in rows if r["status"] == "matched"),
            "pending_discrepancies": pending,
            "freshness":             _match_freshness(conn, session_id, bid, pm),
        }


@app.get("/api/sessions/{session_id}/categories")
def session_categories(session_id: int):
    """Persisted per-category results from the last run-match."""
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM recon_category_result WHERE session_id=? ORDER BY category, bank",
            (session_id,))]


@app.get("/api/sessions/{session_id}/daily")
def session_daily(session_id: int):
    """Day-by-day reconciliation board with per-day pending tracking for late
    documents. Computed live so it reflects current uploads + fee/VAT edits."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        return daily_reconciliation(conn, s["branch_id"], s["period_month"])


@app.get("/api/sessions/{session_id}/bank-coverage")
def session_bank_coverage(session_id: int):
    """Bank-first coverage: every statement credit line classified matched /
    unmatched / excluded / informational, plus the CREDIT 3-document chain (POS gross
    ↔ settlement gross, settlement net ↔ bank credit, fee+VAT auto-pulled). Computed
    live so it reflects current uploads, manual matches and fee/VAT edits."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        return bank_statement_coverage(conn, s["branch_id"], s["period_month"])


# ── Reconcile v2 — channel-first composed view (no session required) ──────────

@app.get("/api/v2/overview")
def v2_overview(month: str):
    """All master branches for one month: progress, per-channel variance, status."""
    _validate_period(month)
    with get_db() as conn:
        return reconcile_overview(conn, month)


@app.get("/api/v2/reconcile")
def v2_reconcile(branch_id: int, month: str):
    """One branch, one month, channel-first: documents checklist, channel cards
    with per-day detail, non-revenue (internal transfers), setup questions, todo."""
    _validate_period(month)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        return reconcile_month(conn, branch_id, month)


@app.get("/api/v2/lane")
def v2_lane(branch_id: int, month: str, category: str, bank: str, day: str | None = None):
    """Line-level drill for v2 (same data as the session lane endpoint, keyed by
    branch+month directly so the UI needs no session to browse)."""
    _validate_period(month)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        lines = lane_lines(conn, branch_id, month, category, bank, day)
        return {"category": category, "bank": bank, "day": day,
                "count": len(lines), "total": round(sum(l["amount"] for l in lines), 2),
                "lines": lines}


@app.get("/api/v2/pos-lines")
def v2_pos_lines(branch_id: int, month: str, category: str,
                 bank: str | None = None, day: str | None = None,
                 code: str | None = None):
    """Bill-level POS drill: the individual receipts (sale_payment) behind a
    channel's POS total for a day/month, grouped by payment code. Ties out to the
    POS number shown on the channel/day card — the POS side of the day drill."""
    _validate_period(month)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        return pos_lane_lines(conn, branch_id, month, category, bank, day, code)


@app.get("/api/sessions/{session_id}/lane")
def session_lane(session_id: int, category: str, bank: str, day: str | None = None):
    """Drill-down: the external lines that make up one (category, bank) lane,
    optionally for a single day. POS side stays an aggregate (no per-receipt
    tender split exists)."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        lines = lane_lines(conn, s["branch_id"], s["period_month"], category, bank, day)
        return {"category": category, "bank": bank, "day": day,
                "count": len(lines), "total": round(sum(l["amount"] for l in lines), 2),
                "lines": lines}


# ── Excel month-pack export ───────────────────────────────────────────────────

def _build_export(conn, session) -> bytes:
    """One workbook per session: Summary / Categories / Payout chain / Daily /
    Discrepancies — the artifact an accountant signs off and files."""
    import io
    import openpyxl
    from openpyxl.styles import Font

    bid, pm = session["branch_id"], session["period_month"]
    rows, unmapped = category_reconciliation(conn, bid, pm)
    payouts = payout_reconciliation(conn, bid, pm)
    daily = daily_reconciliation(conn, bid, pm)
    discs = conn.execute(
        "SELECT * FROM recon_discrepancies WHERE session_id=? ORDER BY resolution, ABS(COALESCE(variance,0)) DESC",
        (session["session_id"],)).fetchall()

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    def sheet(title, headers, data_rows, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        if first:
            ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for r in data_rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[chr(64 + i)].width = max(12, len(str(h)) + 2)
        return ws

    pos_total = round(sum(r["pos_amount"] for r in rows), 2)
    ext_total = round(sum(r["ext_amount"] for r in rows), 2)
    pending = sum(1 for d in discs if d["resolution"] == "pending")
    # AUD-7: total variance "absorbed" by tolerance — the sum of |variance| over
    # lanes that passed as matched. Auditors want this quantified, not invisible.
    tol_absorbed = round(sum(abs(r["variance"]) for r in rows if r["status"] == "matched"), 2)
    accepts = conn.execute(
        "SELECT category, bank, note, created_by, created_at FROM recon_lane_adj "
        "WHERE branch_id=? AND period_month=? AND adj_type='accept' ORDER BY category, bank",
        (bid, pm)).fetchall()
    sheet("Summary",
          ["Field", "Value"],
          [["Branch", session["branch_name"]], ["Period", pm], ["Status", session["status"]],
           ["POS tender (reconcilable)", pos_total], ["External total", ext_total],
           ["Variance", round(ext_total - pos_total, 2)],
           ["Matched lanes", sum(1 for r in rows if r["status"] == "matched")],
           ["Variance absorbed within tolerance", tol_absorbed],   # AUD-7
           ["Open discrepancies", pending],
           ["Acknowledged variances (รับทราบ)", len(accepts)],     # AUD-8
           ["Unmapped POS codes", ", ".join(u["code"] for u in
                                            [{"code": k} for k in sorted(unmapped)]) or "—"]],
          first=True)

    sheet("Categories",
          ["Category", "Bank", "POS", "Statement", "Net (after fee/VAT)", "Variance", "Status"],
          [[r["category"], r["bank"], r["pos_amount"], r["ext_amount"], r["ext_net"],
            r["variance"], r["status"]] for r in rows])

    sheet("Payout chain",
          ["Category", "Report net", "Bank payout", "Variance", "Status"],
          [[p["category"], p["report_net"], p["bank_payout"], p["variance"], p["status"]] for p in payouts])

    daily_rows = []
    for d in daily.get("days", []):
        for l in d["lanes"]:
            daily_rows.append([d["day"], l["category"], l["bank"], l["pos"], l["ext"],
                               l["variance"], l["status"]])
    sheet("Daily",
          ["Day", "Category", "Bank", "POS", "External", "Variance", "Status"], daily_rows)

    sheet("Discrepancies",
          ["Type", "Category", "Bank", "External", "POS", "Variance", "Resolution", "Assignee", "Note"],
          [[d["disc_type"], d["category"], d["bank"], d["amount_external"], d["amount_pos"],
            d["variance"], d["resolution"], d["assigned_to"] or "", d["resolution_note"] or ""]
           for d in discs])

    # AUD-8: every acknowledged variance (รับทราบ) with its reason + who/when, so
    # the owner can sign off the list and an auditor can sample it.
    sheet("Acknowledged variances",
          ["Category", "Bank", "Reason (เหตุผล)", "Accepted by", "At"],
          [[a["category"], a["bank"], a["note"] or "", a["created_by"] or "", a["created_at"] or ""]
           for a in accepts])

    # Per-day acknowledged variances (reconcile BY DATE): each differing day staff
    # cleared with a reason — the granular trail behind a channel's residual.
    day_acks = conn.execute(
        "SELECT category, day, variance, note, created_by, created_at FROM recon_day_ack "
        "WHERE branch_id=? AND period_month=? ORDER BY category, day", (bid, pm)).fetchall()
    sheet("Daily acknowledged variances",
          ["Category", "Day", "Variance", "Reason (เหตุผล)", "Accepted by", "At"],
          [[a["category"], a["day"], a["variance"], a["note"] or "",
            a["created_by"] or "", a["created_at"] or ""] for a in day_acks])

    # Cash 3-way (POS ↔ typed deposit slip ↔ bank deposit) per the owner's table
    tw = cash_three_way(conn, bid, pm)
    sheet("Cash 3-way",
          ["Day", "POS cash", "Slip amount", "Slip day", "Bank deposit", "Deposit date",
           "Slip status", "Bank status", "Slip var", "Bank var"],
          [[r["day"], r["pos_cash"], r["slip_amount"], r["slip_day"] or "", r["bank_deposit"],
            r["deposit_date"] or "", r["slip_leg"], r["bank_leg"],
            r["slip_variance"], r["bank_variance"]] for r in tw["rows"]])

    # Coupon issuer points (KTC/BAY) — issuer_points method
    ip = coupon_reconciliation(conn, bid, pm)
    sheet("Coupon points",
          ["Issuer", "POS points", "Summary gross", "Bank payout", "Variance", "Status"],
          [[r["issuer"], r["pos_point"], r["settle_gross"],
            ("in KTC deposit" if r["bank_in_credit"] else r["bank_payout"]),
            r["variance"], r["status"]] for r in ip])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.get("/api/sessions/{session_id}/export.xlsx")
def export_session(session_id: int):
    with get_db() as conn:
        session = conn.execute("""
            SELECT s.*, b.name AS branch_name
            FROM recon_sessions s JOIN branches b ON b.branch_id = s.branch_id
            WHERE s.session_id=?""", (session_id,)).fetchone()
        if not session:
            raise HTTPException(404)
        data = _build_export(conn, session)
    from urllib.parse import quote
    fname = f"recon_{session['branch_name']}_{session['period_month']}.xlsx".replace(" ", "_")
    # Branch names are Thai → not latin-1 safe for a raw header. RFC 5987: an ASCII
    # fallback plus a UTF-8 filename* the browser prefers.
    ascii_fallback = f"recon_{session['period_month']}.xlsx"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(fname)}"})


# ── Immutable close snapshot (AUD-5) ──────────────────────────────────────────

def _write_close_snapshot(conn, session, status: str, by: str) -> dict:
    """Freeze the working-paper export at sign-off so 'what management signed'
    survives later edits: the .xlsx bytes are written to the volume and a hashed
    row recorded. `session` must carry branch_name (join). Returns the new row."""
    data = _build_export(conn, session)
    ts = _dt.now().strftime("%Y%m%d-%H%M%S")
    pm = session["period_month"]
    rel = f"{pm}/{session['branch_id']}_{pm}_{status}_{ts}.xlsx"
    dest = _SNAPSHOT_DIR / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    digest = hashlib.sha256(data).hexdigest()
    # Source-document manifest: name + sha256 of every confirmed upload that fed this
    # period, frozen alongside the working paper so the inputs are provable later.
    manifest = json.dumps([
        {"file": u["file_name"], "sha256": u["file_hash"],
         "type": u["source_type"], "provider": u["provider"]}
        for u in conn.execute(
            "SELECT file_name, file_hash, source_type, provider FROM recon_uploads "
            "WHERE branch_id=? AND period_month=? AND status='confirmed' "
            "ORDER BY source_type, file_name", (session["branch_id"], pm))],
        ensure_ascii=False)
    cur = conn.execute(
        "INSERT INTO recon_close_snapshot"
        "(session_id, branch_id, period_month, status, file_name, file_path, sha256, size_bytes, created_by, source_manifest) "
        "VALUES (?,?,?,?,?,?,?,?,?,?)",
        (session["session_id"], session["branch_id"], pm, status,
         Path(rel).name, rel, digest, len(data), by or "staff", manifest))
    _audit(conn, "close_snapshot", cur.lastrowid, "create", by or "staff",
           [{"field": "sha256", "old": None, "new": digest}])
    return {"id": cur.lastrowid, "sha256": digest, "size_bytes": len(data)}


@app.get("/api/sessions/{session_id}/snapshots")
def list_close_snapshots(session_id: int):
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT id, status, file_name, sha256, size_bytes, created_by, created_at "
            "FROM recon_close_snapshot WHERE session_id=? ORDER BY created_at DESC",
            (session_id,))]


@app.get("/api/snapshots/{snap_id}/file")
def get_close_snapshot_file(snap_id: int):
    with get_db() as conn:
        r = conn.execute("SELECT * FROM recon_close_snapshot WHERE id=?", (snap_id,)).fetchone()
    if not r:
        raise HTTPException(404, "snapshot not found")
    p = _SNAPSHOT_DIR / r["file_path"]
    if not p.exists():
        raise HTTPException(404, "snapshot file missing on disk")
    return Response(
        content=p.read_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{r['file_name']}\""})


# ── Branch / staff error report ───────────────────────────────────────────────
# Built from the acknowledged-variance markers (รับทราบ / recon_lane_adj accepts):
# each is a recorded "this lane was off and here's why", with who + when. Aggregated
# per branch (+month) it shows which branch mis-records most often + the baht at
# stake — a fair, evidence-based input to staff evaluation.

def _branch_error_report(conn, period_month: str | None = None) -> dict:
    from collections import defaultdict
    q = ("SELECT branch_id, period_month, category, note, created_by, created_at "
         "FROM recon_lane_adj WHERE adj_type='accept'")
    params: list = []
    if period_month:
        q += " AND period_month=?"
        params.append(period_month)
    accepts = conn.execute(q + " ORDER BY branch_id, period_month, category", params).fetchall()
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}

    var_cache: dict = {}        # (branch, period) -> {channel_key: variance}
    items = []
    for a in accepts:
        key = (a["branch_id"], a["period_month"])
        if key not in var_cache:
            try:
                rm = reconcile_month(conn, a["branch_id"], a["period_month"])
                var_cache[key] = {c["key"]: c["variance"] for c in rm["channels"]}
            except Exception:
                var_cache[key] = {}
        items.append({
            "branch_id":    a["branch_id"],
            "branch_name":  bnames.get(a["branch_id"], str(a["branch_id"])),
            "period_month": a["period_month"],
            "category":     a["category"],
            "variance":     round(var_cache[key].get(a["category"], 0.0), 2),
            "reason":       a["note"] or "",
            "by":           a["created_by"] or "",
            "at":           a["created_at"] or "",
        })

    agg: dict = defaultdict(lambda: {"count": 0, "total_abs_variance": 0.0, "months": set()})
    for it in items:
        g = agg[(it["branch_id"], it["branch_name"])]
        g["count"] += 1
        g["total_abs_variance"] += abs(it["variance"])
        g["months"].add(it["period_month"])
    summary = sorted(
        [{"branch_id": bid, "branch_name": bn, "count": v["count"],
          "months": len(v["months"]), "total_abs_variance": round(v["total_abs_variance"], 2)}
         for (bid, bn), v in agg.items()],
        key=lambda x: -x["total_abs_variance"])
    return {"period_month": period_month, "summary": summary, "items": items}


@app.get("/api/reports/branch-errors")
def branch_error_report(period_month: str | None = None):
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        return _branch_error_report(conn, period_month)


@app.get("/api/reports/branch-errors.xlsx")
def branch_error_report_xlsx(period_month: str | None = None):
    import io
    import openpyxl
    from openpyxl.styles import Font
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _branch_error_report(conn, period_month)
    wb = openpyxl.Workbook(); bold = Font(bold=True)
    ws = wb.active; ws.title = "สรุปรายสาขา"
    ws.append(["สาขา", "จำนวนครั้งที่รับทราบ", "จำนวนเดือน", "ยอดผลต่างรวม (abs)"])
    for c in ws[1]:
        c.font = bold
    for s in rep["summary"]:
        ws.append([s["branch_name"], s["count"], s["months"], s["total_abs_variance"]])
    ws2 = wb.create_sheet("รายการ")
    ws2.append(["สาขา", "เดือน", "ช่องทาง", "ผลต่าง", "เหตุผล", "รับทราบโดย", "เมื่อ"])
    for c in ws2[1]:
        c.font = bold
    for it in rep["items"]:
        ws2.append([it["branch_name"], it["period_month"], it["category"],
                    it["variance"], it["reason"], it["by"], it["at"]])
    for w in (ws, ws2):
        for i, _h in enumerate(w[1], 1):
            w.column_dimensions[chr(64 + i)].width = 22
    buf = io.BytesIO(); wb.save(buf)
    fname = f"branch_errors_{period_month or 'all'}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{fname}\""})


# ── Wrong-channel mis-entry report (รายงานเงินบันทึกผิดช่อง) ────────────────────
# A "เงินอาจถูกบันทึกผิดช่อง" suggestion, once confirmed, becomes a manual match:
#   • category      = the POS lane the cashier rang the money into BY MISTAKE
#   • matched lines = the statement entries showing where the money ACTUALLY landed
# So every such match is a logged correction (with reason + who + when). We surface
# them per branch as a สรุป + an acknowledgment sheet the branch signs and fixes.

# bank statement channel codes → our category vocabulary
_BANK_CH_TO_CAT = {
    "QR": "QR", "TRANSFER": "TRANSFER", "CASH_DEP": "CASH",
    "CC_SETTLE": "CREDIT", "GRAB_PAYOUT": "GRAB", "LINEMAN_PAYOUT": "LINEMAN",
}
_EXT_TABLE_CAT = {
    "recon_cc_lines": "CREDIT", "recon_grab_lines": "GRAB",
    "recon_lineman_lines": "LINEMAN",
}


def _real_channel_of_match(conn, match_id: int) -> str | None:
    """Where the money actually landed, inferred from the matched statement lines
    (dominant channel by baht). None when no lines are attached."""
    cats: dict = {}
    for ln in conn.execute(
        "SELECT ext_table, ext_id, amount FROM recon_manual_match_line WHERE match_id=?",
        (match_id,)):
        tbl = ln["ext_table"]; amt = abs(float(ln["amount"] or 0))
        if tbl == "recon_bank_lines":
            row = conn.execute("SELECT channel FROM recon_bank_lines WHERE id=?",
                               (ln["ext_id"],)).fetchone()
            cat = _BANK_CH_TO_CAT.get((row["channel"] or "") if row else "", "OTHER")
        else:
            cat = _EXT_TABLE_CAT.get(tbl, "OTHER")
        cats[cat] = cats.get(cat, 0.0) + amt
    return max(cats, key=cats.get) if cats else None


def _wrong_channel_items(conn, period_month: str | None = None) -> list[dict]:
    """Confirmed wrong-channel mis-entries (true cross-channel manual matches)."""
    q = ("SELECT match_id, branch_id, category, bank, period_month, pos_amount, "
         "ext_amount, pos_days, note, matched_by, matched_at FROM recon_manual_match")
    params: list = []
    if period_month:
        q += " WHERE period_month=?"; params.append(period_month)
    rows = conn.execute(q + " ORDER BY branch_id, matched_at", params).fetchall()
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}
    items = []
    for m in rows:
        to_cat = _real_channel_of_match(conn, m["match_id"])
        from_cat = m["category"]
        if not to_cat or to_cat == from_cat:        # only true cross-channel mis-entries
            continue
        try:
            days = sorted({d.get("day") for d in json.loads(m["pos_days"] or "[]") if d.get("day")})
        except Exception:
            days = []
        amount = round(max(float(m["pos_amount"] or 0), float(m["ext_amount"] or 0)), 2)
        items.append({
            "match_id":     m["match_id"],
            "source":       "match",
            "branch_id":    m["branch_id"],
            "branch_name":  bnames.get(m["branch_id"], str(m["branch_id"])),
            "period_month": m["period_month"],
            "days":         days,
            "from_cat":     from_cat,  "from_th": _CH_LABEL.get(from_cat, from_cat),
            "to_cat":       to_cat,    "to_th":   _CH_LABEL.get(to_cat, to_cat),
            "to_bank_th":   _BANK_LABEL.get(m["bank"], "") if m["bank"] in _BANK_LABEL else "",
            "amount":       amount,
            "reason":       m["note"] or "",
            "by":           m["matched_by"] or "",
            "at":           m["matched_at"] or "",
        })
    return items


def _channel_move_items(conn, period_month: str | None = None,
                        branch_id: int | None = None) -> list[dict]:
    """Manual cross-channel MOVES (recon_channel_move), shaped like _wrong_channel_items
    so the wrong-channel report lists hand-entered corrections alongside detected ones.
    Tagged source='move' (the detected ones are source='match')."""
    q = ("SELECT id, branch_id, period_month, from_category, from_bank, to_category, "
         "to_bank, amount, note, created_by, created_at FROM recon_channel_move")
    conds, params = [], []
    if period_month:
        conds.append("period_month=?"); params.append(period_month)
    if branch_id:
        conds.append("branch_id=?"); params.append(branch_id)
    if conds:
        q += " WHERE " + " AND ".join(conds)
    try:
        rows = conn.execute(q + " ORDER BY branch_id, created_at", params).fetchall()
    except sqlite3.OperationalError:
        return []
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}
    items = []
    for m in rows:
        fc, tc = m["from_category"], m["to_category"]
        items.append({
            "match_id":     None,
            "move_id":      m["id"],
            "source":       "move",
            "branch_id":    m["branch_id"],
            "branch_name":  bnames.get(m["branch_id"], str(m["branch_id"])),
            "period_month": m["period_month"],
            "days":         [],
            "from_cat":     fc, "from_th": _CH_LABEL.get(fc, fc),
            "to_cat":       tc, "to_th":   _CH_LABEL.get(tc, tc),
            "to_bank_th":   _BANK_LABEL.get(m["to_bank"], "") if m["to_bank"] in _BANK_LABEL else "",
            "amount":       round(float(m["amount"] or 0), 2),
            "reason":       m["note"] or "",
            "by":           m["created_by"] or "",
            "at":           m["created_at"] or "",
        })
    return items


# Concrete coaching guidance per (rung-as → really-was) channel confusion. Plain Thai,
# phrased as "what to do next time", so the sheet helps the branch improve — not blame.
_TIP_BY_PAIR = {
    ("CASH", "QR"):       "ลูกค้าสแกนจ่าย QR แต่ถูกกดเป็นเงินสด — เวลารับ QR ให้กดปุ่ม “QR พร้อมเพย์” และดูยอดเด้งบนแอป/สลิปให้ตรงก่อนกดยืนยัน",
    ("QR", "CASH"):       "เงินสดถูกกดเป็น QR — รับเงินสดให้กดปุ่ม “เงินสด” ทุกครั้ง อย่ากดค้างจากบิลก่อน",
    ("CASH", "TRANSFER"): "เงินโอนเข้าบัญชีถูกกดเป็นเงินสด — ถ้าลูกค้าโอน ให้กดปุ่ม “เงินโอน” และยืนยันยอดเข้าบัญชีก่อนส่งของ",
    ("TRANSFER", "CASH"): "เงินสดถูกกดเป็นเงินโอน — รับเงินสดให้กดปุ่ม “เงินสด”",
    ("QR", "TRANSFER"):   "QR กับเงินโอนสลับกัน — QR = สแกนคิวอาร์, เงินโอน = โอนเข้าเลขบัญชี เลือกปุ่มให้ตรงวิธีจ่ายจริง",
    ("TRANSFER", "QR"):   "เงินโอนกับ QR สลับกัน — ตรวจว่าลูกค้าสแกน QR หรือโอนเข้าบัญชี แล้วเลือกปุ่มให้ตรง",
    ("CASH", "CREDIT"):   "ยอดรูดบัตรถูกกดเป็นเงินสด — รูดบัตรให้กดปุ่ม “บัตรเครดิต”",
    ("CREDIT", "CASH"):   "เงินสดถูกกดเป็นบัตรเครดิต — รับเงินสดให้กดปุ่ม “เงินสด”",
}


def _coaching_tips(cat_pairs) -> list[str]:
    """One concrete tip per confusion pattern (most frequent first) + a closing habit."""
    tips, seen = [], set()
    for (fr, to), _n in cat_pairs.most_common():
        t = _TIP_BY_PAIR.get((fr, to)) or (
            f"ช่องทาง {_CH_LABEL.get(fr, fr)} → {_CH_LABEL.get(to, to)}: "
            "ก่อนกดยืนยันรับเงิน ตรวจว่าวิธีจ่ายของลูกค้าตรงกับปุ่มที่กด")
        if t not in seen:
            seen.add(t); tips.append(t)
    tips.append("ทุกสิ้นวัน ตรวจยอดแต่ละช่องทางใน POS เทียบกับเงินเข้าจริงก่อนปิดยอด — จับผิดได้ตั้งแต่วันนั้น")
    return tips


def _wrong_channel_branch_detail(conn, branch_id: int, period_month: str | None,
                                 all_items: list | None = None) -> dict:
    """Self-improvement metrics for one branch: monthly trend, next-month target,
    confusion patterns, and coaching tips."""
    from collections import defaultdict, Counter
    if all_items is None:
        all_items = _wrong_channel_items(conn, None) + _channel_move_items(conn, None)
    mine = [it for it in all_items if it["branch_id"] == branch_id]

    by_m: dict = defaultdict(lambda: {"count": 0, "amount": 0.0})
    for it in mine:
        g = by_m[it["period_month"]]; g["count"] += 1; g["amount"] += it["amount"]
    months = sorted(by_m)
    if period_month:
        months = [m for m in months if m <= period_month]
    trend = [{"month": m, "count": by_m[m]["count"], "amount": round(by_m[m]["amount"], 2)}
             for m in months[-6:]]

    scope = [it for it in mine if (not period_month or it["period_month"] == period_month)]
    pat: dict = defaultdict(lambda: {"count": 0, "amount": 0.0})
    for it in scope:
        g = pat[(it["from_th"], it["to_th"])]; g["count"] += 1; g["amount"] += it["amount"]
    patterns = sorted(
        [{"from_th": k[0], "to_th": k[1], "count": v["count"], "amount": round(v["amount"], 2)}
         for k, v in pat.items()], key=lambda x: (-x["count"], -x["amount"]))
    cat_pairs = Counter((it["from_cat"], it["to_cat"]) for it in scope)

    cur = len(scope)
    prev = trend[-2]["count"] if len(trend) >= 2 and trend[-1]["month"] == period_month else None
    if prev is None:
        direction = "flat"
    elif cur < prev:
        direction = "down"          # ดีขึ้น (น้อยลง)
    elif cur > prev:
        direction = "up"            # แย่ลง
    else:
        direction = "flat"
    target = max(0, cur // 2)       # ตั้งเป้าลดลงราวครึ่งหนึ่งของเดือนนี้
    return {"trend": trend, "current": cur, "prev": prev, "direction": direction,
            "target": target, "patterns": patterns, "tips": _coaching_tips(cat_pairs)}


def _wrong_channel_report(conn, period_month: str | None = None,
                          branch_id: int | None = None) -> dict:
    from collections import defaultdict
    items = _wrong_channel_items(conn, period_month) + _channel_move_items(conn, period_month)
    if branch_id:                                       # optional single-branch filter
        items = [it for it in items if it["branch_id"] == branch_id]
    # for cross-month trend (per-branch inside) — detected matches + manual moves
    all_items = _wrong_channel_items(conn, None) + _channel_move_items(conn, None)

    agg: dict = defaultdict(lambda: {"count": 0, "total": 0.0, "months": set()})
    for it in items:
        g = agg[(it["branch_id"], it["branch_name"])]
        g["count"] += 1; g["total"] += it["amount"]; g["months"].add(it["period_month"])
    summary = []
    for (bid, bn), v in agg.items():
        detail = _wrong_channel_branch_detail(conn, bid, period_month, all_items)
        summary.append({"branch_id": bid, "branch_name": bn, "count": v["count"],
                        "months": len(v["months"]), "total": round(v["total"], 2),
                        "trend": detail["trend"], "direction": detail["direction"],
                        "target": detail["target"], "patterns": detail["patterns"],
                        "tips": detail["tips"]})
    summary.sort(key=lambda x: -x["total"])
    return {"period_month": period_month, "summary": summary, "items": items,
            "total_count": len(items),
            "total_amount": round(sum(it["amount"] for it in items), 2)}


# ── All reconciliation errors of the month (superset of wrong-channel) ────────
# The owner wants the monthly error report to cover EVERY error found, not only
# wrong-channel mis-entries: open variances, missing documents, unmatched
# statement credits, payout shortfalls, and unmapped POS codes — each tagged by
# type + status (open / acknowledged / resolved) with its amount + reason.
_ERR_TYPE_TH = {
    "wrong_channel": "บันทึกผิดช่อง",
    "variance":      "ยอดต่าง (POS≠เงินเข้า)",
    "ext_only":      "เงินเข้าเกิน (ไม่มียอด POS)",
    "missing_doc":   "ขาดเอกสาร",
    "unmapped_code": "ยังไม่ตั้งธนาคารของรหัส",
    "payout_short":  "เงินโอนเข้าไม่ครบ/ช้า",
}


def _all_errors_items(conn, period_month: str, branch_id: int | None = None) -> list[dict]:
    """Every reconciliation error in the month across branches, tagged by type."""
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}
    if branch_id:
        branches = [(branch_id, bnames.get(branch_id, str(branch_id)))]
    else:
        branches = [(r["branch_id"], r["name"]) for r in conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name")]

    items: list[dict] = []

    def add(bid, bn, typ, channel, channel_th, detail, amount, status, reason="", by="", days=None):
        items.append({"branch_id": bid, "branch_name": bn, "period_month": period_month,
                      "type": typ, "type_th": _ERR_TYPE_TH.get(typ, typ),
                      "channel": channel, "channel_th": channel_th, "detail": detail,
                      "amount": round(float(amount or 0), 2), "status": status,
                      "reason": reason or "", "by": by or "", "days": days or []})

    # confirmed wrong-channel mis-entries — detected matches + manual moves
    for w in _wrong_channel_items(conn, period_month) + _channel_move_items(conn, period_month):
        if branch_id and w["branch_id"] != branch_id:
            continue
        tag = " (ย้ายด้วยมือ)" if w.get("source") == "move" else ""
        add(w["branch_id"], w["branch_name"], "wrong_channel", w["from_cat"], w["from_th"],
            f"{w['from_th']} → {w['to_th']}" + (f" ({w['to_bank_th']})" if w["to_bank_th"] else "") + tag,
            w["amount"], "resolved", w["reason"], w["by"], w["days"])

    for bid, bn in branches:
        try:
            rm = reconcile_month(conn, bid, period_month)
        except Exception:
            continue
        for ch in rm.get("channels", []):
            ch_th = ch["label"]
            acc = ch.get("accepted")
            ext_only = ch.get("pos", 0) <= 0.5 and ch.get("ext", 0) > 0.5
            typ = "ext_only" if ext_only else "variance"
            if ch["status"] == "diff":
                add(bid, bn, typ, ch["key"], ch_th, ch.get("detail", ""),
                    ch.get("residual", ch.get("variance", 0)), "open")
            elif acc:
                add(bid, bn, "variance", ch["key"], ch_th, "รับทราบทั้งช่องทาง",
                    ch.get("variance", 0), "acknowledged", acc.get("note", ""), acc.get("by", ""))
            elif ch.get("acked_days"):
                add(bid, bn, "variance", ch["key"], ch_th,
                    f"รับทราบรายวัน {ch['acked_days']} วัน", ch.get("acked_var", 0), "acknowledged")
        for d in rm.get("documents", []):
            if d["status"] == "missing":
                add(bid, bn, "missing_doc", d.get("channel"), d["label"],
                    d.get("detail", "ยังไม่อัปโหลด"), 0, "open")
        for q in rm.get("setup_questions", []):
            add(bid, bn, "unmapped_code", None, q["code"],
                f"รหัส {q['code']} ยังไม่ระบุธนาคาร", q.get("amount", 0), "open")
        try:
            for p in payout_reconciliation(conn, bid, period_month):
                if p["status"] == "mismatch":
                    add(bid, bn, "payout_short", p["category"],
                        _CH_LABEL.get(p["category"], p["category"]),
                        f"สุทธิควรได้ {p['report_net']:,.0f} · เข้าจริง {p['bank_payout']:,.0f}",
                        p["variance"], "open")
        except Exception:
            pass

    items.sort(key=lambda x: (x["branch_name"], x["status"] != "open", x["type"]))
    return items


def _all_errors_report(conn, period_month: str, branch_id: int | None = None) -> dict:
    from collections import Counter
    errs = _all_errors_items(conn, period_month, branch_id)
    by_type = Counter(e["type"] for e in errs)
    by_branch: dict = {}
    for e in errs:
        g = by_branch.setdefault((e["branch_id"], e["branch_name"]),
                                 {"branch_id": e["branch_id"], "branch_name": e["branch_name"],
                                  "count": 0, "open": 0, "amount": 0.0})
        g["count"] += 1
        g["open"] += 1 if e["status"] == "open" else 0
        g["amount"] += abs(e["amount"])
    branch_summary = sorted(by_branch.values(), key=lambda x: -x["open"])
    for b in branch_summary:
        b["amount"] = round(b["amount"], 2)
    return {
        "errors": errs,
        "error_counts": {t: by_type.get(t, 0) for t in _ERR_TYPE_TH},
        "error_type_th": _ERR_TYPE_TH,
        "error_branch_summary": branch_summary,
        "error_total_count": len(errs),
        "error_open_count": sum(1 for e in errs if e["status"] == "open"),
        "error_total_amount": round(sum(abs(e["amount"]) for e in errs), 2),
    }


@app.get("/api/reports/wrong-channel")
def wrong_channel_report(period_month: str | None = None, branch_id: int | None = None):
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _wrong_channel_report(conn, period_month, branch_id)
        if period_month:        # all-errors register needs a concrete month to compose
            rep.update(_all_errors_report(conn, period_month, branch_id))
        return rep


@app.get("/api/reports/wrong-channel.xlsx")
def wrong_channel_report_xlsx(period_month: str | None = None, branch_id: int | None = None):
    import io
    import openpyxl
    from openpyxl.styles import Font
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _wrong_channel_report(conn, period_month, branch_id)
        all_err = _all_errors_report(conn, period_month, branch_id) if period_month else None
    wb = openpyxl.Workbook(); bold = Font(bold=True)
    if all_err is not None:
        # Lead sheet: EVERY error of the month (the broadened report), then the
        # wrong-channel-specific sheets follow for the coaching detail.
        wse = wb.active; wse.title = "ข้อผิดพลาดทั้งหมด"
        wse.append(["สาขา", "ประเภทข้อผิดพลาด", "ช่องทาง/รหัส", "รายละเอียด",
                    "ยอด (บาท)", "สถานะ", "เหตุผล/หมายเหตุ", "โดย"])
        for cc in wse[1]:
            cc.font = bold
        _st_th = {"open": "ค้าง", "acknowledged": "รับทราบแล้ว", "resolved": "แก้ไขแล้ว"}
        for e in all_err["errors"]:
            wse.append([e["branch_name"], e["type_th"], e["channel_th"] or "—", e["detail"],
                        e["amount"], _st_th.get(e["status"], e["status"]),
                        e["reason"], e["by"]])
        for col, w in zip("ABCDEFGH", (20, 22, 18, 34, 14, 14, 38, 14)):
            wse.column_dimensions[col].width = w
        ws = wb.create_sheet("บันทึกผิดช่อง — สรุปรายสาขา")
    else:
        ws = wb.active; ws.title = "สรุปรายสาขา"
    ws.append(["สาขา", "จำนวนครั้ง", "จำนวนเดือน", "ยอดรวม (บาท)"])
    for c in ws[1]:
        c.font = bold
    for s in rep["summary"]:
        ws.append([s["branch_name"], s["count"], s["months"], s["total"]])
    ws2 = wb.create_sheet("รายการ")
    ws2.append(["สาขา", "เดือน", "วันที่", "บันทึกผิดเป็น", "ที่จริงเข้า", "ยอด (บาท)",
                "เหตุผล / หมายเหตุ", "ยืนยันโดย", "เมื่อ"])
    for c in ws2[1]:
        c.font = bold
    for it in rep["items"]:
        to_th = it["to_th"] + (f" ({it['to_bank_th']})" if it["to_bank_th"] else "")
        ws2.append([it["branch_name"], it["period_month"], ", ".join(it["days"]),
                    it["from_th"], to_th, it["amount"], it["reason"], it["by"],
                    (it["at"] or "")[:16]])
    ws.column_dimensions["A"].width = 22
    for col, w in zip("ABCDEFGHI", (22, 10, 22, 16, 18, 14, 40, 16, 18)):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf)
    fname = f"wrong_channel_{period_month or 'all'}{('_b'+str(branch_id)) if branch_id else ''}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{fname}\""})


@app.get("/api/reports/wrong-channel.print")
def wrong_channel_report_print(period_month: str | None = None,
                               branch_id: int | None = None):
    """Print-to-PDF acknowledgment sheet — one signed page per branch listing the
    wrong-channel corrections, for staff to รับทราบ + commit to fix."""
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _wrong_channel_report(conn, period_month)
        all_err = _all_errors_report(conn, period_month, branch_id) if period_month else {"errors": []}
    items = rep["items"]
    if branch_id:
        items = [it for it in items if it["branch_id"] == branch_id]

    # all errors grouped by branch (the broadened report); wrong-channel items below
    err_by_branch: dict = {}
    err_bname: dict = {}
    for e in all_err["errors"]:
        err_by_branch.setdefault(e["branch_id"], []).append(e)
        err_bname[e["branch_id"]] = e["branch_name"]

    # group wrong-channel items by branch (preserve summary order: biggest first)
    order = [s["branch_id"] for s in rep["summary"]]
    by_branch: dict = {}
    for it in items:
        by_branch.setdefault(it["branch_id"], []).append(it)
        err_bname.setdefault(it["branch_id"], it["branch_name"])
    # one sheet per branch that has ANY error (wrong-channel or otherwise)
    bids = [b for b in order if b in err_by_branch or b in by_branch]
    bids += [b for b in err_by_branch if b not in bids]
    bids += [b for b in by_branch if b not in bids]

    import html as _html
    _TH_MO = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
              "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]

    def _m(v):
        return f"{float(v or 0):,.0f}"

    def _esc(s):
        return _html.escape(str(s or ""))

    def _mth(ym):
        try:
            y, mo = ym.split("-"); return f"{_TH_MO[int(mo)]} {(int(y) + 543) % 100:02d}"
        except Exception:
            return ym

    detail_by = {s["branch_id"]: s for s in rep["summary"]}

    css = """<style>
      @page{size:A4;margin:14mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12px;margin:0;padding:24px;background:#f8fafc}
      .sheet{background:#fff;max-width:900px;margin:0 auto 18px;padding:26px 30px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:19px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:4px}
      h2{font-size:15px;margin:18px 0 8px;border-left:4px solid #ea580c;padding-left:8px}
      .muted{color:#64748b;font-size:11px} .right{text-align:right} .mono{font-variant-numeric:tabular-nums;font-family:ui-monospace,monospace}
      table{width:100%;border-collapse:collapse;margin:6px 0 10px;font-size:11.5px}
      th,td{padding:6px 8px;border-bottom:1px solid #eef2f7;text-align:left;vertical-align:top}
      th{color:#64748b;font-size:10px;border-bottom:1.5px solid #e2e8f0}
      tfoot td{border-top:1.5px solid #cbd5e1;font-weight:700}
      .arr{color:#ea580c;font-weight:700} .amt{font-weight:600}
      .pill{display:inline-block;padding:1px 9px;border-radius:99px;font-size:11px;font-weight:600}
      .p-ok{background:#dcfce7;color:#166534}.p-up{background:#fee2e2;color:#b91c1c}.p-flat{background:#f1f5f9;color:#475569}
      .trend{display:flex;gap:6px;align-items:flex-end;margin:8px 0 4px;height:64px}
      .tb{flex:1;display:flex;flex-direction:column;justify-content:flex-end;align-items:center;gap:3px;font-size:10px;color:#64748b}
      .tb .bar{width:60%;min-height:3px;background:#fdba74;border-radius:3px 3px 0 0}
      .tb.cur .bar{background:#ea580c}
      .target{background:#eff6ff;border:1px solid #bfdbfe;border-radius:8px;padding:9px 13px;margin:6px 0;font-size:12px}
      .tips{margin:6px 0 2px;padding-left:0;list-style:none}
      .tips li{position:relative;padding:5px 0 5px 24px;font-size:12px;line-height:1.5;border-bottom:1px dashed #eef2f7}
      .tips li::before{content:'💡';position:absolute;left:0;top:4px}
      .ack{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:12px 14px;margin-top:14px;font-size:12px;line-height:1.6}
      .sign{display:flex;gap:40px;margin-top:30px} .sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;font-size:11px;color:#475569}
      .toolbar{max-width:900px;margin:0 auto 14px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:8px;padding:9px 18px;font-size:13px;cursor:pointer}
      .empty{max-width:900px;margin:40px auto;text-align:center;color:#64748b}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;margin:0;max-width:none;border-radius:0;page-break-after:always}}
    </style>"""

    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
             "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
             "<title>รายงานข้อผิดพลาดประจำเดือน</title>", css, "</head><body>",
             "<div class='toolbar'><button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>"]

    if not bids:
        parts.append("<div class='empty'>ไม่พบข้อผิดพลาดในการกระทบยอด"
                     f"{' สำหรับเดือน ' + _esc(period_month) if period_month else ''} — "
                     "ทุกช่องทางตรง ไม่มียอดต่าง เอกสารครบ และไม่มีการบันทึกผิดช่อง 🎉</div>")
    period_label = period_month or "ทุกเดือน"
    _ST_PILL = {"open": ("p-up", "ค้าง"), "acknowledged": ("p-flat", "รับทราบแล้ว"),
                "resolved": ("p-ok", "แก้ไขแล้ว")}
    for bid in bids:
        rows = by_branch.get(bid, [])
        errs = err_by_branch.get(bid, [])
        bn = err_bname.get(bid, str(bid))
        n_open = sum(1 for e in errs if e["status"] == "open")
        parts.append("<div class='sheet'>")
        parts.append(f"<h1>รายงานข้อผิดพลาดประจำเดือน — สาขา{_esc(bn)}</h1>")
        parts.append(f"<div class='sub'>งวด {_esc(period_label)} · ข้อผิดพลาดทั้งหมด {len(errs)} รายการ "
                     f"(ค้าง {n_open}) · บันทึกผิดช่อง {len(rows)} รายการ</div>")

        # ── ข้อผิดพลาดทั้งหมดในเดือนนี้ (ทุกประเภท) ──
        if errs:
            parts.append("<h2>ข้อผิดพลาดทั้งหมดในเดือนนี้ (ทุกประเภท)</h2>")
            parts.append("<table><thead><tr><th>ประเภท</th><th>ช่องทาง/รหัส</th><th>รายละเอียด</th>"
                         "<th class='right'>ยอด (บาท)</th><th>สถานะ</th><th>เหตุผล/หมายเหตุ</th></tr></thead><tbody>")
            for e in errs:
                pill = _ST_PILL.get(e["status"], ("p-flat", e["status"]))
                parts.append(
                    f"<tr><td>{_esc(e['type_th'])}</td><td>{_esc(e['channel_th'] or '—')}</td>"
                    f"<td>{_esc(e['detail'])}</td>"
                    f"<td class='right mono amt'>{_m(e['amount']) if e['amount'] else '—'}</td>"
                    f"<td><span class='pill {pill[0]}'>{pill[1]}</span></td>"
                    f"<td class='muted'>{_esc(e['reason'] or '—')}</td></tr>")
            parts.append("</tbody></table>")

        if not rows:                          # branch has other errors but no wrong-channel mis-entries
            parts.append("<div class='sign'><div>ผู้จัดการสาขา</div><div>แคชเชียร์ / ผู้บันทึก</div>"
                         "<div>ฝ่ายบัญชี / ผู้ตรวจสอบ</div></div></div>")
            continue

        tot = sum(r["amount"] for r in rows)
        parts.append("<h2>รายการที่บันทึกผิดช่อง (รายละเอียด + แนวทางพัฒนา)</h2>")
        parts.append(f"<div class='sub'>มูลค่ารวมบันทึกผิดช่อง ฿{_m(tot)}</div>")
        parts.append("<table><thead><tr><th>วันที่</th><th>บันทึกผิดเป็น</th>"
                     "<th>ที่จริงเข้า</th><th class='right'>ยอด (บาท)</th>"
                     "<th>เหตุผล / หมายเหตุ</th><th>ยืนยันโดย</th></tr></thead><tbody>")
        for r in rows:
            to_th = _esc(r["to_th"]) + (f" ({_esc(r['to_bank_th'])})" if r["to_bank_th"] else "")
            days = ", ".join(d[8:].lstrip("0") for d in r["days"]) or "—"
            parts.append(
                f"<tr><td>{days}</td><td>{_esc(r['from_th'])}</td>"
                f"<td><span class='arr'>→</span> {to_th}</td>"
                f"<td class='right mono amt'>{_m(r['amount'])}</td>"
                f"<td>{_esc(r['reason'] or '—')}</td>"
                f"<td class='muted'>{_esc(r['by'] or '—')}<br>{_esc((r['at'] or '')[:16])}</td></tr>")
        parts.append(f"</tbody><tfoot><tr><td colspan='3'>รวม</td>"
                     f"<td class='right mono'>{_m(tot)}</td><td colspan='2'></td></tr></tfoot></table>")

        det = detail_by.get(bid, {})
        # ① แนวโน้มรายเดือน + เป้าหมาย
        trend = det.get("trend") or []
        if trend:
            mx = max((t["count"] for t in trend), default=0) or 1
            dirn = det.get("direction", "flat")
            badge = ({"down": ("p-ok", "✓ ดีขึ้น (ผิดน้อยลง)"),
                      "up": ("p-up", "▲ ผิดมากขึ้น — ต้องระวัง"),
                      "flat": ("p-flat", "— ทรงตัว")}).get(dirn, ("p-flat", "—"))
            parts.append(f"<h2>① แนวโน้มการกดผิดช่อง <span class='pill {badge[0]}'>{badge[1]}</span></h2>")
            parts.append("<div class='trend'>")
            for t in trend:
                cur = "cur" if t["month"] == period_month else ""
                h = int(8 + 52 * t["count"] / mx)
                parts.append(f"<div class='tb {cur}'><b>{t['count']}</b>"
                             f"<div class='bar' style='height:{h}px'></div>{_esc(_mth(t['month']))}</div>")
            parts.append("</div>")
            tg = det.get("target", 0)
            parts.append(f"<div class='target'>🎯 <b>เป้าหมายเดือนถัดไป: ไม่เกิน {tg} ครั้ง</b> "
                         f"(เดือนนี้ {det.get('current', len(rows))} ครั้ง) — "
                         "ตั้งเป้าลดลงครึ่งหนึ่ง โดยตรวจปุ่มรับเงินให้ตรงทุกบิล</div>")
        # ② ช่องที่สับสนบ่อย
        patterns = det.get("patterns") or []
        if patterns:
            parts.append("<h2>② ช่องทางที่สับสนบ่อย (จุดที่ต้องระวัง)</h2>")
            parts.append("<table><thead><tr><th>กดผิดเป็น</th><th>ที่จริงคือ</th>"
                         "<th class='right'>จำนวนครั้ง</th><th class='right'>ยอดรวม</th></tr></thead><tbody>")
            for p in patterns:
                parts.append(f"<tr><td>{_esc(p['from_th'])}</td>"
                             f"<td><span class='arr'>→</span> {_esc(p['to_th'])}</td>"
                             f"<td class='right mono amt'>{p['count']}</td>"
                             f"<td class='right mono'>{_m(p['amount'])}</td></tr>")
            parts.append("</tbody></table>")
        # ③ คำแนะนำสำหรับสาขา (โค้ช)
        tips = det.get("tips") or []
        if tips:
            parts.append("<h2>③ คำแนะนำเพื่อพัฒนา</h2><ul class='tips'>")
            for t in tips:
                parts.append(f"<li>{_esc(t)}</li>")
            parts.append("</ul>")

        parts.append(
            "<div class='ack'>ข้าพเจ้าในนามสาขา <b>รับทราบ</b> ว่ารายการข้างต้นเกิดจากการบันทึก"
            "ช่องทางการรับเงินผิดในเครื่อง POS และ <b>รับว่าจะแก้ไข</b>ขั้นตอนการกดรับเงินให้ถูกต้อง "
            "เพื่อไม่ให้เกิดซ้ำในรอบถัดไป (ยอดเงินรวมถูกต้อง ไม่มีเงินขาด — เป็นการลงผิดช่องเท่านั้น)</div>")
        parts.append("<div class='sign'>"
                     "<div>ผู้จัดการสาขา (เซ็นรับทราบ)</div>"
                     "<div>แคชเชียร์ / ผู้บันทึก</div>"
                     "<div>ฝ่ายบัญชี / ผู้ตรวจสอบ</div></div>")
        parts.append("</div>")

    parts.append("</body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


# ── Statement-code dictionary (รหัสในสเตทเมนต์ → ความหมาย → ช่องทาง) ────────────
# Surfaces parsers.bank.STATEMENT_CODE_MAP (the single source the parser uses) so the
# team can review every abbreviation and flag wrong meanings (the SDCH=cash case).

# parser channel → (recon category, plain-Thai label)
_PARSE_CH_TH = {
    "QR":             ("QR",      "QR พร้อมเพย์"),
    "TRANSFER":       ("TRANSFER", "เงินโอน"),
    "CASH_DEP":       ("CASH",    "เงินสด (ฝากเข้าบัญชี)"),
    "CC_SETTLE":      ("CREDIT",  "บัตรเครดิต (settlement)"),
    "GRAB_PAYOUT":    ("GRAB",    "Grab (เงินโอนยอด)"),
    "LINEMAN_PAYOUT": ("LINEMAN", "Line Man (เงินโอนยอด)"),
    "FEE":            ("—",       "ค่าธรรมเนียม (ไม่ใช่รายได้)"),
    "OTHER":          ("—",       "อื่น ๆ / ไม่จัดประเภท"),
}


_OVERRIDE_CHANNELS = {"QR", "TRANSFER", "CASH_DEP", "CC_SETTLE",
                      "GRAB_PAYOUT", "LINEMAN_PAYOUT", "FEE", "OTHER"}


def _statement_code_overrides(conn) -> dict:
    """{CODE(upper): channel} — manager remaps that win over the hardcoded map."""
    try:
        return {r["code"].upper(): r["channel"]
                for r in conn.execute("SELECT code, channel FROM recon_statement_code_map")}
    except sqlite3.OperationalError:
        return {}


def _override_channel(description: str | None, base_channel: str | None, overrides: dict) -> str | None:
    """Override wins when its code is a substring of the line description (longest code
    first); otherwise the parser's base channel stands."""
    if overrides:
        d = (description or "").upper()
        hits = [(code, ch) for code, ch in overrides.items() if code and code in d]
        if hits:
            return max(hits, key=lambda x: len(x[0]))[1]
    return base_channel


def _reapply_statement_overrides(conn, code: str) -> int:
    """Re-classify stored bank lines whose description contains `code`, from the
    parser base + current overrides. Used on save (apply) and delete (revert)."""
    overrides = _statement_code_overrides(conn)
    like = f"%{code.upper()}%"
    rows = conn.execute(
        "SELECT id, description FROM recon_bank_lines "
        "WHERE UPPER(COALESCE(description,'')) LIKE ?", (like,)).fetchall()
    n = 0
    for r in rows:
        base = _normalise_channel(r["description"])
        newch = _override_channel(r["description"], base, overrides)
        conn.execute("UPDATE recon_bank_lines SET channel=? WHERE id=?", (newch, r["id"]))
        n += 1
    return n


def _statement_codes(conn=None) -> list[dict]:
    overrides = _statement_code_overrides(conn) if conn is not None else {}
    out, matched = [], set()
    for e in STATEMENT_CODE_MAP:
        cat, ch_th = _PARSE_CH_TH.get(e["channel"], (e["channel"], e["channel"]))
        eff, ov_code = e["channel"], None
        for code, ch in overrides.items():       # an override whose code is in this entry's pattern
            if code in e["pattern"].upper():
                eff, ov_code = ch, code; matched.add(code)
        eff_cat, eff_th = _PARSE_CH_TH.get(eff, (eff, eff))
        out.append({**e, "category": cat, "channel_th": ch_th,
                    "effective_channel": eff, "effective_channel_th": eff_th,
                    "effective_category": eff_cat,
                    "overridden": eff != e["channel"], "override_code": ov_code})
    for code, ch in overrides.items():           # codes the team added that aren't in the base map
        if code not in matched:
            eff_cat, eff_th = _PARSE_CH_TH.get(ch, (ch, ch))
            out.append({"pattern": code, "code": code, "kind": "code", "bank": "*",
                        "meaning_th": "(เพิ่มโดยทีม)", "status": "confirmed", "note": "",
                        "channel": ch, "category": eff_cat, "channel_th": eff_th,
                        "effective_channel": ch, "effective_channel_th": eff_th,
                        "effective_category": eff_cat, "overridden": True, "override_code": code})
    return out


class CodeOverride(BaseModel):
    code: str
    channel: str
    note: str | None = None
    by: str | None = None


@app.get("/api/statement-codes")
def statement_codes():
    with get_db() as conn:
        codes = _statement_codes(conn)
    return {"codes": codes,
            "review_count": sum(1 for c in codes if c["status"] == "review"),
            "override_count": sum(1 for c in codes if c["overridden"]),
            "channel_options": sorted(_OVERRIDE_CHANNELS)}


@app.post("/api/statement-codes/override", dependencies=[Depends(require_manager_token)])
def set_statement_override(body: CodeOverride):
    code = (body.code or "").strip().upper()
    if not code:
        raise HTTPException(400, "ต้องระบุรหัส")
    if body.channel not in _OVERRIDE_CHANNELS:
        raise HTTPException(400, f"ช่องทางต้องเป็นหนึ่งใน {sorted(_OVERRIDE_CHANNELS)}")
    with get_db() as conn:
        conn.execute(
            "INSERT INTO recon_statement_code_map (code, channel, note, updated_by, updated_at) "
            "VALUES (?,?,?,?,datetime('now')) "
            "ON CONFLICT(code) DO UPDATE SET channel=excluded.channel, note=excluded.note, "
            "updated_by=excluded.updated_by, updated_at=datetime('now')",
            (code, body.channel, body.note or "", body.by or "manager"))
        n = _reapply_statement_overrides(conn, code)
        _audit(conn, "statement_code", 0, "override", body.by or "manager",
               [{"field": code, "old": None, "new": body.channel}])
        conn.commit()
    return {"code": code, "channel": body.channel, "lines_reclassified": n}


@app.delete("/api/statement-codes/override/{code}", dependencies=[Depends(require_manager_token)])
def del_statement_override(code: str):
    code_u = (code or "").strip().upper()
    with get_db() as conn:
        conn.execute("DELETE FROM recon_statement_code_map WHERE code=?", (code_u,))
        n = _reapply_statement_overrides(conn, code_u)   # revert lines to base classification
        _audit(conn, "statement_code", 0, "revert", "manager",
               [{"field": code_u, "old": None, "new": None}])
        conn.commit()
    return {"code": code_u, "removed": True, "lines_reclassified": n}


# ── Reclass-pattern guard (catch a misclassified statement code early) ─────────
# The SDCH bug showed up as a flood of identical wrong-channel suggestions: same
# from→to, same day-gap, cross-bank, one direction — a structural fingerprint, not
# scattered human error. This diagnostic clusters live suggestions and flags those
# that look systematic, so the next mis-mapped code is caught before staff "confirm".

@app.get("/api/diagnostics/reclass-pattern")
def reclass_pattern_diagnostic(month: str, min_cluster: int = 6):
    _validate_period(month)
    from collections import defaultdict
    groups: dict = defaultdict(lambda: {"count": 0, "amount": 0.0, "branches": set()})
    with get_db() as conn:
        bids = [r["branch_id"] for r in
                conn.execute(f"SELECT branch_id FROM branches WHERE {_MASTER_ONLY}")]
        for bid in bids:
            try:
                sugs = suggest_reclassifications(conn, bid, month)
            except Exception:
                sugs = []
            for s in sugs:
                k = (s.get("pos_cat"), s.get("ext_cat"), s.get("day_gap"),
                     s.get("pos_bank") == s.get("ext_bank"))
                g = groups[k]
                g["count"] += 1; g["amount"] += s.get("amount", 0) or 0; g["branches"].add(bid)
    clusters = []
    for (fr, to, gap, same), g in groups.items():
        clusters.append({
            "from_cat": fr, "from_th": _CH_LABEL.get(fr, fr),
            "to_cat": to,   "to_th": _CH_LABEL.get(to, to),
            "day_gap": gap, "same_bank": same,
            "count": g["count"], "amount": round(g["amount"], 2), "branches": len(g["branches"]),
            # systematic smell: many identical cross-bank pairs across branches
            "suspicious": g["count"] >= min_cluster and not same and len(g["branches"]) >= 2,
        })
    clusters.sort(key=lambda x: -x["count"])
    flagged = [c for c in clusters if c["suspicious"]]
    warning = None
    if flagged:
        t = flagged[0]
        warning = (f"พบ {t['count']} คู่รูปแบบเดียวกัน ({t['from_th']} → {t['to_th']}, "
                   f"ห่าง {t['day_gap']} วัน, คนละธนาคาร) ใน {t['branches']} สาขา — "
                   "ดูเป็นรูปแบบของระบบมากกว่าคนกดผิด อาจมีรหัสสเตทเมนต์จัดประเภทผิด (เช่นเคส SDCH) "
                   "ควรตรวจ “พจนานุกรมรหัสในสเตทเมนต์” ก่อนให้สาขายืนยัน")
    return {"month": month, "min_cluster": min_cluster,
            "clusters": clusters, "flagged": flagged, "warning": warning}


@app.get("/api/statement-codes.xlsx")
def statement_codes_xlsx():
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); bold = Font(bold=True)
    warn = PatternFill("solid", fgColor="FFF3CD")
    ws = wb.active; ws.title = "รหัสในสเตทเมนต์"
    ws.append(["รหัส / คำสำคัญ", "ชนิด", "ธนาคาร", "ความหมาย (ไทย)", "ระบบจัดเป็นช่องทาง",
               "สถานะ", "หมายเหตุระบบ", "ทีมยืนยัน (✓/✗)", "ความเห็น/แก้ไขของทีม"])
    for c in ws[1]:
        c.font = bold
    with get_db() as conn:
        codes = _statement_codes(conn)
    for c in codes:
        kind = "รหัสธนาคาร" if c["kind"] == "code" else "คำทั่วไป"
        st = ("✏️ ทีมแก้แล้ว" if c["overridden"]
              else "✅ ยืนยันแล้ว" if c["status"] == "confirmed" else "⚠️ รอตรวจ")
        ch_th = c.get("effective_channel_th", c["channel_th"])
        r = [c["code"], kind, c["bank"], c["meaning_th"], ch_th, st, c["note"], "", ""]
        ws.append(r)
        if c["status"] == "review":
            for cell in ws[ws.max_row]:
                cell.fill = warn
    for col, w in zip("ABCDEFGHI", (30, 12, 8, 46, 24, 14, 50, 14, 36)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=\"statement_codes.xlsx\""})


@app.get("/api/statement-codes.html")
def statement_codes_html():
    """Printable / shareable review sheet of every statement code the parser knows,
    grouped by channel, with a column for the team to confirm or correct."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    with get_db() as conn:
        codes = _statement_codes(conn)
    review_n = sum(1 for c in codes if c["status"] == "review")
    css = """<style>
      @page{size:A4 landscape;margin:12mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12px;margin:0;padding:24px;background:#f8fafc}
      .sheet{background:#fff;max-width:1100px;margin:0 auto;padding:26px 30px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:20px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:14px}
      table{width:100%;border-collapse:collapse;font-size:11.5px}
      th,td{padding:6px 9px;border-bottom:1px solid #eef2f7;text-align:left;vertical-align:top}
      th{color:#475569;font-size:10.5px;border-bottom:1.5px solid #e2e8f0;background:#f8fafc}
      code{background:#f1f5f9;padding:1px 6px;border-radius:5px;font-family:ui-monospace,monospace;font-weight:600}
      .badge{padding:1px 8px;border-radius:99px;font-size:10px;font-weight:600;white-space:nowrap}
      .b-ok{background:#dcfce7;color:#166534} .b-warn{background:#fef3c7;color:#854d0e}
      .b-edit{background:#dbeafe;color:#1e40af}
      .b-code{background:#e0e7ff;color:#3730a3} .b-kw{background:#f1f5f9;color:#475569}
      tr.review{background:#fffbeb} .note{color:#64748b;font-size:11px}
      .fill{color:#cbd5e1} .toolbar{max-width:1100px;margin:0 auto 14px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:8px;padding:9px 18px;font-size:13px;cursor:pointer;text-decoration:none}
      .lead{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:10px 14px;margin-bottom:14px;font-size:12px;line-height:1.6}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;border-radius:0;max-width:none}}
    </style>"""
    rows = []
    for c in codes:
        kind = ('<span class="badge b-code">รหัสธนาคาร</span>' if c["kind"] == "code"
                else '<span class="badge b-kw">คำทั่วไป</span>')
        st = ('<span class="badge b-edit">✏️ ทีมแก้แล้ว</span>' if c["overridden"]
              else '<span class="badge b-ok">✅ ยืนยันแล้ว</span>' if c["status"] == "confirmed"
              else '<span class="badge b-warn">⚠️ รอตรวจ</span>')
        ch_th = c.get("effective_channel_th", c["channel_th"])
        ch_cell = (f'{esc(ch_th)} <span class="note">(เดิม {esc(c["channel_th"])})</span>'
                   if c["overridden"] else esc(ch_th))
        rows.append(
            f'<tr class="{ "review" if c["status"]=="review" and not c["overridden"] else "" }">'
            f'<td><code>{esc(c["code"])}</code></td><td>{kind}</td><td>{esc(c["bank"])}</td>'
            f'<td>{esc(c["meaning_th"])}{("<div class=note>"+esc(c["note"])+"</div>") if c["note"] else ""}</td>'
            f'<td>{ch_cell}</td><td>{st}</td>'
            f'<td class="fill">▢ ใช่ &nbsp; ▢ ไม่ใช่</td><td class="fill">…………………………</td></tr>')
    html = (
        "<!doctype html><html lang='th'><head><meta charset='utf-8'>"
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>"
        "<title>พจนานุกรมรหัสในสเตทเมนต์</title>" + css + "</head><body>"
        "<div class='toolbar'>"
        "<a class='btn' href='/api/statement-codes.xlsx' style='text-decoration:none;margin-right:8px;background:#16794f'>⬇ ดาวน์โหลด Excel</a>"
        "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>"
        "<div class='sheet'>"
        "<h1>พจนานุกรมรหัสในสเตทเมนต์ธนาคาร</h1>"
        f"<div class='sub'>ระบบใช้ตารางนี้แยกประเภทรายการในสเตทเมนต์อัตโนมัติ · {len(codes)} รายการ · "
        f"<b>{review_n} รายการรอทีมตรวจ</b></div>"
        "<div class='lead'>ช่วยตรวจ <b>คอลัมน์ \"ความหมาย\" กับ \"ระบบจัดเป็นช่องทาง\"</b> ว่าตรงกับความจริงไหม "
        "โดยเฉพาะแถวสีเหลือง (⚠️ รอตรวจ) ที่เป็นรหัสเฉพาะของธนาคาร เช่น <code>SDCH</code> "
        "(เจ้าของยืนยันว่า = นำเงินสดไปฝาก ไม่ใช่ QR) · ทำเครื่องหมายในช่อง \"ทีมยืนยัน\" แล้วเขียนสิ่งที่ถูกต้องในช่องสุดท้าย</div>"
        "<table><thead><tr><th>รหัส / คำสำคัญ</th><th>ชนิด</th><th>ธนาคาร</th>"
        "<th>ความหมาย</th><th>ระบบจัดเป็นช่องทาง</th><th>สถานะ</th>"
        "<th>ทีมยืนยัน</th><th>ที่ถูกต้องคือ</th></tr></thead><tbody>"
        + "".join(rows) +
        "</tbody></table>"
        "<div style='margin-top:22px;display:flex;gap:50px'>"
        "<div style='border-top:1px solid #94a3b8;padding-top:5px;flex:1;text-align:center;font-size:11px;color:#475569'>ผู้ตรวจสอบ (ทีมบัญชี)</div>"
        "<div style='border-top:1px solid #94a3b8;padding-top:5px;flex:1;text-align:center;font-size:11px;color:#475569'>วันที่</div>"
        "</div></div></body></html>")
    return Response(content=html, media_type="text/html; charset=utf-8")


# ── Discrepancies ─────────────────────────────────────────────────────────────

@app.get("/api/sessions/{session_id}/discrepancies")
def list_discrepancies(
    session_id: int,
    resolution: str | None = "pending",
    category:   str | None = None,
    bank:       str | None = None,
):
    with get_db() as conn:
        filters: list[str] = ["session_id=?"]
        params:  list      = [session_id]
        if resolution and resolution != "all":
            filters.append("resolution=?"); params.append(resolution)
        if category:
            filters.append("category=?"); params.append(category)
        if bank:
            filters.append("bank=?"); params.append(bank)
        rows = conn.execute(f"""
            SELECT * FROM recon_discrepancies
            WHERE {' AND '.join(filters)}
            ORDER BY ABS(COALESCE(variance, amount_external)) DESC
        """, params).fetchall()
        return [dict(r) for r in rows]


class ResolveBody(BaseModel):
    resolution:      Literal["accepted", "adjusted", "rejected"]
    resolution_note: str = ""
    resolved_by:     str = ""


def _assert_disc_period_open(conn, disc_id: int):
    row = conn.execute("""
        SELECT s.branch_id, s.period_month
        FROM recon_discrepancies d JOIN recon_sessions s ON s.session_id = d.session_id
        WHERE d.id=?""", (disc_id,)).fetchone()
    if row:
        _assert_period_open(conn, row["branch_id"], row["period_month"])


def _audit(conn, entity: str, entity_id: int, action: str, by: str,
           diffs: list[dict] | None = None):
    """Record create/update/delete with optional per-field diffs.
    diffs = [{"field": "status", "old": "draft", "new": "approved"}, ...]
    """
    ts = _dt.utcnow().isoformat(timespec="seconds")
    if not diffs:
        conn.execute(
            "INSERT INTO recon_audit_log(ts,entity,entity_id,action,field,old_val,new_val,by)"
            " VALUES(?,?,?,?,?,?,?,?)",
            (ts, entity, entity_id, action, None, None, None, by or "staff"),
        )
    else:
        for d in diffs:
            old = None if d.get("old") is None else str(d["old"])
            new = None if d.get("new") is None else str(d["new"])
            conn.execute(
                "INSERT INTO recon_audit_log(ts,entity,entity_id,action,field,old_val,new_val,by)"
                " VALUES(?,?,?,?,?,?,?,?)",
                (ts, entity, entity_id, action, d["field"], old, new, by or "staff"),
            )


@app.get("/api/discrepancies/{disc_id}/days")
def disc_daily_breakdown(disc_id: int):
    """Per-day POS + external breakdown for one discrepancy (same lane, filtered daily)."""
    with get_db() as conn:
        disc = conn.execute(
            """SELECT d.category, d.bank, s.branch_id, s.period_month
               FROM recon_discrepancies d
               JOIN recon_sessions s ON s.session_id = d.session_id
               WHERE d.id=?""", (disc_id,)).fetchone()
        if not disc:
            raise HTTPException(404)
        cat, bank, bid, pm = disc["category"], disc["bank"], disc["branch_id"], disc["period_month"]
        daily = daily_reconciliation(conn, bid, pm)
        days_out = []
        for day in daily.get("days", []):
            for lane in day.get("lanes", []):
                if lane["category"] == cat and lane["bank"] == (bank or "NONE"):
                    days_out.append({
                        "day":      day["day"],
                        "pos":      lane["pos"],
                        "ext":      lane["ext"],
                        "ext_net":  lane["ext_net"],
                        "variance": round((lane["ext"] or 0) - (lane["pos"] or 0), 2),
                        "status":   lane["status"],
                    })
        return {"category": cat, "bank": bank, "period_month": pm, "days": days_out}


@app.patch("/api/discrepancies/{disc_id}/resolve", dependencies=[Depends(require_write_token)])
def resolve_discrepancy(disc_id: int, body: ResolveBody):
    with get_db() as conn:
        _assert_disc_period_open(conn, disc_id)
        old = conn.execute(
            "SELECT resolution, resolution_note FROM recon_discrepancies WHERE id=?",
            (disc_id,)).fetchone()
        conn.execute("""
            UPDATE recon_discrepancies
            SET resolution=?, resolution_note=?, resolved_by=?, resolved_at=datetime('now')
            WHERE id=?
        """, (body.resolution, body.resolution_note, body.resolved_by, disc_id))
        diffs = []
        if old:
            if old["resolution"] != body.resolution:
                diffs.append({"field": "resolution", "old": old["resolution"], "new": body.resolution})
            if (old["resolution_note"] or "") != (body.resolution_note or ""):
                diffs.append({"field": "resolution_note",
                               "old": old["resolution_note"], "new": body.resolution_note})
        _audit(conn, "discrepancy", disc_id, "update", body.resolved_by or "staff", diffs or None)
        conn.commit()
    return {"id": disc_id, "resolution": body.resolution}


@app.patch("/api/discrepancies/{disc_id}/assign", dependencies=[Depends(require_write_token)])
def assign_discrepancy(disc_id: int, assigned_to: str = Form(...)):
    with get_db() as conn:
        _assert_disc_period_open(conn, disc_id)
        old = conn.execute(
            "SELECT assigned_to FROM recon_discrepancies WHERE id=?", (disc_id,)).fetchone()
        conn.execute("UPDATE recon_discrepancies SET assigned_to=? WHERE id=?", (assigned_to, disc_id))
        _audit(conn, "discrepancy", disc_id, "update", "staff",
               [{"field": "assigned_to",
                 "old": old["assigned_to"] if old else None, "new": assigned_to}])
        conn.commit()
    return {"id": disc_id, "assigned_to": assigned_to}


class BulkResolveBody(BaseModel):
    ids:             list[int]
    resolution:      Literal["accepted", "adjusted", "rejected"]
    resolution_note: str = ""
    resolved_by:     str = ""


@app.post("/api/sessions/{session_id}/discrepancies/bulk-resolve",
          dependencies=[Depends(require_write_token)])
def bulk_resolve_discrepancies(session_id: int, body: BulkResolveBody):
    if not body.ids:
        raise HTTPException(400, "ids list must not be empty")
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        _assert_period_open(conn, s["branch_id"], s["period_month"])
        ph = ",".join("?" * len(body.ids))
        conn.execute(f"""
            UPDATE recon_discrepancies
            SET resolution=?, resolution_note=?, resolved_by=?, resolved_at=datetime('now')
            WHERE session_id=? AND id IN ({ph}) AND resolution='pending'
        """, [body.resolution, body.resolution_note, body.resolved_by or "staff", session_id]
             + list(body.ids))
        _audit(conn, "session", session_id, "update", body.resolved_by or "staff",
               [{"field": "bulk_resolve",
                 "old": "pending",
                 "new": f"{body.resolution} × {len(body.ids)} items"}])
        conn.commit()
        updated = conn.execute(
            f"SELECT COUNT(*) FROM recon_discrepancies WHERE id IN ({ph}) AND resolution=?",
            list(body.ids) + [body.resolution]).fetchone()[0]
    return {"updated": updated, "resolution": body.resolution}


# ── Manual fee / VAT per line (Grab / Line Man / Credit Card) ──────────────────

_FEE_VAT_TABLE = {
    "grab":        "recon_grab_lines",
    "lineman":     "recon_lineman_lines",
    "credit_card": "recon_cc_lines",
}


class FeeVat(BaseModel):
    fee_manual: float | None = None
    vat_manual: float | None = None
    by:         str = ""


@app.get("/api/lines/{source}")
def list_lines(source: str, branch_id: int, period_month: str):
    """Lines for the fee/VAT editor, with effective fee/VAT and computed net."""
    if source not in _FEE_VAT_TABLE:
        raise HTTPException(400, f"source must be one of {sorted(_FEE_VAT_TABLE)}")
    table = _FEE_VAT_TABLE[source]
    spec = {
        "grab":        ("order_date",      "order_id",  "food_amount",  "commission_amt", "vat_amount",       "promo_deduct"),
        "lineman":     ("order_date",      "order_id",  "food_amount",  "commission_amt", "vat_on_commission","promo_deduct"),
        "credit_card": ("settlement_date", "batch_no",  "gross_amount", "mdr_fee",        "vat_amount",       None),
    }[source]
    date_col, ref_col, gross_col, fee_col, vat_col, promo_col = spec
    out = []
    with get_db() as conn:
        extra = " AND order_status='Completed'" if source in ("grab", "lineman") else ""
        rows = conn.execute(f"""
            SELECT * FROM {table}
            WHERE branch_id=? AND strftime('%Y-%m', {date_col})=?{extra}
            ORDER BY {date_col}
        """, (branch_id, period_month)).fetchall()
        for r in rows:
            gross  = float(r[gross_col] or 0)
            fee    = r["fee_manual"] if r["fee_manual"] is not None else float(r[fee_col] or 0)
            if r["vat_manual"] is not None:
                vat = r["vat_manual"]
            else:
                vat = float(r[vat_col] or 0) if vat_col else 0.0
            promo  = float(r[promo_col] or 0) if promo_col else 0.0
            refund = float(r["refund_amount"] or 0) if source == "credit_card" else 0.0
            out.append({
                "id":         r["id"],
                "date":       r[date_col],
                "ref":        r[ref_col] or "",
                "gross":      round(gross, 2),
                "fee":        r["fee_manual"],
                "vat":        r["vat_manual"],
                "fee_parsed": round(float(r[fee_col] or 0), 2),
                "vat_parsed": round(float(r[vat_col] or 0) if vat_col else 0.0, 2),
                "other":      round(refund + promo, 2),   # non-fee deductions → client live-recompute
                "net":        round(gross - refund - fee - vat - promo, 2),
            })
    return out


@app.patch("/api/lines/{source}/{line_id}/fee-vat", dependencies=[Depends(require_write_token)])
def set_fee_vat(source: str, line_id: int, body: FeeVat):
    table = _FEE_VAT_TABLE.get(source)
    if not table:
        raise HTTPException(400, f"source must be one of {sorted(_FEE_VAT_TABLE)}")
    date_col = "settlement_date" if source == "credit_card" else "order_date"
    with get_db() as conn:
        line = conn.execute(
            f"SELECT branch_id, {date_col} AS d, fee_manual, vat_manual FROM {table} WHERE id=?",
            (line_id,)).fetchone()
        if line and line["d"]:
            _assert_period_open(conn, line["branch_id"], str(line["d"])[:7])
        conn.execute(f"""
            UPDATE {table}
            SET fee_manual=?, vat_manual=?, fee_vat_by=?, fee_vat_at=datetime('now')
            WHERE id=?
        """, (body.fee_manual, body.vat_manual, body.by or "staff", line_id))
        if line:
            diffs = []
            if line["fee_manual"] != body.fee_manual:
                diffs.append({"field": "fee_manual", "old": line["fee_manual"], "new": body.fee_manual})
            if line["vat_manual"] != body.vat_manual:
                diffs.append({"field": "vat_manual", "old": line["vat_manual"], "new": body.vat_manual})
            if diffs:
                _audit(conn, "line", line_id, "update", body.by or "staff", diffs)
        conn.commit()
    return {"source": source, "line_id": line_id,
            "fee_manual": body.fee_manual, "vat_manual": body.vat_manual}


# ── Matches (reconciled category buckets from the last run-match) ─────────────

@app.get("/api/sessions/{session_id}/matches")
def list_matches(session_id: int, category: str | None = None):
    with get_db() as conn:
        filters = ["session_id=?", "status='matched'"]
        params: list = [session_id]
        if category:
            filters.append("category=?"); params.append(category)
        rows = conn.execute(f"""
            SELECT * FROM recon_category_result
            WHERE {' AND '.join(filters)}
            ORDER BY category, bank
        """, params).fetchall()
        return [dict(r) for r in rows]


# ── Cross-branch daily board ──────────────────────────────────────────────────

_DAY_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])-(0[1-9]|[12]\d|3[01])$")


@app.get("/api/daily-board/dates")
def daily_board_dates():
    """Min / max day with POS daily data — drives the date picker bounds + default.

    min_day is floored to RECON_START_DAY so the board never surfaces pre-cutoff
    data (the historical rows stay in the DB for the dashboard)."""
    with get_db() as conn:
        try:
            row = conn.execute(
                "SELECT MIN(day) mn, MAX(day) mx FROM pos_payment_day WHERE day >= ?",
                (RECON_START_DAY,)).fetchone()
            return {"min_day": row["mn"] or RECON_START_DAY, "max_day": row["mx"]}
        except sqlite3.OperationalError:
            return {"min_day": None, "max_day": None}


def _coupon_range(gran: str, period_month: str, day: str | None):
    """Resolve (date_from, date_to, label) for a coupon view. gran = day|week|month.
    Day strings compare lexicographically, so '-31' is a safe month upper bound."""
    if gran == "day":
        if not (day and _DAY_RE.match(day)):
            raise HTTPException(400, "gran=day needs a valid 'day' (YYYY-MM-DD)")
        return day, day, day
    if gran == "week":
        if not (day and _DAY_RE.match(day)):
            raise HTTPException(400, "gran=week needs a valid 'day' (YYYY-MM-DD)")
        d = _dt.strptime(day, "%Y-%m-%d").date()
        mon = d - _td(days=d.weekday())
        sun = mon + _td(days=6)
        return mon.isoformat(), sun.isoformat(), f"{mon.isoformat()} – {sun.isoformat()}"
    _validate_period(period_month)
    return period_month + "-01", period_month + "-31", period_month


def _coupon_prev(gran: str, period_month: str, day: str | None):
    """Previous comparable period (prev day / prev week / prev month) for MoM."""
    if gran == "day":
        d = _dt.strptime(day, "%Y-%m-%d").date() - _td(days=1)
        return d.isoformat(), d.isoformat(), d.isoformat()
    if gran == "week":
        d = _dt.strptime(day, "%Y-%m-%d").date()
        mon = d - _td(days=d.weekday() + 7)
        sun = mon + _td(days=6)
        return mon.isoformat(), sun.isoformat(), f"{mon.isoformat()} – {sun.isoformat()}"
    y, m = int(period_month[:4]), int(period_month[5:7])
    pm = f"{y-1:04d}-12" if m == 1 else f"{y:04d}-{m-1:02d}"
    return pm + "-01", pm + "-31", pm


def _coupon_series_map(conn, d_from, d_to, branch_id):
    """series_id → discount total for a date range (+ optional branch)."""
    w = ("WHERE cd.day BETWEEN ? AND ? AND COALESCE(b.branch_type,'master')='master' AND b.active=1")
    p = [d_from, d_to]
    if branch_id:
        w += " AND cd.branch_id=?"; p.append(branch_id)
    out = {}
    for r in conn.execute(f"""SELECT cd.coupon_series_id sid, SUM(cd.discount_amount) amt
                              FROM pos_coupon_day cd JOIN branches b ON b.branch_id=cd.branch_id
                              {w} GROUP BY cd.coupon_series_id""", p):
        out[r["sid"]] = r["amt"] or 0
    return out


@app.get("/api/coupon-series")
def coupon_series(period_month: str = "", gran: str = "month",
                  day: str | None = None, branch_id: int | None = None,
                  series_id: int | None = None):
    """COUPON detail by coupon series, at day / week / month grain (from
    pos_coupon_day). Floored to RECON_START_MONTH; master only. Pass series_id to
    drill one series (by_branch = that series per branch)."""
    if gran not in ("day", "week", "month"):
        raise HTTPException(400, "gran must be day, week or month")
    if gran == "month" and not period_month:
        raise HTTPException(400, "period_month required for gran=month")
    d_from, d_to, label = _coupon_range(gran, period_month, day)
    empty = {"gran": gran, "label": label, "range": [d_from, d_to],
             "series": [], "by_branch": [], "total": 0.0, "total_used": 0}
    if d_to < RECON_START_DAY:
        return empty
    if d_from < RECON_START_DAY:
        d_from = RECON_START_DAY
    with get_db() as conn:
        params = [d_from, d_to]
        where = ("WHERE cd.day BETWEEN ? AND ? "
                 "AND COALESCE(b.branch_type,'master')='master' AND b.active=1")
        if branch_id:
            where += " AND cd.branch_id=?"; params.append(branch_id)
        if series_id is not None:
            where += " AND cd.coupon_series_id=?"; params.append(series_id)
        try:
            rows = [dict(r) for r in conn.execute(f"""
                SELECT cd.coupon_series_id AS series_id,
                       SUM(cd.used_count)      AS used_count,
                       SUM(cd.discount_amount) AS discount_amount
                FROM pos_coupon_day cd JOIN branches b ON b.branch_id = cd.branch_id
                {where}
                GROUP BY cd.coupon_series_id
                ORDER BY discount_amount DESC
            """, params)]
        except sqlite3.OperationalError:
            return {**empty, "hint": "pos_coupon_day not found — run sync_mssql.py"}
        total = round(sum(r["discount_amount"] or 0 for r in rows), 2)
        total_used = int(sum(r["used_count"] or 0 for r in rows))
        name_map = {}
        try:
            for nr in conn.execute("SELECT coupon_series_id, name, code FROM coupon_series"):
                name_map[nr["coupon_series_id"]] = (nr["name"] or "", nr["code"] or "")
        except sqlite3.OperationalError:
            pass
        # Previous comparable period → MoM per series
        pf, pt, plabel = _coupon_prev(gran, period_month, day)
        if pf < RECON_START_DAY:
            pf = RECON_START_DAY
        prev_map = _coupon_series_map(conn, pf, pt, branch_id) if pt >= RECON_START_DAY else {}
        for r in rows:
            r["discount_amount"] = round(r["discount_amount"] or 0, 2)
            nm = name_map.get(r["series_id"], ("", ""))
            r["series_name"], r["series_code"] = nm[0], nm[1]
            prev = prev_map.get(r["series_id"], 0)
            r["prev_amount"] = round(prev, 2)
            r["mom_pct"] = round((r["discount_amount"] - prev) / prev * 100, 1) if prev else None
        by_branch = []
        if not branch_id:
            by_branch = [dict(r) for r in conn.execute(f"""
                SELECT b.branch_id, b.name AS branch_name,
                       SUM(cd.used_count)      AS used_count,
                       SUM(cd.discount_amount) AS discount_amount
                FROM pos_coupon_day cd JOIN branches b ON b.branch_id = cd.branch_id
                {where}
                GROUP BY b.branch_id
                ORDER BY discount_amount DESC
            """, params)]
            for r in by_branch:
                r["discount_amount"] = round(r["discount_amount"] or 0, 2)
        # Drill: per-day breakdown for a single series
        by_day = []
        if series_id is not None:
            by_day = [dict(r) for r in conn.execute(f"""
                SELECT cd.day,
                       SUM(cd.used_count)      AS used_count,
                       SUM(cd.discount_amount) AS discount_amount
                FROM pos_coupon_day cd JOIN branches b ON b.branch_id = cd.branch_id
                {where}
                GROUP BY cd.day
                ORDER BY cd.day
            """, params)]
            for r in by_day:
                r["discount_amount"] = round(r["discount_amount"] or 0, 2)
        return {"gran": gran, "label": label, "range": [d_from, d_to],
                "prev_label": plabel, "prev_range": [pf, pt],
                "period_month": period_month, "branch_id": branch_id,
                "series": rows, "by_branch": by_branch, "by_day": by_day,
                "total": total, "total_used": total_used,
                "prev_total": round(sum(prev_map.values()), 2)}


@app.get("/api/daily-board")
def daily_board_endpoint(day: str):
    if not _DAY_RE.match(day or ""):
        raise HTTPException(400, "day must be in 'YYYY-MM-DD' format.")
    with get_db() as conn:
        return daily_board(conn, day)


# ── Daily reconciliation report (print-to-PDF evidence) ───────────────────────
_STATUS_TH = {"matched": "ตรงกับเอกสาร", "variance": "มีส่วนต่าง", "pos_only": "รอเอกสาร",
              "ext_only": "เกินจากเอกสาร", "no_source": "ไม่มีแหล่งอ้างอิง", "manual": "กระทบมือ"}

def _m(v):
    return f"{(v or 0):,.2f}"

def _daily_report_html(conn, day: str, branch_id: int | None) -> str:
    board = daily_board(conn, day)
    cmeta = board.get("column_meta", {})
    rows = board["branches"]
    if branch_id:
        rows = [r for r in rows if r["branch_id"] == branch_id]
    keep = {r["branch_id"] for r in rows}
    bf = " AND branch_id=?" if branch_id else ""
    p = [day, branch_id] if branch_id else [day]

    def grouped(sql, params):
        out = {}
        for r in conn.execute(sql, params):
            out.setdefault(r["branch_id"], []).append(r)
        return out
    cc   = grouped(f"SELECT * FROM recon_cc_lines WHERE settlement_date=?{bf}", p)
    grab = grouped(f"SELECT * FROM recon_grab_lines WHERE order_date=? AND order_status='Completed'{bf}", p)
    lnmn = grouped(f"SELECT * FROM recon_lineman_lines WHERE order_date=? AND order_status='Completed'{bf}", p)
    bank = grouped(f"SELECT * FROM recon_bank_lines WHERE txn_date=? AND credit>0{bf}", p)
    # manual matches for the month, keep those that touch this day
    mm = {}
    mp = [day[:7], branch_id] if branch_id else [day[:7]]
    for r in conn.execute(f"SELECT * FROM recon_manual_match WHERE period_month=?{bf}", mp):
        try:
            days = json.loads(r["pos_days"] or "[]")
        except Exception:
            days = []
        amt = sum(float(x.get("amount") or 0) for x in days if x.get("day") == day)
        if amt:
            mm.setdefault(r["branch_id"], []).append((r, amt))

    css = """<style>
      @page{size:A4;margin:14mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12px;margin:0;padding:24px;background:#f8fafc}
      .sheet{background:#fff;max-width:900px;margin:0 auto 18px;padding:26px 30px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:19px;margin:0 0 2px} h2{font-size:15px;margin:22px 0 8px;border-left:4px solid #ea580c;padding-left:8px}
      h3{font-size:13px;margin:14px 0 6px;color:#475569}
      .muted{color:#64748b;font-size:11px} .right{text-align:right} .mono{font-variant-numeric:tabular-nums;font-family:ui-monospace,monospace}
      table{width:100%;border-collapse:collapse;margin:6px 0 10px;font-size:11.5px}
      th,td{padding:5px 7px;border-bottom:1px solid #eef2f7;text-align:left}
      th{color:#64748b;font-size:10px;text-transform:uppercase;border-bottom:1.5px solid #e2e8f0}
      tfoot td{border-top:1.5px solid #cbd5e1;font-weight:700}
      .kpis{display:flex;gap:10px;flex-wrap:wrap;margin:10px 0}
      .kpi{flex:1;min-width:120px;border:1px solid #e5e7eb;border-radius:8px;padding:8px 10px}
      .kpi .l{font-size:10px;color:#64748b} .kpi .v{font-size:17px;font-weight:700}
      .pos{color:#059669} .neg{color:#dc2626} .badge{padding:1px 7px;border-radius:99px;font-size:10px;font-weight:600}
      .b-ok{background:#dcfce7;color:#166534}.b-warn{background:#fef9c3;color:#854d0e}.b-err{background:#fee2e2;color:#b91c1c}.b-mut{background:#f1f5f9;color:#475569}
      .toolbar{max-width:900px;margin:0 auto 14px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:8px;padding:9px 18px;font-size:13px;cursor:pointer}
      .note{background:#fff7ed;border-left:3px solid #f59e0b;padding:6px 10px;border-radius:4px;margin:4px 0}
      .sign{display:flex;gap:40px;margin-top:34px} .sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;font-size:11px;color:#475569}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;margin:0;max-width:none;border-radius:0;page-break-after:always}}
    </style>"""

    def status_badge(s):
        cls = {"matched": "b-ok", "manual": "b-ok", "variance": "b-err", "ext_only": "b-warn",
               "pos_only": "b-warn", "no_source": "b-mut"}.get(s, "b-mut")
        return f'<span class="badge {cls}">{_STATUS_TH.get(s, s)}</span>'

    scope = "ทุกสาขา" if not branch_id else (rows[0]["branch_name"] if rows else "—")
    parts = [f"<!doctype html><html lang='th'><head><meta charset='utf-8'>",
             "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
             css, "</head><body>",
             "<div class='toolbar'><button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึกเป็น PDF</button></div>"]

    if not rows:
        parts.append(f"<div class='sheet'><h1>รายงานกระทบยอดรายวัน</h1><p class='muted'>{day} · {scope}</p>"
                     "<p>ไม่มีข้อมูลการขายสำหรับวันนี้</p></div>")
        parts.append("</body></html>")
        return "".join(parts)

    for b in rows:
        bid = b["branch_id"]
        parts.append("<div class='sheet'>")
        parts.append(f"<h1>รายงานกระทบยอดรายวัน — หลักฐานการกระทบยอด</h1>"
                     f"<p class='muted'>สาขา <b>{b['branch_name']}</b> · วันที่ <b>{day}</b> · "
                     f"ออกรายงานเมื่อ {_dt.now().strftime('%Y-%m-%d %H:%M')}</p>")
        var = b["variance"]; vc = "pos" if abs(var) < 1 else "neg"
        parts.append("<div class='kpis'>"
                     f"<div class='kpi'><div class='l'>ยอด POS (ที่กระทบได้)</div><div class='v mono'>{_m(b['pos_total'])}</div></div>"
                     f"<div class='kpi'><div class='l'>ยอดจากเอกสาร</div><div class='v mono'>{_m(b['ext_total'])}</div></div>"
                     f"<div class='kpi'><div class='l'>ส่วนต่าง</div><div class='v mono {vc}'>{_m(var)}</div></div>"
                     f"<div class='kpi'><div class='l'>รายการมีปัญหา</div><div class='v'>{b['issues']}</div></div>"
                     "</div>")

        # 1) Category reconciliation (POS vs documents)
        parts.append("<h2>① กระทบยอดตามช่องทาง (POS เทียบเอกสาร)</h2>"
                     "<table><thead><tr><th>ช่องทาง</th><th class='right'>POS</th><th class='right'>เอกสาร</th>"
                     "<th class='right'>กระทบมือ</th><th class='right'>ส่วนต่าง</th><th>สถานะ</th></tr></thead><tbody>")
        for ck, cell in b["cells"].items():
            lbl = (cmeta.get(ck) or {}).get("label", ck)
            v = cell["variance"]
            parts.append(f"<tr><td>{lbl}</td><td class='right mono'>{_m(cell['pos'])}</td>"
                         f"<td class='right mono'>{_m(cell['ext'])}</td><td class='right mono'>{_m(cell.get('manual'))}</td>"
                         f"<td class='right mono {'pos' if abs(v)<1 else 'neg'}'>{_m(v)}</td><td>{status_badge(cell['status'])}</td></tr>")
        parts.append(f"<tr><td colspan='6' class='muted' style='border:0'></td></tr></tbody><tfoot><tr><td>รวม</td>"
                     f"<td class='right mono'>{_m(b['pos_total'])}</td><td class='right mono'>{_m(b['ext_total'])}</td>"
                     f"<td></td><td class='right mono {vc}'>{_m(var)}</td><td></td></tr></tfoot></table>")

        # 2) Fees & VAT — line level
        cc_l, gr_l, ln_l = cc.get(bid, []), grab.get(bid, []), lnmn.get(bid, [])
        tot = {"gross": 0.0, "fee": 0.0, "vat": 0.0, "promo": 0.0, "refund": 0.0, "net": 0.0}
        feerows = []
        for r in cc_l:
            fee, vat = _eff_fee_vat(r, "mdr_fee", "vat_amount")
            gross = float(r["gross_amount"] or 0); refund = float(r["refund_amount"] or 0)
            net = gross - refund - fee - vat
            tot["gross"] += gross; tot["fee"] += fee; tot["vat"] += vat; tot["refund"] += refund; tot["net"] += net
            feerows.append(("CREDIT " + (r["acquirer"] or ""), gross, refund, fee, vat, 0.0, net,
                            "MDR" + (" (แก้มือ)" if r["fee_manual"] is not None else "")))
        for r in gr_l:
            fee, vat = _eff_fee_vat(r, "commission_amt", "vat_amount")
            gross = float(r["food_amount"] or 0); promo = float(r["promo_deduct"] or 0)
            net = gross - fee - vat - promo
            tot["gross"] += gross; tot["fee"] += fee; tot["vat"] += vat; tot["promo"] += promo; tot["net"] += net
            feerows.append(("GRAB", gross, 0.0, fee, vat, promo, net, "ค่าคอมมิชชั่น"))
        for r in ln_l:
            fee, vat = _eff_fee_vat(r, "commission_amt", "vat_on_commission")
            gross = float(r["food_amount"] or 0); promo = float(r["promo_deduct"] or 0)
            net = gross - fee - vat - promo
            tot["gross"] += gross; tot["fee"] += fee; tot["vat"] += vat; tot["promo"] += promo; tot["net"] += net
            feerows.append(("LINEMAN", gross, 0.0, fee, vat, promo, net, "ค่าคอมมิชชั่น + VAT"))
        if feerows:
            parts.append("<h2>② ค่าธรรมเนียม &amp; VAT (รายบรรทัดจากเอกสาร)</h2>"
                         "<table><thead><tr><th>ช่องทาง</th><th class='right'>ยอดเต็ม</th><th class='right'>คืนเงิน</th>"
                         "<th class='right'>ค่าธรรมเนียม</th><th class='right'>VAT</th><th class='right'>โปรโมชัน</th>"
                         "<th class='right'>สุทธิ</th><th>หมายเหตุ</th></tr></thead><tbody>")
            for nm, g, rf, fe, va, pr, nt, note in feerows:
                parts.append(f"<tr><td>{nm}</td><td class='right mono'>{_m(g)}</td><td class='right mono'>{_m(rf)}</td>"
                             f"<td class='right mono neg'>{_m(fe)}</td><td class='right mono neg'>{_m(va)}</td>"
                             f"<td class='right mono neg'>{_m(pr)}</td><td class='right mono'>{_m(nt)}</td><td class='muted'>{note}</td></tr>")
            parts.append(f"</tbody><tfoot><tr><td>รวม</td><td class='right mono'>{_m(tot['gross'])}</td>"
                         f"<td class='right mono'>{_m(tot['refund'])}</td><td class='right mono neg'>{_m(tot['fee'])}</td>"
                         f"<td class='right mono neg'>{_m(tot['vat'])}</td><td class='right mono neg'>{_m(tot['promo'])}</td>"
                         f"<td class='right mono'>{_m(tot['net'])}</td><td></td></tr></tfoot></table>"
                         f"<p class='muted'>สุทธิ = ยอดเต็ม − คืนเงิน − ค่าธรรมเนียม − VAT − โปรโมชัน · "
                         f"ค่าธรรมเนียมรวม {_m(tot['fee'])} · VAT รวม {_m(tot['vat'])}</p>")

        # 3) Manual matches — wrong-category reclassification with reason
        ms = mm.get(bid, [])
        if ms:
            parts.append("<h2>③ การกระทบมือ (สาขาบันทึกผิดประเภท / จับคู่เอง)</h2>")
            parts.append("<table><thead><tr><th>ช่องทาง</th><th>ธนาคาร</th><th class='right'>ยอดที่กระทบ (วันนี้)</th>"
                         "<th>เหตุผล / หมายเหตุ</th><th>โดย</th></tr></thead><tbody>")
            for r, amt in ms:
                parts.append(f"<tr><td>{r['category']}</td><td>{r['bank'] or '—'}</td>"
                             f"<td class='right mono'>{_m(amt)}</td><td>{(r['note'] or '—')}</td>"
                             f"<td class='muted'>{(r['matched_by'] or '—')}<br>{(r['matched_at'] or '')[:16]}</td></tr>")
            parts.append("</tbody></table>")

        # 4) Document evidence — the actual lines used
        parts.append("<h2>④ เอกสารที่ใช้อ้างอิง (หลักฐาน)</h2>")
        bk = bank.get(bid, [])
        if bk:
            parts.append("<h3>รายการเดินบัญชีธนาคาร</h3><table><thead><tr><th>ช่องทาง</th><th>บัญชี</th>"
                         "<th class='right'>เงินเข้า</th></tr></thead><tbody>")
            for r in bk:
                parts.append(f"<tr><td>{r['channel']}</td><td class='mono'>{r['account_no'] or '—'}</td>"
                             f"<td class='right mono'>{_m(r['credit'])}</td></tr>")
            parts.append("</tbody></table>")
        if cc_l:
            parts.append("<h3>สรุปยอดบัตรเครดิต (settlement)</h3><table><thead><tr><th>ผู้รับบัตร</th>"
                         "<th>MID/Terminal</th><th class='right'>ยอดเต็ม</th><th class='right'>สุทธิ</th></tr></thead><tbody>")
            for r in cc_l:
                parts.append(f"<tr><td>{r['acquirer'] or '—'}</td><td class='mono'>{(r['merchant_id'] or '—')}/{(r['terminal_id'] or '—')}</td>"
                             f"<td class='right mono'>{_m(r['gross_amount'])}</td><td class='right mono'>{_m(r['net_amount'])}</td></tr>")
            parts.append("</tbody></table>")
        if gr_l or ln_l:
            parts.append("<h3>รายงาน Grab / Line Man (ออเดอร์สำเร็จ)</h3><table><thead><tr><th>ผู้ให้บริการ</th>"
                         "<th class='right'>จำนวนออเดอร์</th><th class='right'>ยอดอาหาร</th></tr></thead><tbody>")
            if gr_l: parts.append(f"<tr><td>Grab</td><td class='right'>{len(gr_l)}</td><td class='right mono'>{_m(sum(float(r['food_amount'] or 0) for r in gr_l))}</td></tr>")
            if ln_l: parts.append(f"<tr><td>Line Man</td><td class='right'>{len(ln_l)}</td><td class='right mono'>{_m(sum(float(r['food_amount'] or 0) for r in ln_l))}</td></tr>")
            parts.append("</tbody></table>")
        if not (bk or cc_l or gr_l or ln_l):
            parts.append("<p class='note'>ยังไม่มีเอกสารอัปโหลดสำหรับวันนี้ — ยอดที่แสดงเป็นฝั่ง POS เท่านั้น</p>")

        parts.append("<div class='sign'><div>ผู้จัดทำ (กระทบยอด)</div><div>ผู้ตรวจสอบ</div><div>ผู้อนุมัติ</div></div>")
        parts.append("</div>")  # sheet

    parts.append("</body></html>")
    return "".join(parts)


@app.get("/api/daily-report")
def daily_report(day: str, branch_id: int | None = None):
    """Print-to-PDF daily reconciliation evidence: POS vs documents, manual
    matches (wrong-category reclass + reason), all fees, and VAT — per the
    selected day. Omit branch_id for an all-branches pack (one page each)."""
    if not _DAY_RE.match(day or ""):
        raise HTTPException(400, "day must be in 'YYYY-MM-DD' format.")
    with get_db() as conn:
        html = _daily_report_html(conn, day, branch_id)
    return Response(content=html, media_type="text/html; charset=utf-8")


# ── Monthly matching audit (รายงานการจับคู่รายวันทั้งเดือน) ────────────────────
# One printable sheet per branch: every day × channel with POS vs document, the
# variance, and the match status — so the owner/auditor can confirm the whole month
# reconciled correctly in one pass. Deep line-level evidence stays in /api/daily-report.

_MATCH_CH_TH = {"QR": "QR พร้อมเพย์", "CASH": "เงินสด", "TRANSFER": "เงินโอน",
                "CREDIT": "บัตรเครดิต", "GRAB": "Grab", "LINEMAN": "Line Man", "COUPON": "คูปอง",
                "DISCOUNT": "ส่วนลด (คูปอง)"}
# status → (Thai label, css class, bucket)  bucket: ok | issue | wait | none
_MATCH_ST = {
    "matched":       ("✓ ตรง",            "m-ok",   "ok"),
    "manual":        ("✓ จับคู่มือ",       "m-ok",   "ok"),
    "variance":      ("✗ ต่าง",           "m-err",  "issue"),
    "missing":       ("✗ ขาดเอกสาร",      "m-err",  "issue"),
    "ext_only":      ("เกินเอกสาร",        "m-warn", "issue"),
    "pending":       ("รอเอกสารตามรอบ",    "m-wait", "wait"),
    "unmapped_bank": ("ยังไม่ตั้งธนาคาร",  "m-info", "issue"),
    "no_source":     ("ไม่มีแหล่งอ้างอิง",  "m-mut",  "none"),
}


def _account_map(conn):
    """slot → (account_code, account_name) from recon_account_map; {} if unset.
    Reports render blank cells for slots with no code yet (owner fills in later)."""
    try:
        return {r["slot"]: (r["account_code"] or "", r["account_name"] or "")
                for r in conn.execute(
                    "SELECT slot, account_code, account_name FROM recon_account_map")}
    except sqlite3.OperationalError:
        return {}


# Channels carrying a platform fee/VAT + a separate discount line (gross from the
# document). cols: (channel, table, date_col, gross, fee, vat, discount, extra-where)
_FEE_CH_SPEC = [
    ("CREDIT",  "recon_cc_lines",      "settlement_date", "gross_amount", "mdr_fee",        "vat_amount",        "refund_amount", ""),
    ("GRAB",    "recon_grab_lines",    "order_date",      "food_amount",  "commission_amt", "vat_amount",        "promo_deduct",  " AND order_status='Completed'"),
    ("LINEMAN", "recon_lineman_lines", "order_date",      "food_amount",  "commission_amt", "vat_on_commission", "promo_deduct",  " AND order_status='Completed'"),
]
_FINANCE_CH_ORDER = {"CASH": 0, "QR": 1, "TRANSFER": 2, "CREDIT": 3, "GRAB": 4, "LINEMAN": 5,
                     "COUPON": 6, "DISCOUNT": 7}


def _daily_finance_data(conn, branch_id: int, period_month: str, daily=None) -> dict:
    """Per-DAY accounting breakdown as structured data: each day → its channels, each
    recorded at GROSS with discount/fee/VAT split out (net = gross − discount − fee −
    VAT) + account code/name from recon_account_map (blank until filled in). Single
    source for the table inside the monthly report (_daily_finance_html), the in-app
    daybook cards, and the printable per-day card report ('1 วัน = 1 ใบ').
    `daily` lets a caller pass a precomputed daily_reconciliation result so the dashboard
    (which also needs the POS↔money side) doesn't recompute it."""
    acct = _account_map(conn)

    # (day, channel) -> {gross, discount, fee, vat} from the fee-bearing documents
    agg: dict = {}
    def cell(day, ch):
        return agg.setdefault((day, ch), {"gross": 0.0, "discount": 0.0, "fee": 0.0, "vat": 0.0})
    for ch, tbl, dcol, gcol, fcol, vcol, discol, extra in _FEE_CH_SPEC:
        try:
            lines = conn.execute(
                f"SELECT * FROM {tbl} WHERE branch_id=? AND substr({dcol},1,7)=?{extra}",
                (branch_id, period_month)).fetchall()
        except sqlite3.OperationalError:
            continue
        for r in lines:
            day = (r[dcol] or "")[:10]
            if not day:
                continue
            fee, vat = _eff_fee_vat(r, fcol, vcol)
            a = cell(day, ch)
            a["gross"] += float(r[gcol] or 0); a["fee"] += fee; a["vat"] += vat
            a["discount"] += float(r[discol] or 0)

    # POS-tender channels (cash/QR/transfer/coupon): gross = POS that day, no fee/discount.
    try:
        d = daily if daily is not None else daily_reconciliation(conn, branch_id, period_month)
    except Exception:
        d = {"days": []}
    for day in d.get("days", []):
        for l in day["lanes"]:
            if l["category"] in ("CASH", "QR", "TRANSFER", "COUPON") and abs(l["pos"]) > 0.005:
                cell(day["day"], l["category"])["gross"] += l["pos"]

    # Split the COUPON tender: KTC/BAY loyalty points settle as REAL money via the
    # card issuer → keep in the COUPON line; everything else (ผู้สูงอายุ/เด็ก/อื่นๆ
    # discount coupons) is NOT money — book the full price as revenue and the discount
    # portion as a deduction (net 0). So revenue stays gross and net = cash actually in.
    points_by_day = point_pos_total_daily(conn, branch_id, period_month)
    for day in [k[0] for k in list(agg) if k[1] == "COUPON"]:
        cp = agg[(day, "COUPON")]
        coupon_total = cp["gross"]
        pts = min(max(points_by_day.get(day, 0.0), 0.0), coupon_total)
        disc = round(coupon_total - pts, 2)
        cp["gross"] = round(pts, 2)
        if disc > 0.005:
            dd = cell(day, "DISCOUNT")
            dd["gross"] += disc; dd["discount"] += disc
        if pts <= 0.005:                       # all discount, no real coupon money — drop empty lane
            agg.pop((day, "COUPON"), None)

    # Bill-level coupon discounts (ส่วนลดเด็ก / ผู้สูงอายุ / วันเกิด …) reduce the bill;
    # the customer then pays the NET via cash/QR/card. They are NOT a payment tender, so
    # without this they're invisible here and ส่วนลด shows 0. Book revenue at GROSS and
    # the discount on its own line (net 0) — same shape as the coupon-points remainder.
    for day, cdisc in coupon_discount_total_daily(conn, branch_id, period_month).items():
        if cdisc > 0.005:
            dd = cell(day, "DISCOUNT")
            dd["gross"] += cdisc
            dd["discount"] += cdisc

    def dom(dstr):
        try:
            return int(dstr[8:10])
        except Exception:
            return dstr

    days = []
    tot = {"gross": 0.0, "discount": 0.0, "fee": 0.0, "vat": 0.0, "net": 0.0}
    for day in sorted({k[0] for k in agg}):
        chans = sorted({c for (dy, c) in agg if dy == day},
                       key=lambda c: _FINANCE_CH_ORDER.get(c, 9))
        rows, dt = [], {"gross": 0.0, "discount": 0.0, "fee": 0.0, "vat": 0.0, "net": 0.0}
        for ch in chans:
            a = agg[(day, ch)]
            net = a["gross"] - a["discount"] - a["fee"] - a["vat"]
            code, name = acct.get(ch, ("", ""))
            rows.append({"ch": ch, "ch_th": _MATCH_CH_TH.get(ch, ch), "code": code, "name": name,
                         "gross": round(a["gross"], 2), "discount": round(a["discount"], 2),
                         "fee": round(a["fee"], 2), "vat": round(a["vat"], 2), "net": round(net, 2)})
            dt["gross"] += a["gross"]; dt["discount"] += a["discount"]; dt["fee"] += a["fee"]
            dt["vat"] += a["vat"]; dt["net"] += net
        days.append({"day": day, "dom": dom(day), "channels": rows,
                     "totals": {k: round(v, 2) for k, v in dt.items()}})
        for k in tot:
            tot[k] += dt[k]
    return {"days": days, "totals": {k: round(v, 2) for k, v in tot.items()}}


def _daily_finance_html(conn, branch_id: int, period_month: str) -> str:
    """Per-day accounting breakdown rendered as the table inside the monthly report.
    Same numbers as the daybook cards — see _daily_finance_data. The in-app section
    (📅 สรุปบัญชีรายวัน รายใบ) renders one card per day from the same data."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.2f}"
    data = _daily_finance_data(conn, branch_id, period_month)
    if not data["days"]:
        return ""

    P = ["<h2 style='border-left-color:#0f6e56'>📅 สรุปบัญชีรายวัน (ยอดเต็ม / ส่วนลดแยกบรรทัด)</h2>",
         "<table><thead><tr><th class='l'>วันที่</th><th class='l'>รหัสบัญชี</th><th class='l'>ชื่อบัญชี</th>"
         "<th class='l'>ช่องทาง</th><th>ยอดเต็ม</th><th>ส่วนลด</th><th>ค่าธรรมเนียม</th>"
         "<th>VAT</th><th>สุทธิ</th></tr></thead><tbody>"]
    for day in data["days"]:
        for i, a in enumerate(day["channels"]):
            daycell = f"<span class='day'>{day['dom']}</span>" if i == 0 else ""
            P.append(
                f"<tr class='{'daystart' if i == 0 else ''}'><td class='l'>{daycell}</td>"
                f"<td class='l mono'>{esc(a['code']) or '—'}</td><td class='l'>{esc(a['name']) or '—'}</td>"
                f"<td class='l'>{esc(a['ch_th'])}</td>"
                f"<td class='mono'>{m(a['gross'])}</td><td class='mono'>{m(a['discount'])}</td>"
                f"<td class='mono'>{m(a['fee'])}</td><td class='mono'>{m(a['vat'])}</td>"
                f"<td class='mono'>{m(a['net'])}</td></tr>")
    t = data["totals"]
    P.append(
        f"</tbody><tfoot><tr><td class='l' colspan='4'>รวมทั้งเดือน</td>"
        f"<td class='mono'>{m(t['gross'])}</td><td class='mono'>{m(t['discount'])}</td><td class='mono'>{m(t['fee'])}</td>"
        f"<td class='mono'>{m(t['vat'])}</td><td class='mono'>{m(t['net'])}</td></tr></tfoot></table>"
        "<p class='legend'>ยอดเต็ม = ยอดขายก่อนหักส่วนลด/ค่าธรรมเนียม · ส่วนลด = โปรโมชัน/คืนเงิน (บันทึกแยกบรรทัด) · "
        "สุทธิ = ยอดเต็ม − ส่วนลด − ค่าธรรมเนียม − VAT · รหัส/ชื่อบัญชีเติมภายหลัง</p>")
    return "".join(P)


def _monthly_finance_html(conn, branch_id: int, period_month: str) -> str:
    """Monthly receipts by channel (gross → discount → fee → VAT → net) + account
    code/name + a draft journal, so the books tie to what came in. Amounts are GROSS;
    discount is recorded separately. Fee/VAT use the effective (override-aware) figures."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.2f}"
    acct = _account_map(conn)

    rows = []   # (channel, gross, discount, fee, vat, net)
    for ch, tbl, dcol, gcol, fcol, vcol, discol, extra in _FEE_CH_SPEC:
        try:
            lines = conn.execute(
                f"SELECT * FROM {tbl} WHERE branch_id=? AND substr({dcol},1,7)=?{extra}",
                (branch_id, period_month)).fetchall()
        except sqlite3.OperationalError:
            continue
        g = disc = f = v = 0.0
        for r in lines:
            g += float(r[gcol] or 0)
            fee, vat = _eff_fee_vat(r, fcol, vcol)
            f += fee; v += vat; disc += float(r[discol] or 0)
        if g > 0.005:
            rows.append((ch, g, disc, f, v, g - disc - f - v))

    # no-fee channels (cash / QR / transfer): money received = POS, no fee/VAT.
    # COUPON splits: KTC/BAY points = real money (kept in COUPON); other discount
    # coupons (ผู้สูงอายุ/เด็ก/อื่นๆ) book the full price as revenue + a discount
    # deduction (net 0) so REVENUE stays gross and the journal's discount line carries them.
    try:
        rm = reconcile_month(conn, branch_id, period_month)
        pts_total = round(sum(point_pos_total_daily(conn, branch_id, period_month).values()), 2)
        for c in rm.get("channels", []):
            key = c["key"]; pos = float(c.get("pos", 0))
            if key in ("CASH", "QR", "TRANSFER") and pos > 0.005:
                rows.append((key, pos, 0.0, 0.0, 0.0, pos))
            elif key == "COUPON" and pos > 0.005:
                pts = min(max(pts_total, 0.0), pos); disc = round(pos - pts, 2)
                if pts > 0.005:
                    rows.append(("COUPON", round(pts, 2), 0.0, 0.0, 0.0, round(pts, 2)))
                if disc > 0.005:
                    rows.append(("DISCOUNT", disc, disc, 0.0, 0.0, 0.0))
    except Exception:
        pass

    # Bill-level coupon discounts (ส่วนลดเด็ก/ผู้สูงอายุ/วันเกิด …) from pos_coupon_day —
    # not a payment tender, so otherwise missing here. Gross up revenue + book the discount
    # (net 0), merged into the single ส่วนลด row so month and per-day journals agree.
    coupon_disc = round(sum(coupon_discount_total_daily(conn, branch_id, period_month).values()), 2)
    if coupon_disc > 0.005:
        for i, r in enumerate(rows):
            if r[0] == "DISCOUNT":
                rows[i] = ("DISCOUNT", r[1] + coupon_disc, r[2] + coupon_disc, 0.0, 0.0, 0.0)
                break
        else:
            rows.append(("DISCOUNT", coupon_disc, coupon_disc, 0.0, 0.0, 0.0))

    if not rows:
        return ""

    rows.sort(key=lambda r: _FINANCE_CH_ORDER.get(r[0], 9))
    G = sum(r[1] for r in rows); D = sum(r[2] for r in rows)
    F = sum(r[3] for r in rows); V = sum(r[4] for r in rows); N = sum(r[5] for r in rows)
    out_vat = round(G * 7 / 107, 2)            # output VAT estimate if prices are VAT-inclusive

    def acc_cells(slot):
        code, name = acct.get(slot, ("", ""))
        return f"<td class='l mono'>{esc(code) or '—'}</td><td class='l'>{esc(name) or '—'}</td>"

    P = ["<div style='margin-top:14px;border-top:2px solid #0f6e56;padding-top:6px'>"
         "<h2 style='border-left-color:#0f6e56'>💰 สรุปการรับเงินทั้งเดือน · ส่วนลด · ค่าธรรมเนียม · VAT</h2>",
         "<table><thead><tr><th class='l'>รหัสบัญชี</th><th class='l'>ชื่อบัญชี</th><th class='l'>ช่องทาง</th>"
         "<th>ยอดรับเต็ม</th><th>ส่วนลด</th><th>ค่าธรรมเนียม</th>"
         "<th>VAT (ค่าธรรมเนียม)</th><th>สุทธิเข้าบัญชี</th></tr></thead><tbody>"]
    for ch, g, disc, f, v, n in rows:
        P.append(f"<tr>{acc_cells(ch)}<td class='l'>{esc(_MATCH_CH_TH.get(ch, ch))}</td>"
                 f"<td class='mono'>{m(g)}</td><td class='mono'>{m(disc)}</td><td class='mono'>{m(f)}</td>"
                 f"<td class='mono'>{m(v)}</td><td class='mono'>{m(n)}</td></tr>")
    P.append(f"</tbody><tfoot><tr><td class='l' colspan='3'>รวม</td><td class='mono'>{m(G)}</td>"
             f"<td class='mono'>{m(D)}</td><td class='mono'>{m(F)}</td><td class='mono'>{m(V)}</td>"
             f"<td class='mono'>{m(N)}</td></tr></tfoot></table>")

    # draft journal entry (money-received basis) — gross revenue, discount booked apart
    P.append("<h2 style='border-left-color:#0f6e56'>📒 ร่างการบันทึกบัญชี (รับชำระเงิน)</h2>")
    P.append("<table><thead><tr><th class='l'>รหัสบัญชี</th><th class='l'>บัญชี</th>"
             "<th>เดบิต (Dr)</th><th>เครดิต (Cr)</th></tr></thead><tbody>")
    je = [("BANK", "เงินสด / ธนาคาร (รับสุทธิ)", N, 0), ("DISCOUNT", "ส่วนลด / โปรโมชัน", D, 0),
          ("FEE", "ค่าธรรมเนียมรับชำระเงิน", F, 0), ("VAT", "ภาษีซื้อ – VAT ค่าธรรมเนียม", V, 0),
          ("REVENUE", "รายได้จากการขาย (ยอดรับเต็ม)", 0, G)]
    for slot, acc, dr, cr in je:
        code, _name = acct.get(slot, ("", ""))
        P.append(f"<tr><td class='l mono'>{esc(code) or '—'}</td><td class='l'>{esc(acc)}</td>"
                 f"<td class='mono'>{m(dr) if dr else ''}</td><td class='mono'>{m(cr) if cr else ''}</td></tr>")
    P.append(f"</tbody><tfoot><tr><td class='l' colspan='2'>รวม</td><td class='mono'>{m(N + D + F + V)}</td>"
             f"<td class='mono'>{m(G)}</td></tr></tfoot></table>")
    P.append("<div class='legend'>หมายเหตุ: VAT ในตารางคือ <b>ภาษีซื้อ</b>จากค่าธรรมเนียมรับชำระ (บัตร/เดลิเวอรี) · "
             f"หากราคาขายรวม VAT แล้ว <b>ภาษีขาย (output) โดยประมาณ = ฿{m(out_vat)}</b> (7/107 ของยอดรับเต็ม) "
             "ซึ่งฝ่ายบัญชีแยกออกจากรายได้ตามเกณฑ์ ex-VAT · ตัวเลขนี้เป็นร่างตามข้อมูลกระทบยอด ฝ่ายบัญชีปรับตามผังบัญชีจริง</div>")
    P.append("</div>")
    return "".join(P)


def _day_finance_html(conn, branch_id: int, day: str) -> str:
    """Receipts summary + draft journal entry (Dr/Cr) for ONE calendar day, so the day's
    accounting can be posted straight into the daily books. Same gross → discount → fee →
    VAT → net split (incl. the KTC/BAY coupon-points vs discount-coupon handling) as the
    monthly journal — see _monthly_finance_html — but scoped to a single day via
    _daily_finance_data so screen, monthly report and paper all tie to the baht."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.2f}"
    acct = _account_map(conn)
    data = _daily_finance_data(conn, branch_id, day[:7])
    dd = next((x for x in data["days"] if x["day"] == day), None)
    if not dd or not dd["channels"]:
        return ("<h2 style='border-left-color:#0f6e56'>📒 ร่างการบันทึกบัญชี (รับชำระเงิน) — "
                f"วันที่ {esc(day)}</h2><div class='empty'>ไม่มีรายการรับเงินในวันนี้</div>")
    chans = dd["channels"]; t = dd["totals"]
    G, D, F, V, N = t["gross"], t["discount"], t["fee"], t["vat"], t["net"]
    out_vat = round(G * 7 / 107, 2)

    def acc_cells(slot):
        code, name = acct.get(slot, ("", ""))
        return f"<td class='l mono'>{esc(code) or '—'}</td><td class='l'>{esc(name) or '—'}</td>"

    P = ["<div style='margin-top:14px;border-top:2px solid #0f6e56;padding-top:6px'>"
         f"<h2 style='border-left-color:#0f6e56'>💰 สรุปการรับเงินวันที่ {esc(day)} · ส่วนลด · ค่าธรรมเนียม · VAT</h2>",
         "<table><thead><tr><th class='l'>รหัสบัญชี</th><th class='l'>ชื่อบัญชี</th><th class='l'>ช่องทาง</th>"
         "<th>ยอดรับเต็ม</th><th>ส่วนลด</th><th>ค่าธรรมเนียม</th>"
         "<th>VAT (ค่าธรรมเนียม)</th><th>สุทธิเข้าบัญชี</th></tr></thead><tbody>"]
    for a in chans:
        P.append(f"<tr>{acc_cells(a['ch'])}<td class='l'>{esc(a['ch_th'])}</td>"
                 f"<td class='mono'>{m(a['gross'])}</td><td class='mono'>{m(a['discount'])}</td><td class='mono'>{m(a['fee'])}</td>"
                 f"<td class='mono'>{m(a['vat'])}</td><td class='mono'>{m(a['net'])}</td></tr>")
    P.append(f"</tbody><tfoot><tr><td class='l' colspan='3'>รวม</td><td class='mono'>{m(G)}</td>"
             f"<td class='mono'>{m(D)}</td><td class='mono'>{m(F)}</td><td class='mono'>{m(V)}</td>"
             f"<td class='mono'>{m(N)}</td></tr></tfoot></table>")

    # draft journal entry for THIS day (money-received basis) — gross revenue, discount apart
    P.append(f"<h2 style='border-left-color:#0f6e56'>📒 ร่างการบันทึกบัญชี (รับชำระเงิน) — วันที่ {esc(day)}</h2>")
    P.append("<table><thead><tr><th class='l'>รหัสบัญชี</th><th class='l'>บัญชี</th>"
             "<th>เดบิต (Dr)</th><th>เครดิต (Cr)</th></tr></thead><tbody>")
    je = [("BANK", "เงินสด / ธนาคาร (รับสุทธิ)", N, 0), ("DISCOUNT", "ส่วนลด / โปรโมชัน", D, 0),
          ("FEE", "ค่าธรรมเนียมรับชำระเงิน", F, 0), ("VAT", "ภาษีซื้อ – VAT ค่าธรรมเนียม", V, 0),
          ("REVENUE", "รายได้จากการขาย (ยอดรับเต็ม)", 0, G)]
    for slot, acc, dr, cr in je:
        code, _name = acct.get(slot, ("", ""))
        P.append(f"<tr><td class='l mono'>{esc(code) or '—'}</td><td class='l'>{esc(acc)}</td>"
                 f"<td class='mono'>{m(dr) if dr else ''}</td><td class='mono'>{m(cr) if cr else ''}</td></tr>")
    P.append(f"</tbody><tfoot><tr><td class='l' colspan='2'>รวม</td><td class='mono'>{m(N + D + F + V)}</td>"
             f"<td class='mono'>{m(G)}</td></tr></tfoot></table>")
    P.append("<div class='legend'>หมายเหตุ: VAT ในตารางคือ <b>ภาษีซื้อ</b>จากค่าธรรมเนียมรับชำระ (บัตร/เดลิเวอรี) · "
             f"หากราคาขายรวม VAT แล้ว <b>ภาษีขาย (output) โดยประมาณ = ฿{m(out_vat)}</b> (7/107 ของยอดรับเต็ม) "
             "ซึ่งฝ่ายบัญชีแยกออกจากรายได้ตามเกณฑ์ ex-VAT · ตัวเลขนี้เป็นร่างตามข้อมูลกระทบยอด ฝ่ายบัญชีปรับตามผังบัญชีจริง</div>")
    P.append("</div>")
    return "".join(P)


def _audit_detail_html(conn, branch_id: int, period_month: str) -> str:
    """Audit-grade evidence for ONE branch+month so an auditor can trace every baht:
    POS composition (codes), credit settlement by acquirer (gross/fee/VAT/net), every
    bank credit line (date/description/ref), Grab/LM payout, and manual matches."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.2f}"
    P = []

    # 1) POS composition — every payment code, its channel + acquirer/bank
    try:
        codes = pos_codes_scope(conn, branch_id, period_month=period_month)
    except sqlite3.OperationalError:
        codes = []
    if codes:
        P.append("<h2>① ที่มาของยอด POS (รายรหัสรับเงิน)</h2>")
        P.append("<table><thead><tr><th class='l'>รหัสรับเงิน</th><th class='l'>ช่องทาง</th>"
                 "<th class='l'>ธนาคาร/ผู้รับบัตร</th><th>จำนวนรายการ</th><th>ยอด (บาท)</th></tr></thead><tbody>")
        tot = 0.0
        for c in sorted(codes, key=lambda x: (x["category"] or "zz", -x["amount"])):
            tot += c["amount"]
            ch = _MATCH_CH_TH.get(c["category"], c["category"] or "—")
            bk = _BANK_LABEL.get(c["bank"], c["bank"] or "—")
            P.append(f"<tr><td class='l mono'>{esc(c['code'])}</td><td class='l'>{esc(ch)}</td>"
                     f"<td class='l'>{esc(bk)}</td><td class='mono'>{c['txn_count']:,}</td>"
                     f"<td class='mono'>{m(c['amount'])}</td></tr>")
        P.append(f"</tbody><tfoot><tr><td colspan='4' class='l'>รวม POS</td>"
                 f"<td class='mono'>{m(tot)}</td></tr></tfoot></table>")

    # 2) Credit settlement by acquirer (the document that backs the CREDIT lane)
    cc = conn.execute(
        "SELECT acquirer, COUNT(*) n, SUM(COALESCE(gross_amount,0)) g, SUM(COALESCE(mdr_fee,0)) f, "
        "SUM(COALESCE(vat_amount,0)) v, SUM(COALESCE(net_amount,0)) net "
        "FROM recon_cc_lines WHERE branch_id=? AND substr(settlement_date,1,7)=? GROUP BY acquirer",
        (branch_id, period_month)).fetchall()
    if cc:
        P.append("<h2>② สรุปยอดบัตรเครดิต ตามผู้รับบัตร (acquirer)</h2>")
        P.append("<table><thead><tr><th class='l'>ผู้รับบัตร</th><th>จำนวนแบตช์</th><th>ยอดเต็ม</th>"
                 "<th>ค่าธรรมเนียม</th><th>VAT</th><th>สุทธิ (เข้าบัญชี)</th></tr></thead><tbody>")
        for r in cc:
            P.append(f"<tr><td class='l'>{esc(_BANK_LABEL.get(r['acquirer'], r['acquirer'] or '—'))}</td>"
                     f"<td class='mono'>{r['n']:,}</td><td class='mono'>{m(r['g'])}</td>"
                     f"<td class='mono'>{m(r['f'])}</td><td class='mono'>{m(r['v'])}</td>"
                     f"<td class='mono'>{m(r['net'])}</td></tr>")
        P.append("</tbody></table>")

    # 3) Every bank credit line — the actual money received (grouped by channel)
    bl = conn.execute(
        "SELECT txn_date, description, channel, ref_no, account_no, COALESCE(credit,0) credit "
        "FROM recon_bank_lines WHERE branch_id=? AND substr(txn_date,1,7)=? AND credit>0 "
        "ORDER BY channel, txn_date", (branch_id, period_month)).fetchall()
    if bl:
        P.append("<h2>③ รายการเงินเข้าบัญชีธนาคาร (ทุกบรรทัด)</h2>")
        P.append("<table><thead><tr><th class='l'>วันที่</th><th class='l'>ช่องทาง</th>"
                 "<th class='l'>รายละเอียด</th><th class='l'>อ้างอิง</th><th class='l'>บัญชี</th>"
                 "<th>เงินเข้า (บาท)</th></tr></thead><tbody>")
        tot = 0.0
        for r in bl:
            tot += float(r["credit"])
            cat = _BANK_CH_TO_CAT.get((r["channel"] or "").upper())
            ch = _MATCH_CH_TH.get(cat, r["channel"] or "—")
            P.append(f"<tr><td class='l mono'>{esc(r['txn_date'])}</td><td class='l'>{esc(ch)}</td>"
                     f"<td class='l'>{esc((r['description'] or '')[:48])}</td>"
                     f"<td class='l mono'>{esc(r['ref_no'] or '—')}</td>"
                     f"<td class='l mono'>{esc(r['account_no'] or '—')}</td>"
                     f"<td class='mono'>{m(r['credit'])}</td></tr>")
        P.append(f"</tbody><tfoot><tr><td colspan='5' class='l'>รวมเงินเข้า</td>"
                 f"<td class='mono'>{m(tot)}</td></tr></tfoot></table>")

    # 4) Grab / Line Man payout summary (gross → fee+VAT → net). Note the VAT column
    # differs: grab='vat_amount', lineman='vat_on_commission'.
    for tbl, label, vatcol in (("recon_grab_lines", "Grab", "vat_amount"),
                               ("recon_lineman_lines", "Line Man", "vat_on_commission")):
        try:
            gr = conn.execute(
                f"SELECT COUNT(*) n, SUM(COALESCE(food_amount,0)) g, SUM(COALESCE(commission_amt,0)) f, "
                f"SUM(COALESCE({vatcol},0)) v, SUM(COALESCE(net_payout,0)) net "
                f"FROM {tbl} WHERE branch_id=? AND substr(order_date,1,7)=? AND order_status='Completed'",
                (branch_id, period_month)).fetchone()
        except sqlite3.OperationalError:
            gr = None
        if gr and gr["n"]:
            P.append(f"<h2>④ {esc(label)} — สรุปออเดอร์สำเร็จ</h2>")
            P.append("<table><thead><tr><th>จำนวนออเดอร์</th><th>ยอดอาหาร</th><th>ค่าคอมมิชชั่น</th>"
                     "<th>VAT</th><th>สุทธิ (payout)</th></tr></thead><tbody>")
            P.append(f"<tr><td class='mono'>{gr['n']:,}</td><td class='mono'>{m(gr['g'])}</td>"
                     f"<td class='mono'>{m(gr['f'])}</td><td class='mono'>{m(gr['v'])}</td>"
                     f"<td class='mono'>{m(gr['net'])}</td></tr></tbody></table>")

    # 5) Manual matches — every hand correction, with reason + who
    mm = conn.execute(
        "SELECT category, bank, COALESCE(pos_amount,0) p, COALESCE(ext_amount,0) e, pos_days, note, matched_by "
        "FROM recon_manual_match WHERE branch_id=? AND period_month=? ORDER BY matched_at",
        (branch_id, period_month)).fetchall()
    if mm:
        P.append("<h2>⑤ การจับคู่ด้วยมือ (พร้อมเหตุผล)</h2>")
        P.append("<table><thead><tr><th class='l'>ช่องทาง</th><th class='l'>ธนาคาร</th><th class='l'>วันที่</th>"
                 "<th>POS</th><th>เอกสาร</th><th class='l'>เหตุผล</th><th class='l'>โดย</th></tr></thead><tbody>")
        for r in mm:
            try:
                days = ", ".join(x.get("day", "")[8:] for x in json.loads(r["pos_days"] or "[]"))
            except Exception:
                days = ""
            P.append(f"<tr><td class='l'>{esc(_MATCH_CH_TH.get(r['category'], r['category']))}</td>"
                     f"<td class='l'>{esc(r['bank'] or '—')}</td><td class='l mono'>{esc(days)}</td>"
                     f"<td class='mono'>{m(r['p'])}</td><td class='mono'>{m(r['e'])}</td>"
                     f"<td class='l'>{esc(r['note'] or '—')}</td><td class='l'>{esc(r['matched_by'] or '—')}</td></tr>")
        P.append("</tbody></table>")

    if not P:
        return "<div class='empty'>ไม่มีข้อมูลหลักฐานสำหรับเดือนนี้</div>"
    return ("<div style='margin-top:14px;border-top:2px solid #ea580c;padding-top:6px'>"
            "<h2 style='border:0;padding:0;color:#9a3412'>📑 หลักฐานประกอบ (สำหรับผู้ตรวจสอบ) — "
            "ตรวจทานได้ทุกบรรทัด</h2>" + "".join(P) + "</div>")


def _monthly_matching_html(conn, period_month: str, branch_id: int | None,
                           detail: bool = False, day_sel: str | None = None) -> str:
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.0f}"

    if branch_id:
        brs = conn.execute("SELECT branch_id, name FROM branches WHERE branch_id=?",
                           (branch_id,)).fetchall()
    else:
        # all-branches pack = the 5 master (active) branches only, like every other view
        brs = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    # master list drives the on-page branch picker (always all 5, regardless of filter)
    master = conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()

    css = """<style>
      @page{size:A4;margin:13mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:11.5px;margin:0;padding:22px;background:#f8fafc}
      .sheet{background:#fff;max-width:1000px;margin:0 auto 18px;padding:24px 28px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:18px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:12px}
      .kpis{display:flex;gap:10px;flex-wrap:wrap;margin:10px 0 14px}
      .kpi{border:1px solid #e5e7eb;border-radius:8px;padding:7px 12px;min-width:96px}
      .kpi .l{font-size:10px;color:#64748b} .kpi .v{font-size:16px;font-weight:700}
      .kpi.ok .v{color:#166534} .kpi.issue .v{color:#b91c1c} .kpi.wait .v{color:#92610e}
      table{width:100%;border-collapse:collapse;font-size:11px;margin-top:4px}
      th,td{padding:4px 8px;border-bottom:1px solid #eef2f7;text-align:right}
      th{color:#64748b;font-size:9.5px;text-transform:uppercase;border-bottom:1.5px solid #e2e8f0;background:#f8fafc;position:sticky;top:0}
      td:first-child,th:first-child,td.l,th.l{text-align:left}
      .mono{font-variant-numeric:tabular-nums}
      tr.daystart td{border-top:2px solid #cbd5e1}
      tr.issue{background:#fef2f2} tr.wait{background:#fffbeb}
      .day{font-weight:700;color:#334155}
      .badge{padding:1px 7px;border-radius:99px;font-size:9.5px;font-weight:700;white-space:nowrap}
      .m-ok{background:#dcfce7;color:#166534}.m-err{background:#fee2e2;color:#b91c1c}
      .m-warn{background:#ffedd5;color:#9a3412}.m-wait{background:#fef9c3;color:#854d0e}
      .m-info{background:#dbeafe;color:#1e40af}.m-mut{background:#f1f5f9;color:#475569}
      .neg{color:#b91c1c} .legend{font-size:10px;color:#64748b;margin:2px 0 10px}
      h2{font-size:13px;margin:18px 0 6px;border-left:4px solid #ea580c;padding-left:8px}
      .sign{display:flex;gap:40px;margin-top:26px}.sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;font-size:11px;color:#475569}
      .toolbar{max-width:1000px;margin:0 auto 12px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:8px;padding:9px 18px;font-size:13px;cursor:pointer}
      .ok-all{background:#f0fdf4;border:1px solid #bbf7d0;border-radius:8px;padding:9px 13px;color:#166534;font-size:12px;margin:8px 0}
      .empty{color:#64748b;font-size:12px;padding:10px 0}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;border-radius:0;max-width:none;margin:0;page-break-after:always}th{position:static}}
    </style>"""

    def th_date(d):
        try:
            return f"{int(d[8:])}"
        except Exception:
            return d

    # query-string fragments preserved when a control changes (so detail/day/branch stick)
    _detail_q = "&detail=1" if detail else ""
    _day_q = f"&day={day_sel}" if day_sel else ""
    _br_q = f"&branch_id={branch_id}" if branch_id else ""

    # on-page branch picker (hidden when printing) — choose 1 of the 5 or all
    bopts = (f"<option value=''{'' if branch_id else ' selected'}>ทุกสาขา (5 สาขาหลัก)</option>"
             + "".join(f"<option value='{b['branch_id']}'"
                       f"{' selected' if branch_id == b['branch_id'] else ''}>{esc(b['name'])}</option>"
                       for b in master))
    picker = (f"<select onchange=\"location.href='?period_month={period_month}'"
              f"+(this.value?'&branch_id='+this.value:'')+'{_detail_q}{_day_q}'\" "
              "style='padding:8px 10px;border:1px solid #cbd5e1;border-radius:8px;"
              "font:inherit;margin-right:8px'>" + bopts + "</select>")
    # date filter — pick ONE day to post that day's accounting; clear to view the whole month
    daypick = (f"<input type='date' value='{esc(day_sel or '')}' "
               f"min='{period_month}-01' max='{period_month}-31' title='เลือกวันเดียว' "
               f"onchange=\"location.href='?period_month={period_month}{_br_q}{_detail_q}'"
               "+(this.value?'&day='+this.value:'')\" "
               "style='padding:7px 10px;border:1px solid #cbd5e1;border-radius:8px;font:inherit;margin-right:8px'>")
    clear = (f"<a href='?period_month={period_month}{_br_q}{_detail_q}' "
             "style='margin-right:10px;font-size:12px;color:#0f6e56'>↺ ทั้งเดือน</a>"
             if day_sel else "")

    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
             "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
             "<title>รายงานการจับคู่รายวัน</title>", css, "</head><body>",
             "<div class='toolbar'>" + picker + daypick + clear
             + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>"]

    for b in brs:
        bid, bname = b["branch_id"], b["name"]
        try:
            d = daily_reconciliation(conn, bid, period_month)
        except Exception:
            d = {"pos_daily_available": False, "days": []}
        parts.append("<div class='sheet'>")
        parts.append(f"<h1>รายงานการจับคู่รายวัน — สาขา{esc(bname)}</h1>")
        if day_sel:
            parts.append(f"<div class='sub'>เฉพาะวันที่ <b>{esc(day_sel)}</b> · "
                         "สำหรับนำการบันทึกบัญชีของวันนี้เข้าระบบบัญชีรายวัน</div>")
        else:
            parts.append(f"<div class='sub'>เดือน {esc(period_month)} · ตรวจสอบว่าทุกช่องทางในแต่ละวันจับคู่กับเอกสารถูกต้อง</div>")
        if not d.get("pos_daily_available"):
            parts.append("<div class='empty'>ยังไม่มีข้อมูล POS รายวันสำหรับเดือนนี้ — ต้องซิงก์ข้อมูลก่อน</div></div>")
            continue

        # active lanes only (POS or document present)
        rows, n_ok, n_issue, n_wait = [], 0, 0, 0
        tot_pos = tot_ext = tot_absvar = 0.0
        prev_day = None
        for day in d["days"]:
            if day_sel and day["day"] != day_sel:   # date filter — one day only
                continue
            day_lanes = [l for l in day["lanes"]
                         if abs(l["pos"]) > 0.005 or abs(l["ext"]) > 0.005]
            for i, l in enumerate(day_lanes):
                lab, cls, bucket = _MATCH_ST.get(l["status"], (l["status"], "m-mut", "none"))
                if bucket == "ok": n_ok += 1
                elif bucket == "issue": n_issue += 1
                elif bucket == "wait": n_wait += 1
                tot_pos += l["pos"]; tot_ext += l["ext"]; tot_absvar += abs(l["variance"])
                trcls = (" daystart" if i == 0 else "") + (f" {bucket}" if bucket in ("issue", "wait") else "")
                daycell = f"<span class='day'>{th_date(day['day'])}</span>" if i == 0 else ""
                bank = "" if l["bank"] in ("NONE", "ANY") else esc(l["bank"])
                rows.append(
                    f"<tr class='{trcls.strip()}'>"
                    f"<td class='l'>{daycell}</td>"
                    f"<td class='l'>{esc(_MATCH_CH_TH.get(l['category'], l['category']))}</td>"
                    f"<td class='l'>{bank}</td>"
                    f"<td class='mono'>{m(l['pos'])}</td>"
                    f"<td class='mono'>{m(l['ext'])}</td>"
                    f"<td class='mono{' neg' if abs(l['variance'])>0.5 else ''}'>{m(l['variance'])}</td>"
                    f"<td class='l'><span class='badge {cls}'>{esc(lab)}</span></td></tr>")
        total_lanes = n_ok + n_issue + n_wait

        parts.append("<div class='kpis'>")
        parts.append(f"<div class='kpi ok'><div class='l'>จับคู่ตรง</div><div class='v'>{n_ok}/{total_lanes}</div></div>")
        parts.append(f"<div class='kpi issue'><div class='l'>ต้องตรวจ (ต่าง/ขาด)</div><div class='v'>{n_issue}</div></div>")
        parts.append(f"<div class='kpi wait'><div class='l'>รอเอกสารตามรอบ</div><div class='v'>{n_wait}</div></div>")
        parts.append(f"<div class='kpi'><div class='l'>POS รวม</div><div class='v mono'>{m(tot_pos)}</div></div>")
        parts.append(f"<div class='kpi'><div class='l'>เอกสาร/เงินเข้ารวม</div><div class='v mono'>{m(tot_ext)}</div></div>")
        parts.append(f"<div class='kpi'><div class='l'>ผลต่างสะสม (abs)</div><div class='v mono'>{m(tot_absvar)}</div></div>")
        parts.append("</div>")

        if n_issue == 0 and total_lanes:
            parts.append("<div class='ok-all'>✓ ทุกช่องทางที่มียอดในเดือนนี้จับคู่กับเอกสารครบ ไม่มีรายการที่ต้องตรวจ"
                         + (f" (มี {n_wait} รายการรอเอกสารเข้าตามรอบ — ปกติ)" if n_wait else "") + "</div>")
        parts.append("<div class='legend'>สถานะ: <b>✓ ตรง/จับคู่มือ</b> = ยอด POS ตรงกับเอกสาร · "
                     "<b>✗ ต่าง/ขาดเอกสาร</b> = ต้องตรวจ · <b>เกินเอกสาร</b> = มีเงินเข้าเกินยอด POS · "
                     "<b>รอเอกสารตามรอบ</b> = ยังไม่ถึงรอบเงินเข้า (ไม่ใช่ปัญหา) · ผลต่าง = เอกสาร − POS</div>")

        if not rows and not d["days"]:
            parts.append("<div class='empty'>ไม่มีรายการในเดือนนี้</div>")
        parts.append("<table><thead><tr><th class='l'>วันที่</th><th class='l'>ช่องทาง</th>"
                     "<th class='l'>ธนาคาร</th><th>ยอด POS</th><th>เอกสาร/เงินเข้า</th>"
                     "<th>ผลต่าง</th><th class='l'>สถานะการจับคู่</th></tr></thead><tbody>")
        parts.extend(rows)
        parts.append("</tbody></table>")

        # hand-confirmed matches (reclass corrections) — the human overrides, with reason
        mm = conn.execute(
            "SELECT category, bank, pos_amount, ext_amount, note, matched_by, matched_at "
            "FROM recon_manual_match WHERE branch_id=? AND period_month=? ORDER BY matched_at",
            (bid, period_month)).fetchall()
        if mm and not detail:        # in detail mode the audit pack (⑤) lists these in full
            parts.append("<h2>การจับคู่ด้วยมือ (เจ้าหน้าที่ยืนยันเอง) — มีเหตุผลกำกับ</h2>")
            parts.append("<table><thead><tr><th class='l'>ช่องทาง</th><th class='l'>ธนาคาร</th>"
                         "<th>ยอด POS</th><th>เอกสาร</th><th class='l'>เหตุผล/หมายเหตุ</th>"
                         "<th class='l'>โดย</th></tr></thead><tbody>")
            for r in mm:
                parts.append(
                    f"<tr><td class='l'>{esc(_MATCH_CH_TH.get(r['category'], r['category']))}</td>"
                    f"<td class='l'>{esc(r['bank'] or '—')}</td>"
                    f"<td class='mono'>{m(r['pos_amount'])}</td><td class='mono'>{m(r['ext_amount'])}</td>"
                    f"<td class='l'>{esc(r['note'] or '—')}</td>"
                    f"<td class='l'>{esc(r['matched_by'] or '—')}</td></tr>")
            parts.append("</tbody></table>")

        if day_sel:
            # date-filtered: just that one day's receipts + its own draft journal (Dr/Cr)
            parts.append(_day_finance_html(conn, bid, day_sel))
        else:
            # per-day accounting breakdown (the daily summary: gross + discount split)
            parts.append(_daily_finance_html(conn, bid, period_month))
            # monthly receipts roll-up + fees/VAT + draft journal entry (accounting tie-out)
            parts.append(_monthly_finance_html(conn, bid, period_month))

        if detail and not day_sel:   # audit-grade line-level evidence for this branch (month-wide)
            parts.append(_audit_detail_html(conn, bid, period_month))

        parts.append("<div class='sign'><div>ผู้จัดทำ (กระทบยอด)</div>"
                     "<div>ผู้ตรวจสอบ</div><div>ผู้อนุมัติ</div></div>")
        parts.append("</div>")

    parts.append("</body></html>")
    return "".join(parts)


@app.get("/api/monthly-report")
def monthly_matching_report(period_month: str, branch_id: int | None = None,
                            detail: int = 0, day: str | None = None):
    """Print-to-PDF monthly matching audit: every day × channel with POS vs document,
    variance, and match status — confirm the whole month reconciled. Omit branch_id
    for an all-branches pack (one page each). detail=1 appends audit-grade line-level
    evidence (POS codes, credit-by-acquirer, every bank line, payouts, manual matches).
    day=YYYY-MM-DD (within period_month) scopes the report to one calendar day and shows
    that day's own draft journal entry (Dr/Cr) so it can be posted to the daily books."""
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    if day is not None:
        if not _DAY_RE.match(day) or day[:7] != period_month:
            raise HTTPException(400, "day must be 'YYYY-MM-DD' within period_month.")
    if branch_id:
        with get_db() as conn:
            _validate_master_branch(conn, branch_id)
    with get_db() as conn:
        html = _monthly_matching_html(conn, period_month, branch_id,
                                      detail=bool(detail), day_sel=day)
    return Response(content=html, media_type="text/html; charset=utf-8")


# ── Daily accounting summary, one card per day (สรุปบัญชีรายวัน รายใบ) ──────────
# Stand-alone section: the same per-day GROSS / discount-on-its-own-line breakdown that
# the monthly report carries, but pulled out as its own browsable view — each day is a
# card ('1 วัน = 1 ใบ'), one branch at a time, with a print-to-PDF version.

@app.get("/api/daily-accounts")
def daily_accounts(period_month: str, branch_id: int | None = None):
    """สรุปบัญชีรายวัน (รายใบ) for ONE branch: per-day accounting where each day lists its
    channels at GROSS with discount/fee/VAT split out + account code/name, plus a day
    total. Returns the master-branch list so the UI can switch branches; omitting
    branch_id defaults to the first master branch."""
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        masters = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
        branches = [{"branch_id": b["branch_id"], "name": b["name"]} for b in masters]
        if not branches:
            return {"period_month": period_month, "branches": [], "branch_id": None,
                    "branch_name": None, "days": [], "totals": {}}
        if branch_id is None:
            branch_id = branches[0]["branch_id"]
        else:
            _validate_master_branch(conn, branch_id)
        bname = next((b["name"] for b in branches if b["branch_id"] == branch_id), None)
        data = _daily_finance_data(conn, branch_id, period_month)
        acct = _account_map(conn)
        # account codes/names for the per-day draft journal (Dr/Cr) the UI renders per card
        je_accounts = {s: list(acct.get(s, ("", "")))
                       for s in ("BANK", "DISCOUNT", "FEE", "VAT", "REVENUE")}
    return {"period_month": period_month, "branches": branches,
            "branch_id": branch_id, "branch_name": bname,
            "days": data["days"], "totals": data["totals"], "je_accounts": je_accounts}


def _daily_dashboard_data(conn, day: str) -> dict:
    """Revenue by channel for ONE calendar day across ALL master branches: each branch's
    channels (gross/discount/fee/VAT/net) for that day + per-branch totals, plus grand
    totals across branches. Reuses _daily_finance_data so the KTC/BAY coupon split + the
    ส่วนลด(คูปอง) deduction line are identical to the daybook and monthly report — then
    extracts the single matching day. Also attaches the POS↔money-in view per channel
    (pos_channels): what the register rang up (POS) vs what landed in the bank/settlement
    (ext) + the variance — the reconciliation lens, from the same daily_reconciliation."""
    period_month = day[:7]
    masters = conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    zero = {"gross": 0.0, "discount": 0.0, "fee": 0.0, "vat": 0.0, "net": 0.0}
    pzero = {"pos": 0.0, "money": 0.0}
    branches, grand, pgrand = [], dict(zero), dict(pzero)
    for b in masters:
        bid = b["branch_id"]
        try:
            daily = daily_reconciliation(conn, bid, period_month)
        except Exception:
            daily = {"days": []}
        data = _daily_finance_data(conn, bid, period_month, daily=daily)   # one daily compute
        dd = next((x for x in data["days"] if x["day"] == day), None)
        channels = dd["channels"] if dd else []
        totals = dd["totals"] if dd else dict(zero)
        # POS ↔ money-in per channel for the day (sum lanes over banks)
        day_lanes = next((x["lanes"] for x in daily.get("days", []) if x["day"] == day), [])
        pos_agg: dict = {}
        for l in day_lanes:
            e = pos_agg.setdefault(l["category"], {"pos": 0.0, "money": 0.0})
            e["pos"] += float(l.get("pos") or 0); e["money"] += float(l.get("ext") or 0)
        pos_channels, ptot = [], dict(pzero)
        for ch in sorted(pos_agg, key=lambda c: _FINANCE_CH_ORDER.get(c, 9)):
            p, mny = round(pos_agg[ch]["pos"], 2), round(pos_agg[ch]["money"], 2)
            if abs(p) <= 0.005 and abs(mny) <= 0.005:
                continue
            pos_channels.append({"ch": ch, "ch_th": _MATCH_CH_TH.get(ch, ch),
                                 "pos": p, "money": mny, "variance": round(mny - p, 2)})
            ptot["pos"] += p; ptot["money"] += mny
        ptot = {"pos": round(ptot["pos"], 2), "money": round(ptot["money"], 2),
                "variance": round(ptot["money"] - ptot["pos"], 2)}
        branches.append({"branch_id": bid, "name": b["name"],
                         "channels": channels, "totals": totals,
                         "pos_channels": pos_channels, "pos_totals": ptot})
        for k in grand:
            grand[k] += float(totals.get(k, 0) or 0)
        pgrand["pos"] += ptot["pos"]; pgrand["money"] += ptot["money"]
    return {"day": day, "branches": branches,
            "grand_totals": {k: round(v, 2) for k, v in grand.items()},
            "pos_grand": {"pos": round(pgrand["pos"], 2), "money": round(pgrand["money"], 2),
                          "variance": round(pgrand["money"] - pgrand["pos"], 2)}}


@app.get("/api/daily-dashboard")
def daily_dashboard(day: str):
    """Dashboard: รายได้ตามช่องทางของทุกสาขาในวันเดียว (เลือกวันจากปฏิทิน). Returns each
    master branch's per-channel breakdown for `day` + per-branch totals + grand totals
    across branches. Inherits the coupon split (KTC/BAY points kept as เงินรับจริง; other
    discount coupons surfaced as a ส่วนลด deduction line) from _daily_finance_data."""
    try:
        _dt.strptime(day or "", "%Y-%m-%d")
    except ValueError:
        raise HTTPException(400, "day must be in 'YYYY-MM-DD' format.")
    with get_db() as conn:
        return _daily_dashboard_data(conn, day)


def _daily_accounts_print_html(conn, period_month: str, branch_id: int | None) -> str:
    """Printable per-day card report. Each day is a card (ใบ) with its channels at GROSS
    + discount/fee/VAT split + day total; one branch per sheet (all masters if no
    branch_id). Mirrors _daily_finance_data so screen and paper agree to the baht."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.2f}"
    def neg(v):  # discount/fee/VAT print as a deduction; blank when zero
        return f"−{m(v)}" if abs(float(v or 0)) > 0.005 else "—"
    acct = _account_map(conn)

    def journal_block(t):  # per-day draft journal (Dr/Cr) for posting to the daily books
        def code(slot):
            c = acct.get(slot, ("", ""))[0]
            return f" <span class='code'>{esc(c)}</span>" if c else ""
        rows = [f"<tr><td>เงินสด/ธนาคาร (สุทธิ){code('BANK')}</td><td class='mono'>{m(t['net'])}</td><td class='mono'>—</td></tr>"]
        if abs(float(t.get("discount") or 0)) > 0.005:
            rows.append(f"<tr><td>ส่วนลด/โปรโมชัน{code('DISCOUNT')}</td><td class='mono'>{m(t['discount'])}</td><td class='mono'>—</td></tr>")
        if abs(float(t.get("fee") or 0)) > 0.005:
            rows.append(f"<tr><td>ค่าธรรมเนียมรับชำระเงิน{code('FEE')}</td><td class='mono'>{m(t['fee'])}</td><td class='mono'>—</td></tr>")
        if abs(float(t.get("vat") or 0)) > 0.005:
            rows.append(f"<tr><td>ภาษีซื้อ–VAT ค่าธรรมเนียม{code('VAT')}</td><td class='mono'>{m(t['vat'])}</td><td class='mono'>—</td></tr>")
        rows.append(f"<tr><td>รายได้จากการขาย (เต็ม){code('REVENUE')}</td><td class='mono'>—</td><td class='mono'>{m(t['gross'])}</td></tr>")
        dtot = float(t["net"]) + float(t["discount"]) + float(t["fee"]) + float(t["vat"])
        return ("<table class='je'><thead><tr><th>ร่างบันทึกบัญชี (Dr/Cr)</th><th>เดบิต</th><th>เครดิต</th></tr></thead><tbody>"
                + "".join(rows)
                + f"</tbody><tfoot><tr><td>รวม</td><td class='mono'>{m(dtot)}</td><td class='mono'>{m(t['gross'])}</td></tr></tfoot></table>")

    if branch_id:
        brs = conn.execute("SELECT branch_id, name FROM branches WHERE branch_id=?",
                           (branch_id,)).fetchall()
    else:
        brs = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    master = conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()

    css = """<style>
      @page{size:A4;margin:13mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:11.5px;margin:0;padding:22px;background:#f8fafc}
      .sheet{background:#fff;max-width:900px;margin:0 auto 18px;padding:24px 28px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:18px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:14px}
      .legend{font-size:10px;color:#64748b;margin:2px 0 14px}
      .cards{display:grid;grid-template-columns:repeat(2,1fr);gap:12px}
      .dc{border:1px solid #e2e8f0;border-radius:9px;overflow:hidden;break-inside:avoid}
      .dc-h{display:flex;justify-content:space-between;align-items:center;padding:7px 12px;background:#ecfdf5;border-bottom:1px solid #d1fae5}
      .dc-date{font-weight:700;color:#0a523f;font-size:12.5px}
      .dc-net{font-weight:700;font-size:12.5px}
      table{width:100%;border-collapse:collapse;font-size:10.5px}
      th,td{padding:3px 12px;border-bottom:1px solid #f1f5f9;text-align:right}
      th{color:#64748b;font-size:9px;text-transform:uppercase;background:#fff}
      td:first-child,th:first-child{text-align:left}
      .mono{font-variant-numeric:tabular-nums}
      .disc{color:#9a3412} .code{color:#94a3b8;font-size:9px}
      tfoot td{border-top:1.4px solid #cbd5e1;border-bottom:none;font-weight:700;background:#f8fafc}
      table.je{border-top:1px solid #d1fae5} table.je th{background:#ecfdf5;color:#0a523f}
      .mtot{margin-top:14px;border:1.5px solid #0f6e56;border-radius:9px;overflow:hidden}
      .mtot .dc-h{background:#0f6e56;border-bottom:none} .mtot .dc-date,.mtot .dc-net{color:#fff}
      .empty{color:#64748b;font-size:12px;padding:10px 0}
      .toolbar{max-width:900px;margin:0 auto 12px;text-align:right}
      .btn{background:#0f6e56;color:#fff;border:0;border-radius:8px;padding:9px 18px;font-size:13px;cursor:pointer}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;border-radius:0;max-width:none;margin:0;page-break-after:always}}
    </style>"""

    bopts = (f"<option value=''{'' if branch_id else ' selected'}>ทุกสาขา ({len(master)} สาขา)</option>"
             + "".join(f"<option value='{b['branch_id']}'"
                       f"{' selected' if branch_id == b['branch_id'] else ''}>{esc(b['name'])}</option>"
                       for b in master))
    picker = (f"<select onchange=\"location.href='?period_month={period_month}'"
              f"+(this.value?'&branch_id='+this.value:'')\" "
              "style='padding:8px 10px;border:1px solid #cbd5e1;border-radius:8px;"
              "font:inherit;margin-right:8px'>" + bopts + "</select>")

    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
             "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
             "<title>สรุปบัญชีรายวัน (รายใบ)</title>", css, "</head><body>",
             "<div class='toolbar'>" + picker
             + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>"]

    def card_table(channels, foot, foot_label):
        rows = ["<table><thead><tr><th>ช่องทาง</th><th>ยอดเต็ม</th><th>ส่วนลด</th>"
                "<th>ค่าธรรมเนียม</th><th>VAT</th><th>สุทธิ</th></tr></thead><tbody>"]
        for a in channels:
            code = f" <span class='code'>{esc(a['code'])}</span>" if a.get("code") else ""
            rows.append(
                f"<tr><td>{esc(a['ch_th'])}{code}</td><td class='mono'>{m(a['gross'])}</td>"
                f"<td class='mono disc'>{neg(a['discount'])}</td><td class='mono'>{neg(a['fee'])}</td>"
                f"<td class='mono'>{neg(a['vat'])}</td><td class='mono'>{m(a['net'])}</td></tr>")
        rows.append(
            f"</tbody><tfoot><tr><td>{foot_label}</td><td class='mono'>{m(foot['gross'])}</td>"
            f"<td class='mono disc'>{neg(foot['discount'])}</td><td class='mono'>{neg(foot['fee'])}</td>"
            f"<td class='mono'>{neg(foot['vat'])}</td><td class='mono'>{m(foot['net'])}</td></tr></tfoot></table>")
        return "".join(rows)

    for b in brs:
        data = _daily_finance_data(conn, b["branch_id"], period_month)
        parts.append("<div class='sheet'>")
        parts.append(f"<h1>สรุปบัญชีรายวัน — สาขา{esc(b['name'])}</h1>")
        parts.append(f"<div class='sub'>เดือน {esc(period_month)} · แต่ละวัน = 1 ใบบันทึกบัญชี "
                     "(บันทึกยอดเต็ม แล้วแยกบรรทัดส่วนลด)</div>")
        if not data["days"]:
            parts.append("<div class='empty'>ไม่มีรายการบัญชีในเดือนนี้</div></div>")
            continue
        parts.append("<div class='cards'>")
        for day in data["days"]:
            parts.append("<div class='dc'><div class='dc-h'>"
                         f"<span class='dc-date'>วันที่ {day['dom']}</span>"
                         f"<span class='dc-net'>สุทธิ ฿{m(day['totals']['net'])}</span></div>"
                         + card_table(day["channels"], day["totals"], "รวมวันนี้")
                         + journal_block(day["totals"]) + "</div>")
        parts.append("</div>")
        parts.append("<div class='mtot'><div class='dc-h'>"
                     "<span class='dc-date'>รวมทั้งเดือน</span>"
                     f"<span class='dc-net'>สุทธิ ฿{m(data['totals']['net'])}</span></div>"
                     + card_table([], data["totals"], "รวมทั้งเดือน") + "</div>")
        parts.append("<div class='legend'>ยอดเต็ม = ยอดขายก่อนหักส่วนลด/ค่าธรรมเนียม · "
                     "ส่วนลด/ค่าธรรมเนียม/VAT แสดงเป็นยอดหัก (−) แยกบรรทัด · "
                     "สุทธิ = ยอดเต็ม − ส่วนลด − ค่าธรรมเนียม − VAT · รหัสบัญชีตั้งที่หน้าตั้งค่าระบบ</div>")
        parts.append("</div>")  # sheet
    parts.append("</body></html>")
    return "".join(parts)


@app.get("/api/daily-accounts.print")
def daily_accounts_print(period_month: str, branch_id: int | None = None):
    """Print-to-PDF of สรุปบัญชีรายวัน (รายใบ): one card per day. Omit branch_id for an
    all-branches pack (one sheet each)."""
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    if branch_id:
        with get_db() as conn:
            _validate_master_branch(conn, branch_id)
    with get_db() as conn:
        html = _daily_accounts_print_html(conn, period_month, branch_id)
    return Response(content=html, media_type="text/html; charset=utf-8")


# ── Outstanding-documents follow-up (รายงานติดตามเอกสารที่ขาด รายสาขา/เดือน) ────
# "Which documents does each branch still need to upload this month, and how much
# money is waiting on each?" — works for ANY period so the team can chase month-to-month.

def _doc_channel(label: str) -> str | None:
    """Map a required-document label to the channel it backs (for the baht-at-stake)."""
    l = (label or "")
    if "บัตร" in l:                                   return "CREDIT"
    if "Grab" in l or "แกร็บ" in l:                   return "GRAB"
    if "Line Man" in l or "ไลน์แมน" in l or "LM" in l.upper(): return "LINEMAN"
    if "คะแนน" in l or "คูปอง" in l:                 return "COUPON"
    if "สเตทเมนต์" in l or "statement" in l.lower():  return "BANK"   # backs QR/CASH/TRANSFER
    return None


def _outstanding_docs_report(conn, period_month: str, branch_id: int | None = None) -> dict:
    from collections import defaultdict
    if branch_id:
        masters = conn.execute("SELECT branch_id, name FROM branches WHERE branch_id=?",
                               (branch_id,)).fetchall()
    else:
        masters = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()

    branches, grand = [], 0.0
    for b in masters:
        try:
            rm = reconcile_month(conn, b["branch_id"], period_month)
        except Exception:
            continue
        chmap = {c["key"]: c for c in rm.get("channels", [])}
        miss = defaultdict(list)            # channel -> [missing/partial doc labels]
        for doc in rm.get("documents", []):
            if doc.get("status") in ("missing", "partial"):
                miss[_doc_channel(doc["label"])].append(
                    {"label": doc["label"], "status": doc["status"]})

        # POS composition (for the detailed per-acquirer / per-code breakdown of what's waiting)
        try:
            codes = pos_codes_scope(conn, b["branch_id"], period_month=period_month)
        except sqlite3.OperationalError:
            codes = []
        by_ch = defaultdict(list)
        for cd in codes:
            if cd["category"] and cd["amount"] > 0:
                by_ch[cd["category"]].append(cd)

        rows = []
        for key, docs in miss.items():
            if key == "BANK":               # statements back several lanes; report at branch level
                continue
            c = chmap.get(key, {})
            unmatched = round(max(0.0, float(c.get("pos", 0)) - float(c.get("ext", 0))), 2)
            # detail: for CREDIT split by acquirer (which settlement to chase); else by POS code
            agg = defaultdict(lambda: {"amount": 0.0, "txn": 0})
            for cd in by_ch.get(key, []):
                k = (_BANK_LABEL.get(cd["bank"], cd["bank"]) if key == "CREDIT"
                     else cd["code"])
                agg[k]["amount"] += cd["amount"]; agg[k]["txn"] += cd["txn_count"]
            breakdown = sorted(
                ({"label": k, "amount": round(v["amount"], 2), "txn": v["txn"]}
                 for k, v in agg.items()), key=lambda x: -x["amount"])
            rows.append({"channel": key, "channel_th": _MATCH_CH_TH.get(key, key),
                         "unmatched": unmatched, "breakdown": breakdown,
                         "docs": [d["label"] for d in docs],
                         "doc_status": "partial" if any(d["status"] == "partial" for d in docs) else "missing"})
        # missing bank statements (rare — usually uploaded) listed separately
        bank_missing = [d["label"] for d in miss.get("BANK", [])]
        rows.sort(key=lambda r: -r["unmatched"])
        branch_total = round(sum(r["unmatched"] for r in rows), 2)
        grand += branch_total
        if rows or bank_missing:
            branches.append({"branch_id": b["branch_id"], "branch_name": b["name"],
                             "progress_pct": rm.get("progress_pct"),
                             "rows": rows, "bank_missing": bank_missing,
                             "branch_total": branch_total})
    branches.sort(key=lambda x: -x["branch_total"])
    return {"period_month": period_month, "branches": branches,
            "grand_total": round(grand, 2)}


@app.get("/api/reports/outstanding-docs")
def outstanding_docs_report(period_month: str, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _outstanding_docs_report(conn, period_month, branch_id)


@app.get("/api/reports/outstanding-docs.print")
def outstanding_docs_report_print(period_month: str, branch_id: int | None = None):
    """Print-to-PDF follow-up sheet: per branch, the documents still missing this month
    + the baht waiting on each. Month + branch picker on the page (hidden when printing)."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        rep = _outstanding_docs_report(conn, period_month, branch_id)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()

    css = """<style>
      @page{size:A4;margin:14mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12px;margin:0;padding:22px;background:#f8fafc}
      .sheet{background:#fff;max-width:920px;margin:0 auto 16px;padding:24px 28px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:19px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:6px}
      h2{font-size:14px;margin:16px 0 6px;border-left:4px solid #ea580c;padding-left:8px}
      table{width:100%;border-collapse:collapse;font-size:12px;margin:4px 0 6px}
      th,td{padding:6px 9px;border-bottom:1px solid #eef2f7;text-align:right}
      th{color:#64748b;font-size:10px;border-bottom:1.5px solid #e2e8f0;background:#f8fafc}
      td:first-child,th:first-child,td.l,th.l{text-align:left}
      .mono{font-variant-numeric:tabular-nums} .big{color:#b91c1c;font-weight:700}
      .muted{color:#64748b;font-size:11px} tr.sub td{border-bottom:1px dashed #f1f5f9;padding:3px 9px}
      tfoot td{border-top:1.5px solid #cbd5e1;font-weight:700}
      .chk{width:16px;height:16px;border:1.5px solid #94a3b8;border-radius:3px;display:inline-block;vertical-align:middle}
      .partial{color:#92610e}.lead{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:9px 13px;font-size:12px;margin-bottom:12px}
      .none{background:#f0fdf4;border:1px solid #bbf7d0;border-radius:8px;padding:9px 13px;color:#166534;font-size:12px}
      .toolbar{max-width:920px;margin:0 auto 12px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:8px;padding:9px 18px;font-size:13px;cursor:pointer}
      .ctl{padding:8px 10px;border:1px solid #cbd5e1;border-radius:8px;font:inherit;margin-right:8px}
      .sign{display:flex;gap:40px;margin-top:22px}.sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;font-size:11px;color:#475569}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;border-radius:0;max-width:none;margin:0}}
    </style>"""

    bopts = (f"<option value=''{'' if branch_id else ' selected'}>ทุกสาขา (5 สาขาหลัก)</option>"
             + "".join(f"<option value='{b['branch_id']}'"
                       f"{' selected' if branch_id == b['branch_id'] else ''}>{esc(b['name'])}</option>"
                       for b in master))
    bsuffix = f"&branch_id={branch_id}" if branch_id else ""
    month_oc = "location.href='?period_month='+this.value+'" + bsuffix + "'"
    sel_oc = (f"location.href='?period_month={period_month}'"
              "+(this.value?'&branch_id='+this.value:'')")
    picker = (f"<input type='month' value='{esc(period_month)}' class='ctl' onchange=\"{month_oc}\">"
              f"<select class='ctl' onchange=\"{sel_oc}\">" + bopts + "</select>")

    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
             "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
             "<title>ติดตามเอกสารที่ขาด</title>", css, "</head><body>",
             "<div class='toolbar'>" + picker
             + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>",
             "<div class='sheet'>",
             f"<h1>รายงานติดตามเอกสารที่ยังขาด — งวด {esc(period_month)}</h1>",
             f"<div class='sub'>เอกสารที่แต่ละสาขายังต้องอัปโหลด + ยอดเงินที่รอเอกสารนั้นจับคู่ · "
             f"รวมเงินที่รอเอกสารทั้งหมด <b class='big'>฿{m(rep['grand_total'])}</b></div>"]

    if not rep["branches"]:
        parts.append("<div class='none'>✓ ทุกสาขาอัปเอกสารครบสำหรับงวดนี้ — ไม่มีเอกสารค้าง</div>")
    for b in rep["branches"]:
        parts.append(f"<h2>{esc(b['branch_name'])} — เหลือเอกสารคิดเป็นเงิน ฿{m(b['branch_total'])}"
                     + (f" · ความคืบหน้า {b['progress_pct']}%" if b.get('progress_pct') is not None else "") + "</h2>")
        if b["rows"]:
            parts.append("<table><thead><tr><th class='l'>อัปแล้ว</th><th class='l'>ช่องทาง</th>"
                         "<th class='l'>เอกสารที่ต้องอัปโหลด</th><th>ยอดที่รอจับคู่ (บาท)</th></tr></thead><tbody>")
            for r in b["rows"]:
                docs = " · ".join(esc(x) for x in r["docs"])
                st = " <span class='partial'>(อัปบางส่วน)</span>" if r["doc_status"] == "partial" else ""
                parts.append(f"<tr><td class='l'><span class='chk'></span></td>"
                             f"<td class='l'><b>{esc(r['channel_th'])}</b></td>"
                             f"<td class='l'>{docs}{st}</td>"
                             f"<td class='mono{' big' if r['unmatched']>20000 else ''}'>{m(r['unmatched'])}</td></tr>")
                # detailed breakdown: per acquirer (CREDIT) or per POS code — so the team
                # knows exactly which document to get and the expected amount it should total
                for d in r.get("breakdown", []):
                    note = ("เอกสารต้องมียอดราว ฿" + m(d["amount"])) if r["channel"] == "CREDIT" else ""
                    parts.append(
                        f"<tr class='sub'><td></td><td class='l muted'>↳ {esc(d['label'])}</td>"
                        f"<td class='l muted'>{esc(note)} <span class='muted'>({d['txn']:,} รายการ)</span></td>"
                        f"<td class='mono muted'>{m(d['amount'])}</td></tr>")
            parts.append(f"</tbody><tfoot><tr><td colspan='3' class='l'>รวมเงินที่รอเอกสาร</td>"
                         f"<td class='mono'>{m(b['branch_total'])}</td></tr></tfoot></table>")
        if b["bank_missing"]:
            parts.append("<div class='lead'>⚠ สเตทเมนต์ธนาคารที่ยังขาด: "
                         + " · ".join(esc(x) for x in b["bank_missing"])
                         + " (จำเป็นต่อ QR/เงินสด/เงินโอน)</div>")
        parts.append("<div class='sign'><div>ผู้รับผิดชอบสาขา</div><div>วันที่อัปครบ</div></div>")
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


# ════════════════════════════════════════════════════════════════════════════
#  Internal-control reports (Big4 audit evidence) — all READ-ONLY.
#  #1 Management-override register · #2 Reconciliation sign-off certificate ·
#  #3 Open-items aging.  None of these mutate state or touch accounting/recon
#  logic; they only READ the existing reversible-adjustment tables and
#  reconcile_month(). Master-branch (company-owned) scope only.
# ════════════════════════════════════════════════════════════════════════════

_IC_BANK_TH = {"KTB": "กรุงไทย", "KBANK": "กสิกร", "KTC": "KTC", "BAY": "กรุงศรี",
               "SCB": "ไทยพาณิชย์", "BBL": "กรุงเทพ", "TTB": "ทีทีบี",
               "PROMPTPAY": "พร้อมเพย์", "ANY": "ทุกธนาคาร", "NONE": "—", "ALL": "ทุกช่อง"}
_IC_SESS_TH = {"draft": "ร่าง (ยังไม่สอบทาน)", "in_review": "รอสอบทาน",
               "approved": "อนุมัติแล้ว", "closed": "ปิดงวดแล้ว"}
_IC_CHST_TH = {"ok": "ตรงกัน", "diff": "มีผลต่าง", "waiting_doc": "รอเอกสาร",
               "no_source": "ไม่มีคู่เทียบ", "setup": "รอตั้งค่า"}


def _ic_lane_th(cat, bank):
    base = _CH_LABEL.get(cat, cat or "—")
    if bank and bank not in ("NONE", "ALL", "ANY", "", None):
        return f"{base}·{_IC_BANK_TH.get(bank, bank)}"
    return base


def _override_register(conn, period_month, branch_id=None):
    """Every manual intervention (management override) in the month across master
    branches: manual tick-to-match, cross-channel move, lane adjustment, and per-day
    acknowledged variance — the four reversible tables. One normalized row each:
    branch · when · who · kind · lane · amount · reason. Read-only."""
    bmap = {r["branch_id"]: r["name"] for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY}")}

    def keep(bid):
        return (bid in bmap) and (branch_id is None or bid == branch_id)

    rows = []
    for r in conn.execute("SELECT * FROM recon_manual_match WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["matched_at"], "by": r["matched_by"] or "—",
            "kind": "จับคู่ด้วยมือ", "kind_key": "manual_match",
            "lane": _ic_lane_th(r["category"], r["bank"]),
            "amount": round(float(r["ext_amount"] or 0), 2),
            "detail": (f"POS {float(r['pos_amount'] or 0):,.0f} ↔ เงินเข้า "
                       f"{float(r['ext_amount'] or 0):,.0f} "
                       f"(ต่าง {float(r['variance'] or 0):+,.0f})"),
            "reason": r["note"] or ""})
    for r in conn.execute("SELECT * FROM recon_channel_move WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["created_at"], "by": r["created_by"] or "—",
            "kind": "ย้ายช่อง", "kind_key": "channel_move",
            "lane": (f"{_ic_lane_th(r['from_category'], r['from_bank'])} → "
                     f"{_ic_lane_th(r['to_category'], r['to_bank'])}"),
            "amount": round(float(r["amount"] or 0), 2),
            "detail": f"ย้าย {float(r['amount'] or 0):,.0f} บาท",
            "reason": r["note"] or ""})
    _adj_kind = {"accept": "รับทราบผลต่าง (ทั้งช่อง)", "off_pos": "เพิ่มยอด POS",
                 "deposit_note": "บันทึกเงินฝาก", "other": "ปรับปรุงอื่น ๆ"}
    for r in conn.execute("SELECT * FROM recon_lane_adj WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        amt = float(r["adj_pos"] or 0) + float(r["adj_ext"] or 0)
        det = ((f"POS +{float(r['adj_pos'] or 0):,.0f}" if r["adj_pos"] else "")
               + (f" เงินเข้า +{float(r['adj_ext'] or 0):,.0f}" if r["adj_ext"] else "")).strip()
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["created_at"], "by": r["created_by"] or "—",
            "kind": _adj_kind.get(r["adj_type"], "ปรับปรุง"), "kind_key": "lane_adj",
            "lane": _ic_lane_th(r["category"], r["bank"]),
            "amount": round(amt, 2),
            "detail": det or "รับทราบผลต่างทั้งช่องให้ปิดงวดได้",
            "reason": r["note"] or ""})
    for r in conn.execute("SELECT * FROM recon_day_ack WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["created_at"], "by": r["created_by"] or "—",
            "kind": "รับทราบผลต่างรายวัน", "kind_key": "day_ack",
            "lane": f"{_CH_LABEL.get(r['category'], r['category'])} · {r['day']}",
            "amount": round(float(r["variance"] or 0), 2),
            "detail": f"วันที่ {r['day']} ต่าง {float(r['variance'] or 0):+,.0f} บาท",
            "reason": r["note"] or ""})
    rows.sort(key=lambda x: (x["at"] or ""), reverse=True)
    for r in rows:
        r["has_reason"] = bool((r["reason"] or "").strip())
    return {"period_month": period_month, "branch_id": branch_id,
            "rows": rows, "count": len(rows),
            "total_abs_amount": round(sum(abs(r["amount"]) for r in rows), 2),
            "missing_reason": sum(0 if r["has_reason"] else 1 for r in rows)}


def _ic_period_end(m):
    import calendar
    from datetime import date
    y, mo = int(m[:4]), int(m[5:7])
    return date(y, mo, calendar.monthrange(y, mo)[1])


def _ic_prev_month(m, k):
    y, mo = int(m[:4]), int(m[5:7])
    idx = y * 12 + (mo - 1) - k
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def _open_items_aging(conn, period_month, branch_id=None, lookback=2):
    """Open (unresolved) reconciling items, aged from period-end to today, for the
    selected month + `lookback` prior months. Skips approved/closed (locked) sessions —
    those are signed off, not open. Each item = one reconcile_month todo. Read-only."""
    from datetime import date
    today = date.today()
    months = [period_month] + [_ic_prev_month(period_month, k)
                               for k in range(1, lookback + 1)]
    branches = [(r["branch_id"], r["name"]) for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name")
        if (branch_id is None or r["branch_id"] == branch_id)]
    items = []
    for m in months:
        pend = _ic_period_end(m)
        age = max(0, (today - pend).days)
        bucket = "0-7" if age <= 7 else ("8-30" if age <= 30 else ">30")
        for bid, bname in branches:
            sess = conn.execute(
                "SELECT status FROM recon_sessions WHERE branch_id=? AND period_month=?",
                (bid, m)).fetchone()
            if sess and sess["status"] in ("approved", "closed"):
                continue
            try:
                r = reconcile_month(conn, bid, m)
            except Exception:
                continue
            for t in r["todo"]:
                items.append({"branch_id": bid, "branch_name": bname,
                    "period_month": m, "age_days": age, "bucket": bucket,
                    "type": t.get("type"), "label": t.get("label"),
                    "channel": t.get("channel")})
    items.sort(key=lambda x: (-x["age_days"], x["branch_name"]))
    buckets = {"0-7": 0, "8-30": 0, ">30": 0}
    for it in items:
        buckets[it["bucket"]] += 1
    return {"as_of": today.isoformat(), "period_month": period_month,
            "lookback": lookback, "months": months,
            "items": items, "count": len(items), "buckets": buckets}


def _recon_certificate(conn, branch_id, period_month):
    """Sign-off certificate data for one branch-month: per-channel POS vs money-in
    with residual after acknowledgements, preparer/reviewer + approval date, the
    frozen working-paper snapshot hash, and the count of manual overrides. Read-only."""
    r = reconcile_month(conn, branch_id, period_month)
    chs = r["channels"]
    rows, resid_total = [], 0.0
    for c in chs:
        resid = c.get("residual", c.get("variance", 0)) or 0
        if c["key"] != "COUPON":          # COUPON variance is loyalty points, not baht
            resid_total += resid
        rows.append({"key": c["key"], "label": c["label"],
            "pos": c["pos"], "ext": c["ext"], "variance": c.get("variance", 0),
            "residual": resid, "status": c["status"],
            "accepted": bool(c.get("accepted")), "detail": c.get("detail", "")})
    reg = _override_register(conn, period_month, branch_id)
    ack_rows = [x for x in reg["rows"] if x["kind_key"] in ("day_ack", "lane_adj")]
    snap = conn.execute(
        "SELECT status, sha256, created_at, file_name, source_manifest FROM recon_close_snapshot "
        "WHERE session_id=(SELECT session_id FROM recon_sessions "
        "                  WHERE branch_id=? AND period_month=?) "
        "ORDER BY id DESC LIMIT 1", (branch_id, period_month)).fetchone()
    snap_d = dict(snap) if snap else None
    if snap_d:
        try:
            snap_d["source_doc_count"] = len(json.loads(snap_d.get("source_manifest") or "[]"))
        except Exception:
            snap_d["source_doc_count"] = 0
    return {"branch_id": branch_id, "branch_name": r["branch_name"],
        "period_month": period_month, "session": r.get("session") or {},
        "progress_pct": r["progress_pct"], "todo_count": len(r["todo"]),
        "channels": rows,
        "pos_total": round(sum(c["pos"] for c in chs), 2),
        "ext_total": round(sum(c["ext"] for c in chs), 2),
        "residual_total": round(resid_total, 2),
        "ack_count": len(ack_rows),
        "ack_amount": round(sum(abs(x["amount"]) for x in ack_rows), 2),
        "override_count": reg["count"], "override_missing_reason": reg["missing_reason"],
        "snapshot": snap_d}


# ── JSON endpoints (drive the in-app views) ───────────────────────────────────
@app.get("/api/reports/override-register")
def override_register_api(period_month: str, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _override_register(conn, period_month, branch_id)


@app.get("/api/reports/open-items-aging")
def open_items_aging_api(period_month: str, branch_id: int | None = None,
                         lookback: int = 2):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _open_items_aging(conn, period_month, branch_id, max(0, min(6, lookback)))


@app.get("/api/reports/recon-certificate")
def recon_certificate_api(branch_id: int, period_month: str):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        return _recon_certificate(conn, branch_id, period_month)


# ── Shared print CSS (matches the outstanding-docs sheet look) ────────────────
def _ic_print_css(accent):
    return ("<style>@page{size:A4;margin:14mm}"
        "*{box-sizing:border-box}"
        "body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;"
        "font-size:12px;margin:0;padding:22px;background:#f8fafc}"
        ".sheet{background:#fff;max-width:960px;margin:0 auto 16px;padding:24px 28px;"
        "border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}"
        "h1{font-size:19px;margin:0 0 2px}.sub{color:#64748b;font-size:12px;margin-bottom:8px}"
        f"h2{{font-size:14px;margin:16px 0 6px;border-left:4px solid {accent};padding-left:8px}}"
        "table{width:100%;border-collapse:collapse;font-size:12px;margin:4px 0 8px}"
        "th,td{padding:6px 9px;border-bottom:1px solid #eef2f7;text-align:right}"
        "th{color:#64748b;font-size:10px;border-bottom:1.5px solid #e2e8f0;background:#f8fafc}"
        "td:first-child,th:first-child,td.l,th.l{text-align:left}"
        ".mono{font-variant-numeric:tabular-nums}.muted{color:#64748b;font-size:11px}"
        ".big{color:#b91c1c;font-weight:700}.ok{color:#047857}.warn{color:#b45309}"
        "tfoot td{border-top:1.5px solid #cbd5e1;font-weight:700}"
        f".lead{{background:#f8fafc;border:1px solid #e2e8f0;border-left:4px solid {accent};"
        "border-radius:8px;padding:9px 13px;font-size:12px;margin-bottom:12px}"
        ".none{background:#f0fdf4;border:1px solid #bbf7d0;border-radius:8px;padding:9px 13px;"
        "color:#166534;font-size:12px}"
        ".chip{display:inline-block;padding:1px 8px;border-radius:999px;font-size:10px;font-weight:600}"
        ".c-ok{background:#ecfdf5;color:#047857}.c-warn{background:#fef3c7;color:#92610e}"
        ".c-err{background:#fef2f2;color:#b91c1c}.c-mut{background:#f1f5f9;color:#475569}"
        ".toolbar{max-width:960px;margin:0 auto 12px;text-align:right}"
        f".btn{{background:{accent};color:#fff;border:0;border-radius:8px;padding:9px 18px;"
        "font-size:13px;cursor:pointer}"
        ".ctl{padding:8px 10px;border:1px solid #cbd5e1;border-radius:8px;font:inherit;margin-right:8px}"
        ".sign{display:flex;gap:40px;margin-top:26px}"
        ".sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;"
        "font-size:11px;color:#475569}"
        ".kv{display:flex;flex-wrap:wrap;gap:6px 26px;margin:8px 0}"
        ".kv div{font-size:12px}.kv b{color:#0f172a}"
        "@media print{body{background:#fff;padding:0}.toolbar{display:none}"
        ".sheet{box-shadow:none;border-radius:0;max-width:none;margin:0}"
        ".pageb{page-break-before:always}}</style>")


def _ic_picker(period_month, branch_id, master, esc, all_label="ทุกสาขา (สาขาหลัก)",
               extra=""):
    bopts = (f"<option value=''{'' if branch_id else ' selected'}>{all_label}</option>"
             + "".join(f"<option value='{b['branch_id']}'"
                       f"{' selected' if branch_id == b['branch_id'] else ''}>"
                       f"{esc(b['name'])}</option>" for b in master))
    bsuffix = f"&branch_id={branch_id}" if branch_id else ""
    month_oc = f"location.href='?period_month='+this.value+'{bsuffix}{extra}'"
    sel_oc = (f"location.href='?period_month={period_month}'"
              f"+(this.value?'&branch_id='+this.value:'')+'{extra}'")
    return (f"<input type='month' value='{esc(period_month)}' class='ctl' onchange=\"{month_oc}\">"
            f"<select class='ctl' onchange=\"{sel_oc}\">" + bopts + "</select>")


@app.get("/api/reports/override-register.print")
def override_register_print(period_month: str, branch_id: int | None = None):
    """Print-to-PDF management-override register: every manual match / move / lane
    adjustment / day-acknowledgement in the month, with who·when·reason — the log a
    Big4 auditor pulls first. Rows with no reason are flagged."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v): return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _override_register(conn, period_month, branch_id)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc)
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>ทะเบียนการปรับปรุงด้วยมือ</title>", _ic_print_css("#4f46e5"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>ทะเบียนการปรับปรุงด้วยมือ (Management Override Register) — งวด {esc(period_month)}</h1>",
        f"<div class='sub'>ทุกรายการที่คนแก้ไขด้วยมือในเดือนนี้: จับคู่เอง · ย้ายช่อง · "
        f"รับทราบผลต่าง · ปรับปรุง lane — รวม <b>{reg['count']}</b> รายการ คิดเป็นเงิน "
        f"<b>฿{m(reg['total_abs_amount'])}</b>"
        + (f" · <span class='big'>{reg['missing_reason']} รายการไม่มีเหตุผลกำกับ</span>"
           if reg["missing_reason"] else " · ทุกรายการมีเหตุผลกำกับ ✓") + "</div>"]
    if reg["missing_reason"]:
        parts.append("<div class='lead'>⚠ รายการที่ไม่มีเหตุผล (เน้นสีแดง) คือจุดที่ผู้สอบบัญชี"
                     "จะตั้งคำถาม — ควรให้ผู้บันทึกกลับไปใส่เหตุผลทุกครั้งที่แก้ด้วยมือ</div>")
    if not reg["rows"]:
        parts.append("<div class='none'>✓ ไม่มีการปรับปรุงด้วยมือในงวดนี้ — ระบบจับคู่อัตโนมัติทั้งหมด</div>")
    else:
        parts.append("<table><thead><tr><th class='l'>วันเวลา</th><th class='l'>สาขา</th>"
            "<th class='l'>ประเภท</th><th class='l'>ช่องทาง / รายการ</th><th>จำนวน (บาท)</th>"
            "<th class='l'>โดย</th><th class='l'>เหตุผล</th></tr></thead><tbody>")
        for r in reg["rows"]:
            rsn = (esc(r["reason"]) if r["has_reason"]
                   else "<span class='big'>— ไม่ระบุเหตุผล —</span>")
            parts.append(f"<tr><td class='l muted'>{esc((r['at'] or '')[:16])}</td>"
                f"<td class='l'>{esc(r['branch_name'])}</td>"
                f"<td class='l'>{esc(r['kind'])}</td>"
                f"<td class='l'>{esc(r['lane'])}<div class='muted'>{esc(r['detail'])}</div></td>"
                f"<td class='mono'>{m(abs(r['amount']))}</td>"
                f"<td class='l'>{esc(r['by'])}</td>"
                f"<td class='l'>{rsn}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้สอบทาน (ผู้จัดการ)</div><div>วันที่</div></div>")
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


@app.get("/api/reports/open-items-aging.print")
def open_items_aging_print(period_month: str, branch_id: int | None = None,
                           lookback: int = 2):
    """Print-to-PDF aging of open reconciling items (งวดที่เลือก + ย้อนหลัง),
    bucketed 0-7 / 8-30 / >30 days from period-end. Stale items (>30d) flagged."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    lookback = max(0, min(6, lookback))
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        ag = _open_items_aging(conn, period_month, branch_id, lookback)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc, extra=f"&lookback={lookback}")
    b = ag["buckets"]
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>อายุรายการค้าง</title>", _ic_print_css("#b45309"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>อายุรายการค้าง (Aging of Open Items) — ณ {esc(ag['as_of'])}</h1>",
        f"<div class='sub'>รายการกระทบยอดที่ยังไม่ปิด นับอายุจากสิ้นงวด · "
        f"ตรวจงวด {esc(period_month)} ย้อนหลัง {lookback} เดือน "
        f"(ไม่รวมงวดที่อนุมัติ/ปิดแล้ว) · รวม <b>{ag['count']}</b> รายการ</div>",
        "<div class='kv'>"
        f"<div>0–7 วัน: <b class='ok'>{b['0-7']}</b></div>"
        f"<div>8–30 วัน: <b class='warn'>{b['8-30']}</b></div>"
        f"<div>เกิน 30 วัน: <b class='big'>{b['>30']}</b></div></div>"]
    if b[">30"]:
        parts.append("<div class='lead'>⚠ มีรายการค้างเกิน 30 วัน — ตามหลักการควบคุมภายใน "
                     "รายการกระทบยอดควรเคลียร์ภายในรอบถัดไป ของที่ค้างนานต้องมีคำอธิบาย/แผนปิด</div>")
    if not ag["items"]:
        parts.append("<div class='none'>✓ ไม่มีรายการค้าง — ทุกงวดที่ตรวจปิดครบหรืออยู่ในเกณฑ์</div>")
    else:
        parts.append("<table><thead><tr><th class='l'>อายุ (วัน)</th><th class='l'>ช่วง</th>"
            "<th class='l'>งวด</th><th class='l'>สาขา</th><th class='l'>ประเภท</th>"
            "<th class='l'>รายการที่ต้องทำ</th></tr></thead><tbody>")
        _bk_chip = {"0-7": "c-ok", "8-30": "c-warn", ">30": "c-err"}
        _ty_th = {"upload": "อัปโหลดเอกสาร", "setup": "ตั้งค่า", "diff": "ตรวจผลต่าง",
                  "waiting": "รอเอกสาร"}
        for it in ag["items"]:
            parts.append(f"<tr><td class='l mono'>{it['age_days']}</td>"
                f"<td class='l'><span class='chip {_bk_chip[it['bucket']]}'>{esc(it['bucket'])}</span></td>"
                f"<td class='l muted'>{esc(it['period_month'])}</td>"
                f"<td class='l'>{esc(it['branch_name'])}</td>"
                f"<td class='l'>{esc(_ty_th.get(it['type'], it['type']))}</td>"
                f"<td class='l'>{esc(it['label'])}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้จัดทำ</div><div>ผู้สอบทาน</div><div>วันที่</div></div>")
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


@app.get("/api/reports/recon-certificate.print")
def recon_certificate_print(period_month: str, branch_id: int | None = None):
    """Print-to-PDF reconciliation sign-off certificate — one signed sheet per branch:
    per-channel POS vs money-in, residual after acknowledgements, preparer/reviewer +
    approval date, frozen working-paper hash, override count. No branch_id = all
    master branches, each on its own page."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v): return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
        if branch_id:
            _validate_master_branch(conn, branch_id)
            targets = [branch_id]
        else:
            targets = [b["branch_id"] for b in master]
        certs = [_recon_certificate(conn, bid, period_month) for bid in targets]
    picker = _ic_picker(period_month, branch_id, master, esc,
                        all_label="ทุกสาขา (พิมพ์ทีละหน้า)")

    def sheet(c, first):
        sess = c["session"] or {}
        st = sess.get("status") or "draft"
        st_th = _IC_SESS_TH.get(st, st)
        st_chip = ("c-ok" if st in ("approved", "closed")
                   else "c-warn" if st == "in_review" else "c-mut")
        prep = esc(sess.get("created_by") or "—")
        rev = esc(sess.get("approved_by") or "—")
        appd = esc((sess.get("approved_at") or "")[:16] or "—")
        snap = c["snapshot"]
        snap_line = (f"<div>หลักฐานแช่แข็ง (SHA-256): <b class='mono'>{esc(snap['sha256'][:16])}…</b> "
                     f"· {esc((snap['created_at'] or '')[:16])} · "
                     f"ผูกเอกสารต้นทาง {snap.get('source_doc_count', 0)} ไฟล์</div>" if snap else
                     "<div class='muted'>ยังไม่มีหลักฐานแช่แข็ง (อนุมัติงวดเพื่อสร้าง snapshot)</div>")
        out = [f"<div class='sheet{'' if first else ' pageb'}'>",
            f"<h1>ใบรับรองการกระทบยอด (Reconciliation Certificate)</h1>",
            f"<div class='sub'>สาขา <b>{esc(c['branch_name'])}</b> · งวด <b>{esc(c['period_month'])}</b> "
            f"· สถานะ <span class='chip {st_chip}'>{esc(st_th)}</span> "
            f"· กระทบยอดแล้ว {c['progress_pct']}%</div>",
            "<div class='kv'>"
            f"<div>POS รวม: <b>฿{m(c['pos_total'])}</b></div>"
            f"<div>เงินเข้าจริงรวม: <b>฿{m(c['ext_total'])}</b></div>"
            f"<div>ผลต่างคงเหลือ (หลังรับทราบ): <b class='{'big' if abs(c['residual_total'])>100 else 'ok'}'>"
            f"฿{m(c['residual_total'])}</b></div>"
            f"<div>รายการรับทราบ: <b>{c['ack_count']}</b> (฿{m(c['ack_amount'])})</div>"
            f"<div>แก้ด้วยมือทั้งหมด: <b>{c['override_count']}</b> รายการ</div></div>",
            "<table><thead><tr><th class='l'>ช่องทาง</th><th>POS</th><th>เงินเข้าจริง</th>"
            "<th>ผลต่าง</th><th>คงเหลือ</th><th class='l'>สถานะ</th></tr></thead><tbody>"]
        for ch in c["channels"]:
            stx = "c-ok" if ch["status"] == "ok" else ("c-mut" if ch["status"] == "no_source"
                  else "c-warn" if ch["status"] == "waiting_doc" else "c-err")
            stl = ("รับทราบแล้ว" if ch["accepted"] else _IC_CHST_TH.get(ch["status"], ch["status"]))
            resid = ch["residual"]
            rcls = "" if abs(resid) <= 1 or ch["key"] == "COUPON" else (" big" if abs(resid) > 100 else " warn")
            rtxt = "—" if ch["key"] == "COUPON" else f"{resid:+,.0f}"
            out.append(f"<tr><td class='l'><b>{esc(ch['label'])}</b></td>"
                f"<td class='mono'>{m(ch['pos'])}</td><td class='mono'>{m(ch['ext'])}</td>"
                f"<td class='mono'>{ch['variance']:+,.0f}</td>"
                f"<td class='mono{rcls}'>{rtxt}</td>"
                f"<td class='l'><span class='chip {stx}'>{esc(stl)}</span></td></tr>")
        out.append(f"</tbody><tfoot><tr><td class='l'>รวม</td>"
            f"<td class='mono'>{m(c['pos_total'])}</td><td class='mono'>{m(c['ext_total'])}</td>"
            f"<td></td><td class='mono'>{c['residual_total']:+,.0f}</td><td></td></tr></tfoot></table>")
        out.append("<div class='lead'>" + snap_line + "</div>")
        out.append("<div class='sub' style='margin-top:14px'>ข้าพเจ้าขอรับรองว่าการกระทบยอด"
            "รายได้และเงินรับของสาขา/งวดข้างต้นถูกต้องและครบถ้วน ผลต่างที่คงเหลือได้รับการ"
            "ตรวจสอบและรับทราบตามเหตุผลที่บันทึกไว้ในทะเบียนการปรับปรุง</div>")
        out.append("<div class='sign'>"
            f"<div>ผู้จัดทำ<br><b>{prep}</b></div>"
            f"<div>ผู้สอบทาน/อนุมัติ<br><b>{rev}</b></div>"
            f"<div>วันที่อนุมัติ<br><b>{appd}</b></div></div>")
        out.append("</div>")
        return "".join(out)

    body = "".join(sheet(c, i == 0) for i, c in enumerate(certs)) or \
        "<div class='sheet'><div class='none'>ไม่มีข้อมูลสาขาสำหรับงวดนี้</div></div>"
    html_doc = ("<!doctype html><html lang='th'><head><meta charset='utf-8'>"
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>"
        "<title>ใบรับรองการกระทบยอด</title>" + _ic_print_css("#0F6E56") + "</head><body>"
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>"
        + body + "</body></html>")
    return Response(content=html_doc, media_type="text/html; charset=utf-8")


# ── #4 Close register + #5 Accepted-variance summary (audit evidence) ─────────

def _close_register(conn, period_month, months=6, branch_id=None):
    """Cross-branch × multi-month close monitoring: status (draft/in_review/approved/
    closed), %reconciled, todo count, approver + approval date, and whether the frozen
    working-paper snapshot exists. Read-only; reuses reconcile_overview per month."""
    months = max(1, min(12, months))
    win = list(dict.fromkeys(
        [period_month] + [_ic_prev_month(period_month, k) for k in range(1, months)]))
    ph = ",".join("?" * len(win))
    sess = {(r["branch_id"], r["period_month"]): dict(r) for r in conn.execute(
        "SELECT branch_id, period_month, approved_by, approved_at "
        f"FROM recon_sessions WHERE period_month IN ({ph})", win)}
    snaps = {(r["branch_id"], r["period_month"]) for r in conn.execute(
        f"SELECT branch_id, period_month FROM recon_close_snapshot WHERE period_month IN ({ph})", win)}
    rows, roll = [], {}
    for mth in win:
        ov = reconcile_overview(conn, mth)
        for b in ov["branches"]:
            if branch_id and b["branch_id"] != branch_id:
                continue
            s = sess.get((b["branch_id"], mth), {})
            st = b.get("status") or "draft"
            rows.append({"period_month": mth, "branch_id": b["branch_id"],
                "branch_name": b["branch_name"], "status": st,
                "progress_pct": b["progress_pct"], "todo_count": b["todo_count"],
                "approved_by": s.get("approved_by"), "approved_at": s.get("approved_at"),
                "has_snapshot": (b["branch_id"], mth) in snaps})
            r = roll.setdefault(mth, {"closed": 0, "approved": 0, "open": 0, "total": 0})
            r["total"] += 1
            r["closed" if st == "closed" else "approved" if st == "approved" else "open"] += 1
    return {"period_month": period_month, "months": win, "rows": rows,
            "month_rollup": roll,
            "open_total": sum(1 for r in rows if r["status"] not in ("approved", "closed"))}


def _accepted_variance_summary(conn, period_month, branch_id=None, threshold=1000.0):
    """Write-off view: every ACKNOWLEDGED variance (recon_lane_adj accept = whole channel,
    recon_day_ack = per day) for the month, grouped by branch and by reason, with the
    grand total accepted and items over a policy threshold flagged. Read-only."""
    from collections import defaultdict
    bmap = {r["branch_id"]: r["name"] for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY}")}

    def keep(bid):
        return bid in bmap and (branch_id is None or bid == branch_id)

    items, var_cache = [], {}
    for a in conn.execute(
            "SELECT branch_id, period_month, category, bank, note, created_by, created_at "
            "FROM recon_lane_adj WHERE adj_type='accept' AND period_month=?", (period_month,)):
        if not keep(a["branch_id"]):
            continue
        key = (a["branch_id"], a["period_month"])
        if key not in var_cache:
            try:
                var_cache[key] = {c["key"]: c.get("residual", c["variance"])
                                  for c in reconcile_month(conn, a["branch_id"], a["period_month"])["channels"]}
            except Exception:
                var_cache[key] = {}
        items.append({"branch_id": a["branch_id"], "branch_name": bmap[a["branch_id"]],
            "scope": _ic_lane_th(a["category"], a["bank"]), "channel": a["category"],
            "amount": round(var_cache[key].get(a["category"], 0.0), 2),
            "reason": a["note"] or "", "by": a["created_by"] or "—",
            "at": a["created_at"] or "", "source": "channel"})
    for a in conn.execute(
            "SELECT branch_id, category, day, variance, note, created_by, created_at "
            "FROM recon_day_ack WHERE period_month=?", (period_month,)):
        if not keep(a["branch_id"]):
            continue
        items.append({"branch_id": a["branch_id"], "branch_name": bmap[a["branch_id"]],
            "scope": f"{_CH_LABEL.get(a['category'], a['category'])} · {a['day']}",
            "channel": a["category"], "amount": round(float(a["variance"] or 0), 2),
            "reason": a["note"] or "", "by": a["created_by"] or "—",
            "at": a["created_at"] or "", "source": "day"})
    items.sort(key=lambda x: -abs(x["amount"]))
    by_branch = defaultdict(lambda: {"count": 0, "total_abs": 0.0})
    by_reason = defaultdict(lambda: {"count": 0, "total_abs": 0.0})
    for it in items:
        gb = by_branch[(it["branch_id"], it["branch_name"])]
        gb["count"] += 1
        gb["total_abs"] += abs(it["amount"])
        gr = by_reason[(it["reason"].strip() or "ไม่ระบุเหตุผล")]
        gr["count"] += 1
        gr["total_abs"] += abs(it["amount"])
    over = [it for it in items if abs(it["amount"]) > threshold]
    return {"period_month": period_month, "threshold": threshold,
        "items": items, "count": len(items),
        "grand_total_abs": round(sum(abs(it["amount"]) for it in items), 2),
        "by_branch": sorted([{"branch_id": k[0], "branch_name": k[1], "count": v["count"],
                              "total_abs": round(v["total_abs"], 2)} for k, v in by_branch.items()],
                            key=lambda x: -x["total_abs"]),
        "by_reason": sorted([{"reason": k, "count": v["count"],
                              "total_abs": round(v["total_abs"], 2)} for k, v in by_reason.items()],
                            key=lambda x: -x["total_abs"]),
        "over_threshold": over, "over_count": len(over)}


@app.get("/api/reports/close-register")
def close_register_api(period_month: str, months: int = 6, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _close_register(conn, period_month, months, branch_id)


@app.get("/api/reports/accepted-variance")
def accepted_variance_api(period_month: str, branch_id: int | None = None,
                          threshold: float = 1000.0):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _accepted_variance_summary(conn, period_month, branch_id, max(0.0, threshold))


@app.get("/api/reports/close-register.print")
def close_register_print(period_month: str, months: int = 6, branch_id: int | None = None):
    """Print-to-PDF close register: which branch-months are closed / approved / still
    open, who signed off + when, %reconciled, and whether the frozen snapshot exists."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    months = max(1, min(12, months))
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _close_register(conn, period_month, months, branch_id)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc, extra=f"&months={months}")
    st_chip = {"closed": "c-ok", "approved": "c-ok", "in_review": "c-warn",
               "draft": "c-mut"}
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>ทะเบียนการปิดเดือน</title>", _ic_print_css("#1d4ed8"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>ทะเบียนการปิดเดือน (Close Register) — {esc(months)} เดือนล่าสุดถึง {esc(period_month)}</h1>",
        f"<div class='sub'>สถานะการปิดงวดของทุกสาขาหลัก · ใครอนุมัติ · เมื่อไหร่ · กระทบยอดแล้วกี่ % · "
        f"มีหลักฐานแช่แข็งหรือยัง · <b>ยังไม่ปิดรวม {reg['open_total']}</b> สาขา-เดือน</div>"]
    if not reg["rows"]:
        parts.append("<div class='none'>ไม่มีข้อมูลสาขาสำหรับช่วงนี้</div>")
    for mth in reg["months"]:
        mr = reg["month_rollup"].get(mth)
        if not mr:
            continue
        parts.append(f"<h2>{esc(mth)} — ปิดแล้ว {mr['closed']} · อนุมัติ {mr['approved']} · "
                     f"<span class='warn'>ยังเปิด {mr['open']}</span> / {mr['total']} สาขา</h2>")
        parts.append("<table><thead><tr><th class='l'>สาขา</th><th class='l'>สถานะ</th>"
            "<th>กระทบยอด</th><th>คงค้าง</th><th class='l'>ผู้อนุมัติ</th>"
            "<th class='l'>วันที่อนุมัติ</th><th class='l'>หลักฐาน</th></tr></thead><tbody>")
        for r in [x for x in reg["rows"] if x["period_month"] == mth]:
            parts.append(f"<tr><td class='l'><b>{esc(r['branch_name'])}</b></td>"
                f"<td class='l'><span class='chip {st_chip.get(r['status'], 'c-mut')}'>"
                f"{esc(_IC_SESS_TH.get(r['status'], r['status']))}</span></td>"
                f"<td class='mono'>{r['progress_pct']}%</td>"
                f"<td class='mono'>{r['todo_count'] or ''}</td>"
                f"<td class='l'>{esc(r['approved_by'] or '—')}</td>"
                f"<td class='l muted'>{esc((r['approved_at'] or '')[:16] or '—')}</td>"
                f"<td class='l'>{'✔' if r['has_snapshot'] else '—'}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้สอบทาน</div><div>วันที่</div></div>")
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


@app.get("/api/reports/accepted-variance.print")
def accepted_variance_print(period_month: str, branch_id: int | None = None,
                            threshold: float = 1000.0):
    """Print-to-PDF accepted-variance (write-off) summary: total baht acknowledged,
    grouped by reason and by branch, with single items over the policy threshold flagged."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v): return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    threshold = max(0.0, threshold)
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        s = _accepted_variance_summary(conn, period_month, branch_id, threshold)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc,
                        extra=f"&threshold={int(threshold)}")
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>สรุปผลต่างที่รับทราบ</title>", _ic_print_css("#b91c1c"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>สรุปผลต่างที่รับทราบ (Accepted-Variance / Write-off) — งวด {esc(period_month)}</h1>",
        f"<div class='sub'>ยอดผลต่างที่ถูก “รับทราบ” แล้วปล่อยผ่าน (ทั้งช่อง + รายวัน) · "
        f"รวม <b class='big'>฿{m(s['grand_total_abs'])}</b> จาก <b>{s['count']}</b> รายการ · "
        f"เกินเกณฑ์ ฿{m(threshold)}/รายการ: <b class='big'>{s['over_count']}</b> รายการ</div>"]
    if not s["items"]:
        parts.append("<div class='none'>✓ ไม่มีผลต่างที่รับทราบในงวดนี้ — ทุกช่องทางตรงตามจริง</div>")
    else:
        if s["over_threshold"]:
            parts.append(f"<div class='lead'>⚠ มี {s['over_count']} รายการที่รับทราบเกิน "
                         f"฿{m(threshold)} ต่อรายการ — ผู้สอบบัญชีจะขอหลักฐาน/การอนุมัติเพิ่มเป็นพิเศษ</div>")
        parts.append("<h2>แยกตามเหตุผล</h2><table><thead><tr><th class='l'>เหตุผล</th>"
            "<th>จำนวนครั้ง</th><th>ยอดรวม (บาท)</th></tr></thead><tbody>")
        for r in s["by_reason"]:
            parts.append(f"<tr><td class='l'>{esc(r['reason'])}</td><td class='mono'>{r['count']}</td>"
                         f"<td class='mono'>{m(r['total_abs'])}</td></tr>")
        parts.append(f"</tbody><tfoot><tr><td class='l'>รวม</td><td class='mono'>{s['count']}</td>"
                     f"<td class='mono'>{m(s['grand_total_abs'])}</td></tr></tfoot></table>")
        parts.append("<h2>แยกตามสาขา</h2><table><thead><tr><th class='l'>สาขา</th>"
            "<th>จำนวนครั้ง</th><th>ยอดรวม (บาท)</th></tr></thead><tbody>")
        for r in s["by_branch"]:
            parts.append(f"<tr><td class='l'><b>{esc(r['branch_name'])}</b></td>"
                         f"<td class='mono'>{r['count']}</td><td class='mono'>{m(r['total_abs'])}</td></tr>")
        parts.append("</tbody></table>")
        parts.append("<h2>รายการทั้งหมด (เรียงยอดมากสุด)</h2><table><thead><tr>"
            "<th class='l'>สาขา</th><th class='l'>ช่องทาง / วัน</th><th>ยอด</th>"
            "<th class='l'>เหตุผล</th><th class='l'>โดย</th></tr></thead><tbody>")
        for it in s["items"]:
            big = " big" if abs(it["amount"]) > threshold else ""
            parts.append(f"<tr><td class='l'>{esc(it['branch_name'])}</td>"
                f"<td class='l'>{esc(it['scope'])}</td>"
                f"<td class='mono{big}'>{it['amount']:+,.0f}</td>"
                f"<td class='l'>{esc(it['reason']) or '<span class=\"big\">— ไม่ระบุ —</span>'}</td>"
                f"<td class='l'>{esc(it['by'])}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้จัดทำ</div><div>ผู้สอบทาน (ผู้จัดการ)</div><div>วันที่</div></div>")
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


# ── #6 Audit-trail viewer (change history) — READ-ONLY ────────────────────────
_AUDIT_ENTITY_TH = {"session": "งวด / สถานะ", "manual_match": "จับคู่ด้วยมือ",
    "channel_move": "ย้ายช่อง", "lane_adj": "ปรับ lane", "day_ack": "รับทราบรายวัน",
    "upload": "เอกสารอัปโหลด", "close_snapshot": "ปิดงวด / หลักฐาน"}
_AUDIT_ACTION_TH = {"create": "สร้าง", "update": "แก้ไข", "delete": "ลบ", "accept": "รับทราบ"}
# entity → (parent table, primary-key column) so an audit row can be resolved to a
# branch+period. Deleted parents drop out (their row is gone); the append-only audit
# log itself is never mutated.
_AUDIT_SRC = [("session", "recon_sessions", "session_id"),
              ("manual_match", "recon_manual_match", "match_id"),
              ("channel_move", "recon_channel_move", "id"),
              ("lane_adj", "recon_lane_adj", "id"),
              ("day_ack", "recon_day_ack", "id"),
              ("upload", "recon_uploads", "upload_id"),
              ("close_snapshot", "recon_close_snapshot", "id")]


def _audit_trail(conn, period_month, branch_id=None, limit=400):
    """Branch+month-scoped change history from the append-only audit log: who changed
    what, when, across sessions/overrides/uploads/snapshots. Read-only."""
    bmap = {r["branch_id"]: r["name"] for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY}")}
    rows = []
    for entity, tbl, pk in _AUDIT_SRC:
        try:
            cur = conn.execute(
                f"SELECT a.ts, a.action, a.field, a.old_val, a.new_val, a.\"by\" AS by_, "
                f"p.branch_id AS bid, p.period_month AS pm "
                f"FROM recon_audit_log a JOIN {tbl} p ON p.{pk}=a.entity_id "
                f"WHERE a.entity=? AND p.period_month=?", (entity, period_month))
        except sqlite3.OperationalError:
            continue
        for r in cur:
            bid = r["bid"]
            if bid not in bmap or (branch_id is not None and bid != branch_id):
                continue
            if r["action"] == "update" and (r["old_val"] is not None or r["new_val"] is not None):
                detail = f"{r['field'] or ''} {r['old_val'] or '—'} → {r['new_val'] or '—'}".strip()
            else:
                detail = r["new_val"] or r["field"] or ""
            rows.append({"ts": r["ts"], "by": r["by_"] or "—", "entity": entity,
                "entity_th": _AUDIT_ENTITY_TH.get(entity, entity),
                "action": r["action"], "action_th": _AUDIT_ACTION_TH.get(r["action"], r["action"]),
                "detail": detail, "branch_id": bid, "branch_name": bmap.get(bid, str(bid))})
    rows.sort(key=lambda x: (x["ts"] or ""), reverse=True)
    return {"period_month": period_month, "branch_id": branch_id,
            "count": len(rows), "rows": rows[:limit], "shown": min(len(rows), limit)}


@app.get("/api/reports/audit-trail")
def audit_trail_api(period_month: str, branch_id: int | None = None, limit: int = 400):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _audit_trail(conn, period_month, branch_id, max(1, min(2000, limit)))


@app.get("/api/reports/audit-trail.print")
def audit_trail_print(period_month: str, branch_id: int | None = None, limit: int = 400):
    """Print-to-PDF change history: every recorded change for the month, newest first."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        tr = _audit_trail(conn, period_month, branch_id, max(1, min(2000, limit)))
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc)
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>ประวัติการแก้ไข</title>", _ic_print_css("#0891b2"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>ประวัติการแก้ไข (Audit Trail) — งวด {esc(period_month)}</h1>",
        f"<div class='sub'>ทุกการเปลี่ยนแปลงที่ระบบบันทึกไว้ (เรียงใหม่สุดก่อน) · บันทึกแบบเพิ่มอย่างเดียว "
        f"แก้/ลบไม่ได้ · แสดง {tr['shown']} จาก {tr['count']} รายการ</div>"]
    if not tr["rows"]:
        parts.append("<div class='none'>ไม่มีการเปลี่ยนแปลงที่บันทึกไว้สำหรับงวดนี้</div>")
    else:
        parts.append("<table><thead><tr><th class='l'>เวลา</th><th class='l'>สาขา</th>"
            "<th class='l'>ประเภท</th><th class='l'>การกระทำ</th><th class='l'>รายละเอียด</th>"
            "<th class='l'>โดย</th></tr></thead><tbody>")
        for r in tr["rows"]:
            parts.append(f"<tr><td class='l muted'>{esc((r['ts'] or '')[:16])}</td>"
                f"<td class='l'>{esc(r['branch_name'])}</td><td class='l'>{esc(r['entity_th'])}</td>"
                f"<td class='l'>{esc(r['action_th'])}</td><td class='l'>{esc(r['detail'])}</td>"
                f"<td class='l'>{esc(r['by'])}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


# ── Excel (.xlsx) exports for the internal-control reports ────────────────────
def _xlsx_response(sheets, fname):
    """sheets = [(title, [headers...], [[cells...], ...]), ...] → styled .xlsx Response."""
    import io
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    first = True
    for title, headers, body in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        first = False
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for row in body:
            ws.append(row)
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{fname}\""})


@app.get("/api/reports/override-register.xlsx")
def override_register_xlsx(period_month: str, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _override_register(conn, period_month, branch_id)
    body = [[(r["at"] or "")[:16], r["branch_name"], r["kind"], r["lane"],
             round(abs(r["amount"]), 2), r["by"],
             r["reason"] if r["has_reason"] else "— ไม่ระบุเหตุผล —"] for r in reg["rows"]]
    return _xlsx_response(
        [("ทะเบียนการปรับปรุงด้วยมือ",
          ["วันเวลา", "สาขา", "ประเภท", "ช่องทาง/รายการ", "จำนวน", "โดย", "เหตุผล"], body)],
        f"override_register_{period_month}.xlsx")


@app.get("/api/reports/open-items-aging.xlsx")
def open_items_aging_xlsx(period_month: str, branch_id: int | None = None, lookback: int = 2):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        ag = _open_items_aging(conn, period_month, branch_id, max(0, min(6, lookback)))
    _ty = {"upload": "อัปโหลดเอกสาร", "setup": "ตั้งค่า", "diff": "ตรวจผลต่าง", "waiting": "รอเอกสาร"}
    body = [[it["age_days"], it["bucket"], it["period_month"], it["branch_name"],
             _ty.get(it["type"], it["type"]), it["label"]] for it in ag["items"]]
    return _xlsx_response(
        [("อายุรายการค้าง",
          ["อายุ(วัน)", "ช่วง", "งวด", "สาขา", "ประเภท", "รายการที่ต้องทำ"], body)],
        f"open_items_aging_{period_month}.xlsx")


@app.get("/api/reports/close-register.xlsx")
def close_register_xlsx(period_month: str, months: int = 6, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _close_register(conn, period_month, max(1, min(12, months)), branch_id)
    body = [[r["period_month"], r["branch_name"], _IC_SESS_TH.get(r["status"], r["status"]),
             r["progress_pct"], r["todo_count"], r["approved_by"] or "—",
             (r["approved_at"] or "")[:16], "มี" if r["has_snapshot"] else "—"] for r in reg["rows"]]
    return _xlsx_response(
        [("ทะเบียนการปิดเดือน",
          ["งวด", "สาขา", "สถานะ", "กระทบยอด%", "คงค้าง", "ผู้อนุมัติ", "วันที่อนุมัติ", "หลักฐาน"], body)],
        f"close_register_{period_month}.xlsx")


@app.get("/api/reports/accepted-variance.xlsx")
def accepted_variance_xlsx(period_month: str, branch_id: int | None = None,
                           threshold: float = 1000.0):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        s = _accepted_variance_summary(conn, period_month, branch_id, max(0.0, threshold))
    by_reason = [[r["reason"], r["count"], r["total_abs"]] for r in s["by_reason"]]
    by_branch = [[r["branch_name"], r["count"], r["total_abs"]] for r in s["by_branch"]]
    items = [[it["branch_name"], it["scope"], round(it["amount"], 2), it["reason"], it["by"]]
             for it in s["items"]]
    return _xlsx_response([
        ("แยกตามเหตุผล", ["เหตุผล", "จำนวนครั้ง", "ยอดรวม"], by_reason),
        ("แยกตามสาขา", ["สาขา", "จำนวนครั้ง", "ยอดรวม"], by_branch),
        ("รายการทั้งหมด", ["สาขา", "ช่องทาง/วัน", "ยอด", "เหตุผล", "โดย"], items)],
        f"accepted_variance_{period_month}.xlsx")


@app.get("/api/reports/audit-trail.xlsx")
def audit_trail_xlsx(period_month: str, branch_id: int | None = None, limit: int = 2000):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        tr = _audit_trail(conn, period_month, branch_id, max(1, min(5000, limit)))
    body = [[(r["ts"] or "")[:16], r["branch_name"], r["entity_th"], r["action_th"],
             r["detail"], r["by"]] for r in tr["rows"]]
    return _xlsx_response(
        [("ประวัติการแก้ไข", ["เวลา", "สาขา", "ประเภท", "การกระทำ", "รายละเอียด", "โดย"], body)],
        f"audit_trail_{period_month}.xlsx")


# ── Optional login enforcement (opt-in via RECON_ENFORCE_AUTH) ────────────────
# Recon shares the dashboard's signed session cookie (see _actor). When this flag is
# set, WRITE requests must carry a valid session — making the actor identity provable,
# not just best-effort. DEFAULT OFF: unset → unchanged (open/token-gated), so nothing
# breaks until the owner turns it on (login lives on the dashboard app). Reads stay open.
@app.middleware("http")
async def _enforce_recon_session(request: Request, call_next):
    if (os.environ.get("RECON_ENFORCE_AUTH", "").strip()
            and _parse_session is not None
            and request.method in ("POST", "PATCH", "PUT", "DELETE")
            and request.url.path.startswith("/api/")):
        try:
            ok = bool(_parse_session(request.cookies.get(_SESSION_COOKIE, "") or ""))
        except Exception:
            ok = False
        if not ok:
            from fastapi.responses import JSONResponse
            return JSONResponse(
                {"detail": "ต้องเข้าสู่ระบบก่อนทำรายการ (ล็อกอินที่หน้าหลัก)"}, status_code=401)
    return await call_next(request)


# ── UAT sign-off sheet (เอกสารทดสอบ + รับรองการใช้งาน) ─────────────────────────
# Printable acceptance-test checklist with a usage narrative beside each case, for the
# team to walk through, mark pass/fail, and sign off before going fully live.

_UAT_CASES = [
    {"id": "TC-01", "title": "หน้าแรก — กล่องงานทุกสาขา + แบนเนอร์เตือนรูปแบบน่าสงสัย",
     "role": "ทุกคน",
     "narrative": "หน้าแรกคืออินบ็อกซ์รวมทุกสาขา เรียงสาขาที่ยังมีงานค้างขึ้นก่อน บอกชัดว่า “เหลือกี่เรื่อง” "
                  "และมีแถบ “รอฉันตรวจ” สำหรับผู้จัดการ ถ้าระบบจับได้ว่าการ์ดเงินผิดช่องมีรูปแบบซ้ำผิดปกติ "
                  "(อาจเกิดจากรหัสในสเตทเมนต์ตั้งผิด ไม่ใช่คนกดผิด) จะมีแบนเนอร์ 🔎 เตือนผู้จัดการให้ไปตรวจพจนานุกรมรหัสก่อน",
     "steps": ["เปิด /app/ เลือกเดือนที่ทดสอบ (มุมขวาบน)",
               "ดูรายการสาขา — สาขาที่มีงานค้างขึ้นก่อน สาขาเรียบร้อยยุบใน “เรียบร้อยแล้ว”",
               "สังเกตแถบสรุปด้านบน (ความคืบหน้าทั้งเครือ / เงินที่อาจผิดช่อง)"],
     "expected": "เห็นรายการสาขาพร้อม “เหลือ N เรื่อง”, กดเข้าสาขาได้, (ถ้ามี) แบนเนอร์เตือนแสดงพร้อมลิงก์ไปพจนานุกรมรหัส"},
    {"id": "TC-02", "title": "หน้าสาขา — ช่องทางรับเงิน + อัปโหลดเอกสาร",
     "role": "เจ้าหน้าที่กระทบยอด",
     "narrative": "เข้าสาขาแล้วระบบแยกช่องทางเป็น “ต้องทำ” (ส้ม/แดง), “เศษเหรียญ ≤฿100” (รวบกดทีเดียว), "
                  "“รอเงินเข้าตามรอบ” (ปกติ) และ “เรียบร้อยแล้ว” (ยุบไว้) แถบเอกสารบอกว่าอัปครบหรือยัง — ป้ายแดง = ยังขาด "
                  "กดที่ป้ายเพื่ออัปได้เลย หรือ “⬆ หลายไฟล์ / 📁 ทั้งโฟลเดอร์” ให้ระบบแยกประเภทไฟล์เอง",
     "steps": ["กดเข้าสาขาที่ทดสอบ",
               "ที่แถบเอกสาร กดป้ายที่ยังขาด (แดง) เลือกไฟล์จริงของเดือนนั้น",
               "รอประมวลผล — ป้ายเปลี่ยนเป็น ✓ และช่องทางอัปเดตทันที",
               "ลอง “⬆ หลายไฟล์” เลือกหลายไฟล์พร้อมกัน"],
     "expected": "อัปโหลดสำเร็จ, ป้ายเป็น ✓ พร้อมช่วงวันที่, ช่องทางที่ตรงเลื่อนไป “เรียบร้อย”, ระบบแยกประเภทไฟล์ถูกต้อง"},
    {"id": "TC-03", "title": "การ์ด “เงินอาจถูกบันทึกผิดช่อง” (จับคู่ข้ามช่องทาง)",
     "role": "เจ้าหน้าที่กระทบยอด",
     "narrative": "เมื่อระบบเจอช่องทางหนึ่ง “เกินมา” และอีกช่องทาง “ขาดไป” เป็นยอดเท่ากันในวันใกล้กัน จะขึ้นการ์ดเล่าเป็นเรื่องว่า "
                  "“น่าจะเป็นเงินก้อนเดียวกัน — POS กดเป็นเงินสด แต่จริง ๆ เข้าเป็น QR” กด “✅ ใช่ เป็นเงินก้อนเดียวกัน” "
                  "ระบบย้ายเข้าช่องที่ถูก (ยอดรวมไม่หาย) ไม่แน่ใจกด “🔍 ขอดูรายการ” ดูบรรทัดจริงก่อน",
     "steps": ["หาการ์ด “⇄ เงินอาจถูกบันทึกผิดช่อง” (ถ้ามี)",
               "กด “🔍 ขอดูรายการ” ตรวจว่าเป็นเงินก้อนเดียวกันจริง",
               "ถ้าใช่ กด “✅ ใช่ เป็นเงินก้อนเดียวกัน”",
               "ตรวจว่ายอดย้ายเข้าช่องที่ถูก การ์ดหายไป"],
     "expected": "เห็นรายละเอียดเมื่อขอดู, ยืนยันแล้วยอดย้ายช่องถูก ยอดรวมวันไม่เปลี่ยน, บันทึกเป็น “การจับคู่ด้วยมือ”"},
    {"id": "TC-04", "title": "รับทราบผลต่าง (Accept Variance) พร้อมเหตุผล",
     "role": "เจ้าหน้าที่กระทบยอด",
     "narrative": "บางช่องทางมีผลต่างที่อธิบายได้ด้วย “สาเหตุจริง” (เช่น ผลต่างจากค่าธรรมเนียม/ปัดเศษที่กระทบยอดแล้ว) กด "
                  "“รับทราบผลต่าง” แล้วเลือก/พิมพ์เหตุผล (อย่างน้อย 10 ตัวอักษร — ต้องระบุสาเหตุจริง ไม่ใช่ “ยอดเล็กไม่คุ้ม”) "
                  "ระบบเคลียร์งานนั้นโดยไม่แก้ตัวเลข และเก็บหลักฐานว่าใครรับทราบ เพราะอะไร เมื่อไร",
     "steps": ["หาช่องทางที่มีผลต่างที่อธิบายสาเหตุได้",
               "กด “รับทราบผลต่าง” เลือก/พิมพ์เหตุผลที่ระบุสาเหตุจริง ยืนยัน",
               "ตรวจว่างานถูกเคลียร์ และมีบันทึกผู้รับทราบ + เวลา"],
     "expected": "เหตุผลสั้นเกิน (<10) ระบบเตือน, ครบแล้วเคลียร์ได้, ตัวเลขไม่เปลี่ยน, มีหลักฐาน"},
    {"id": "TC-05", "title": "ส่งตรวจ → ปิดเดือน (สองชั้น + รหัสผู้จัดการ)",
     "role": "เจ้าหน้าที่ (ส่งตรวจ) + ผู้จัดการ (ปิดเดือน)",
     "narrative": "การปิดเดือนแยกสองชั้น — เจ้าหน้าที่กด “ส่งตรวจ” เมื่อทำครบ ผู้จัดการตรวจสรุปแล้วกด “ปิดเดือน” "
                  "(ต้องใส่รหัสผู้จัดการ) ปิดแล้วตัวเลขล็อกถาวรแก้ไม่ได้ (เปิดใหม่ผ่าน “ส่งกลับแก้”) ป้องกันแก้ย้อนหลังไม่มีอนุมัติ",
     "steps": ["ทำสาขาให้เรียบร้อย กด “ส่งตรวจ” → สถานะ “รอผู้จัดการตรวจ”",
               "(ผู้จัดการ) เปิดสาขา เห็นสรุป กด “ปิดเดือน”",
               "ระบบถามรหัสผู้จัดการ — ผิด=ไม่ผ่าน, ถูก=ปิดได้",
               "ลองแก้หลังปิด — ต้องแก้ไม่ได้"],
     "expected": "ส่งตรวจได้, ปิดเดือนต้องใช้รหัสผู้จัดการ, ปิดแล้วล็อก, แสดงผู้ปิด + เวลา"},
    {"id": "TC-06", "title": "รายงานเงินบันทึกผิดช่อง + ใบเซ็นรับทราบ + แนวทางพัฒนาสาขา",
     "role": "ฝ่ายบัญชี / ผู้จัดการ",
     "narrative": "รายการที่ยืนยันว่ากดผิดช่อง (TC-03) ถูกเก็บเป็น log และออกเป็นรายงานให้สาขาเซ็นรับทราบ + แก้ไข มีส่วนช่วยพัฒนา: "
                  "กราฟแนวโน้มรายเดือน (ดีขึ้น/แย่ลง) + เป้าหมายเดือนหน้า, ช่องที่สับสนบ่อย, คำแนะนำเฉพาะจุด (ยอดเงินไม่หาย เป็นการลงผิดช่อง)",
     "steps": ["เมนู ⋯ เพิ่มเติม → “⇄ รายงานเงินบันทึกผิดช่อง”",
               "ตรวจสรุปต่อสาขา + การ์ดแนวทางพัฒนา (เป้า/ช่องสับสน/คำแนะนำ)",
               "กด “🖨️ ใบเซ็นรับทราบ (พิมพ์/PDF)” ตรวจ ① แนวโน้ม ② ช่องสับสน ③ คำแนะนำ + ข้อความรับทราบ + ช่องเซ็น 3 คน",
               "กด “⬇ Excel” + ลอง “พิมพ์เฉพาะสาขานี้”"],
     "expected": "รายงานแสดงรายการจริง, ใบพิมพ์ครบ 3 ส่วน + ช่องเซ็น, Excel โหลดได้, พิมพ์ PDF ได้"},
    {"id": "TC-07", "title": "พจนานุกรมรหัสในสเตทเมนต์ + ผู้จัดการแก้รหัสเอง",
     "role": "ฝ่ายบัญชี (ตรวจ) + ผู้จัดการ (แก้)",
     "narrative": "สเตทเมนต์มีรหัสย่อที่ระบบใช้แยกประเภทอัตโนมัติ (เช่น SDCH, IORSDT) บางรหัสอาจถูกตีความผิด พจนานุกรมให้ทีมตรวจว่า "
                  "แปลถูกไหม (รหัสที่ยังไม่ยืนยันขึ้น “⚠ รอตรวจ”) ถ้าผิด ผู้จัดการแก้เองได้ในตั้งค่า เลือกช่องทางที่ถูก "
                  "(มีตัวเลือก “ไม่นับ/ไม่ใช่รายได้”) มีผลทันทีกับข้อมูลเดิม + ไฟล์ใหม่ ไม่ต้องรอทีมพัฒนา",
     "steps": ["📊 รายงาน → “📖 พจนานุกรมรหัสในสเตทเมนต์” ตรวจความหมาย/ช่องทาง (พิมพ์ใบให้ทีมตรวจได้)",
               "⚙️ Admin → #setup → ส่วน ⑤ “รหัสในสเตทเมนต์ → ช่องทาง”",
               "เลือกรหัส (เช่น SDCH) เปลี่ยนช่องทาง กด “บันทึก” → ระบบถามรหัสผู้จัดการ",
               "ตรวจว่าขึ้น “✏️ ทีมแก้แล้ว” + จำนวนบรรทัดที่ปรับ, แล้วกด “คืนค่าเดิม”"],
     "expected": "ดูพจนานุกรมได้, แก้รหัสต้องใช้รหัสผู้จัดการ, แก้แล้วมีผล + นับบรรทัด, คืนค่าเดิมได้"},
    {"id": "TC-08", "title": "รายงานการจับคู่รายวันทั้งเดือน (ตรวจการจับคู่ + พิมพ์ PDF)",
     "role": "ฝ่ายบัญชี / ผู้ตรวจสอบ",
     "narrative": "เพื่อมั่นใจว่าจับคู่ถูกครบทั้งเดือน รายงานนี้ไล่ทุกวัน × ทุกช่องทางในใบเดียว แสดง POS เทียบเอกสาร, ผลต่าง, สถานะ "
                  "(✓ ตรง / ✗ ต่าง / ขาดเอกสาร / รอเอกสารตามรอบ) แถวที่ต้องตรวจขึ้นสีแดง ด้านบนมีสรุป “จับคู่ตรง X/Y” "
                  "ถ้าครบมีแบนเนอร์ “✓ ทุกช่องทางจับคู่ครบ” · ท้ายใบมี “💰 สรุปการรับเงินทั้งเดือน + ค่าธรรมเนียม/VAT + "
                  "📒 ร่างการบันทึกบัญชี (Dr/Cr)” และช่องเซ็น 3 คน · ปุ่ม “📑 ฉบับตรวจสอบ (ละเอียด)” เพิ่มหลักฐานระดับบรรทัด "
                  "(รหัส POS, ยอดบัตรราย acquirer, ทุกบรรทัดเงินเข้า, การจับคู่ด้วยมือ) สำหรับ auditor · "
                  "มี “📆 ช่องเลือกวัน (filter by date)” บนหัวรายงาน — เลือกวันเดียวแล้วรายงานย่อเหลือเฉพาะวันนั้น "
                  "พร้อมร่างบันทึกบัญชี (Dr/Cr) ของวันนั้นโดยตรง สำหรับนำไปลงบัญชีรายวัน",
     "steps": ["หน้าสาขา → กด “🖨️ ตรวจการจับคู่รายวัน”",
               "ตรวจ KPI (จับคู่ตรง/ต้องตรวจ/รอเอกสาร/POS รวม/เอกสารรวม/ผลต่างสะสม)",
               "ไล่ตาราง — ต่าง/ขาดเอกสารแดง, รอเอกสารเหลือง, ตรง=✓",
               "เลื่อนลงดู “💰 สรุปการรับเงิน + ค่าธรรมเนียม/VAT” และ “📒 ร่างบันทึกบัญชี” ว่ายอดตรง (Dr = Cr)",
               "เลือกวันเดียวที่ “📆 ช่องเลือกวัน” → รายงานเหลือเฉพาะวันนั้น + ร่างบันทึกบัญชี (Dr/Cr) ของวันนั้น แล้วกด “↺ ทั้งเดือน” เพื่อกลับ",
               "กด “📑 ฉบับตรวจสอบ (ละเอียด)” → ตรวจหลักฐานระดับบรรทัด (รหัส POS, ยอดบัตรราย acquirer, ทุกบรรทัดเงินเข้า)",
               "กด “🖨️ พิมพ์ / บันทึก PDF”; ลองเมนู ⋯ → “รายงานการจับคู่รายวัน (ทุกสาขา)”"],
     "expected": "ทุกวัน × ช่องทางสถานะถูกต้อง, KPI ตรงข้อมูล, สรุปรับเงิน/ค่าธรรมเนียม/VAT ถูก + ร่างบัญชี Dr=Cr, "
                 "เลือกวันแล้วย่อเหลือวันเดียว + ร่างบัญชีของวันนั้นถูกต้อง, ฉบับละเอียดมีหลักฐานครบทุกบรรทัด, พิมพ์ PDF ได้, แพ็กทุกสาขาแยกหน้า"},
    {"id": "TC-09", "title": "รายงานติดตามเอกสารที่ยังขาด (ตามงานรายเดือน)",
     "role": "ฝ่ายบัญชี / ผู้จัดการ",
     "narrative": "ใช้ไล่ตามเอกสารที่แต่ละสาขายังไม่อัป + ยอดเงินที่รอเอกสารนั้นจับคู่ · เลือกเดือน + สาขาได้บนหน้า → ใช้ตามงานทุกเดือน · "
                  "บัตรเครดิตแตกย่อยรายผู้รับบัตร (KBANK/KTC) พร้อมยอดที่ควรได้ + จำนวนรายการ → รู้เลยว่าต้องไปขอ settlement ของค่ายไหน",
     "steps": ["เมนู ⋯ เพิ่มเติม → “📥 ติดตามเอกสารที่ยังขาด”",
               "เลือกเดือน/สาขาจากตัวเลือกบนหน้า",
               "ดูแต่ละสาขา: เอกสารที่ต้องอัป + ยอดที่รอ + แตกย่อยราย acquirer (↳ KTC/กสิกร)",
               "กด “🖨️ พิมพ์ / บันทึก PDF” แจกทีมไปไล่อัป (มีช่อง ☐ ติ๊ก + ช่องเซ็น)"],
     "expected": "แสดงเอกสารที่ขาดต่อสาขา + ยอดรอจับคู่ถูกต้อง, แตกรายผู้รับบัตรพร้อมจำนวนรายการ, เลือกเดือน/สาขาได้, พิมพ์ PDF ได้"},
    {"id": "TC-10", "title": "รายงานหลักฐานรายวัน (ระดับบรรทัด) — เผื่อตรวจลึก",
     "role": "ฝ่ายบัญชี / ผู้ตรวจสอบ",
     "narrative": "เมื่ออยากดูบรรทัดจริงของวันใดวันหนึ่ง (รายการธนาคาร, ยอดบัตร, Grab/LM, ค่าธรรมเนียม+VAT, การจับคู่ด้วยมือพร้อมเหตุผล) "
                  "ใช้รายงานหลักฐานรายวัน — เชิงลึกที่สุดต่อ 1 วัน ใช้คู่กับ TC-08 (ภาพรวมชี้จุด → เจาะวันที่มีปัญหา)",
     "steps": ["เปิด …:8443/api/daily-report?day=YYYY-MM-DD (เว้น branch_id = ทุกสาขา, ใส่ &branch_id=NN = สาขาเดียว)",
               "ตรวจว่ามีรายการธนาคาร/บัตร/Grab/LM + การจับคู่ด้วยมือ + ค่าธรรมเนียม/VAT",
               "พิมพ์ PDF"],
     "expected": "เห็นหลักฐานระดับบรรทัดของวันนั้น, พิมพ์ได้"},
    {"id": "TC-11", "title": "สิทธิ์การเข้าถึง (รหัส admin / manager)",
     "role": "ผู้ดูแลระบบ / ผู้จัดการ",
     "narrative": "ระบบกันการแก้ไขโดยไม่ได้รับอนุญาต 2 ชั้น — รหัสบันทึก (admin) สำหรับอัปโหลด/บันทึก และรหัสผู้จัดการ "
                  "สำหรับปิดเดือน + แก้รหัสสเตทเมนต์ ทั้งคู่บังคับบนเซิร์ฟเวอร์จริง (เครื่องทดสอบเปิดให้สะดวก)",
     "steps": ["ลองอัปโหลด/บันทึกโดยยังไม่กรอกรหัส → ระบบถามรหัส (ผิด=ไม่ผ่าน, ถูก=ผ่าน)",
               "ลองปิดเดือน/แก้รหัสสเตทเมนต์โดยไม่มีรหัสผู้จัดการ → ไม่ผ่าน, ใส่ถูก → ผ่าน"],
     "expected": "รหัสผิดถูกปฏิเสธ, รหัสถูกทำงานได้, สิทธิ์ admin/manager แยกชัดเจน"},
    {"id": "TC-12", "title": "รูปแบบไฟล์ที่รองรับ — KTB (Excel/PDF/OCR) + Line Man Balance-Report",
     "role": "เจ้าหน้าที่กระทบยอด",
     "narrative": "สเตทเมนต์ KTB อัปได้ทั้ง Excel/CSV (แม่นที่สุด), PDF ที่มีตัวอักษร, หรือ PDF ที่ไม่มีตัวอักษร "
                  "(ไฟล์สแกน หรือไฟล์ที่ได้จาก “Print to PDF” ซึ่งแปลงตัวเลขเป็นรูป) — กรณีหลังระบบจะอ่านด้วย OCR "
                  "แล้ว “ตรวจยอดคงเหลือ” อัตโนมัติ (ยอดก่อนหน้า ± จำนวนเงิน = ยอดคงเหลือ) ถ้า OCR อ่านเลขผิดแม้ตัวเดียว "
                  "ห่วงโซ่ยอดจะพังและระบบ “ปฏิเสธ” ไม่บันทึกเลขที่ผิด · รายงาน Line Man แบบ Balance-Report (ไฟล์ Excel "
                  "หัวคอลัมน์อังกฤษ) ก็รองรับเป็นรายงาน settlement ที่ถูกต้องแล้ว (เดิมเคยถูกปฏิเสธ)",
     "steps": ["อัป KTB เป็น Excel/CSV (ถ้ามี) → ตรวจจำนวนรายการ/ยอดถูกต้อง",
               "อัป KTB เป็น PDF ที่ไม่มีตัวอักษร (เช่นไฟล์จาก Print to PDF) → ระบบขึ้น chip “🔍 อ่านด้วย OCR — โปรดตรวจยอดก่อนยืนยัน”",
               "เทียบยอดรวมในตัวอย่างกับยอดสรุปท้ายสเตทเมนต์ว่าตรงกัน ก่อนกดยืนยัน",
               "อัปไฟล์ Line Man “Balance-Report-*.xlsx” ในช่อง Line Man → ต้องอ่านได้ ไม่ถูกปฏิเสธ"],
     "expected": "KTB Excel/PDF อ่านได้ถูกต้อง, PDF ไม่มีตัวอักษรอ่านด้วย OCR + ขึ้น chip เตือน + ยอดผ่านการตรวจยอดคงเหลือ "
                 "(ถ้าเลขเพี้ยนจะถูกปฏิเสธพร้อมข้อความแนะนำให้ใช้ Excel/CSV), Line Man Balance-Report อัปได้"},
    {"id": "TC-13", "title": "Line Man — จ่ายตามยอดสะสม ฿500 + เดือนยอดน้อยตกไปวันที่ 1 เดือนถัดไป",
     "role": "เจ้าหน้าที่กระทบยอด / ฝ่ายบัญชี",
     "narrative": "LINE MAN ไม่ได้โอนรายวันแบบ Grab — เงินสุทธิจะ “สะสม” จนเกิน ฿500 แล้วจึงโอนเข้าบัญชี (วันเล็ก ๆ ต่ำกว่า "
                  "฿50 ก็ถูกนับสะสมด้วย) ถ้าทั้งเดือนยอดยังไม่ถึง ฿500 LINE MAN จะโอนยอดทั้งเดือนใน “วันที่ 1 ของเดือนถัดไป” "
                  "ระบบจะจับคู่เงินก้อนนั้นกลับมาให้เดือนที่เป็นเจ้าของยอดโดยอัตโนมัติ (payout trace ขึ้น ✓ ตรง ไม่ค้างเป็น "
                  "“รอเงินเข้า” ตลอด) และจะไม่นับเงินก้อนเดิมซ้ำในเดือนถัดไป (ไม่เกิดยอดเกินปลอม)",
     "steps": ["เลือกสาขา/เดือนที่ยอด LINE MAN ทั้งเดือนน้อย (< ฿500) — กางการ์ด LINEMAN ดู payout trace",
               "อัปสเตทเมนต์ KBANK ของ “เดือนถัดไป” ด้วย (เงินโอน LINE MAN จะอยู่ราววันที่ 1)",
               "ตรวจว่าเดือนนั้น payout trace ขึ้น ✓ ตรง ผูกกับเงินโอนวันที่ 1 ของเดือนถัดไป (ไม่ค้าง “รอเงินเข้า”)",
               "เปิดเดือนถัดไป — ตรวจว่าเงินก้อนวันที่ 1 ไม่ถูกนับเป็นยอดเกิน/รายได้ของเดือนใหม่"],
     "expected": "ยอดสะสม (รวมวันเล็ก ๆ) จับคู่กับเงินโอนก้อนเดียว, เดือนยอดน้อยผูกกับเงินโอนวันที่ 1 เดือนถัดไป = ✓ ตรง "
                 "(ไม่ค้างรอเงินเข้า), เดือนถัดไปไม่เกิดยอดเกินปลอมจากเงินก้อนเดิม"},
    {"id": "TC-14", "title": "ตรวจผลต่างรายวัน — รับทราบทีละวัน (กระทบยอดตามวันที่)",
     "role": "เจ้าหน้าที่กระทบยอด",
     "narrative": "ผลต่างของช่องทาง (เช่น เงินสด / Line Man) ต้องกระทบยอด “ตามวันที่” ไม่ใช่ยอดรวมทั้งก้อน · กางช่องทางที่ต่าง "
                  "จะมีตาราง “ตรวจผลต่างรายวัน” โชว์ทุกวันที่ต่าง กดปุ่ม “รับทราบวันนี้” ใส่เหตุผลแยกราย วัน (อย่างน้อย 10 ตัวอักษร) "
                  "เมื่อรับทราบครบทุกวัน ช่องทางนั้นจะเคลียร์ออกจาก “เหลือ N เรื่อง” เอง · เหตุผลรายวันติดไปในไฟล์ Excel ปิดเดือน",
     "steps": ["เปิดสาขาที่มีผลต่าง กดการ์ดช่องทางที่ขึ้น “ตรวจผลต่าง...”",
               "ดูตาราง “ตรวจผลต่างรายวัน” — แต่ละวันที่ต่างมีปุ่ม “รับทราบวันนี้”",
               "กดรับทราบทีละวัน ใส่เหตุผล (สั้นกว่า 10 ตัวอักษรต้องถูกปฏิเสธ)",
               "เมื่อครบทุกวัน ตรวจว่าช่องทางหายจาก “เหลือ N เรื่อง” และวันที่รับทราบขึ้น ✓"],
     "expected": "รับทราบแยกรายวันได้, เหตุผล <10 ตัวถูกปฏิเสธ, ครบทุกวัน → ช่องทางเคลียร์ปิดเดือนได้, "
                 "เหตุผลรายวันอยู่ในไฟล์ Excel ปิดเดือน (ชีต Daily acknowledged variances)"},
    {"id": "TC-15", "title": "รายงานการจับคู่รายวัน — สรุปบัญชีรายวัน (ยอดเต็ม/ส่วนลดแยก + รหัสบัญชี)",
     "role": "ฝ่ายบัญชี",
     "narrative": "รายงานการจับคู่รายวันมี “📅 สรุปบัญชีรายวัน” ไล่ทุกวัน × ช่องทาง บันทึกเป็น “ยอดเต็ม (gross)” โดยแยกบรรทัด "
                  "ส่วนลด / ค่าธรรมเนียม / VAT / สุทธิ — ใช้ลงบัญชีรายวันได้ · มีคอลัมน์ รหัสบัญชี / ชื่อบัญชี (เว้นว่างก่อน "
                  "เติมภายหลังผ่าน /api/account-map) · ส่วน “💰 สรุปการรับเงินทั้งเดือน” (ยอดเต็ม·ส่วนลด·ค่าธรรมเนียม·VAT·สุทธิ + "
                  "ร่างบันทึกบัญชี Dr/Cr) ยังอยู่ครบเป็นสรุปทั้งเดือน · บรรทัด “ส่วนลด” รวม <b>ส่วนลดบนบิล (เด็ก/ผู้สูงอายุ/วันเกิด ฯลฯ)</b> "
                  "จาก pos_coupon_day ด้วย — ไม่ใช่แค่ส่วนลดที่จ่ายผ่านคูปอง (แต้ม KTC/BAY ไม่นับซ้ำ)",
     "steps": ["เปิด ⋯ เพิ่มเติม → รายงานการจับคู่รายวัน (หรือหน้าสาขา → 🖨️)",
               "ตรวจส่วน “📅 สรุปบัญชีรายวัน” — มีคอลัมน์ ยอดเต็ม / ส่วนลด แยกกัน + รหัสบัญชี/ชื่อบัญชี",
               "สาขาที่ใช้คูปองเด็ก/ผู้สูงอายุเยอะ (เช่น อุดรธานี): ตรวจว่าบรรทัด “ส่วนลด” มียอด ไม่ใช่ 0",
               "(ผู้จัดการ) ตั้งรหัสบัญชีผ่าน /api/account-map แล้วรีเฟรช — รหัส/ชื่อบัญชีต้องขึ้นในรายงาน",
               "ตรวจ “💰 สรุปการรับเงินทั้งเดือน” ว่ามีคอลัมน์ส่วนลด และร่างบัญชี Dr=Cr"],
     "expected": "มีสรุปบัญชีรายวัน (gross + ส่วนลดแยก), ส่วนลดบนบิล (เด็ก/ผู้สูงอายุ) ปรากฏในบรรทัดส่วนลด (ไม่เป็น 0 ผิดปกติ), "
                 "คอลัมน์รหัส/ชื่อบัญชีว่างจนกว่าจะตั้งค่า แล้วขึ้นตามที่ตั้ง, สรุปทั้งเดือนยังครบ + ร่างบัญชีสมดุล"},
    {"id": "TC-16", "title": "รายงานข้อผิดพลาดประจำเดือน (ทุกประเภท ไม่ใช่แค่บันทึกผิดช่อง)",
     "role": "ฝ่ายบัญชี / ผู้จัดการ",
     "narrative": "รายงานเดิม “เงินบันทึกผิดช่อง” ขยายเป็น “ข้อผิดพลาดประจำเดือน” รวมข้อผิดพลาดทุกชนิดของเดือน: ยอดต่าง, "
                  "เงินเข้าเกิน (ไม่มี POS), ขาดเอกสาร, เงินโอนเข้าไม่ครบ/ช้า, รหัสยังไม่ตั้งธนาคาร และบันทึกผิดช่อง — แต่ละแถวมี "
                  "ประเภท + สถานะ (ค้าง/รับทราบแล้ว/แก้ไขแล้ว) + เหตุผล · รายละเอียดบันทึกผิดช่อง + ใบเซ็นยังอยู่ครบเป็นส่วนท้าย",
     "steps": ["เปิด ⋯ เพิ่มเติม → รายงานข้อผิดพลาดประจำเดือน (เลือกเดือน)",
               "ตรวจตาราง “ข้อผิดพลาดทั้งหมด” — มีหลายประเภท ไม่ใช่แค่บันทึกผิดช่อง + สถานะ ค้าง/รับทราบ/แก้ไข",
               "เลือกสาขาจากดรอปดาวน์ — ตารางกรองเฉพาะสาขานั้น",
               "กด ⬇ Excel (แผ่นแรก “ข้อผิดพลาดทั้งหมด”) และ 🖨️ พิมพ์/PDF"],
     "expected": "เห็นข้อผิดพลาดทุกประเภทของเดือน + สถานะ + เหตุผล, กรองรายสาขาได้, Excel/พิมพ์ออกได้ "
                 "(แผ่น/หัวข้อ “ข้อผิดพลาดทั้งหมด”), รายละเอียดบันทึกผิดช่องยังอยู่ครบ"},
    {"id": "TC-17", "title": "สรุปบัญชีรายวัน (รายใบ) — แต่ละวัน = 1 การ์ด + เลือกสาขา + พิมพ์ PDF",
     "role": "ฝ่ายบัญชี",
     "narrative": "ส่วน “📅 สรุปบัญชีรายวัน” (ยอดเต็ม/ส่วนลดแยกบรรทัด) ถูกดึงออกมาเป็น “หน้าแยกของตัวเอง” และแสดงเป็น "
                  "“การ์ดรายวัน — 1 วัน = 1 ใบบันทึกบัญชี” ดูทีละสาขา (เลือกจากดรอปดาวน์) · ด้านบนมีแถบสรุปทั้งเดือน "
                  "(ยอดเต็มรวม·ส่วนลด·ค่าธรรมเนียม·VAT·สุทธิ·จำนวนใบ) · แต่ละการ์ดไล่ทุกช่องทางของวันนั้นเป็นยอดเต็ม "
                  "แยกบรรทัดส่วนลด/ค่าธรรมเนียม/VAT → สุทธิ + “รวมวันนี้” · รหัสบัญชีมาจากที่ตั้งใน ตั้งค่าระบบ · "
                  "แต่ละการ์ดมีปุ่มพับ “📒 ร่างบันทึกบัญชีวันนี้ (Dr/Cr)” กางออกดูคู่เดบิต/เครดิตของวันนั้น (มีในฉบับพิมพ์ PDF ด้วย) · "
                  "ตัวเลขชุดเดียวกับ TC-15 (จอ ↔ กระดาษตรงกัน) มีปุ่มพิมพ์ PDF เป็นแผ่นการ์ด",
     "steps": ["เปิด ⋯ เพิ่มเติม → “📅 สรุปบัญชีรายวัน (รายใบ)”",
               "เลือกสาขาจากดรอปดาวน์มุมขวา — การ์ดเปลี่ยนตามสาขา (สาขาที่ไม่มีรายการขึ้นข้อความว่าง)",
               "ตรวจการ์ดวันหนึ่ง: มี ยอดเต็ม / ส่วนลด (แยกบรรทัด) / ค่าธรรมเนียม / VAT / สุทธิ + แถว “รวมวันนี้”",
               "กางปุ่ม “📒 ร่างบันทึกบัญชีวันนี้ (Dr/Cr)” → เดบิต (เงินสด/ธนาคารสุทธิ·ส่วนลด·ค่าธรรมเนียม·VAT) = เครดิต (รายได้เต็มยอด)",
               "เทียบยอดรวมทั้งเดือนด้านบนกับ §สรุปบัญชีรายวัน ในรายงานจับคู่ (TC-15) ว่าตรงกัน",
               "กด “🖨️ พิมพ์/PDF” → ออกเป็นแผ่นการ์ดรายวัน (เลือกสาขาเดียว หรือเว้น = ทุกสาขา หน้าละสาขา) มีร่างบัญชี Dr/Cr ในทุกการ์ด"],
     "expected": "หน้าแสดงการ์ด 1 วัน/ใบ ครบทุกวันที่มีรายการ, ยอดเต็ม/ส่วนลดแยกบรรทัดถูกต้อง + รวมวันนี้ + รวมทั้งเดือน, "
                 "ร่างบันทึกบัญชี Dr/Cr ต่อวันสมดุล (Dr=Cr), เลือกสาขาได้, รหัสบัญชีขึ้นตามที่ตั้ง, ตัวเลขตรงกับ TC-15, พิมพ์ PDF เป็นการ์ดได้"},
    {"id": "TC-18", "title": "บันทึกผิดช่อง (ย้ายช่อง) — กระทบสองยอดพร้อมกัน",
     "role": "เจ้าหน้าที่กระทบยอด / ผู้จัดการ",
     "narrative": "ถ้ายอดถูกบันทึกผิดช่อง (เช่น ลูกค้าสแกน QR แต่พนักงานกดเป็นเงินสด) เจ้าหน้าที่ระบุ “ย้ายจากช่องผิด → ช่องถูก + "
                  "จำนวนเงิน” โปรแกรมจะลดยอดช่องผิดและเพิ่มยอดช่องถูกพร้อมกันทันที (กระทบสองยอด) โดยไม่ต้องรันจับคู่ใหม่ · "
                  "ย้อนกลับได้ (กด “↩ ยกเลิก”) · ทุกการย้ายเข้าไปอยู่ในรายงานเงินบันทึกผิดช่อง (TC-06) ติดป้าย “ย้ายด้วยมือ”",
     "steps": ["หน้าสาขา → การ์ด “⇄ บันทึกผิดช่อง (ย้ายช่อง)” → กด “+ บันทึกผิดช่อง”",
               "เลือก จากช่อง (ผิด) → ไปช่อง (ถูก) + ใส่จำนวนเงิน + เหตุผล กด “บันทึกการย้าย”",
               "ตรวจว่ายอดทั้งสองช่องเปลี่ยนพร้อมกันในตารางสรุป + รายงานจับคู่รายวัน + Dashboard",
               "กด “↩ ยกเลิก” → ยอดกลับไปช่องเดิม",
               "เปิด ⋯ → “⇄ รายงานเงินบันทึกผิดช่อง” ตรวจว่ารายการย้ายขึ้นพร้อมป้าย “ย้ายด้วยมือ”"],
     "expected": "ย้ายแล้วช่องผิดลด+ช่องถูกเพิ่มเท่ากัน, รายวันตรงกับรายเดือน (ไม่ทำให้ผลต่างเพี้ยน), ยกเลิกคืนค่าได้, "
                 "เข้าช่องเดียวกัน/จำนวน≤0 ถูกปฏิเสธ, ขึ้นในรายงานพร้อมป้ายย้ายด้วยมือ"},
    {"id": "TC-19", "title": "Dashboard รายได้ตามช่องทางรายวัน (ทุกสาขาในวันเดียว · เลือกวันปฏิทิน)",
     "role": "ผู้จัดการ / ฝ่ายบัญชี",
     "narrative": "หน้า Dashboard ใหม่: เลือกวันจากปฏิทิน → เห็นรายได้แยกช่องทางของ “ทุกสาขา” ในวันนั้นพร้อมกัน + แถบยอดรวมทุกสาขา · "
                  "มีปุ่มสลับ 3 โหมด: (1) เงินเข้าจริง(บัญชี) — ยอดเต็ม/ส่วนลด/สุทธิ, แยกคูปอง KTC/BAY(แต้ม)=เงินรับจริง; "
                  "(2) ยอด POS — ยอดที่เครื่อง POS บันทึกแยกช่องทาง; (3) ประกบ POS↔เงินเข้า — แต่ละช่องทางแสดง POS | เงินเข้าจริง | ต่าง "
                  "(เงินเข้า = ยอดเข้าบัญชี/ค่ายบัตรตาม statement·settlement). สลับโหมดไม่ต้องโหลดใหม่",
     "steps": ["เปิด ⋯ เพิ่มเติม → “📊 รายได้ตามช่องทาง รายวัน (Dashboard)”",
               "เลือกวันจากปฏิทินมุมขวา — การ์ดทุกสาขาปรับตามวันนั้น",
               "โหมด “เงินเข้าจริง(บัญชี)”: ตรวจ ช่องทาง/ยอดเต็ม/ส่วนลด/สุทธิ + แถว “รวมสาขา” + แถบยอดรวม",
               "กดปุ่ม “ยอด POS”: เห็นยอดที่ POS บันทึกแยกช่องทาง",
               "กดปุ่ม “ประกบ POS ↔ เงินเข้า”: เห็น POS | เงินเข้า | ต่าง (ช่องยังไม่อัปเอกสาร เงินเข้า=0 → ต่างติดลบเท่า POS)",
               "ลองเลือกวันที่ไม่มีข้อมูล → ขึ้นข้อความว่าง ไม่ error"],
     "expected": "เลือกวันได้, สลับ 3 โหมดได้ไม่โหลดใหม่, ทุกสาขา (เฉพาะ master) แสดงแยกการ์ด, ยอดรวมเปลี่ยนตามโหมด, "
                 "COUPON = แต้มเท่านั้น + บรรทัดส่วนลดแยก, ประกบ POS↔เงินเข้า + ผลต่างถูกต้อง, วันไม่มีข้อมูลไม่ค้าง"},
    {"id": "TC-20", "title": "บัญชีรายได้ — ยอดเต็มก่อนหักส่วนลด + ส่วนลดจากคูปอง (ไม่รวม KTC/BAY)",
     "role": "ฝ่ายบัญชี",
     "narrative": "การบันทึกบัญชีลง “ยอดเต็มก่อนหักคูปองส่วนลด” แล้วแยกบรรทัด “ส่วนลด” เป็นยอดหัก · ส่วนลดมาจากคูปองส่วนลด "
                  "(ผู้สูงอายุ/เด็ก/อื่น ๆ) เท่านั้น — คูปอง KTC/BAY ที่ใช้แต้มถือเป็น “เงินรับจริง” (ค่ายบัตรจ่ายคืน) ไม่ใช่ส่วนลด · "
                  "เห็นได้ทั้งในหน้าสรุปบัญชีรายวัน (TC-17), รายงานจับคู่รายวัน (§สรุปการรับเงิน + ร่างบัญชี), และ Dashboard (TC-19)",
     "steps": ["เปิด “📅 สรุปบัญชีรายวัน (รายใบ)” ของสาขาที่มีคูปอง",
               "ตรวจว่าช่อง COUPON = เฉพาะแต้ม KTC/BAY (เงินรับจริง) + มีบรรทัด “ส่วนลด (คูปอง)” แยก = ยอดหัก",
               "ตรวจว่า ยอดเต็ม (gross) = ยอดก่อนหักส่วนลด, สุทธิ = ยอดเต็ม − ส่วนลด − ค่าธรรมเนียม − VAT",
               "เปิดรายงานจับคู่รายวัน → §“💰 สรุปการรับเงิน” + “📒 ร่างบันทึกบัญชี”: รายได้ (Cr) = ยอดเต็ม, ส่วนลด (Dr) รวมคูปองส่วนลด, Dr = Cr"],
     "expected": "ยอดเต็มไม่หักส่วนลด, ส่วนลดแยกบรรทัดมาจากคูปองส่วนลด (ไม่รวม KTC/BAY), KTC/BAY อยู่ในเงินรับจริง, สุทธิถูกต้อง, ร่างบัญชี Dr = Cr"},
    # -- Phase 1: diagnostic views migrated from classic.html --
    {"id": "TC-21", "title": "Daily Board (#board) -- ตารางทุกสาขาในวันเดียว",
     "role": "ผู้จัดการ / ฝ่ายบัญชี",
     "narrative": "Daily Board คือมุมมอง 'ภาพรวมทั้งเครือในวันเดียว' -- เปิดใน index.html เป็น route #board "
                  "แทนที่จะต้องเปิด classic.html แยก - เห็นทุกสาขา x ทุกช่องทาง สถานะจับคู่ (สีเขียว/แดง/เหลือง) "
                  "พร้อมยอดรวมด้านบน + ปุ่ม PDF ทุกสาขา - ใช้ตรวจสอบประจำวันว่าเงินเข้าครบทุกสาขาหรือยัง",
     "steps": ["เปิด [เพิ่มเติม] > วินิจฉัย / Diagnostics > Daily Board",
               "ตรวจปฏิทินวันที่ขึ้นค่าล่าสุดอัตโนมัติ -- กด [<] [>] หรือเลือกวันจาก date picker",
               "ตรวจตารางสาขา x ช่องทาง: เซลล์สีเขียว = ตรง, แดง = ต่าง, เหลือง = ขาดเอกสาร",
               "ตรวจแถบสรุปด้านบน: POS รวม / เงินเข้า / ผลต่าง / จำนวนสาขา",
               "กด [PDF ทุกสาขา] -> เปิดรายงานหลักฐานรายวันในแท็บใหม่",
               "ลองวันที่ไม่มีข้อมูล -> ขึ้นข้อความว่างไม่ error"],
     "expected": "เลือกวันได้ (picker + [<][>]), ตารางแสดงทุก master branch x ช่องทาง, สีเซลล์ตรงสถานะ, "
                 "ยอดรวม footer ถูกต้อง, PDF ลิงก์ทำงาน, วันไม่มีข้อมูลขึ้นข้อความว่างไม่ค้าง, "
                 "ไม่ใช้ Tailwind CDN / classic.html CSS"},
    {"id": "TC-22", "title": "Audit Log (#audit) -- ประวัติการแก้ไขระบบ",
     "role": "ผู้จัดการ / ผู้ตรวจสอบ",
     "narrative": "Audit log บันทึกทุกการเปลี่ยนแปลงในระบบโดยอัตโนมัติ -- ใครแก้อะไร เมื่อไร ค่าเดิม -> ใหม่ "
                  "ใช้สืบสาวต้นเหตุเมื่อตัวเลขเปลี่ยนโดยไม่คาดคิด หรือตรวจว่าใครปิดเดือน/แก้รหัส",
     "steps": ["เปิด [เพิ่มเติม] > วินิจฉัย > Audit Log",
               "ตรวจตารางแสดงคอลัมน์: เวลา / entity / ID / action / field / ค่าเก่า->ใหม่ / โดย",
               "กรอง entity จาก dropdown (เช่น session) -> ตารางกรองได้ถูกต้อง",
               "กลับเป็น ทุก entity -> รายการครบ"],
     "expected": "แสดง 200 รายการล่าสุด, กรอง entity ได้, คอลัมน์ครบ 7 ช่อง, ค่าเดิม->ใหม่อ่านได้ชัด, "
                 "ไม่โหลด Tailwind / classic CSS"},
    {"id": "TC-23", "title": "Trend Heatmap (#trend) -- ปฏิทินสีความคืบหน้ารายวัน (session-scoped)",
     "role": "ผู้จัดการ / เจ้าหน้าที่กระทบยอด",
     "narrative": "Trend heatmap แสดงสีปฏิทินรายวันตาม % lane match ของสาขาที่เลือก -- เห็นแนวโน้มการกระทบยอด "
                  "ทั้งเดือนในแผนที่สีเดียว - session-scoped: ต้องมี session (เปิดหน้าสาขาแล้ว) ถึงจะมีข้อมูล",
     "steps": ["เปิด [เพิ่มเติม] > วินิจฉัย > Trend Heatmap",
               "เลือกสาขาจาก dropdown -- ระบบดึง session ของเดือนนั้นให้อัตโนมัติ (ไม่สร้าง session ใหม่)",
               "ถ้าสาขานั้นยังไม่มี session -> ขึ้นข้อความ 'ยังไม่มี session' + ลิงก์เปิดหน้าสาขา",
               "ถ้ามี session -> ตรวจตารางสรุปรายสัปดาห์ (W1-W4): POS / เงินเข้า / % lane / % cov",
               "ตรวจปฏิทินรายวัน: สีเขียว >=99% - เหลือง 60-98% - แดง <60% - เลขในเซลล์ = % + matched/total"],
     "expected": "dropdown branch ทำงานได้, ไม่มี session -> empty state ชัดเจน + ลิงก์หน้าสาขา, "
                 "มี session -> ตารางสัปดาห์ + ปฏิทินสีถูกต้อง, สีตรงตามเกณฑ์ 3 ระดับ, "
                 "เฉพาะ master branch ในรายการ, ไม่สร้าง session ใหม่"},
    {"id": "TC-26", "title": "Payout Matching (#payout) -- ตรวจเงินเข้าธนาคารทุกช่องทาง (เฟส 3a)",
     "role": "ฝ่ายบัญชี / ผู้จัดการ",
     "narrative": "หน้า #payout แสดงการจับคู่ยอดสุทธิรายวัน <-> เงินโอนเข้าบัญชี แยกตามช่องทาง "
                  "(บัตรเครดิต / Grab / Line Man / BAY Point / KTC Point) สถานะ matched/รอเงินเข้า "
                  "พร้อมวันขาย ยอดสุทธิ เงินโอน วันที่โอน ล่าช้ากี่วัน รวมวันอัตโนมัติ "
                  "KTC Point แสดงแบบ bank_in_credit (รวมในยอดบัตร) "
                  "เฟส 3a consolidation 2026-06-19 แทนที่ openPayoutMatch ใน classic.html",
     "steps": ["เปิดหน้าสาขา -> กดปุ่ม เงินเข้าธนาคาร (header สาขา)",
               "ตรวจยอดสรุป: ยอดสุทธิรวม / เข้าบัญชีแล้ว / รอเงินเข้า",
               "ตรวจการ์ด Grab -- วันขาย / ยอดสุทธิ / เงินโอน / วันที่โอน + chip สถานะ",
               "ตรวจการ์ด KTC Point -- ขึ้นป้าย รวมในยอดบัตร + ต่าง POS vs ยอดสุทธิ",
               "เปิดจากเมนู เพิ่มเติม -> บัญชี / งานสาขา -> ได้หน้าเดียวกัน",
               "สาขาไม่มี payout data -> ขึ้น empty state ไม่ error"],
     "expected": "เปิดได้จากปุ่มในหน้าสาขา + เมนู, ยอดสรุปครบ, การ์ดแต่ละช่องทางแสดงถูกต้อง, "
                 "KTC แสดงแบบ bank_in_credit, สาขาไม่มีข้อมูลขึ้น empty state ไม่ error, "
                 "ไม่ใช้ Tailwind / classic CSS"},
    {"id": "TC-30", "title": "classic.html ถูกยกเลิก -- ตรวจรับ parity สุดท้าย (เฟส 6)",
     "role": "ทุกบทบาท",
     "narrative": "เฟส 6 ลบหน้า classic.html (เครื่องมือขั้นสูงหน้าเดิม) ออกจากระบบ — "
                  "ทุกฟีเจอร์ที่เคยมีเฉพาะใน classic ถูกย้ายเข้าแอปเดียวครบในเฟส 1-5 แล้ว "
                  "และเข้าถึงได้จากแถบนำทาง 4 โซน ไม่มีหน้าจอกำพร้า/ลิงก์ตาย "
                  "(เฟส 6 consolidation 2026-06-19 — ปิดงานรวมแอป)",
     "steps": ["เปิด /app/classic.html ตรง ๆ -> ต้องได้ 404 (ไฟล์ถูกลบแล้ว)",
               "ตรวจว่าไม่มีปุ่ม/ลิงก์ 'เครื่องมือขั้นสูง' ใน #reports hub หรือที่ใดในแอป",
               "ไล่เข้าทุกฟีเจอร์ classic เดิมจากนำทาง: #board #trend #audit #bankcov (วินิจฉัย/Admin), "
               "#coupon (รายงาน), #payout #discrepancies #setup (Admin), อัปโหลดเอกสาร+จับคู่เอง (หน้าสาขา)",
               "เปิดคู่มือ /app/help.html -> footer ไม่มีลิงก์ classic แล้ว; Q13 อธิบายว่า classic ถูกยกเลิก",
               "ยืนยัน 130 tests เขียว + canary อุดร 2026-05-01 ส่วนลด ฿35,511.48 ไม่เปลี่ยน"],
     "expected": "classic.html -> 404, ไม่มีลิงก์ classic หลงเหลือ, ทุกฟีเจอร์เดิมเข้าถึงได้จากนำทาง 4 โซน, "
                 "ไม่มีหน้าจอกำพร้า, คู่มืออัปเดต, เทสต์เขียว, ตัวเลขบัญชีไม่ถดถอย"},
    {"id": "TC-31", "title": "รายงานเพื่อการควบคุมภายใน / ถูกตรวจสอบ (Audit & Controls)",
     "role": "ผู้จัดการ / ผู้สอบบัญชี",
     "narrative": "กลุ่มรายงานใหม่ในหน้า 📊 รายงาน สำหรับการควบคุมภายในและการถูกตรวจโดยผู้สอบบัญชี "
                  "ทุกตัวอ่านอย่างเดียว (ไม่แก้ตัวเลข/ตรรกะบัญชี) เปิดในแอป + พิมพ์ PDF ได้ ขอบเขตสาขาหลัก: "
                  "(1) ใบรับรองการกระทบยอด — POS↔เงินเข้า↔ผลต่างคงเหลือ แยกช่องทาง + ผู้จัดทำ/ผู้สอบทาน/วันอนุมัติ "
                  "+ เลข SHA-256 ของงบที่ปิด; (2) ทะเบียนการปรับปรุงด้วยมือ — ทุกการแก้มือ ใคร·เมื่อ·เหตุผล "
                  "รายการไม่มีเหตุผลถูกเน้นแดง; (3) อายุรายการค้าง — แยกช่วง 0-7/8-30/เกิน 30 วันนับจากสิ้นงวด "
                  "(เพิ่ม 2026-06-19 — ปิดช่องว่างตามหลัก Big4/COSO)",
     "steps": ["เปิด #reports -> กลุ่ม 'การควบคุมภายใน / Audit & Controls' เห็น 3 การ์ด",
               "เปิด #certificate เลือกสาขา -> เห็นตารางช่องทาง + ผลต่างคงเหลือ + ช่องผู้จัดทำ/ผู้สอบทาน; "
               "กด 'พิมพ์สาขานี้' และ 'พิมพ์ทุกสาขา' -> ได้ PDF (แต่ละสาขาขึ้นหน้าใหม่)",
               "เปิด #override -> ทุกการแก้มือของเดือนพร้อมเหตุผล; รายการไม่มีเหตุผลขึ้นสีแดง + ตัวนับ; พิมพ์ PDF ได้",
               "เปิด #aging -> ตัวเลขสรุป 3 ช่วงอายุ + ตารางรายการค้าง (งวดที่อนุมัติ/ปิดแล้วไม่นับ); พิมพ์ PDF ได้",
               "ยืนยันสาขาแฟรนไชส์ไม่โผล่ในทั้ง 3 รายงาน + canary อุดร 2026-05-01 ส่วนลด ฿35,511.48 ไม่เปลี่ยน"],
     "expected": "3 รายงานเปิดในแอป + พิมพ์ PDF ได้, ใบรับรองมีช่องเซ็น+SHA-256, ทะเบียนเน้นรายการไร้เหตุผล, "
                 "อายุแยก 3 ช่วงถูกต้อง, เฉพาะสาขาหลัก, ไม่แตะตัวเลขบัญชี"},
    {"id": "TC-32", "title": "ทะเบียนการปิดเดือน + สรุปผลต่างที่รับทราบ (Audit & Controls #2)",
     "role": "ผู้จัดการ / ผู้สอบบัญชี",
     "narrative": "อีกสองรายงานควบคุมภายในในหน้า 📊 รายงาน อ่านอย่างเดียว สาขาหลัก: "
                  "(1) ทะเบียนการปิดเดือน — มุมมองข้ามทุกสาขา × 6 เดือนล่าสุด: ปิด/อนุมัติ/ยังเปิด + "
                  "ผู้อนุมัติ·วันที่ + %กระทบยอด + มี snapshot ไหม (control dashboard ว่ามีงวดไหนค้าง); "
                  "(2) สรุปผลต่างที่รับทราบ/Write-off — รวมยอดที่รับทราบปล่อยผ่าน (ช่อง+รายวัน) แยกตามเหตุผล+สาขา "
                  "ตั้งเกณฑ์ต่อรายการได้ เน้นแดงรายการเกินเกณฑ์ (เพิ่ม 2026-06-19)",
     "steps": ["เปิด #reports -> กลุ่มควบคุมภายในเห็นการ์ด 'ทะเบียนการปิดเดือน' + 'สรุปผลต่างที่รับทราบ'",
               "เปิด #closereg -> เห็นทุกสาขา×6เดือน สถานะ/ผู้อนุมัติ/วันที่/% /หลักฐาน + ตัวนับ 'ยังไม่ปิด'; พิมพ์ PDF ได้",
               "เปิด #accepted -> ยอดรับทราบรวม + ตารางแยกเหตุผล/สาขา + รายการ; เปลี่ยนเกณฑ์ (เช่น 500) แล้วตัวเลขเกินเกณฑ์ปรับตาม; พิมพ์ PDF ได้",
               "ยืนยันแฟรนไชส์ไม่โผล่ทั้งสองรายงาน + canary อุดร 2026-05-01 ฿35,511.48 ไม่เปลี่ยน"],
     "expected": "ปิดเดือนแสดงสถานะ/ผู้อนุมัติ/snapshot ครบ, Write-off รวม+แยกเหตุผล/สาขา+เน้นเกินเกณฑ์ถูกต้อง, "
                 "พิมพ์ PDF ได้, เฉพาะสาขาหลัก, ไม่แตะตัวเลขบัญชี"},
    {"id": "TC-33", "title": "การควบคุมภายในของการแก้ไข — เหตุผลบังคับ / เปิดงวด / ตัวตนจริง / manifest",
     "role": "ผู้จัดการ / ผู้สอบบัญชี",
     "narrative": "ปิดช่องโหว่เชิงโครงสร้างตามหลักการควบคุมภายใน (2026-06-19): "
                  "(1) ทุกการแก้ด้วยมือต้องมีเหตุผล >=10 ตัวอักษร — ย้ายช่อง/ปรับ lane/รับทราบผลต่าง "
                  "(จับคู่ด้วยมือถ้าไม่ใส่ ระบบเติมเหตุผลมาตรฐาน); (2) เปิดงวดที่ปิดแล้ว (ส่งกลับแก้) "
                  "ต้องใช้รหัสผู้จัดการ + ระบุเหตุผล บันทึก audit; (3) ระบุตัวตนผู้ทำจริงจากบัญชีล็อกอิน "
                  "(ไม่เชื่อค่าที่หน้าเว็บส่ง) ทั้งการแก้และอนุมัติ; (4) snapshot ปิดงวดผูกรายชื่อ+SHA-256 "
                  "ของเอกสารต้นทางทุกไฟล์",
     "steps": ["ลองย้ายช่องโดยไม่ใส่เหตุผล (หรือ <10 ตัวอักษร) -> ระบบปฏิเสธ; ใส่เหตุผลชัดเจน -> ผ่าน",
               "ปิดงวด (ผู้จัดการ) แล้วลอง 'ส่งกลับแก้' โดยไม่ใส่รหัสผู้จัดการ -> ถูกปฏิเสธ; "
               "ใส่รหัส+เหตุผล -> เปิดได้ และมีบันทึกใน audit log",
               "ถ้าล็อกอิน dashboard ไว้ -> ตรวจ audit log/ผู้อนุมัติว่าเป็นชื่อบัญชีจริง ไม่ใช่ 'staff'",
               "เปิดใบรับรองการกระทบยอดของงวดที่ปิดแล้ว -> บรรทัดหลักฐานแช่แข็งแสดง 'ผูกเอกสารต้นทาง N ไฟล์'",
               "ยืนยัน 134 tests เขียว + canary อุดร 2026-05-01 ฿35,511.48 ไม่เปลี่ยน"],
     "expected": "แก้ด้วยมือไม่มีเหตุผลถูกปฏิเสธ, เปิดงวดต้องรหัสผู้จัดการ+เหตุผล, ตัวตนจริงถูกบันทึก, "
                 "ใบรับรองแสดงจำนวนเอกสารต้นทางที่ผูก, ไม่กระทบตรรกะบัญชี/การกระทบยอด"},
    {"id": "TC-34", "title": "ประวัติการแก้ไข (Audit Trail) + Excel ทุกรายงาน + ผูกเข้าปิดเดือน",
     "role": "ผู้จัดการ / ผู้สอบบัญชี",
     "narrative": "ส่วนขยายชุดควบคุมภายใน (2026-06-19): (1) หน้า #auditlog ไล่ทุกการเปลี่ยนแปลงในเดือน "
                  "ใคร·เมื่อ·แก้อะไร เรียงใหม่สุดก่อน บันทึกแบบเพิ่มอย่างเดียว; (2) ทุกรายงานควบคุมภายในมีปุ่ม "
                  "⬇ Excel; (3) หน้าสรุปผู้จัดการมีปุ่มเปิดทะเบียนการปรับปรุง + ใบรับรองของสาขานั้น และ "
                  "เดือนที่ปิดแล้วมีปุ่มใบรับรองบนหัวสาขา; (4) option บังคับล็อกอิน env RECON_ENFORCE_AUTH "
                  "(ค่าเริ่ม=ปิด ไม่กระทบเดิม)",
     "steps": ["เปิด #reports -> เห็นการ์ด 'ประวัติการแก้ไข (Audit Trail)' ในกลุ่มควบคุมภายใน",
               "เปิด #auditlog -> เห็นรายการ ใคร·เมื่อ·ประเภท·การกระทำ·รายละเอียด; เลือกสาขา + พิมพ์ PDF + ⬇ Excel ได้",
               "ในแต่ละรายงานควบคุมภายใน กด ⬇ Excel -> ได้ไฟล์ .xlsx เปิดได้",
               "เข้าหน้าสาขาสถานะ 'รอผู้จัดการตรวจ' -> มีปุ่มทะเบียนการปรับปรุง + ใบรับรอง (ก่อนเซ็น); "
               "สาขาที่ปิดแล้ว -> หัวสาขามีปุ่มใบรับรอง",
               "(ออปชัน) ตั้ง RECON_ENFORCE_AUTH=1 บน Fly -> การเขียนต้องล็อกอินก่อน (401 ถ้าไม่ล็อกอิน)",
               "ยืนยัน 135 tests เขียว + canary อุดร 2026-05-01 ฿35,511.48 ไม่เปลี่ยน"],
     "expected": "Audit Trail แสดงครบ+พิมพ์/Excel ได้, ทุกรายงานมี Excel, ปุ่มในหน้าผู้จัดการ/หัวสาขาทำงาน, "
                 "enforcement ทำงานเมื่อเปิด env, เฉพาะสาขาหลัก, ไม่กระทบตัวเลขบัญชี"},
    {"id": "TC-29", "title": "Responsive Nav + Reports Hub (#reports) -- เฟส 5",
     "role": "ทุกบทบาท",
     "narrative": "เฟส 5 เพิ่มแถบนำทาง 4 โซน (หน้าหลัก / รายงาน / Admin / วินิจฉัย) "
                  "แสดงเป็น top-tab บน desktop (≥641px) และ bottom-nav ตรึงด้านล่างบนมือถือ "
                  "#reports = hub รายงานทั้งหมดแทนเมนู ⋯ ที่เคยกระจัดกระจาย "
                  "แท็บ active เปลี่ยนตาม route ปัจจุบันอัตโนมัติ (เฟส 5 consolidation 2026-06-19)",
     "steps": ["บนจอกว้าง: ตรวจ top-tab 4 แท็บปรากฏใต้แถบ header — หน้าหลัก / รายงาน / Admin / วินิจฉัย",
               "กด แท็บ 'รายงาน' หรือนำทางไป #reports -> ตรวจ hub แสดง 3 กลุ่ม: งานประจำ / การตรวจสอบ / เอกสาร&เครื่องมือ",
               "กดการ์ด 'Dashboard รายได้รายวัน' -> route #dashboard; กด back; กดการ์ด 'ผิดช่อง' -> #wrong",
               "กดการ์ด PDF (รายงานการจับคู่รายวัน) -> เปิด tab ใหม่ ไม่ใช่ route ใน SPA",
               "กด แท็บ 'Admin' -> นำทางไป #setup; ตรวจ active state ที่ tab Admin",
               "กด แท็บ 'วินิจฉัย' -> #board; กด 'หน้าหลัก' -> กลับ overview",
               "บนมือถือ (จำลอง viewport <641px): ตรวจ bottom-nav ตรึง 4 icon ด้านล่าง; top-tab ซ่อน",
               "เปลี่ยน route ผ่าน bottom-nav -> active icon เปลี่ยนสีตาม zone"],
     "expected": "top-tab ปรากฏบน desktop / bottom-nav บน mobile, active state ถูกต้องทุก route, "
                 "#reports hub แสดงครบ 3 กลุ่ม, การ์ด in-app นำทาง SPA, การ์ด PDF เปิดใหม่, "
                 "ไม่มีปุ่ม ⋯ เพิ่มเติม แล้ว, 🔑 รหัส ยังอยู่"},
    {"id": "TC-28", "title": "Setup in-app (#setup) -- ตั้งค่าระบบ (เฟส 4)",
     "role": "ผู้จัดการ / เจ้าหน้าที่กระทบยอด",
     "narrative": "หน้า #setup คือ admin panel รวม 5 ส่วน: ① ตารางจับคู่รหัสรับเงิน (payment-map) "
                  "② วิธีกระทบยอดราย lane (method-config) ③ บัญชีธนาคาร ④ MIDs ⑤ รหัสสเตทเมนต์ "
                  "แทนที่ /app/setup.html ใน เฟส 4 consolidation 2026-06-19 "
                  "ส่วนที่ต้องรหัสผู้จัดการ: แก้ lane / seed-requirement / override statement-code",
     "steps": ["เปิด [เพิ่มเติม] > ⚙️ ตั้งค่าระบบ -> route #setup โหลดการตั้งค่า",
               "ตรวจ ① ตาราง payment-map: คอลัมน์รหัส POS / ประเภท / ธนาคาร / เอกสาร / วิธี / พารามิเตอร์ / ยอด / กระทบ / สถานะ",
               "แก้ category/bank ใน dropdown ① -> toast 'บันทึก ... แล้ว' (write token)",
               "คลิก 'แก้ไข ▸' ใน lane column -> form inline ปรากฏ -> แก้ค่า -> บันทึก (ผู้จัดการ) -> prompt รหัสผู้จัดการ",
               "กด '+ เพิ่มบัญชี' ใน ③ -> prompt สาขา/ธนาคาร/เลขบัญชี -> toast 'เพิ่มบัญชีแล้ว'",
               "กด '+ เพิ่ม MID' ใน ④ -> prompt สาขา/acquirer/MID -> toast 'เพิ่ม MID แล้ว'",
               "ใน ⑤ เลือก channel ใหม่สำหรับรหัส -> กด 'บันทึก' -> prompt ผู้จัดการ -> toast 'บันทึก ... (แก้ N บรรทัด)'"],
     "expected": "เปิดได้ใน index.html (route #setup) ไม่ redirect ไป setup.html, "
                 "ทั้ง 5 ส่วนโหลดครบ, dropdown payment-map บันทึกได้ (write token), "
                 "lane editor inline ปรากฏ/ซ่อนได้ (manager token), "
                 "③④ เพิ่ม/ลบได้, ⑤ override/revert ทำงาน (manager token), "
                 "month picker เปลี่ยนเดือนโหลดใหม่, ปุ่ม ← กลับ ออกจาก #setup"},
    {"id": "TC-27", "title": "Discrepancies Admin (#discrepancies) -- จัดการข้อแตกต่าง (เฟส 3b)",
     "role": "ผู้จัดการ / เจ้าหน้าที่กระทบยอด",
     "narrative": "หน้า #discrepancies คือ admin view สำหรับจัดการข้อแตกต่าง (ต้องมี session แล้ว) "
                  "กรองตามสถานะ/ช่องทาง/ธนาคาร จัดเรียงตามขนาดผลต่าง resolve ทีละแถว "
                  "มอบหมายผู้รับผิดชอบ bulk-resolve หลายแถวพร้อมกัน drill รายวัน "
                  "ใช้คู่กับ (ไม่แทนที่) ตาราง ตรวจผลต่างรายวัน ในหน้าสาขา "
                  "เฟส 3b consolidation 2026-06-19 แทนที่ loadDiscrepancies ~2915 ใน classic.html",
     "steps": ["เปิดหน้าสาขาที่มี session -> กดปุ่ม ข้อแตกต่าง (header สาขา)",
               "ตรวจตาราง 9 คอลัมน์: ช่องทาง / ธนาคาร / ประเภท / ต่าง / เงินเข้า / POS / มอบหมาย / การแก้ไข",
               "คลิก header ต่าง -> เรียงตามขนาดผลต่าง (descending ก่อน กดอีกครั้งกลับ)",
               "คลิกเซลล์ยอดต่าง -> แถวขยายแสดงผลต่างรายวัน",
               "เลือก dropdown ยอมรับ ที่แถว -> แถวหายจาก view รอแก้ไข",
               "ติ๊ก checkbox หลายแถว -> แถบ bulk ล่างจอ -> กด ยอมรับทั้งหมด",
               "สาขาไม่มี session -> ขึ้น empty state + ลิงก์เปิดหน้าสาขาก่อน"],
     "expected": "ตารางครบ 9 คอลัมน์, เรียงตามต่างได้ (asc/desc), drill รายวันทำงาน, "
                 "resolve ทีละแถวบันทึกทันที, bulk-resolve อัปเดตจำนวน รอแก้ไข, "
                 "assign บันทึกอัตโนมัติ, สาขาไม่มี session ขึ้น empty state, "
                 "ไม่ใช้ Tailwind / classic CSS"},
    {"id": "TC-25", "title": "Coupon Series (#coupon) -- วิเคราะห์ส่วนลดคูปองแยกซีรีส์ (in-app)",
     "role": "ฝ่ายบัญชี / ผู้จัดการ",
     "narrative": "หน้า #coupon แสดงส่วนลดคูปองแยกตามซีรีส์ (coupon_series_id) จาก pos_coupon_day -- "
                  "ข้อมูลเดียวกับที่ร่างบัญชีใช้คำนวณบรรทัดส่วนลด ยอดจึงสอดคล้องกัน "
                  "เลือกระดับเวลา (เดือน/สัปดาห์/วัน) + กรองสาขาได้ + คลิกซีรีส์เพื่อดูแยกรายสาขาและรายวัน "
                  "แทนที่ลิงก์เดิมที่ส่งไปหน้า classic.html#coupon (เฟส 2 consolidation)",
     "steps": ["เปิด [เพิ่มเติม] > 🎟️ วิเคราะห์คูปองซีรีส์ (หรือกดปุ่ม 🎟️ รายละเอียดคูปอง ในแถว COUPON ของสาขา)",
               "ตรวจ KPI ด้านบน: ส่วนลดรวม / จำนวนใบ / จำนวนซีรีส์ / ช่วงก่อนหน้า",
               "ตรวจตารางซีรีส์: ชื่อซีรีส์, จำนวนใช้, มูลค่าส่วนลด, MoM (▲แดง=มากขึ้น ▼เขียว=น้อยลง), % รวม",
               "คลิกแถวซีรีส์ใด -> ตาราง drill แยกรายสาขา + รายวัน ขึ้นด้านล่าง กด ✕ ปิด drill",
               "กดปุ่ม สัปดาห์ / วัน -> เลือกวันที่ -> ยอดเปลี่ยนตาม grain",
               "เลือกสาขาจาก dropdown -> กรองเฉพาะสาขานั้น (ตาราง by_branch หายไปเพราะ filtered แล้ว)",
               "ตรวจว่ายอดรวม #coupon ตรงกับบรรทัดส่วนลดในสรุปบัญชีรายวัน (TC-17) เดือนเดียวกัน"],
     "expected": "เปิดได้ใน index.html (route #coupon) ไม่ redirect ไป classic.html, "
                 "ตารางซีรีส์แสดงครบ, drill รายสาขา/รายวันทำงาน, gran toggle ทำงาน, "
                 "กรองสาขาทำงาน (master only), ยอดรวมสอดคล้องกับร่างบัญชี, "
                 "ไม่มี Tailwind / classic CSS, ปุ่ม 🎟️ ในแถว COUPON ชี้มา #coupon ไม่ใช่ classic.html"},
    {"id": "TC-24", "title": "Bank Coverage (#bankcov) -- สถานะครอบคลุมสเตทเมนต์ธนาคาร (session-scoped)",
     "role": "ฝ่ายบัญชี / ผู้ตรวจสอบ",
     "narrative": "Bank coverage ตรวจทุกบรรทัดในสเตทเมนต์ธนาคารของสาขาที่เลือก ว่ากระทบยอดครบ/ยังรอ/ตัดออก "
                  "- ส่วน '3 ขั้นบัตรเครดิต' ตรวจว่า POS gross <-> สรุปยอดบัตร <-> เงินเข้าธนาคารตรงกัน "
                  "- session-scoped เช่นเดียวกับ Trend",
     "steps": ["เปิด [เพิ่มเติม] > วินิจฉัย > Bank Coverage",
               "เลือกสาขาจาก dropdown -- ระบบดึง session ให้อัตโนมัติ",
               "ถ้าไม่มี session -> ขึ้น empty state + ลิงก์หน้าสาขา",
               "ถ้ามี session: ตรวจแถบสรุป X/Y ตรงแล้ว - % match - ยอดตรง/ยังรอ/ตัดออก",
               "ตรวจ chip รายช่องทาง (เงินสด/QR/บัตร/Grab/LM) matched/total",
               "กดปุ่มกรอง [ยังรอ] / [ตรงแล้ว] / [ตัดออก] / [ทั้งหมด] -> ตารางกรองได้",
               "ถ้ามีบัตรเครดิต ตรวจส่วน '3 ขั้นบัตรเครดิต': POS gross / สรุปยอด / เงินเข้า + สถานะขั้น 1&2"],
     "expected": "dropdown branch ทำงาน, ไม่มี session -> empty state, มี session -> แถบสรุปถูกต้อง, "
                 "chip รายช่องทาง, ปุ่มกรองทำงาน, 3 ขั้นบัตรเครดิตถ้ามีข้อมูล, "
                 "เฉพาะ master branch, ไม่สร้าง session ใหม่"},
]


def _uat_signoff_html(period_month: str | None) -> str:
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    pm = esc(period_month) if period_month else "____________"
    BL = "______________"                       # write-in blank
    css = """<style>
      @page{size:A4;margin:15mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12.5px;line-height:1.55;margin:0;padding:22px;background:#f8fafc}
      .sheet{background:#fff;max-width:920px;margin:0 auto 18px;padding:28px 32px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:21px;margin:0 0 4px} h2{font-size:15.5px;margin:20px 0 8px;border-left:4px solid #ea580c;padding-left:9px}
      h3{font-size:13.5px;margin:0 0 4px} .sub{color:#64748b;font-size:12px}
      table{width:100%;border-collapse:collapse;margin:6px 0 12px;font-size:12px}
      th,td{padding:6px 9px;border:1px solid #e2e8f0;text-align:left;vertical-align:top}
      th{background:#f8fafc;color:#475569;font-size:11px;font-weight:700}
      ol,ul{margin:4px 0 4px 18px;padding:0} li{margin:2px 0}
      .tc{border:1px solid #e5e7eb;border-radius:10px;padding:13px 16px;margin:11px 0;page-break-inside:avoid}
      .tc-h{display:flex;align-items:baseline;gap:9px;flex-wrap:wrap;margin-bottom:5px}
      .tc-id{font-weight:700;color:#ea580c;font-size:14px} .tc-title{font-weight:700;font-size:13.5px}
      .role{background:#eff6ff;color:#1e40af;border-radius:99px;padding:1px 9px;font-size:10.5px;font-weight:600;white-space:nowrap}
      .narr{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:8px 12px;margin:6px 0;font-size:12px;line-height:1.55}
      .narr b.k{color:#9a3412}
      .lab{font-size:11px;color:#64748b;font-weight:700;margin-top:6px}
      .exp{font-size:12px}
      .result{display:flex;gap:18px;flex-wrap:wrap;align-items:center;margin-top:9px;padding-top:8px;border-top:1px dashed #e2e8f0;font-size:12px}
      .chk{font-size:13px} .fill{color:#94a3b8}
      .sign{margin-top:8px} .sign td{height:34px}
      .note{background:#f1f5f9;border-radius:8px;padding:9px 13px;font-size:12px;margin:8px 0}
      .toolbar{max-width:920px;margin:0 auto 14px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:8px;padding:9px 18px;font-size:13px;cursor:pointer}
      .mut{color:#64748b;font-size:11.5px}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;border-radius:0;max-width:none;margin:0}.sheet+.sheet{page-break-before:always}}
    </style>"""

    p = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
         "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
         "<title>UAT Sign-off — ระบบกระทบยอด</title>", css, "</head><body>",
         "<div class='toolbar'><button class='btn' onclick='window.print()'>🖨️ พิมพ์ / บันทึก PDF</button></div>"]

    # Sheet 1 — header, participants, prerequisites, happy path
    p.append("<div class='sheet'>")
    p.append("<h1>เอกสารทดสอบและรับรองการใช้งาน (UAT Sign-off)</h1>")
    p.append("<div class='sub'>ระบบกระทบยอดรายเดือน Oshinei · เวอร์ชัน 2026-06-15</div>")
    p.append("<h2>1. ข้อมูลเอกสาร</h2><table>"
             f"<tr><th style='width:38%'>URL ทดสอบ</th><td>oshinei-dashboard.fly.dev:8443/app/</td></tr>"
             f"<tr><th>เดือนที่ทดสอบ</th><td>{pm}</td></tr>"
             f"<tr><th>สาขาที่ทดสอบ</th><td>{BL}</td></tr>"
             f"<tr><th>วันที่ทดสอบ</th><td>{BL}</td></tr></table>")
    p.append("<h2>2. ผู้เข้าร่วมทดสอบ</h2><table>"
             "<tr><th style='width:38%'>บทบาท</th><th>ชื่อ-สกุล</th></tr>"
             "<tr><td>เจ้าหน้าที่กระทบยอด</td><td></td></tr>"
             "<tr><td>ผู้จัดการ</td><td></td></tr>"
             "<tr><td>ฝ่ายบัญชี / ผู้ตรวจสอบ</td><td></td></tr></table>")
    p.append("<h2>3. สิ่งที่ต้องเตรียม</h2><ul>"
             "<li><b>รหัสบันทึก (admin)</b> — กรอกผ่านปุ่ม 🔑 รหัส · <b>รหัสผู้จัดการ</b> — ระบบถามตอนปิดเดือน/แก้รหัส (ขอจากเจ้าของระบบ)</li>"
             "<li>เอกสารตัวอย่างของเดือนที่ทดสอบ: สเตทเมนต์ KTB/KBANK, สรุปยอดบัตร, รายงาน Grab/Line Man, สลิปนำฝากเงินสด</li>"
             "<li>เบราว์เซอร์รุ่นใหม่ + เครื่องพิมพ์/Save as PDF · ข้อมูล POS ของเดือนนั้นถูกซิงก์แล้ว</li></ul>")
    p.append("<h2>4. ภาพรวมขั้นตอน 1 เดือน (อ่านก่อนเริ่ม)</h2>")
    p.append("<div class='narr'>ระบบเปรียบ “ยอดที่ POS บันทึก” กับ “เงินที่เข้าบัญชี/เอกสารจริง” ของแต่ละ<b class='k'>ช่องทางรับเงิน</b> "
             "ไล่เป็นรายวันทั้งเดือน เพื่อมั่นใจว่าเงินเข้าครบจริง ไม่ตกหล่น/บันทึกผิดช่อง</div>")
    p.append("<ol>"
             "<li>เปิดหน้าแรก → เห็นทุกสาขาเป็นกล่องงาน</li>"
             "<li>เข้าสาขา → ช่องทางที่ “ยังไม่ตรง” ขึ้นก่อน + บอกว่าขาดเอกสารอะไร</li>"
             "<li>อัปโหลดเอกสารที่ขาด</li>"
             "<li>ระบบจับคู่อัตโนมัติทันที (ไม่มีปุ่ม Run Match)</li>"
             "<li>ถ้ามี “เงินอาจถูกบันทึกผิดช่อง” → ยืนยันปุ่มเดียว</li>"
             "<li>ผลต่างที่อธิบายได้ → “รับทราบผลต่าง” พร้อมเหตุผล</li>"
             "<li>เจ้าหน้าที่ “ส่งตรวจ” → ผู้จัดการ “ปิดเดือน” (ล็อกถาวร)</li>"
             "<li>ออกรายงานเก็บหลักฐาน / ให้สาขาเซ็น / ตรวจการจับคู่</li></ol>")
    p.append("<div class='note'>วิธีบันทึกผล: ทำตาม “ขั้นตอน” → เทียบ “ผลที่คาดหวัง” → ทำเครื่องหมาย ☑ ผ่าน หรือ ☒ ไม่ผ่าน "
             "ถ้าไม่ผ่านบันทึกในช่องหมายเหตุ + ลงเลขที่ปัญหาในตารางบันทึกปัญหา (ท้ายเอกสาร)</div>")
    p.append("</div>")

    # Sheet 2 — test cases
    p.append("<div class='sheet'>")
    p.append("<h2>5. รายการทดสอบ + คำอธิบายการใช้งาน</h2>")
    for c in _UAT_CASES:
        p.append("<div class='tc'>")
        p.append(f"<div class='tc-h'><span class='tc-id'>{esc(c['id'])}</span>"
                 f"<span class='tc-title'>{esc(c['title'])}</span>"
                 f"<span class='role'>{esc(c['role'])}</span></div>")
        p.append(f"<div class='narr'>📖 {esc(c['narrative'])}</div>")
        p.append("<div class='lab'>ขั้นตอนทดสอบ</div><ol>")
        for s in c["steps"]:
            p.append(f"<li>{esc(s)}</li>")
        p.append("</ol>")
        p.append(f"<div class='lab'>ผลที่คาดหวัง</div><div class='exp'>{esc(c['expected'])}</div>")
        p.append("<div class='result'>"
                 "<span class='chk'>☐ ผ่าน &nbsp; ☐ ไม่ผ่าน</span>"
                 f"<span class='fill'>ผู้ทดสอบ: {BL}</span>"
                 f"<span class='fill'>วันที่: __________</span>"
                 f"<span class='fill'>หมายเหตุ: {BL}</span></div>")
        p.append("</div>")
    p.append("</div>")

    # Sheet 3 — issue log + sign-off + quick reference
    p.append("<div class='sheet'>")
    p.append("<h2>6. บันทึกปัญหาที่พบ (Issue Log)</h2><table>"
             "<tr><th style='width:8%'>#</th><th style='width:10%'>TC</th><th>รายละเอียดปัญหา</th>"
             "<th style='width:16%'>ความรุนแรง</th><th style='width:14%'>สถานะ</th></tr>"
             + "".join("<tr><td></td><td></td><td></td><td></td><td></td></tr>" for _ in range(4))
             + "</table>")
    p.append("<h2>7. สรุปผลและการลงนามรับรอง</h2>")
    p.append(f"<p>ผลการทดสอบโดยรวม: จำนวนผ่าน {BL} / {len(_UAT_CASES)} รายการ</p>")
    p.append("<p><b>มติการรับรอง</b> (เลือกหนึ่ง):</p><ul style='list-style:none;margin-left:0'>"
             "<li>☐ <b>รับรองให้ใช้งานจริง</b> (Accepted) — ผ่านทุกข้อสำคัญ</li>"
             "<li>☐ <b>รับรองแบบมีเงื่อนไข</b> — ใช้งานได้ แต่ต้องแก้ข้อ ______ ภายใน ______</li>"
             "<li>☐ <b>ยังไม่รับรอง</b> (Rejected) — มีปัญหาสำคัญข้อ ______</li></ul>")
    p.append("<table class='sign'>"
             "<tr><th style='width:30%'>บทบาท</th><th>ชื่อ-สกุล</th><th>ลายเซ็น</th><th style='width:16%'>วันที่</th></tr>"
             "<tr><td>เจ้าหน้าที่กระทบยอด</td><td></td><td></td><td></td></tr>"
             "<tr><td>ผู้จัดการ</td><td></td><td></td><td></td></tr>"
             "<tr><td>ฝ่ายบัญชี / ผู้ตรวจสอบ</td><td></td><td></td><td></td></tr>"
             "<tr><td>เจ้าของระบบ (อนุมัติ)</td><td></td><td></td><td></td></tr></table>")
    p.append("<h2>ภาคผนวก — ทางลัด</h2><table>"
             "<tr><th style='width:46%'>อยากทำอะไร</th><th>ไปที่</th></tr>"
             "<tr><td>ดูภาพรวมทุกสาขา / ทำสาขา</td><td>🏠 หน้าหลัก (/app/)</td></tr>"
             "<tr><td>ตั้งค่า / แก้รหัสสเตทเมนต์</td><td>⚙️ Admin → #setup (ส่วน ⑤)</td></tr>"
             "<tr><td>รายงานเงินผิดช่อง + ใบเซ็น</td><td>📊 รายงาน → ⇄ ผิดช่อง (#wrong)</td></tr>"
             "<tr><td>ตรวจการจับคู่ทั้งเดือน (PDF)</td><td>หน้าสาขา → 🖨️ ตรวจการจับคู่รายวัน</td></tr>"
             "<tr><td>สรุปบัญชีรายวัน (รายใบ — 1 วัน/การ์ด)</td><td>📊 รายงาน → 📅 สรุปบัญชีรายวัน (#daybook)</td></tr>"
             "<tr><td>หลักฐานรายวันระดับบรรทัด</td><td>/api/daily-report?day=YYYY-MM-DD</td></tr>"
             "<tr><td>พจนานุกรมรหัส</td><td>📊 รายงาน → 📖 พจนานุกรมรหัสในสเตทเมนต์</td></tr>"
             "<tr><td>คู่มือเต็ม</td><td>📊 รายงาน → 📖 คู่มือ (/app/help.html)</td></tr></table>")
    p.append("<div class='mut'>สถานะการจับคู่: ✓ ตรง = POS ตรงเอกสาร · ✗ ต่าง = ยอดไม่ตรงต้องตรวจ · ขาดเอกสาร = POS มีแต่เอกสารไม่มา · "
             "เกินเอกสาร = เงินเข้าเกิน POS · รอเอกสารตามรอบ = ยังไม่ถึงรอบ (ปกติ) · ผลต่าง = เอกสาร − POS</div>")
    p.append("</div></body></html>")
    return "".join(p)


@app.get("/api/uat-signoff")
def uat_signoff(period_month: str | None = None):
    """Printable UAT sign-off checklist (usage narrative + pass/fail + signatures)."""
    if period_month and not _PERIOD_RE.match(period_month):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    return Response(content=_uat_signoff_html(period_month),
                    media_type="text/html; charset=utf-8")


# ── Transaction-level drill-down (mapped / unmapped codes + statement lines) ───

@app.get("/api/drill")
def drill(branch_id: int, category: str | None = None, bank: str | None = None,
          day: str | None = None, period_month: str | None = None):
    """Day-scoped (or month-scoped) drill into one branch's payment codes and the
    matching statement lines, so staff can see which codes are mapped vs unmapped
    and act on the unmapped. `category`/`bank` narrow to a single lane."""
    if day and not _DAY_RE.match(day):
        raise HTTPException(400, "day must be in 'YYYY-MM-DD' format.")
    if not day and not (period_month and _PERIOD_RE.match(period_month)):
        raise HTTPException(400, "pass a valid day (YYYY-MM-DD) or period_month (YYYY-MM).")
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        codes = pos_codes_scope(conn, branch_id, day=day, period_month=period_month)
        for c in codes:
            c["in_lane"] = bool(category) and c["category"] == category \
                and (not bank or bank in ("*", "") or c["bank"] == bank)
        lane_codes = [c for c in codes if c["in_lane"]] if category else codes
        attention  = [c for c in codes if c["needs_map"]]

        ext_lines = []
        if category:
            ext_lines = lane_lines(conn, branch_id, (day or period_month)[:7],
                                   category, bank or "NONE", day)

        pos_total = round(sum(c["amount"] for c in lane_codes if c["reconcile"]), 2)
        ext_total = round(sum(l["amount"] for l in ext_lines), 2)
        bname = conn.execute("SELECT name FROM branches WHERE branch_id=?",
                             (branch_id,)).fetchone()["name"]
        return {
            "branch_id": branch_id, "branch_name": bname,
            "day": day, "period_month": period_month,
            "category": category, "bank": bank,
            "codes": codes, "lane_codes": lane_codes, "attention": attention,
            "ext_lines": ext_lines,
            "pos_total": pos_total, "ext_total": ext_total,
            "variance": round(ext_total - pos_total, 2),
        }


# ── Manual reconciliation (tick-to-match, reversible) ──────────────────────────

_EXT_TABLES = {"recon_bank_lines", "recon_cc_lines", "recon_grab_lines", "recon_lineman_lines"}


class ManualMatch(BaseModel):
    branch_id:    int
    category:     str
    bank:         str
    period_month: str
    pos_days:     list[dict] = []   # [{day, amount}]
    ext:          list[dict] = []   # [{table, id, amount}]
    note:         str = ""
    matched_by:   str = ""


@app.get("/api/payout-matching")
def payout_matching_api(branch_id: int, period_month: str):
    """Per-provider net-day ↔ bank-credit grouping (handles combined-day payouts).
    The 3rd leg of the chain: POS↔report gross, report net↔bank payout, here
    traced per actual transfer."""
    _validate_period(period_month)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        return payout_matching(conn, branch_id, period_month)


@app.get("/api/map-suggestions")
def get_map_suggestions(branch_id: int, period_month: str):
    """Bank suggestions for needs-map POS codes, derived from the uploaded
    documents themselves (daily-series match against statement-lane residuals)."""
    _validate_period(period_month)
    with get_db() as conn:
        return {"suggestions": map_bank_suggestions(conn, branch_id, period_month)}


@app.get("/api/suggestions")
def suggestions(branch_id: int, period_month: str):
    """Wrong-channel reclassification candidates: same bank, equal amount, ±1 day.
    Each item is a ready-to-apply /api/manual payload (POST it to accept)."""
    _validate_period(period_month)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        return {"suggestions": suggest_reclassifications(conn, branch_id, period_month)}


@app.get("/api/manual")
def manual_list(branch_id: int, category: str, bank: str, period_month: str):
    _validate_period(period_month)
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        pos = lane_pos_days(conn, branch_id, period_month, category, bank)
        ext = lane_ext_lines(conn, branch_id, period_month, category, bank)

        matches, consumed = [], {}
        for m in conn.execute(
            "SELECT * FROM recon_manual_match WHERE branch_id=? AND period_month=? "
            "AND category=? AND bank=? ORDER BY matched_at DESC",
            (branch_id, period_month, category, bank)):
            pd = json.loads(m["pos_days"] or "[]")
            for x in pd:
                consumed[x["day"]] = consumed.get(x["day"], 0.0) + float(x.get("amount") or 0)
            lines = [dict(r) for r in conn.execute(
                "SELECT ext_table, ext_id, amount FROM recon_manual_match_line WHERE match_id=?",
                (m["match_id"],))]
            matches.append({**dict(m), "pos_days": pd, "lines": lines})

        for d in pos:
            d["consumed"]  = round(consumed.get(d["day"], 0.0), 2)
            d["remaining"] = round(d["total"] - d["consumed"], 2)

        return {"branch_id": branch_id, "category": category, "bank": bank,
                "period_month": period_month, "pos_days": pos, "ext_lines": ext,
                "matches": matches}


@app.post("/api/manual", dependencies=[Depends(require_write_token)])
def manual_create(body: ManualMatch, request: Request):
    _validate_period(body.period_month)
    if not body.pos_days or not body.ext:
        raise HTTPException(400, "Tick at least one POS day and one statement line.")
    # Every manual match carries a reason for the audit trail; one-tap matches (e.g. the
    # reclass suggestions) pass their own note, and anything still blank is stamped with
    # a system reason rather than rejected (keeps the one-click flow working).
    note = (body.note or "").strip() or "(จับคู่ด้วยมือ — ตามรายการที่ตรงกัน)"
    actor = _actor(request, body.matched_by)
    pos_amount = round(sum(float(x.get("amount") or 0) for x in body.pos_days), 2)
    ext_amount = round(sum(float(x.get("amount") or 0) for x in body.ext), 2)
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        _assert_period_open(conn, body.branch_id, body.period_month)

        # Partial allocation guard: a day can't be allocated beyond what's left.
        totals = {d["day"]: d["total"] for d in
                  lane_pos_days(conn, body.branch_id, body.period_month, body.category, body.bank)}
        already: dict = {}
        for m in conn.execute(
            "SELECT pos_days FROM recon_manual_match WHERE branch_id=? AND period_month=? "
            "AND category=? AND bank=?",
            (body.branch_id, body.period_month, body.category, body.bank)):
            for x in json.loads(m["pos_days"] or "[]"):
                already[x["day"]] = already.get(x["day"], 0.0) + float(x.get("amount") or 0)
        for x in body.pos_days:
            day, amt = x.get("day"), float(x.get("amount") or 0)
            if amt <= 0:
                raise HTTPException(400, f"amount for {day} must be positive.")
            remaining = round(totals.get(day, 0.0) - already.get(day, 0.0), 2)
            if amt > remaining + 0.5:
                raise HTTPException(
                    409, f"{day}: amount {amt:.0f} exceeds remaining {remaining:.0f}.")
        cur = conn.execute("""
            INSERT INTO recon_manual_match
                (branch_id, category, bank, period_month, pos_amount, ext_amount,
                 variance, pos_days, note, matched_by)
            VALUES (?,?,?,?,?,?,?,?,?,?)
        """, (body.branch_id, body.category, body.bank, body.period_month,
              pos_amount, ext_amount, round(ext_amount - pos_amount, 2),
              json.dumps(body.pos_days), note, actor))
        mid = cur.lastrowid
        for e in body.ext:
            tbl = e.get("table")
            if tbl not in _EXT_TABLES:
                raise HTTPException(400, f"bad ext table {tbl!r}")
            eid = int(e.get("id"))
            conn.execute("INSERT INTO recon_manual_match_line (match_id, ext_table, ext_id, amount) "
                         "VALUES (?,?,?,?)", (mid, tbl, eid, float(e.get("amount") or 0)))
            conn.execute(f"UPDATE {tbl} SET match_status='manual' WHERE id=?", (eid,))
        _audit(conn, "manual_match", mid, "create", actor,
               [{"field": "match", "old": None,
                 "new": f"POS {pos_amount:.0f} ↔ EXT {ext_amount:.0f} (var {ext_amount-pos_amount:.0f})"}])
        conn.commit()
    return {"match_id": mid, "pos_amount": pos_amount, "ext_amount": ext_amount,
            "variance": round(ext_amount - pos_amount, 2)}


@app.delete("/api/manual/{match_id}", dependencies=[Depends(require_write_token)])
def manual_delete(match_id: int):
    with get_db() as conn:
        m = conn.execute("SELECT branch_id, period_month FROM recon_manual_match WHERE match_id=?",
                         (match_id,)).fetchone()
        if not m:
            raise HTTPException(404)
        _assert_period_open(conn, m["branch_id"], m["period_month"])
        for ln in conn.execute("SELECT ext_table, ext_id FROM recon_manual_match_line WHERE match_id=?",
                               (match_id,)):
            if ln["ext_table"] in _EXT_TABLES:
                conn.execute(f"UPDATE {ln['ext_table']} SET match_status='unmatched' "
                             f"WHERE id=? AND match_status='manual'", (ln["ext_id"],))
        conn.execute("DELETE FROM recon_manual_match_line WHERE match_id=?", (match_id,))
        conn.execute("DELETE FROM recon_manual_match WHERE match_id=?", (match_id,))
        _audit(conn, "manual_match", match_id, "delete", "staff",
               [{"field": "match", "old": "matched", "new": None}])
        conn.commit()
    return {"deleted": match_id}


# ── Lane adjustments (off-POS revenue / manual deposit entries) ──────────────

class LaneAdj(BaseModel):
    branch_id:    int
    period_month: str
    category:     str
    bank:         str = "NONE"
    adj_pos:      float = 0.0
    adj_ext:      float = 0.0
    note:         str | None = None
    adj_type:     str = "other"
    created_by:   str | None = None


@app.get("/api/lanes/adj")
def lane_adj_list(branch_id: int, period_month: str):
    _validate_period(period_month)
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM recon_lane_adj WHERE branch_id=? AND period_month=? "
            "ORDER BY category, bank",
            (branch_id, period_month)).fetchall()
        return [dict(r) for r in rows]


@app.post("/api/lanes/adj", dependencies=[Depends(require_write_token)])
def lane_adj_upsert(body: LaneAdj, request: Request):
    _validate_period(body.period_month)
    if body.adj_type not in ("off_pos", "deposit_note", "other", "accept"):
        raise HTTPException(400, "adj_type must be off_pos | deposit_note | other | accept")
    # AUD-8 (widened 2026-06-19): EVERY lane adjustment is a manual override of the
    # reconciliation and must carry a MEANINGFUL reason (auditors sample these), not a
    # one-character placeholder — not just adj_type='accept'.
    if len((body.note or "").strip()) < 10:
        raise HTTPException(400, "การปรับปรุง lane ต้องระบุเหตุผลที่ชัดเจน (อย่างน้อย 10 ตัวอักษร)")
    actor = _actor(request, body.created_by)
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        _assert_period_open(conn, body.branch_id, body.period_month)
        existing = conn.execute(
            "SELECT id FROM recon_lane_adj WHERE branch_id=? AND period_month=? "
            "AND category=? AND bank=?",
            (body.branch_id, body.period_month, body.category, body.bank)).fetchone()
        if existing:
            conn.execute(
                "UPDATE recon_lane_adj SET adj_pos=?, adj_ext=?, note=?, adj_type=?, "
                "created_by=?, created_at=datetime('now') WHERE id=?",
                (body.adj_pos, body.adj_ext, body.note, body.adj_type,
                 actor, existing["id"]))
            row_id = existing["id"]
        else:
            cur = conn.execute(
                "INSERT INTO recon_lane_adj "
                "(branch_id,period_month,category,bank,adj_pos,adj_ext,note,adj_type,created_by) "
                "VALUES (?,?,?,?,?,?,?,?,?)",
                (body.branch_id, body.period_month, body.category, body.bank,
                 body.adj_pos, body.adj_ext, body.note, body.adj_type, actor))
            row_id = cur.lastrowid
        _audit(conn, "lane_adj", row_id, "update" if existing else "create", actor,
               [{"field": body.adj_type, "old": None,
                 "new": f"{body.category}·{body.bank} POS+{body.adj_pos:.0f}/EXT+{body.adj_ext:.0f}"}])
        conn.commit()
    return {"id": row_id, "adj_pos": body.adj_pos, "adj_ext": body.adj_ext}


@app.delete("/api/lanes/adj/{adj_id}", dependencies=[Depends(require_write_token)])
def lane_adj_delete(adj_id: int):
    with get_db() as conn:
        row = conn.execute(
            "SELECT branch_id, period_month FROM recon_lane_adj WHERE id=?", (adj_id,)).fetchone()
        if not row:
            raise HTTPException(404)
        _assert_period_open(conn, row["branch_id"], row["period_month"])
        conn.execute("DELETE FROM recon_lane_adj WHERE id=?", (adj_id,))
        conn.commit()
    return {"deleted": adj_id}


# ── Manual cross-channel MOVE (booked-in-wrong-channel correction) ───────────
# Declare that ฿amount was rung up under the WRONG channel and actually belongs in
# another; at read-time the move shifts the POS base OUT of (from) and INTO (to),
# correcting BOTH lanes at once (กระทบสองยอดพร้อมกัน). Reversible by deleting the
# row. The shift itself lives in matcher.channel_move_deltas.

_MOVE_CHANNELS = {"CASH", "QR", "TRANSFER", "CREDIT", "GRAB", "LINEMAN", "COUPON"}


def _norm_move_bank(category: str, bank: str | None) -> str:
    b = (bank or "NONE").strip() or "NONE"
    return "NONE" if category == "CASH" else b   # cash carries no bank dimension


class ChannelMove(BaseModel):
    branch_id:     int
    period_month:  str
    from_category: str
    from_bank:     str = "NONE"
    to_category:   str
    to_bank:       str = "NONE"
    amount:        float
    note:          str | None = None
    created_by:    str | None = None


@app.get("/api/channel-move")
def channel_move_list(branch_id: int, period_month: str):
    _validate_period(period_month)
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM recon_channel_move WHERE branch_id=? AND period_month=? "
            "ORDER BY created_at DESC, id DESC",
            (branch_id, period_month)).fetchall()
        return [dict(r) for r in rows]


@app.post("/api/channel-move", dependencies=[Depends(require_write_token)])
def channel_move_create(body: ChannelMove, request: Request):
    _validate_period(body.period_month)
    # A cross-channel move is a management override of what the POS recorded — it must
    # carry a meaningful reason for the audit trail (auditors sample these).
    if len((body.note or "").strip()) < 10:
        raise HTTPException(400, "การย้ายช่องต้องระบุเหตุผลที่ชัดเจน (อย่างน้อย 10 ตัวอักษร)")
    actor = _actor(request, body.created_by)
    fc, tc = body.from_category, body.to_category
    if fc not in _MOVE_CHANNELS or tc not in _MOVE_CHANNELS:
        raise HTTPException(400, f"channel must be one of {sorted(_MOVE_CHANNELS)}")
    fb, tb = _norm_move_bank(fc, body.from_bank), _norm_move_bank(tc, body.to_bank)
    if (fc, fb) == (tc, tb):
        raise HTTPException(400, "ช่องต้นทางและปลายทางต้องไม่ใช่ช่องเดียวกัน")
    if not (body.amount and body.amount > 0):
        raise HTTPException(400, "จำนวนเงินต้องมากกว่า 0")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        _assert_period_open(conn, body.branch_id, body.period_month)
        # Validate against the branch's CURRENT lanes: the source must exist and hold
        # enough POS (over-moving would drive the lane negative — the monthly view drops
        # the row while the daily view keeps it, so the two diverge). The destination must
        # be a real lane or a known bank (else a typo'd bank spawns an orphan pos_only lane).
        rows, _ = category_reconciliation(conn, body.branch_id, body.period_month)
        lanes = {(r["category"], r["bank"]): r["pos_amount"] for r in rows}
        if (fc, fb) not in lanes:
            raise HTTPException(400, "ช่องต้นทางไม่มียอดในเดือนนี้ — เลือกช่องที่มีรายการจริง")
        if body.amount > lanes[(fc, fb)] + 0.5:
            raise HTTPException(400,
                f"ย้ายได้ไม่เกินยอดที่อยู่ในช่องต้นทาง (คงเหลือ {lanes[(fc, fb)]:,.2f})")
        if (tc, tb) not in lanes and tb not in (set(_BANK_LABEL) | {"NONE", "ANY"}):
            raise HTTPException(400, "ช่อง/ธนาคารปลายทางไม่ถูกต้อง")
        cur = conn.execute(
            "INSERT INTO recon_channel_move "
            "(branch_id,period_month,from_category,from_bank,to_category,to_bank,amount,note,created_by) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (body.branch_id, body.period_month, fc, fb, tc, tb,
             round(body.amount, 2), body.note, actor))
        _audit(conn, "channel_move", cur.lastrowid, "create", actor,
               [{"field": "move", "old": f"{fc}·{fb}", "new": f"{tc}·{tb}"},
                {"field": "amount", "old": None, "new": round(body.amount, 2)}])
        conn.commit()
        return {"id": cur.lastrowid, "from": [fc, fb], "to": [tc, tb],
                "amount": round(body.amount, 2)}


@app.delete("/api/channel-move/{move_id}", dependencies=[Depends(require_write_token)])
def channel_move_delete(move_id: int):
    with get_db() as conn:
        row = conn.execute(
            "SELECT branch_id, period_month FROM recon_channel_move WHERE id=?", (move_id,)).fetchone()
        if not row:
            raise HTTPException(404)
        _assert_period_open(conn, row["branch_id"], row["period_month"])
        conn.execute("DELETE FROM recon_channel_move WHERE id=?", (move_id,))
        _audit(conn, "channel_move", move_id, "delete", "staff")
        conn.commit()
    return {"deleted": move_id}


# ── Per-day acknowledged variances (reconcile BY DATE) ───────────────────────
# Staff work a channel's variance day-by-day: each differing day is either
# fixed/matched or acknowledged here with a reason. The channel's todo clears once
# every differing day is acknowledged (or the whole channel is accepted via
# recon_lane_adj). The reason rides into the close Excel pack + all-errors report.

class DayAck(BaseModel):
    branch_id:    int
    period_month: str
    category:     str
    day:          str
    variance:     float = 0.0
    note:         str | None = None
    created_by:   str | None = None


@app.get("/api/day-ack")
def day_ack_list(branch_id: int, period_month: str):
    _validate_period(period_month)
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM recon_day_ack WHERE branch_id=? AND period_month=? "
            "ORDER BY category, day", (branch_id, period_month)).fetchall()
        return [dict(r) for r in rows]


@app.post("/api/day-ack", dependencies=[Depends(require_write_token)])
def day_ack_upsert(body: DayAck, request: Request):
    _validate_period(body.period_month)
    # Same audit-grade rule as a channel accept: a meaningful reason (≥10 chars).
    if len((body.note or "").strip()) < 10:
        raise HTTPException(400, "การรับทราบผลต่างรายวันต้องใส่เหตุผลที่ชัดเจน (อย่างน้อย 10 ตัวอักษร)")
    if not (body.day or "").startswith(body.period_month):
        raise HTTPException(400, "day must fall inside period_month")
    actor = _actor(request, body.created_by)
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        _assert_period_open(conn, body.branch_id, body.period_month)
        existing = conn.execute(
            "SELECT id FROM recon_day_ack WHERE branch_id=? AND period_month=? "
            "AND category=? AND day=?",
            (body.branch_id, body.period_month, body.category, body.day)).fetchone()
        if existing:
            conn.execute(
                "UPDATE recon_day_ack SET variance=?, note=?, created_by=?, "
                "created_at=datetime('now') WHERE id=?",
                (body.variance, body.note, actor, existing["id"]))
            row_id = existing["id"]
        else:
            cur = conn.execute(
                "INSERT INTO recon_day_ack "
                "(branch_id,period_month,category,day,variance,note,created_by) "
                "VALUES (?,?,?,?,?,?,?)",
                (body.branch_id, body.period_month, body.category, body.day,
                 body.variance, body.note, actor))
            row_id = cur.lastrowid
        _audit(conn, "day_ack", row_id, "accept", actor,
               [{"field": "variance", "old": None, "new": body.variance}])
        conn.commit()
    return {"id": row_id}


@app.delete("/api/day-ack/{ack_id}", dependencies=[Depends(require_write_token)])
def day_ack_delete(ack_id: int):
    with get_db() as conn:
        row = conn.execute(
            "SELECT branch_id, period_month FROM recon_day_ack WHERE id=?", (ack_id,)).fetchone()
        if not row:
            raise HTTPException(404)
        _assert_period_open(conn, row["branch_id"], row["period_month"])
        conn.execute("DELETE FROM recon_day_ack WHERE id=?", (ack_id,))
        conn.commit()
    return {"deleted": ack_id}


# ── Chart-of-accounts map (report account code/name per channel / journal slot) ──
# Reports render these columns blank until the owner fills them in here (no deploy).

class AccountMap(BaseModel):
    slot:         str
    account_code: str | None = None
    account_name: str | None = None
    updated_by:   str | None = None


@app.get("/api/account-map")
def account_map_list():
    with get_db() as conn:
        try:
            rows = conn.execute(
                "SELECT slot, account_code, account_name FROM recon_account_map ORDER BY slot").fetchall()
        except sqlite3.OperationalError:
            return []
        return [dict(r) for r in rows]


@app.put("/api/account-map", dependencies=[Depends(require_write_token)])
def account_map_upsert(body: AccountMap):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO recon_account_map(slot,account_code,account_name,updated_by) "
            "VALUES(?,?,?,?) ON CONFLICT(slot) DO UPDATE SET account_code=excluded.account_code, "
            "account_name=excluded.account_name, updated_at=datetime('now'), updated_by=excluded.updated_by",
            (body.slot, body.account_code, body.account_name, body.updated_by or "staff"))
        conn.commit()
    return {"slot": body.slot}


# ── Cross-branch overview ─────────────────────────────────────────────────────

@app.get("/api/overview")
def overview(period_month: str):
    with get_db() as conn:
        branches = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name"
        ).fetchall()

        result = []
        for b in branches:
            bid = b["branch_id"]

            # Reconcilable POS tender total vs external statement total (same engine
            # as the session view, so overview and detail agree).
            rows, _ = category_reconciliation(conn, bid, period_month)
            pos_total = round(sum(r["pos_amount"] for r in rows), 2)
            ext_total = round(sum(r["ext_amount"] for r in rows), 2)

            session = conn.execute(
                "SELECT session_id, status FROM recon_sessions WHERE branch_id=? AND period_month=?",
                (bid, period_month)
            ).fetchone()

            # Open issues to follow up: pending discrepancies (if a session exists)
            # + lanes that aren't matched in the live reconciliation.
            open_issues = 0
            if session:
                open_issues = conn.execute(
                    "SELECT COUNT(*) FROM recon_discrepancies WHERE session_id=? AND resolution='pending'",
                    (session["session_id"],)).fetchone()[0]
            unmatched = sum(1 for r in rows if r["status"] != "matched")

            result.append({
                "branch_id":   bid,
                "branch_name": b["name"],
                "pos_total":   pos_total,
                "ext_total":   ext_total,
                "variance":    round(ext_total - pos_total, 2),
                "session_id":  session["session_id"] if session else None,
                "status":      session["status"] if session else "no_session",
                "open_issues": open_issues,
                "unmatched":   unmatched,
            })

    return result


# ── Session trend (per-day + weekly rollup) ───────────────────────────────────

@app.get("/api/sessions/{session_id}/trend")
def session_trend(session_id: int):
    """Per-day and per-week recon health for one session.
    Reuses the daily_reconciliation engine so numbers agree with the Daily tab."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        data = daily_reconciliation(conn, s["branch_id"], s["period_month"])

    daily_out: list[dict] = []
    for dr in data.get("days", []):
        lanes       = dr.get("lanes", [])
        matched     = dr.get("matched", 0)
        total_lanes = len(lanes)
        pos         = float(dr.get("pos_total") or 0)
        ext         = float(dr.get("ext_total") or 0)
        lane_pct    = round(matched / total_lanes * 100) if total_lanes else 0
        cov_pct     = min(100, round(ext / pos * 100)) if pos else 0
        daily_out.append({
            "day":         dr["day"],
            "pos":         round(pos, 2),
            "ext":         round(ext, 2),
            "matched":     matched,
            "total_lanes": total_lanes,
            "lane_pct":    lane_pct,
            "cov_pct":     cov_pct,
        })

    # Weekly rollup: W1=days 1-7, W2=8-14, W3=15-21, W4=22+
    week_data: dict = {}
    for d in daily_out:
        day_num = int(d["day"][8:10])
        week_n  = min((day_num - 1) // 7 + 1, 4)
        wk      = f"W{week_n}"
        if wk not in week_data:
            week_data[wk] = {"week": wk, "week_n": week_n,
                              "pos": 0.0, "ext": 0.0, "matched": 0, "lanes": 0}
        w = week_data[wk]
        w["pos"]     += d["pos"]
        w["ext"]     += d["ext"]
        w["matched"] += d["matched"]
        w["lanes"]   += d["total_lanes"]

    weekly_out = []
    for w in sorted(week_data.values(), key=lambda x: x["week_n"]):
        w["pos"]      = round(w["pos"], 2)
        w["ext"]      = round(w["ext"], 2)
        w["lane_pct"] = round(w["matched"] / w["lanes"] * 100) if w["lanes"] else 0
        w["cov_pct"]  = min(100, round(w["ext"] / w["pos"] * 100)) if w["pos"] else 0
        weekly_out.append(w)

    return {"session_id": session_id, "daily": daily_out, "weekly": weekly_out}


# ── Audit log ─────────────────────────────────────────────────────────────────

@app.get("/api/audit")
def get_audit_log(
    entity:     str | None = None,
    entity_id:  int | None = None,
    session_id: int | None = None,
    limit:      int        = 100,
):
    """Field-level change log.
    Pass session_id to get all events scoped to one session (session-level + its discrepancies).
    Pass entity + entity_id for a specific record's history."""
    cap = min(limit, 500)
    with get_db() as conn:
        if session_id and not entity and not entity_id:
            rows = conn.execute("""
                SELECT * FROM recon_audit_log
                WHERE (entity='session' AND entity_id=?)
                   OR (entity='discrepancy'
                       AND entity_id IN (SELECT id FROM recon_discrepancies WHERE session_id=?))
                ORDER BY ts DESC, id DESC LIMIT ?
            """, (session_id, session_id, cap)).fetchall()
        else:
            filters: list[str] = ["entity != 'auth'"]
            params:  list      = []
            if entity:
                filters.append("entity=?");    params.append(entity)
            if entity_id:
                filters.append("entity_id=?"); params.append(entity_id)
            where = "WHERE " + " AND ".join(filters)
            rows = conn.execute(f"""
                SELECT * FROM recon_audit_log {where}
                ORDER BY ts DESC, id DESC LIMIT ?
            """, params + [cap]).fetchall()
        return [dict(r) for r in rows]


# ── Diagnostic endpoints ──────────────────────────────────────────────────────
# Use these to inspect what's actually in the database before tuning matchers.

@app.get("/api/diagnostics/payment-codes", dependencies=[Depends(require_write_token)])
def diag_payment_codes(period_month: str, branch_id: int | None = None):
    """
    Show all PaymentCode values in pos_payment_month for a given month.
    Use this to verify keyword patterns match your actual Easy POS codes.
    Example: GET /api/diagnostics/payment-codes?period_month=2025-12
    """
    with get_db() as conn:
        try:
            filters = ["period_month=?"]
            params: list = [period_month]
            if branch_id:
                filters.append("branch_id=?"); params.append(branch_id)
            rows = conn.execute(f"""
                SELECT payment_code, SUM(amount) amount, SUM(txn_count) txn_count,
                       COUNT(DISTINCT branch_id) branches
                FROM pos_payment_month
                WHERE {' AND '.join(filters)}
                GROUP BY payment_code
                ORDER BY SUM(amount) DESC
            """, params).fetchall()
            return {"period_month": period_month, "codes": [dict(r) for r in rows]}
        except Exception as exc:
            return {"error": str(exc), "hint": "pos_payment_month not found — run sync_mssql.py first"}


@app.get("/api/diagnostics/customer-types", dependencies=[Depends(require_write_token)])
def diag_customer_types(period_month: str, branch_id: int | None = None):
    """
    Show all CustomerTypeName values in pos_customer_type_month for a given month.
    Use this to verify Grab/Lineman type names match your POS configuration.
    Example: GET /api/diagnostics/customer-types?period_month=2025-12
    """
    with get_db() as conn:
        try:
            filters = ["period_month=?"]
            params: list = [period_month]
            if branch_id:
                filters.append("branch_id=?"); params.append(branch_id)
            rows = conn.execute(f"""
                SELECT type_name, SUM(customers) customers, ROUND(SUM(net)/1.07) net_ex_vat,
                       COUNT(DISTINCT branch_id) branches
                FROM pos_customer_type_month
                WHERE {' AND '.join(filters)}
                GROUP BY type_name
                ORDER BY SUM(net) DESC
            """, params).fetchall()
            return {"period_month": period_month, "types": [dict(r) for r in rows]}
        except Exception as exc:
            return {"error": str(exc), "hint": "pos_customer_type_month not found — run sync_mssql.py first"}


@app.get("/api/diagnostics/excel-payment-cols", dependencies=[Depends(require_write_token)])
def diag_excel_payment_cols(period_month: str, branch_id: int | None = None):
    """
    Show Easy POS Excel payment column totals from sales_month_summary.
    cash/credit_card/qr are only populated after re-importing with the fixed ingest.
    Use to verify col 40 (qr field) actually contains QR/PromptPay data.
    """
    with get_db() as conn:
        filters = ["period_month=?", "source='Easy'"]
        params: list = [period_month]
        if branch_id:
            filters.append("branch_id=?"); params.append(branch_id)
        rows = conn.execute(f"""
            SELECT b.name branch, s.cash, s.credit_card, s.transfer, s.qr, s.coupon, s.net_sales
            FROM sales_month_summary s
            JOIN branches b ON b.branch_id = s.branch_id
            WHERE {' AND '.join(filters)}
            ORDER BY b.name
        """, params).fetchall()
        return {"period_month": period_month, "rows": [dict(r) for r in rows]}


# ── Health ────────────────────────────────────────────────────────────────────

@app.get("/health")
def health():
    with get_db() as conn:
        conn.execute("SELECT 1")
    return {"status": "ok"}
