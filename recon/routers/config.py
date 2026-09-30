"""Configuration the manager edits: payment map, method config, bank accounts,
acquirer MIDs, statement codes, account map, setup, branches.

Moved out of recon/main.py in Phase 7c (2026-09) by the AST: these endpoints and
the helpers only they use, unchanged. URLs are unchanged.
"""
from fastapi import APIRouter
from ..core import RECON_START_MONTH
from ..core import _BANKS
from ..core import _CATEGORIES
from ..core import _MASTER_ONLY
from ..core import _audit
from ..core import _override_channel
from ..core import _seed_method_defaults
from ..core import _seed_requirement_codes
from ..core import _statement_code_overrides
from ..core import _statement_codes
from ..core import _validate_master_branch
from ..core import _validate_period
from ..core import get_db
from ..core import require_manager_token
from ..core import require_write_token
from ..matcher import TOL_ABS
from ..matcher import TOL_ABS_DAY
from ..matcher import TOL_PCT
from ..matcher import TOL_PCT_DAY
from ..parsers.bank import _normalise_channel
from ..seed_payment_map import seed_payment_map
from fastapi import Depends
from fastapi import HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
import os
import sqlite3

router = APIRouter()


@router.get("/api/branches")
def list_branches():
    with get_db() as conn:
        rows = conn.execute(
            f"SELECT branch_id, name, branch_type FROM branches "
            f"WHERE {_MASTER_ONLY} ORDER BY name"
        ).fetchall()
        return [dict(r) for r in rows]


class PaymentMapUpsert(BaseModel):
    payment_code_norm: str | None = None   # required on create, ignored on update
    category:          str
    bank:              str = "NONE"
    reconcile:         int = 1
    notes:             str = ""
    updated_by:        str = ""


def _validate_map(body: PaymentMapUpsert):
    if body.category not in _CATEGORIES:
        raise HTTPException(400, f"category must be one of {sorted(_CATEGORIES)}")
    if body.bank not in _BANKS:
        raise HTTPException(400, f"bank must be one of {sorted(_BANKS)}")


@router.get("/api/payment-map")
def list_payment_map(needs_attention: int = 0):
    from ..seed_payment_map import normalize
    with get_db() as conn:
        sql = "SELECT * FROM recon_payment_map"
        if needs_attention:
            # Codes the daily engine can't lane: bank not set, or PROMPTPAY in a
            # category that reconciles per settling bank (matcher → unmapped_bank).
            sql += (" WHERE bank='UNKNOWN'"
                    " OR (bank='PROMPTPAY' AND category IN ('QR','TRANSFER','CREDIT'))")
        sql += " ORDER BY (bank='UNKNOWN') DESC, category, bank, payment_code_norm"
        rows = [dict(r) for r in conn.execute(sql)]
        # Attach POS volume per normalized code (master branches) so staff can see
        # which codes actually matter — e.g. KNNQR is the single biggest tender.
        amounts: dict = {}
        try:
            for r in conn.execute("""
                SELECT p.payment_code, SUM(p.amount) amt, SUM(p.txn_count) n
                FROM pos_payment_month p JOIN branches b ON b.branch_id = p.branch_id
                WHERE COALESCE(b.branch_type,'master')='master' AND b.active=1
                  AND p.period_month >= ?
                GROUP BY p.payment_code
            """, (RECON_START_MONTH,)):
                a = amounts.setdefault(normalize(r["payment_code"]), {"amount": 0.0, "txn": 0})
                a["amount"] += float(r["amt"] or 0)
                a["txn"]    += int(r["n"] or 0)
        except sqlite3.OperationalError:
            pass
        for row in rows:
            a = amounts.get(row["payment_code_norm"], {"amount": 0.0, "txn": 0})
            row["pos_amount"] = round(a["amount"], 2)
            row["pos_txn"]    = a["txn"]
        if needs_attention:
            # Only codes that actually carry money block reconciliation. Dormant
            # seed codes (฿0, 0 txns) are mappable but not urgent — drop them here
            # so the attention badge counts real work, not dead POS codes.
            rows = [r for r in rows if r["pos_txn"] > 0 or r["pos_amount"] > 0]
        return rows


@router.post("/api/payment-map", dependencies=[Depends(require_write_token)])
def create_payment_map(body: PaymentMapUpsert):
    _validate_map(body)
    code = (body.payment_code_norm or "").strip()
    if not code:
        raise HTTPException(400, "payment_code_norm is required")
    with get_db() as conn:
        try:
            cur = conn.execute("""
                INSERT INTO recon_payment_map (payment_code_norm, category, bank, reconcile, notes, updated_by)
                VALUES (?,?,?,?,?,?)
            """, (code, body.category, body.bank, body.reconcile, body.notes, body.updated_by or "manual"))
            conn.commit()
            return {"id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"payment_code_norm {code!r} already mapped")


@router.patch("/api/payment-map/{map_id}", dependencies=[Depends(require_write_token)])
def update_payment_map(map_id: int, body: PaymentMapUpsert):
    _validate_map(body)
    with get_db() as conn:
        old = conn.execute(
            "SELECT category, bank, reconcile FROM recon_payment_map WHERE id=?",
            (map_id,)).fetchone()
        conn.execute("""
            UPDATE recon_payment_map
            SET category=?, bank=?, reconcile=?, notes=?, updated_by=?, updated_at=datetime('now')
            WHERE id=?
        """, (body.category, body.bank, body.reconcile, body.notes, body.updated_by or "manual", map_id))
        if old:
            diffs = []
            if old["category"] != body.category:
                diffs.append({"field": "category", "old": old["category"], "new": body.category})
            if old["bank"] != body.bank:
                diffs.append({"field": "bank", "old": old["bank"], "new": body.bank})
            if old["reconcile"] != body.reconcile:
                diffs.append({"field": "reconcile", "old": old["reconcile"], "new": body.reconcile})
            if diffs:
                _audit(conn, "payment_map", map_id, "update", body.updated_by or "staff", diffs)
        conn.commit()
    return {"id": map_id}


@router.delete("/api/payment-map/{map_id}", dependencies=[Depends(require_write_token)])
def delete_payment_map(map_id: int):
    with get_db() as conn:
        conn.execute("DELETE FROM recon_payment_map WHERE id=?", (map_id,))
        conn.commit()
    return {"deleted": map_id}


@router.post("/api/payment-map/seed", dependencies=[Depends(require_write_token)])
def reseed_payment_map():
    with get_db() as conn:
        result = seed_payment_map(conn)
    return result


_METHODS = {"direct", "two_stage", "cash_3way", "itemized", "issuer_points"}


_DOC1S   = {"deposit_slip", "cc_summary", "grab_report", "lineman_report", "points_summary"}


_DOC2S   = {"bank_stmt_KTB", "bank_stmt_KBANK", "bank_stmt_*"}


class MethodConfig(BaseModel):
    category:         str
    bank:             str = "*"
    doc1:             str | None = None
    doc2:             str | None = None
    method:           str = "direct"
    tol_abs:          float | None = None
    tol_abs_day:      float | None = None
    lag_days:         int | None = None
    merge_days:       int | None = None
    fee_vat_expected: int = 0
    notes:            str = ""
    updated_by:       str = ""


def _validate_method_config(body: MethodConfig):
    if body.category not in _CATEGORIES:
        raise HTTPException(400, f"category must be one of {sorted(_CATEGORIES)}")
    if body.bank != "*" and body.bank not in _BANKS:
        raise HTTPException(400, f"bank must be '*' or one of {sorted(_BANKS)}")
    if body.method not in _METHODS:
        raise HTTPException(400, f"method must be one of {sorted(_METHODS)}")
    if body.doc1 and body.doc1 not in _DOC1S:
        raise HTTPException(400, f"doc1 must be one of {sorted(_DOC1S)}")
    if body.doc2 and body.doc2 not in _DOC2S:
        raise HTTPException(400, f"doc2 must be one of {sorted(_DOC2S)}")
    # Clamp parameters server-side — a typo must not detonate the matcher.
    if body.tol_abs     is not None: body.tol_abs     = max(0.0, min(100_000.0, body.tol_abs))
    if body.tol_abs_day is not None: body.tol_abs_day = max(0.0, min(100_000.0, body.tol_abs_day))
    if body.lag_days    is not None: body.lag_days    = max(0,   min(14,        body.lag_days))
    if body.merge_days  is not None: body.merge_days  = max(1,   min(7,         body.merge_days))


@router.get("/api/method-config")
def list_method_config():
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM recon_method_config ORDER BY category, bank")]


@router.post("/api/method-config", dependencies=[Depends(require_manager_token)])
def create_method_config(body: MethodConfig):
    _validate_method_config(body)
    with get_db() as conn:
        try:
            cur = conn.execute("""
                INSERT INTO recon_method_config
                    (category, bank, doc1, doc2, method, tol_abs, tol_abs_day,
                     lag_days, merge_days, fee_vat_expected, notes, updated_by)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
            """, (body.category, body.bank, body.doc1, body.doc2, body.method,
                  body.tol_abs, body.tol_abs_day, body.lag_days, body.merge_days,
                  body.fee_vat_expected, body.notes or None, body.updated_by or "manager"))
            _audit(conn, "method_config", cur.lastrowid, "create", body.updated_by or "manager")
            conn.commit()
            return {"id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"lane {body.category}/{body.bank} already configured")


@router.patch("/api/method-config/{cfg_id}", dependencies=[Depends(require_manager_token)])
def update_method_config(cfg_id: int, body: MethodConfig):
    _validate_method_config(body)
    with get_db() as conn:
        old = conn.execute("SELECT * FROM recon_method_config WHERE id=?", (cfg_id,)).fetchone()
        if not old:
            raise HTTPException(404, f"method config {cfg_id} not found")
        conn.execute("""
            UPDATE recon_method_config
            SET category=?, bank=?, doc1=?, doc2=?, method=?, tol_abs=?, tol_abs_day=?,
                lag_days=?, merge_days=?, fee_vat_expected=?, notes=?,
                updated_by=?, updated_at=datetime('now')
            WHERE id=?
        """, (body.category, body.bank, body.doc1, body.doc2, body.method,
              body.tol_abs, body.tol_abs_day, body.lag_days, body.merge_days,
              body.fee_vat_expected, body.notes or None,
              body.updated_by or "manager", cfg_id))
        diffs = [{"field": f, "old": old[f], "new": getattr(body, f)}
                 for f in ("doc1", "doc2", "method", "tol_abs", "tol_abs_day",
                           "lag_days", "merge_days", "fee_vat_expected")
                 if old[f] != getattr(body, f)]
        if diffs:
            _audit(conn, "method_config", cfg_id, "update", body.updated_by or "manager", diffs)
        conn.commit()
    return {"id": cfg_id}


@router.delete("/api/method-config/{cfg_id}", dependencies=[Depends(require_manager_token)])
def delete_method_config(cfg_id: int):
    with get_db() as conn:
        cur = conn.execute("DELETE FROM recon_method_config WHERE id=?", (cfg_id,))
        if cur.rowcount == 0:
            raise HTTPException(404, f"method config {cfg_id} not found")
        _audit(conn, "method_config", cfg_id, "delete", "manager")
        conn.commit()
    return {"deleted": cfg_id}


# Thai label per method for the master table (the spreadsheet's กระทบยอด column)
_METHOD_TH = {
    "direct":        "เทียบยอดตรง POS ↔ เอกสาร",
    "two_stage":     "2 ขั้น: สรุปยอด↔สเตทเมนต์ แล้ว gross↔POS",
    "cash_3way":     "3 ทาง: POS ↔ สลิปนำฝาก ↔ สเตทเมนต์",
    "itemized":      "รายตามรายการในสเตทเมนต์ (วันเดียวกัน/วันถัดไป)",
    "issuer_points": "แต้มบัตร: ใบสรุปคะแนน ↔ สเตทเมนต์",
}


def _doc_presence(conn, period_month: str) -> dict:
    """Company-wide 'is this document type in for the month?' map for the master
    table's doc chips: {doc_key: True/False}."""
    have: dict = {}
    try:
        for r in conn.execute(
            "SELECT DISTINCT source_type, provider FROM recon_uploads "
            "WHERE period_month=? AND status='confirmed'", (period_month,)):
            st, prov = r["source_type"], (r["provider"] or "")
            if st == "bank":
                from ..matcher import _provider_bank
                have[f"bank_stmt_{_provider_bank(prov)}"] = True
                have["bank_stmt_*"] = True
            elif st == "credit_card":
                have["cc_summary"] = True
                if "POINT" in prov.upper():
                    have["points_summary"] = True
            elif st == "grab":
                have["grab_report"] = True
            elif st == "lineman":
                have["lineman_report"] = True
        if conn.execute("SELECT 1 FROM recon_deposit_slip WHERE substr(sale_day,1,7)=? LIMIT 1",
                        (period_month,)).fetchone():
            have["deposit_slip"] = True
    except sqlite3.OperationalError:
        pass
    return have


@router.get("/api/setup/master")
def setup_master(period_month: str):
    """Everything the advanced Setup page renders in one call: the master
    payment-code table (the owner's 'ทุกรหัสในตารางเดียว'), joined with each
    lane's method config + POS volume + document presence, plus the raw
    method_configs / bank_accounts / acquirer_mids and the global tolerances."""
    _validate_period(period_month)
    rows = list_payment_map()          # includes pos_amount / pos_txn
    with get_db() as conn:
        cfgs = {(r["category"], r["bank"]): dict(r) for r in conn.execute(
            "SELECT * FROM recon_method_config")}
        have_docs = _doc_presence(conn, period_month)
        accounts = [dict(r) for r in conn.execute(
            """SELECT a.*, b.name AS branch_name FROM recon_bank_account a
               JOIN branches b ON b.branch_id=a.branch_id ORDER BY b.name, a.bank""")]
        mids = [dict(r) for r in conn.execute(
            """SELECT m.*, b.name AS branch_name FROM recon_acquirer_mid m
               JOIN branches b ON b.branch_id=m.branch_id ORDER BY b.name, m.acquirer""")]

    out_rows = []
    for r in rows:
        mc = cfgs.get((r["category"], r["bank"])) or cfgs.get((r["category"], "*")) or {}
        needs_bank = (r["bank"] == "UNKNOWN" or
                      (r["bank"] == "PROMPTPAY" and r["category"] in ("QR", "TRANSFER", "CREDIT")))
        status = ("off" if not r["reconcile"]
                  else "needs_bank" if needs_bank
                  else "dormant" if (r["pos_amount"] == 0 and r["pos_txn"] == 0)
                  else "ok")
        out_rows.append({
            **r,
            "method": {
                "id":               mc.get("id"),
                "doc1":             mc.get("doc1"),
                "doc2":             mc.get("doc2"),
                "method":           mc.get("method"),
                "method_label":     _METHOD_TH.get(mc.get("method"), mc.get("method")),
                "tol_abs":          mc.get("tol_abs"),
                "tol_abs_day":      mc.get("tol_abs_day"),
                "lag_days":         mc.get("lag_days"),
                "merge_days":       mc.get("merge_days"),
                "fee_vat_expected": mc.get("fee_vat_expected", 0),
                "lane_bank":        mc.get("bank"),   # '*' = applies to the whole category
            },
            "doc_status": {
                "doc1": (None if not mc.get("doc1")
                         else "uploaded" if have_docs.get(mc["doc1"]) else "missing"),
                "doc2": (None if not mc.get("doc2")
                         else "uploaded" if have_docs.get(mc["doc2"]) else "missing"),
            },
            "status": status,
        })
    return {
        "period_month":   period_month,
        "rows":           out_rows,
        "method_configs": [dict(v) for v in sorted(cfgs.values(), key=lambda x: (x["category"], x["bank"]))],
        "bank_accounts":  accounts,
        "acquirer_mids":  mids,
        "globals": {"TOL_ABS": TOL_ABS, "TOL_PCT": TOL_PCT,
                    "TOL_ABS_DAY": TOL_ABS_DAY, "TOL_PCT_DAY": TOL_PCT_DAY},
        "manager_required": bool((os.environ.get("OSHINEI_MANAGER_TOKEN") or "").strip()),
    }


@router.post("/api/setup/seed-requirement", dependencies=[Depends(require_manager_token)])
def seed_requirement():
    """Re-apply the owner's requirement table (codes + lane methods) on demand.
    Idempotent; never overwrites staff/manager edits."""
    with get_db() as conn:
        codes = _seed_requirement_codes(conn)
        methods = _seed_method_defaults(conn)
        conn.commit()
    return {"codes": codes, "methods": methods}


class BankAccount(BaseModel):
    branch_id:    int
    bank:         str
    account_no:   str
    account_name: str = ""
    active:       int = 1


@router.get("/api/bank-accounts")
def list_bank_accounts(branch_id: int | None = None):
    with get_db() as conn:
        sql = """SELECT a.*, b.name AS branch_name FROM recon_bank_account a
                 JOIN branches b ON b.branch_id = a.branch_id"""
        params: list = []
        if branch_id:
            sql += " WHERE a.branch_id=?"; params.append(branch_id)
        sql += " ORDER BY b.name, a.bank"
        return [dict(r) for r in conn.execute(sql, params)]


@router.post("/api/bank-accounts", dependencies=[Depends(require_write_token)])
def create_bank_account(body: BankAccount):
    if body.bank not in ("KTB", "KBANK"):
        raise HTTPException(400, "bank must be KTB or KBANK")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        try:
            cur = conn.execute("""
                INSERT INTO recon_bank_account (branch_id, bank, account_no, account_name, active)
                VALUES (?,?,?,?,?)
            """, (body.branch_id, body.bank, body.account_no, body.account_name, body.active))
            conn.commit()
            return {"account_id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"account {body.bank}/{body.account_no} already exists")


@router.patch("/api/bank-accounts/{account_id}", dependencies=[Depends(require_write_token)])
def update_bank_account(account_id: int, body: BankAccount):
    if body.bank not in ("KTB", "KBANK"):
        raise HTTPException(400, "bank must be KTB or KBANK")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        cur = conn.execute("""
            UPDATE recon_bank_account SET branch_id=?, bank=?, account_no=?, account_name=?, active=?
            WHERE account_id=?
        """, (body.branch_id, body.bank, body.account_no, body.account_name, body.active, account_id))
        if cur.rowcount == 0:
            raise HTTPException(404, f"bank account {account_id} not found")
        conn.commit()
    return {"account_id": account_id}


@router.delete("/api/bank-accounts/{account_id}", dependencies=[Depends(require_write_token)])
def delete_bank_account(account_id: int):
    with get_db() as conn:
        cur = conn.execute("DELETE FROM recon_bank_account WHERE account_id=?", (account_id,))
        if cur.rowcount == 0:
            raise HTTPException(404, f"bank account {account_id} not found")
        conn.commit()
    return {"deleted": account_id}


class AcquirerMid(BaseModel):
    branch_id:   int
    acquirer:    str
    merchant_id: str = ""
    terminal_id: str = ""
    active:      int = 1


@router.get("/api/acquirer-mids")
def list_acquirer_mids(branch_id: int | None = None):
    with get_db() as conn:
        sql = """SELECT m.*, b.name AS branch_name FROM recon_acquirer_mid m
                 JOIN branches b ON b.branch_id = m.branch_id"""
        params: list = []
        if branch_id:
            sql += " WHERE m.branch_id=?"; params.append(branch_id)
        sql += " ORDER BY b.name, m.acquirer"
        return [dict(r) for r in conn.execute(sql, params)]


@router.post("/api/acquirer-mids", dependencies=[Depends(require_write_token)])
def create_acquirer_mid(body: AcquirerMid):
    if body.acquirer not in ("KTB", "KBANK", "KTC", "BAY", "SCB", "BBL", "KTCPOINT", "BAYPOINT"):
        raise HTTPException(400, "acquirer must be KTB, KBANK, KTC, BAY, SCB, BBL, KTCPOINT or BAYPOINT")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        cur = conn.execute("""
            INSERT INTO recon_acquirer_mid (branch_id, acquirer, merchant_id, terminal_id, active)
            VALUES (?,?,?,?,?)
        """, (body.branch_id, body.acquirer, body.merchant_id, body.terminal_id, body.active))
        conn.commit()
        return {"mid_id": cur.lastrowid}


@router.patch("/api/acquirer-mids/{mid_id}", dependencies=[Depends(require_write_token)])
def update_acquirer_mid(mid_id: int, body: AcquirerMid):
    if body.acquirer not in ("KTB", "KBANK", "KTC", "BAY", "SCB", "BBL", "KTCPOINT", "BAYPOINT"):
        raise HTTPException(400, "acquirer must be KTB, KBANK, KTC, BAY, SCB, BBL, KTCPOINT or BAYPOINT")
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        cur = conn.execute("""
            UPDATE recon_acquirer_mid SET branch_id=?, acquirer=?, merchant_id=?, terminal_id=?, active=?
            WHERE mid_id=?
        """, (body.branch_id, body.acquirer, body.merchant_id, body.terminal_id, body.active, mid_id))
        if cur.rowcount == 0:
            raise HTTPException(404, f"acquirer MID {mid_id} not found")
        conn.commit()
    return {"mid_id": mid_id}


@router.delete("/api/acquirer-mids/{mid_id}", dependencies=[Depends(require_write_token)])
def delete_acquirer_mid(mid_id: int):
    with get_db() as conn:
        cur = conn.execute("DELETE FROM recon_acquirer_mid WHERE mid_id=?", (mid_id,))
        if cur.rowcount == 0:
            raise HTTPException(404, f"acquirer MID {mid_id} not found")
        conn.commit()
    return {"deleted": mid_id}


_OVERRIDE_CHANNELS = {"QR", "TRANSFER", "CASH_DEP", "CC_SETTLE",
                      "GRAB_PAYOUT", "LINEMAN_PAYOUT", "FEE", "OTHER"}


def _reapply_statement_overrides(conn, code: str) -> int:
    """Re-classify stored bank lines whose description contains `code`, from the
    parser base + current overrides. Used on save (apply) and delete (revert)."""
    overrides = _statement_code_overrides(conn)
    like = f"%{code.upper()}%"
    rows = conn.execute(
        "SELECT id, description FROM recon_bank_lines "
        "WHERE UPPER(COALESCE(description,'')) LIKE ?", (like,)).fetchall()
    n = 0
    for r in rows:
        base = _normalise_channel(r["description"])
        newch = _override_channel(r["description"], base, overrides)
        conn.execute("UPDATE recon_bank_lines SET channel=? WHERE id=?", (newch, r["id"]))
        n += 1
    return n


class CodeOverride(BaseModel):
    code: str
    channel: str
    note: str | None = None
    by: str | None = None


@router.get("/api/statement-codes")
def statement_codes():
    with get_db() as conn:
        codes = _statement_codes(conn)
    return {"codes": codes,
            "review_count": sum(1 for c in codes if c["status"] == "review"),
            "override_count": sum(1 for c in codes if c["overridden"]),
            "channel_options": sorted(_OVERRIDE_CHANNELS)}


@router.post("/api/statement-codes/override", dependencies=[Depends(require_manager_token)])
def set_statement_override(body: CodeOverride):
    code = (body.code or "").strip().upper()
    if not code:
        raise HTTPException(400, "ต้องระบุรหัส")
    if body.channel not in _OVERRIDE_CHANNELS:
        raise HTTPException(400, f"ช่องทางต้องเป็นหนึ่งใน {sorted(_OVERRIDE_CHANNELS)}")
    with get_db() as conn:
        conn.execute(
            "INSERT INTO recon_statement_code_map (code, channel, note, updated_by, updated_at) "
            "VALUES (?,?,?,?,datetime('now')) "
            "ON CONFLICT(code) DO UPDATE SET channel=excluded.channel, note=excluded.note, "
            "updated_by=excluded.updated_by, updated_at=datetime('now')",
            (code, body.channel, body.note or "", body.by or "manager"))
        n = _reapply_statement_overrides(conn, code)
        _audit(conn, "statement_code", 0, "override", body.by or "manager",
               [{"field": code, "old": None, "new": body.channel}])
        conn.commit()
    return {"code": code, "channel": body.channel, "lines_reclassified": n}


@router.delete("/api/statement-codes/override/{code}", dependencies=[Depends(require_manager_token)])
def del_statement_override(code: str):
    code_u = (code or "").strip().upper()
    with get_db() as conn:
        conn.execute("DELETE FROM recon_statement_code_map WHERE code=?", (code_u,))
        n = _reapply_statement_overrides(conn, code_u)   # revert lines to base classification
        _audit(conn, "statement_code", 0, "revert", "manager",
               [{"field": code_u, "old": None, "new": None}])
        conn.commit()
    return {"code": code_u, "removed": True, "lines_reclassified": n}


@router.get("/api/statement-codes.xlsx")
def statement_codes_xlsx():
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); bold = Font(bold=True)
    warn = PatternFill("solid", fgColor="FFF3CD")
    ws = wb.active; ws.title = "รหัสในสเตทเมนต์"
    ws.append(["รหัส / คำสำคัญ", "ชนิด", "ธนาคาร", "ความหมาย (ไทย)", "ระบบจัดเป็นช่องทาง",
               "สถานะ", "หมายเหตุระบบ", "ทีมยืนยัน (✓/✗)", "ความเห็น/แก้ไขของทีม"])
    for c in ws[1]:
        c.font = bold
    with get_db() as conn:
        codes = _statement_codes(conn)
    for c in codes:
        kind = "รหัสธนาคาร" if c["kind"] == "code" else "คำทั่วไป"
        st = ("ทีมแก้แล้ว" if c["overridden"]
              else "ยืนยันแล้ว" if c["status"] == "confirmed" else "⚠ รอตรวจ")
        ch_th = c.get("effective_channel_th", c["channel_th"])
        r = [c["code"], kind, c["bank"], c["meaning_th"], ch_th, st, c["note"], "", ""]
        ws.append(r)
        if c["status"] == "review":
            for cell in ws[ws.max_row]:
                cell.fill = warn
    for col, w in zip("ABCDEFGHI", (30, 12, 8, 46, 24, 14, 50, 14, 36)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=\"statement_codes.xlsx\""})


@router.get("/api/statement-codes.html")
def statement_codes_html():
    """Printable / shareable review sheet of every statement code the parser knows,
    grouped by channel, with a column for the team to confirm or correct."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    with get_db() as conn:
        codes = _statement_codes(conn)
    review_n = sum(1 for c in codes if c["status"] == "review")
    css = """<style>
      @page{size:A4 landscape;margin:12mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12px;margin:0;padding:24px;background:#f8fafc}
      .sheet{background:#fff;max-width:1100px;margin:0 auto;padding:26px 30px;border-radius:4px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:20px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:14px}
      table{width:100%;border-collapse:collapse;font-size:11.5px}
      th,td{padding:6px 9px;border-bottom:1px solid #eef2f7;text-align:left;vertical-align:top}
      th{color:#475569;font-size:10.5px;border-bottom:1.5px solid #e2e8f0;background:#f8fafc}
      code{background:#f1f5f9;padding:1px 6px;border-radius:4px;font-family:ui-monospace,monospace;font-weight:600}
      .badge{padding:1px 8px;border-radius:99px;font-size:10px;font-weight:600;white-space:nowrap}
      .b-ok{background:#dcfce7;color:#166534} .b-warn{background:#fef3c7;color:#854d0e}
      .b-edit{background:#dbeafe;color:#1e40af}
      .b-code{background:#e0e7ff;color:#3730a3} .b-kw{background:#f1f5f9;color:#475569}
      tr.review{background:#fffbeb} .note{color:#64748b;font-size:11px}
      .fill{color:#cbd5e1} .toolbar{max-width:1100px;margin:0 auto 14px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:4px;padding:9px 18px;font-size:13px;cursor:pointer;text-decoration:none}
      .lead{background:#fff7ed;border:1px solid #fed7aa;border-radius:4px;padding:10px 14px;margin-bottom:14px;font-size:12px;line-height:1.6}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;border-radius:0;max-width:none}}
    </style>"""
    rows = []
    for c in codes:
        kind = ('<span class="badge b-code">รหัสธนาคาร</span>' if c["kind"] == "code"
                else '<span class="badge b-kw">คำทั่วไป</span>')
        st = ('<span class="badge b-edit">ทีมแก้แล้ว</span>' if c["overridden"]
              else '<span class="badge b-ok">ยืนยันแล้ว</span>' if c["status"] == "confirmed"
              else '<span class="badge b-warn">⚠ รอตรวจ</span>')
        ch_th = c.get("effective_channel_th", c["channel_th"])
        ch_cell = (f'{esc(ch_th)} <span class="note">(เดิม {esc(c["channel_th"])})</span>'
                   if c["overridden"] else esc(ch_th))
        rows.append(
            f'<tr class="{ "review" if c["status"]=="review" and not c["overridden"] else "" }">'
            f'<td><code>{esc(c["code"])}</code></td><td>{kind}</td><td>{esc(c["bank"])}</td>'
            f'<td>{esc(c["meaning_th"])}{("<div class=note>"+esc(c["note"])+"</div>") if c["note"] else ""}</td>'
            f'<td>{ch_cell}</td><td>{st}</td>'
            f'<td class="fill">▢ ใช่ &nbsp; ▢ ไม่ใช่</td><td class="fill">…………………………</td></tr>')
    html = (
        "<!doctype html><html lang='th'><head><meta charset='utf-8'>"
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>"
        "<title>พจนานุกรมรหัสในสเตทเมนต์</title>" + css + "</head><body>"
        "<div class='toolbar'>"
        "<a class='btn' href='/api/statement-codes.xlsx' style='text-decoration:none;margin-right:8px;background:#16794f'>ดาวน์โหลด Excel</a>"
        "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>"
        "<div class='sheet'>"
        "<h1>พจนานุกรมรหัสในสเตทเมนต์ธนาคาร</h1>"
        f"<div class='sub'>ระบบใช้ตารางนี้แยกประเภทรายการในสเตทเมนต์อัตโนมัติ · {len(codes)} รายการ · "
        f"<b>{review_n} รายการรอทีมตรวจ</b></div>"
        "<div class='lead'>ช่วยตรวจ <b>คอลัมน์ \"ความหมาย\" กับ \"ระบบจัดเป็นช่องทาง\"</b> ว่าตรงกับความจริงไหม "
        "โดยเฉพาะแถวสีเหลือง (⚠ รอตรวจ) ที่เป็นรหัสเฉพาะของธนาคาร เช่น <code>SDCH</code> "
        "(เจ้าของยืนยันว่า = นำเงินสดไปฝาก ไม่ใช่ QR) · ทำเครื่องหมายในช่อง \"ทีมยืนยัน\" แล้วเขียนสิ่งที่ถูกต้องในช่องสุดท้าย</div>"
        "<table><thead><tr><th>รหัส / คำสำคัญ</th><th>ชนิด</th><th>ธนาคาร</th>"
        "<th>ความหมาย</th><th>ระบบจัดเป็นช่องทาง</th><th>สถานะ</th>"
        "<th>ทีมยืนยัน</th><th>ที่ถูกต้องคือ</th></tr></thead><tbody>"
        + "".join(rows) +
        "</tbody></table>"
        "<div style='margin-top:22px;display:flex;gap:50px'>"
        "<div style='border-top:1px solid #94a3b8;padding-top:5px;flex:1;text-align:center;font-size:11px;color:#475569'>ผู้ตรวจสอบ (ทีมบัญชี)</div>"
        "<div style='border-top:1px solid #94a3b8;padding-top:5px;flex:1;text-align:center;font-size:11px;color:#475569'>วันที่</div>"
        "</div></div></body></html>")
    return Response(content=html, media_type="text/html; charset=utf-8")


class AccountMap(BaseModel):
    slot:         str
    account_code: str | None = None
    account_name: str | None = None
    updated_by:   str | None = None


@router.get("/api/account-map")
def account_map_list():
    with get_db() as conn:
        try:
            rows = conn.execute(
                "SELECT slot, account_code, account_name FROM recon_account_map ORDER BY slot").fetchall()
        except sqlite3.OperationalError:
            return []
        return [dict(r) for r in rows]


@router.put("/api/account-map", dependencies=[Depends(require_write_token)])
def account_map_upsert(body: AccountMap):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO recon_account_map(slot,account_code,account_name,updated_by) "
            "VALUES(?,?,?,?) ON CONFLICT(slot) DO UPDATE SET account_code=excluded.account_code, "
            "account_name=excluded.account_name, updated_at=datetime('now'), updated_by=excluded.updated_by",
            (body.slot, body.account_code, body.account_name, body.updated_by or "staff"))
        conn.commit()
    return {"slot": body.slot}
