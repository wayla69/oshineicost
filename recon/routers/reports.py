"""Audit reports: variance, audit trail, close register, open items, overrides,
wrong-channel, certificate, auditor pack — JSON, print and workbook forms.

Moved out of recon/main.py in Phase 7c (2026-09) by the AST: these endpoints and
the helpers only they use, unchanged. URLs are unchanged.
"""
from fastapi import APIRouter
from ..core import _BANK_CH_TO_CAT
from ..core import _MASTER_ONLY
from ..core import _MATCH_CH_TH
from ..core import _PERIOD_RE
from ..core import _actor
from ..core import _validate_master_branch
from ..core import _validate_period
from ..core import get_db
from ..matcher import _BANK_LABEL
from ..matcher import payout_reconciliation
from ..matcher import pos_codes_scope
from ..reconcile_v2 import _CH_LABEL
from ..reconcile_v2 import reconcile_month
from ..reconcile_v2 import reconcile_overview
from datetime import datetime as _dt
from fastapi import HTTPException
from fastapi import Request
from fastapi.responses import Response
import json
import sqlite3

router = APIRouter()


# Text a spreadsheet would read as a formula (OWASP CSV/formula injection). Notes and
# reasons are typed by staff and reach the Auditor Pack, so a cell starting with one
# of these is forced to plain text (quote-prefixed) — strings only, numbers untouched.
_XL_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _xl_append(ws, row):
    """ws.append(row), then neutralise any str cell that would act as a formula."""
    ws.append(row)
    for c in ws[ws.max_row]:
        v = c.value
        if isinstance(v, str) and v.startswith(_XL_FORMULA_LEAD):
            c.data_type = "s"
            c.quotePrefix = True


def _branch_error_report(conn, period_month: str | None = None) -> dict:
    from collections import defaultdict
    q = ("SELECT branch_id, period_month, category, note, created_by, created_at "
         "FROM recon_lane_adj WHERE adj_type='accept'")
    params: list = []
    if period_month:
        q += " AND period_month=?"
        params.append(period_month)
    accepts = conn.execute(q + " ORDER BY branch_id, period_month, category", params).fetchall()
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}

    var_cache: dict = {}        # (branch, period) -> {channel_key: variance}
    items = []
    for a in accepts:
        key = (a["branch_id"], a["period_month"])
        if key not in var_cache:
            try:
                rm = reconcile_month(conn, a["branch_id"], a["period_month"])
                var_cache[key] = {c["key"]: c["variance"] for c in rm["channels"]}
            except Exception:
                var_cache[key] = {}
        items.append({
            "branch_id":    a["branch_id"],
            "branch_name":  bnames.get(a["branch_id"], str(a["branch_id"])),
            "period_month": a["period_month"],
            "category":     a["category"],
            "variance":     round(var_cache[key].get(a["category"], 0.0), 2),
            "reason":       a["note"] or "",
            "by":           a["created_by"] or "",
            "at":           a["created_at"] or "",
        })

    agg: dict = defaultdict(lambda: {"count": 0, "total_abs_variance": 0.0, "months": set()})
    for it in items:
        g = agg[(it["branch_id"], it["branch_name"])]
        g["count"] += 1
        g["total_abs_variance"] += abs(it["variance"])
        g["months"].add(it["period_month"])
    summary = sorted(
        [{"branch_id": bid, "branch_name": bn, "count": v["count"],
          "months": len(v["months"]), "total_abs_variance": round(v["total_abs_variance"], 2)}
         for (bid, bn), v in agg.items()],
        key=lambda x: -x["total_abs_variance"])
    return {"period_month": period_month, "summary": summary, "items": items}


@router.get("/api/reports/branch-errors")
def branch_error_report(period_month: str | None = None):
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        return _branch_error_report(conn, period_month)


@router.get("/api/reports/branch-errors.xlsx")
def branch_error_report_xlsx(period_month: str | None = None):
    import io
    import openpyxl
    from openpyxl.styles import Font
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _branch_error_report(conn, period_month)
    wb = openpyxl.Workbook(); bold = Font(bold=True)
    ws = wb.active; ws.title = "สรุปรายสาขา"
    ws.append(["สาขา", "จำนวนครั้งที่รับทราบ", "จำนวนเดือน", "ยอดผลต่างรวม (abs)"])
    for c in ws[1]:
        c.font = bold
    for s in rep["summary"]:
        _xl_append(ws, [s["branch_name"], s["count"], s["months"], s["total_abs_variance"]])
    ws2 = wb.create_sheet("รายการ")
    ws2.append(["สาขา", "เดือน", "ช่องทาง", "ผลต่าง", "เหตุผล", "รับทราบโดย", "เมื่อ"])
    for c in ws2[1]:
        c.font = bold
    for it in rep["items"]:
        _xl_append(ws2, [it["branch_name"], it["period_month"], it["category"],
                         it["variance"], it["reason"], it["by"], it["at"]])
    for w in (ws, ws2):
        for i, _h in enumerate(w[1], 1):
            w.column_dimensions[chr(64 + i)].width = 22
    buf = io.BytesIO(); wb.save(buf)
    fname = f"branch_errors_{period_month or 'all'}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{fname}\""})


_EXT_TABLE_CAT = {
    "recon_cc_lines": "CREDIT", "recon_grab_lines": "GRAB",
    "recon_lineman_lines": "LINEMAN",
}


def _real_channel_of_match(conn, match_id: int) -> str | None:
    """Where the money actually landed, inferred from the matched statement lines
    (dominant channel by baht). None when no lines are attached."""
    cats: dict = {}
    for ln in conn.execute(
        "SELECT ext_table, ext_id, amount FROM recon_manual_match_line WHERE match_id=?",
        (match_id,)):
        tbl = ln["ext_table"]; amt = abs(float(ln["amount"] or 0))
        if tbl == "recon_bank_lines":
            row = conn.execute("SELECT channel FROM recon_bank_lines WHERE id=?",
                               (ln["ext_id"],)).fetchone()
            cat = _BANK_CH_TO_CAT.get((row["channel"] or "") if row else "", "OTHER")
        else:
            cat = _EXT_TABLE_CAT.get(tbl, "OTHER")
        cats[cat] = cats.get(cat, 0.0) + amt
    return max(cats, key=cats.get) if cats else None


def _wrong_channel_items(conn, period_month: str | None = None) -> list[dict]:
    """Confirmed wrong-channel mis-entries (true cross-channel manual matches)."""
    q = ("SELECT match_id, branch_id, category, bank, period_month, pos_amount, "
         "ext_amount, pos_days, note, matched_by, matched_at FROM recon_manual_match")
    params: list = []
    if period_month:
        q += " WHERE period_month=?"; params.append(period_month)
    rows = conn.execute(q + " ORDER BY branch_id, matched_at", params).fetchall()
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}
    items = []
    for m in rows:
        to_cat = _real_channel_of_match(conn, m["match_id"])
        from_cat = m["category"]
        if not to_cat or to_cat == from_cat:        # only true cross-channel mis-entries
            continue
        try:
            days = sorted({d.get("day") for d in json.loads(m["pos_days"] or "[]") if d.get("day")})
        except Exception:
            days = []
        amount = round(max(float(m["pos_amount"] or 0), float(m["ext_amount"] or 0)), 2)
        items.append({
            "match_id":     m["match_id"],
            "source":       "match",
            "branch_id":    m["branch_id"],
            "branch_name":  bnames.get(m["branch_id"], str(m["branch_id"])),
            "period_month": m["period_month"],
            "days":         days,
            "from_cat":     from_cat,  "from_th": _CH_LABEL.get(from_cat, from_cat),
            "to_cat":       to_cat,    "to_th":   _CH_LABEL.get(to_cat, to_cat),
            "to_bank_th":   _BANK_LABEL.get(m["bank"], "") if m["bank"] in _BANK_LABEL else "",
            "amount":       amount,
            "reason":       m["note"] or "",
            "by":           m["matched_by"] or "",
            "at":           m["matched_at"] or "",
        })
    return items


def _channel_move_items(conn, period_month: str | None = None,
                        branch_id: int | None = None) -> list[dict]:
    """Manual cross-channel MOVES (recon_channel_move), shaped like _wrong_channel_items
    so the wrong-channel report lists hand-entered corrections alongside detected ones.
    Tagged source='move' (the detected ones are source='match')."""
    q = ("SELECT id, branch_id, period_month, from_category, from_bank, to_category, "
         "to_bank, amount, note, created_by, created_at FROM recon_channel_move")
    conds, params = [], []
    if period_month:
        conds.append("period_month=?"); params.append(period_month)
    if branch_id:
        conds.append("branch_id=?"); params.append(branch_id)
    if conds:
        q += " WHERE " + " AND ".join(conds)
    try:
        rows = conn.execute(q + " ORDER BY branch_id, created_at", params).fetchall()
    except sqlite3.OperationalError:
        return []
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}
    items = []
    for m in rows:
        fc, tc = m["from_category"], m["to_category"]
        items.append({
            "match_id":     None,
            "move_id":      m["id"],
            "source":       "move",
            "branch_id":    m["branch_id"],
            "branch_name":  bnames.get(m["branch_id"], str(m["branch_id"])),
            "period_month": m["period_month"],
            "days":         [],
            "from_cat":     fc, "from_th": _CH_LABEL.get(fc, fc),
            "to_cat":       tc, "to_th":   _CH_LABEL.get(tc, tc),
            "to_bank_th":   _BANK_LABEL.get(m["to_bank"], "") if m["to_bank"] in _BANK_LABEL else "",
            "amount":       round(float(m["amount"] or 0), 2),
            "reason":       m["note"] or "",
            "by":           m["created_by"] or "",
            "at":           m["created_at"] or "",
        })
    return items


# Concrete coaching guidance per (rung-as → really-was) channel confusion. Plain Thai,
# phrased as "what to do next time", so the sheet helps the branch improve — not blame.
_TIP_BY_PAIR = {
    ("CASH", "QR"):       "ลูกค้าสแกนจ่าย QR แต่ถูกกดเป็นเงินสด — เวลารับ QR ให้กดปุ่ม “QR พร้อมเพย์” และดูยอดเด้งบนแอป/สลิปให้ตรงก่อนกดยืนยัน",
    ("QR", "CASH"):       "เงินสดถูกกดเป็น QR — รับเงินสดให้กดปุ่ม “เงินสด” ทุกครั้ง อย่ากดค้างจากบิลก่อน",
    ("CASH", "TRANSFER"): "เงินโอนเข้าบัญชีถูกกดเป็นเงินสด — ถ้าลูกค้าโอน ให้กดปุ่ม “เงินโอน” และยืนยันยอดเข้าบัญชีก่อนส่งของ",
    ("TRANSFER", "CASH"): "เงินสดถูกกดเป็นเงินโอน — รับเงินสดให้กดปุ่ม “เงินสด”",
    ("QR", "TRANSFER"):   "QR กับเงินโอนสลับกัน — QR = สแกนคิวอาร์, เงินโอน = โอนเข้าเลขบัญชี เลือกปุ่มให้ตรงวิธีจ่ายจริง",
    ("TRANSFER", "QR"):   "เงินโอนกับ QR สลับกัน — ตรวจว่าลูกค้าสแกน QR หรือโอนเข้าบัญชี แล้วเลือกปุ่มให้ตรง",
    ("CASH", "CREDIT"):   "ยอดรูดบัตรถูกกดเป็นเงินสด — รูดบัตรให้กดปุ่ม “บัตรเครดิต”",
    ("CREDIT", "CASH"):   "เงินสดถูกกดเป็นบัตรเครดิต — รับเงินสดให้กดปุ่ม “เงินสด”",
}


def _coaching_tips(cat_pairs) -> list[str]:
    """One concrete tip per confusion pattern (most frequent first) + a closing habit."""
    tips, seen = [], set()
    for (fr, to), _n in cat_pairs.most_common():
        t = _TIP_BY_PAIR.get((fr, to)) or (
            f"ช่องทาง {_CH_LABEL.get(fr, fr)} → {_CH_LABEL.get(to, to)}: "
            "ก่อนกดยืนยันรับเงิน ตรวจว่าวิธีจ่ายของลูกค้าตรงกับปุ่มที่กด")
        if t not in seen:
            seen.add(t); tips.append(t)
    tips.append("ทุกสิ้นวัน ตรวจยอดแต่ละช่องทางใน POS เทียบกับเงินเข้าจริงก่อนปิดยอด — จับผิดได้ตั้งแต่วันนั้น")
    return tips


def _wrong_channel_branch_detail(conn, branch_id: int, period_month: str | None,
                                 all_items: list | None = None) -> dict:
    """Self-improvement metrics for one branch: monthly trend, next-month target,
    confusion patterns, and coaching tips."""
    from collections import defaultdict, Counter
    if all_items is None:
        all_items = _wrong_channel_items(conn, None) + _channel_move_items(conn, None)
    mine = [it for it in all_items if it["branch_id"] == branch_id]

    by_m: dict = defaultdict(lambda: {"count": 0, "amount": 0.0})
    for it in mine:
        g = by_m[it["period_month"]]; g["count"] += 1; g["amount"] += it["amount"]
    months = sorted(by_m)
    if period_month:
        months = [m for m in months if m <= period_month]
    trend = [{"month": m, "count": by_m[m]["count"], "amount": round(by_m[m]["amount"], 2)}
             for m in months[-6:]]

    scope = [it for it in mine if (not period_month or it["period_month"] == period_month)]
    pat: dict = defaultdict(lambda: {"count": 0, "amount": 0.0})
    for it in scope:
        g = pat[(it["from_th"], it["to_th"])]; g["count"] += 1; g["amount"] += it["amount"]
    patterns = sorted(
        [{"from_th": k[0], "to_th": k[1], "count": v["count"], "amount": round(v["amount"], 2)}
         for k, v in pat.items()], key=lambda x: (-x["count"], -x["amount"]))
    cat_pairs = Counter((it["from_cat"], it["to_cat"]) for it in scope)

    cur = len(scope)
    prev = trend[-2]["count"] if len(trend) >= 2 and trend[-1]["month"] == period_month else None
    if prev is None:
        direction = "flat"
    elif cur < prev:
        direction = "down"          # ดีขึ้น (น้อยลง)
    elif cur > prev:
        direction = "up"            # แย่ลง
    else:
        direction = "flat"
    target = max(0, cur // 2)       # ตั้งเป้าลดลงราวครึ่งหนึ่งของเดือนนี้
    return {"trend": trend, "current": cur, "prev": prev, "direction": direction,
            "target": target, "patterns": patterns, "tips": _coaching_tips(cat_pairs)}


def _wrong_channel_report(conn, period_month: str | None = None,
                          branch_id: int | None = None) -> dict:
    from collections import defaultdict
    items = _wrong_channel_items(conn, period_month) + _channel_move_items(conn, period_month)
    if branch_id:                                       # optional single-branch filter
        items = [it for it in items if it["branch_id"] == branch_id]
    # for cross-month trend (per-branch inside) — detected matches + manual moves
    all_items = _wrong_channel_items(conn, None) + _channel_move_items(conn, None)

    agg: dict = defaultdict(lambda: {"count": 0, "total": 0.0, "months": set()})
    for it in items:
        g = agg[(it["branch_id"], it["branch_name"])]
        g["count"] += 1; g["total"] += it["amount"]; g["months"].add(it["period_month"])
    summary = []
    for (bid, bn), v in agg.items():
        detail = _wrong_channel_branch_detail(conn, bid, period_month, all_items)
        summary.append({"branch_id": bid, "branch_name": bn, "count": v["count"],
                        "months": len(v["months"]), "total": round(v["total"], 2),
                        "trend": detail["trend"], "direction": detail["direction"],
                        "target": detail["target"], "patterns": detail["patterns"],
                        "tips": detail["tips"]})
    summary.sort(key=lambda x: -x["total"])
    return {"period_month": period_month, "summary": summary, "items": items,
            "total_count": len(items),
            "total_amount": round(sum(it["amount"] for it in items), 2)}


# ── All reconciliation errors of the month (superset of wrong-channel) ────────
# The owner wants the monthly error report to cover EVERY error found, not only
# wrong-channel mis-entries: open variances, missing documents, unmatched
# statement credits, payout shortfalls, and unmapped POS codes — each tagged by
# type + status (open / acknowledged / resolved) with its amount + reason.
_ERR_TYPE_TH = {
    "wrong_channel": "บันทึกผิดช่อง",
    "variance":      "ยอดต่าง (POS≠เงินเข้า)",
    "ext_only":      "เงินเข้าเกิน (ไม่มียอด POS)",
    "missing_doc":   "ขาดเอกสาร",
    "unmapped_code": "ยังไม่ตั้งธนาคารของรหัส",
    "payout_short":  "เงินโอนเข้าไม่ครบ/ช้า",
}


def _all_errors_items(conn, period_month: str, branch_id: int | None = None) -> list[dict]:
    """Every reconciliation error in the month across branches, tagged by type."""
    bnames = {r["branch_id"]: r["name"] for r in conn.execute("SELECT branch_id, name FROM branches")}
    if branch_id:
        branches = [(branch_id, bnames.get(branch_id, str(branch_id)))]
    else:
        branches = [(r["branch_id"], r["name"]) for r in conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name")]

    items: list[dict] = []

    def add(bid, bn, typ, channel, channel_th, detail, amount, status, reason="", by="", days=None):
        items.append({"branch_id": bid, "branch_name": bn, "period_month": period_month,
                      "type": typ, "type_th": _ERR_TYPE_TH.get(typ, typ),
                      "channel": channel, "channel_th": channel_th, "detail": detail,
                      "amount": round(float(amount or 0), 2), "status": status,
                      "reason": reason or "", "by": by or "", "days": days or []})

    # confirmed wrong-channel mis-entries — detected matches + manual moves
    for w in _wrong_channel_items(conn, period_month) + _channel_move_items(conn, period_month):
        if branch_id and w["branch_id"] != branch_id:
            continue
        tag = " (ย้ายด้วยมือ)" if w.get("source") == "move" else ""
        add(w["branch_id"], w["branch_name"], "wrong_channel", w["from_cat"], w["from_th"],
            f"{w['from_th']} → {w['to_th']}" + (f" ({w['to_bank_th']})" if w["to_bank_th"] else "") + tag,
            w["amount"], "resolved", w["reason"], w["by"], w["days"])

    for bid, bn in branches:
        try:
            rm = reconcile_month(conn, bid, period_month)
        except Exception as e:
            print(f"[recon-report] reconcile_month branch={bid} {period_month} failed, branch skipped: {e}")
            continue
        for ch in rm.get("channels", []):
            ch_th = ch["label"]
            acc = ch.get("accepted")
            ext_only = ch.get("pos", 0) <= 0.5 and ch.get("ext", 0) > 0.5
            typ = "ext_only" if ext_only else "variance"
            if ch["status"] == "diff":
                add(bid, bn, typ, ch["key"], ch_th, ch.get("detail", ""),
                    ch.get("residual", ch.get("variance", 0)), "open")
            elif acc:
                add(bid, bn, "variance", ch["key"], ch_th, "รับทราบทั้งช่องทาง",
                    ch.get("variance", 0), "acknowledged", acc.get("note", ""), acc.get("by", ""))
            elif ch.get("acked_days"):
                add(bid, bn, "variance", ch["key"], ch_th,
                    f"รับทราบรายวัน {ch['acked_days']} วัน", ch.get("acked_var", 0), "acknowledged")
        for d in rm.get("documents", []):
            if d["status"] == "missing":
                add(bid, bn, "missing_doc", d.get("channel"), d["label"],
                    d.get("detail", "ยังไม่อัปโหลด"), 0, "open")
        for q in rm.get("setup_questions", []):
            add(bid, bn, "unmapped_code", None, q["code"],
                f"รหัส {q['code']} ยังไม่ระบุธนาคาร", q.get("amount", 0), "open")
        try:
            for p in payout_reconciliation(conn, bid, period_month):
                if p["status"] == "mismatch":
                    add(bid, bn, "payout_short", p["category"],
                        _CH_LABEL.get(p["category"], p["category"]),
                        f"สุทธิควรได้ {p['report_net']:,.0f} · เข้าจริง {p['bank_payout']:,.0f}",
                        p["variance"], "open")
        except Exception as e:
            print(f"[recon-report] payout rows skipped: {e}")

    items.sort(key=lambda x: (x["branch_name"], x["status"] != "open", x["type"]))
    return items


def _all_errors_report(conn, period_month: str, branch_id: int | None = None) -> dict:
    from collections import Counter
    errs = _all_errors_items(conn, period_month, branch_id)
    by_type = Counter(e["type"] for e in errs)
    by_branch: dict = {}
    for e in errs:
        g = by_branch.setdefault((e["branch_id"], e["branch_name"]),
                                 {"branch_id": e["branch_id"], "branch_name": e["branch_name"],
                                  "count": 0, "open": 0, "amount": 0.0})
        g["count"] += 1
        g["open"] += 1 if e["status"] == "open" else 0
        g["amount"] += abs(e["amount"])
    branch_summary = sorted(by_branch.values(), key=lambda x: -x["open"])
    for b in branch_summary:
        b["amount"] = round(b["amount"], 2)
    return {
        "errors": errs,
        "error_counts": {t: by_type.get(t, 0) for t in _ERR_TYPE_TH},
        "error_type_th": _ERR_TYPE_TH,
        "error_branch_summary": branch_summary,
        "error_total_count": len(errs),
        "error_open_count": sum(1 for e in errs if e["status"] == "open"),
        "error_total_amount": round(sum(abs(e["amount"]) for e in errs), 2),
    }


@router.get("/api/reports/wrong-channel")
def wrong_channel_report(period_month: str | None = None, branch_id: int | None = None):
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _wrong_channel_report(conn, period_month, branch_id)
        if period_month:        # all-errors register needs a concrete month to compose
            rep.update(_all_errors_report(conn, period_month, branch_id))
        return rep


@router.get("/api/reports/wrong-channel.xlsx")
def wrong_channel_report_xlsx(period_month: str | None = None, branch_id: int | None = None):
    import io
    import openpyxl
    from openpyxl.styles import Font
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _wrong_channel_report(conn, period_month, branch_id)
        all_err = _all_errors_report(conn, period_month, branch_id) if period_month else None
    wb = openpyxl.Workbook(); bold = Font(bold=True)
    if all_err is not None:
        # Lead sheet: EVERY error of the month (the broadened report), then the
        # wrong-channel-specific sheets follow for the coaching detail.
        wse = wb.active; wse.title = "ข้อผิดพลาดทั้งหมด"
        wse.append(["สาขา", "ประเภทข้อผิดพลาด", "ช่องทาง/รหัส", "รายละเอียด",
                    "ยอด (บาท)", "สถานะ", "เหตุผล/หมายเหตุ", "โดย"])
        for cc in wse[1]:
            cc.font = bold
        _st_th = {"open": "ค้าง", "acknowledged": "รับทราบแล้ว", "resolved": "แก้ไขแล้ว"}
        for e in all_err["errors"]:
            _xl_append(wse, [e["branch_name"], e["type_th"], e["channel_th"] or "—", e["detail"],
                             e["amount"], _st_th.get(e["status"], e["status"]),
                             e["reason"], e["by"]])
        for col, w in zip("ABCDEFGH", (20, 22, 18, 34, 14, 14, 38, 14)):
            wse.column_dimensions[col].width = w
        ws = wb.create_sheet("บันทึกผิดช่อง — สรุปรายสาขา")
    else:
        ws = wb.active; ws.title = "สรุปรายสาขา"
    ws.append(["สาขา", "จำนวนครั้ง", "จำนวนเดือน", "ยอดรวม (บาท)"])
    for c in ws[1]:
        c.font = bold
    for s in rep["summary"]:
        _xl_append(ws, [s["branch_name"], s["count"], s["months"], s["total"]])
    ws2 = wb.create_sheet("รายการ")
    ws2.append(["สาขา", "เดือน", "วันที่", "บันทึกผิดเป็น", "ที่จริงเข้า", "ยอด (บาท)",
                "เหตุผล / หมายเหตุ", "ยืนยันโดย", "เมื่อ"])
    for c in ws2[1]:
        c.font = bold
    for it in rep["items"]:
        to_th = it["to_th"] + (f" ({it['to_bank_th']})" if it["to_bank_th"] else "")
        _xl_append(ws2, [it["branch_name"], it["period_month"], ", ".join(str(d) for d in it["days"]),
                         it["from_th"], to_th, it["amount"], it["reason"], it["by"],
                         (it["at"] or "")[:16]])
    ws.column_dimensions["A"].width = 22
    for col, w in zip("ABCDEFGHI", (22, 10, 22, 16, 18, 14, 40, 16, 18)):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf)
    fname = f"wrong_channel_{period_month or 'all'}{('_b'+str(branch_id)) if branch_id else ''}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{fname}\""})


@router.get("/api/reports/wrong-channel.print")
def wrong_channel_report_print(period_month: str | None = None,
                               branch_id: int | None = None):
    """Print-to-PDF acknowledgment sheet — one signed page per branch listing the
    wrong-channel corrections, for staff to รับทราบ + commit to fix."""
    if period_month:
        _validate_period(period_month)
    with get_db() as conn:
        rep = _wrong_channel_report(conn, period_month)
        all_err = _all_errors_report(conn, period_month, branch_id) if period_month else {"errors": []}
    items = rep["items"]
    if branch_id:
        items = [it for it in items if it["branch_id"] == branch_id]

    # all errors grouped by branch (the broadened report); wrong-channel items below
    err_by_branch: dict = {}
    err_bname: dict = {}
    for e in all_err["errors"]:
        err_by_branch.setdefault(e["branch_id"], []).append(e)
        err_bname[e["branch_id"]] = e["branch_name"]

    # group wrong-channel items by branch (preserve summary order: biggest first)
    order = [s["branch_id"] for s in rep["summary"]]
    by_branch: dict = {}
    for it in items:
        by_branch.setdefault(it["branch_id"], []).append(it)
        err_bname.setdefault(it["branch_id"], it["branch_name"])
    # one sheet per branch that has ANY error (wrong-channel or otherwise)
    bids = [b for b in order if b in err_by_branch or b in by_branch]
    bids += [b for b in err_by_branch if b not in bids]
    bids += [b for b in by_branch if b not in bids]

    import html as _html
    _TH_MO = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
              "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]

    def _m(v):
        return f"{float(v or 0):,.0f}"

    def _esc(s):
        return _html.escape(str(s or ""))

    def _mth(ym):
        try:
            y, mo = ym.split("-"); return f"{_TH_MO[int(mo)]} {(int(y) + 543) % 100:02d}"
        except Exception:
            return ym

    detail_by = {s["branch_id"]: s for s in rep["summary"]}

    css = """<style>
      @page{size:A4;margin:14mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12px;margin:0;padding:24px;background:#f8fafc}
      .sheet{background:#fff;max-width:900px;margin:0 auto 18px;padding:26px 30px;border-radius:4px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:19px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:4px}
      h2{font-size:15px;margin:18px 0 8px;border:1px solid #ea580c;padding-left:8px}
      .muted{color:#64748b;font-size:11px} .right{text-align:right} .mono{font-variant-numeric:tabular-nums;font-family:ui-monospace,monospace}
      table{width:100%;border-collapse:collapse;margin:6px 0 10px;font-size:11.5px}
      th,td{padding:6px 8px;border-bottom:1px solid #eef2f7;text-align:left;vertical-align:top}
      th{color:#64748b;font-size:10px;border-bottom:1.5px solid #e2e8f0}
      tfoot td{border-top:1.5px solid #cbd5e1;font-weight:700}
      .arr{color:#ea580c;font-weight:700} .amt{font-weight:600}
      .pill{display:inline-block;padding:1px 9px;border-radius:99px;font-size:11px;font-weight:600}
      .p-ok{background:#dcfce7;color:#166534}.p-up{background:#fee2e2;color:#b91c1c}.p-flat{background:#f1f5f9;color:#475569}
      .trend{display:flex;gap:6px;align-items:flex-end;margin:8px 0 4px;height:64px}
      .tb{flex:1;display:flex;flex-direction:column;justify-content:flex-end;align-items:center;gap:3px;font-size:10px;color:#64748b}
      .tb .bar{width:60%;min-height:3px;background:#fdba74;border-radius:3px 3px 0 0}
      .tb.cur .bar{background:#ea580c}
      .target{background:#eff6ff;border:1px solid #bfdbfe;border-radius:4px;padding:9px 13px;margin:6px 0;font-size:12px}
      .tips{margin:6px 0 2px;padding-left:0;list-style:none}
      .tips li{position:relative;padding:5px 0 5px 24px;font-size:12px;line-height:1.5;border-bottom:1px dashed #eef2f7}
      .tips li::before{content:'';position:absolute;left:0;top:4px}
      .ack{background:#fff7ed;border:1px solid #fed7aa;border-radius:4px;padding:12px 14px;margin-top:14px;font-size:12px;line-height:1.6}
      .sign{display:flex;gap:40px;margin-top:30px} .sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;font-size:11px;color:#475569}
      .toolbar{max-width:900px;margin:0 auto 14px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:4px;padding:9px 18px;font-size:13px;cursor:pointer}
      .empty{max-width:900px;margin:40px auto;text-align:center;color:#64748b}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;margin:0;max-width:none;border-radius:0;page-break-after:always}}
    </style>"""

    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
             "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
             "<title>รายงานข้อผิดพลาดประจำเดือน</title>", css, "</head><body>",
             "<div class='toolbar'><button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>"]

    if not bids:
        parts.append("<div class='empty'>ไม่พบข้อผิดพลาดในการกระทบยอด"
                     f"{' สำหรับเดือน ' + _esc(period_month) if period_month else ''} — "
                     "ทุกช่องทางตรง ไม่มียอดต่าง เอกสารครบ และไม่มีการบันทึกผิดช่อง</div>")
    period_label = period_month or "ทุกเดือน"
    _ST_PILL = {"open": ("p-up", "ค้าง"), "acknowledged": ("p-flat", "รับทราบแล้ว"),
                "resolved": ("p-ok", "แก้ไขแล้ว")}
    for bid in bids:
        rows = by_branch.get(bid, [])
        errs = err_by_branch.get(bid, [])
        bn = err_bname.get(bid, str(bid))
        n_open = sum(1 for e in errs if e["status"] == "open")
        parts.append("<div class='sheet'>")
        parts.append(f"<h1>รายงานข้อผิดพลาดประจำเดือน — สาขา{_esc(bn)}</h1>")
        parts.append(f"<div class='sub'>งวด {_esc(period_label)} · ข้อผิดพลาดทั้งหมด {len(errs)} รายการ "
                     f"(ค้าง {n_open}) · บันทึกผิดช่อง {len(rows)} รายการ</div>")

        # ── ข้อผิดพลาดทั้งหมดในเดือนนี้ (ทุกประเภท) ──
        if errs:
            parts.append("<h2>ข้อผิดพลาดทั้งหมดในเดือนนี้ (ทุกประเภท)</h2>")
            parts.append("<table><thead><tr><th>ประเภท</th><th>ช่องทาง/รหัส</th><th>รายละเอียด</th>"
                         "<th class='right'>ยอด (บาท)</th><th>สถานะ</th><th>เหตุผล/หมายเหตุ</th></tr></thead><tbody>")
            for e in errs:
                pill = _ST_PILL.get(e["status"], ("p-flat", e["status"]))
                parts.append(
                    f"<tr><td>{_esc(e['type_th'])}</td><td>{_esc(e['channel_th'] or '—')}</td>"
                    f"<td>{_esc(e['detail'])}</td>"
                    f"<td class='right mono amt'>{_m(e['amount']) if e['amount'] else '—'}</td>"
                    f"<td><span class='pill {pill[0]}'>{pill[1]}</span></td>"
                    f"<td class='muted'>{_esc(e['reason'] or '—')}</td></tr>")
            parts.append("</tbody></table>")

        if not rows:                          # branch has other errors but no wrong-channel mis-entries
            parts.append("<div class='sign'><div>ผู้จัดการสาขา</div><div>แคชเชียร์ / ผู้บันทึก</div>"
                         "<div>ฝ่ายบัญชี / ผู้ตรวจสอบ</div></div></div>")
            continue

        tot = sum(r["amount"] for r in rows)
        parts.append("<h2>รายการที่บันทึกผิดช่อง (รายละเอียด + แนวทางพัฒนา)</h2>")
        parts.append(f"<div class='sub'>มูลค่ารวมบันทึกผิดช่อง ฿{_m(tot)}</div>")
        parts.append("<table><thead><tr><th>วันที่</th><th>บันทึกผิดเป็น</th>"
                     "<th>ที่จริงเข้า</th><th class='right'>ยอด (บาท)</th>"
                     "<th>เหตุผล / หมายเหตุ</th><th>ยืนยันโดย</th></tr></thead><tbody>")
        for r in rows:
            to_th = _esc(r["to_th"]) + (f" ({_esc(r['to_bank_th'])})" if r["to_bank_th"] else "")
            days = ", ".join(_esc(str(d)[8:].lstrip("0")) for d in r["days"]) or "—"
            parts.append(
                f"<tr><td>{days}</td><td>{_esc(r['from_th'])}</td>"
                f"<td><span class='arr'>→</span> {to_th}</td>"
                f"<td class='right mono amt'>{_m(r['amount'])}</td>"
                f"<td>{_esc(r['reason'] or '—')}</td>"
                f"<td class='muted'>{_esc(r['by'] or '—')}<br>{_esc((r['at'] or '')[:16])}</td></tr>")
        parts.append(f"</tbody><tfoot><tr><td colspan='3'>รวม</td>"
                     f"<td class='right mono'>{_m(tot)}</td><td colspan='2'></td></tr></tfoot></table>")

        det = detail_by.get(bid, {})
        # ① แนวโน้มรายเดือน + เป้าหมาย
        trend = det.get("trend") or []
        if trend:
            mx = max((t["count"] for t in trend), default=0) or 1
            dirn = det.get("direction", "flat")
            badge = ({"down": ("p-ok", "✓ ดีขึ้น (ผิดน้อยลง)"),
                      "up": ("p-up", "▲ ผิดมากขึ้น — ต้องระวัง"),
                      "flat": ("p-flat", "— ทรงตัว")}).get(dirn, ("p-flat", "—"))
            parts.append(f"<h2>① แนวโน้มการกดผิดช่อง <span class='pill {badge[0]}'>{badge[1]}</span></h2>")
            parts.append("<div class='trend'>")
            for t in trend:
                cur = "cur" if t["month"] == period_month else ""
                h = int(8 + 52 * t["count"] / mx)
                parts.append(f"<div class='tb {cur}'><b>{t['count']}</b>"
                             f"<div class='bar' style='height:{h}px'></div>{_esc(_mth(t['month']))}</div>")
            parts.append("</div>")
            tg = det.get("target", 0)
            parts.append(f"<div class='target'><b>เป้าหมายเดือนถัดไป: ไม่เกิน {tg} ครั้ง</b> "
                         f"(เดือนนี้ {det.get('current', len(rows))} ครั้ง) — "
                         "ตั้งเป้าลดลงครึ่งหนึ่ง โดยตรวจปุ่มรับเงินให้ตรงทุกบิล</div>")
        # ② ช่องที่สับสนบ่อย
        patterns = det.get("patterns") or []
        if patterns:
            parts.append("<h2>② ช่องทางที่สับสนบ่อย (จุดที่ต้องระวัง)</h2>")
            parts.append("<table><thead><tr><th>กดผิดเป็น</th><th>ที่จริงคือ</th>"
                         "<th class='right'>จำนวนครั้ง</th><th class='right'>ยอดรวม</th></tr></thead><tbody>")
            for p in patterns:
                parts.append(f"<tr><td>{_esc(p['from_th'])}</td>"
                             f"<td><span class='arr'>→</span> {_esc(p['to_th'])}</td>"
                             f"<td class='right mono amt'>{p['count']}</td>"
                             f"<td class='right mono'>{_m(p['amount'])}</td></tr>")
            parts.append("</tbody></table>")
        # ③ คำแนะนำสำหรับสาขา (โค้ช)
        tips = det.get("tips") or []
        if tips:
            parts.append("<h2>③ คำแนะนำเพื่อพัฒนา</h2><ul class='tips'>")
            for t in tips:
                parts.append(f"<li>{_esc(t)}</li>")
            parts.append("</ul>")

        parts.append(
            "<div class='ack'>ข้าพเจ้าในนามสาขา <b>รับทราบ</b> ว่ารายการข้างต้นเกิดจากการบันทึก"
            "ช่องทางการรับเงินผิดในเครื่อง POS และ <b>รับว่าจะแก้ไข</b>ขั้นตอนการกดรับเงินให้ถูกต้อง "
            "เพื่อไม่ให้เกิดซ้ำในรอบถัดไป (ยอดเงินรวมถูกต้อง ไม่มีเงินขาด — เป็นการลงผิดช่องเท่านั้น)</div>")
        parts.append("<div class='sign'>"
                     "<div>ผู้จัดการสาขา (เซ็นรับทราบ)</div>"
                     "<div>แคชเชียร์ / ผู้บันทึก</div>"
                     "<div>ฝ่ายบัญชี / ผู้ตรวจสอบ</div></div>")
        parts.append("</div>")

    parts.append("</body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


def _doc_channel(label: str) -> str | None:
    """Map a required-document label to the channel it backs (for the baht-at-stake)."""
    l = (label or "")
    if "บัตร" in l:                                   return "CREDIT"
    if "Grab" in l or "แกร็บ" in l:                   return "GRAB"
    if "Line Man" in l or "ไลน์แมน" in l or "LM" in l.upper(): return "LINEMAN"
    if "คะแนน" in l or "คูปอง" in l:                 return "COUPON"
    if "สเตทเมนต์" in l or "statement" in l.lower():  return "BANK"   # backs QR/CASH/TRANSFER
    return None


def _outstanding_docs_report(conn, period_month: str, branch_id: int | None = None) -> dict:
    from collections import defaultdict
    if branch_id:
        masters = conn.execute("SELECT branch_id, name FROM branches WHERE branch_id=?",
                               (branch_id,)).fetchall()
    else:
        masters = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()

    branches, grand = [], 0.0
    for b in masters:
        try:
            rm = reconcile_month(conn, b["branch_id"], period_month)
        except Exception as e:
            print(f"[recon-report] reconcile_month branch={b['branch_id']} {period_month} failed, branch skipped: {e}")
            continue
        chmap = {c["key"]: c for c in rm.get("channels", [])}
        miss = defaultdict(list)            # channel -> [missing/partial doc labels]
        for doc in rm.get("documents", []):
            if doc.get("status") in ("missing", "partial"):
                miss[_doc_channel(doc["label"])].append(
                    {"label": doc["label"], "status": doc["status"]})

        # POS composition (for the detailed per-acquirer / per-code breakdown of what's waiting)
        try:
            codes = pos_codes_scope(conn, b["branch_id"], period_month=period_month)
        except sqlite3.OperationalError:
            codes = []
        by_ch = defaultdict(list)
        for cd in codes:
            if cd["category"] and cd["amount"] > 0:
                by_ch[cd["category"]].append(cd)

        rows = []
        for key, docs in miss.items():
            if key == "BANK":               # statements back several lanes; report at branch level
                continue
            c = chmap.get(key, {})
            unmatched = round(max(0.0, float(c.get("pos", 0)) - float(c.get("ext", 0))), 2)
            # detail: for CREDIT split by acquirer (which settlement to chase); else by POS code
            agg = defaultdict(lambda: {"amount": 0.0, "txn": 0})
            for cd in by_ch.get(key, []):
                k = (_BANK_LABEL.get(cd["bank"], cd["bank"]) if key == "CREDIT"
                     else cd["code"])
                agg[k]["amount"] += cd["amount"]; agg[k]["txn"] += cd["txn_count"]
            breakdown = sorted(
                ({"label": k, "amount": round(v["amount"], 2), "txn": v["txn"]}
                 for k, v in agg.items()), key=lambda x: -x["amount"])
            rows.append({"channel": key, "channel_th": _MATCH_CH_TH.get(key, key),
                         "unmatched": unmatched, "breakdown": breakdown,
                         "docs": [d["label"] for d in docs],
                         "doc_status": "partial" if any(d["status"] == "partial" for d in docs) else "missing"})
        # missing bank statements (rare — usually uploaded) listed separately
        bank_missing = [d["label"] for d in miss.get("BANK", [])]
        rows.sort(key=lambda r: -r["unmatched"])
        branch_total = round(sum(r["unmatched"] for r in rows), 2)
        grand += branch_total
        if rows or bank_missing:
            branches.append({"branch_id": b["branch_id"], "branch_name": b["name"],
                             "progress_pct": rm.get("progress_pct"),
                             "rows": rows, "bank_missing": bank_missing,
                             "branch_total": branch_total})
    branches.sort(key=lambda x: -x["branch_total"])
    return {"period_month": period_month, "branches": branches,
            "grand_total": round(grand, 2)}


@router.get("/api/reports/outstanding-docs")
def outstanding_docs_report(period_month: str, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _outstanding_docs_report(conn, period_month, branch_id)


@router.get("/api/reports/outstanding-docs.print")
def outstanding_docs_report_print(period_month: str, branch_id: int | None = None):
    """Print-to-PDF follow-up sheet: per branch, the documents still missing this month
    + the baht waiting on each. Month + branch picker on the page (hidden when printing)."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v):  return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        rep = _outstanding_docs_report(conn, period_month, branch_id)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()

    css = """<style>
      @page{size:A4;margin:14mm}
      *{box-sizing:border-box} body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;font-size:12px;margin:0;padding:22px;background:#f8fafc}
      .sheet{background:#fff;max-width:920px;margin:0 auto 16px;padding:24px 28px;border-radius:4px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
      h1{font-size:19px;margin:0 0 2px} .sub{color:#64748b;font-size:12px;margin-bottom:6px}
      h2{font-size:14px;margin:16px 0 6px;border:1px solid #ea580c;padding-left:8px}
      table{width:100%;border-collapse:collapse;font-size:12px;margin:4px 0 6px}
      th,td{padding:6px 9px;border-bottom:1px solid #eef2f7;text-align:right}
      th{color:#64748b;font-size:10px;border-bottom:1.5px solid #e2e8f0;background:#f8fafc}
      td:first-child,th:first-child,td.l,th.l{text-align:left}
      .mono{font-variant-numeric:tabular-nums} .big{color:#b91c1c;font-weight:700}
      .muted{color:#64748b;font-size:11px} tr.sub td{border-bottom:1px dashed #f1f5f9;padding:3px 9px}
      tfoot td{border-top:1.5px solid #cbd5e1;font-weight:700}
      .chk{width:16px;height:16px;border:1.5px solid #94a3b8;border-radius:3px;display:inline-block;vertical-align:middle}
      .partial{color:#92610e}.lead{background:#fff7ed;border:1px solid #fed7aa;border-radius:4px;padding:9px 13px;font-size:12px;margin-bottom:12px}
      .none{background:#f0fdf4;border:1px solid #bbf7d0;border-radius:4px;padding:9px 13px;color:#166534;font-size:12px}
      .toolbar{max-width:920px;margin:0 auto 12px;text-align:right}
      .btn{background:#ea580c;color:#fff;border:0;border-radius:4px;padding:9px 18px;font-size:13px;cursor:pointer}
      .ctl{padding:8px 10px;border:1px solid #cbd5e1;border-radius:4px;font:inherit;margin-right:8px}
      .sign{display:flex;gap:40px;margin-top:22px}.sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;font-size:11px;color:#475569}
      @media print{body{background:#fff;padding:0}.toolbar{display:none}.sheet{box-shadow:none;border-radius:0;max-width:none;margin:0}}
    </style>"""

    bopts = (f"<option value=''{'' if branch_id else ' selected'}>ทุกสาขา (5 สาขาหลัก)</option>"
             + "".join(f"<option value='{b['branch_id']}'"
                       f"{' selected' if branch_id == b['branch_id'] else ''}>{esc(b['name'])}</option>"
                       for b in master))
    bsuffix = f"&branch_id={branch_id}" if branch_id else ""
    month_oc = "location.href='?period_month='+this.value+'" + bsuffix + "'"
    sel_oc = (f"location.href='?period_month={period_month}'"
              "+(this.value?'&branch_id='+this.value:'')")
    picker = (f"<input type='month' value='{esc(period_month)}' class='ctl' onchange=\"{month_oc}\">"
              f"<select class='ctl' onchange=\"{sel_oc}\">" + bopts + "</select>")

    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
             "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
             "<title>ติดตามเอกสารที่ขาด</title>", css, "</head><body>",
             "<div class='toolbar'>" + picker
             + "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>",
             "<div class='sheet'>",
             f"<h1>รายงานติดตามเอกสารที่ยังขาด — งวด {esc(period_month)}</h1>",
             f"<div class='sub'>เอกสารที่แต่ละสาขายังต้องอัปโหลด + ยอดเงินที่รอเอกสารนั้นจับคู่ · "
             f"รวมเงินที่รอเอกสารทั้งหมด <b class='big'>฿{m(rep['grand_total'])}</b></div>"]

    if not rep["branches"]:
        parts.append("<div class='none'>✓ ทุกสาขาอัปเอกสารครบสำหรับงวดนี้ — ไม่มีเอกสารค้าง</div>")
    for b in rep["branches"]:
        parts.append(f"<h2>{esc(b['branch_name'])} — เหลือเอกสารคิดเป็นเงิน ฿{m(b['branch_total'])}"
                     + (f" · ความคืบหน้า {b['progress_pct']}%" if b.get('progress_pct') is not None else "") + "</h2>")
        if b["rows"]:
            parts.append("<table><thead><tr><th class='l'>อัปแล้ว</th><th class='l'>ช่องทาง</th>"
                         "<th class='l'>เอกสารที่ต้องอัปโหลด</th><th>ยอดที่รอจับคู่ (บาท)</th></tr></thead><tbody>")
            for r in b["rows"]:
                docs = " · ".join(esc(x) for x in r["docs"])
                st = " <span class='partial'>(อัปบางส่วน)</span>" if r["doc_status"] == "partial" else ""
                parts.append(f"<tr><td class='l'><span class='chk'></span></td>"
                             f"<td class='l'><b>{esc(r['channel_th'])}</b></td>"
                             f"<td class='l'>{docs}{st}</td>"
                             f"<td class='mono{' big' if r['unmatched']>20000 else ''}'>{m(r['unmatched'])}</td></tr>")
                # detailed breakdown: per acquirer (CREDIT) or per POS code — so the team
                # knows exactly which document to get and the expected amount it should total
                for d in r.get("breakdown", []):
                    note = ("เอกสารต้องมียอดราว ฿" + m(d["amount"])) if r["channel"] == "CREDIT" else ""
                    parts.append(
                        f"<tr class='sub'><td></td><td class='l muted'>↳ {esc(d['label'])}</td>"
                        f"<td class='l muted'>{esc(note)} <span class='muted'>({d['txn']:,} รายการ)</span></td>"
                        f"<td class='mono muted'>{m(d['amount'])}</td></tr>")
            parts.append(f"</tbody><tfoot><tr><td colspan='3' class='l'>รวมเงินที่รอเอกสาร</td>"
                         f"<td class='mono'>{m(b['branch_total'])}</td></tr></tfoot></table>")
        if b["bank_missing"]:
            parts.append("<div class='lead'>⚠ สเตทเมนต์ธนาคารที่ยังขาด: "
                         + " · ".join(esc(x) for x in b["bank_missing"])
                         + " (จำเป็นต่อ QR/เงินสด/เงินโอน)</div>")
        parts.append("<div class='sign'><div>ผู้รับผิดชอบสาขา</div><div>วันที่อัปครบ</div></div>")
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


_IC_BANK_TH = {"KTB": "กรุงไทย", "KBANK": "กสิกร", "KTC": "KTC", "BAY": "กรุงศรี",
               "SCB": "ไทยพาณิชย์", "BBL": "กรุงเทพ", "TTB": "ทีทีบี",
               "PROMPTPAY": "พร้อมเพย์", "ANY": "ทุกธนาคาร", "NONE": "—", "ALL": "ทุกช่อง"}


_IC_SESS_TH = {"draft": "ร่าง (ยังไม่สอบทาน)", "in_review": "รอสอบทาน",
               "approved": "อนุมัติแล้ว", "closed": "ปิดงวดแล้ว"}


_IC_CHST_TH = {"ok": "ตรงกัน", "diff": "มีผลต่าง", "waiting_doc": "รอเอกสาร",
               "no_source": "ไม่มีคู่เทียบ", "setup": "รอตั้งค่า"}


def _ic_lane_th(cat, bank):
    base = _CH_LABEL.get(cat, cat or "—")
    if bank and bank not in ("NONE", "ALL", "ANY", "", None):
        return f"{base}·{_IC_BANK_TH.get(bank, bank)}"
    return base


def _override_register(conn, period_month, branch_id=None):
    """Every manual intervention (management override) in the month across master
    branches: manual tick-to-match, cross-channel move, lane adjustment, and per-day
    acknowledged variance — the four reversible tables. One normalized row each:
    branch · when · who · kind · lane · amount · reason. Read-only."""
    bmap = {r["branch_id"]: r["name"] for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY}")}

    def keep(bid):
        return (bid in bmap) and (branch_id is None or bid == branch_id)

    rows = []
    for r in conn.execute("SELECT * FROM recon_manual_match WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["matched_at"], "by": r["matched_by"] or "—",
            "kind": "จับคู่ด้วยมือ", "kind_key": "manual_match",
            "lane": _ic_lane_th(r["category"], r["bank"]),
            "amount": round(float(r["ext_amount"] or 0), 2),
            "detail": (f"POS {float(r['pos_amount'] or 0):,.0f} ↔ เงินเข้า "
                       f"{float(r['ext_amount'] or 0):,.0f} "
                       f"(ต่าง {float(r['variance'] or 0):+,.0f})"),
            "reason": r["note"] or ""})
    for r in conn.execute("SELECT * FROM recon_channel_move WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["created_at"], "by": r["created_by"] or "—",
            "kind": "ย้ายช่อง", "kind_key": "channel_move",
            "lane": (f"{_ic_lane_th(r['from_category'], r['from_bank'])} → "
                     f"{_ic_lane_th(r['to_category'], r['to_bank'])}"),
            "amount": round(float(r["amount"] or 0), 2),
            "detail": f"ย้าย {float(r['amount'] or 0):,.0f} บาท",
            "reason": r["note"] or ""})
    _adj_kind = {"accept": "รับทราบผลต่าง (ทั้งช่อง)", "off_pos": "เพิ่มยอด POS",
                 "deposit_note": "บันทึกเงินฝาก", "other": "ปรับปรุงอื่น ๆ"}
    for r in conn.execute("SELECT * FROM recon_lane_adj WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        amt = float(r["adj_pos"] or 0) + float(r["adj_ext"] or 0)
        det = ((f"POS +{float(r['adj_pos'] or 0):,.0f}" if r["adj_pos"] else "")
               + (f" เงินเข้า +{float(r['adj_ext'] or 0):,.0f}" if r["adj_ext"] else "")).strip()
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["created_at"], "by": r["created_by"] or "—",
            "kind": _adj_kind.get(r["adj_type"], "ปรับปรุง"), "kind_key": "lane_adj",
            "lane": _ic_lane_th(r["category"], r["bank"]),
            "amount": round(amt, 2),
            "detail": det or "รับทราบผลต่างทั้งช่องให้ปิดงวดได้",
            "reason": r["note"] or ""})
    for r in conn.execute("SELECT * FROM recon_day_ack WHERE period_month=?",
                          (period_month,)):
        if not keep(r["branch_id"]):
            continue
        rows.append({"branch_id": r["branch_id"], "branch_name": bmap[r["branch_id"]],
            "at": r["created_at"], "by": r["created_by"] or "—",
            "kind": "รับทราบผลต่างรายวัน", "kind_key": "day_ack",
            "lane": f"{_CH_LABEL.get(r['category'], r['category'])} · {r['day']}",
            "amount": round(float(r["variance"] or 0), 2),
            "detail": f"วันที่ {r['day']} ต่าง {float(r['variance'] or 0):+,.0f} บาท",
            "reason": r["note"] or ""})
    rows.sort(key=lambda x: (x["at"] or ""), reverse=True)
    for r in rows:
        r["has_reason"] = bool((r["reason"] or "").strip())
    return {"period_month": period_month, "branch_id": branch_id,
            "rows": rows, "count": len(rows),
            "total_abs_amount": round(sum(abs(r["amount"]) for r in rows), 2),
            "missing_reason": sum(0 if r["has_reason"] else 1 for r in rows)}


def _ic_period_end(m):
    import calendar
    from datetime import date
    y, mo = int(m[:4]), int(m[5:7])
    return date(y, mo, calendar.monthrange(y, mo)[1])


def _ic_prev_month(m, k):
    y, mo = int(m[:4]), int(m[5:7])
    idx = y * 12 + (mo - 1) - k
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def _today():
    """The day the aging is measured from — one seam so the golden test can pin it
    (the report changes every midnight otherwise)."""
    from datetime import date
    return date.today()


def _open_items_aging(conn, period_month, branch_id=None, lookback=2):
    """Open (unresolved) reconciling items, aged from period-end to today, for the
    selected month + `lookback` prior months. Skips approved/closed (locked) sessions —
    those are signed off, not open. Each item = one reconcile_month todo. Read-only."""
    today = _today()
    months = [period_month] + [_ic_prev_month(period_month, k)
                               for k in range(1, lookback + 1)]
    branches = [(r["branch_id"], r["name"]) for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name")
        if (branch_id is None or r["branch_id"] == branch_id)]
    items = []
    for m in months:
        pend = _ic_period_end(m)
        age = max(0, (today - pend).days)
        bucket = "0-7" if age <= 7 else ("8-30" if age <= 30 else ">30")
        for bid, bname in branches:
            sess = conn.execute(
                "SELECT status FROM recon_sessions WHERE branch_id=? AND period_month=?",
                (bid, m)).fetchone()
            if sess and sess["status"] in ("approved", "closed"):
                continue
            try:
                r = reconcile_month(conn, bid, m)
            except Exception as e:
                print(f"[recon-report] reconcile_month branch={bid} {m} failed, month skipped: {e}")
                continue
            for t in r["todo"]:
                items.append({"branch_id": bid, "branch_name": bname,
                    "period_month": m, "age_days": age, "bucket": bucket,
                    "type": t.get("type"), "label": t.get("label"),
                    "channel": t.get("channel")})
    items.sort(key=lambda x: (-x["age_days"], x["branch_name"]))
    buckets = {"0-7": 0, "8-30": 0, ">30": 0}
    for it in items:
        buckets[it["bucket"]] += 1
    return {"as_of": today.isoformat(), "period_month": period_month,
            "lookback": lookback, "months": months,
            "items": items, "count": len(items), "buckets": buckets}


def _recon_certificate(conn, branch_id, period_month):
    """Sign-off certificate data for one branch-month: per-channel POS vs money-in
    with residual after acknowledgements, preparer/reviewer + approval date, the
    frozen working-paper snapshot hash, and the count of manual overrides. Read-only."""
    r = reconcile_month(conn, branch_id, period_month)
    chs = r["channels"]
    rows, resid_total = [], 0.0
    for c in chs:
        resid = c.get("residual", c.get("variance", 0)) or 0
        if c["key"] != "COUPON":          # COUPON variance is loyalty points, not baht
            resid_total += resid
        rows.append({"key": c["key"], "label": c["label"],
            "pos": c["pos"], "ext": c["ext"], "variance": c.get("variance", 0),
            "residual": resid, "status": c["status"],
            "accepted": bool(c.get("accepted")), "detail": c.get("detail", "")})
    reg = _override_register(conn, period_month, branch_id)
    ack_rows = [x for x in reg["rows"] if x["kind_key"] in ("day_ack", "lane_adj")]
    snap = conn.execute(
        "SELECT status, sha256, created_at, file_name, source_manifest FROM recon_close_snapshot "
        "WHERE session_id=(SELECT session_id FROM recon_sessions "
        "                  WHERE branch_id=? AND period_month=?) "
        "ORDER BY id DESC LIMIT 1", (branch_id, period_month)).fetchone()
    snap_d = dict(snap) if snap else None
    if snap_d:
        try:
            snap_d["source_doc_count"] = len(json.loads(snap_d.get("source_manifest") or "[]"))
        except Exception:
            snap_d["source_doc_count"] = 0
    return {"branch_id": branch_id, "branch_name": r["branch_name"],
        "period_month": period_month, "session": r.get("session") or {},
        "progress_pct": r["progress_pct"], "todo_count": len(r["todo"]),
        "channels": rows,
        "pos_total": round(sum(c["pos"] for c in chs), 2),
        "ext_total": round(sum(c["ext"] for c in chs), 2),
        "residual_total": round(resid_total, 2),
        "ack_count": len(ack_rows),
        "ack_amount": round(sum(abs(x["amount"]) for x in ack_rows), 2),
        "override_count": reg["count"], "override_missing_reason": reg["missing_reason"],
        "snapshot": snap_d}


# ── JSON endpoints (drive the in-app views) ───────────────────────────────────
@router.get("/api/reports/override-register")
def override_register_api(period_month: str, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _override_register(conn, period_month, branch_id)


@router.get("/api/reports/open-items-aging")
def open_items_aging_api(period_month: str, branch_id: int | None = None,
                         lookback: int = 2):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _open_items_aging(conn, period_month, branch_id, max(0, min(6, lookback)))


@router.get("/api/reports/recon-certificate")
def recon_certificate_api(branch_id: int, period_month: str):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        _validate_master_branch(conn, branch_id)
        return _recon_certificate(conn, branch_id, period_month)


# ── Shared print CSS (matches the outstanding-docs sheet look) ────────────────
def _ic_print_css(accent):
    return ("<style>@page{size:A4;margin:14mm}"
        "*{box-sizing:border-box}"
        "body{font-family:'Sarabun','TH Sarabun New',system-ui,sans-serif;color:#1e293b;"
        "font-size:12px;margin:0;padding:22px;background:#f8fafc}"
        ".sheet{background:#fff;max-width:960px;margin:0 auto 16px;padding:24px 28px;"
        "border-radius:4px;box-shadow:0 1px 4px rgba(0,0,0,.08)}"
        "h1{font-size:19px;margin:0 0 2px}.sub{color:#64748b;font-size:12px;margin-bottom:8px}"
        f"h2{{font-size:14px;margin:16px 0 6px;border:1px solid {accent};padding-left:8px}}"
        "table{width:100%;border-collapse:collapse;font-size:12px;margin:4px 0 8px}"
        "th,td{padding:6px 9px;border-bottom:1px solid #eef2f7;text-align:right}"
        "th{color:#64748b;font-size:10px;border-bottom:1.5px solid #e2e8f0;background:#f8fafc}"
        "td:first-child,th:first-child,td.l,th.l{text-align:left}"
        ".mono{font-variant-numeric:tabular-nums}.muted{color:#64748b;font-size:11px}"
        ".big{color:#b91c1c;font-weight:700}.ok{color:#047857}.warn{color:#b45309}"
        "tfoot td{border-top:1.5px solid #cbd5e1;font-weight:700}"
        f".lead{{background:#f8fafc;border:1px solid #e2e8f0;border:1px solid {accent};"
        "border-radius:4px;padding:9px 13px;font-size:12px;margin-bottom:12px}"
        ".none{background:#f0fdf4;border:1px solid #bbf7d0;border-radius:4px;padding:9px 13px;"
        "color:#166534;font-size:12px}"
        ".chip{display:inline-block;padding:1px 8px;border-radius:999px;font-size:10px;font-weight:600}"
        ".c-ok{background:#ecfdf5;color:#047857}.c-warn{background:#fef3c7;color:#92610e}"
        ".c-err{background:#fef2f2;color:#b91c1c}.c-mut{background:#f1f5f9;color:#475569}"
        ".toolbar{max-width:960px;margin:0 auto 12px;text-align:right}"
        f".btn{{background:{accent};color:#fff;border:0;border-radius:4px;padding:9px 18px;"
        "font-size:13px;cursor:pointer}"
        ".ctl{padding:8px 10px;border:1px solid #cbd5e1;border-radius:4px;font:inherit;margin-right:8px}"
        ".sign{display:flex;gap:40px;margin-top:26px}"
        ".sign div{flex:1;border-top:1px solid #94a3b8;padding-top:5px;text-align:center;"
        "font-size:11px;color:#475569}"
        ".kv{display:flex;flex-wrap:wrap;gap:6px 26px;margin:8px 0}"
        ".kv div{font-size:12px}.kv b{color:#0f172a}"
        "@media print{body{background:#fff;padding:0}.toolbar{display:none}"
        ".sheet{box-shadow:none;border-radius:0;max-width:none;margin:0}"
        ".pageb{page-break-before:always}}</style>")


def _ic_picker(period_month, branch_id, master, esc, all_label="ทุกสาขา (สาขาหลัก)",
               extra=""):
    bopts = (f"<option value=''{'' if branch_id else ' selected'}>{all_label}</option>"
             + "".join(f"<option value='{b['branch_id']}'"
                       f"{' selected' if branch_id == b['branch_id'] else ''}>"
                       f"{esc(b['name'])}</option>" for b in master))
    bsuffix = f"&branch_id={branch_id}" if branch_id else ""
    month_oc = f"location.href='?period_month='+this.value+'{bsuffix}{extra}'"
    sel_oc = (f"location.href='?period_month={period_month}'"
              f"+(this.value?'&branch_id='+this.value:'')+'{extra}'")
    return (f"<input type='month' value='{esc(period_month)}' class='ctl' onchange=\"{month_oc}\">"
            f"<select class='ctl' onchange=\"{sel_oc}\">" + bopts + "</select>")


def _ic_signed_footer(request):
    """A provenance line for printed control evidence: the logged-in account that
    generated the document + a server timestamp — an electronic 'signed by' stamp."""
    import html as _html
    who = _actor(request, default="")
    now = _dt.now().strftime("%Y-%m-%d %H:%M")
    if who:
        return ("<div class='lbl' style='margin-top:18px;color:#475569'>"
                f"จัดทำ/พิมพ์โดย (บัญชีที่ล็อกอิน): <b>{_html.escape(who)}</b> · {now}"
                " · ระบบกระทบยอด Oshinei</div>")
    return ("<div class='lbl' style='margin-top:18px;color:#94a3b8'>"
            f"พิมพ์เมื่อ {now} · (ไม่พบบัญชีที่ล็อกอิน — เปิดผ่านระบบหลังล็อกอินเพื่อบันทึกชื่อผู้จัดทำ)</div>")


@router.get("/api/reports/override-register.print")
def override_register_print(request: Request, period_month: str, branch_id: int | None = None):
    """Print-to-PDF management-override register: every manual match / move / lane
    adjustment / day-acknowledgement in the month, with who·when·reason — the log a
    Big4 auditor pulls first. Rows with no reason are flagged."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v): return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _override_register(conn, period_month, branch_id)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc)
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>ทะเบียนการปรับปรุงด้วยมือ</title>", _ic_print_css("#4f46e5"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>ทะเบียนการปรับปรุงด้วยมือ (Management Override Register) — งวด {esc(period_month)}</h1>",
        f"<div class='sub'>ทุกรายการที่คนแก้ไขด้วยมือในเดือนนี้: จับคู่เอง · ย้ายช่อง · "
        f"รับทราบผลต่าง · ปรับปรุง lane — รวม <b>{reg['count']}</b> รายการ คิดเป็นเงิน "
        f"<b>฿{m(reg['total_abs_amount'])}</b>"
        + (f" · <span class='big'>{reg['missing_reason']} รายการไม่มีเหตุผลกำกับ</span>"
           if reg["missing_reason"] else " · ทุกรายการมีเหตุผลกำกับ ✓") + "</div>"]
    if reg["missing_reason"]:
        parts.append("<div class='lead'>⚠ รายการที่ไม่มีเหตุผล (เน้นสีแดง) คือจุดที่ผู้สอบบัญชี"
                     "จะตั้งคำถาม — ควรให้ผู้บันทึกกลับไปใส่เหตุผลทุกครั้งที่แก้ด้วยมือ</div>")
    if not reg["rows"]:
        parts.append("<div class='none'>✓ ไม่มีการปรับปรุงด้วยมือในงวดนี้ — ระบบจับคู่อัตโนมัติทั้งหมด</div>")
    else:
        parts.append("<table><thead><tr><th class='l'>วันเวลา</th><th class='l'>สาขา</th>"
            "<th class='l'>ประเภท</th><th class='l'>ช่องทาง / รายการ</th><th>จำนวน (บาท)</th>"
            "<th class='l'>โดย</th><th class='l'>เหตุผล</th></tr></thead><tbody>")
        for r in reg["rows"]:
            rsn = (esc(r["reason"]) if r["has_reason"]
                   else "<span class='big'>— ไม่ระบุเหตุผล —</span>")
            parts.append(f"<tr><td class='l muted'>{esc((r['at'] or '')[:16])}</td>"
                f"<td class='l'>{esc(r['branch_name'])}</td>"
                f"<td class='l'>{esc(r['kind'])}</td>"
                f"<td class='l'>{esc(r['lane'])}<div class='muted'>{esc(r['detail'])}</div></td>"
                f"<td class='mono'>{m(abs(r['amount']))}</td>"
                f"<td class='l'>{esc(r['by'])}</td>"
                f"<td class='l'>{rsn}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้สอบทาน (ผู้จัดการ)</div><div>วันที่</div></div>")
    parts.append(_ic_signed_footer(request))
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


@router.get("/api/reports/open-items-aging.print")
def open_items_aging_print(request: Request, period_month: str, branch_id: int | None = None,
                           lookback: int = 2):
    """Print-to-PDF aging of open reconciling items (งวดที่เลือก + ย้อนหลัง),
    bucketed 0-7 / 8-30 / >30 days from period-end. Stale items (>30d) flagged."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    lookback = max(0, min(6, lookback))
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        ag = _open_items_aging(conn, period_month, branch_id, lookback)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc, extra=f"&lookback={lookback}")
    b = ag["buckets"]
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>อายุรายการค้าง</title>", _ic_print_css("#b45309"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>อายุรายการค้าง (Aging of Open Items) — ณ {esc(ag['as_of'])}</h1>",
        f"<div class='sub'>รายการกระทบยอดที่ยังไม่ปิด นับอายุจากสิ้นงวด · "
        f"ตรวจงวด {esc(period_month)} ย้อนหลัง {lookback} เดือน "
        f"(ไม่รวมงวดที่อนุมัติ/ปิดแล้ว) · รวม <b>{ag['count']}</b> รายการ</div>",
        "<div class='kv'>"
        f"<div>0–7 วัน: <b class='ok'>{b['0-7']}</b></div>"
        f"<div>8–30 วัน: <b class='warn'>{b['8-30']}</b></div>"
        f"<div>เกิน 30 วัน: <b class='big'>{b['>30']}</b></div></div>"]
    if b[">30"]:
        parts.append("<div class='lead'>⚠ มีรายการค้างเกิน 30 วัน — ตามหลักการควบคุมภายใน "
                     "รายการกระทบยอดควรเคลียร์ภายในรอบถัดไป ของที่ค้างนานต้องมีคำอธิบาย/แผนปิด</div>")
    if not ag["items"]:
        parts.append("<div class='none'>✓ ไม่มีรายการค้าง — ทุกงวดที่ตรวจปิดครบหรืออยู่ในเกณฑ์</div>")
    else:
        parts.append("<table><thead><tr><th class='l'>อายุ (วัน)</th><th class='l'>ช่วง</th>"
            "<th class='l'>งวด</th><th class='l'>สาขา</th><th class='l'>ประเภท</th>"
            "<th class='l'>รายการที่ต้องทำ</th></tr></thead><tbody>")
        _bk_chip = {"0-7": "c-ok", "8-30": "c-warn", ">30": "c-err"}
        _ty_th = {"upload": "อัปโหลดเอกสาร", "setup": "ตั้งค่า", "diff": "ตรวจผลต่าง",
                  "waiting": "รอเอกสาร"}
        for it in ag["items"]:
            parts.append(f"<tr><td class='l mono'>{it['age_days']}</td>"
                f"<td class='l'><span class='chip {_bk_chip[it['bucket']]}'>{esc(it['bucket'])}</span></td>"
                f"<td class='l muted'>{esc(it['period_month'])}</td>"
                f"<td class='l'>{esc(it['branch_name'])}</td>"
                f"<td class='l'>{esc(_ty_th.get(it['type'], it['type']))}</td>"
                f"<td class='l'>{esc(it['label'])}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้จัดทำ</div><div>ผู้สอบทาน</div><div>วันที่</div></div>")
    parts.append(_ic_signed_footer(request))
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


@router.get("/api/reports/recon-certificate.print")
def recon_certificate_print(period_month: str, branch_id: int | None = None,
                            request: Request = None):
    """Print-to-PDF reconciliation sign-off certificate — one signed sheet per branch:
    per-channel POS vs money-in, residual after acknowledgements, preparer/reviewer +
    approval date, frozen working-paper hash, override count. No branch_id = all
    master branches, each on its own page."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v): return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
        if branch_id:
            _validate_master_branch(conn, branch_id)
            targets = [branch_id]
        else:
            targets = [b["branch_id"] for b in master]
        certs = [_recon_certificate(conn, bid, period_month) for bid in targets]
    picker = _ic_picker(period_month, branch_id, master, esc,
                        all_label="ทุกสาขา (พิมพ์ทีละหน้า)")

    def sheet(c, first):
        sess = c["session"] or {}
        st = sess.get("status") or "draft"
        st_th = _IC_SESS_TH.get(st, st)
        st_chip = ("c-ok" if st in ("approved", "closed")
                   else "c-warn" if st == "in_review" else "c-mut")
        prep = esc(sess.get("created_by") or "—")
        rev = esc(sess.get("approved_by") or "—")
        appd = esc((sess.get("approved_at") or "")[:16] or "—")
        snap = c["snapshot"]
        snap_line = (f"<div>หลักฐานแช่แข็ง (SHA-256): <b class='mono'>{esc(snap['sha256'][:16])}…</b> "
                     f"· {esc((snap['created_at'] or '')[:16])} · "
                     f"ผูกเอกสารต้นทาง {snap.get('source_doc_count', 0)} ไฟล์</div>" if snap else
                     "<div class='muted'>ยังไม่มีหลักฐานแช่แข็ง (อนุมัติงวดเพื่อสร้าง snapshot)</div>")
        out = [f"<div class='sheet{'' if first else ' pageb'}'>",
            "<h1>ใบรับรองการกระทบยอด (Reconciliation Certificate)</h1>",
            f"<div class='sub'>สาขา <b>{esc(c['branch_name'])}</b> · งวด <b>{esc(c['period_month'])}</b> "
            f"· สถานะ <span class='chip {st_chip}'>{esc(st_th)}</span> "
            f"· กระทบยอดแล้ว {c['progress_pct']}%</div>",
            "<div class='kv'>"
            f"<div>POS รวม: <b>฿{m(c['pos_total'])}</b></div>"
            f"<div>เงินเข้าจริงรวม: <b>฿{m(c['ext_total'])}</b></div>"
            f"<div>ผลต่างคงเหลือ (หลังรับทราบ): <b class='{'big' if abs(c['residual_total'])>100 else 'ok'}'>"
            f"฿{m(c['residual_total'])}</b></div>"
            f"<div>รายการรับทราบ: <b>{c['ack_count']}</b> (฿{m(c['ack_amount'])})</div>"
            f"<div>แก้ด้วยมือทั้งหมด: <b>{c['override_count']}</b> รายการ</div></div>",
            "<table><thead><tr><th class='l'>ช่องทาง</th><th>POS</th><th>เงินเข้าจริง</th>"
            "<th>ผลต่าง</th><th>คงเหลือ</th><th class='l'>สถานะ</th></tr></thead><tbody>"]
        for ch in c["channels"]:
            stx = "c-ok" if ch["status"] == "ok" else ("c-mut" if ch["status"] == "no_source"
                  else "c-warn" if ch["status"] == "waiting_doc" else "c-err")
            stl = ("รับทราบแล้ว" if ch["accepted"] else _IC_CHST_TH.get(ch["status"], ch["status"]))
            resid = ch["residual"]
            rcls = "" if abs(resid) <= 1 or ch["key"] == "COUPON" else (" big" if abs(resid) > 100 else " warn")
            rtxt = "—" if ch["key"] == "COUPON" else f"{resid:+,.0f}"
            out.append(f"<tr><td class='l'><b>{esc(ch['label'])}</b></td>"
                f"<td class='mono'>{m(ch['pos'])}</td><td class='mono'>{m(ch['ext'])}</td>"
                f"<td class='mono'>{ch['variance']:+,.0f}</td>"
                f"<td class='mono{rcls}'>{rtxt}</td>"
                f"<td class='l'><span class='chip {stx}'>{esc(stl)}</span></td></tr>")
        out.append(f"</tbody><tfoot><tr><td class='l'>รวม</td>"
            f"<td class='mono'>{m(c['pos_total'])}</td><td class='mono'>{m(c['ext_total'])}</td>"
            f"<td></td><td class='mono'>{c['residual_total']:+,.0f}</td><td></td></tr></tfoot></table>")
        out.append("<div class='lead'>" + snap_line + "</div>")
        out.append("<div class='sub' style='margin-top:14px'>ข้าพเจ้าขอรับรองว่าการกระทบยอด"
            "รายได้และเงินรับของสาขา/งวดข้างต้นถูกต้องและครบถ้วน ผลต่างที่คงเหลือได้รับการ"
            "ตรวจสอบและรับทราบตามเหตุผลที่บันทึกไว้ในทะเบียนการปรับปรุง</div>")
        out.append("<div class='sign'>"
            f"<div>ผู้จัดทำ<br><b>{prep}</b></div>"
            f"<div>ผู้สอบทาน/อนุมัติ<br><b>{rev}</b></div>"
            f"<div>วันที่อนุมัติ<br><b>{appd}</b></div></div>")
        out.append(_ic_signed_footer(request))
        out.append("</div>")
        return "".join(out)

    body = "".join(sheet(c, i == 0) for i, c in enumerate(certs)) or \
        "<div class='sheet'><div class='none'>ไม่มีข้อมูลสาขาสำหรับงวดนี้</div></div>"
    html_doc = ("<!doctype html><html lang='th'><head><meta charset='utf-8'>"
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>"
        "<title>ใบรับรองการกระทบยอด</title>" + _ic_print_css("#1f3a6b") + "</head><body>"
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>"
        + body + "</body></html>")
    return Response(content=html_doc, media_type="text/html; charset=utf-8")


def _close_register(conn, period_month, months=6, branch_id=None):
    """Cross-branch × multi-month close monitoring: status (draft/in_review/approved/
    closed), %reconciled, todo count, approver + approval date, and whether the frozen
    working-paper snapshot exists. Read-only; reuses reconcile_overview per month."""
    months = max(1, min(12, months))
    win = list(dict.fromkeys(
        [period_month] + [_ic_prev_month(period_month, k) for k in range(1, months)]))
    ph = ",".join("?" * len(win))
    sess = {(r["branch_id"], r["period_month"]): dict(r) for r in conn.execute(
        "SELECT branch_id, period_month, approved_by, approved_at "
        f"FROM recon_sessions WHERE period_month IN ({ph})", win)}
    snaps = {(r["branch_id"], r["period_month"]) for r in conn.execute(
        f"SELECT branch_id, period_month FROM recon_close_snapshot WHERE period_month IN ({ph})", win)}
    rows, roll = [], {}
    for mth in win:
        ov = reconcile_overview(conn, mth)
        for b in ov["branches"]:
            if branch_id and b["branch_id"] != branch_id:
                continue
            s = sess.get((b["branch_id"], mth), {})
            st = b.get("status") or "draft"
            rows.append({"period_month": mth, "branch_id": b["branch_id"],
                "branch_name": b["branch_name"], "status": st,
                "progress_pct": b["progress_pct"], "todo_count": b["todo_count"],
                "approved_by": s.get("approved_by"), "approved_at": s.get("approved_at"),
                "has_snapshot": (b["branch_id"], mth) in snaps})
            r = roll.setdefault(mth, {"closed": 0, "approved": 0, "open": 0, "total": 0})
            r["total"] += 1
            r["closed" if st == "closed" else "approved" if st == "approved" else "open"] += 1
    return {"period_month": period_month, "months": win, "rows": rows,
            "month_rollup": roll,
            "open_total": sum(1 for r in rows if r["status"] not in ("approved", "closed"))}


def _accepted_variance_summary(conn, period_month, branch_id=None, threshold=1000.0):
    """Write-off view: every ACKNOWLEDGED variance (recon_lane_adj accept = whole channel,
    recon_day_ack = per day) for the month, grouped by branch and by reason, with the
    grand total accepted and items over a policy threshold flagged. Read-only."""
    from collections import defaultdict
    bmap = {r["branch_id"]: r["name"] for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY}")}

    def keep(bid):
        return bid in bmap and (branch_id is None or bid == branch_id)

    items, var_cache = [], {}
    for a in conn.execute(
            "SELECT branch_id, period_month, category, bank, note, created_by, created_at "
            "FROM recon_lane_adj WHERE adj_type='accept' AND period_month=?", (period_month,)):
        if not keep(a["branch_id"]):
            continue
        key = (a["branch_id"], a["period_month"])
        if key not in var_cache:
            try:
                var_cache[key] = {c["key"]: c.get("residual", c["variance"])
                                  for c in reconcile_month(conn, a["branch_id"], a["period_month"])["channels"]}
            except Exception:
                var_cache[key] = {}
        items.append({"branch_id": a["branch_id"], "branch_name": bmap[a["branch_id"]],
            "scope": _ic_lane_th(a["category"], a["bank"]), "channel": a["category"],
            "amount": round(var_cache[key].get(a["category"], 0.0), 2),
            "reason": a["note"] or "", "by": a["created_by"] or "—",
            "at": a["created_at"] or "", "source": "channel"})
    for a in conn.execute(
            "SELECT branch_id, category, day, variance, note, created_by, created_at "
            "FROM recon_day_ack WHERE period_month=?", (period_month,)):
        if not keep(a["branch_id"]):
            continue
        items.append({"branch_id": a["branch_id"], "branch_name": bmap[a["branch_id"]],
            "scope": f"{_CH_LABEL.get(a['category'], a['category'])} · {a['day']}",
            "channel": a["category"], "amount": round(float(a["variance"] or 0), 2),
            "reason": a["note"] or "", "by": a["created_by"] or "—",
            "at": a["created_at"] or "", "source": "day"})
    items.sort(key=lambda x: -abs(x["amount"]))
    by_branch = defaultdict(lambda: {"count": 0, "total_abs": 0.0})
    by_reason = defaultdict(lambda: {"count": 0, "total_abs": 0.0})
    for it in items:
        gb = by_branch[(it["branch_id"], it["branch_name"])]
        gb["count"] += 1
        gb["total_abs"] += abs(it["amount"])
        gr = by_reason[(it["reason"].strip() or "ไม่ระบุเหตุผล")]
        gr["count"] += 1
        gr["total_abs"] += abs(it["amount"])
    over = [it for it in items if abs(it["amount"]) > threshold]
    return {"period_month": period_month, "threshold": threshold,
        "items": items, "count": len(items),
        "grand_total_abs": round(sum(abs(it["amount"]) for it in items), 2),
        "by_branch": sorted([{"branch_id": k[0], "branch_name": k[1], "count": v["count"],
                              "total_abs": round(v["total_abs"], 2)} for k, v in by_branch.items()],
                            key=lambda x: -x["total_abs"]),
        "by_reason": sorted([{"reason": k, "count": v["count"],
                              "total_abs": round(v["total_abs"], 2)} for k, v in by_reason.items()],
                            key=lambda x: -x["total_abs"]),
        "over_threshold": over, "over_count": len(over)}


@router.get("/api/reports/close-register")
def close_register_api(period_month: str, months: int = 6, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _close_register(conn, period_month, months, branch_id)


@router.get("/api/reports/accepted-variance")
def accepted_variance_api(period_month: str, branch_id: int | None = None,
                          threshold: float = 1000.0):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _accepted_variance_summary(conn, period_month, branch_id, max(0.0, threshold))


@router.get("/api/reports/close-register.print")
def close_register_print(request: Request, period_month: str, months: int = 6, branch_id: int | None = None):
    """Print-to-PDF close register: which branch-months are closed / approved / still
    open, who signed off + when, %reconciled, and whether the frozen snapshot exists."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    months = max(1, min(12, months))
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _close_register(conn, period_month, months, branch_id)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc, extra=f"&months={months}")
    st_chip = {"closed": "c-ok", "approved": "c-ok", "in_review": "c-warn",
               "draft": "c-mut"}
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>ทะเบียนการปิดเดือน</title>", _ic_print_css("#1d4ed8"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>ทะเบียนการปิดเดือน (Close Register) — {esc(months)} เดือนล่าสุดถึง {esc(period_month)}</h1>",
        f"<div class='sub'>สถานะการปิดงวดของทุกสาขาหลัก · ใครอนุมัติ · เมื่อไหร่ · กระทบยอดแล้วกี่ % · "
        f"มีหลักฐานแช่แข็งหรือยัง · <b>ยังไม่ปิดรวม {reg['open_total']}</b> สาขา-เดือน</div>"]
    if not reg["rows"]:
        parts.append("<div class='none'>ไม่มีข้อมูลสาขาสำหรับช่วงนี้</div>")
    for mth in reg["months"]:
        mr = reg["month_rollup"].get(mth)
        if not mr:
            continue
        parts.append(f"<h2>{esc(mth)} — ปิดแล้ว {mr['closed']} · อนุมัติ {mr['approved']} · "
                     f"<span class='warn'>ยังเปิด {mr['open']}</span> / {mr['total']} สาขา</h2>")
        parts.append("<table><thead><tr><th class='l'>สาขา</th><th class='l'>สถานะ</th>"
            "<th>กระทบยอด</th><th>คงค้าง</th><th class='l'>ผู้อนุมัติ</th>"
            "<th class='l'>วันที่อนุมัติ</th><th class='l'>หลักฐาน</th></tr></thead><tbody>")
        for r in [x for x in reg["rows"] if x["period_month"] == mth]:
            parts.append(f"<tr><td class='l'><b>{esc(r['branch_name'])}</b></td>"
                f"<td class='l'><span class='chip {st_chip.get(r['status'], 'c-mut')}'>"
                f"{esc(_IC_SESS_TH.get(r['status'], r['status']))}</span></td>"
                f"<td class='mono'>{r['progress_pct']}%</td>"
                f"<td class='mono'>{r['todo_count'] or ''}</td>"
                f"<td class='l'>{esc(r['approved_by'] or '—')}</td>"
                f"<td class='l muted'>{esc((r['approved_at'] or '')[:16] or '—')}</td>"
                f"<td class='l'>{'✔' if r['has_snapshot'] else '—'}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้สอบทาน</div><div>วันที่</div></div>")
    parts.append(_ic_signed_footer(request))
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


@router.get("/api/reports/accepted-variance.print")
def accepted_variance_print(request: Request, period_month: str, branch_id: int | None = None,
                            threshold: float = 1000.0):
    """Print-to-PDF accepted-variance (write-off) summary: total baht acknowledged,
    grouped by reason and by branch, with single items over the policy threshold flagged."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    def m(v): return f"{float(v or 0):,.0f}"
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    threshold = max(0.0, threshold)
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        s = _accepted_variance_summary(conn, period_month, branch_id, threshold)
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc,
                        extra=f"&threshold={int(threshold)}")
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>สรุปผลต่างที่รับทราบ</title>", _ic_print_css("#b91c1c"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>สรุปผลต่างที่รับทราบ (Accepted-Variance / Write-off) — งวด {esc(period_month)}</h1>",
        f"<div class='sub'>ยอดผลต่างที่ถูก “รับทราบ” แล้วปล่อยผ่าน (ทั้งช่อง + รายวัน) · "
        f"รวม <b class='big'>฿{m(s['grand_total_abs'])}</b> จาก <b>{s['count']}</b> รายการ · "
        f"เกินเกณฑ์ ฿{m(threshold)}/รายการ: <b class='big'>{s['over_count']}</b> รายการ</div>"]
    if not s["items"]:
        parts.append("<div class='none'>✓ ไม่มีผลต่างที่รับทราบในงวดนี้ — ทุกช่องทางตรงตามจริง</div>")
    else:
        if s["over_threshold"]:
            parts.append(f"<div class='lead'>⚠ มี {s['over_count']} รายการที่รับทราบเกิน "
                         f"฿{m(threshold)} ต่อรายการ — ผู้สอบบัญชีจะขอหลักฐาน/การอนุมัติเพิ่มเป็นพิเศษ</div>")
        parts.append("<h2>แยกตามเหตุผล</h2><table><thead><tr><th class='l'>เหตุผล</th>"
            "<th>จำนวนครั้ง</th><th>ยอดรวม (บาท)</th></tr></thead><tbody>")
        for r in s["by_reason"]:
            parts.append(f"<tr><td class='l'>{esc(r['reason'])}</td><td class='mono'>{r['count']}</td>"
                         f"<td class='mono'>{m(r['total_abs'])}</td></tr>")
        parts.append(f"</tbody><tfoot><tr><td class='l'>รวม</td><td class='mono'>{s['count']}</td>"
                     f"<td class='mono'>{m(s['grand_total_abs'])}</td></tr></tfoot></table>")
        parts.append("<h2>แยกตามสาขา</h2><table><thead><tr><th class='l'>สาขา</th>"
            "<th>จำนวนครั้ง</th><th>ยอดรวม (บาท)</th></tr></thead><tbody>")
        for r in s["by_branch"]:
            parts.append(f"<tr><td class='l'><b>{esc(r['branch_name'])}</b></td>"
                         f"<td class='mono'>{r['count']}</td><td class='mono'>{m(r['total_abs'])}</td></tr>")
        parts.append("</tbody></table>")
        parts.append("<h2>รายการทั้งหมด (เรียงยอดมากสุด)</h2><table><thead><tr>"
            "<th class='l'>สาขา</th><th class='l'>ช่องทาง / วัน</th><th>ยอด</th>"
            "<th class='l'>เหตุผล</th><th class='l'>โดย</th></tr></thead><tbody>")
        for it in s["items"]:
            big = " big" if abs(it["amount"]) > threshold else ""
            parts.append(f"<tr><td class='l'>{esc(it['branch_name'])}</td>"
                f"<td class='l'>{esc(it['scope'])}</td>"
                f"<td class='mono{big}'>{it['amount']:+,.0f}</td>"
                f"<td class='l'>{esc(it['reason']) or '<span class=\"big\">— ไม่ระบุ —</span>'}</td>"
                f"<td class='l'>{esc(it['by'])}</td></tr>")
        parts.append("</tbody></table>")
    parts.append("<div class='sign'><div>ผู้จัดทำ</div><div>ผู้สอบทาน (ผู้จัดการ)</div><div>วันที่</div></div>")
    parts.append(_ic_signed_footer(request))
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


# ── #6 Audit-trail viewer (change history) — READ-ONLY ────────────────────────
_AUDIT_ENTITY_TH = {"session": "งวด / สถานะ", "manual_match": "จับคู่ด้วยมือ",
    "channel_move": "ย้ายช่อง", "lane_adj": "ปรับ lane", "day_ack": "รับทราบรายวัน",
    "upload": "เอกสารอัปโหลด", "close_snapshot": "ปิดงวด / หลักฐาน"}


_AUDIT_ACTION_TH = {"create": "สร้าง", "update": "แก้ไข", "delete": "ลบ", "accept": "รับทราบ"}


# entity → (parent table, primary-key column) so an audit row can be resolved to a
# branch+period. Deleted parents drop out (their row is gone); the append-only audit
# log itself is never mutated.
_AUDIT_SRC = [("session", "recon_sessions", "session_id"),
              ("manual_match", "recon_manual_match", "match_id"),
              ("channel_move", "recon_channel_move", "id"),
              ("lane_adj", "recon_lane_adj", "id"),
              ("day_ack", "recon_day_ack", "id"),
              ("upload", "recon_uploads", "upload_id"),
              ("close_snapshot", "recon_close_snapshot", "id")]


def _audit_trail(conn, period_month, branch_id=None, limit=400):
    """Branch+month-scoped change history from the append-only audit log: who changed
    what, when, across sessions/overrides/uploads/snapshots. Read-only."""
    bmap = {r["branch_id"]: r["name"] for r in conn.execute(
        f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY}")}
    rows = []
    for entity, tbl, pk in _AUDIT_SRC:
        try:
            cur = conn.execute(
                f"SELECT a.ts, a.action, a.field, a.old_val, a.new_val, a.\"by\" AS by_, "
                f"p.branch_id AS bid, p.period_month AS pm "
                f"FROM recon_audit_log a JOIN {tbl} p ON p.{pk}=a.entity_id "
                f"WHERE a.entity=? AND p.period_month=?", (entity, period_month))
        except sqlite3.OperationalError:
            continue
        for r in cur:
            bid = r["bid"]
            if bid not in bmap or (branch_id is not None and bid != branch_id):
                continue
            if r["action"] == "update" and (r["old_val"] is not None or r["new_val"] is not None):
                detail = f"{r['field'] or ''} {r['old_val'] or '—'} → {r['new_val'] or '—'}".strip()
            else:
                detail = r["new_val"] or r["field"] or ""
            rows.append({"ts": r["ts"], "by": r["by_"] or "—", "entity": entity,
                "entity_th": _AUDIT_ENTITY_TH.get(entity, entity),
                "action": r["action"], "action_th": _AUDIT_ACTION_TH.get(r["action"], r["action"]),
                "detail": detail, "branch_id": bid, "branch_name": bmap.get(bid, str(bid))})
    rows.sort(key=lambda x: (x["ts"] or ""), reverse=True)
    return {"period_month": period_month, "branch_id": branch_id,
            "count": len(rows), "rows": rows[:limit], "shown": min(len(rows), limit)}


@router.get("/api/reports/audit-trail")
def audit_trail_api(period_month: str, branch_id: int | None = None, limit: int = 400):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        return _audit_trail(conn, period_month, branch_id, max(1, min(2000, limit)))


@router.get("/api/reports/audit-trail.print")
def audit_trail_print(request: Request, period_month: str, branch_id: int | None = None, limit: int = 400):
    """Print-to-PDF change history: every recorded change for the month, newest first."""
    import html as _html
    def esc(s): return _html.escape(str(s or ""))
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        tr = _audit_trail(conn, period_month, branch_id, max(1, min(2000, limit)))
        master = conn.execute(
            f"SELECT branch_id, name FROM branches WHERE {_MASTER_ONLY} ORDER BY name").fetchall()
    picker = _ic_picker(period_month, branch_id, master, esc)
    parts = ["<!doctype html><html lang='th'><head><meta charset='utf-8'>",
        "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
        "<title>ประวัติการแก้ไข</title>", _ic_print_css("#0891b2"), "</head><body>",
        "<div class='toolbar'>" + picker
        + "<button class='btn' onclick='window.print()'>พิมพ์ / บันทึก PDF</button></div>",
        "<div class='sheet'>",
        f"<h1>ประวัติการแก้ไข (Audit Trail) — งวด {esc(period_month)}</h1>",
        f"<div class='sub'>ทุกการเปลี่ยนแปลงที่ระบบบันทึกไว้ (เรียงใหม่สุดก่อน) · บันทึกแบบเพิ่มอย่างเดียว "
        f"แก้/ลบไม่ได้ · แสดง {tr['shown']} จาก {tr['count']} รายการ</div>"]
    if not tr["rows"]:
        parts.append("<div class='none'>ไม่มีการเปลี่ยนแปลงที่บันทึกไว้สำหรับงวดนี้</div>")
    else:
        parts.append("<table><thead><tr><th class='l'>เวลา</th><th class='l'>สาขา</th>"
            "<th class='l'>ประเภท</th><th class='l'>การกระทำ</th><th class='l'>รายละเอียด</th>"
            "<th class='l'>โดย</th></tr></thead><tbody>")
        for r in tr["rows"]:
            parts.append(f"<tr><td class='l muted'>{esc((r['ts'] or '')[:16])}</td>"
                f"<td class='l'>{esc(r['branch_name'])}</td><td class='l'>{esc(r['entity_th'])}</td>"
                f"<td class='l'>{esc(r['action_th'])}</td><td class='l'>{esc(r['detail'])}</td>"
                f"<td class='l'>{esc(r['by'])}</td></tr>")
        parts.append("</tbody></table>")
    parts.append(_ic_signed_footer(request))
    parts.append("</div></body></html>")
    return Response(content="".join(parts), media_type="text/html; charset=utf-8")


# ── Excel (.xlsx) exports for the internal-control reports ────────────────────
def _xlsx_response(sheets, fname):
    """sheets = [(title, [headers...], [[cells...], ...]), ...] → styled .xlsx Response."""
    import io
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    first = True
    for title, headers, body in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        first = False
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for row in body:
            _xl_append(ws, row)
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{fname}\""})


@router.get("/api/reports/override-register.xlsx")
def override_register_xlsx(period_month: str, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _override_register(conn, period_month, branch_id)
    body = [[(r["at"] or "")[:16], r["branch_name"], r["kind"], r["lane"],
             round(abs(r["amount"]), 2), r["by"],
             r["reason"] if r["has_reason"] else "— ไม่ระบุเหตุผล —"] for r in reg["rows"]]
    return _xlsx_response(
        [("ทะเบียนการปรับปรุงด้วยมือ",
          ["วันเวลา", "สาขา", "ประเภท", "ช่องทาง/รายการ", "จำนวน", "โดย", "เหตุผล"], body)],
        f"override_register_{period_month}.xlsx")


@router.get("/api/reports/open-items-aging.xlsx")
def open_items_aging_xlsx(period_month: str, branch_id: int | None = None, lookback: int = 2):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        ag = _open_items_aging(conn, period_month, branch_id, max(0, min(6, lookback)))
    _ty = {"upload": "อัปโหลดเอกสาร", "setup": "ตั้งค่า", "diff": "ตรวจผลต่าง", "waiting": "รอเอกสาร"}
    body = [[it["age_days"], it["bucket"], it["period_month"], it["branch_name"],
             _ty.get(it["type"], it["type"]), it["label"]] for it in ag["items"]]
    return _xlsx_response(
        [("อายุรายการค้าง",
          ["อายุ(วัน)", "ช่วง", "งวด", "สาขา", "ประเภท", "รายการที่ต้องทำ"], body)],
        f"open_items_aging_{period_month}.xlsx")


@router.get("/api/reports/close-register.xlsx")
def close_register_xlsx(period_month: str, months: int = 6, branch_id: int | None = None):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        reg = _close_register(conn, period_month, max(1, min(12, months)), branch_id)
    body = [[r["period_month"], r["branch_name"], _IC_SESS_TH.get(r["status"], r["status"]),
             r["progress_pct"], r["todo_count"], r["approved_by"] or "—",
             (r["approved_at"] or "")[:16], "มี" if r["has_snapshot"] else "—"] for r in reg["rows"]]
    return _xlsx_response(
        [("ทะเบียนการปิดเดือน",
          ["งวด", "สาขา", "สถานะ", "กระทบยอด%", "คงค้าง", "ผู้อนุมัติ", "วันที่อนุมัติ", "หลักฐาน"], body)],
        f"close_register_{period_month}.xlsx")


@router.get("/api/reports/accepted-variance.xlsx")
def accepted_variance_xlsx(period_month: str, branch_id: int | None = None,
                           threshold: float = 1000.0):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        s = _accepted_variance_summary(conn, period_month, branch_id, max(0.0, threshold))
    by_reason = [[r["reason"], r["count"], r["total_abs"]] for r in s["by_reason"]]
    by_branch = [[r["branch_name"], r["count"], r["total_abs"]] for r in s["by_branch"]]
    items = [[it["branch_name"], it["scope"], round(it["amount"], 2), it["reason"], it["by"]]
             for it in s["items"]]
    return _xlsx_response([
        ("แยกตามเหตุผล", ["เหตุผล", "จำนวนครั้ง", "ยอดรวม"], by_reason),
        ("แยกตามสาขา", ["สาขา", "จำนวนครั้ง", "ยอดรวม"], by_branch),
        ("รายการทั้งหมด", ["สาขา", "ช่องทาง/วัน", "ยอด", "เหตุผล", "โดย"], items)],
        f"accepted_variance_{period_month}.xlsx")


@router.get("/api/reports/audit-trail.xlsx")
def audit_trail_xlsx(period_month: str, branch_id: int | None = None, limit: int = 2000):
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
        tr = _audit_trail(conn, period_month, branch_id, max(1, min(5000, limit)))
    body = [[(r["ts"] or "")[:16], r["branch_name"], r["entity_th"], r["action_th"],
             r["detail"], r["by"]] for r in tr["rows"]]
    return _xlsx_response(
        [("ประวัติการแก้ไข", ["เวลา", "สาขา", "ประเภท", "การกระทำ", "รายละเอียด", "โดย"], body)],
        f"audit_trail_{period_month}.xlsx")


# ── Auditor Pack: one-click ZIP of all control evidence for a period ──────────
@router.get("/api/reports/auditor-pack.zip")
def auditor_pack_zip(period_month: str, branch_id: int | None = None):
    """Bundle every internal-control report for the period into one ZIP to hand to the
    auditor: sign-off certificate, override register, open-items aging, close register,
    accepted-variance (write-off), and the audit trail — Excel + the signable HTML."""
    import io
    import zipfile
    if not _PERIOD_RE.match(period_month or ""):
        raise HTTPException(400, "period_month must be in 'YYYY-MM' format.")
    with get_db() as conn:
        if branch_id:
            _validate_master_branch(conn, branch_id)
    pm, bsuf = period_month, (f"_b{branch_id}" if branch_id else "")

    def _body(resp):
        b = resp.body
        return b if isinstance(b, (bytes, bytearray)) else str(b).encode("utf-8")

    # (filename, builder thunk) — each wrapped so one failure can't sink the pack
    parts = [
        (f"01_ใบรับรองการกระทบยอด_{pm}.html", lambda: recon_certificate_print(pm, branch_id)),
        (f"02_ทะเบียนการปรับปรุงด้วยมือ_{pm}.xlsx", lambda: override_register_xlsx(pm, branch_id)),
        (f"03_อายุรายการค้าง_{pm}.xlsx", lambda: open_items_aging_xlsx(pm, branch_id, 2)),
        (f"04_ทะเบียนการปิดเดือน_{pm}.xlsx", lambda: close_register_xlsx(pm, 6, branch_id)),
        (f"05_สรุปผลต่างที่รับทราบ_{pm}.xlsx", lambda: accepted_variance_xlsx(pm, branch_id, 1000.0)),
        (f"06_ประวัติการแก้ไข_{pm}.xlsx", lambda: audit_trail_xlsx(pm, branch_id, 5000)),
    ]
    readme = (
        "ชุดเอกสารผู้สอบบัญชี (Auditor Pack)\n"
        f"งวด: {pm}{(' · สาขา ' + str(branch_id)) if branch_id else ' · ทุกสาขาหลัก'}\n"
        "สร้างจากระบบกระทบยอดรายได้ (Oshinei recon) — อ่านอย่างเดียว ไม่แก้ตัวเลข\n\n"
        "ในชุดนี้:\n"
        "  01 ใบรับรองการกระทบยอด (มีช่องเซ็น + เลข SHA-256 ของงบที่ปิด)\n"
        "  02 ทะเบียนการปรับปรุงด้วยมือ (ทุกการแก้มือ ใคร·เมื่อ·เหตุผล)\n"
        "  03 อายุรายการค้าง (0-7/8-30/เกิน 30 วัน)\n"
        "  04 ทะเบียนการปิดเดือน (ทุกสาขา × 6 เดือน · ใครอนุมัติ)\n"
        "  05 สรุปผลต่างที่รับทราบ / write-off (แยกตามเหตุผล·สาขา)\n"
        "  06 ประวัติการแก้ไข (audit trail — เพิ่มอย่างเดียว แก้ไม่ได้)\n")
    buf = io.BytesIO()
    failed = []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, thunk in parts:
            try:
                z.writestr(name, _body(thunk()))
            except Exception as e:                       # never let one report sink the pack
                failed.append(f"{name}: {e}")
        if failed:
            readme += "\nหมายเหตุ — รายงานที่สร้างไม่สำเร็จ:\n  " + "\n  ".join(failed) + "\n"
        z.writestr("00_README.txt", readme)
    return Response(
        content=buf.getvalue(), media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename=\"auditor_pack_{pm}{bsuf}.zip\""})
