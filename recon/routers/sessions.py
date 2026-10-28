"""Reconciliation sessions and what hangs off them: discrepancies, lines,
bank lines, snapshots, day acknowledgements.

Moved out of recon/main.py in Phase 7c (2026-09) by the AST: these endpoints and
the helpers only they use, unchanged. URLs are unchanged.
"""
from fastapi import APIRouter
from .. import auth as _A
from .auth import _ip
from ..core import _DAY_RE
from ..core import _SNAPSHOT_DIR
from ..core import _actor
from ..core import _assert_period_open
from ..core import _assert_bank_line_open
from ..core import _assert_bank_line_release_ok
from ..core import _audit
from ..core import _require_manager
from ..core import _validate_master_branch
from ..core import _validate_period
from ..core import get_db
from ..core import require_write_token
from ..matcher import bank_statement_coverage
from ..matcher import cash_three_way
from ..matcher import category_reconciliation
from ..matcher import coupon_reconciliation
from ..matcher import daily_reconciliation
from ..matcher import lane_lines
from ..matcher import payout_reconciliation
from ..matcher import run_auto_match
from ..reconcile_v2 import CHANNELS as _V2_CHANNELS
from ..reconcile_v2 import reconcile_month
from datetime import datetime as _dt
from fastapi import Depends
from fastapi import Form
from fastapi import HTTPException
from fastapi import Header
from fastapi import Request
from fastapi.responses import Response
from pathlib import Path
from pydantic import BaseModel
from typing import Literal
import hashlib
import json
import sqlite3

router = APIRouter()


class _IgnoreBody(BaseModel):
    ignored: bool
    note:    str = ""
    by:      str = ""


class _BulkIgnoreBody(BaseModel):
    line_ids: list[int]
    ignored:  bool = True
    note:     str = ""
    by:       str = ""


def _set_ignore(conn, line_id: int, ignored: bool, by: str, note: str,
                lock_cache: dict | None = None) -> str:
    """Flag one bank credit as not-revenue (or undo it). Raises HTTPException with
    the reason a line cannot be changed; the caller commits. Shared by the
    single PATCH and the bulk POST so both apply exactly the same rules.
    Lock: the line's own month must be open, and a signed-off PREVIOUS month must
    not rest on it (ignoring a credit it claimed) nor start claiming it (restoring)."""
    line = conn.execute(
        "SELECT branch_id, txn_date, description, credit, match_status "
        "FROM recon_bank_lines WHERE id=?", (line_id,)).fetchone()
    if not line:
        raise HTTPException(404, "bank line not found")
    _assert_bank_line_open(conn, line["branch_id"], line_id, line["txn_date"], lock_cache)
    if ignored and line["match_status"] == "manual":
        raise HTTPException(409, "line is manually matched — undo the manual match first")
    if ignored and line["match_status"] == "excluded":
        raise HTTPException(409, "line is already ignored")
    if not ignored and line["match_status"] != "excluded":
        raise HTTPException(409, "line is not ignored")
    new_status = "excluded" if ignored else "unmatched"
    if not conn.in_transaction:
        conn.execute("BEGIN")                  # so RELEASE below never commits on its own
    conn.execute("SAVEPOINT set_ignore")
    conn.execute("UPDATE recon_bank_lines SET match_status=? WHERE id=?",
                 (new_status, line_id))
    if not ignored:
        try:
            _assert_bank_line_release_ok(conn, line["branch_id"], line_id, line["txn_date"])
        except HTTPException:
            conn.execute("ROLLBACK TO set_ignore")
            conn.execute("RELEASE set_ignore")
            raise
    conn.execute("RELEASE set_ignore")
    _audit(conn, "bank_line", line_id, "update", by or "staff",
           [{"field": "match_status", "old": line["match_status"], "new": new_status},
            {"field": "note", "old": None,
             "new": note or f"{line['txn_date']} {line['description']} {line['credit']}"}])
    return new_status


@router.patch("/api/bank-lines/{line_id}/ignore", dependencies=[Depends(require_write_token)])
def set_bank_line_ignore(line_id: int, body: _IgnoreBody, request: Request):
    """Flag a bank credit as an internal transfer (not revenue) so it stops
    polluting the lanes — e.g. a ฿2M intercompany sweep showing as TRANSFER
    ext_only. Reversible; survives Run Match; audited with the note."""
    with get_db() as conn:
        new_status = _set_ignore(conn, line_id, body.ignored, _actor(request, body.by), body.note)
        conn.commit()
    return {"line_id": line_id, "match_status": new_status}


@router.post("/api/bank-lines/bulk-ignore", dependencies=[Depends(require_write_token)])
def bulk_bank_line_ignore(body: _BulkIgnoreBody, request: Request):
    """Exclude (or restore) many bank credits in one request — the suggestion
    card's "ตัดออกทั้งหมด" and the month-end IORSDT sweep. Same rules as the
    single PATCH, applied per line: a line that cannot change is reported in
    `skipped` with its reason and the rest still go through; one commit."""
    ids = list(dict.fromkeys(int(i) for i in body.line_ids))
    if not ids:
        raise HTTPException(400, "line_ids is empty")
    if len(ids) > 500:
        raise HTTPException(400, "at most 500 lines per request")
    done, skipped = [], []
    actor = _actor(request, body.by)
    lock_cache: dict = {}
    with get_db() as conn:
        for line_id in ids:
            try:
                _set_ignore(conn, line_id, body.ignored, actor, body.note, lock_cache)
                done.append(line_id)
            except HTTPException as e:
                skipped.append({"line_id": line_id, "reason": e.detail})
        conn.commit()
    return {"ignored" if body.ignored else "restored": done, "skipped": skipped,
            "match_status": "excluded" if body.ignored else "unmatched"}


class SessionCreate(BaseModel):
    branch_id:    int
    period_month: str
    notes:        str = ""
    created_by:   str = ""


@router.post("/api/sessions", dependencies=[Depends(require_write_token)])
def create_session(body: SessionCreate, request: Request):
    _validate_period(body.period_month)
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        try:
            cur = conn.execute("""
                INSERT INTO recon_sessions (branch_id, period_month, notes, created_by)
                VALUES (?,?,?,?)
            """, (body.branch_id, body.period_month, body.notes, _actor(request, body.created_by, default="")))
            conn.commit()
            return {"session_id": cur.lastrowid}
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"Session already exists for branch {body.branch_id} / {body.period_month}")


@router.get("/api/sessions")
def list_sessions(branch_id: int | None = None, period_month: str | None = None):
    with get_db() as conn:
        filters, params = [], []
        if branch_id:
            filters.append("s.branch_id=?");    params.append(branch_id)
        if period_month:
            filters.append("s.period_month=?"); params.append(period_month)
        where = ("WHERE " + " AND ".join(filters)) if filters else ""
        rows = conn.execute(f"""
            SELECT s.*, b.name AS branch_name
            FROM recon_sessions s
            JOIN branches b ON b.branch_id = s.branch_id
            {where}
            ORDER BY s.period_month DESC, b.name
        """, params).fetchall()
        return [dict(r) for r in rows]


@router.get("/api/sessions/{session_id}")
def get_session(session_id: int):
    with get_db() as conn:
        row = conn.execute("""
            SELECT s.*, b.name AS branch_name
            FROM recon_sessions s
            JOIN branches b ON b.branch_id = s.branch_id
            WHERE s.session_id=?
        """, (session_id,)).fetchone()
        if not row:
            raise HTTPException(404)
        return dict(row)


@router.post("/api/sessions/{session_id}/run-match", dependencies=[Depends(require_write_token)])
def run_match(session_id: int):
    with get_db() as conn:
        session = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)
        ).fetchone()
        if not session:
            raise HTTPException(404)
        _assert_period_open(conn, session["branch_id"], session["period_month"])
        result = run_auto_match(conn, session_id, session["branch_id"], session["period_month"])
    return result


@router.patch("/api/sessions/{session_id}/status", dependencies=[Depends(require_write_token)])
def update_session_status(
    session_id:  int,
    request:     Request,
    status:      str = Form(...),
    approved_by: str = Form(""),
    reason:      str = Form(""),
    password:    str = Form(""),
    x_manager_token: str = Header(None),
):
    """Move a month along draft → in_review → approved (→ closed), or reopen it.
    Approving / closing / reopening is a sign-off: login OFF → the manager code;
    login ON → the manager ROLE plus a step-up re-authentication (the signed-in
    user's own `password` in the form, checked last, after every other gate, and
    counted toward the login lockout)."""
    valid = {"draft", "in_review", "approved", "closed"}
    if status not in valid:
        raise HTTPException(400, f"status must be one of {valid}")
    if status in ("approved", "closed"):
        _require_manager(x_manager_token, request)
    actor = _actor(request, approved_by, default="")        # blank when unknown → SoD-safe
    if status == "approved":
        # Freshness: the approval snapshot, Excel export and the pending gate below
        # read the PERSISTED match results (recon_category_result / discrepancies),
        # which only Run Match refreshes. If data changed since (an upload, manual
        # match, adjustment, ignore) re-run it first; the locked section re-checks.
        with get_db() as conn:
            s0 = conn.execute("SELECT branch_id, period_month, status FROM recon_sessions "
                              "WHERE session_id=?", (session_id,)).fetchone()
            if s0 and s0["status"] not in ("approved", "closed") and \
                    _match_freshness(conn, session_id, s0["branch_id"], s0["period_month"])["stale"]:
                run_auto_match(conn, session_id, s0["branch_id"], s0["period_month"])
    with get_db() as conn:
        # Read-check-write under one write lock: two managers approving at the same
        # moment must not both pass the checks below (the UPDATEs are guarded too).
        conn.execute("BEGIN IMMEDIATE")
        old_row = conn.execute(
            "SELECT status, created_by, branch_id, period_month FROM recon_sessions "
            "WHERE session_id=?", (session_id,)).fetchone()
        if not old_row:
            raise HTTPException(404, f"session {session_id} not found")
        old_status = old_row["status"]
        # Re-approving would silently replace the approver (no audit, no snapshot).
        if status == old_status and status in ("approved", "closed"):
            raise HTTPException(409, f"งวดนี้อยู่ในสถานะ {status} แล้ว")
        # 'closed' is the archive step AFTER sign-off: SoD and the approver stamp
        # happen at 'approved', so a month can't skip past them straight to closed.
        if status == "closed" and old_status != "approved":
            raise HTTPException(
                409, "ปิดงวดได้หลังอนุมัติแล้วเท่านั้น — ให้ผู้อนุมัติ (ไม่ใช่ผู้จัดทำ) อนุมัติก่อน")
        # Reopen gate (control fix 2026-06-19): unlocking an already-signed-off month
        # (approved/closed → draft/in_review) is sensitive — require the MANAGER code +
        # a logged reason, just like approval. The next re-approval re-freezes a fresh
        # snapshot, so the working paper always matches the final figures.
        reopening = old_status in ("approved", "closed") and status not in ("approved", "closed")
        if reopening:
            _require_manager(x_manager_token, request)
            if len((reason or "").strip()) < 10:
                raise HTTPException(
                    400, "การเปิดงวดที่ปิดแล้วต้องระบุเหตุผล (อย่างน้อย 10 ตัวอักษร)")
        # Approval gate: don't let a month be signed off while issues are open.
        if status in ("approved", "closed"):
            pending = conn.execute(
                "SELECT COUNT(*) FROM recon_discrepancies "
                "WHERE session_id=? AND resolution='pending'",
                (session_id,)).fetchone()[0]
            if pending:
                raise HTTPException(
                    409, f"ยังปิดไม่ได้ — มีข้อแตกต่างค้าง {pending} รายการ (ระบบเพิ่งจับคู่ใหม่ก่อนลงนาม) · "
                         "จัดการหรือรับทราบในหน้า ข้อแตกต่าง ก่อน แล้วลงนามอีกครั้ง")
        if status == "approved":
            # Data changed between the re-run above and this lock → refuse, never
            # sign off figures the persisted results no longer reflect.
            if _match_freshness(conn, session_id, old_row["branch_id"],
                                old_row["period_month"])["stale"]:
                raise HTTPException(
                    409, "ข้อมูลของเดือนนี้เปลี่ยนหลังการจับคู่ล่าสุด — กดจับคู่ใหม่ แล้วอนุมัติอีกครั้ง")
            # Segregation of duties (AUD-4, widened 2026-10): nobody who worked on the
            # month — opened it, submitted it for review, matched/moved/adjusted/
            # acknowledged/resolved/ignored anything in it — may approve it. And the
            # month must HAVE an identifiable preparer: a blank or machine-token
            # creator with no other named worker is refused (send it for review first).
            _assert_sod(conn, session_id, old_row, actor)
        # Step-up (login ON): a sign-off is the user's own act — re-enter the password.
        # Last gate, so a refused attempt never burns a try on a month that could not
        # be signed anyway; a wrong one counts toward the same lockout as the login.
        if _A.login_enabled() and (status in ("approved", "closed") or reopening):
            _stepup(conn, request, password, session_id, status)
        # Only stamp approval metadata when the session is actually approved;
        # moving back to draft/in_review must not record a (false) approval.
        if status == "approved":
            cur = conn.execute("""
                UPDATE recon_sessions
                SET status=?, approved_by=?, approved_at=datetime('now')
                WHERE session_id=? AND status=?
            """, (status, actor, session_id, old_status))
        else:
            cur = conn.execute(
                "UPDATE recon_sessions SET status=? WHERE session_id=? AND status=?",
                (status, session_id, old_status),
            )
        if cur.rowcount != 1:
            raise HTTPException(409, "สถานะงวดถูกเปลี่ยนระหว่างทำรายการ — โหลดหน้าใหม่แล้วลองอีกครั้ง")
        if old_status != status:
            diffs = [{"field": "status", "old": old_status, "new": status}]
            if reopening:
                diffs.append({"field": "reopen_reason", "old": None, "new": reason.strip()})
            _audit(conn, "session", session_id, "update", actor or "staff", diffs)
        conn.commit()
    # AUD-5: freeze the working paper at each sign-off transition. Best-effort — the
    # export can be regenerated, so a snapshot hiccup must not block close. Built
    # after the status commit so the workbook render doesn't hold the write lock.
    snap = None
    if status in ("approved", "closed") and old_status != status:
        try:
            with get_db() as conn:
                full = conn.execute(
                    "SELECT s.*, b.name AS branch_name FROM recon_sessions s "
                    "JOIN branches b ON b.branch_id=s.branch_id WHERE s.session_id=?",
                    (session_id,)).fetchone()
                snap = _write_close_snapshot(conn, full, status, actor)
                conn.commit()
        except Exception:
            snap = None
    return {"session_id": session_id, "status": status, "snapshot": snap}


# Placeholder names older code stamped when the actor was unknown — not identities.
_NO_IDENTITY = {"", "staff", "manual", "manager", "system"}


def _ident(v) -> str:
    """Normalised identity, or '' for a placeholder / machine-token actor."""
    s = str(v or "").strip().lower()
    return "" if (s in _NO_IDENTITY or s.startswith("token")) else s


def _month_preparers(conn, session_id: int, branch_id: int, period_month: str,
                     created_by) -> set:
    """Every named person who did work on this branch-month (see _assert_sod)."""
    names = {created_by}
    q = [
        ("SELECT matched_by FROM recon_manual_match WHERE branch_id=? AND period_month=?",
         (branch_id, period_month)),
        ("SELECT created_by FROM recon_lane_adj WHERE branch_id=? AND period_month=?",
         (branch_id, period_month)),
        ("SELECT created_by FROM recon_day_ack WHERE branch_id=? AND period_month=?",
         (branch_id, period_month)),
        ("SELECT created_by FROM recon_channel_move WHERE branch_id=? AND period_month=?",
         (branch_id, period_month)),
        ("SELECT created_by FROM recon_deposit_slip WHERE branch_id=? AND substr(sale_day,1,7)=?",
         (branch_id, period_month)),
        ("SELECT resolved_by FROM recon_discrepancies WHERE session_id=? AND resolution!='pending'",
         (session_id,)),
        # who sent it for review (draft → in_review)
        ("SELECT by FROM recon_audit_log WHERE entity='session' AND entity_id=? "
         "AND field='status' AND new_val='in_review'", (session_id,)),
        # who ignored / restored a bank line of the month
        ("SELECT a.by FROM recon_audit_log a JOIN recon_bank_lines l ON l.id=a.entity_id "
         "WHERE a.entity='bank_line' AND l.branch_id=? AND substr(l.txn_date,1,7)=?",
         (branch_id, period_month)),
    ]
    for sql, params in q:
        try:
            names |= {r[0] for r in conn.execute(sql, params)}
        except sqlite3.OperationalError:          # table absent on an old DB
            pass
    return {n for n in (_ident(x) for x in names) if n}


def _assert_sod(conn, session_id: int, row, actor: str) -> None:
    appr = _ident(actor)
    if not appr:
        raise HTTPException(
            409, "แบ่งแยกหน้าที่ (SoD): ไม่ทราบตัวผู้อนุมัติ — เข้าสู่ระบบหรือระบุชื่อผู้อนุมัติ")
    preparers = _month_preparers(conn, session_id, row["branch_id"], row["period_month"],
                                 row["created_by"])
    if appr in preparers:
        raise HTTPException(
            409, "แบ่งแยกหน้าที่ (SoD): ผู้อนุมัติได้ทำรายการในเดือนนี้แล้ว (เป็นผู้จัดทำ) — "
                 "ต้องให้ผู้อื่นเป็นผู้อนุมัติ")
    if not preparers:
        raise HTTPException(
            409, "แบ่งแยกหน้าที่ (SoD): เดือนนี้ไม่มีผู้จัดทำที่ระบุตัวได้ — "
                 "ให้ผู้จัดทำกด ส่งตรวจ ก่อน แล้วจึงอนุมัติ")


def _stepup(conn, request, password: str, session_id: int, status: str) -> None:
    """403 unless `password` is the signed-in user's own. A counted failure (or an
    attempt while locked) is audited as 'stepup_fail' and committed before the
    refusal; the caller's transaction has written nothing else yet."""
    u = _A.current_user(request)
    if u is None:                                     # require_role ran first; belt and braces
        raise HTTPException(401, _A.MSG_NEED_LOGIN)
    res = _A.verify_stepup(conn, u["id"], password or "")
    if res == "ok":
        return
    if res != "blank":
        _audit(conn, "recon_auth", u["id"], "stepup_fail", u["username"],
               [{"field": "reason", "new": res}, {"field": "session_id", "new": session_id},
                {"field": "status", "new": status}, {"field": "ip", "new": _ip(request)}])
        conn.commit()
    raise HTTPException(403, _A.stepup_message(res))


def _norm_ts(s):
    """Normalise a stored timestamp for lexical comparison. computed_at /
    uploaded_at use 'YYYY-MM-DD HH:MM:SS' while audit_log.ts uses a 'T' separator
    — line them up so a plain string `>` is correct."""
    return s.replace("T", " ")[:19] if s else None


def _match_freshness(conn, session_id, branch_id, period_month):
    """When did Run Match last persist this session's category snapshot, vs. the
    latest data change (upload, manual match, lane adjustment, internal-transfer
    ignore)? The live `/summary` cards are always recomputed, but the *persisted*
    `recon_category_result` (Excel export, `/categories`, approval snapshot) only
    refreshes on Run Match — `stale=True` means it lags and staff should re-run."""
    computed_at = _norm_ts(conn.execute(
        "SELECT MAX(computed_at) FROM recon_category_result WHERE session_id=?",
        (session_id,)).fetchone()[0])
    changes = [
        conn.execute(
            "SELECT MAX(COALESCE(confirmed_at, uploaded_at)) FROM recon_uploads "
            "WHERE branch_id=? AND period_month=? AND status!='error'",
            (branch_id, period_month)).fetchone()[0],
        conn.execute(
            "SELECT MAX(matched_at) FROM recon_manual_match WHERE branch_id=? AND period_month=?",
            (branch_id, period_month)).fetchone()[0],
        conn.execute(
            "SELECT MAX(created_at) FROM recon_lane_adj WHERE branch_id=? AND period_month=?",
            (branch_id, period_month)).fetchone()[0],
        conn.execute(
            "SELECT MAX(a.ts) FROM recon_audit_log a "
            "JOIN recon_bank_lines bl ON bl.id = a.entity_id "
            "WHERE a.entity='bank_line' AND bl.branch_id=? "
            "AND strftime('%Y-%m', bl.txn_date)=?",
            (branch_id, period_month)).fetchone()[0],
    ]
    norm = [_norm_ts(x) for x in changes if x]
    last_change_at = max(norm) if norm else None
    # Stale when results were never computed, or a change post-dates the last run.
    stale = computed_at is None or (last_change_at is not None and last_change_at > computed_at)
    return {"computed_at": computed_at, "last_change_at": last_change_at, "stale": stale}


@router.get("/api/sessions/{session_id}/summary")
def session_summary(session_id: int):
    """Live per-category reconciliation (POS tender vs the right external source),
    computed fresh so it reflects current uploads/fee-VAT edits."""
    with get_db() as conn:
        session = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)
        ).fetchone()
        if not session:
            raise HTTPException(404)
        bid, pm = session["branch_id"], session["period_month"]

        rows, unmapped = category_reconciliation(conn, bid, pm)
        # per-lane manual-match tally (so the cards can flag manually-cleared lanes)
        manual: dict = {}
        for mr in conn.execute(
            "SELECT category, bank, COUNT(*) n, COALESCE(SUM(pos_amount),0) amt "
            "FROM recon_manual_match WHERE branch_id=? AND period_month=? GROUP BY category, bank",
            (bid, pm)):
            manual[f"{mr['category']}|{mr['bank']}"] = {"count": mr["n"], "amount": round(mr["amt"], 2)}
        categories = []
        for r in rows:
            c = {k: r[k] for k in
                 ("category", "bank", "pos_amount", "ext_amount", "ext_net", "variance", "status",
                  "adj_pos", "adj_ext", "adj_note", "adj_type", "adj_id")}
            c["manual"] = manual.get(f"{r['category']}|{r['bank']}")
            categories.append(c)
        payouts = payout_reconciliation(conn, bid, pm)

        pending = conn.execute(
            "SELECT COUNT(*) FROM recon_discrepancies WHERE session_id=? AND resolution='pending'",
            (session_id,)
        ).fetchone()[0]

        pos_total = sum(r["pos_amount"] for r in rows)
        ext_total = sum(r["ext_amount"] for r in rows)
        return {
            "categories":            categories,
            "payouts":               payouts,
            "unmapped":              [{"code": k, "amount": round(v, 2)} for k, v in sorted(unmapped.items())],
            "pos_total":             round(pos_total, 2),
            "ext_total":             round(ext_total, 2),
            "variance":              round(ext_total - pos_total, 2),
            "matched":               sum(1 for r in rows if r["status"] == "matched"),
            "pending_discrepancies": pending,
            "freshness":             _match_freshness(conn, session_id, bid, pm),
        }


@router.get("/api/sessions/{session_id}/categories")
def session_categories(session_id: int):
    """Persisted per-category results from the last run-match."""
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM recon_category_result WHERE session_id=? ORDER BY category, bank",
            (session_id,))]


@router.get("/api/sessions/{session_id}/daily")
def session_daily(session_id: int):
    """Day-by-day reconciliation board with per-day pending tracking for late
    documents. Computed live so it reflects current uploads + fee/VAT edits."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        return daily_reconciliation(conn, s["branch_id"], s["period_month"])


@router.get("/api/sessions/{session_id}/bank-coverage")
def session_bank_coverage(session_id: int):
    """Bank-first coverage: every statement credit line classified matched /
    unmatched / excluded / informational, plus the CREDIT 3-document chain (POS gross
    ↔ settlement gross, settlement net ↔ bank credit, fee+VAT auto-pulled). Computed
    live so it reflects current uploads, manual matches and fee/VAT edits."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        return bank_statement_coverage(conn, s["branch_id"], s["period_month"])


@router.get("/api/sessions/{session_id}/lane")
def session_lane(session_id: int, category: str, bank: str, day: str | None = None):
    """Drill-down: the external lines that make up one (category, bank) lane,
    optionally for a single day. POS side stays an aggregate (no per-receipt
    tender split exists)."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        lines = lane_lines(conn, s["branch_id"], s["period_month"], category, bank, day)
        return {"category": category, "bank": bank, "day": day,
                "count": len(lines), "total": round(sum(l["amount"] for l in lines), 2),
                "lines": lines}


def _build_export(conn, session) -> bytes:
    """One workbook per session — the working paper an accountant signs off and files:
    Summary / Tie-out & Sign-off / Categories / Payout chain / Daily / Discrepancies /
    acknowledged variances / Cash 3-way / Coupon points, then the auditor's line-level
    sheets: Line detail / Unmatched / Manual matches / Source documents (with SHA-256)."""
    import io
    import openpyxl
    from openpyxl.styles import Font

    bid, pm = session["branch_id"], session["period_month"]
    rows, unmapped = category_reconciliation(conn, bid, pm)
    payouts = payout_reconciliation(conn, bid, pm)
    daily = daily_reconciliation(conn, bid, pm)
    discs = conn.execute(
        "SELECT * FROM recon_discrepancies WHERE session_id=? ORDER BY resolution, ABS(COALESCE(variance,0)) DESC",
        (session["session_id"],)).fetchall()

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    def sheet(title, headers, data_rows, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        if first:
            ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for r in data_rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[chr(64 + i)].width = max(12, len(str(h)) + 2)
        return ws

    pos_total = round(sum(r["pos_amount"] for r in rows), 2)
    ext_total = round(sum(r["ext_amount"] for r in rows), 2)
    pending = sum(1 for d in discs if d["resolution"] == "pending")
    # AUD-7: total variance "absorbed" by tolerance — the sum of |variance| over
    # lanes that passed as matched. Auditors want this quantified, not invisible.
    tol_absorbed = round(sum(abs(r["variance"]) for r in rows if r["status"] == "matched"), 2)
    accepts = conn.execute(
        "SELECT category, bank, note, created_by, created_at FROM recon_lane_adj "
        "WHERE branch_id=? AND period_month=? AND adj_type='accept' ORDER BY category, bank",
        (bid, pm)).fetchall()
    sheet("Summary",
          ["Field", "Value"],
          [["Branch", session["branch_name"]], ["Period", pm], ["Status", session["status"]],
           ["POS tender (reconcilable)", pos_total], ["External total", ext_total],
           ["Variance", round(ext_total - pos_total, 2)],
           ["Matched lanes", sum(1 for r in rows if r["status"] == "matched")],
           ["Variance absorbed within tolerance", tol_absorbed],   # AUD-7
           ["Open discrepancies", pending],
           ["Acknowledged variances (รับทราบ)", len(accepts)],     # AUD-8
           ["Unmapped POS codes", ", ".join(u["code"] for u in
                                            [{"code": k} for k in sorted(unmapped)]) or "—"]],
          first=True)

    # Tie-out & sign-off: channel-level POS → external → variance (the live
    # reconcile_month, incl. acknowledged variances and the residual after per-day
    # acks) + preparer/approver block — the front page an auditor reads first.
    v2 = reconcile_month(conn, bid, pm)
    _ST = {"ok": "ตรงกัน / Tied", "diff": "มีผลต่าง / Difference", "waiting_doc": "รอเอกสาร / Awaiting document",
           "setup": "ต้องตั้งค่า / Setup needed"}
    ws_so = sheet("Tie-out & Sign-off",
                  ["ช่องทาง / Channel", "POS", "เงินเข้า / External", "ผลต่าง / Variance",
                   "ผลต่างคงเหลือ / Residual", "สถานะ / Status", "รับทราบโดย / Accepted by",
                   "เหตุผล / Reason"],
                  [[c["label"], c["pos"], c["ext"], c["variance"], c.get("residual", c["variance"]),
                    _ST.get(c["status"], c["status"]),
                    (c.get("accepted") or {}).get("by") or "", (c.get("accepted") or {}).get("note") or ""]
                   for c in v2["channels"]])
    ws_so.append(["รวม / Total", round(sum(c["pos"] for c in v2["channels"]), 2),
                  round(sum(c["ext"] for c in v2["channels"]), 2),
                  round(sum(c["variance"] for c in v2["channels"]), 2), "", "", "", ""])
    nr = v2["non_revenue"]
    ws_so.append(["ไม่ใช่รายได้ (โอนภายใน/ตัดออก) / Non-revenue excluded credits",
                  "", nr["amount"], "", "", f"{nr['count']} รายการ", "", ""])
    ws_so.append(["ความคืบหน้า / Reconciled progress", "", "", "", "", f"{v2['progress_pct']}%", "", ""])
    ws_so.append([])
    ws_so.append(["สถานะงวด / Session status", session["status"]])
    ws_so.append(["จัดทำโดย / Prepared by", session["created_by"] or "—",
                  "วันที่ / Date", session["created_at"] or ""])
    ws_so.append(["อนุมัติโดย / Approved by", session["approved_by"] or "—",
                  "วันที่ / Date", session["approved_at"] or ""])
    ws_so.append(["พิมพ์เมื่อ / Exported at", _dt.now().strftime("%Y-%m-%d %H:%M")])
    ws_so.append([])
    for th, en in (("ผู้จัดทำ", "Prepared by"), ("ผู้สอบทาน", "Reviewed by"),
                   ("ผู้อนุมัติ", "Approved by")):
        ws_so.append([f"ลงชื่อ{th} / Signature ({en})",
                      "________________________", "วันที่ / Date", "____________"])
    ws_so.column_dimensions["A"].width = 44

    sheet("Categories",
          ["Category", "Bank", "POS", "Statement", "Net (after fee/VAT)", "Variance", "Status"],
          [[r["category"], r["bank"], r["pos_amount"], r["ext_amount"], r["ext_net"],
            r["variance"], r["status"]] for r in rows])

    sheet("Payout chain",
          ["Category", "Report net", "Bank payout", "Variance", "Status"],
          [[p["category"], p["report_net"], p["bank_payout"], p["variance"], p["status"]] for p in payouts])

    daily_rows = []
    for d in daily.get("days", []):
        for l in d["lanes"]:
            daily_rows.append([d["day"], l["category"], l["bank"], l["pos"], l["ext"],
                               l["variance"], l["status"]])
    sheet("Daily",
          ["Day", "Category", "Bank", "POS", "External", "Variance", "Status"], daily_rows)

    sheet("Discrepancies",
          ["Type", "Category", "Bank", "External", "POS", "Variance", "Resolution", "Assignee", "Note"],
          [[d["disc_type"], d["category"], d["bank"], d["amount_external"], d["amount_pos"],
            d["variance"], d["resolution"], d["assigned_to"] or "", d["resolution_note"] or ""]
           for d in discs])

    # AUD-8: every acknowledged variance (รับทราบ) with its reason + who/when, so
    # the owner can sign off the list and an auditor can sample it.
    sheet("Acknowledged variances",
          ["Category", "Bank", "Reason (เหตุผล)", "Accepted by", "At"],
          [[a["category"], a["bank"], a["note"] or "", a["created_by"] or "", a["created_at"] or ""]
           for a in accepts])

    # Per-day acknowledged variances (reconcile BY DATE): each differing day staff
    # cleared with a reason — the granular trail behind a channel's residual.
    day_acks = conn.execute(
        "SELECT category, day, variance, note, created_by, created_at FROM recon_day_ack "
        "WHERE branch_id=? AND period_month=? ORDER BY category, day", (bid, pm)).fetchall()
    sheet("Daily acknowledged variances",
          ["Category", "Day", "Variance", "Reason (เหตุผล)", "Accepted by", "At"],
          [[a["category"], a["day"], a["variance"], a["note"] or "",
            a["created_by"] or "", a["created_at"] or ""] for a in day_acks])

    # Cash 3-way (POS ↔ typed deposit slip ↔ bank deposit) per the owner's table
    tw = cash_three_way(conn, bid, pm)
    sheet("Cash 3-way",
          ["Day", "POS cash", "Slip amount", "Slip day", "Bank deposit", "Deposit date",
           "Slip status", "Bank status", "Slip var", "Bank var"],
          [[r["day"], r["pos_cash"], r["slip_amount"], r["slip_day"] or "", r["bank_deposit"],
            r["deposit_date"] or "", r["slip_leg"], r["bank_leg"],
            r["slip_variance"], r["bank_variance"]] for r in tw["rows"]])

    # Coupon issuer points (KTC/BAY) — issuer_points method
    ip = coupon_reconciliation(conn, bid, pm)
    sheet("Coupon points",
          ["Issuer", "POS points", "Summary gross", "Bank payout", "Variance", "Status"],
          [[r["issuer"], r["pos_point"], r["settle_gross"],
            ("in KTC deposit" if r["bank_in_credit"] else r["bank_payout"]),
            r["variance"], r["status"]] for r in ip])

    # ── Working-paper detail sheets (auditor handoff) ─────────────────────────
    # Line detail: every external document line behind the month with its match
    # status — the population an auditor samples from. Statuses are the snapshot
    # persisted by run-match (ส่งตรวจ/ปิดเดือน); 'manual' and 'excluded' are staff
    # decisions, kept across re-runs.
    def _f(v):
        return round(float(v or 0), 2)

    detail = []   # unified columns across the 4 line tables
    for r in conn.execute("""
        SELECT l.txn_date d, l.description ref, l.channel ch, l.debit, l.credit,
               l.match_status ms, l.notes, u.provider, u.file_name
        FROM recon_bank_lines l JOIN recon_uploads u ON u.upload_id=l.upload_id
        WHERE l.branch_id=? AND strftime('%Y-%m', l.txn_date)=?
        ORDER BY l.txn_date, l.id""", (bid, pm)):
        detail.append(["bank", r["provider"] or "", r["d"], r["ref"] or "", r["ch"] or "",
                       _f(r["credit"]), _f(r["debit"]), "", "", "",
                       r["ms"], r["notes"] or "", r["file_name"]])
    for r in conn.execute("""
        SELECT l.settlement_date d, l.batch_no ref, l.acquirer, l.card_type,
               l.gross_amount, l.refund_amount,
               COALESCE(l.fee_manual, l.mdr_fee) fee, COALESCE(l.vat_manual, l.vat_amount) vat,
               l.match_status ms, l.notes, u.provider, u.file_name
        FROM recon_cc_lines l JOIN recon_uploads u ON u.upload_id=l.upload_id
        WHERE l.branch_id=? AND strftime('%Y-%m', l.settlement_date)=?
        ORDER BY l.settlement_date, l.id""", (bid, pm)):
        gross, refund, fee, vat = _f(r["gross_amount"]), _f(r["refund_amount"]), _f(r["fee"]), _f(r["vat"])
        detail.append(["credit_card", r["provider"] or "", r["d"], r["ref"] or "",
                       r["acquirer"] or r["card_type"] or "", gross, refund, fee, vat,
                       round(gross - refund - fee - vat, 2),
                       r["ms"], r["notes"] or "", r["file_name"]])
    for src, table, vat_col in (("grab", "recon_grab_lines", "vat_amount"),
                                ("lineman", "recon_lineman_lines", "vat_on_commission")):
        for r in conn.execute(f"""
            SELECT l.order_date d, l.order_id ref, l.order_status,
                   l.food_amount, l.promo_deduct,
                   COALESCE(l.fee_manual, l.commission_amt) fee,
                   COALESCE(l.vat_manual, l.{vat_col}) vat,
                   l.match_status ms, l.notes, u.file_name
            FROM {table} l JOIN recon_uploads u ON u.upload_id=l.upload_id
            WHERE l.branch_id=? AND strftime('%Y-%m', l.order_date)=?
            ORDER BY l.order_date, l.id""", (bid, pm)):
            food, promo, fee, vat = _f(r["food_amount"]), _f(r["promo_deduct"]), _f(r["fee"]), _f(r["vat"])
            detail.append([src, "Grab" if src == "grab" else "LINE MAN", r["d"], r["ref"] or "",
                           r["order_status"] or "", food, promo, fee, vat,
                           round(food - promo - fee - vat, 2),
                           r["ms"], r["notes"] or "", r["file_name"]])

    detail_headers = ["แหล่ง / Source", "ผู้ให้บริการ / Provider", "วันที่ / Date",
                      "อ้างอิง / Ref", "ช่องทาง / Channel", "เงินเข้า / Gross in",
                      "เงินออก-คืน / Out-Refund", "ค่าธรรมเนียม / Fee", "VAT",
                      "สุทธิ / Net", "สถานะจับคู่ / Match status",
                      "หมายเหตุ / Note", "ไฟล์ / File"]
    sheet("Line detail", detail_headers, detail)

    # Unmatched & excluded: the exception listing auditors work through.
    sheet("Unmatched", detail_headers,
          [row for row in detail if row[10] in ("unmatched", "excluded")])

    # Manual matches: every staff tick-to-match decision, one row per statement
    # line, with the POS days it covers and who matched it.
    ref_cols = {"recon_bank_lines":    ("txn_date", "description"),
                "recon_cc_lines":      ("settlement_date", "batch_no"),
                "recon_grab_lines":    ("order_date", "order_id"),
                "recon_lineman_lines": ("order_date", "order_id")}
    manual_rows = []
    for m in conn.execute(
            "SELECT * FROM recon_manual_match WHERE branch_id=? AND period_month=? "
            "ORDER BY matched_at", (bid, pm)):
        pos_days = "; ".join(f"{x.get('day')}: {float(x.get('amount') or 0):,.2f}"
                             for x in json.loads(m["pos_days"] or "[]"))
        lines = conn.execute(
            "SELECT ext_table, ext_id, amount FROM recon_manual_match_line WHERE match_id=?",
            (m["match_id"],)).fetchall() or [None]
        for ln in lines:
            d = ref = ""
            if ln is not None and ln["ext_table"] in ref_cols:
                dcol, rcol = ref_cols[ln["ext_table"]]
                hit = conn.execute(f"SELECT {dcol} d, {rcol} ref FROM {ln['ext_table']} WHERE id=?",
                                   (ln["ext_id"],)).fetchone()
                if hit:
                    d, ref = hit["d"], hit["ref"] or ""
            manual_rows.append([m["match_id"], m["category"], m["bank"], pos_days,
                                m["pos_amount"], d, ref,
                                (ln["amount"] if ln is not None else m["ext_amount"]),
                                m["variance"], m["note"] or "",
                                m["matched_by"] or "", m["matched_at"]])
    sheet("Manual matches",
          ["Match", "ช่องทาง / Category", "ธนาคาร / Bank", "วัน POS / POS days",
           "ยอด POS / POS amount", "วันที่เอกสาร / Doc date", "อ้างอิง / Ref",
           "ยอดเอกสาร / Doc amount", "ผลต่าง / Variance", "หมายเหตุ / Note",
           "จับคู่โดย / By", "เมื่อ / At"],
          manual_rows)

    # Source documents: which uploads the numbers come from. The SHA-256 hash
    # proves a file handed to the auditor is bit-identical to what was ingested.
    ws_docs = sheet("Source documents",
                    ["ID", "ประเภท / Type", "ผู้ให้บริการ / Provider", "ไฟล์ / File",
                     "จำนวนแถว / Rows", "สถานะ / Status", "อัปโหลดเมื่อ / Uploaded at",
                     "SHA-256"],
                    [[u["upload_id"], u["source_type"], u["provider"] or "", u["file_name"],
                      u["row_count"], u["status"], u["uploaded_at"], u["file_hash"] or ""]
                     for u in conn.execute(
                        "SELECT * FROM recon_uploads WHERE branch_id=? AND period_month=? "
                        "ORDER BY uploaded_at", (bid, pm))])
    atts = conn.execute(
        "SELECT * FROM recon_attachment WHERE branch_id=? AND strftime('%Y-%m', day)=? "
        "ORDER BY day", (bid, pm)).fetchall()
    if atts:
        ws_docs.append([])
        ws_docs.append(["หลักฐานแนบ / Evidence attachments"])
        ws_docs.append(["วันที่ / Day", "ช่องทาง / Category", "ไฟล์ / File",
                        "หมายเหตุ / Note", "โดย / By", "เมื่อ / At"])
        for a in atts:
            ws_docs.append([a["day"], a["category"] or "", a["file_name"],
                            a["note"] or "", a["uploaded_by"] or "", a["uploaded_at"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/api/sessions/{session_id}/export.xlsx")
def export_session(session_id: int):
    with get_db() as conn:
        session = conn.execute("""
            SELECT s.*, b.name AS branch_name
            FROM recon_sessions s JOIN branches b ON b.branch_id = s.branch_id
            WHERE s.session_id=?""", (session_id,)).fetchone()
        if not session:
            raise HTTPException(404)
        data = _build_export(conn, session)
    from urllib.parse import quote
    fname = f"recon_{session['branch_name']}_{session['period_month']}.xlsx".replace(" ", "_")
    # Branch names are Thai → not latin-1 safe for a raw header. RFC 5987: an ASCII
    # fallback plus a UTF-8 filename* the browser prefers.
    ascii_fallback = f"recon_{session['period_month']}.xlsx"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(fname)}"})


def _write_close_snapshot(conn, session, status: str, by: str) -> dict:
    """Freeze the working-paper export at sign-off so 'what management signed'
    survives later edits: the .xlsx bytes are written to the volume and a hashed
    row recorded. `session` must carry branch_name (join). Returns the new row."""
    data = _build_export(conn, session)
    ts = _dt.now().strftime("%Y%m%d-%H%M%S")
    pm = session["period_month"]
    rel = f"{pm}/{session['branch_id']}_{pm}_{status}_{ts}.xlsx"
    dest = _SNAPSHOT_DIR / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    digest = hashlib.sha256(data).hexdigest()
    # Source-document manifest: name + sha256 of every confirmed upload that fed this
    # period, frozen alongside the working paper so the inputs are provable later.
    manifest = json.dumps([
        {"file": u["file_name"], "sha256": u["file_hash"],
         "type": u["source_type"], "provider": u["provider"]}
        for u in conn.execute(
            "SELECT file_name, file_hash, source_type, provider FROM recon_uploads "
            "WHERE branch_id=? AND period_month=? AND status='confirmed' "
            "ORDER BY source_type, file_name", (session["branch_id"], pm))],
        ensure_ascii=False)
    cur = conn.execute(
        "INSERT INTO recon_close_snapshot"
        "(session_id, branch_id, period_month, status, file_name, file_path, sha256, size_bytes, created_by, source_manifest) "
        "VALUES (?,?,?,?,?,?,?,?,?,?)",
        (session["session_id"], session["branch_id"], pm, status,
         Path(rel).name, rel, digest, len(data), by or "staff", manifest))
    _audit(conn, "close_snapshot", cur.lastrowid, "create", by or "staff",
           [{"field": "sha256", "old": None, "new": digest}])
    return {"id": cur.lastrowid, "sha256": digest, "size_bytes": len(data)}


@router.get("/api/sessions/{session_id}/snapshots")
def list_close_snapshots(session_id: int):
    with get_db() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT id, status, file_name, sha256, size_bytes, created_by, created_at "
            "FROM recon_close_snapshot WHERE session_id=? ORDER BY created_at DESC",
            (session_id,))]


@router.get("/api/snapshots/{snap_id}/file")
def get_close_snapshot_file(snap_id: int):
    with get_db() as conn:
        r = conn.execute("SELECT * FROM recon_close_snapshot WHERE id=?", (snap_id,)).fetchone()
    if not r:
        raise HTTPException(404, "snapshot not found")
    p = _SNAPSHOT_DIR / r["file_path"]
    if not p.exists():
        raise HTTPException(404, "snapshot file missing on disk")
    return Response(
        content=p.read_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{r['file_name']}\""})


@router.get("/api/sessions/{session_id}/discrepancies")
def list_discrepancies(
    session_id: int,
    resolution: str | None = "pending",
    category:   str | None = None,
    bank:       str | None = None,
):
    with get_db() as conn:
        filters: list[str] = ["session_id=?"]
        params:  list      = [session_id]
        if resolution and resolution != "all":
            filters.append("resolution=?"); params.append(resolution)
        if category:
            filters.append("category=?"); params.append(category)
        if bank:
            filters.append("bank=?"); params.append(bank)
        rows = conn.execute(f"""
            SELECT * FROM recon_discrepancies
            WHERE {' AND '.join(filters)}
            ORDER BY ABS(COALESCE(variance, amount_external)) DESC
        """, params).fetchall()
        return [dict(r) for r in rows]


class ResolveBody(BaseModel):
    resolution:      Literal["accepted", "adjusted", "rejected"]
    resolution_note: str = ""
    resolved_by:     str = ""


def _assert_disc_period_open(conn, disc_id: int):
    row = conn.execute("""
        SELECT s.branch_id, s.period_month
        FROM recon_discrepancies d JOIN recon_sessions s ON s.session_id = d.session_id
        WHERE d.id=?""", (disc_id,)).fetchone()
    if row:
        _assert_period_open(conn, row["branch_id"], row["period_month"])


@router.get("/api/discrepancies/{disc_id}/days")
def disc_daily_breakdown(disc_id: int):
    """Per-day POS + external breakdown for one discrepancy (same lane, filtered daily)."""
    with get_db() as conn:
        disc = conn.execute(
            """SELECT d.category, d.bank, s.branch_id, s.period_month
               FROM recon_discrepancies d
               JOIN recon_sessions s ON s.session_id = d.session_id
               WHERE d.id=?""", (disc_id,)).fetchone()
        if not disc:
            raise HTTPException(404)
        cat, bank, bid, pm = disc["category"], disc["bank"], disc["branch_id"], disc["period_month"]
        daily = daily_reconciliation(conn, bid, pm)
        days_out = []
        for day in daily.get("days", []):
            for lane in day.get("lanes", []):
                if lane["category"] == cat and lane["bank"] == (bank or "NONE"):
                    days_out.append({
                        "day":      day["day"],
                        "pos":      lane["pos"],
                        "ext":      lane["ext"],
                        "ext_net":  lane["ext_net"],
                        "variance": round((lane["ext"] or 0) - (lane["pos"] or 0), 2),
                        "status":   lane["status"],
                    })
        return {"category": cat, "bank": bank, "period_month": pm, "days": days_out}


@router.patch("/api/discrepancies/{disc_id}/resolve", dependencies=[Depends(require_write_token)])
def resolve_discrepancy(disc_id: int, body: ResolveBody, request: Request):
    actor = _actor(request, body.resolved_by)          # the logged-in user wins over the body
    with get_db() as conn:
        old = conn.execute(
            "SELECT resolution, resolution_note FROM recon_discrepancies WHERE id=?",
            (disc_id,)).fetchone()
        if not old:
            raise HTTPException(404, f"discrepancy {disc_id} not found")
        _assert_disc_period_open(conn, disc_id)
        conn.execute("""
            UPDATE recon_discrepancies
            SET resolution=?, resolution_note=?, resolved_by=?, resolved_at=datetime('now')
            WHERE id=?
        """, (body.resolution, body.resolution_note, actor, disc_id))
        diffs = []
        if old["resolution"] != body.resolution:
            diffs.append({"field": "resolution", "old": old["resolution"], "new": body.resolution})
        if (old["resolution_note"] or "") != (body.resolution_note or ""):
            diffs.append({"field": "resolution_note",
                           "old": old["resolution_note"], "new": body.resolution_note})
        _audit(conn, "discrepancy", disc_id, "update", actor, diffs or None)
        conn.commit()
    return {"id": disc_id, "resolution": body.resolution}


@router.patch("/api/discrepancies/{disc_id}/assign", dependencies=[Depends(require_write_token)])
def assign_discrepancy(disc_id: int, request: Request, assigned_to: str = Form(...)):
    with get_db() as conn:
        old = conn.execute(
            "SELECT assigned_to FROM recon_discrepancies WHERE id=?", (disc_id,)).fetchone()
        if not old:
            raise HTTPException(404, f"discrepancy {disc_id} not found")
        _assert_disc_period_open(conn, disc_id)
        conn.execute("UPDATE recon_discrepancies SET assigned_to=? WHERE id=?", (assigned_to, disc_id))
        _audit(conn, "discrepancy", disc_id, "update", _actor(request),
               [{"field": "assigned_to", "old": old["assigned_to"], "new": assigned_to}])
        conn.commit()
    return {"id": disc_id, "assigned_to": assigned_to}


class BulkResolveBody(BaseModel):
    ids:             list[int]
    resolution:      Literal["accepted", "adjusted", "rejected"]
    resolution_note: str = ""
    resolved_by:     str = ""


@router.post("/api/sessions/{session_id}/discrepancies/bulk-resolve",
          dependencies=[Depends(require_write_token)])
def bulk_resolve_discrepancies(session_id: int, body: BulkResolveBody, request: Request):
    if not body.ids:
        raise HTTPException(400, "ids list must not be empty")
    actor = _actor(request, body.resolved_by)
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        _assert_period_open(conn, s["branch_id"], s["period_month"])
        ph = ",".join("?" * len(body.ids))
        # rowcount = what THIS request changed (ids of another session, unknown or
        # already-resolved ids are skipped by the WHERE and must not be counted).
        updated = conn.execute(f"""
            UPDATE recon_discrepancies
            SET resolution=?, resolution_note=?, resolved_by=?, resolved_at=datetime('now')
            WHERE session_id=? AND id IN ({ph}) AND resolution='pending'
        """, [body.resolution, body.resolution_note, actor, session_id]
             + list(body.ids)).rowcount
        _audit(conn, "session", session_id, "update", actor,
               [{"field": "bulk_resolve",
                 "old": "pending",
                 "new": f"{body.resolution} × {updated} items"}])
        conn.commit()
    return {"updated": updated, "resolution": body.resolution}


_FEE_VAT_TABLE = {
    "grab":        "recon_grab_lines",
    "lineman":     "recon_lineman_lines",
    "credit_card": "recon_cc_lines",
}


class FeeVat(BaseModel):
    fee_manual: float | None = None
    vat_manual: float | None = None
    by:         str = ""


@router.get("/api/lines/{source}")
def list_lines(source: str, branch_id: int, period_month: str):
    """Lines for the fee/VAT editor, with effective fee/VAT and computed net."""
    if source not in _FEE_VAT_TABLE:
        raise HTTPException(400, f"source must be one of {sorted(_FEE_VAT_TABLE)}")
    table = _FEE_VAT_TABLE[source]
    spec = {
        "grab":        ("order_date",      "order_id",  "food_amount",  "commission_amt", "vat_amount",       "promo_deduct"),
        "lineman":     ("order_date",      "order_id",  "food_amount",  "commission_amt", "vat_on_commission","promo_deduct"),
        "credit_card": ("settlement_date", "batch_no",  "gross_amount", "mdr_fee",        "vat_amount",       None),
    }[source]
    date_col, ref_col, gross_col, fee_col, vat_col, promo_col = spec
    out = []
    with get_db() as conn:
        extra = " AND order_status='Completed'" if source in ("grab", "lineman") else ""
        rows = conn.execute(f"""
            SELECT * FROM {table}
            WHERE branch_id=? AND strftime('%Y-%m', {date_col})=?{extra}
            ORDER BY {date_col}
        """, (branch_id, period_month)).fetchall()
        for r in rows:
            gross  = float(r[gross_col] or 0)
            fee    = r["fee_manual"] if r["fee_manual"] is not None else float(r[fee_col] or 0)
            if r["vat_manual"] is not None:
                vat = r["vat_manual"]
            else:
                vat = float(r[vat_col] or 0) if vat_col else 0.0
            promo  = float(r[promo_col] or 0) if promo_col else 0.0
            refund = float(r["refund_amount"] or 0) if source == "credit_card" else 0.0
            out.append({
                "id":         r["id"],
                "date":       r[date_col],
                "ref":        r[ref_col] or "",
                "gross":      round(gross, 2),
                "fee":        r["fee_manual"],
                "vat":        r["vat_manual"],
                "fee_parsed": round(float(r[fee_col] or 0), 2),
                "vat_parsed": round(float(r[vat_col] or 0) if vat_col else 0.0, 2),
                "other":      round(refund + promo, 2),   # non-fee deductions → client live-recompute
                "net":        round(gross - refund - fee - vat - promo, 2),
            })
    return out


@router.patch("/api/lines/{source}/{line_id}/fee-vat", dependencies=[Depends(require_write_token)])
def set_fee_vat(source: str, line_id: int, body: FeeVat, request: Request):
    table = _FEE_VAT_TABLE.get(source)
    if not table:
        raise HTTPException(400, f"source must be one of {sorted(_FEE_VAT_TABLE)}")
    date_col = "settlement_date" if source == "credit_card" else "order_date"
    actor = _actor(request, body.by)
    with get_db() as conn:
        line = conn.execute(
            f"SELECT branch_id, {date_col} AS d, fee_manual, vat_manual FROM {table} WHERE id=?",
            (line_id,)).fetchone()
        if not line:
            raise HTTPException(404, f"{source} line {line_id} not found")
        if line["d"]:
            _assert_period_open(conn, line["branch_id"], str(line["d"])[:7])
        conn.execute(f"""
            UPDATE {table}
            SET fee_manual=?, vat_manual=?, fee_vat_by=?, fee_vat_at=datetime('now')
            WHERE id=?
        """, (body.fee_manual, body.vat_manual, actor, line_id))
        diffs = []
        if line["fee_manual"] != body.fee_manual:
            diffs.append({"field": "fee_manual", "old": line["fee_manual"], "new": body.fee_manual})
        if line["vat_manual"] != body.vat_manual:
            diffs.append({"field": "vat_manual", "old": line["vat_manual"], "new": body.vat_manual})
        if diffs:
            _audit(conn, "line", line_id, "update", actor, diffs)
        conn.commit()
    return {"source": source, "line_id": line_id,
            "fee_manual": body.fee_manual, "vat_manual": body.vat_manual}


@router.get("/api/sessions/{session_id}/matches")
def list_matches(session_id: int, category: str | None = None):
    with get_db() as conn:
        filters = ["session_id=?", "status='matched'"]
        params: list = [session_id]
        if category:
            filters.append("category=?"); params.append(category)
        rows = conn.execute(f"""
            SELECT * FROM recon_category_result
            WHERE {' AND '.join(filters)}
            ORDER BY category, bank
        """, params).fetchall()
        return [dict(r) for r in rows]


class DayAck(BaseModel):
    branch_id:    int
    period_month: str
    category:     str
    day:          str
    variance:     float = 0.0
    note:         str | None = None
    created_by:   str | None = None


_ACK_CHANNELS = {k for k, _ in _V2_CHANNELS}      # the channel cards a day can be acked on


def _live_day_variance(conn, branch_id: int, period_month: str, category: str, day: str):
    """The day's variance exactly as the branch page shows it: reconcile_month sums
    daily_reconciliation's lanes of the channel for the day, ext − pos. None when
    the engine has no lane for that channel/day (or cannot run)."""
    try:
        daily = daily_reconciliation(conn, branch_id, period_month)
    except Exception:
        return None
    lanes = [l for d in daily.get("days", []) if d["day"] == day
             for l in d["lanes"] if l["category"] == category]
    if not lanes:
        return None
    return round(sum(l["ext"] for l in lanes) - sum(l["pos"] for l in lanes), 2)


@router.get("/api/day-ack")
def day_ack_list(branch_id: int, period_month: str):
    _validate_period(period_month)
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM recon_day_ack WHERE branch_id=? AND period_month=? "
            "ORDER BY category, day", (branch_id, period_month)).fetchall()
        return [dict(r) for r in rows]


@router.post("/api/day-ack", dependencies=[Depends(require_write_token)])
def day_ack_upsert(body: DayAck, request: Request):
    _validate_period(body.period_month)
    # Same audit-grade rule as a channel accept: a meaningful reason (≥10 chars).
    if len((body.note or "").strip()) < 10:
        raise HTTPException(400, "การรับทราบผลต่างรายวันต้องใส่เหตุผลที่ชัดเจน (อย่างน้อย 10 ตัวอักษร)")
    if not (_DAY_RE.match(body.day or "") and body.day.startswith(body.period_month)):
        raise HTTPException(400, "day must be a YYYY-MM-DD inside period_month")
    if body.category not in _ACK_CHANNELS:
        raise HTTPException(400, f"category must be one of {sorted(_ACK_CHANNELS)}")
    actor = _actor(request, body.created_by)
    with get_db() as conn:
        _validate_master_branch(conn, body.branch_id)
        _assert_period_open(conn, body.branch_id, body.period_month)
        # Store the variance the branch page shows for that day, not the client's
        # number; the body value is kept only when the engine has no lane that day.
        variance = _live_day_variance(conn, body.branch_id, body.period_month,
                                      body.category, body.day)
        source = "engine"
        if variance is None:
            variance, source = body.variance, "client"
        existing = conn.execute(
            "SELECT id FROM recon_day_ack WHERE branch_id=? AND period_month=? "
            "AND category=? AND day=?",
            (body.branch_id, body.period_month, body.category, body.day)).fetchone()
        if existing:
            conn.execute(
                "UPDATE recon_day_ack SET variance=?, note=?, created_by=?, "
                "created_at=datetime('now') WHERE id=?",
                (variance, body.note, actor, existing["id"]))
            row_id = existing["id"]
        else:
            cur = conn.execute(
                "INSERT INTO recon_day_ack "
                "(branch_id,period_month,category,day,variance,note,created_by) "
                "VALUES (?,?,?,?,?,?,?)",
                (body.branch_id, body.period_month, body.category, body.day,
                 variance, body.note, actor))
            row_id = cur.lastrowid
        diffs = [{"field": "variance", "old": None, "new": variance}]
        if source == "client":
            diffs.append({"field": "variance_source", "old": None, "new": "client"})
        _audit(conn, "day_ack", row_id, "accept", actor, diffs)
        conn.commit()
    return {"id": row_id, "variance": variance, "variance_source": source}


@router.delete("/api/day-ack/{ack_id}", dependencies=[Depends(require_write_token)])
def day_ack_delete(ack_id: int, request: Request):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM recon_day_ack WHERE id=?", (ack_id,)).fetchone()
        if not row:
            raise HTTPException(404)
        _assert_period_open(conn, row["branch_id"], row["period_month"])
        conn.execute("DELETE FROM recon_day_ack WHERE id=?", (ack_id,))
        # Un-acknowledging reopens the day's variance — audit it with what was removed.
        _audit(conn, "day_ack", ack_id, "delete", _actor(request), [
            {"field": "variance", "old": row["variance"], "new": None},
            {"field": "branch_id", "old": row["branch_id"], "new": None},
            {"field": "period_month", "old": row["period_month"], "new": None},
            {"field": "day", "old": f"{row['category']} {row['day']}", "new": None},
            {"field": "note", "old": row["note"], "new": None}])
        conn.commit()
    return {"deleted": ack_id}


@router.get("/api/sessions/{session_id}/trend")
def session_trend(session_id: int):
    """Per-day and per-week recon health for one session.
    Reuses the daily_reconciliation engine so numbers agree with the Daily tab."""
    with get_db() as conn:
        s = conn.execute(
            "SELECT branch_id, period_month FROM recon_sessions WHERE session_id=?",
            (session_id,)).fetchone()
        if not s:
            raise HTTPException(404)
        data = daily_reconciliation(conn, s["branch_id"], s["period_month"])

    daily_out: list[dict] = []
    for dr in data.get("days", []):
        lanes       = dr.get("lanes", [])
        matched     = dr.get("matched", 0)
        total_lanes = len(lanes)
        pos         = float(dr.get("pos_total") or 0)
        ext         = float(dr.get("ext_total") or 0)
        lane_pct    = round(matched / total_lanes * 100) if total_lanes else 0
        cov_pct     = min(100, round(ext / pos * 100)) if pos else 0
        daily_out.append({
            "day":         dr["day"],
            "pos":         round(pos, 2),
            "ext":         round(ext, 2),
            "matched":     matched,
            "total_lanes": total_lanes,
            "lane_pct":    lane_pct,
            "cov_pct":     cov_pct,
        })

    # Weekly rollup: W1=days 1-7, W2=8-14, W3=15-21, W4=22+
    week_data: dict = {}
    for d in daily_out:
        day_num = int(d["day"][8:10])
        week_n  = min((day_num - 1) // 7 + 1, 4)
        wk      = f"W{week_n}"
        if wk not in week_data:
            week_data[wk] = {"week": wk, "week_n": week_n,
                              "pos": 0.0, "ext": 0.0, "matched": 0, "lanes": 0}
        w = week_data[wk]
        w["pos"]     += d["pos"]
        w["ext"]     += d["ext"]
        w["matched"] += d["matched"]
        w["lanes"]   += d["total_lanes"]

    weekly_out = []
    for w in sorted(week_data.values(), key=lambda x: x["week_n"]):
        w["pos"]      = round(w["pos"], 2)
        w["ext"]      = round(w["ext"], 2)
        w["lane_pct"] = round(w["matched"] / w["lanes"] * 100) if w["lanes"] else 0
        w["cov_pct"]  = min(100, round(w["ext"] / w["pos"] * 100)) if w["pos"] else 0
        weekly_out.append(w)

    return {"session_id": session_id, "daily": daily_out, "weekly": weekly_out}
