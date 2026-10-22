# -*- coding: utf-8 -*-
"""PR-B4b - the endpoints batch 4b added answer, validate and refuse (the app's own client, auth off)."""
import io

from openpyxl import load_workbook


def test_the_endpoints_answer_and_refuse(client):
    r = client.get("/api/bom-center/close-settings").json()
    assert r["reason_min_baht"] == 500 and r["cycle_counts"] is False and r["reason_min_baht_max"] == 100000
    assert client.post("/api/bom-center/close-settings", json={"reason_min_baht": 100001}).status_code == 400
    assert client.post("/api/bom-center/close-settings", json={"reason_min_baht": -5}).status_code == 400
    ok = client.post("/api/bom-center/close-settings", json={"reason_min_baht": 800, "cycle_counts": False})
    assert ok.status_code == 200 and ok.json()["reason_min_baht"] == 800 and ok.json()["cycle_counts"] is False
    assert client.post("/api/bom-center/close-settings", json={"reason_min_baht": 500, "cycle_counts": True}).json()["cycle_counts"] is True
    plan = client.get("/api/bom-center/stock/plan?branch_id=1")
    assert plan.status_code == 200 and plan.json()["mode"] == "full" and plan.json()["first"] is True
    assert client.get("/api/bom-center/stock/plan").status_code == 400                    # the centre must name a branch
    assert client.get("/api/bom-center/stock/sheet?branch_id=1&mode=weekly").status_code == 400
    assert client.post("/api/bom-center/stock/partial-close?branch_id=1", json={"month": "2026-13"}).status_code == 400
    for m in ("2026-09", "2026-10"):                                                       # withdrawn 2026-10-13: no month is closed on a partial count
        x = client.post("/api/bom-center/stock/partial-close?branch_id=1", json={"month": m})
        assert x.status_code == 400 and "ปิดเดือนต้องนับของครบทุกรายการ" in x.json()["detail"]
    b = client.get("/api/bom-center/backlog?top=3").json()
    assert set(b) >= {"window", "items", "top", "counts", "total", "kind_th", "note"} and len(b["top"]) <= 3
    assert client.get("/api/bom-center/feedback").json() == {"tasks": []}
    assert client.post("/api/bom-center/feedback/ack", json={"rule": "miscount", "subject": "1", "month": "2026-09"}).status_code == 400
    assert client.post("/api/bom-center/feedback/ack", json={"rule": "nonsense", "subject": "1", "month": "2026-09"}).status_code == 400
    assert client.post("/api/bom-center/feedback/ack", json={"rule": "miscount", "subject": "1", "month": "9"}).status_code == 400
    o = client.get("/api/bom-center/hub/outcome?month=2026-09").json()
    assert o["month"] == "2026-09" and o["branches"] == [] and o["total"]["deliveries"] == 0
    assert client.get("/api/bom-center/hub/outcome?month=junk").status_code == 400
    x = client.get("/api/bom-center/hub/compliance.xlsx?date_from=2026-09-01&date_to=2026-09-30")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["ตรวจสอบการสั่ง", "ใบตัดเบิก", "ผลของการสั่งจริง"]
    assert list(next(wb["ผลของการสั่งจริง"].iter_rows(min_row=2, max_row=2, values_only=True))) == [
        "สาขา", "รอบส่ง", "วันที่ส่ง", "รอบส่งถัดไปถึง", "รายการที่วัด", "วันที่ของพอใช้ (มัธยฐาน)", "รายการที่หมดก่อนรอบถัดไป", "รายการที่หมด",
        "วันที่ประมาณจากยอดขาย", "หมายเหตุ"]
    # the settings the test changed are put back: the shared test database stays as it was
    import sqlite3
    import main
    c = sqlite3.connect(str(main.DB_PATH))
    c.execute("DELETE FROM app_settings WHERE key IN ('bc_reason_min_baht', 'bc_cycle_counts')")
    c.execute("DELETE FROM bc_audit WHERE action IN ('reason_min_baht', 'cycle_counts')")
    c.commit()
    c.close()
