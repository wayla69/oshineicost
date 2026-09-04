"""Golden for the Management Summary workbook — /api/reports/management.xlsx.

Every cell of every sheet (value, number format, bold, wrap, fill) plus the
column widths and the download filename, pinned per parameter set. This is
the net under which report_management_xlsx (312 lines, one function for
eleven sheets) is cut into one function per sheet (Phase 8b, 2026-09): the
JSON golden proves the report is deterministic on the fixture, and this
proves the workbook built from it is too.

  regenerate:  GOLDEN_UPDATE=1 python -m pytest backend/test_mgt_xlsx_golden.py
"""
import io
import os
import re
from pathlib import Path

import openpyxl
import pytest

import test_endpoints as te

GOLDEN = Path(__file__).parent / "golden" / "mgt_xlsx"
UPDATE = os.environ.get("GOLDEN_UPDATE") == "1"

CASES = [q for p, q in te.SMOKE if p == "/api/reports/management"] + [
    {"period": "weekly", "lang": "th"},
    {"period": "daily", "lang": "th"},
    {"period": "monthly", "lang": "th"},
]


def _slug(params):
    return (re.sub(r"[^A-Za-z0-9._-]+", "_", "_".join(f"{k}-{v}" for k, v in sorted(params.items())))
            or "default") + ".txt"


def _cell(c):
    v = round(c.value, 6) if isinstance(c.value, float) else c.value
    flags = ""
    if c.font is not None and c.font.bold:
        flags += "b"
    if c.alignment is not None and c.alignment.wrap_text:
        flags += "w"
    if c.alignment is not None and c.alignment.horizontal:
        flags += c.alignment.horizontal[0]
    fill = ""
    if c.fill is not None and c.fill.fill_type == "solid":
        fill = str(c.fill.fgColor.rgb)[-6:]
    nf = c.number_format if c.number_format != "General" else ""
    return f"{c.coordinate}\t{v!r}\t{nf}\t{flags}\t{fill}"


def _dump(r) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    out = [r.headers.get("content-disposition", "")]
    for ws in wb.worksheets:
        widths = {k: v.width for k, v in sorted(ws.column_dimensions.items()) if v.width}
        out.append(f"## {ws.title} {ws.max_row}x{ws.max_column} widths={widths}")
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out.append(_cell(c))
    return "\n".join(out) + "\n"


@pytest.mark.parametrize("params", CASES, ids=[_slug(q)[:-4] for q in CASES])
def test_management_xlsx_golden(client, params):
    r = client.get("/api/reports/management.xlsx", params=params)
    assert r.status_code == 200, r.text[:300]
    assert "spreadsheetml" in r.headers.get("content-type", "")
    got = _dump(r)
    f = GOLDEN / _slug(params)
    if UPDATE or not f.exists():
        GOLDEN.mkdir(parents=True, exist_ok=True)
        f.write_text(got, encoding="utf-8")
        if not UPDATE:
            pytest.skip(f"golden file created: {f.name}")
        return
    want = f.read_text(encoding="utf-8")
    if got != want:
        gl, wl = got.splitlines(), want.splitlines()
        for i, (a, b) in enumerate(zip(gl, wl)):
            if a != b:
                pytest.fail(f"{f.name} differs at line {i + 1}:\n  golden: {b[:160]}\n  now:    {a[:160]}")
        pytest.fail(f"{f.name}: length differs ({len(wl)} -> {len(gl)} lines)")


def test_every_case_has_a_golden_file():
    missing = [_slug(q) for q in CASES if not (GOLDEN / _slug(q)).exists()]
    assert not missing, f"run with GOLDEN_UPDATE=1 to capture: {missing}"
