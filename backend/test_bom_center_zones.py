"""PR-K · editable storage zones (โซนจัดเก็บ) and the ingredient sheet's two new columns.

The promises:
  * the zones live in app_settings (bc_zones): the five defaults until the centre changes them; a code is stable, a
    label may be renamed, the list order is the order a branch counts in;
  * a zone that still holds ingredients cannot be deleted (rename it or move the items first); the fallback zone
    "other" cannot be deleted at all; duplicate labels, empty labels and odd codes are refused;
  * the count sheet follows the configured order and a guess never lands in a zone that no longer exists;
  * the ingredient template / import / editor carry "โซนจัดเก็บ" (a label) and "หน่วยที่นับ"; an old file that still has
    the supplier column is accepted and the column is ignored; the write is the centre's (a branch is refused).
"""
import io
import sqlite3
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

import auth
from bom_center import stock as ST
from test_bom_center_s4b import con  # noqa: F401  (five weeks of sales at branch 1)

AS_OF = date(2026, 10, 5)
DEFAULT_LABELS = ["แช่แข็ง", "แช่เย็น", "ของแห้ง", "บาร์น้ำ", "อื่น ๆ"]


def _zl(con):
    return [(z, t) for z, t in ST.zones(con)]


# ── the setting ───────────────────────────────────────────────────────────────
def test_defaults_are_the_five_current_zones(con):  # noqa: F811
    assert [c for c, _ in ST.zones(con)] == ["frozen", "chilled", "dry", "bar", "other"]
    assert [t for _, t in ST.zones(con)] == DEFAULT_LABELS
    con.execute("CREATE TABLE IF NOT EXISTS app_settings(key TEXT PRIMARY KEY, value TEXT NOT NULL DEFAULT '')")
    con.execute("INSERT OR REPLACE INTO app_settings(key, value) VALUES('bc_zones', 'not json')")
    assert ST.zones(con) == ST.DEFAULT_ZONES                       # a broken value falls back to the defaults


def test_round_trip_rename_add_reorder_and_delete(con):  # noqa: F811
    new = ST.set_zones(con, [{"code": "dry", "label": "ของแห้ง / ชั้นวาง"}, {"code": "frozen", "label": "แช่แข็ง"},
                             {"label": "ตู้โชว์"}, {"code": "other", "label": "อื่น ๆ"}], "ss")
    assert [c for c, _ in new][:2] == ["dry", "frozen"] and new[2] == ("zone1", "ตู้โชว์")      # a new zone gets a code
    assert _zl(con) == list(new)                                     # what is read back is what was saved, in order
    again = ST.set_zones(con, [{"code": c, "label": t} for c, t in new if c != "zone1"], "ss")  # an empty zone may go
    assert [c for c, _ in again] == ["dry", "frozen", "other"]
    assert con.execute("SELECT COUNT(*) FROM bc_audit WHERE action='zones_set'").fetchone()[0] == 2
    assert dict(again)["dry"] == "ของแห้ง / ชั้นวาง"                    # a renamed label keeps its code


def test_delete_guard_and_validation(con):  # noqa: F811
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, storage_zone) VALUES('Z-1', 'ของในตู้แช่แข็ง', 'raw', 'กรัม', 'frozen')")
    con.commit()
    keep = [{"code": c, "label": t} for c, t in ST.zones(con)]
    without = [z for z in keep if z["code"] != "frozen"]
    with pytest.raises(ValueError) as e:
        ST.set_zones(con, without, "ss")
    assert "แช่แข็ง" in str(e.value) and "1 รายการ" in str(e.value) and "เปลี่ยนชื่อ" in str(e.value)
    assert [c for c, _ in ST.zones(con)] == ["frozen", "chilled", "dry", "bar", "other"]   # nothing changed
    con.execute("UPDATE bc_ingredient SET storage_zone='dry' WHERE code='Z-1'")            # move the item, then it may go
    ST.set_zones(con, without, "ss")
    assert "frozen" not in dict(ST.zones(con))
    with pytest.raises(ValueError) as e:                                                    # the fallback zone stays
        ST.set_zones(con, [z for z in without if z["code"] != "other"], "ss")
    assert "โซนสำรอง" in str(e.value)
    for bad, msg in (([], "อย่างน้อย 1 โซน"),
                     ([{"code": "a", "label": ""}], "ชื่อโซน"),
                     ([{"code": "a", "label": "x"}, {"code": "b", "label": "x"}], "ซ้ำ"),
                     ([{"code": "a", "label": "x"}, {"code": "a", "label": "y"}], "ซ้ำ"),
                     ([{"code": "ก", "label": "x"}], "รหัสโซน"),
                     ([{"code": "a", "label": "ย" * 41}], "40")):
        with pytest.raises(ValueError) as e:
            ST.set_zones(con, bad, "ss")
        assert msg in str(e.value), (bad, str(e.value))


def test_sheet_groups_by_the_configured_order_and_guesses_stay_inside_it(con):  # noqa: F811
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv, monitored, storage_zone)"
                " VALUES('S-011', 'เกลือป่น 1 กก.', 'raw', 'กรัม', 'ถุง 1 กก.', 1000, 1, NULL)")
    con.commit()
    base = [(z["zone"], [i["code"] for i in z["items"]]) for z in ST.sheet(con, 1, AS_OF)["zones"]]
    assert [z for z, _ in base] == [z for z, _ in base if z in ("frozen", "chilled", "dry", "bar", "other")]
    assert any("S-011" in codes for _, codes in base)
    # reversed order: the same zones, listed the other way round
    order = [{"code": c, "label": t} for c, t in reversed(ST.zones(con))]
    ST.set_zones(con, order, "ss")
    now = [z["zone"] for z in ST.sheet(con, 1, AS_OF)["zones"]]
    assert now == [z for z, _ in reversed(base)]
    # the guessed zone of every item is dropped from the configuration: it lands in "other" instead
    guessed = next(z for z, codes in base if "S-011" in codes)
    if guessed != "other":
        ST.set_zones(con, [z for z in order if z["code"] != guessed], "ss")
        sh = ST.sheet(con, 1, AS_OF)
        assert guessed not in [z["zone"] for z in sh["zones"]]
        assert any("S-011" in [i["code"] for i in z["items"]] for z in sh["zones"] if z["zone"] == "other")
    assert ST.guess_zone("S-011", "เกลือ", "S", "raw", valid=["frozen", "bar"]) == "bar"       # no "other": the last zone
    assert ST.guess_zone("S-011", "เกลือ", "S", "raw") == "dry"


def test_layout_accepts_only_configured_zones(con):  # noqa: F811
    ST.set_zones(con, [{"code": "dry", "label": "ของแห้ง"}, {"code": "display", "label": "ตู้โชว์"}, {"code": "other", "label": "อื่น ๆ"}], "ss")
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('Z-2', 'ของ', 'raw', 'กรัม')")
    assert ST.set_layout(con, [{"code": "Z-2", "zone": "display"}], "ss") == 1
    assert con.execute("SELECT storage_zone FROM bc_ingredient WHERE code='Z-2'").fetchone()[0] == "display"
    with pytest.raises(ValueError):
        ST.set_layout(con, [{"code": "Z-2", "zone": "frozen"}], "ss")                         # no longer a zone


def test_zone_code_for_reads_a_label_or_a_code(con):  # noqa: F811
    assert ST.zone_code_for(con, "แช่เย็น") == "chilled" and ST.zone_code_for(con, "CHILLED") == "chilled"
    assert ST.zone_code_for(con, "  ") == ""
    with pytest.raises(ValueError) as e:
        ST.zone_code_for(con, "ห้องใต้ดิน")
    assert "ห้องใต้ดิน" in str(e.value) and "แช่แข็ง" in str(e.value)


# ── the API ───────────────────────────────────────────────────────────────────
def _clean(db_path, codes=(), zones=True):
    c = sqlite3.connect(str(db_path))
    if zones:
        c.execute("DELETE FROM app_settings WHERE key='bc_zones'")
    for code in codes:
        c.execute("DELETE FROM bc_ingredient WHERE code=?", (code,))
    c.commit()
    c.close()


def test_zones_endpoint_round_trip_and_guard(client):
    import main
    client.get("/api/bom-center/zones")
    try:
        d = client.get("/api/bom-center/zones").json()
        assert [z["code"] for z in d["zones"]] == ["frozen", "chilled", "dry", "bar", "other"]
        assert d["zones"][-1]["fallback"] is True and [z["label"] for z in d["defaults"]] == DEFAULT_LABELS
        rows = [{"code": z["code"], "label": z["label"]} for z in d["zones"]]
        rows[1]["label"] = "ตู้เย็น"
        rows.insert(0, {"label": "ตู้โชว์"})
        r = client.post("/api/bom-center/zones", json=rows)
        assert r.status_code == 200, r.text
        got = r.json()["zones"]
        assert [z["label"] for z in got][:3] == ["ตู้โชว์", "แช่แข็ง", "ตู้เย็น"] and got[0]["code"] == "zone1"
        assert client.get("/api/bom-center/stock/me").json()["zones"][0] == {"zone": "zone1", "label": "ตู้โชว์"}
        # an ingredient in a zone blocks its deletion with the plain message
        c = sqlite3.connect(str(main.DB_PATH))
        c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, storage_zone, active) VALUES('ZT-1', 'ทดสอบโซน', 'raw', 'กรัม', 'zone1', 1)")
        c.commit()
        c.close()
        r = client.post("/api/bom-center/zones", json=[x for x in rows if x.get("code")])
        assert r.status_code == 400 and "ตู้โชว์" in r.json()["detail"] and "1 รายการ" in r.json()["detail"]
        assert client.get("/api/bom-center/zones").json()["zones"][0]["items"] == 1
        assert client.post("/api/bom-center/zones", json=[]).status_code == 400
    finally:
        _clean(main.DB_PATH, ["ZT-1"])


def test_zone_writes_are_the_centres_not_the_branchs():
    mods = lambda role: auth.modules_for_role(role)  # noqa: E731
    for role in ("salessupport", "analyst"):
        assert auth.path_allowed("POST", "/api/bom-center/zones", mods(role)), role
    for role in ("branch", "warehouse", "accounting", "innovation", "marketing"):
        assert not auth.path_allowed("POST", "/api/bom-center/zones", mods(role)), role
    assert not auth.path_allowed("GET", "/api/bom-center/zones", mods("branch"))             # the branch reads them with the sheet


# ── the ingredient sheet: two new columns, the supplier column gone ─────────────
def _xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, data, dry=False):
    return client.post(f"/api/bom-center/ingredients/import?dry_run={'true' if dry else 'false'}",
                       files={"file": ("m.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_template_has_the_two_columns_and_no_supplier(client):
    import main
    c = sqlite3.connect(str(main.DB_PATH))
    client.get("/api/bom-center/zones")
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv, storage_zone, count_unit, active)"
              " VALUES('ZT-2', 'ของทดสอบ', 'raw', 'กรัม', 'แพ็ค', 500, 'bar', 'แพ็ค', 1)")
    c.commit()
    try:
        r = client.get("/api/bom-center/ingredients/template.xlsx")
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        assert "ผู้ขาย" not in rows[1] and "ผู้จำหน่าย" not in rows[1]
        assert rows[1][-2:] == ("โซนจัดเก็บ", "หน่วยที่นับ") and "บริษัท ไดกิโตะ กรุ๊ป จำกัด" in rows[0][0]
        mine = next(x for x in rows[2:] if x[0] == "ZT-2")
        assert mine[-2:] == ("บาร์น้ำ", "แพ็ค")                                          # the label, not the code
    finally:
        c.execute("DELETE FROM bc_ingredient WHERE code='ZT-2'")
        c.commit()
        c.close()


def test_import_reads_zone_and_count_unit_and_ignores_the_old_supplier_column(client):
    import main
    client.get("/api/bom-center/zones")
    heads = ["รหัส", "ชื่อวัตถุดิบ", "ผู้ขาย", "หน่วยฐาน", "หน่วยรับ", "ขนาดบรรจุ", "โซนจัดเก็บ", "หน่วยที่นับ"]
    data = _xlsx(heads, [
        ["ZT-3", "ปลาทดสอบ", "01_DAIKI FOOD", "กรัม", "แพ็ค", 500, "แช่แข็ง", "แพ็ค"],
        ["ZT-4", "ผักทดสอบ", "บริษัท ก", "กรัม", "ถุง", 1000, "", ""],                       # blank cells: nothing is set
        ["ZT-5", "ซอสทดสอบ", "", "กรัม", "ขวด", 700, "ห้องใต้ดิน", "ขวด"],                    # an unknown zone: the row is skipped
        ["ZT-6", "น้ำมันทดสอบ", "", "กรัม", "ขวด", 900, "ของแห้ง", "เมตร"],                    # a count unit that cannot be converted
    ])
    try:
        d = _post(client, data, dry=True).json()
        assert d["dry_run"] is True and "zone" in d["columns"] and "count_unit" in d["columns"] and "supplier" not in d["columns"]
        c = sqlite3.connect(str(main.DB_PATH))
        assert c.execute("SELECT COUNT(*) FROM bc_ingredient WHERE code LIKE 'ZT-%'").fetchone()[0] == 0        # a preview writes nothing
        d = _post(client, data).json()
        assert d["inserted"] >= 3
        why = {s.get("code"): s["why"] for s in d["skipped"]}
        assert "ห้องใต้ดิน" in why["ZT-5"] and "เมตร" in why["ZT-6"]
        got = {r[0]: r[1:] for r in c.execute("SELECT code, storage_zone, count_unit, supplier FROM bc_ingredient WHERE code LIKE 'ZT-%'")}
        assert got["ZT-3"] == ("frozen", "แพ็ค", None)                                    # label -> code; supplier ignored
        assert got["ZT-4"] == (None, None, None)
        assert "ZT-5" not in got
        assert got["ZT-6"][1] is None                                                      # the unit it could not convert is not kept
        # a second import with blanks does not clear what is set
        _post(client, _xlsx(heads, [["ZT-3", "ปลาทดสอบ", "", "กรัม", "แพ็ค", 500, "", ""]]))
        assert c.execute("SELECT storage_zone, count_unit FROM bc_ingredient WHERE code='ZT-3'").fetchone() == ("frozen", "แพ็ค")
        c.close()
    finally:
        _clean(main.DB_PATH, ["ZT-3", "ZT-4", "ZT-5", "ZT-6"], zones=False)


def test_editor_saves_the_zone_and_count_unit_and_can_clear_them(client):
    import main
    c = sqlite3.connect(str(main.DB_PATH))
    client.get("/api/bom-center/zones")
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv, supplier, active)"
              " VALUES('ZT-7', 'ทดสอบ', 'raw', 'กรัม', 'แพ็ค', 500, 'ผู้จำหน่าย ก', 1)")
    c.commit()
    try:
        r = client.post("/api/bom-center/ingredients", json=[{"code": "ZT-7", "storage_zone": "chilled", "count_unit": "แพ็ค", "supplier": "ไม่มีผล"}])
        assert r.status_code == 200, r.text
        row = c.execute("SELECT storage_zone, count_unit, supplier FROM bc_ingredient WHERE code='ZT-7'").fetchone()
        assert row == ("chilled", "แพ็ค", "ผู้จำหน่าย ก")                                 # the old column is neither shown nor edited
        got = next(x for x in client.get("/api/bom-center/ingredients?q=ZT-7").json()["rows"] if x["code"] == "ZT-7")
        assert got["storage_zone"] == "chilled" and got["count_unit"] == "แพ็ค" and "supplier" not in got
        assert client.post("/api/bom-center/ingredients", json=[{"code": "ZT-7", "storage_zone": ""}]).status_code == 200
        assert c.execute("SELECT storage_zone FROM bc_ingredient WHERE code='ZT-7'").fetchone()[0] is None     # back to the guess
        assert client.post("/api/bom-center/ingredients", json=[{"code": "ZT-7", "storage_zone": "ห้องใต้ดิน"}]).status_code == 400
        r = client.post("/api/bom-center/ingredients", json=[{"code": "ZT-7", "count_unit": "เมตร"}])
        assert r.status_code == 400 and "เมตร" in r.json()["detail"]
        assert c.execute("SELECT count_unit FROM bc_ingredient WHERE code='ZT-7'").fetchone()[0] == "แพ็ค"    # rolled back
    finally:
        c.execute("DELETE FROM bc_ingredient WHERE code='ZT-7'")
        c.commit()
        c.close()
