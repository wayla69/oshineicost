# -*- coding: utf-8 -*-
"""The kitchen builds the BoM from a multi-menu template and submits it; the analyst reviews (PR-C).

The promises:
  * a blank template holds many menus in one file: sheet “สูตร” (the only one `import_flat` reads), the POS
    menus, the ingredient codes and the instructions; a filled copy imports back as recipes;
  * the kitchen submits a monthly draft for review (needs at least one recipe); while it is in review only
    a holder of bom.center (the analyst) may still change a recipe, the kitchen may not;
  * the analyst sends it back with a reason, which reopens it for the kitchen and is kept in the notes;
  * the analyst can still release a draft or a release in review;
  * หน้างาน follows the review: the kitchen's item, then the analyst's;
  * permissions: the kitchen (bom.recipes) may open a draft, import, download the template and submit, but not
    send back, release or run the checks; the analyst sends back; Sales Support has no part in it.
"""
import io
import sqlite3

import openpyxl
import pytest
from fastapi.testclient import TestClient

import auth
from bom_center import recipes as R
from bom_center import release as RL
from bom_center import todo as TD
from bom_center.clock import today as bkk_today
from test_bom_center_s2 import _block
from test_bom_center_s3 import _draft, con  # noqa: F401  (v0 in force; crab-stick sashimi ฿5 today)

KITCHEN = {"bom.recipes", "bom.edit"}
ANALYST = {"bom.center", "bom.release"}


def _fill(wb, rows):
    """The template with its sample rows replaced by `rows`, as bytes ready for import_flat."""
    ws = wb["สูตร"]
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── the template ─────────────────────────────────────────────────────────────
def test_the_template_holds_many_menus_and_imports_back(con):
    c = con
    rid = R.get_or_create_draft(c, "2026-11 draft", "2026-11", "2026-07", "innovation")
    wb = R.template_flat(c, rid)
    assert wb.sheetnames == ["สูตร", "รหัส POS", "รหัสวัตถุดิบ", "วิธีใช้"]
    # only “สูตร” opens with “รหัส POS” in A1: the other sheets are for reading
    assert [ws.title for ws in wb.worksheets if str(ws["A1"].value).strip() == R.FLAT_HEADERS[0]
            and [x.value for x in ws[1]][:len(R.FLAT_HEADERS)] == R.FLAT_HEADERS] == ["สูตร"]
    pos = {r[0].value: r for r in list(wb["รหัส POS"].iter_rows())[1:]}
    assert {"OTHBFSSM00020", "OTHBFGNT00010", "OTHBFSUS00033"} <= set(pos)
    assert [x.value for x in wb["รหัส POS"][1]] == ["รหัสเมนู POS", "ชื่อเมนู", "Tier", "มีสูตรในร่างแล้ว"]
    assert [x.value for x in wb["รหัสวัตถุดิบ"][1]] == ["รหัสวัตถุดิบ", "ชื่อวัตถุดิบ", "หน่วยฐาน", "ชนิด"]
    assert "13-01-009" in {r[0].value for r in list(wb["รหัสวัตถุดิบ"].iter_rows())[1:]}
    samples = [r for r in wb["สูตร"].iter_rows(min_row=2, values_only=True)]
    assert len(samples) == 3 and all("(ตัวอย่าง ลบก่อนนำเข้า)" in r[1] for r in samples)
    assert len({r[0] for r in samples}) == 2                                    # two menus, three lines
    # QA L4d: the sample rows left in by mistake are skipped (no recipe, no error)
    buf = io.BytesIO()
    wb.save(buf)
    out = R.import_flat(c, buf.getvalue(), rid, "innovation", modules=KITCHEN)
    assert out["recipes"] == 0 and out["errors"] == [] and out["lines"] == 0
    # filled in: two menus, three lines
    filled = _fill(R.template_flat(c, rid), [
        ["OTHBFSSM00020", "ซาชิมิปูอัด", 459, 1, "13-01-009", "ปูอัด", 45, "กรัม", None, "", None],
        ["OTHBFSSM00020", "ซาชิมิปูอัด", 459, 2, "V-007", "หัวไชเท้า", 10, "กรัม", None, "", None],
        ["OTHBFGNT00010", "ไก่ย่างเทริยากิ", 459, 1, "13-01-009", "ปูอัด", 20, "กรัม", None, "", None]])
    out = R.import_flat(c, filled, rid, "innovation", modules=KITCHEN)
    assert (out["recipes"], out["lines"], out["errors"]) == (2, 3, [])
    assert len(R.recipe_lines(c, rid, "OTHBFSSM00020")) == 2
    # a menu that has a recipe is marked in the POS sheet next time
    again = {r[0].value: r[3].value for r in list(R.template_flat(c, rid)["รหัส POS"].iter_rows())[1:]}
    assert again["OTHBFSSM00020"] and again["OTHBFGNT00010"] and not again["OTHBFSUS00033"]


# ── submit and the review lock ───────────────────────────────────────────────
def test_submit_review_locks_the_kitchen_out(con):
    c = con
    rid = R.get_or_create_draft(c, "2026-11 draft", "2026-11", "2026-07", "innovation")
    with pytest.raises(ValueError, match="ยังไม่มีสูตร"):
        RL.submit_review(c, rid, "innovation")
    c.rollback()
    rid = _draft(c, [_block(459, "ซาชิมิปูอัด")], label="2026-11 draft")
    line = [{"ingredient_code": "13-01-009", "qty": 50, "unit": "กรัม"}]
    RL.submit_review(c, rid, "innovation")
    rel = R.release_row(c, rid)
    assert rel["status"] == "review" and rel["submitted_by"] == "innovation" and rel["submitted_at"]
    with pytest.raises(ValueError, match="ส่งตรวจได้เฉพาะร่าง"):
        RL.submit_review(c, rid, "innovation")                                  # already submitted
    with pytest.raises(ValueError, match="อยู่ระหว่างตรวจของฝ่ายวิเคราะห์"):
        R.save_lines(c, rid, "OTHBFSSM00020", line, "innovation", modules=KITCHEN)
    with pytest.raises(ValueError, match="อยู่ระหว่างตรวจของฝ่ายวิเคราะห์"):
        R.get_or_create_draft(c, "2026-11 draft", "2026-11", "2026-07", "innovation", modules=KITCHEN)
    with pytest.raises(ValueError, match="อยู่ระหว่างตรวจของฝ่ายวิเคราะห์"):
        R.import_flat(c, _fill(R.template_flat(c, rid), []), rid, "innovation", modules=KITCHEN)
    R.save_lines(c, rid, "OTHBFSSM00020", line, "analyst", modules=ANALYST)     # the analyst may still edit
    R.save_lines(c, rid, "OTHBFSSM00020", line, "system", modules=None)         # no permission context: internal
    assert R.get_or_create_draft(c, "2026-11 draft", "2026-11", "2026-07", "analyst", modules=ANALYST) == rid
    assert c.execute("SELECT action FROM bc_audit WHERE action='release_submit_review' AND row_key=?",
                     (str(rid),)).fetchone()


def test_send_back_reopens_the_draft_with_the_reason(con):
    c = con
    rid = _draft(c, [_block(459, "ซาชิมิปูอัด")], label="2026-11 draft")
    with pytest.raises(ValueError, match="ส่งกลับได้เฉพาะเวอร์ชันที่รอตรวจ"):
        RL.send_back(c, rid, "x", "analyst")                                    # still a draft
    RL.submit_review(c, rid, "innovation")
    with pytest.raises(ValueError, match="ต้องระบุเหตุผลที่ส่งกลับ"):
        RL.send_back(c, rid, "  ", "analyst")
    RL.send_back(c, rid, "ปริมาณปูอัดสูงเกินไป", "analyst")
    rel = R.release_row(c, rid)
    assert rel["status"] == "draft" and rel["submitted_by"] is None and rel["submitted_at"] is None
    assert "ส่งกลับแก้ไขโดย analyst: ปริมาณปูอัดสูงเกินไป" in rel["notes"]
    R.save_lines(c, rid, "OTHBFSSM00020", [{"ingredient_code": "13-01-009", "qty": 40, "unit": "กรัม"}],
                 "innovation", modules=KITCHEN)                                 # the kitchen edits again
    assert c.execute("SELECT action FROM bc_audit WHERE action='release_send_back' AND row_key=?",
                     (str(rid),)).fetchone()


def test_release_still_accepts_a_draft_or_a_release_in_review(con):
    c = con
    rid = _draft(c, [_block(459, "ซาชิมิปูอัด")], label="2026-11 draft")
    RL.submit_review(c, rid, "innovation")
    assert RL.run_checks(c, rid)["can_release"]
    assert RL.do_release(c, rid, "analyst", backdate_reason="ทดสอบ")["released"]
    assert R.release_row(c, rid)["status"] == "released"
    rid2 = _draft(c, [_block(459, "ซาชิมิปูอัด")], label="2026-12 draft")        # straight from draft, as before
    c.execute("UPDATE bc_release SET effective_month='2026-12' WHERE release_id=?", (rid2,))
    assert R.release_row(c, rid2)["status"] == "draft"
    assert RL.do_release(c, rid2, "analyst", backdate_reason="ทดสอบ")["released"]


# ── หน้างาน follows the review ───────────────────────────────────────────────
def _items(c):
    return {i["key"]: i for i in TD._month_items(c, "2026-09", bkk_today())}


def test_todo_items_follow_the_review(con):
    c = con
    rid = R.get_or_create_draft(c, "2026-11 draft", "2026-11", "2026-07", "innovation")
    it = _items(c)["draft"]
    assert it["status"] == "todo" and "ส่งตรวจ" in it["title"] and "ครัว" in it["owner_th"]
    assert it["owner_modules"] == ["bom.recipes"] and it["go"]["pane"] == "rec" and it["go"]["extra"]["release_id"] == rid
    assert "review_pending" not in _items(c)
    c.rollback()
    rid = _draft(c, [_block(459, "ซาชิมิปูอัด")], label="2026-11 draft")
    RL.submit_review(c, rid, "innovation")
    items = _items(c)
    assert items["draft"]["status"] == "done" and items["draft"]["detail"].startswith("ส่งตรวจแล้ว")
    assert "innovation" in items["draft"]["detail"]
    rp = items["review_pending"]
    assert rp["status"] == "todo" and "ฝ่ายวิเคราะห์" in rp["owner_th"] and rp["owner_modules"] == ["bom.release"]
    assert rp["go"]["pane"] == "rel" and rp["go"]["extra"]["release_id"] == rid
    RL.send_back(c, rid, "แก้ไข", "analyst")
    items = _items(c)
    assert items["draft"]["status"] == "todo" and "review_pending" not in items


def test_todo_without_a_draft_points_to_the_template(con):
    c = con
    c.execute("DELETE FROM bc_release WHERE kind='monthly' AND status IN ('draft','review')")
    it = _items(c)["draft"]
    if it["status"] == "todo":                                                  # (a released version of that month would say done)
        assert "แม่แบบ" in it["detail"] and it["owner_th"].startswith("ครัว") and it["owner_modules"] == ["bom.recipes"]


# ── who may call what ────────────────────────────────────────────────────────
def test_the_permission_matrix():
    innovation, analyst = auth.modules_for_role("innovation"), auth.modules_for_role("analyst")
    sales = auth.modules_for_role("salessupport")
    kitchen_ok = [("POST", "/api/bom-center/release-act/new-draft"), ("POST", "/api/bom-center/release-act/submit-review"),
                  ("POST", "/api/bom-center/recipes/import-flat"), ("GET", "/api/bom-center/recipes/template.xlsx"),
                  ("PUT", "/api/bom-center/recipes/OTHBFSSM00020")]
    for m, p in kitchen_ok:
        assert auth.path_allowed(m, p, innovation), p
    for p in ("send-back", "release", "check", "accept", "acknowledge", "retire"):
        assert not auth.path_allowed("POST", f"/api/bom-center/release-act/{p}", innovation), p
    assert auth.path_allowed("POST", "/api/bom-center/release-act/send-back", analyst)
    assert auth.path_allowed("POST", "/api/bom-center/release-act/release", analyst)
    assert auth.path_allowed("POST", "/api/bom-center/release-act/submit-review", analyst)   # bom.center
    for m, p in kitchen_ok[:2] + [("POST", "/api/bom-center/release-act/send-back")]:
        assert not auth.path_allowed(m, p, sales), p


@pytest.fixture
def auth_on(monkeypatch):
    monkeypatch.setenv("OSHINEI_AUTH", "1")
    monkeypatch.setenv("OSHINEI_AUTH_SECRET", "test-secret-stable")
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "owner-token-xyz")
    yield


TOKEN = {"X-Upload-Token": "owner-token-xyz"}


def _login(user):
    import main
    c = TestClient(main.app)
    assert c.post("/api/auth/login", json={"username": user, "password": auth.DEFAULT_PASSWORD}).status_code == 200
    return c


def _drop(label):
    import core
    with core.get_db() as c:
        ids = [r[0] for r in c.execute("SELECT release_id FROM bc_release WHERE label=?", (label,))]
        for i in ids:
            for t in ("bc_recipe_line", "bc_recipe", "bc_release_check"):
                c.execute(f"DELETE FROM {t} WHERE release_id=?", (i,))
            c.execute("DELETE FROM bc_audit WHERE table_name='bc_release' AND row_key=?", (str(i),))
            c.execute("DELETE FROM bc_release WHERE release_id=?", (i,))
        c.commit()


def test_http_the_kitchen_submits_the_analyst_sends_back(client, app_db, auth_on):
    label = "2099-01 pr-c review flow"
    try:
        _http_flow(label)
    finally:
        _drop(label)


def _http_flow(label):
    kitchen, analyst, sales = _login("innovation1"), _login("analyst1"), _login("salessupport1")
    r = kitchen.get("/api/bom-center/recipes/template.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["สูตร", "รหัส POS", "รหัสวัตถุดิบ", "วิธีใช้"]
    # the kitchen opens a fresh draft (no source release) and cannot send it back, release it or run checks
    body = {"effective_month": "2099-01", "price_month": "2026-07", "label": label}
    assert kitchen.post("/api/bom-center/release-act/new-draft", json=body).headers.get("X-Denied") == "token"
    r = kitchen.post("/api/bom-center/release-act/new-draft", json=body, headers=TOKEN)
    assert r.status_code == 200, r.text
    rid = r.json()["release_id"]
    rel = next(x for x in kitchen.get("/api/bom-center/releases").json()["releases"] if x["release_id"] == rid)
    assert rel["status"] == "draft" and rel["next_step"] == "submit"
    for p in ("send-back", "release", "check"):
        d = kitchen.post(f"/api/bom-center/release-act/{p}?release_id={rid}", json={"release_id": rid, "note": "x"},
                         headers=TOKEN)
        assert d.status_code == 403 and d.headers.get("X-Denied") == "role", p
    assert sales.post("/api/bom-center/release-act/submit-review", json={"release_id": rid},
                      headers=TOKEN).headers.get("X-Denied") == "role"
    # an empty draft cannot be submitted
    r = kitchen.post("/api/bom-center/release-act/submit-review", json={"release_id": rid}, headers=TOKEN)
    assert r.status_code == 400 and "ยังไม่มีสูตร" in r.json()["detail"]
    # the analyst cannot send back what is not in review
    r = analyst.post("/api/bom-center/release-act/send-back", json={"release_id": rid, "note": "x"}, headers=TOKEN)
    assert r.status_code == 400 and "ส่งกลับได้เฉพาะเวอร์ชันที่รอตรวจ" in r.json()["detail"]
    # put it in review as the kitchen's submission would (a recipe is needed first)
    import core
    with core.get_db() as c:
        c.execute("INSERT INTO bc_recipe(release_id, sku_code, kind) VALUES(?, 'OTHBFX', 'recipe')", (rid,))
        c.commit()
    assert kitchen.post("/api/bom-center/release-act/submit-review", json={"release_id": rid},
                        headers=TOKEN).status_code == 200
    rel = next(x for x in kitchen.get("/api/bom-center/releases").json()["releases"] if x["release_id"] == rid)
    assert rel["status"] == "review" and rel["next_step"] == "release" and rel["submitted_by"] == "innovation1"
    # in review the kitchen cannot change a recipe, nor open the draft again; the analyst is not stopped by the lock
    put = kitchen.put(f"/api/bom-center/recipes/OTHBFX?release_id={rid}", json=[], headers=TOKEN)
    assert put.status_code == 400 and "อยู่ระหว่างตรวจ" in put.json()["detail"]
    again = kitchen.post("/api/bom-center/release-act/new-draft", json=body, headers=TOKEN)
    assert again.status_code == 400 and "อยู่ระหว่างตรวจ" in again.json()["detail"]
    imp = kitchen.post(f"/api/bom-center/recipes/import-flat?release_id={rid}&dry_run=true",
                       files={"file": ("t.xlsx", kitchen.get("/api/bom-center/recipes/template.xlsx").content)})
    assert imp.status_code == 400 and "อยู่ระหว่างตรวจ" in imp.json()["detail"]
    ok = analyst.put(f"/api/bom-center/recipes/OTHBFX?release_id={rid}", json=[], headers=TOKEN)
    assert "อยู่ระหว่างตรวจ" not in ok.text
    # send back needs a reason, then reopens the draft for the kitchen
    r = analyst.post("/api/bom-center/release-act/send-back", json={"release_id": rid, "note": " "}, headers=TOKEN)
    assert r.status_code == 400 and "ต้องระบุเหตุผล" in r.json()["detail"]
    r = analyst.post("/api/bom-center/release-act/send-back", json={"release_id": rid, "note": "ปริมาณสูงเกินไป"},
                     headers=TOKEN)
    assert r.status_code == 200, r.text
    rel = next(x for x in kitchen.get("/api/bom-center/releases").json()["releases"] if x["release_id"] == rid)
    assert rel["status"] == "draft" and rel["next_step"] == "submit" and rel["submitted_by"] is None
    assert "ส่งกลับแก้ไขโดย analyst1: ปริมาณสูงเกินไป" in rel["notes"]          # the kitchen can read why
    put = kitchen.put(f"/api/bom-center/recipes/OTHBFX?release_id={rid}", json=[], headers=TOKEN)
    assert "อยู่ระหว่างตรวจ" not in put.text
