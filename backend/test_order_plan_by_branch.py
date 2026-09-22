"""ปริมาณที่ควรสั่ง รายวัตถุดิบ — per branch, whole numbers, always rounded up.

Owner, 2026-09-22: "ต้องทำเป็นรายงานแยกสาขาครับ จะได้แนะนำสาขาได้ถูก และปัดเศษให้เป็นจำนวนเต็ม
(ปัดขึ้นทุกกรณี)". Runs against the shared fixture DB (two branches, salmon plate + steak).
"""
import io
import math
import sqlite3

import openpyxl
import pytest

import conftest as _cf
import main
import routers.ingredients as ing

END = f"{_cf.LATEST_MONTH}-03"
PLAN = {"date_to": END, "days": 28}
BOM = [
    (2, 1, "ING-SAL", "แซลมอน", 0.12, "กก.", 500.0, 60.0, "protein"),
    (2, 2, "ING-RICE", "ข้าวญี่ปุ่น", 0.05, "กก.", 40.0, 2.0, "carb"),
    (3, 1, "ING-BEEF", "เนื้อ", 0.20, "กก.", 600.0, 120.0, "protein"),
]


@pytest.fixture(scope="module", autouse=True)
def _seed(client):
    con = sqlite3.connect(str(main.DB_PATH))
    keys = [(b[0], b[1]) for b in BOM]
    before = [r for k in keys for r in con.execute("SELECT * FROM bom_ingredients WHERE sku_id=? AND seq=?", k)]
    con.executemany(
        "INSERT OR REPLACE INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name,"
        " qty_per_serving, unit, price_per_unit, ingredient_cost, ingredient_type, source_file)"
        " VALUES(?,?,?,?,?,?,?,?,?,'test')", BOM)
    for tbl in ("ingredient_order_unit", "ingredient_sourcing"):      # another module may have left rows (test_ingredients does)
        try:
            con.execute(f"DELETE FROM {tbl}")
        except sqlite3.OperationalError:
            pass
    con.commit(); con.close()
    main.invalidate_cache()
    yield
    con = sqlite3.connect(str(main.DB_PATH))
    con.executemany("DELETE FROM bom_ingredients WHERE sku_id=? AND seq=?", keys)
    if before:
        con.executemany(f"INSERT INTO bom_ingredients VALUES ({','.join('?' * len(before[0]))})", before)
    for t in ("ingredient_order_unit", "ingredient_sourcing"):
        try:
            con.execute(f"DELETE FROM {t}")
        except sqlite3.OperationalError:
            pass
    con.commit(); con.close()
    main.invalidate_cache()


# ── the rounding rule, on its own ────────────────────────────────────────────
def test_every_order_quantity_is_a_whole_number_rounded_up():
    o = lambda need, unit="กรัม", **pk: ing._ing_to_order(need, unit, pk or None)
    assert o(12_345.6)["units"] == 13 and o(12_345.6)["order_unit"] == "กก."      # fallback: whole kilos
    assert o(1.0)["units"] == 1 and o(0)["units"] == 0
    assert o(7, "แผ่น")["units"] == 7 and o(7.01, "ชิ้น")["units"] == 8            # a unit with no fallback: itself, whole
    assert o(25, "กก.", qty_per_order=10, order_unit="ลัง")["units"] == 3           # ceil(2.5)
    assert o(25, "กก.", qty_per_order=10, round_step=0.5)["units"] == 3            # half-cartons are not a thing any more
    assert o(25, "กก.", qty_per_order=10, round_step=6)["units"] == 6              # cartons come in sixes: up to the next six
    assert o(25, "กก.", qty_per_order=10, moq=5)["units"] == 5                     # minimum order still applies
    assert o(1, "กก.", qty_per_order=10, moq=2.5)["units"] == 3                    # even a fractional MOQ rounds up
    for r in (o(12_345.6), o(25, "กก.", qty_per_order=10, round_step=0.5), o(1, "กก.", qty_per_order=10, moq=2.5)):
        assert isinstance(r["units"], int)


def test_sheet_names_are_excel_safe_and_unique():
    taken = set()
    assert ing._sheet_title("ขอนแก่น", taken) == "ขอนแก่น"
    assert ing._sheet_title("โคราช/สาขา 2 [ใหม่]", taken) == "โคราชสาขา 2 ใหม่"
    long = "x" * 40
    assert ing._sheet_title(long, taken) == "x" * 31
    assert ing._sheet_title(long, taken) == "x" * 27 + " (2)"                   # still 31, still unique


# ── the workbook ─────────────────────────────────────────────────────────────
def _wb(client, **kw):
    r = client.get("/api/ingredients/order-plan-by-branch.xlsx", params={**PLAN, **kw})
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"].startswith("attachment; filename=ORDER_PLAN_BY_BRANCH_")
    return openpyxl.load_workbook(io.BytesIO(r.content))


def _rows(ws):
    head = [c.value for c in ws[4]]
    out = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[1] and r[0] and isinstance(r[0], int):
            out.append(dict(zip(head, r)))
    return head, out


def _tab(client, branch, include="all"):
    """What the tab shows for one branch - the report must say the same."""
    r = client.get("/api/ingredients/order-plan", params={**PLAN, "branch": branch, "include": include})
    assert r.status_code == 200, r.text
    return {x["code"]: x for x in r.json()["rows"]}


def test_one_sheet_per_branch_with_whole_numbers(client, monkeypatch):
    monkeypatch.delenv("OSHINEI_SEED_TOKEN", raising=False)
    assert client.post("/api/ingredients/order-units", json=[{
        "ingredient_code": "ING-SAL", "base_unit": "กก.", "order_unit": "ลัง", "qty_per_order": 10, "price_per_order": 4800}]).status_code == 200
    main.invalidate_cache()
    wb = _wb(client, include="all")
    assert wb.sheetnames == ["สรุป", "BranchA", "BranchB"]                    # summary first, then a sheet per branch
    tab = {b: _tab(client, b) for b in ("BranchA", "BranchB")}
    for b in ("BranchA", "BranchB"):
        head, rows = _rows(wb[b])
        assert head[:4] == ["ลำดับ", "รหัส", "วัตถุดิบ", "บังคับสั่ง"] and "ควรสั่ง/สัปดาห์" in head and "ควรสั่ง/เดือน" in head
        assert {r["รหัส"] for r in rows} == set(tab[b])                        # exactly what that branch sells
        for r in rows:
            assert isinstance(r["ควรสั่ง/สัปดาห์"], int) and isinstance(r["ควรสั่ง/เดือน"], int), r
            assert r["ควรสั่ง/สัปดาห์"] > 0 and r["ควรสั่ง/เดือน"] >= r["ควรสั่ง/สัปดาห์"]
            assert r["ควรสั่ง/สัปดาห์"] == tab[b][r["รหัส"]]["order_week_units"]     # the same number the tab shows
        assert [r["รหัส"] for r in rows] == sorted(tab[b], key=lambda c: -tab[b][c]["value_month"])   # biggest first
    _, rows = _rows(wb["BranchA"])
    sal = next(r for r in rows if r["รหัส"] == "ING-SAL")
    assert sal["หน่วยสั่ง"] == "ลัง" and sal["ขนาดบรรจุ (1 หน่วยสั่ง = ? หน่วยนับ)"] == 10
    assert sal["ควรสั่ง/สัปดาห์"] == math.ceil(tab["BranchA"]["ING-SAL"]["order_week"] / 10)
    assert wb["BranchA"].cell(row=5, column=9).number_format == "#,##0"
    # the summary adds up the branch sheets
    summ = {r[0]: r for r in wb["สรุป"].iter_rows(min_row=5, values_only=True) if r[0]}
    assert summ["BranchA"][1] == len(tab["BranchA"]) and summ["BranchA"][5] == "BranchA"
    assert summ["รวม"][1] == len(tab["BranchA"]) + len(tab["BranchB"])
    assert abs(summ["รวม"][3] - (summ["BranchA"][3] + summ["BranchB"][3])) < 0.01


def test_default_is_the_central_list_and_a_single_branch_is_one_sheet(client, monkeypatch):
    monkeypatch.delenv("OSHINEI_SEED_TOKEN", raising=False)
    assert client.post("/api/ingredients/sourcing", json=[{"ingredient_code": "ING-RICE", "branch": "", "mode": "local"}]).status_code == 200
    main.invalidate_cache()
    try:
        every = set(_tab(client, "BranchA"))
        assert "ING-RICE" in every
        _, rows = _rows(_wb(client)["BranchA"])                                  # no include= -> central
        assert {r["รหัส"] for r in rows} == every - {"ING-RICE"}, "rice is bought by the branch, so it is not advised"
        _, rows = _rows(_wb(client, include="local")["BranchA"])
        assert {r["รหัส"] for r in rows} == {"ING-RICE"} and "สาขาซื้อเอง" in rows[0]["หมายเหตุ"]
        wb = _wb(client, branch="BranchB", include="all")
        assert wb.sheetnames == ["สรุป", "BranchB"]
        _, rows = _rows(wb["BranchB"])
        assert {r["รหัส"] for r in rows} == set(_tab(client, "BranchB"))
    finally:
        client.post("/api/ingredients/sourcing", json=[{"ingredient_code": "ING-RICE", "branch": "", "mode": "default"}])
        main.invalidate_cache()


def test_top_n_and_the_fallback_unit_note(client):
    tab = _tab(client, "BranchA")
    biggest = max(tab, key=lambda c: tab[c]["value_month"])
    _, rows = _rows(_wb(client, include="all", top=1)["BranchA"])
    assert [r["รหัส"] for r in rows] == [biggest]
    _, rows = _rows(_wb(client, include="all")["BranchA"])
    rice = next(r for r in rows if r["รหัส"] == "ING-RICE")
    assert rice["หน่วยสั่ง"] == "กก." and "ยังไม่ตั้งขนาดบรรจุ" in rice["หมายเหตุ"]
