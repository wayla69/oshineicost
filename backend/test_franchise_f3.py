"""PR-F3 - ร้านของฉัน: วัตถุดิบจากส่วนกลาง (what the branch buys from the centre).

  * GET /api/franchise/supply: the payload shape; the delivery rounds and their ordered / delivered baht are the supply module's
    ONE order value (hub.with_values, as orders.history shows them); purchases against sales = delivered / the report's net sales
    of the same window, against last month; claims raised in the month by status with the approved value; late orders counted;
    the latest counted period's over-use from variance.period, in the item's readable unit;
  * a branch that has not started the supply module gets one neutral empty state;
  * the verdict line on purchases against sales, from its module constants;
  * the sheet วัตถุดิบ in summary.xlsx;
  * 403 for another branch and for a franchise login without a branch; no write; the forbidden-word and other-branch scan on the
    payload and the sheet; the router still imports its neighbours relatively.
"""
import io
import json
import sqlite3

import openpyxl
import pytest

import auth
from test_franchise_role import TOKEN, _walk, auth_on, world      # noqa: F401  (fixtures used by name)

FORBIDDEN = ("recipe", "price", "cashier", "staff", "supplier", "receipt_no", "bill_no", "username", "password", "token", "bom")
MONTH = "2026-06"
SAL, RICE = "F3-SAL", "F3-RICE"
TABLES = ("bc_ingredient", "bc_ingredient_price", "bc_order", "bc_order_line", "bc_claim", "bc_claim_line", "bc_stock_count",
          "bc_stock_count_line", "bc_receipt", "bc_receipt_line", "bc_issue_slip")


def _db():
    import main
    return sqlite3.connect(str(main.DB_PATH))


def _cols(c, t):
    return [r[1] for r in c.execute(f"PRAGMA table_info({t})")]


@pytest.fixture
def f3data(world, monkeypatch):
    """A supply cycle for BranchB (id 2): orders delivered in May and June, claims, a counted period. Every bc_* table it touches
    is saved first and put back exactly afterwards; the branches' own rows from earlier modules are cleared at setup."""
    import main
    from bom_center import forecast as FC
    c = _db()
    saved = {t: (_cols(c, t), c.execute(f"SELECT * FROM {t}").fetchall()) for t in TABLES}
    try:
        for t in ("bc_order", "bc_claim", "bc_stock_count", "bc_receipt", "bc_issue_slip"):
            c.execute(f"DELETE FROM {t} WHERE branch_id IN (1, 2)")
        c.execute("DELETE FROM bc_ingredient WHERE code IN (?, ?)", (SAL, RICE))
        c.execute("DELETE FROM bc_ingredient_price WHERE ingredient_code IN (?, ?)", (SAL, RICE))
        c.executemany("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv, loss_pct) VALUES(?,?,?,?,?,?,0)",
                      [(SAL, "แซลมอนทดสอบ", "raw", "กรัม", "แพ็ค 1 กก.", 1000), (RICE, "ข้าวทดสอบ", "raw", "กรัม", None, None)])
        c.executemany("INSERT INTO bc_ingredient_price(ingredient_code, month, price_per_base, source) VALUES(?,?,?,'manual')",
                      [(SAL, "2026-05", 0.4), (RICE, "2026-05", 0.05), (SAL, "2026-06", 0.4), (RICE, "2026-06", 0.05)])
        orders = [  # (key, round, cycle, delivery, status, late, [(code, confirmed, processed, received, unit)])
            ("may", "1", "2026-05-10", "2026-05-12", "received", 0, [(SAL, 5, 5, 5, "pack")]),
            ("o1", "1", "2026-05-31", "2026-06-02", "received", 0, [(SAL, 3, 2, 2, "pack"), (RICE, 5000, 5000, 5000, "base")]),
            ("o2", "2", "2026-06-03", "2026-06-05", "processed", 1, [(SAL, 1, 1, None, "pack")]),
            ("o3", "salmon", "2026-06-08", "2026-06-10", "submitted", 1, [(SAL, 2, None, None, "pack")]),
            ("draft", "1", "2026-06-10", "2026-06-12", "confirmed", 0, [(SAL, 9, None, None, "pack")]),
            ("gone", "2", "2026-06-10", "2026-06-12", "cancelled", 0, [(SAL, 9, None, None, "pack")]),
            ("a1", "1", "2026-06-01", "2026-06-03", "confirmed", 0, [(SAL, 1, None, None, "pack")]),  # BranchA's draft: not a start
        ]
        ids = {}
        for key, rnd, cyc, dlv, st, late, lines in orders:
            bid = 1 if key == "a1" else 2
            cur = c.execute("INSERT INTO bc_order(branch_id, cycle_date, round, status, submitted_at, delivery_date, late)"
                            " VALUES(?,?,?,?,?,?,?)", (bid, cyc, rnd, st, cyc + "T05:00:00", dlv, late))
            ids[key] = cur.lastrowid
            c.executemany("INSERT INTO bc_order_line(order_id, ingredient_code, suggested_units, confirmed_units, processed_units,"
                          " received_units, order_unit) VALUES(?,?,?,?,?,?,?)",
                          [(cur.lastrowid, code, cf, cf, pr, rc, u) for code, cf, pr, rc, u in lines])
        claims = [  # (order, status, created_at, decision note, [(code, claim, approved, reason)])
            ("o1", "approved", "2026-06-02T03:00:00", "รับเคลมแล้ว ส่งทดแทนรอบหน้า ตามที่สาขาBranchA แจ้ง", [(SAL, 1, 1, "damaged")]),
            ("o1", "rejected", "2026-06-02T04:00:00", "ไม่พบส่วนต่าง", [(RICE, 500, None, "short")]),
            ("o2", "submitted", "2026-06-05T02:00:00", None, [(SAL, 1, None, "spoiled")]),
            ("o2", "cancelled", "2026-06-05T03:00:00", None, [(SAL, 1, None, "other")]),
            ("may", "approved", "2026-05-20T03:00:00", "ok", [(SAL, 1, 1, "damaged")]),
            ("o1", "submitted", "2026-05-31T18:30:00", None, [(SAL, 1, None, "expired")]),       # 1 June 01:30 in Bangkok
        ]
        for key, st, at, dn, lines in claims:
            cur = c.execute("INSERT INTO bc_claim(branch_id, order_id, status, note, created_by, created_at, decision_note, decided_by, origin)"
                            " VALUES(2,?,?,?,?,?,?,?,'after')", (ids[key], st, "บันทึกภายในของสาขา", "own_t_b", at, dn, "salessupport1"))
            c.executemany("INSERT INTO bc_claim_line(claim_id, ingredient_code, claim_units, reason_code, approved_units) VALUES(?,?,?,?,?)",
                          [(cur.lastrowid, code, q, rs, a) for code, q, a, rs in lines])
        cnt = []
        for d, kind, sal, rice in (("2026-05-01", "begin", 5000, 10000), ("2026-06-01", "end", 1000, 2000)):
            cur = c.execute("INSERT INTO bc_stock_count(branch_id, count_date, kind, counted_by, created_at, status)"
                            " VALUES(2,?,?,'own_t_b',?,'submitted')", (d, kind, d + "T10:00:00"))
            cnt.append(cur.lastrowid)
            c.executemany("INSERT INTO bc_stock_count_line(count_id, ingredient_code, qty_kitchen, unit, qty_base) VALUES(?,?,?,'กรัม',?)",
                          [(cur.lastrowid, SAL, sal, sal), (cur.lastrowid, RICE, rice, rice)])
        cur = c.execute("INSERT INTO bc_receipt(branch_id, receipt_date, source, ref) VALUES(2,'2026-05-12','order',?)", (f"ORDER-{ids['may']}",))
        c.execute("INSERT INTO bc_receipt_line(receipt_id, ingredient_code, qty, unit, qty_base) VALUES(?,?,5,'แพ็ค 1 กก.',5000)",
                  (cur.lastrowid, SAL))
        c.commit()
        # the recipes' use over the period: salmon 7,000 g against 9,000 g used (over), rice 7,900 g against 8,000 g (within)
        real = FC.usage_history

        def fake(con, branch_id, start, end, month):
            if branch_id == 2:
                return {SAL: {start.isoformat(): 7000.0}, RICE: {start.isoformat(): 7900.0}}, [start.isoformat()], {}
            return real(con, branch_id, start, end, month)
        monkeypatch.setattr(FC, "usage_history", fake)
        yield {"ids": ids, "counts": cnt}
    finally:
        for t in reversed(TABLES):
            cols, rows = saved[t]
            c.execute(f"DELETE FROM {t}")
            if rows:
                c.executemany(f"INSERT INTO {t}({', '.join(cols)}) VALUES({', '.join('?' * len(cols))})", rows)
        c.commit()
        c.close()
        if hasattr(main, "invalidate_cache"):
            try:
                main.invalidate_cache()
            except Exception:
                pass


def _supply(c, **q):
    r = c.get("/api/franchise/supply", params={"month": MONTH, **q})
    assert r.status_code == 200, r.text
    return r.json()


def _report_sales(month):
    from routers import franchise as F
    rep = F._report("BranchB", month)
    assert rep["selected_month"] == month
    return rep["headline"]["revenue"], rep["range"][1], not rep.get("partial")


# ── the rule ─────────────────────────────────────────────────────────────────
def test_the_supply_read_is_under_the_one_get_rule():
    mods = auth.ROLE_MODULES["franchise"]
    assert auth.path_allowed("GET", "/api/franchise/supply", mods)
    for m in ("POST", "PUT", "PATCH", "DELETE"):
        assert not auth.path_allowed(m, "/api/franchise/supply", mods), m


# ── the payload ──────────────────────────────────────────────────────────────
def test_the_payload_has_the_agreed_shape(f3data, world):
    d = _supply(world["b"])
    assert d["branch"] == "BranchB" and d["month"] == MONTH and d["started"] is True and d["not_started_th"] is None
    assert set(d) == {"branch", "month", "label", "started", "start_date", "not_started_th", "rounds", "purchases", "claims",
                      "late", "overuse"}
    for r in d["rounds"]:
        assert set(r) == {"delivery", "round_th", "ordered", "delivered", "status", "status_th", "late"}
    assert {"ordered", "delivered", "rounds", "rounds_delivered", "pct", "pct_prev", "change_pp", "through",
            "month_complete", "delivered_to_date", "sales_to_date", "delivered_prev", "sales_prev", "month_prev", "days"} == set(d["purchases"])
    assert set(d["claims"]) == {"count", "by_status", "approved_value", "latest"}
    assert [x["status_th"] for x in d["claims"]["by_status"]] == ["รอพิจารณา", "อนุมัติ", "ไม่อนุมัติ"]
    for x in d["claims"]["latest"]:
        assert set(x) == {"date", "round_th", "items", "status", "status_th", "decision_note"}
        assert all(set(i) == {"name", "qty", "unit", "reason_th"} for i in x["items"])
    assert set(d["late"]) == {"count", "text"}
    assert set(d["overuse"]) == {"begin", "end", "rows", "lines_over"}
    assert all(set(r) == {"name", "qty", "unit", "value", "pct"} for r in d["overuse"]["rows"])


def test_the_rounds_are_the_order_screens_values(f3data, world):
    """The month's SENT orders by delivery day, valued as the branch's order history values them (hub.with_values)."""
    import main
    from bom_center import orders as OR
    d = _supply(world["b"])
    assert [(r["delivery"], r["round_th"], r["status"]) for r in d["rounds"]] == [
        ("2026-06-02", "รอบ 1 · อ. 2/6", "received"), ("2026-06-05", "รอบ 2 · ศ. 5/6", "processed"),
        ("2026-06-10", "แซลมอนสด · พ. 10/6", "submitted")]                        # no draft, no cancelled order, nothing of May
    with main.get_db() as con:
        hist = {h["delivery"]: h for h in OR.history(con, 2, limit=50)}
    for r in d["rounds"]:
        h = hist[r["delivery"]]
        assert r["ordered"] == h["value_ordered"] and r["delivered"] == h["value_sent"] and r["status_th"] == h["status_th"]
    # by hand: 3 packs of 1 kg at ฿0.40/g + 5,000 g of rice at ฿0.05/g ordered; 2 packs + the rice sent
    assert [(r["ordered"], r["delivered"]) for r in d["rounds"]] == [(1450.0, 1050.0), (400.0, 400.0), (800.0, None)]
    assert [r["status_th"] for r in d["rounds"]] == ["รับสินค้าแล้ว", "คลังจัดส่งแล้ว", "ส่งแล้ว (รอคลังจัดส่ง)"]
    p = d["purchases"]
    assert p["ordered"] == 2650.0 and p["delivered"] == 1450.0 and p["rounds"] == 3 and p["rounds_delivered"] == 2
    assert d["late"] == {"count": 2, "text": "ส่งหลังเวลาตัดรอบ 2 ใบ (มีค่าปรับ)"} and [r["late"] for r in d["rounds"]] == [False, True, True]


def test_purchases_against_sales_is_the_same_window_and_last_month(f3data, world):
    p = _supply(world["b"])["purchases"]
    sales, through, complete = _report_sales(MONTH)
    assert p["through"] == through == "2026-06-03" and p["month_complete"] is complete is False
    # the month is open: deliveries up to the report's last finished day only (5 June is after it)
    assert p["delivered_to_date"] == 1050.0 and p["sales_to_date"] == round(sales, 2)
    assert p["pct"] == round(100.0 * 1050.0 / sales, 1)
    may, _, _ = _report_sales("2026-05")
    assert p["month_prev"] == "2026-05" and p["delivered_prev"] == 2000.0 and p["sales_prev"] == round(may, 2)
    assert p["pct_prev"] == round(100.0 * 2000.0 / may, 1)
    assert p["change_pp"] == round(p["pct"] - p["pct_prev"], 1)
    # the summary carries the same ratio for its verdict
    s = world["b"].get("/api/franchise/summary", params={"month": MONTH}).json()
    assert s["supply_ratio"]["pct"] == p["pct"] and s["supply_ratio"]["pct_prev"] == p["pct_prev"]


def test_the_claims_of_the_month(f3data, world):
    c = _supply(world["b"])["claims"]
    # raised in June (Bangkok day): approved, rejected, submitted, and the one keyed at 01:30 on 1 June; cancelled and May left out
    assert c["count"] == 4
    assert {x["status"]: x["count"] for x in c["by_status"]} == {"submitted": 2, "approved": 1, "rejected": 1}
    assert c["approved_value"] == 400.0                                            # 1 pack of 1 kg at ฿0.40/g
    assert len(c["latest"]) == 4 and [x["date"] for x in c["latest"]] == ["2026-06-05", "2026-06-02", "2026-06-02", "2026-06-01"]
    top = c["latest"][0]
    assert top["items"] == [{"name": "แซลมอนทดสอบ", "qty": 1.0, "unit": "แพ็ค", "reason_th": "สินค้าเสีย เน่า หรือคุณภาพไม่ได้มาตรฐาน"}]
    assert top["status_th"] == "รอพิจารณา" and top["decision_note"] is None and top["round_th"] == "รอบ 2 (ส่ง 5/6)"
    rej = [x for x in c["latest"] if x["status"] == "rejected"][0]
    assert rej["items"][0]["unit"] == "กรัม" and rej["items"][0]["qty"] == 500.0 and rej["decision_note"] == "ไม่พบส่วนต่าง"
    ok = [x for x in c["latest"] if x["status"] == "approved"][0]
    assert ok["decision_note"] == "รับเคลมแล้ว ส่งทดแทนรอบหน้า ตามที่สาขาอื่น แจ้ง"                     # another branch scrubbed


def test_the_overuse_is_the_latest_counted_period(f3data, world):
    import main
    from bom_center import variance as VA
    o = _supply(world["b"])["overuse"]
    assert (o["begin"], o["end"]) == ("2026-05-01", "2026-06-01")
    with main.get_db() as con:
        v = VA.period(con, 2, f3data["counts"][1])
    sal = [r for r in v["rows"] if r["code"] == SAL][0]
    assert sal["status"] == "over" and [r["code"] for r in v["rows"] if r["status"] == "over"] == [SAL]
    assert o["rows"] == [{"name": "แซลมอนทดสอบ", "qty": 2.0, "unit": "กิโลกรัม", "value": round(sal["gap_value"]), "pct": sal["gap_pct"]}]
    assert o["rows"][0]["value"] == 800 and o["lines_over"] == 1                 # 2,000 g over at ฿0.40/g; rice within its band


def test_a_branch_that_has_not_started_gets_one_neutral_line(f3data, world):
    d = _supply(world["a"])
    assert d["started"] is False and d["not_started_th"] == "สาขายังไม่เริ่มใช้ระบบสั่งวัตถุดิบ"
    assert d["rounds"] == [] and d["purchases"] is None and d["claims"] is None and d["late"] is None and d["overuse"] is None
    assert world["mg"].get("/api/franchise/summary", params={"branch_id": 1, "month": MONTH}).json()["supply_ratio"] is None


def test_the_verdict_line_follows_its_constants():
    from routers import franchise as F
    base = {"pct": 38.0, "pct_prev": 33.0, "change_pp": 5.0, "month_complete": True, "days": 30}
    lvl, text = F._supply_verdict(base)
    assert lvl == "watch" and "ร้อยละ 38" in text and "5 จุด" in text and "ควรตรวจสอบ" in text
    assert F._supply_verdict({**base, "change_pp": F.T_PURCHASE_RISE_WATCH}) is None              # not more than the band
    assert F._supply_verdict({**base, "change_pp": -6.0}) is None                                   # a fall is not judged
    assert F._supply_verdict({**base, "month_complete": False, "days": F.T_PURCHASE_MIN_DAYS - 1}) is None
    assert F._supply_verdict({**base, "month_complete": False, "days": F.T_PURCHASE_MIN_DAYS})[0] == "watch"
    assert F._supply_verdict({**base, "pct_prev": None, "change_pp": None}) is None and F._supply_verdict(None) is None
    v = F._verdict(None, None, None, None, None, None, supply=base)
    assert v["watch"] == [text]
    # the claims line stays the discipline block's own: one line, not two
    v = F._verdict(None, None, None, {"module_active": True, "claims_month": 4}, None, None, supply=base)
    assert sum("เคลม" in t for t in v["watch"]) == 1


def test_the_streak_line_names_what_the_usage_is_compared_with():
    """P5: in stock mode the streak is against the branch's normal / the all-branch median (no recipe); in full mode it stays "สูตร"."""
    from routers import franchise as F
    ing = {"latest_period": None, "streak_flag": True, "streak_days": 3, "waste_pct_month": None}
    stock = F._verdict(None, None, {**ing, "stock_mode": True}, None, None, None)["fix"]
    full = F._verdict(None, None, {**ing, "stock_mode": False}, None, None, None)["fix"]
    assert stock == ["มีวัตถุดิบหลักที่เบิกเกินปกติของสาขาหรือค่ากลางทุกสาขาติดต่อกันหลายวัน (3 วันในเดือนนี้) ควรตรวจสอบหน้างาน"]
    assert full == ["มีวัตถุดิบหลักที่เบิกเกินสูตรติดต่อกันหลายวัน (3 วันในเดือนนี้) ควรตรวจสอบหน้างาน"]


def test_the_download_has_the_supply_sheet(f3data, world):
    r = world["b"].get("/api/franchise/summary.xlsx", params={"month": MONTH})
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["สรุป", "รายวัน", "12 เดือน", "เมนูขายดี", "วัตถุดิบ"]
    ws = wb["วัตถุดิบ"]
    vals = {row[0]: row[1] for row in ws.iter_rows(min_row=3, values_only=True) if row and row[0]}
    d = _supply(world["b"])
    assert vals["ยอดสั่งจากส่วนกลาง (บาท)"] == d["purchases"]["ordered"] and vals["ยอดที่คลังจัดส่ง (บาท)"] == d["purchases"]["delivered"]
    assert vals["ยอดซื้อเทียบยอดขาย (%)"] == d["purchases"]["pct"] and vals["ส่งใบสั่งซื้อหลังเวลาตัดรอบ (ใบ)"] == 2
    assert vals["มูลค่าเคลมที่อนุมัติ (บาท)"] == 400.0
    text = " ".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)
    for w in ("รอบ 1 · อ. 2/6", "แซลมอนทดสอบ", "กิโลกรัม", "วัตถุดิบที่ใช้เกินสูตร 5 อันดับ"):
        assert w in text, w
    cells = " ".join(str(v) for s in wb.worksheets for row in s.iter_rows(values_only=True) for v in row if v is not None)
    assert not [w for w in FORBIDDEN if w in cells.lower()] and "BranchA" not in cells
    assert "บันทึกภายในของสาขา" not in cells and "salessupport1" not in cells and "own_t_b" not in cells


def test_the_payload_holds_no_forbidden_word_person_or_other_branch(f3data, world):
    for who, other in (("b", "BranchA"), ("a", "BranchB")):
        d = _supply(world[who])
        hits = [(p, kind, s) for p, kind, s in _walk(d) if any(w in s.lower() for w in FORBIDDEN)]
        assert not hits, hits
        text = json.dumps(d, ensure_ascii=False)
        assert other not in text
        for w in ("บันทึกภายในของสาขา", "salessupport1", "own_t_b", "order_id", "claim_id", "F3-SAL", "F3-RICE", "ORDER-"):
            assert w not in text, (who, w)


def test_another_branch_and_a_login_without_a_branch_are_refused(f3data, world):
    r = world["b"].get("/api/franchise/supply", params={"branch_id": 1, "month": MONTH})
    assert r.status_code == 403 and r.headers.get("X-Denied") == "role"
    for q in ({}, {"branch_id": 1}, {"branch_id": 2}):
        r = world["nobranch"].get("/api/franchise/supply", params=q)
        assert r.status_code == 403 and r.headers.get("X-Denied") == "role", q
    assert world["mg"].get("/api/franchise/supply").status_code == 400
    assert world["mg"].get("/api/franchise/supply", params={"branch_id": 2, "month": MONTH}).json()["branch"] == "BranchB"
    for m in ("POST", "PUT", "DELETE"):
        r = world["b"].request(m, "/api/franchise/supply", headers={"X-Upload-Token": TOKEN})
        assert r.status_code == 403 and r.headers.get("X-Denied") == "role", m


def test_the_franchise_module_still_imports_its_neighbours_relatively():
    from pathlib import Path
    import re
    src = (Path(__file__).resolve().parent / "routers" / "franchise.py").read_text(encoding="utf-8")
    assert not re.search(r"^\s*from routers[. ]", src, re.M) and not re.search(r"^\s*import routers", src, re.M)
