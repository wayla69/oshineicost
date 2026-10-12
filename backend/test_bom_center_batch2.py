"""BoM improvement proposal · batch 2 (owner approved 2026-10-12): make branch work faster.

The promises:
  * 3.1 a branch issues and counts in the SMALLEST SEALED unit (แพ็ค, not กล่อง (10 แพ็ค)); an analyst's own count unit is kept; the
    month-end count has two boxes (กล่อง + แพ็ค) that sum in base units; a figure read off the system is rounded to one decimal and
    "น้อยกว่า 1 แพ็ค" under one; a placeholder never shows 0.005;
  * 3.2 the slip opens on the items in use (14 days, system stock, ต้องนับ, or already typed); "แสดงทั้งหมด (N รายการ)" shows the rest and
    is remembered on the device; the zone chips follow the list; a search looks at everything;
  * 2.3 หน้างาน: the strip names the NEXT order due (never "ส่งแล้ว" while another is open), the days left behind are one row, a branch
    has no month to pick;
  * 1.9 the confirm of an order says "N รายการ · ฿x · รอบก่อน ฿y (+z%)" and the 3 lines that changed most;
  * 4.1 a short name in the register (screen and file), shown on the branch screens with the code under "ดูที่มา";
  * 4.2 no record number on a worker's screen; 4.4 dialog buttons are 44 px; 4.5 no pickers where there is nothing to filter, and the
    unsteady items are a link; 4.6 a gap is written in words ("ใช้น้อยกว่าสูตร ฿16,197 (8.1%)").
"""
import asyncio
import io
import json
import re
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest

from bom_center import central as CE
from bom_center import fmt
from bom_center import issues as IS
from bom_center import master as M
from bom_center import orders as OR
from bom_center import stock as ST
from bom_center import todo as TD
from bom_center import variance as VA
from bom_center.units import Converter
from test_bom_center_issues import CRAB, NOW, at  # noqa: F401
from test_bom_center_s11 import ITEMS, _report, udon  # noqa: F401
from test_bom_center_s4b import con  # noqa: F401

REPO = Path(__file__).resolve().parent.parent
NODE = shutil.which("node")
needs_node = pytest.mark.skipif(NODE is None, reason="node is not installed")
BOX = "กล่อง (20 แพ็ค)"


def _src(rel):
    return (REPO / rel).read_text(encoding="utf-8")


def _fn(src, name):
    """The source of `function name(…) {…}` (braces counted) — the screens are one shared global scope, so a test lifts what it needs."""
    m = re.search(r"(?:async )?function %s\(" % re.escape(name), src)
    assert m, name
    i = src.index("{", m.end() - 1 if src[m.end() - 1] == "{" else src.index(")", m.end()))
    depth, j = 0, i
    while j < len(src):
        c = src[j]
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return src[m.start():j + 1]
        j += 1
    raise AssertionError("unbalanced " + name)


def _const(src, name):
    m = re.search(r"^(?:const|let) %s = .*?;\s*$" % re.escape(name), src, re.M | re.S)
    assert m, name
    return m.group(0)


def _node(prelude, expr):
    out = subprocess.run([NODE, "-e", prelude + f"\nconsole.log(JSON.stringify({expr}));"], capture_output=True, text=True,
                         encoding="utf-8", check=True).stdout
    return json.loads(out)


BMC = _src("js/13-bom-center.js")
BRN = _src("js/14-branch.js")
SHM = _src("js/16-supplyhome.js")
OHB = _src("js/15-orderhub.js")
NUM = _fn(BMC, "_bmcNum")


def _box_item(con, count_unit=None):
    """The crab stick as the centre sells it: a box of 20 packs of 500 g (the central catalogue sets the unit and the pack rule)."""
    CE.apply_catalogue(con, CE.parse_report(con, _report(ITEMS[:1])), "ss")
    if count_unit:
        con.execute("UPDATE bc_ingredient SET count_unit=? WHERE code=?", (count_unit, CRAB))
    con.commit()


# ── 3.1 · the unit of entry ──────────────────────────────────────────────────────────────────────────────────────
def test_a_box_item_is_issued_and_counted_in_packs(udon):  # noqa: F811
    _box_item(udon)
    assert udon.execute("SELECT recv_unit, qty_per_recv FROM bc_ingredient WHERE code=?", (CRAB,)).fetchone() == (BOX, 10000.0)
    assert ST.count_units(udon, [CRAB])[CRAB] == ("แพ็ค", 500.0)                    # not กล่อง (20 แพ็ค)
    it = {i["code"]: i for z in ST.sheet(udon, 1, date(2026, 10, 5))["zones"] for i in z["items"]}[CRAB]
    assert (it["count_unit"], it["per_count_unit"]) == ("แพ็ค", 500.0)
    assert it["unit_big"] == {"unit": "กล่อง", "label": BOX, "per": 20, "factor": 10000.0}
    sl = {i["code"]: i for z in IS.sheet(udon, 1, "2026-10-05", now=NOW)["zones"] for i in z["items"]}[CRAB]
    assert (sl["count_unit"], sl["per_count_unit"]) == ("แพ็ค", 500.0)
    # 3 packs typed on the slip are 1,500 g
    IS.save(udon, 1, "2026-10-05", [{"code": CRAB, "qty_used": 3}], "udon", now=NOW)
    assert udon.execute("SELECT unit, qty_used_base FROM bc_issue_line").fetchone() == ("แพ็ค", 1500.0)
    sl = {i["code"]: i for z in IS.sheet(udon, 1, "2026-10-05", now=NOW)["zones"] for i in z["items"]}[CRAB]
    assert sl["used"] == 3.0


def test_the_orders_and_the_variance_read_the_same_unit(udon):  # noqa: F811
    _box_item(udon)
    p = OR.plan(udon, 1, "1", date(2026, 10, 7), as_of=date(2026, 10, 5))
    crab = {r["code"]: r for r in p["rows"]}[CRAB]
    assert (crab["count_unit"], crab["system_unit"]) == ("แพ็ค", "แพ็ค")
    rows = VA.with_count_units(udon, [{"code": CRAB, "base_unit": "กรัม"}])
    assert rows[0]["count_unit"] == "แพ็ค" and rows[0]["per_count_unit"] == 500.0


def test_the_analysts_own_count_unit_is_kept(udon):  # noqa: F811
    _box_item(udon, count_unit=BOX)
    assert ST.count_units(udon, [CRAB])[CRAB] == (BOX, 10000.0)
    it = {i["code"]: i for z in ST.sheet(udon, 1, date(2026, 10, 5))["zones"] for i in z["items"]}[CRAB]
    assert it["count_unit"] == BOX and it["unit_big"] is None                          # the box is the top: no second box
    udon.execute("UPDATE bc_ingredient SET count_unit='กรัม' WHERE code=?", (CRAB,))
    assert ST.count_units(udon, [CRAB])[CRAB] == ("กรัม", 1.0)                         # a smaller unit set on purpose stays too


def test_a_loose_receiving_unit_or_no_pack_size_keeps_the_unit_it_had(udon):  # noqa: F811
    udon.execute("UPDATE bc_ingredient SET recv_unit='กิโลกรัม', qty_per_recv=1000 WHERE code=?", (CRAB,))
    udon.execute("INSERT INTO bc_unit_rule(ingredient_code, unit, factor_to_base) VALUES(?, 'แพ็ค', 5000)", (CRAB,))
    assert ST.count_units(udon, [CRAB])[CRAB] == ("กิโลกรัม", 1000.0)                  # a loose weight is not a sealed unit
    udon.execute("UPDATE bc_ingredient SET recv_unit=NULL, qty_per_recv=NULL WHERE code=?", (CRAB,))
    assert ST.count_units(udon, [CRAB])[CRAB] == ("กรัม", 1.0)


def test_the_smallest_sealed_unit_of_the_ladder(udon):  # noqa: F811
    _box_item(udon)
    udon.execute("INSERT INTO bc_unit_rule(ingredient_code, unit, factor_to_base) VALUES(?, 'ถุง', 250)", (CRAB,))
    conv = Converter(udon)
    assert [(u["unit"], u["factor"], u["source"]) for u in conv.sealed(CRAB)] == [("ถุง", 250.0, "ladder"), ("แพ็ค", 500.0, "ladder"), ("กล่อง", 10000.0, "recv")]
    assert ST.count_units(udon, [CRAB])[CRAB] == ("ถุง", 250.0)
    # a piece that differs in weight and a loose weight are not sealed
    udon.execute("INSERT INTO bc_unit_rule(ingredient_code, unit, factor_to_base, parent_unit, qty_in_parent, conv_kind)"
                 " VALUES(?, 'ตัว', 100, NULL, 100, 'avg')", (CRAB,))
    assert "ตัว" not in [u["unit"] for u in Converter(udon).sealed(CRAB)]


def test_an_item_stocked_in_packs_has_its_own_unit_as_the_smallest(udon):  # noqa: F811
    udon.execute("UPDATE bc_ingredient SET base_unit='แพ็ค', recv_unit='กล่อง (20 แพ็ค)', qty_per_recv=20 WHERE code=?", (CRAB,))
    assert ST.count_units(udon, [CRAB])[CRAB] == ("แพ็ค", 1.0)


def test_a_count_of_two_boxes_sums_in_base_units(udon):  # noqa: F811
    _box_item(udon)
    r = ST.save(udon, 1, [{"code": CRAB, "parts": [{"qty": 2, "unit": BOX}, {"qty": 3, "unit": "แพ็ค"}]}], "udon", count_date="2026-10-01")
    assert r["lines"] == 1 and r["unconverted"] == []
    assert udon.execute("SELECT qty_kitchen, unit, qty_base FROM bc_stock_count_line").fetchone() == (43.0, "แพ็ค", 21500.0)    # 2 × 10,000 + 3 × 500
    it = {i["code"]: i for z in ST.sheet(udon, 1, date(2026, 10, 5))["zones"] for i in z["items"]}[CRAB]
    assert it["draft_qty"] == 43.0
    # one box only, an empty box counts as 0, and no box with a number removes the line
    ST.save(udon, 1, [{"code": CRAB, "parts": [{"qty": 1, "unit": BOX}, {"qty": None, "unit": "แพ็ค"}]}], "udon", count_date="2026-10-01")
    assert udon.execute("SELECT qty_base FROM bc_stock_count_line").fetchone()[0] == 10000.0
    ST.save(udon, 1, [{"code": CRAB, "parts": [{"qty": None, "unit": BOX}, {"qty": "", "unit": "แพ็ค"}]}], "udon", count_date="2026-10-01")
    assert udon.execute("SELECT COUNT(*) FROM bc_stock_count_line").fetchone()[0] == 0
    with pytest.raises(ValueError, match="จำนวนต้องไม่ติดลบ"):
        ST.save(udon, 1, [{"code": CRAB, "parts": [{"qty": -1, "unit": BOX}]}], "udon", count_date="2026-10-01")
    bad = ST.save(udon, 1, [{"code": CRAB, "parts": [{"qty": 1, "unit": "ช้อนโต๊ะ"}]}], "udon", count_date="2026-10-01")
    assert bad["unconverted"] and "ยังไม่ได้ตั้งหน่วย" in bad["unconverted"][0]


def test_a_draft_typed_in_the_old_unit_is_read_in_the_new_one(udon):  # noqa: F811
    _box_item(udon)
    ST.save(udon, 1, [{"code": CRAB, "qty": 2, "unit": BOX}], "udon", count_date="2026-10-01")          # typed as 2 boxes before this change
    it = {i["code"]: i for z in ST.sheet(udon, 1, date(2026, 10, 5))["zones"] for i in z["items"]}[CRAB]
    assert it["count_unit"] == "แพ็ค" and it["draft_qty"] == 40.0


def test_the_count_line_model_takes_parts(client):
    from routers.bom_center import StockLine
    ln = StockLine(code="X", parts=[{"qty": 2, "unit": "กล่อง"}, {"qty": 3, "unit": "แพ็ค"}])
    assert ln.dict()["parts"] == [{"qty": 2.0, "unit": "กล่อง"}, {"qty": 3.0, "unit": "แพ็ค"}]
    assert StockLine(code="X", qty=1).dict()["parts"] is None


@needs_node
def test_figures_on_the_slip_and_the_count_are_tidy():
    pre = NUM + _fn(BRN, "_brnNice") + _fn(BRN, "_brnPh") + _fn(BRN, "_brnUnitW") + _fn(BRN, "_brnSysTxt")
    got = _node(pre, "[_brnSysTxt(9.01,'แพ็ค'), _brnSysTxt(9.46,'แพ็ค 1 กก.'), _brnSysTxt(0.4,'แพ็ค 1 กก.'), _brnSysTxt(0,'แพ็ค'), "
                     "_brnSysTxt(-3,'กล่อง (10 แพ็ค)'), _brnSysTxt(null,'แพ็ค'), _brnNice(0.005), _brnNice(0.34), _brnNice(12.04), _brnPh(0.005), _brnPh(0.34), _brnPh(null)]")
    assert got == ["9 แพ็ค", "9.5 แพ็ค", "น้อยกว่า 1 แพ็ค", "0 แพ็ค", "-3 กล่อง", "", "<0.1", "0.3", "12", "0", "0.3", "0"]


@needs_node
def test_the_two_boxes_split_and_add_up_in_the_entry_unit():
    pre = NUM + _fn(BRN, "_brnNice") + _fn(BRN, "_brnUnitW") + _fn(BRN, "_brnSysTxt") + "let _brnQty = {}, _brnQty2 = {};\n" \
        + _fn(BRN, "_brnSplit") + _fn(BRN, "_brnPair") + _fn(BRN, "_brnTotal2")
    it = "{code:'A', unit_big:{per:20}}"
    got = _node(pre, f"[_brnSplit(43,20), _brnSplit(40,20), _brnSplit(7,20), _brnSplit(0,20), _brnTotal2({it},{{big:'2',small:'3'}}), "
                     f"_brnTotal2({it},{{big:'',small:'3'}}), _brnTotal2({it},{{big:'',small:''}}), _brnTotal2({it},{{big:'1',small:''}})]")
    assert got == [{"big": "2", "small": "3"}, {"big": "2", "small": ""}, {"big": "", "small": "7"}, {"big": "", "small": "0"}, "43", "3", "", "20"]
    src = BRN
    assert "l.parts = [{ qty:" in src and "unit_big" in src and "oninput=\"brnSet2(" in src


# ── 3.2 · the short slip ─────────────────────────────────────────────────────────────────────────────────────────
def test_the_default_slip_rule():
    f = IS.default_list
    assert not f({})
    assert f({"recent": True}) and f({"used_recent": True}) and f({"required": True}) and f({"system_qty": 0.4})
    assert not f({"system_qty": 0}) and not f({"system_qty": -2}) and not f({"system_qty": None})
    assert f({"used": 2}) and f({"waste": 1})                                       # a quantity typed on this slip is never hidden


def test_the_sheet_marks_what_the_short_slip_keeps(udon):  # noqa: F811
    CE.apply_catalogue(udon, CE.parse_report(udon, _report(ITEMS[:2])), "ss")       # crab + salmon are on the centre's list …
    udon.executemany("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv, monitored) VALUES(?,?,?,?,?,?,?)",
                     [("S-901", "น้ำปลา ไม่ได้ใช้", "raw", "กรัม", "แพ็ค 1 กก.", 1000, 0),
                      ("S-902", "เกลือ มีคงเหลือ", "raw", "กรัม", "แพ็ค 1 กก.", 1000, 0),
                      ("S-903", "พริกไทย ต้องนับ", "raw", "กรัม", "แพ็ค 1 กก.", 1000, 1)])
    udon.executemany("INSERT INTO bc_central_item(code, name, active) VALUES(?,?,1)", [("S-901", "x"), ("S-902", "x")])
    ST.save(udon, 1, [{"code": "S-902", "qty": 4, "unit": "แพ็ค 1 กก."}, {"code": "S-901", "qty": 0, "unit": "แพ็ค 1 กก."}],
            "udon", submit=True, count_date="2026-10-01", require_all=False)
    udon.commit()
    sh = IS.sheet(udon, 1, "2026-10-05", now=NOW)
    it = {i["code"]: i for z in sh["zones"] for i in z["items"]}
    assert it[CRAB]["used_recent"] and it[CRAB]["in_default"]                       # the recipes used it in the last 14 days
    assert it["S-902"]["in_default"] and it["S-902"]["system_qty"] == 4.0           # system stock above 0
    assert it["S-903"]["in_default"] and it["S-903"]["required"]                    # ต้องนับ
    assert it["S-901"]["in_default"] is False                                       # not used, no stock, not required
    assert sh["items"] == len(it) and sh["items_default"] == len(it) - 1
    # a quantity typed on the draft keeps the item on the list
    IS.save(udon, 1, "2026-10-05", [{"code": "S-901", "qty_used": 1, "unit": "แพ็ค 1 กก."}], "udon", now=NOW)
    it = {i["code"]: i for z in IS.sheet(udon, 1, "2026-10-05", now=NOW)["zones"] for i in z["items"]}
    assert it["S-901"]["in_default"] is True
    # an item a confirmed slip used within 14 days (the slip's own `recent`) is on the list
    IS.save(udon, 1, "2026-10-07", [{"code": "S-901", "qty_used": 1, "unit": "แพ็ค 1 กก."}], "udon", submit=True, now=at(7, 21))
    it = {i["code"]: i for z in IS.sheet(udon, 1, "2026-10-08", now=NOW)["zones"] for i in z["items"]}
    assert it["S-901"]["recent"] and it["S-901"]["in_default"]


@needs_node
def test_the_toggle_and_the_chips_follow_the_short_list():
    pre = ("let _brnIssAll = false; const _brnIssQty = {};\n"
           "let _brnIss = {items_default: 3, zones: [{zone:'frozen', label:'แช่แข็ง', items:[{code:'A', in_default:true},{code:'B', in_default:false}]},"
           "{zone:'dry', label:'ของแห้ง', items:[{code:'C', in_default:false}]}, {zone:'bar', label:'บาร์', items:[{code:'D', in_default:true},{code:'E'}]}]};\n"
           + _fn(BRN, "_brnIssAny") + _fn(BRN, "_brnIssAllOn") + _fn(BRN, "_brnIssShow") + _fn(BRN, "_brnIssZonesShown"))
    got = _node(pre, "(() => { const z = () => _brnIssZonesShown().map(x => x.zone + ':' + x.items.map(i => i.code).join('')); const a = z();"
                     "_brnIssQty.B = {used:'2'}; const b = z(); delete _brnIssQty.B; _brnIssAll = true; const c = z(); _brnIssAll = false;"
                     "_brnIss.items_default = 0; const d = z(); return [a, b, c, d]; })()")
    assert got[0] == ["frozen:A", "bar:DE"]                                         # the zone with nothing left has no chip; an item without the flag is kept
    assert got[1] == ["frozen:AB", "bar:DE"]                                        # typed = shown
    assert got[2] == ["frozen:AB", "dry:C", "bar:DE"]                               # แสดงทั้งหมด
    assert got[3] == got[2]                                                         # a branch with nothing in the short list sees everything


def test_the_toggle_is_remembered_on_the_device_and_a_search_reads_everything():
    assert "oshinei_brn_iss_all" in BRN and "localStorage.setItem(_BRN_ISS_ALL_KEY" in BRN and "try {" in _fn(BRN, "brnIssToggleAll")
    assert "แสดงทั้งหมด (${all.length} รายการ)" in BRN
    body = _fn(BRN, "brnIssRender")
    assert "(q ? all : zone.items)" in body                                         # a search looks at all items, a zone at its short list
    assert 'id="brnIssFilter"' in _src("tabs/53-branchstock.html")


# ── 2.3 · หน้างาน ────────────────────────────────────────────────────────────────────────────────────────────────
STUBS = ("const fmtLeft = h => `อีก ${Math.round(h)} ชม.`, fmtCutoff = s => 'CUT ' + String(s).slice(0, 10), fmtDayShort = s => 'D ' + String(s).slice(5);\n"
         "const _SHM_WD = []; const _shmRole = () => 'branch';\n")


def _strip_rows(items):
    pre = STUBS + _fn(SHM, "_shmTodayRows")
    return _node(pre, f"_shmTodayRows({{items: {json.dumps(items, ensure_ascii=False)}}})")


def _due(r, status, cutoff, late=False, today=False, delivery="2026-10-17"):
    return {"key": f"order_due_{r}:7", "status": status, "cutoff_at": cutoff, "late": late, "due_today": today, "delivery": delivery,
            "go": {"tab": "branchstock", "pane": "order", "extra": {"round": r}}}


def test_the_strip_names_the_next_order_due_not_a_tick_for_the_one_that_went():
    far = "2999-01-01T13:15:00+07:00"
    items = [{"key": "issue_today:7", "status": "done", "confirmed_hhmm": "19:52", "go": {}},
             _due("2", "done", "2026-10-12T13:15:00+07:00", today=True),
             _due("1", "todo", far, delivery="2026-10-17"), _due("salmon", "todo", "2999-02-02T13:15:00+07:00", delivery="2026-10-19")]
    order = next(r for r in _strip_rows(items) if r["k"] == "order")
    assert order["ok"] is False and "ส่งแล้ว" not in order["text"]
    assert "รอบส่ง D 10-17" in order["text"] and "CUT 2999-01-01" in order["text"] and order["text"].endswith("รออีก 1 รอบ")
    assert order["go"]["extra"]["round"] == "1"                                     # the one whose cutoff comes first
    late = next(r for r in _strip_rows(items + [_due("2b", "todo", "2026-10-12T13:15:00+07:00", late=True)]) if r["k"] == "order")
    assert late["bad"] is True and "เลยเวลาตัดรอบ" in late["text"]                  # a round past its cutoff comes first
    done = next(r for r in _strip_rows([i for i in items if i["status"] == "done"]) if r["k"] == "order")
    assert done["ok"] is True and done["text"] == "ส่งแล้ว"
    none = next(r for r in _strip_rows([items[0]]) if r["k"] == "order")
    assert none["ok"] is True and none["text"] == "ไม่มีรอบที่ต้องส่ง"


@needs_node
def test_the_days_left_behind_are_one_row():
    pre = _fn(BRN, "_brnIssDm") + _fn(BRN, "_brnIssRunTxt") + _fn(BRN, "_brnIssRuns") + _fn(SHM, "_shmCollapseBacklog")
    gaps = [{"key": f"issue_missing_2026-10-{d:02d}:7", "slip_date": f"2026-10-{d:02d}", "status": "todo", "title": f"ย้อนหลัง {d}", "seq": 306,
             "go": {"extra": {"date": f"2026-10-{d:02d}"}}, "count": None} for d in (8, 9, 10, 11)]
    other = {"key": "issue_today:7", "status": "todo", "title": "วันนี้", "seq": 305}
    got = _node(pre, f"_shmCollapseBacklog({json.dumps([other] + gaps, ensure_ascii=False)})")
    assert [g["title"] for g in got] == ["วันนี้", "ใบตัดเบิกค้าง 4 วัน (8–11/10)"]
    assert got[1]["key"] == "issue_missing_2026-10-08:7" and got[1]["go"]["extra"]["date"] == "2026-10-08"   # the button opens the oldest day
    assert _node(pre, "_shmCollapseBacklog([])") == []
    two = _node(pre, f"_shmCollapseBacklog({json.dumps([gaps[0], gaps[3]], ensure_ascii=False)})")
    assert [g["title"] for g in two] == ["ใบตัดเบิกค้าง 2 วัน (8/10, 11/10)"]
    done = dict(gaps[0], status="done")                                              # a day already made is not a debt
    assert _node(pre, f"_shmCollapseBacklog({json.dumps([done], ensure_ascii=False)})")[0]["title"] == "ย้อนหลัง 8"
    assert "_shmCollapseBacklog(items.filter(_shmMine))" in SHM


def test_a_branch_login_has_no_month_to_pick_and_the_warehouse_and_accounting_no_pickers():
    assert 'id="shmMonthWrap"' in _src("tabs/55-supplyhome.html")
    assert "const _SHM_NO_PICKER_ROLES = ['warehouse', 'accounting'];" in SHM
    assert "function _shmNoPickers() { return isPinnedSession() ||" in SHM
    assert "mw.style.display = bare ? 'none' : ''" in SHM
    assert "if (!bare) supplyCtxSet(" in SHM                                          # a hidden picker does not clear the module's shared choice


def test_the_receive_task_names_the_round_and_the_day(udon):  # noqa: F811
    udon.execute("INSERT INTO bc_order(branch_id, cycle_date, round, status, delivery_date) VALUES(1, '2026-10-08', '2', 'processed', '2026-10-10')")
    items = {i["key"].split(":")[0]: i for i in TD.build(udon, "2026-09", None, None, today=date(2026, 10, 12), now=NOW)["items"]}
    assert items["receive_pending"]["detail"] == "รอบ 2 (ส่ง 10/10)"
    assert fmt.round_day_th("salmon", "2026-10-19") == "แซลมอนสด (ส่ง 19/10)" and fmt.round_day_th("1", None) == "รอบ 1"


# ── 1.9 · what the confirm says before an order goes ─────────────────────────────────────────────────────────────
def _order(con, cycle, delivery, status, lines):
    cur = con.execute("INSERT INTO bc_order(branch_id, cycle_date, round, status, delivery_date) VALUES(1, ?, '1', ?, ?)", (cycle, status, delivery))
    for code, units in lines:
        con.execute("INSERT INTO bc_order_line(order_id, ingredient_code, suggested_units, confirmed_units, order_unit) VALUES(?,?,?,?,'pack')",
                    (cur.lastrowid, code, units, units))
    con.commit()
    return cur.lastrowid


def test_the_plan_carries_the_previous_order_of_the_same_round(udon):  # noqa: F811
    _box_item(udon)                                                                  # crab: ฿1,500 a box of 10,000 g → ฿0.15 a gram
    _order(udon, "2026-09-30", "2026-10-02", "received", [(CRAB, 2)])                # two rounds back
    _order(udon, "2026-10-02", "2026-10-04", "processed", [(CRAB, 3)])               # the last order sent
    _order(udon, "2026-10-03", "2026-10-05", "suggested", [(CRAB, 9)])               # a draft is no earlier order
    p = OR.plan(udon, 1, "1", date(2026, 10, 7), as_of=date(2026, 10, 5))
    pv = p["previous"]
    assert pv["round"] == "1" and pv["round_th"] == "รอบ 1" and pv["delivery"] == "2026-10-04" and pv["delivery_th"] == "4/10/2026"
    assert pv["n"] == 1 and pv["value"] == pytest.approx(4500.0) and pv["lines"] == [
        {"code": CRAB, "name": "ปูอัด", "units": 3.0, "unit_word": "กล่อง", "value": pytest.approx(4500.0)}]
    assert OR.previous(udon, 1, "2", "2026-10-05") is None and OR.previous(udon, 1, "1", "2026-09-30") is None


@needs_node
def test_the_confirm_says_what_changed_most():
    pre = NUM + _fn(BRN, "_brnNm") + _fn(BRN, "_brnUnitW") + _fn(BRN, "_brnUnitLabel") + _fn(BRN, "_brnOrdSummary")
    rows = [{"code": c, "name": n, "short_name": s, "order_unit": "pack", "recv_unit": "กล่อง (10 แพ็ค)", "qty_per_recv": 1000, "price_per_base": 0.1}
            for c, n, s in (("A", "กุ้ง เต็ม", "กุ้ง"), ("B", "ปลา", None), ("C", "หมึก", None), ("D", "ไข่", None), ("E", "งา", None))]
    prev = {"value": 1000, "lines": [{"code": "A", "name": "กุ้ง เต็ม", "units": 10, "unit_word": "กล่อง", "value": 1000},
                                      {"code": "Z", "name": "ปู", "units": 5, "unit_word": "กล่อง", "value": 500}]}
    qty = {"A": "4", "B": "20", "C": "1", "D": "0", "E": ""}
    got = _node(pre, f"_brnOrdSummary({{rows: {json.dumps(rows, ensure_ascii=False)}, previous: {json.dumps(prev, ensure_ascii=False)}}}, {json.dumps(qty)})")
    assert got["n"] == 3 and got["value"] == pytest.approx(2500.0)
    assert got["head"] == "3 รายการ · ฿2,500 · รอบก่อน ฿1,000 (+150%)"
    assert [t["name"] for t in got["top"]] == ["ปลา", "กุ้ง", "ปู"]                       # |Δ฿| 2000, 600, 500; the short name is used; a line dropped counts
    assert got["text"].splitlines()[2] == "• ปลา: 0 → 20 กล่อง (+฿2,000)" and got["text"].splitlines()[4] == "• ปู: 5 → 0 กล่อง (−฿500)"
    nope = _node(pre, f"_brnOrdSummary({{rows: {json.dumps(rows, ensure_ascii=False)}, previous: null}}, {json.dumps(qty)})")
    assert nope["text"] == "3 รายการ · ฿2,500" and nope["top"] == []
    assert "_brnOrdSummary(p, _brnOrdQty).text" in BRN


# ── 4.1 · the short name ─────────────────────────────────────────────────────────────────────────────────────────
@pytest.fixture
def reg(tmp_path):
    import sqlite3
    import conftest as _cf
    from bom_center import schema as S
    p = tmp_path / "reg.db"
    _cf._build_db(str(p))
    c = sqlite3.connect(str(p))
    S.ensure(c, commit=True)
    c.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv) VALUES('02-01-001','ปลาแซลมอนเทร้านอร์เวย์ 4-5 KG/PCS SalmonFresh (KG)','raw','กรัม','กิโลกรัม',1000)")
    c.commit()
    yield c
    c.close()


def _short(c, code="02-01-001"):
    return c.execute("SELECT short_name FROM bc_ingredient WHERE code=?", (code,)).fetchone()[0]


def test_the_migration_is_append_only_and_idempotent(reg):
    from bom_center import schema as S
    assert "short_name" in {r[1] for r in reg.execute("PRAGMA table_info(bc_ingredient)")}
    S.ensure(reg, commit=True)
    S.ensure(reg, commit=True)
    assert ("bc_ingredient", "short_name", "TEXT") in S.MIGRATIONS and _short(reg) is None             # nothing is seeded


def test_a_short_name_is_set_and_cleared_on_the_screen(reg):
    M.update_rows(reg, [{"code": "02-01-001", "short_name": "  แซลมอน   สด  "}], "analyst1")
    assert _short(reg) == "แซลมอน สด"
    assert reg.execute("SELECT name_th FROM bc_ingredient WHERE code='02-01-001'").fetchone()[0].startswith("ปลาแซลมอน")   # the full name stands
    M.update_rows(reg, [{"code": "02-01-001", "short_name": "ก" * 60}], "analyst1")
    assert _short(reg) == "ก" * 40                                                                       # one line, 40 characters
    M.update_rows(reg, [{"code": "02-01-001", "note": "x"}], "analyst1")                                 # an edit of another field leaves it
    assert _short(reg) == "ก" * 40
    M.update_rows(reg, [{"code": "02-01-001", "short_name": ""}], "analyst1")                            # an empty box takes it away
    assert _short(reg) is None


def _read_file(rows):
    """The register file as the importer reads it (the real _read_master_file, on an uploaded workbook)."""
    from openpyxl import Workbook
    from starlette.datastructures import UploadFile
    from routers.bom_center import _read_master_file
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return asyncio.run(_read_master_file(UploadFile(buf, filename="m.xlsx")))


def test_the_register_file_may_carry_the_short_name_or_not(reg):
    head = ["รหัส", "ชื่อวัตถุดิบ", "หน่วยฐาน", "หน่วยรับ", "ขนาดบรรจุ"]
    rows, _sheet, cols = _read_file([head + ["ชื่อสั้น"], ["02-01-001", "ปลาแซลมอน (ชื่อเต็ม)", "กรัม", "กิโลกรัม", 1000, "แซลมอน"],
                                     ["S-777", "เกลือ", "กรัม", "ถุง", 500, None]])
    assert "short_name" in cols and rows[0]["short_name"] == "แซลมอน"
    st = M.import_rows(reg, rows, "analyst1")
    assert _short(reg) == "แซลมอน" and _short(reg, "S-777") is None and st["updated"] == 1 and st["inserted"] == 1
    # an old file with no such column: nothing changes, and a blank cell does not clear what the screen set
    rows, _s, cols = _read_file([head, ["02-01-001", "ปลาแซลมอน (ชื่อเต็ม)", "กรัม", "กิโลกรัม", 1000]])
    assert "short_name" not in cols
    M.import_rows(reg, rows, "analyst1")
    assert _short(reg) == "แซลมอน"
    rows, _s, _c = _read_file([head + ["ชื่อสั้น"], ["02-01-001", "ปลาแซลมอน (ชื่อเต็ม)", "กรัม", "กิโลกรัม", 1000, None]])
    M.import_rows(reg, rows, "analyst1")
    assert _short(reg) == "แซลมอน"
    rows, _s, _c = _read_file([head + ["ชื่อสั้น"], ["02-01-001", "ปลาแซลมอน (ชื่อเต็ม)", "กรัม", "กิโลกรัม", 1000, "แซลมอนใหม่"]])
    M.import_rows(reg, rows, "analyst1")
    assert _short(reg) == "แซลมอนใหม่"


def test_the_register_template_has_the_column_and_the_list_returns_it(client):
    from openpyxl import load_workbook
    r = client.get("/api/bom-center/ingredients/template.xlsx")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    heads = [c.value for row in ws.iter_rows(min_row=1, max_row=3) for c in row if c.value]
    assert "ชื่อสั้น" in heads and "หน่วยที่นับ" in heads
    from routers.bom_center import IngredientEdit
    assert IngredientEdit(code="X", short_name="กุ้ง").dict(exclude_unset=True) == {"code": "X", "short_name": "กุ้ง"}


def test_the_branch_payloads_carry_the_short_name_beside_the_full_one(udon):  # noqa: F811
    udon.execute("UPDATE bc_ingredient SET short_name='ปูอัด สั้น' WHERE code=?", (CRAB,))
    payload = {"rows": [{"code": CRAB, "name": "ปูอัดซาชิมิ Kani Kamoboko"}, {"code": "02-01-001", "name": "แซลมอน"}],
               "x": {"lines": [{"code": CRAB, "name": "ปูอัดซาชิมิ Kani Kamoboko"}]}}
    out = fmt.add_short(udon, payload)
    assert out["rows"][0] == {"code": CRAB, "name": "ปูอัดซาชิมิ Kani Kamoboko", "short_name": "ปูอัด สั้น"}
    assert "short_name" not in out["rows"][1] and out["x"]["lines"][0]["short_name"] == "ปูอัด สั้น"
    sh = fmt.add_short(udon, ST.sheet(udon, 1, date(2026, 10, 5)))
    assert {i["code"]: i for z in sh["zones"] for i in z["items"]}[CRAB]["short_name"] == "ปูอัด สั้น"
    # a message to the floor says it too, and the full name is back when the short one goes
    assert fmt.item(udon, CRAB) == "ปูอัด สั้น" and fmt.item(Converter(udon), CRAB) == "ปูอัด สั้น"
    udon.execute("UPDATE bc_ingredient SET short_name=NULL WHERE code=?", (CRAB,))
    assert fmt.item(udon, CRAB) != "ปูอัด สั้น" and fmt.add_short(udon, {"code": CRAB, "name": "x"}) == {"code": CRAB, "name": "x"}


def test_every_branch_screen_writes_the_short_name_and_hides_the_code_behind_the_source_button():
    for fn in ("_brnNm", "_brnSrc", "brnSrc"):
        assert f"function {fn}(" in BRN
    assert BRN.count("_brnNm(") >= 40 and BRN.count("_brnSrc(") >= 8
    for rel, needle in (("backend/routers/bom_center.py", "_short(con, _st.sheet(con, bid, mode=mode))"), ("backend/routers/bom_center.py", "_short(con, _is.sheet("),
                        ("backend/routers/bom_center.py", "_short(con, out)"), ("backend/routers/bom_center.py", "_short(con, _va.period("),
                        ("backend/routers/bom_center.py", "_short(con, _cl.detail("), ("backend/routers/bom_center.py", "_short(con, o)")):
        assert needle in _src(rel), needle
    assert "inp(r, 'short_name', 130)" in BMC and 'title="ชื่อที่หน้าสาขาแสดง' in _src("tabs/52-bomcenter.html")
    # no row of a branch list prints the code in the open any more
    assert "${esc(it.code)}${sys}${avg}" not in BRN and "${c} · ส่งจริง" not in BRN


# ── 4.2 · no record number on a worker's screen ──────────────────────────────────────────────────────────────────
def test_no_branch_screen_writes_a_record_number():
    for rel in ("js/14-branch.js", "js/16-supplyhome.js", "tabs/53-branchstock.html", "tabs/55-supplyhome.html"):
        src = _src(rel)
        assert not re.search(r"ORDER-|CLAIM-|CENTRAL-", src), rel
        assert not re.search(r"#\$\{|#\{|ใบสั่งซื้อ #\d|เคลม #\d", src), rel
    assert "_brnRoundDay(" in BRN and BRN.count("_brnRoundDay(") >= 5
    todo = _src("backend/bom_center/todo.py")
    assert not re.search(r"f\"[^\"\n]*#\{", todo)                                        # no "#9" in a sentence the floor reads
    claims = _src("backend/bom_center/claims.py")
    assert "#{claim_id}" not in claims and "{lbl}" in claims


@needs_node
def test_an_order_is_named_by_its_round_and_day():
    pre = _fn(BRN, "_brnIssDm") + "const _BRN_ROUNDS = [['1','รอบ 1'],['2','รอบ 2'],['salmon','แซลมอนสด']];\n" + _fn(BRN, "_brnRoundLabel") + _fn(BRN, "_brnRoundDay")
    assert _node(pre, "[_brnRoundDay({round:'2', delivery:'2026-10-10'}), _brnRoundDay({round:'salmon', cycle_date:'2026-10-17'}), _brnRoundDay({round:'1'})]") \
        == ["รอบ 2 (ส่ง 10/10)", "แซลมอนสด (ส่ง 17/10)", "รอบ 1"]


# ── 4.4 · the buttons of a question ──────────────────────────────────────────────────────────────────────────────
def test_the_dialog_buttons_are_44_px_on_a_phone():
    core = _src("js/01-core.js")
    assert "matchMedia('(max-width: 700px), (pointer: coarse)')" in core and core.count("min-height:${btnH}px") == 2
    assert re.search(r"btnH = \(.*\) \? 44 : 0", core)


# ── 4.5 · pickers that do not apply ──────────────────────────────────────────────────────────────────────────────
@needs_node
def test_the_unsteady_items_are_a_link_to_the_five_of_the_most_baht():
    pre = NUM + "const esc = s => String(s == null ? '' : s);\nconst box = {innerHTML: ''}; const document = {getElementById: () => box};\n" \
        + "const items = [" + ",".join("{code:'I%d', name:'N%d', steady:%s, value_per_100k:%d, swing_pct:%d}" % (i, i, "false" if i % 4 else "true", i * 10, 20 + i)
                                        for i in range(1, 15)) + "];\n" \
        + "const _ohbPs = {branches: [{branch_id: 7, branch: 'ขอนแก่น', items}]};\n" + _fn(OHB, "ohbPsUnsteady")
    got = _node(pre, "(() => { ohbPsUnsteady(7); return box.innerHTML; })()")
    names = re.findall(r"<td class=\"pr-2 py-1\" title=\"[^\"]*\">(N\d+)</td>", got)
    assert names == ["N14", "N13", "N11", "N10", "N9"]                                   # steady ones (every 4th) left out; the most baht first; five of them
    assert "ใช้ไม่สม่ำเสมอ 11 รายการ" in got and "ohbPsPick(7)" in got
    assert "ohbPsUnsteady(${b.branch_id})" in OHB and '<div id="ohbPsTop"' in OHB


# ── 4.6 · a gap in words ─────────────────────────────────────────────────────────────────────────────────────────
GAPS = [(-16197, -8.1, "use"), (4200, 12.34, "issue"), (-1500, None, "count"), (None, 5.3, "use"), (0.2, 0.01, "use"), (None, None, "issue"),
        (320, 0, "use"), (-0.4, -9.9, "issue"), (88000, 31.7, "count")]


def test_a_gap_is_written_in_words():
    assert fmt.gap_th(-16197, -8.1) == "ใช้น้อยกว่าสูตร ฿16,197 (8.1%)"
    assert fmt.gap_th(16197, 8.1) == "ใช้มากกว่าสูตร ฿16,197 (8.1%)"
    assert fmt.gap_th(4200, 12.34, "issue") == "เบิกมากกว่าสูตร ฿4,200 (12%)"
    assert fmt.gap_th(-1500, None, "count") == "นับได้น้อยกว่าระบบ ฿1,500" and fmt.gap_th(1500, None, "count") == "นับได้มากกว่าระบบ ฿1,500"
    assert fmt.gap_th(None, 5.3) == "ใช้มากกว่าสูตร 5.3%"
    assert fmt.gap_th(0.2, 0.01) == "ใช้ตามสูตร" and fmt.gap_th(None, None, "issue") == "เบิกตามสูตร"
    assert fmt.gap_th(-0.4, -9.9, "issue") == "เบิกน้อยกว่าสูตร 9.9%"                              # ฿0 is not written
    assert "−" not in fmt.gap_th(-16197, -8.1) and "-" not in fmt.gap_th(-16197, -8.1)           # the sign is in the words


@needs_node
def test_the_screen_writes_the_gap_the_way_the_server_does():
    pre = NUM + _fn(BMC, "bmcGapDir") + _const(BMC, "_BMC_GAP_WORDS") + _const(BMC, "_BMC_GAP_SHORT") + _fn(BMC, "bmcGapParts") + _fn(BMC, "bmcGapWords")
    cases = json.dumps([[v, p, k] for v, p, k in GAPS])
    got = _node(pre, f"{cases}.map(([v, p, k]) => bmcGapWords(v, p, k))")
    assert got == [fmt.gap_th(v, p, k) for v, p, k in GAPS]
    ink = _node(pre, "[bmcGapParts(-5, null, 'use').ink, bmcGapParts(5, null, 'use').ink, bmcGapParts(5, null, 'count').ink, bmcGapParts(-5, null, 'count').ink, bmcGapParts(0, 0, 'use').ink]")
    assert ink == ["var(--amber-ink)", "var(--bad-ink)", "var(--amber-ink)", "var(--bad-ink)", "var(--good-ink)"]       # theme inks, by direction
    assert _node(pre, "bmcGapWords(-9, -8.4, 'issue', true)") == "น้อยกว่า ฿9 (8.4%)"


def test_the_branch_and_ops_read_the_gap_in_words_and_excel_keeps_the_signs():
    for needle in ("bmcGapParts(v.gap_value, pct, 'use')", "bmcGapHtml(b.gap_value, b.gap_pct, 'use')", "bmcGapHtml(r.var_value, r.var_pct, 'issue')",
                   "bmcGapParts(r.gap, r.gap_pct, 'use')"):
        assert needle in BRN, needle
    for needle in ("bmcGapHtml(t.var_value, t.var_pct, 'issue')", "bmcGapHtml(r.var_value, r.var_pct, 'issue')", "bmcGapHtml(t.diff_value, null, 'count')",
                   "bmcGapParts(x.diff, x.diff_pct, 'count')", "bmcGapHtml(t.gap_value, t.gap_pct, 'use')", "bmcGapHtml(b.gap_value, b.gap_pct, 'use')"):
        assert needle in OHB, needle
    assert "ผลต่าง <b" not in BRN and "th('ผลต่าง ฿'" not in OHB and "th('ผลต่าง %'" not in OHB          # no signed column left on the screens
    router = _src("backend/routers/bom_center.py")
    assert '"ผลต่าง (มูลค่า)"' in router and '"ผลต่าง %"' in router                                     # the Excel files keep the signed numbers
