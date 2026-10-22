# -*- coding: utf-8 -*-
"""Excel forms for the branch pages (PR-SI): the daily issue slip, the order and the receipt can each be keyed in the app OR downloaded
as a template, filled in Excel and imported back. The branch chooses; the result is the same, because an import goes through the very
code the app's save goes through (issues.save, orders.save_flow, orders.receive) - this module only reads the file and reports.

  * a template has a note row (what it is, the branch's id, the date / round), a header row and one row per item; the columns a branch
    must not edit are grey and locked (sheet protection without a password), the reason columns are drop-downs of the same labels the app
    shows (a second sheet lists them);
  * an import first answers as a dry run (a table per date / order, nothing written), then writes on the branch's say-so;
  * a dry run runs the real path inside a savepoint that is rolled back - what it reports is what the real call would do.

Every Thai string of the forms is in TXT below, so the wording can be changed in one place."""
import io
import re
from contextlib import contextmanager
from datetime import date, datetime, timedelta

from . import claims as _cl
from . import issues as _is
from . import orders as _or
from .fmt import add_short as _short
from .fmt import item as _item
from .master import audit

TXT = {
    "bad_file": "อ่านไฟล์ไม่ได้ · ใช้ไฟล์ .xlsx จากปุ่ม “ดาวน์โหลดแบบฟอร์ม”",
    "no_note": "ไม่พบบรรทัดหมายเหตุของแบบฟอร์ม · ใช้ไฟล์จากปุ่ม “ดาวน์โหลดแบบฟอร์ม”",
    "no_head": "ไม่พบหัวตาราง (ต้องมีคอลัมน์ “รหัส”) · ใช้ไฟล์จากปุ่ม “ดาวน์โหลดแบบฟอร์ม”",
    "wrong_kind": "ไฟล์นี้ไม่ใช่แบบฟอร์ม{what} · ใช้ไฟล์จากปุ่ม “ดาวน์โหลดแบบฟอร์ม” ของหน้านี้",
    "other_branch": "ไฟล์นี้เป็นของสาขาอื่น (รหัสสาขา {file}) จึงนำเข้าที่สาขานี้ไม่ได้",
    "days_range": "จำนวนวันต้องอยู่ระหว่าง 1 ถึง {max}",
    "note_tail": "กรอกเฉพาะช่องเบิกใช้ ของเสีย และสาเหตุ · ห้ามแก้รหัส",
    "slip_locked": "ใบตัดเบิกวันที่ {day} ยืนยันแล้ว หากต้องแก้ไข แจ้งประสานงานขาย",
    "slip_note": "นำเข้าจากไฟล์ {name}",
    "bad_date": "วันที่ไม่ถูกต้อง (ต้องเป็น YYYY-MM-DD)",
    "dup_code": "{name}: รหัสซ้ำในไฟล์",
    "unknown_code": "ไม่พบรหัส {code} ในรายการของสาขานี้",
    "bad_reason": "{name}: สาเหตุ “{value}” ไม่อยู่ในรายการ",
    "ord_note_tail": "กรอกเฉพาะช่องสั่ง เหตุผล และหมายเหตุ · ห้ามแก้รหัส",
    "ord_other_round": "ไฟล์นี้เป็นของรอบหรือวันส่งอื่น ({file}) ไม่ตรงกับที่เลือก ({want})",
    "ord_unknown": "ไม่พบรหัส {code} ในรายการสั่งของรอบนี้",
    "ord_errors": "ไฟล์มี {n} แถวที่ต้องแก้ไข: {first}",          # also the slip and the receipt
    "ord_none": "ไม่พบจำนวนที่สั่งในไฟล์",
    "rcv_note_tail": "กรอกเฉพาะช่องรับจริง สาเหตุส่วนต่าง และหมายเหตุ · ห้ามแก้รหัส",
    "rcv_other_order": "ไฟล์นี้เป็นของใบสั่งซื้อเลขที่ {file} ไม่ตรงกับใบที่เลือก ({want})",
    "rcv_unknown": "ไม่พบรหัส {code} ในใบสั่งซื้อนี้",
    "rcv_state": "รับสินค้าได้เมื่อส่งใบสั่งซื้อแล้วเท่านั้น (สถานะปัจจุบัน: {st})",
    "rcv_weight_note": "รายการที่ชั่งน้ำหนักและนับจำนวนตัว (เช่น ปลาทั้งตัว) รับในแอป",
    "rcv_weight": "{name}: รับรายการนี้ในแอป (ต้องชั่งน้ำหนักและนับจำนวน{unit}) จึงนำเข้าใบรับสินค้านี้จากไฟล์ไม่ได้",
    "no_order": "ไม่พบใบสั่งซื้อของสาขานี้",
}

SLIP_SHEET = "ใบตัดเบิก"
SLIP_REASONS_SHEET = "สาเหตุของเสีย"
SLIP_HEAD_FIX = ["วันที่", "รหัส", "วัตถุดิบ", "หน่วย"]
SLIP_HEAD_EDIT = ["เบิกใช้", "ของเสีย", "สาเหตุของเสีย", "หมายเหตุ"]
MAX_DAYS = 7
ORDER_SHEET = "ใบสั่งซื้อ"
ORDER_REASONS_SHEET = "เหตุผลการสั่ง"
ORDER_HEAD = ["รหัส", "วัตถุดิบ", "หน่วย", "คงเหลือตามระบบ", "ยอดแนะนำ", "สั่งครั้งก่อน", "สั่ง", "เหตุผล", "หมายเหตุ"]
RECEIVE_SHEET = "รับสินค้า"
RECEIVE_REASONS_SHEET = "สาเหตุส่วนต่าง"
RECEIVE_HEAD = ["รหัส", "วัตถุดิบ", "หน่วย", "ส่งจริง", "รับจริง", "สาเหตุส่วนต่าง", "หมายเหตุ"]


# ── workbook pieces ──────────────────────────────────────────────────────────
def _make(sheet: str, note: str, headers: list, widths: list):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append([note])
    ws["A1"].font = Font(italic=True, color="5F6C80")
    ws.append(headers)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A6B")
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A3"
    return wb, ws


def _add(ws, values: list, editable: set) -> None:
    """One row: the columns in `editable` (0-based) are open, the rest grey and locked."""
    from openpyxl.styles import PatternFill, Protection
    ws.append(values)
    r = ws.max_row
    grey = PatternFill("solid", fgColor="E5E7EB")
    for i in range(len(values)):
        c = ws.cell(r, i + 1)
        if i in editable:
            c.protection = Protection(locked=False)
        else:
            c.fill = grey


def _finish(wb, ws, reasons_title: str, pairs: list, col: str, label_head: str) -> None:
    """Lock the sheet (no password), add the reasons sheet and the drop-down of its labels on column `col`."""
    from openpyxl.worksheet.datavalidation import DataValidation
    rs = wb.create_sheet(reasons_title)
    rs.append([label_head, "รหัส"])
    for code, label in pairs:
        rs.append([label, code])
    rs.column_dimensions["A"].width = 44
    rs.column_dimensions["B"].width = 14
    if ws.max_row >= 3 and pairs:
        dv = DataValidation(type="list", formula1=f"'{reasons_title}'!$A$2:$A${len(pairs) + 1}", allow_blank=True,
                            showErrorMessage=True, errorTitle="เลือกจากรายการ", error="เลือกสาเหตุจากรายการ")
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}{ws.max_row}")
    for sh in (ws, rs):
        sh.protection.sheet = True
        sh.protection.formatColumns = False
        sh.protection.formatRows = False
        sh.protection.autoFilter = False
        sh.protection.sort = False


def _blank(v) -> bool:
    return v is None or str(v).strip() == ""


def _num(v):
    """A cell as the app would receive the number: numbers as they are, text with thousands commas removed, blank None."""
    if _blank(v):
        return None
    if isinstance(v, str):
        return v.strip().replace(",", "")
    return v


def _table(raw: bytes, sheet: str) -> tuple:
    """(note text, header names, [(excel row, cells)]) of the form's sheet."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        raise ValueError(TXT["bad_file"])
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(grid[:10]) if "รหัส" in [str(c or "").strip() for c in r]), None)
    if hi is None:
        raise ValueError(TXT["no_head"])
    note = " ".join(str(c) for r in grid[:hi] for c in r if c)
    if not note:
        raise ValueError(TXT["no_note"])
    head = [str(c or "").strip() for c in grid[hi]]
    return note, head, [(i + 1, r) for i, r in enumerate(grid) if i > hi]


def _cell(head: list, row: list, name: str):
    return row[head.index(name)] if name in head and head.index(name) < len(row) else None


def _file_branch(note: str):
    m = re.search(r"รหัสสาขา\s*(\d+)", note)
    return int(m.group(1)) if m else None


def _reason_code(value, pairs: list):
    """The reason code for what a cell holds - the label the drop-down offers, or the code itself. Blank: None; unknown: KeyError."""
    s = str(value or "").strip()
    if not s:
        return None
    low = s.lower()
    for code, label in pairs:
        if low == label.lower() or low == code.lower():
            return code
    raise KeyError(s)


@contextmanager
def trial(con, keep: bool):
    """Run the real write path in a savepoint. keep=False: whatever it wrote is rolled back (a dry run); an error always rolls back."""
    con.execute("SAVEPOINT bf")
    try:
        yield
    except BaseException:
        con.execute("ROLLBACK TO bf")
        con.execute("RELEASE bf")
        raise
    if keep:
        con.execute("RELEASE bf")
    else:
        con.execute("ROLLBACK TO bf")
        con.execute("RELEASE bf")


def _kind(msg: str) -> str:
    for key, kind in (("ตัวเลข", "number"), ("ติดลบ", "negative"), ("เลือกสาเหตุ", "reason"), ("อื่น ๆ", "note"),
                      ("แปลง", "unit")):
        if key in msg:
            return kind
    return "line"


def _bname(con, branch_id: int) -> str:
    return con.execute("SELECT name FROM branches WHERE branch_id=?", (branch_id,)).fetchone()[0]


def _wrong_branch(parsed: dict) -> ValueError:
    return ValueError(TXT["other_branch"].format(file=parsed["branch_id"] if parsed["branch_id"] is not None else "-"))


def _num_or_zero(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ── 1 · the issue slip ───────────────────────────────────────────────────────
def slip_template(con, branch_id: int, start: date, days: int = 1, now: datetime = None):
    """-> (workbook, branch name, first date, last date). `days` consecutive dates from `start`, one block each."""
    from .forecast import HISTORY_WEEKS
    if not 1 <= days <= MAX_DAYS:
        raise ValueError(TXT["days_range"].format(max=MAX_DAYS))
    today = _is._day(now)
    dates = [start + timedelta(days=k) for k in range(days)]
    for d in dates:
        prob = _is.day_problem(con, branch_id, d, today)
        if prob:
            raise ValueError(prob[1])
    name = _bname(con, branch_id)
    span = dates[0].isoformat() + (f" ถึง {dates[-1].isoformat()}" if days > 1 else "")
    note = f"ใบตัดเบิก สาขา{name} · รหัสสาขา {branch_id} · วันที่ {span} · {TXT['note_tail']}"
    head = SLIP_HEAD_FIX + [f"ปกติ ({HISTORY_WEEKS} สัปดาห์)", "คงเหลือตามระบบ"] + SLIP_HEAD_EDIT
    wb, ws = _make(SLIP_SHEET, note, head, [12, 14, 44, 16, 16, 16, 12, 12, 30, 24])
    edit = {6, 7, 8, 9}
    for d in dates:
        sh = _short(con, _is.sheet(con, branch_id, d, now))
        for z in sh["zones"]:
            for it in z["items"]:
                _add(ws, [d.isoformat(), it["code"], it.get("short_name") or it["name"], it["count_unit"],
                          _q(it.get("normal_qty") or None, 2), _q(it.get("system_qty"), 2),
                          it.get("used") or None, it.get("waste") or None,
                          _is.WASTE_REASON_TH.get(it.get("waste_reason")) if it.get("waste_reason") else None,
                          it.get("line_note")], edit)
    _finish(wb, ws, SLIP_REASONS_SHEET, [(c, t) for c, t in _is.WASTE_REASONS], "I", "สาเหตุของเสีย")
    return wb, name, dates[0], dates[-1]


def read_slip_file(raw: bytes) -> dict:
    """{branch_id, rows: [(row, date, line dict)], ignored, file_errors} from a slip template. A row with both quantities blank is ignored."""
    note, head, body = _table(raw, SLIP_SHEET)
    if not note.startswith("ใบตัดเบิก "):
        raise ValueError(TXT["wrong_kind"].format(what="ใบตัดเบิก"))
    pairs = [(c, t) for c, t in _is.WASTE_REASONS]
    rows, ignored, errs = [], 0, []
    for rno, r in body:
        code = str(_cell(head, r, "รหัส") or "").strip()
        if not code:
            continue
        used, waste = _cell(head, r, "เบิกใช้"), _cell(head, r, "ของเสีย")
        if _blank(used) and _blank(waste):
            ignored += 1
            continue
        d = _cell(head, r, "วันที่")
        if isinstance(d, datetime):
            iso = d.date().isoformat()
        elif isinstance(d, date):
            iso = d.isoformat()
        else:
            iso = str(d or "").strip()[:10]
            try:
                date.fromisoformat(iso)
            except ValueError:
                errs.append({"row": rno, "code": "date", "message": TXT["bad_date"]})
                continue
        bad = None
        try:
            reason = _reason_code(_cell(head, r, "สาเหตุของเสีย"), pairs)
        except KeyError as e:                       # the row stays with its date; the import reports the reason as that row's error
            reason, bad = None, e.args[0]
        rows.append((rno, iso, {"code": code, "qty_used": _num(used), "qty_waste": _num(waste), "waste_reason": reason,
                                "note": str(_cell(head, r, "หมายเหตุ") or "").strip() or None, "_bad_reason": bad}))
    return {"branch_id": _file_branch(note), "rows": rows, "ignored": ignored, "file_errors": errs}


def import_slips(con, branch_id: int, parsed: dict, who: str, filename: str, dry_run: bool = True, submit: bool = False,
                 now: datetime = None) -> dict:
    """Check the file date by date and, unless `dry_run`, save each date that has no error (a draft, or confirmed with `submit`).
    A date with an error, a locked (confirmed) date, a future / too old / posted date is refused; the others go through. A confirmed slip is
    never appended to: `append` ADDS to the lines of the day, a file carries the whole day, so adding would count it twice."""
    if parsed["branch_id"] != branch_id:
        raise _wrong_branch(parsed)
    if parsed["file_errors"] and not dry_run:             # a row with no readable date belongs to no day: the file is fixed first
        raise ValueError(TXT["ord_errors"].format(n=len(parsed["file_errors"]), first=parsed["file_errors"][0]["message"]))
    from .stock import template_rows
    today = _is._day(now)
    listed = {r["code"] for r in template_rows(con, branch_id)}
    by_day = {}
    for rno, iso, ln in parsed["rows"]:
        by_day.setdefault(iso, []).append((rno, ln))
    note = TXT["slip_note"].format(name=(filename or "").replace("\\", "/").rsplit("/", 1)[-1][:120])
    out = []
    for iso in sorted(by_day):
        rows = by_day[iso]
        d = date.fromisoformat(iso)
        res = {"date": iso, "state": "new", "lines": len(rows),
               "used_lines": sum(1 for _r, ln in rows if _num_or_zero(ln["qty_used"]) > 0),
               "waste_lines": sum(1 for _r, ln in rows if _num_or_zero(ln["qty_waste"]) > 0),
               "errors": [], "result": None, "message": None, "late": 1 if today > d else 0}
        prob = _is.day_problem(con, branch_id, d, today, check_posted=True)
        row = _is._row(con, branch_id, iso)
        if prob:
            res["state"] = prob[0]
            res["errors"].append({"row": None, "code": prob[0], "message": prob[1]})
        elif row and row["status"] == "confirmed":
            res["state"] = "confirmed"
            res["errors"].append({"row": None, "code": "confirmed", "message": TXT["slip_locked"].format(day=_is._fmt_day(iso))})
        else:
            res["state"] = "draft" if row else "new"
            seen = set()
            for rno, ln in rows:
                c = ln["code"]
                if ln.get("_bad_reason"):
                    res["errors"].append({"row": rno, "code": "reason",
                                          "message": TXT["bad_reason"].format(name=_item(con, c), value=ln["_bad_reason"])})
                if c not in listed:
                    res["errors"].append({"row": rno, "code": "code", "message": TXT["unknown_code"].format(code=c)})
                elif c in seen:
                    res["errors"].append({"row": rno, "code": "dup", "message": TXT["dup_code"].format(name=_item(con, c))})
                seen.add(c)
            if not res["errors"]:
                try:
                    _clean, bad = _is._clean_lines(con, [ln for _r, ln in rows])
                    for rno, ln in rows:
                        hit = next((b for b in bad if _item(con, ln["code"]) in b or ln["code"] in b), None)
                        if hit:
                            res["errors"].append({"row": rno, "code": "unit", "message": hit})
                except ValueError:                      # find every bad row, not only the first
                    for rno, ln in rows:
                        try:
                            _is._clean_lines(con, [ln])
                        except ValueError as e:
                            res["errors"].append({"row": rno, "code": _kind(str(e)), "message": str(e)})
        if not res["errors"]:
            try:
                with trial(con, keep=not dry_run):
                    r = _is.save(con, branch_id, iso, [ln for _r, ln in rows], who, submit=submit, now=now, log_note=note)
                    if not dry_run:
                        audit(con, who, "issue_import", "bc_issue_slip", str(r["slip_id"]), None,
                              {"branch_id": branch_id, "date": iso, "lines": len(rows), "submit": bool(submit), "file": note})
                res["result"] = None if dry_run else ("confirmed" if submit else "saved")
                if submit:
                    res["late"] = int(r.get("late") or 0)
            except ValueError as e:
                res["errors"].append({"row": None, "code": "save", "message": str(e)})
        if res["errors"]:
            res["result"] = "refused"
            res["message"] = res["errors"][0]["message"]
        out.append(res)
    return {"dry_run": dry_run, "submit": bool(submit), "filename": filename, "branch_id": branch_id, "dates": out,
            "ignored": parsed["ignored"], "file_errors": parsed["file_errors"],
            "ok": not parsed["file_errors"] and bool(out) and all(not x["errors"] for x in out),
            "saved": sum(1 for x in out if x["result"] == "saved"), "confirmed": sum(1 for x in out if x["result"] == "confirmed"),
            "refused": sum(1 for x in out if x["result"] == "refused")}


# ── 2 · the order ────────────────────────────────────────────────────────────
def _factor(r: dict) -> float:
    """How many base units one typed unit is: a line with no pack size is typed in กิโลกรัม / ลิตร (show_factor), the rest as they are."""
    return float(r["show_factor"]) if r.get("no_pack") and float(r.get("show_factor") or 1) > 1 else 1.0


def _unit_label(r: dict) -> str:
    if _factor(r) > 1:
        return r.get("show_unit") or ""
    return (r.get("recv_unit") or "แพ็ค") if r.get("order_unit") == "pack" else (r.get("base_unit") or "")


def _q(v, nd=3):
    return None if v is None else round(float(v), nd)


def order_template(con, plan: dict, branch_id: int, rnd: str, delivery: date):
    """-> (workbook, branch name). The central list of the plan, with the plan's own numbers beside the boxes."""
    from .forecast import ROUND_TH
    name = plan.get("branch") or _bname(con, branch_id)
    fb = plan.get("fallback") or {}
    why = f" · {fb['message']}" if fb.get("message") else ""
    note = (f"ใบสั่งซื้อ สาขา{name} · รหัสสาขา {branch_id} · รอบ {rnd} ({ROUND_TH.get(rnd, rnd)}) · วันส่ง {delivery.isoformat()}"
            f" · {TXT['ord_note_tail']}{why}")
    wb, ws = _make(ORDER_SHEET, note, ORDER_HEAD, [14, 44, 16, 22, 14, 14, 12, 38, 24])
    plan = _short(con, plan)
    prev = {x["code"]: x["units"] for x in ((plan.get("previous") or {}).get("lines") or [])}
    for r in plan["rows"]:
        f = _factor(r)
        sys_txt = None
        if r.get("system_qty") is not None:
            v = float(r["system_qty"])
            if f > 1 and (not r.get("system_unit") or r.get("system_unit") == r.get("base_unit")):
                v, u = v / f, _unit_label(r)
            else:
                u = r.get("system_unit") or r.get("count_unit") or r.get("base_unit") or ""
            sys_txt = f"{round(v, 2):g} {u}".strip()
        no_sug = r.get("suggested", 0) is None
        sugg = None if no_sug else _q(r["order_units"] / f)
        last = r.get("last_order_qty") if no_sug else prev.get(r["code"])
        _add(ws, [r["code"], r.get("short_name") or r["name"], _unit_label(r), sys_txt, sugg, None if last is None else _q(last / f), None, None, None],
             {6, 7, 8})
    _finish(wb, ws, ORDER_REASONS_SHEET, [(c["code"], c["label"]) for c in plan["reasons"]], "H", "เหตุผล")
    return wb, name


def read_order_file(raw: bytes) -> dict:
    note, head, body = _table(raw, ORDER_SHEET)
    if not note.startswith("ใบสั่งซื้อ "):
        raise ValueError(TXT["wrong_kind"].format(what="ใบสั่งซื้อ"))
    m_r, m_d = re.search(r"รอบ\s*(\S+)", note), re.search(r"วันส่ง\s*(\d{4}-\d{2}-\d{2})", note)
    pairs = [(c, t) for c, t in _or.REASONS]
    rows, errs = [], []
    for rno, r in body:
        code = str(_cell(head, r, "รหัส") or "").strip()
        q = _cell(head, r, "สั่ง")
        if not code or _blank(q):
            continue
        try:
            rs = _reason_code(_cell(head, r, "เหตุผล"), pairs)
        except KeyError as e:
            errs.append({"row": rno, "code": "reason", "message": TXT["bad_reason"].format(name=code, value=e.args[0])})
            continue
        rows.append((rno, {"code": code, "confirmed_units": _num(q), "reason_code": rs,
                           "note": str(_cell(head, r, "หมายเหตุ") or "").strip() or None}))
    return {"branch_id": _file_branch(note), "round": m_r.group(1) if m_r else None, "delivery": m_d.group(1) if m_d else None,
            "rows": rows, "file_errors": errs}


def import_order(con, plan: dict, branch_id: int, rnd: str, delivery: date, parsed: dict, who: str, filename: str,
                 dry_run: bool = True, submit: bool = False, ack_late: bool = False, now: datetime = None, as_of: date = None) -> dict:
    """Check the file against the plan and save it as the app's save does (orders.save_flow): a draft, or submitted with `submit`
    (the cut-off, today's confirmed slip and the reason rule are the app's, raised by the same code). A dry run rolls everything back
    and never stops at the late fee (it reports `late`; the real call needs the branch's `ack_late`)."""
    if parsed["branch_id"] != branch_id:
        raise _wrong_branch(parsed)
    if parsed["round"] != rnd or parsed["delivery"] != delivery.isoformat():
        raise ValueError(TXT["ord_other_round"].format(file=f"{parsed['round']} · {parsed['delivery']}",
                                                       want=f"{rnd} · {delivery.isoformat()}"))
    by = {r["code"]: r for r in plan["rows"]}
    errors, lines, seen = list(parsed["file_errors"]), [], set()
    for rno, ln in parsed["rows"]:
        c = ln["code"]
        r = by.get(c)
        if r is None:
            errors.append({"row": rno, "code": "code", "message": TXT["ord_unknown"].format(code=c)})
            continue
        if c in seen:
            errors.append({"row": rno, "code": "dup", "message": TXT["dup_code"].format(name=_item(con, c))})
            continue
        seen.add(c)
        try:
            _or.clean_line(con, ln)
        except ValueError as e:
            errors.append({"row": rno, "code": _kind(str(e)), "message": str(e)})
            continue
        f = _factor(r)
        q = None if ln["confirmed_units"] is None else float(ln["confirmed_units"])
        lines.append({**ln, "confirmed_units": q if q is None or f == 1 else round(q * f * 1000) / 1000})
    res = {"dry_run": dry_run, "submit": bool(submit), "filename": filename, "branch_id": branch_id, "round": rnd,
           "delivery": delivery.isoformat(), "lines": len(lines), "errors": errors, "ok": not errors and bool(lines),
           "fallback": plan.get("fallback"), "order_id": None, "status": None, "needs_reason": [], "needs_reason_names": [],
           "value": None, "late": 0, "cutoff_th": plan.get("cutoff_th"), "slip_today": None}
    if errors or not lines:
        if not dry_run:
            raise ValueError(TXT["ord_errors"].format(n=len(errors), first=errors[0]["message"]) if errors else TXT["ord_none"])
        if not errors:
            res["errors"].append({"row": None, "code": "none", "message": TXT["ord_none"]})
        return res
    try:
        with trial(con, keep=not dry_run):
            out = _or.save_flow(con, branch_id, rnd, delivery, lines, who, submit=submit, ack_late=bool(ack_late or dry_run),
                                now=now, as_of=as_of)
            if not dry_run:
                audit(con, who, "order_import", "bc_order", str(out["order_id"]), None,
                      {"branch_id": branch_id, "round": rnd, "delivery": delivery.isoformat(), "lines": len(lines),
                       "submit": bool(submit), "file": (filename or "")[:120]})
    except ValueError as e:
        if not dry_run:                                      # a real import raises what the app raises (LateSubmit included)
            raise
        res["errors"].append({"row": None, "code": _kind(str(e)), "message": str(e)})
        res["ok"] = False
        return res
    res.update(order_id=None if dry_run else out["order_id"], status=out["status"], needs_reason=out["needs_reason"],
               needs_reason_names=sorted({_item(con, c) for c in out["needs_reason"]}), value=out["value"],
               late=int(out.get("late") or 0), slip_today=out.get("slip_today"))
    return res


# ── 3 · the receipt ──────────────────────────────────────────────────────────
def _rcv_state(o: dict) -> None:
    if o["status"] not in ("submitted", "processed"):
        raise ValueError(TXT["rcv_state"].format(st=_or.STATUS_TH.get(o["status"])))


def _rcv_factor(ln: dict) -> float:
    if ln.get("pieces_unit"):                               # a line received by weight keeps its own boxes
        return 1.0
    return _factor(ln)


def _rcv_unit(ln: dict) -> str:
    if _rcv_factor(ln) > 1:
        return ln.get("show_unit") or ln.get("unit_word") or ""
    return ln.get("unit_word") or ("" if ln.get("order_unit") == "base" else "แพ็ค")


def _rcv_sent(ln: dict) -> float:
    return float(ln["processed_units"] if ln.get("processed_units") is not None else (ln.get("confirmed_units") or 0))


def _by_weight(o: dict) -> list:
    """The lines of the order the branch receives by WEIGHT and counts in pieces (a whole salmon): the form has no box for the count."""
    return [ln for ln in o["lines"] if ln.get("pieces_unit") and _rcv_sent(ln) > 0]


def receive_template(con, branch_id: int, order_id: int):
    """-> (workbook, branch name). One row per line the warehouse sent, รับจริง prefilled with ส่งจริง."""
    o = _short(con, _or.detail(con, order_id, receive=True))
    if o["branch_id"] != branch_id:
        raise ValueError(TXT["no_order"])
    _rcv_state(o)
    name = _bname(con, branch_id)
    by_weight = _by_weight(o)
    note = (f"รับสินค้า ใบสั่งซื้อเลขที่ {order_id} · สาขา{name} · รหัสสาขา {branch_id} · {TXT['rcv_note_tail']}"
            + (f" · {TXT['rcv_weight_note']}" if by_weight else ""))
    wb, ws = _make(RECEIVE_SHEET, note, RECEIVE_HEAD, [14, 44, 16, 14, 14, 40, 24])
    for ln in o["lines"]:
        sent = _rcv_sent(ln)
        if not (sent > 0 or ln.get("received_units")):
            continue
        f = _rcv_factor(ln)
        _add(ws, [ln["code"], ln.get("short_name") or ln["name"], _rcv_unit(ln), _q(sent / f), _q(sent / f), None, None],
             set() if ln in by_weight else {4, 5, 6})                  # a line received by weight is received in the app
    _finish(wb, ws, RECEIVE_REASONS_SHEET, [(c, t) for c, t in _cl.REASONS], "F", "สาเหตุส่วนต่าง")
    return wb, name


def read_receive_file(raw: bytes) -> dict:
    note, head, body = _table(raw, RECEIVE_SHEET)
    if not note.startswith("รับสินค้า "):
        raise ValueError(TXT["wrong_kind"].format(what="รับสินค้า"))
    m = re.search(r"ใบสั่งซื้อเลขที่\s*(\d+)", note)
    pairs = [(c, t) for c, t in _cl.REASONS]
    rows, errs = [], []
    for rno, r in body:
        code = str(_cell(head, r, "รหัส") or "").strip()
        if not code:
            continue
        try:
            rs = _reason_code(_cell(head, r, "สาเหตุส่วนต่าง"), pairs)
        except KeyError as e:
            errs.append({"row": rno, "code": "reason", "message": TXT["bad_reason"].format(name=code, value=e.args[0])})
            continue
        rows.append((rno, {"code": code, "received_units": _num(_cell(head, r, "รับจริง")), "reason": rs,
                           "note": str(_cell(head, r, "หมายเหตุ") or "").strip() or None}))
    return {"branch_id": _file_branch(note), "order_id": int(m.group(1)) if m else None, "rows": rows, "file_errors": errs}


def import_receive(con, branch_id: int, order_id: int, parsed: dict, who: str, filename: str, dry_run: bool = True,
                   ack: bool = False, receipt_date: str = None, claim_note: str = None) -> dict:
    """Check the file and book the receipt as the app does (orders.receive): a line short of what was sent needs its reason and becomes a
    claim in the same save. The branch's acknowledgement (`ack`) is never assumed on a real import: orders.receive refuses without it
    (a dry run only previews, so it does not ask)."""
    if parsed["branch_id"] != branch_id:
        raise _wrong_branch(parsed)
    if parsed["order_id"] != order_id:
        raise ValueError(TXT["rcv_other_order"].format(file=parsed["order_id"] if parsed["order_id"] is not None else "-", want=order_id))
    o = _or.detail(con, order_id, receive=True)
    if o["branch_id"] != branch_id:
        raise ValueError(TXT["no_order"])
    _rcv_state(o)
    by = {ln["code"]: ln for ln in o["lines"]}
    errors, lines, seen, rowof = list(parsed["file_errors"]), [], set(), {}
    for w in _by_weight(o):              # the pieces of a line received by weight are counted in the app: a file would book it without them
        errors.append({"row": None, "code": "weight", "message": TXT["rcv_weight"].format(name=w.get("short_name") or w["name"],
                                                                                          unit=w["pieces_unit"])})
    for rno, ln in parsed["rows"]:
        c = ln["code"]
        line = by.get(c)
        if line is None:
            errors.append({"row": rno, "code": "code", "message": TXT["rcv_unknown"].format(code=c)})
            continue
        if c in seen:
            errors.append({"row": rno, "code": "dup", "message": TXT["dup_code"].format(name=_item(con, c))})
            continue
        seen.add(c)
        rowof[_item(con, c)] = rno
        q = ln["received_units"]
        if q is not None:
            try:
                q = float(q)
            except (TypeError, ValueError):
                errors.append({"row": rno, "code": "number", "message": f"{_item(con, c)}: จำนวนต้องเป็นตัวเลข"})
                continue
            f = _rcv_factor(line)
            q = q if f == 1 else round(q * f * 1000) / 1000
        lines.append({"code": c, "received_units": q, "reason": ln["reason"] or "", "note": ln["note"]})
    res = {"dry_run": dry_run, "filename": filename, "branch_id": branch_id, "order_id": order_id, "lines": len(lines),
           "errors": errors, "ok": False, "over": [], "claim_lines": 0, "claim_id": None, "within_tolerance": [], "received": False}
    if errors:
        if not dry_run:
            raise ValueError(TXT["ord_errors"].format(n=len(errors), first=errors[0]["message"]))
        return res
    try:
        with trial(con, keep=not dry_run):
            out = _or.receive(con, branch_id, order_id, lines, who, receipt_date=receipt_date, ack=bool(ack or dry_run),
                              claim_note=claim_note)
            if not dry_run:
                audit(con, who, "receive_import", "bc_order", str(order_id), None,
                      {"branch_id": branch_id, "lines": len(lines), "file": (filename or "")[:120]})
    except ValueError as e:
        if not dry_run:
            raise
        msg = str(e)
        res["errors"].append({"row": next((n for nm, n in rowof.items() if msg.startswith(nm)), None),
                              "code": _kind(msg), "message": msg})
        return res
    res.update(ok=True, received=not dry_run, over=out["over"], claim_lines=out["claim_lines"],
               claim_id=None if dry_run else out["claim_id"], within_tolerance=out["within_tolerance"])
    return res
