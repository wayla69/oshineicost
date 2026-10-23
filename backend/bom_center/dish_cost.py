# -*- coding: utf-8 -*-
"""bom_center.dish_cost - ต้นทุนต่อจานรายเดือน: the analyst imports the dish cost, month by month (PR-S2).

Owner 2026-10-22 (docs/STOCK_FIRST_PLAN.md, decisions O2 and O5): while recipes are Phase 2, the
cost of a dish in every report is the number the analyst imports here. One row per POS code per month
(bc_sku_cost_list); a month with no import carries the previous imported month, per SKU. It beats the
cost of a released recipe version (rules.dated_cost_sql) and never reaches before the Grand Menu month.
There is no release or acknowledge step: the analyst imports, the reports follow.

Buffet package codes (rules.is_tier_code) have no cost per plate - the dishes eaten are costed instead -
so they are refused at import and `_cogs_case` never counts a package twice.
"""
import io
import re
from datetime import datetime, timedelta, timezone

import rules

from . import dated
from .master import audit

MIN_MONTH = dated.CUTOVER_MONTH                                    # '2026-09'
_MONTH = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
SHEET = "ต้นทุนต่อจาน"
HEADERS = ["รหัส POS", "ชื่อเมนู", "ต้นทุนต่อจาน (บาท)", "หมายเหตุ"]
CHANGE_PCT = 30.0                                                  # a cost that moves by more than this is shown for a second look
SRC_IMPORTED, SRC_NONE = "นำเข้า", "ไม่มีในตาราง (ใช้ต้นทุนเดิม)"

_ALIASES = {
    "code": ("รหัส pos", "รหัสpos", "รหัส", "sku_code", "sku", "code"),
    "name": ("ชื่อเมนู", "ชื่อ", "name"),
    "cost": ("ต้นทุนต่อจาน (บาท)", "ต้นทุนต่อจาน", "ต้นทุน", "cost"),
    "note": ("หมายเหตุ", "note"),
}


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S")


def bkk_day(at):
    """The Bangkok calendar day of an imported_at stamp (UTC, no zone suffix): an import at 03:00 Bangkok is stamped the day before."""
    if not at:
        return None
    try:
        return (datetime.strptime(str(at)[:19], "%Y-%m-%dT%H:%M:%S").replace(tzinfo=timezone.utc) + timedelta(hours=7)).date().isoformat()
    except ValueError:
        return str(at)[:10]


def _next(month: str) -> str:
    return dated._next(month)


def _prev(month: str) -> str:
    y, m = int(month[:4]), int(month[5:7])
    return f"{y - (m == 1):04d}-{12 if m == 1 else m - 1:02d}"


def max_month(today=None) -> str:
    """The latest month a cost may be entered for: next month (a list is prepared before the month starts), never further (PL1 #16)."""
    from .clock import today as _today
    t = today or _today()
    return _next(f"{t.year:04d}-{t.month:02d}")


def check_month(month, today=None) -> str:
    if not month or not _MONTH.match(str(month)):
        raise ValueError("เดือนไม่ถูกต้อง (เช่น 2026-09)")
    if month < MIN_MONTH:
        raise ValueError(f"ต้นทุนต่อจานเริ่มได้ตั้งแต่เดือน {MIN_MONTH}")
    if month > max_month(today):
        raise ValueError(f"ต้นทุนต่อจานบันทึกล่วงหน้าได้ไม่เกินเดือน {max_month(today)}")
    return month


def facts_note(now=None) -> str:
    """The line shown after a write: the dated reports follow at once, the fact tables (menu engineering, per-bill P&L, Top 60 per head)
    are built by the sync, so they take the cost at its next run (PL1 #6)."""
    try:
        import cloud_sync as _cs
        from .clock import now_bkk
        n = now or now_bkk()
        nxt = _cs.next_slot(n, _cs.configured_times())
    except Exception:
        return "รายงานเมนูและรายงานรายบิลใช้ต้นทุนนี้หลังการซิงก์รอบถัดไป"
    day = "วันนี้" if nxt.date() == n.astimezone(nxt.tzinfo).date() else "พรุ่งนี้"
    return f"รายงานเมนูและรายงานรายบิลใช้ต้นทุนนี้หลังการซิงก์รอบถัดไป เวลา {nxt:%H:%M} น. ({day})"


def _num(v):
    """float >= 0 or raise ValueError; a blank is the caller's to skip."""
    if isinstance(v, bool):
        raise ValueError("ต้นทุนต้องเป็นตัวเลขไม่ติดลบ")
    try:
        x = float(str(v).replace(",", "").replace("฿", "").strip())
    except (TypeError, ValueError):
        raise ValueError("ต้นทุนต้องเป็นตัวเลขไม่ติดลบ")
    if x != x or x < 0 or x in (float("inf"),):
        raise ValueError("ต้นทุนต้องเป็นตัวเลขไม่ติดลบ")
    return x


# ── reading the file ─────────────────────────────────────────────────────────
def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def _grid(raw: bytes, filename: str) -> list:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        import csv
        text = raw.decode("utf-8-sig", errors="replace")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"อ่านไฟล์ Excel ไม่ได้: {e}")
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def parse(raw: bytes, filename: str = "") -> dict:
    """-> {"format", "rows": [{code, name, cost, note}]}. Two layouts:
      (a) the template: sheet `ต้นทุนต่อจาน`, headers รหัส POS . ชื่อเมนู . ต้นทุนต่อจาน (บาท) . หมายเหตุ;
      (b) BOM Grand Menu.xlsx: `New Item Code AL` + `ต้นทุน AL` and `New Item Code BF` + `ต้นทุน BF` (each code its own cost).
    A row without a cost is left out (a blank is "not provided"); the cost itself is judged by preview()."""
    if not raw:
        raise ValueError("ไฟล์ไม่มีข้อมูล")
    grid = _grid(raw, filename)
    for i, row in enumerate(grid[:15]):
        cells = [_norm(c) for c in row]
        if "new item code al" in cells or "new item code bf" in cells:
            return {"format": "BOM Grand Menu", "rows": _parse_grand_menu(grid, i, cells)}
        hit = {k: next((j for j, c in enumerate(cells) if c in al), None) for k, al in _ALIASES.items()}
        if hit["code"] is not None and hit["cost"] is not None:
            return {"format": "แม่แบบต้นทุนต่อจาน", "rows": _parse_template(grid, i, hit)}
    raise ValueError("ไม่พบหัวคอลัมน์ รหัส POS และ ต้นทุนต่อจาน (บาท) ในไฟล์")


def _parse_template(grid, head, hit) -> list:
    out = []
    for row in grid[head + 1:]:
        def at(k):
            j = hit.get(k)
            return row[j] if j is not None and j < len(row) else None
        code = _text(at("code"))
        cost = at("cost")
        if not code or cost is None or _text(cost) == "":
            continue
        out.append({"code": code, "name": _text(at("name")), "cost": cost, "note": _text(at("note"))})
    return out


def _parse_grand_menu(grid, head, cells) -> list:
    """Each of the two code columns pairs with its own cost column (AL with ต้นทุน AL, BF with ต้นทุน BF)."""
    def col(*names):
        return next((j for j, c in enumerate(cells) if c in names), None)
    pairs = [(col("new item code al"), col("ต้นทุน al")), (col("new item code bf"), col("ต้นทุน bf"))]
    name_c = col("ชื่อไทย", "ชื่อเมนู")
    out = []
    for row in grid[head + 1:]:
        for cc, vc in pairs:
            if cc is None or vc is None or cc >= len(row) or vc >= len(row):
                continue
            code, cost = _text(row[cc]), row[vc]
            if not code or cost is None or _text(cost) == "":
                continue
            out.append({"code": code, "name": _text(row[name_c]) if name_c is not None and name_c < len(row) else "",
                        "cost": cost, "note": ""})
    return out


# ── what the app uses today ──────────────────────────────────────────────────
def _sku_rows(con, codes=None):
    q = "SELECT sku_id, sku_code, COALESCE(description, '') FROM skus"
    rows = con.execute(q).fetchall()
    if codes is not None:
        want = set(codes)
        rows = [r for r in rows if r[1] in want]
    return rows


def _effective(con, ids, month: str) -> dict:
    """{sku_id: the unit cost the reports multiply by in `month`} - rules.effective_unit_cost_sql itself."""
    ids = list(ids)
    out = {}
    expr = rules.effective_unit_cost_sql("ss", f"'{month}'")
    for i in range(0, len(ids), 800):
        chunk = ids[i:i + 800]
        q = f"SELECT ss.sku_id, {expr} FROM skus ss WHERE ss.sku_id IN ({','.join('?' * len(chunk))})"
        out.update({sid: float(c or 0.0) for sid, c in con.execute(q, chunk)})
    return out


def _sold(con, month: str) -> dict:
    """{sku_id: plates} sold in `month` by the lines that SHOULD carry a cost (rules.needs_cost_sql: not a package, not a service line)."""
    q = ("SELECT d.sku_id, SUM(d.qty) FROM sales_sku_daily d JOIN skus sk ON sk.sku_id = d.sku_id"
         f" WHERE d.day >= ? AND d.day < ? AND {rules.needs_cost_sql('sk.sku_code')} GROUP BY d.sku_id HAVING SUM(d.qty) > 0")
    return {sid: float(n) for sid, n in con.execute(q, (f"{month}-01", f"{_next(month)}-01"))}


def _list_rows(con, month: str) -> dict:
    """{sku_id: (list month, cost)} - for each SKU the latest list row on or before `month`."""
    out = {}
    for sid, m, c in con.execute("SELECT sku_id, month, unit_cost FROM bc_sku_cost_list WHERE month <= ? ORDER BY month", (month,)):
        out[sid] = (m, c)
    return out


def _source_of(con, month: str):
    """(sku_id -> source label) maker for the cost the app uses now, when the list does not hold the SKU."""
    rel = dated._latest(dated.month_map(con), month) if month >= dated.CUTOVER_MONTH else {}

    def label(sid, cost):
        if sid in rel:
            return "สูตรที่ประกาศใช้"
        return "ต้นทุนเดิม" if cost > 0 else "ไม่มีต้นทุน"
    return label


# ── preview and commit ───────────────────────────────────────────────────────
def _classify(con, rows):
    """-> (ok [{sku_id, code, name, cost, note}], unknown_codes, tier_codes, bad_values). The last of a repeated code wins."""
    sku = {c: (sid, d) for sid, c, d in _sku_rows(con)}
    ok, unknown, tiers, bad = {}, [], [], []
    for r in rows:
        code = _text(r.get("code"))
        if not code:
            continue
        if rules.is_tier_code(code):
            if code not in tiers:
                tiers.append(code)
            continue
        if code not in sku:
            if code not in unknown:
                unknown.append(code)
            continue
        try:
            cost = _num(r.get("cost"))
        except ValueError as e:
            bad.append({"code": code, "value": _text(r.get("cost")), "why": str(e)})
            continue
        ok[code] = {"sku_id": sku[code][0], "code": code, "name": _text(r.get("name")) or sku[code][1], "cost": cost,
                    "note": _text(r.get("note"))}
    return list(ok.values()), unknown, tiers, bad


def preview(con, month: str, rows: list) -> dict:
    """What an import of `rows` into `month` would do. Nothing is written."""
    check_month(month)
    dated.ensure(con)
    ok, unknown, tiers, bad = _classify(con, rows)
    ids = [o["sku_id"] for o in ok]
    before = _effective(con, ids, _prev(month))               # the month before: what the cost was moving from
    changed = []
    for o in ok:
        old = before.get(o["sku_id"], 0.0)
        if old > 0:
            pct = (o["cost"] - old) / old * 100.0
            if abs(pct) > CHANGE_PCT:
                changed.append({"code": o["code"], "name": o["name"], "old": round(old, 2), "new": round(o["cost"], 2),
                                "pct": round(pct, 1)})
    changed.sort(key=lambda x: -abs(x["pct"]))
    # sold in the month, and still without a list row on or before it once this import is in
    have = {sid for sid in _list_rows(con, month)} | set(ids)
    sold = _sold(con, month)
    miss = [sid for sid in sold if sid not in have]
    cur = _effective(con, miss, month)
    label = _source_of(con, month)
    missing = set(miss)
    info = {sid: (c, d) for sid, c, d in _sku_rows(con) if sid in missing}
    without = sorted(({"code": info[sid][0], "name": info[sid][1], "plates": round(sold[sid], 1),
                       "current_cost": round(cur.get(sid, 0.0), 2), "current_source": label(sid, cur.get(sid, 0.0))}
                      for sid in miss), key=lambda x: -x["plates"])
    return {"month": month, "rows": [{k: o[k] for k in ("code", "name", "cost", "note")} for o in ok],
            "imported": len(ok), "unknown_codes": unknown, "tier_codes": tiers, "bad_values": bad,
            "zero_costs": [o["code"] for o in ok if o["cost"] == 0],
            "changed_30pct": changed, "sold_without_cost": without}


ZERO_CONFIRM_TH = "ต้นทุน 0 บาทจะทำให้เมนูนั้นไม่มีต้นทุนในรายงาน · ยืนยันอีกครั้งหากต้องการบันทึกเป็น 0 บาท"


def commit(con, month: str, rows: list, who: str, confirm_zero: bool = False) -> dict:
    """UPSERT the rows of `rows` into `month` (the caller commits) and audit it. Lines the preview calls
    unknown / package / bad are skipped, exactly as it said. A 0 baht cost wins over the cost table and the recipe,
    so it is written only when the caller confirms it (PL1 #16)."""
    pre = preview(con, month, rows)
    if not pre["imported"]:
        raise ValueError("ไม่มีรายการที่นำเข้าได้")
    if pre["zero_costs"] and not confirm_zero:
        raise ValueError(f"มี {len(pre['zero_costs'])} เมนูที่ต้นทุน 0 บาท ({', '.join(pre['zero_costs'][:5])}) · {ZERO_CONFIRM_TH}")
    ok, _, _, _ = _classify(con, rows)
    at = _now()
    con.executemany(
        "INSERT INTO bc_sku_cost_list(sku_id, month, unit_cost, note, imported_by, imported_at) VALUES(?,?,?,?,?,?)"
        " ON CONFLICT(sku_id, month) DO UPDATE SET unit_cost=excluded.unit_cost, note=excluded.note,"
        " imported_by=excluded.imported_by, imported_at=excluded.imported_at",
        [(o["sku_id"], month, o["cost"], o["note"] or None, who, at) for o in ok])
    audit(con, who, "dish_cost_import", "bc_sku_cost_list", month, None, f"{len(ok)} รายการ")
    return pre


def set_one(con, month: str, sku_code: str, unit_cost, note: str, who: str, confirm_zero: bool = False) -> dict:
    """Edit one row of the month (the caller commits)."""
    check_month(month)
    dated.ensure(con)
    code = _text(sku_code)
    if rules.is_tier_code(code):
        raise ValueError("รหัสแพ็กเกจบุฟเฟ่ต์ไม่มีต้นทุนต่อจาน")
    r = con.execute("SELECT sku_id FROM skus WHERE sku_code=?", (code,)).fetchone()
    if not r:
        raise ValueError("ไม่พบรหัส POS นี้")
    cost = _num(unit_cost)
    if cost == 0 and not confirm_zero:
        raise ValueError(ZERO_CONFIRM_TH)
    old = con.execute("SELECT unit_cost FROM bc_sku_cost_list WHERE sku_id=? AND month=?", (r[0], month)).fetchone()
    con.execute(
        "INSERT INTO bc_sku_cost_list(sku_id, month, unit_cost, note, imported_by, imported_at) VALUES(?,?,?,?,?,?)"
        " ON CONFLICT(sku_id, month) DO UPDATE SET unit_cost=excluded.unit_cost, note=excluded.note,"
        " imported_by=excluded.imported_by, imported_at=excluded.imported_at",
        (r[0], month, cost, _text(note) or None, who, _now()))
    audit(con, who, "dish_cost_edit", "bc_sku_cost_list", f"{month}:{code}", None if not old else f"{old[0]:g}", f"{cost:g}")
    return {"month": month, "sku_code": code, "unit_cost": cost}


# ── reading ──────────────────────────────────────────────────────────────────
def month_table(con, month: str) -> dict:
    """Every SKU sold in `month` or holding a list cost on or before it: the cost, the month before, the change,
    the plates and where the cost comes from."""
    check_month(month)
    dated.ensure(con)
    lst = _list_rows(con, month)
    sold = _sold(con, month)
    ids = sorted(set(lst) | set(sold))
    info = {sid: (c, d) for sid, c, d in _sku_rows(con)}
    ids = [i for i in ids if i in info and not rules.is_tier_code(info[i][0])]
    now_c = _effective(con, ids, month)
    prev_c = _effective(con, ids, _prev(month))
    label = _source_of(con, month)
    rows = []
    for sid in ids:
        code, name = info[sid]
        cost, prev = now_c.get(sid, 0.0), prev_c.get(sid, 0.0)
        if sid in lst:
            lm, cost = lst[sid]
            kind, src = ("imported", SRC_IMPORTED) if lm == month else ("carried", f"ยกมาจากเดือน {lm}")
        else:
            kind, src = "none", SRC_NONE
        pct = round((cost - prev) / prev * 100.0, 1) if prev > 0 else None
        rows.append({"code": code, "name": name, "unit_cost": round(cost, 4), "prev_cost": round(prev, 4) if prev > 0 else None,
                     "change_pct": pct, "changed30": bool(pct is not None and abs(pct) > CHANGE_PCT),
                     "plates": round(sold.get(sid, 0.0), 1), "source": src, "source_kind": kind,
                     "cost_source": label(sid, cost) if kind == "none" else None})
    rows.sort(key=lambda r: (r["source_kind"] != "none", -r["plates"], r["code"]))
    counts = {"all": len(rows), "imported": sum(r["source_kind"] == "imported" for r in rows),
              "carried": sum(r["source_kind"] == "carried" for r in rows),
              "none": sum(r["source_kind"] == "none" for r in rows), "changed30": sum(r["changed30"] for r in rows)}
    last = con.execute("SELECT MAX(imported_at) FROM bc_sku_cost_list WHERE month=?", (month,)).fetchone()[0]
    who = con.execute("SELECT imported_by FROM bc_sku_cost_list WHERE month=? AND imported_at=? LIMIT 1", (month, last)).fetchone() if last else None
    return {"month": month, "rows": rows, "counts": counts, "last_import_at": last, "last_import_day": bkk_day(last),
            "last_import_by": who[0] if who else None}


def months(con, upto: str = None) -> list:
    """The months the page offers: the Grand Menu month through `upto` (this month) and any month that holds a list."""
    dated.ensure(con)
    have = {m: (n, at, by) for m, n, at, by in con.execute(
        "SELECT month, COUNT(*), MAX(imported_at), (SELECT imported_by FROM bc_sku_cost_list b WHERE b.month = a.month"
        " ORDER BY imported_at DESC LIMIT 1) FROM bc_sku_cost_list a GROUP BY month")}
    last = max([upto or MIN_MONTH, MIN_MONTH] + list(have))
    out = []
    for m in dated._months(MIN_MONTH, last):
        n, at, by = have.get(m, (0, None, None))
        out.append({"month": m, "listed": n, "last_import_at": at, "last_import_by": by})
    return out


def template_rows(con, month: str) -> list:
    """The template for `month`: every SKU sold in the month before it (that should carry a cost), with the cost the app uses
    in `month` - blank when it has none."""
    check_month(month)
    sold = _sold(con, _prev(month))
    info = {sid: (c, d) for sid, c, d in _sku_rows(con)}
    ids = [i for i in sold if i in info]
    cur = _effective(con, ids, month)
    return sorted(([info[i][0], info[i][1], round(cur.get(i, 0.0), 4) if cur.get(i, 0.0) > 0 else None, ""] for i in ids),
                  key=lambda r: r[0])


def sold_without_cost(con, month: str) -> list:
    """For หน้างาน: the menus sold in `month` with no list row on or before it."""
    lst = _list_rows(con, month)
    return [sid for sid in _sold(con, month) if sid not in lst]
