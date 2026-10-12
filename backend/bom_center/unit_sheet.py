# -*- coding: utf-8 -*-
"""bom_center.unit_sheet — the unit ladder and the portion recipes as an Excel file (PR-U).

ONE file, laid out for people to fill in:

  worksheet « หน่วยและการแปลง »   one row per ladder unit.   A1 = a note, row 2 = the headers
      รหัส · ชื่อวัตถุดิบ · หน่วยฐาน · หน่วย · 1 หน่วยนี้ = · หน่วยที่อ้างอิง · ชนิดการแปลง · ต่ำสุด · สูงสุด · หมายเหตุ
      `หน่วยที่อ้างอิง` blank = the base unit; `ชนิดการแปลง` = แน่นอน | เฉลี่ย; ต่ำสุด / สูงสุด only for เฉลี่ย.
      The name and the base unit are for reading: they are ignored on import (a different base unit is reported).
  worksheet « สูตรแบ่งชิ้น » (optional)   one row per portion recipe.   A1 = a note, row 2 = the headers
      รหัสสินค้ากึ่งสำเร็จ · ชื่อสินค้ากึ่งสำเร็จ · รหัสวัตถุดิบ · จำนวนที่ใช้ · หน่วยที่ใช้ · ได้ (ชิ้น) · หน่วยที่ได้ · หมายเหตุ

An import REPLACES the whole ladder of every code that has a unit in the file (codes not in the file are untouched; a code
whose rows have no unit keeps its ladder). The file is applied all or nothing: one code with a problem stops the lot.
"""
import io

from .master import _num as _mnum
from .master import plan_ladder, set_ladder
from .prep import MODE_PORTION, _clean, latest_price_month, save_portion
from .units import KIND_TH, Converter, norm

SHEET = "หน่วยและการแปลง"
HEADERS = ["รหัส", "ชื่อวัตถุดิบ", "หน่วยฐาน", "หน่วย", "1 หน่วยนี้ =", "หน่วยที่อ้างอิง", "ชนิดการแปลง", "ต่ำสุด", "สูงสุด", "หมายเหตุ"]
NOTE = ("บันไดหน่วยของวัตถุดิบ · 1 หน่วยนี้ = จำนวนของหน่วยที่อ้างอิง (เว้นว่าง = หน่วยฐาน) · แปลงได้ไม่เกิน 3 ขั้นจากหน่วยฐาน · "
        "ชนิดการแปลง “เฉลี่ย” ใช้กับของที่แต่ละชิ้นน้ำหนักไม่เท่ากัน ใส่ค่าเฉลี่ยในช่อง “1 หน่วยนี้ =” และช่วงต่ำสุด–สูงสุด · "
        "นำเข้าแล้วระบบแทนที่บันไดหน่วยของรหัสที่อยู่ในไฟล์ทั้งรายการ")
PORTION_SHEET = "สูตรแบ่งชิ้น"
PORTION_HEADERS = ["รหัสสินค้ากึ่งสำเร็จ", "ชื่อสินค้ากึ่งสำเร็จ", "รหัสวัตถุดิบ", "จำนวนที่ใช้", "หน่วยที่ใช้", "ได้ (ชิ้น)", "หน่วยที่ได้",
                   "หมายเหตุ"]
PORTION_NOTE = ("สูตรแบ่งชิ้นจากวัตถุดิบทั้งตัว · 1 แถวต่อ 1 สินค้ากึ่งสำเร็จ · “จำนวนที่ใช้” และ “หน่วยที่ใช้” คือของทั้งตัวที่นำมาแบ่ง "
                "(หน่วยต้องอยู่ในบันไดหน่วยของวัตถุดิบ เช่น ตัว) · “ได้ (ชิ้น)” คือจำนวนชิ้นที่ใช้ได้จริงจากของที่ใช้ "
                "ซึ่งนับส่วนสูญเสียไว้แล้ว ระบบจึงไม่บวกส่วนสูญเสียซ้ำ · รหัสที่ยังไม่มีในทะเบียนจะถูกสร้างเป็นสินค้ากึ่งสำเร็จ")

_KEY = {"รหัส": "code", "ชื่อวัตถุดิบ": "name", "หน่วยฐาน": "base", "หน่วย": "unit", "1หน่วยนี้=": "qty", "หน่วยที่อ้างอิง": "parent",
        "ชนิดการแปลง": "kind", "ต่ำสุด": "rmin", "สูงสุด": "rmax", "หมายเหตุ": "note"}
_PKEY = {"รหัสสินค้ากึ่งสำเร็จ": "code", "ชื่อสินค้ากึ่งสำเร็จ": "name", "รหัสวัตถุดิบ": "raw", "จำนวนที่ใช้": "qty",
         "หน่วยที่ใช้": "unit", "ได้(ชิ้น)": "pieces", "หน่วยที่ได้": "out_unit", "หมายเหตุ": "note"}


def _h(v) -> str:
    return str(v or "").replace(" ", "").replace(" ", "").strip()


# ── export ───────────────────────────────────────────────────────────────────
def ladder_rows(con, all_items: bool = False) -> list:
    """Rows of the first worksheet: every item that has a ladder (lowest step first); with `all_items` also one empty row for
    every active item that has none, so a person can fill a new ladder in the same file."""
    conv = Converter(con)
    names = {c: (n, k, a) for c, n, k, a in con.execute("SELECT code, name_th, kind, active FROM bc_ingredient")}
    labels = {(c, norm(u)): lb for c, u, lb in con.execute("SELECT ingredient_code, unit, label FROM bc_unit_rule")}
    codes = sorted(set(conv.ladders) | ({c for c, v in names.items() if v[2]} if all_items else set()))
    out = []
    for c in codes:
        if c not in names:
            continue
        lad = conv.ladder(c)
        if not lad:
            if names[c][2]:
                out.append([c, names[c][0], conv.base_unit(c), None, None, None, None, None, None, None])
            continue
        for r in lad:
            out.append([c, names[c][0], conv.base_unit(c), r["unit"], r["qty_in_parent"],
                        None if r["parent_unit"] == conv.base_unit(c) else r["parent_unit"], KIND_TH.get(r["kind"], "แน่นอน"),
                        r["range_min"], r["range_max"], labels.get((c, r["unit"]))])
    return out


def portion_rows(con) -> list:
    from .prep import book
    names = dict(con.execute("SELECT code, name_th FROM bc_ingredient"))
    out = []
    for code, r in sorted(book(con).items()):
        if r.get("mode") != MODE_PORTION or not r["lines"]:
            continue
        ln = r["lines"][0]
        note = con.execute("SELECT note FROM bc_prep_recipe WHERE prepared_code=?", (code,)).fetchone()
        out.append([code, names.get(code), ln["ingredient_code"], ln["qty"], ln["unit"], r["yield_qty"], r["yield_unit"],
                    note[0] if note else None])
    return out


# ── import: read ─────────────────────────────────────────────────────────────
def _find_header(rows: list, keys: dict, must: tuple) -> tuple:
    for i, r in enumerate(rows[:6]):
        m = {}
        for j, v in enumerate(r):
            k = keys.get(_h(v))
            if k and k not in m:
                m[k] = j
        if all(x in m for x in must):
            return i, m
    return None, {}


def parse(data: bytes) -> dict:
    """-> {ladders: {code: {rows: [...], row_nos: [...], base_in_file: set}}, portions: [..], problems: [..], sheets: [..]}"""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                                  # not an Excel file
        raise ValueError(f"อ่านไฟล์ Excel ไม่ได้: {e}")
    out = {"ladders": {}, "portions": [], "problems": [], "sheets": []}
    ws_main = wb[SHEET] if SHEET in wb.sheetnames else None
    ws_port = wb[PORTION_SHEET] if PORTION_SHEET in wb.sheetnames else None
    if ws_main is None:                                                     # a renamed worksheet: take the first that has the headers
        for ws in wb.worksheets:
            if ws is ws_port:
                continue
            rows = list(ws.iter_rows(values_only=True, max_row=8))
            if _find_header(rows, _KEY, ("code", "unit", "qty"))[0] is not None:
                ws_main = ws
                break
    if ws_main is None and ws_port is None:
        raise ValueError(f"ไม่พบแผ่นงาน “{SHEET}” หรือ “{PORTION_SHEET}” · ใช้ไฟล์จากปุ่ม “ดาวน์โหลดแม่แบบหน่วย (Excel)”")
    if ws_main is not None:
        rows = list(ws_main.iter_rows(values_only=True))
        hi, m = _find_header(rows, _KEY, ("code", "unit", "qty"))
        if hi is None:
            raise ValueError(f"แผ่นงาน “{ws_main.title}” ไม่มีหัวคอลัมน์ รหัส · หน่วย · 1 หน่วยนี้ =")
        out["sheets"].append(ws_main.title)
        for n, r in enumerate(rows[hi + 1:], start=hi + 2):
            def g(k, r=r):
                j = m.get(k)
                return r[j] if j is not None and j < len(r) else None
            code, unit = _clean(g("code")), _clean(g("unit"))
            if not code or not unit:
                continue                                                    # an empty row (or an item with no unit) is not a ladder row
            e = out["ladders"].setdefault(code, {"rows": [], "row_nos": [], "base_in_file": set(), "name": None})
            e["name"] = e["name"] or _clean(g("name")) or None
            if _clean(g("base")):
                e["base_in_file"].add(norm(g("base")))
            if _clean(g("qty")) and _mnum(g("qty")) is None:
                out["problems"].append({"code": code, "row": n, "problem": f"แถว {n}: “{_clean(g('qty'))}” ไม่ใช่ตัวเลข"})
            e["rows"].append({"unit": unit, "qty_in_parent": g("qty"), "parent_unit": g("parent"), "conv_kind": g("kind"),
                              "range_min": g("rmin"), "range_max": g("rmax"), "label": _clean(g("note")) or None})
            e["row_nos"].append(n)
    if ws_port is not None:
        rows = list(ws_port.iter_rows(values_only=True))
        hi, m = _find_header(rows, _PKEY, ("code", "raw", "qty", "pieces"))
        if hi is None:
            raise ValueError(f"แผ่นงาน “{PORTION_SHEET}” ไม่มีหัวคอลัมน์ที่ต้องใช้ · ใช้ไฟล์จากปุ่ม “ดาวน์โหลดแม่แบบหน่วย (Excel)”")
        out["sheets"].append(ws_port.title)
        for n, r in enumerate(rows[hi + 1:], start=hi + 2):
            def g(k, r=r):
                j = m.get(k)
                return r[j] if j is not None and j < len(r) else None
            code, raw = _clean(g("code")), _clean(g("raw"))
            if not code and not raw:
                continue
            out["portions"].append({"row": n, "code": code, "name": _clean(g("name")), "raw": raw, "qty": g("qty"),
                                    "unit": _clean(g("unit")), "pieces": g("pieces"), "out_unit": _clean(g("out_unit")) or "ชิ้น",
                                    "note": _clean(g("note")) or None})
    return out


# ── import: plan / apply (the caller commits, or rolls back a dry run) ──────────
def apply(con, parsed: dict, who: str, month: str = None) -> dict:
    """Write what the file says INSIDE the caller's transaction and report per code. -> {ladders: [..], portions: [..],
    problems: N, warnings: [..]}. A code with a problem is not written; the caller refuses the commit when problems > 0."""
    month = month or latest_price_month(con)
    pre = {}                                                              # parse-time problems by code
    for p in parsed["problems"]:
        pre.setdefault(p["code"], []).append(p["problem"])
    res, warnings, total = [], [], 0
    from .aliases import active_map, suggest_items          # PR-Y: a code that is an alias is its item; an unknown code is never added here
    known = {r[0] for r in con.execute("SELECT code FROM bc_ingredient")}
    amap = active_map(con)
    for file_code, e in parsed["ladders"].items():
        code = file_code if file_code in known else amap.get(file_code, file_code)
        plan = plan_ladder(con, code, e["rows"])
        if code not in known:
            hint = suggest_items(con, e.get("name"), code=file_code)
            more = (" · อาจเป็นรหัสใหม่ของ " + " / ".join(f"{h['code']} {h['name']}" for h in hint) if hint
                    else " · ถ้าเป็นรหัสใหม่ของวัตถุดิบเดิม") + " ให้ฝ่ายวิเคราะห์บันทึกเป็นรหัสในระบบอื่นก่อน"
            plan["problems"] = [f"ไม่พบรหัส {file_code} ในทะเบียน{more}"]
        for pb in pre.get(file_code, []):
            if pb not in plan["problems"]:
                plan["problems"].insert(0, pb)
        base = plan.get("base_unit")
        for b in e["base_in_file"]:
            if base and b and b != base:
                warnings.append(f"{code}: หน่วยฐานในไฟล์ “{b}” ต่างจากทะเบียน “{base}” · ใช้ตามทะเบียน")
        if not plan["problems"]:
            set_ladder(con, code, e["rows"], who)
        total += 1 if plan["problems"] else 0
        res.append({"code": code, "file_code": file_code if file_code != code else None, "name": plan["name"], "base_unit": base,
                    "ladder": plan["ladder"], "problems": plan["problems"], "changes": plan["changes"], "qty_per_recv_change": plan["qty_per_recv_change"], "rows": e["row_nos"]})
    seen, port = set(), []
    for p in parsed["portions"]:
        item = {"row": p["row"], "prepared_code": p["code"], "prepared_name": p["name"], "raw_code": p["raw"], "qty": _mnum(p["qty"]),
                "unit": p["unit"], "pieces": _mnum(p["pieces"]), "out_unit": p["out_unit"], "problems": []}
        try:
            if not p["code"]:
                raise ValueError("ระบุรหัสสินค้ากึ่งสำเร็จ")
            if p["code"] in seen:
                raise ValueError(f"{p['code']}: สูตรแบ่งชิ้นมีวัตถุดิบได้ 1 บรรทัด · รหัสนี้อยู่ในไฟล์มากกว่า 1 แถว")
            seen.add(p["code"])
            # no savepoint: the caller rolls the whole file back when anything has a problem (and always on a dry run)
            raw_code = p["raw"] if p["raw"] in known else amap.get(p["raw"], p["raw"])      # PR-Y: an outside code of the raw item
            r = save_portion(con, p["code"], raw_code, p["qty"], p["unit"], p["pieces"], p["out_unit"], who, note=p["note"],
                             name=p["name"], month=month)
            item.update(factor_base=r["factor_base"], approx=r["approx"], how=r["how"], base_unit=r["base_unit"],
                        base_per_piece=r["base_per_piece"], cost_per_piece=r["cost_per_piece"], price_month=r["price_month"],
                        range=r["range"], created=r["created"], replaced=r["replaced"])
        except ValueError as ex:
            item["problems"].append(str(ex))
            total += 1
        port.append(item)
    return {"ladders": res, "portions": port, "problems": total, "warnings": warnings, "month": month}
