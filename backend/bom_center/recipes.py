# -*- coding: utf-8 -*-
"""bom_center.recipes — recipes by tier (plan §4.2 M3, Sprint S2).

A monthly release holds QUANTITIES only. A line costs

    qty in base units × price_per_base(price_month) ÷ (1 − loss %)

where loss % is the line's own trim loss when the source sheet set one (Cost-RW
writes it per line: radish 60.2 g lost per 1,000 g), else the ingredient's.
A line marked zero_by_design costs 0 on purpose (salmon "ส่วนบราวน์" made from
the trimmings, owner decision §10-9). A line that cannot be costed — no price,
a unit the converter cannot turn into base units, no ingredient code — is a
PROBLEM, never a silent zero: the dish lists it, and the release checks (S3)
block on it.

Tier is data: a dish's min_tier is the tier sheet it is defined on, and the
basket of tier T is every dish with min_tier <= T (a ฿1,299 guest eats the ฿459
dishes too).

Importing the Cost-RW workbook: its recipe blocks carry NO POS code, only the
dish's English and Thai name, so each block is matched to a buffet SKU by
  1. a person's earlier answer (bc_dish_name_map),
  2. the same name on a SKU of the same tier (skus.description + sku_tier),
  3. the same name on a buffet SKU of any tier,
  4. a close spelling on the same tier (difflib >= 0.88) — tagged 'fuzzy' so
     the analyst sees it,
and otherwise kept as an UNMATCHED block (bc_release_check) that the analyst
maps on screen; mapping it applies the block at once and remembers the answer.
"""
import difflib
import io
import json
import re
from datetime import datetime, timezone

from .master import add_from_purchase, audit, family_of, kind_of
from .prices import price_for
from .schema import ensure
from .units import Converter, norm

TIERS = (459, 629, 789, 999, 1299)
FUZZY_MIN = 0.88
SKIP = "-"            # bc_dish_name_map.sku_code for a block the analyst ruled "not on sale"


def now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S")


def name_key(name) -> str:
    """A dish name reduced to what identifies it across the workbook and the POS:
    no (BF)/(B) prefix, no "(6 ชิ้น ต่อ 1 เสิร์ฟ)" serving note, no spaces, and the
    two spellings of teriyaki made one."""
    s = str(name or "").replace("​", "")
    s = re.sub(r"^\s*\((?:BF|B|AL|A|T)\)\s*", "", s)
    s = re.sub(r"\([^)]*\)", "", s)
    s = s.lower().replace("เทอริ", "เทริ").replace("เทอรี", "เทริ")
    return re.sub(r"[\s\-_.,·]+", "", s)


# ── Cost-RW workbook ─────────────────────────────────────────────────────────
def parse_cost_rw(data: bytes, kinds=("BF",)) -> dict:
    """-> {'blocks': [...], 'sheets': [...], 'skipped': [...]} for the BF_<tier>NET sheets."""
    import openpyxl
    from prep_grand_menu_bom import parse_recipe_sheet
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks, sheets, skipped = [], [], []
    for sn in wb.sheetnames:
        m = re.match(r"^(BF|AC)_([\d,]+)NET", sn.strip())
        if not m:
            if re.search(r"(BF|AC)_[\d,]+NET", sn):
                skipped.append({"sheet": sn, "why": "เป็นสำเนาหรือชื่อไม่ตรงรูปแบบ จึงไม่นำเข้า"})
            continue
        kind, tier = m.group(1), int(m.group(2).replace(",", ""))
        if kind not in kinds:
            continue
        got = [b for b in parse_recipe_sheet(wb[sn])
               if b.get("ings") and (b.get("th") or "").strip() not in ("ชื่อภาษาไทย", "MENU NAME")]
        for b in got:
            b.update({"sheet": sn, "tier": tier, "kind": kind})
        blocks += got
        sheets.append({"sheet": sn, "tier": tier, "blocks": len(got)})
    return {"blocks": blocks, "sheets": sheets, "skipped": skipped}


def _candidates(con):
    """Buffet SKUs by (tier, name_key) and by name_key alone."""
    tier_of = dict(con.execute("SELECT s.sku_code, t.tier FROM sku_tier t JOIN skus s USING(sku_id)")) \
        if con.execute("SELECT 1 FROM sqlite_master WHERE name='sku_tier'").fetchone() else {}
    by_tier, by_name = {}, {}
    for code, desc in con.execute("SELECT sku_code, description FROM skus WHERE sku_code LIKE 'OTHBF%'"):
        k = name_key(desc)
        if not k:
            continue
        t = str(tier_of.get(code) or "")
        by_tier.setdefault((t, k), []).append(code)
        by_name.setdefault(k, []).append(code)
    return by_tier, by_name


def match_blocks(con, blocks: list) -> list:
    """Annotate each block with sku_code / method / score (in place) and return it."""
    ensure(con)
    by_tier, by_name = _candidates(con)
    mapped = {(t, k): c for t, k, c in con.execute("SELECT tier, name_key, sku_code FROM bc_dish_name_map")}
    for b in blocks:
        tier = str(b["tier"])
        keys = [k for k in (name_key(b.get("th")), name_key(b.get("en"))) if k]
        b["name_key"] = keys[0] if keys else ""
        b.update(sku_code=None, method=None, score=None)
        for k in keys:
            if (tier, k) in mapped:
                if mapped[(tier, k)] == SKIP:          # "not on sale" — remembered, never re-asked
                    b.update(method="skip")
                else:
                    b.update(sku_code=mapped[(tier, k)], method="map")
                break
        if b["sku_code"] or b["method"] == "skip":
            continue
        for k in keys:
            hit = by_tier.get((tier, k))
            if hit and len(hit) == 1:
                b.update(sku_code=hit[0], method="exact")
                break
        if b["sku_code"]:
            continue
        for k in keys:
            hit = by_name.get(k)
            if hit and len(hit) == 1:
                b.update(sku_code=hit[0], method="exact-anytier")
                break
        if b["sku_code"] or not keys:
            continue
        pool = [(k, codes[0]) for (t, k), codes in by_tier.items() if t == tier and len(codes) == 1]
        best = max(((difflib.SequenceMatcher(None, keys[0], k).ratio(), c) for k, c in pool), default=(0, None))
        if best[0] >= FUZZY_MIN:
            b.update(sku_code=best[1], method="fuzzy", score=round(best[0], 3))
    # one SKU claimed by two blocks: keep the stronger claim, unmatch the other
    rank = {"map": 4, "exact": 3, "exact-anytier": 2, "fuzzy": 1}
    owner = {}
    for b in blocks:
        if not b["sku_code"]:
            continue
        o = owner.get(b["sku_code"])
        if o is None or rank[b["method"]] > rank[o["method"]]:
            if o is not None:
                o.update(sku_code=None, method=None, why=f"รหัสนี้ถูกใช้กับสูตรอื่นแล้ว ({b['th'] or b['en']})")
            owner[b["sku_code"]] = b
        else:
            b.update(sku_code=None, method=None, why=f"รหัสนี้ถูกใช้กับสูตรอื่นแล้ว ({o['th'] or o['en']})")
    return blocks


# ── release rows ─────────────────────────────────────────────────────────────
def get_or_create_draft(con, label: str, effective_month: str, price_month: str, who: str, modules=None) -> int:
    ensure(con)
    from .dated import CUTOVER_MONTH
    if effective_month < CUTOVER_MONTH:
        raise ValueError(f"เวอร์ชันต้องมีผลตั้งแต่ {CUTOVER_MONTH} (เดือนที่เริ่มเมนูใหม่) เดือนก่อนหน้านั้นใช้ต้นทุนจากระบบเดิม")
    r = con.execute("SELECT release_id, status FROM bc_release WHERE label=?", (label,)).fetchone()
    if r:
        if r[1] not in ("draft", "review"):
            raise ValueError(f"release {label} ประกาศใช้แล้ว แก้ไขไม่ได้ · สร้างร่างใหม่")
        if r[1] == "review" and modules is not None and "bom.center" not in modules:
            raise ValueError(REVIEW_LOCKED)
        con.execute("UPDATE bc_release SET effective_month=?, price_month=? WHERE release_id=?",
                    (effective_month, price_month, r[0]))
        return r[0]
    parent = con.execute("SELECT release_id FROM bc_release WHERE status='released' ORDER BY effective_month DESC,"
                         " release_id DESC LIMIT 1").fetchone()
    con.execute("INSERT INTO bc_release(label, kind, effective_month, status, scope, created_by, created_at,"
                " parent_release_id, price_month) VALUES(?, 'monthly', ?, 'draft', 'buffet', ?, ?, ?, ?)",
                (label, effective_month, who, now(), parent[0] if parent else None, price_month))
    return con.execute("SELECT release_id FROM bc_release WHERE label=?", (label,)).fetchone()[0]


def _serve(text: str):
    m = re.match(r"\s*([\d.]+)\s*(.*)$", text or "")
    if not m:
        return None, (text or "").strip() or None
    try:
        return float(m.group(1)), (text or "").strip()
    except ValueError:
        return None, (text or "").strip()


def _ensure_dish(con, sku_code: str, b: dict) -> int:
    r = con.execute("SELECT dish_id FROM bc_dish_sku WHERE sku_code=?", (sku_code,)).fetchone()
    q, unit = _serve(b.get("serve"))
    if r:
        con.execute("UPDATE bc_dish SET name_th=COALESCE(NULLIF(?,''), name_th), name_en=COALESCE(NULLIF(?,''), name_en),"
                    " min_tier=? WHERE dish_id=?", (b.get("th"), b.get("en"), b["tier"], r[0]))
        con.execute("UPDATE bc_dish_sku SET serve_qty=COALESCE(?, serve_qty), serve_unit=COALESCE(?, serve_unit)"
                    " WHERE sku_code=?", (q, unit, sku_code))
        return r[0]
    con.execute("INSERT INTO bc_dish(name_th, name_en, channel_default, min_tier) VALUES(?,?, 'BF', ?)",
                (b.get("th") or b.get("en") or sku_code, b.get("en"), b["tier"]))
    did = con.execute("SELECT last_insert_rowid()").fetchone()[0]
    con.execute("INSERT INTO bc_dish_sku(sku_code, dish_id, channel, serve_qty, serve_unit) VALUES(?,?, 'BF', ?, ?)",
                (sku_code, did, q, unit))
    return did


def _block_lines(b: dict) -> list:
    """Cost-RW ingredient rows -> recipe lines (quantities only)."""
    out = []
    for i, g in enumerate(b.get("ings") or [], start=1):
        code = (g.get("code") or "").strip()
        if code in ("#N/A", "None"):
            code = ""
        loss = None
        if g.get("full_qty") and g["full_qty"] > 0 and g.get("loss") is not None:
            loss = round(100.0 * g["loss"] / g["full_qty"], 6)
        out.append({"seq": i, "ingredient_code": code or None,
                    "ingredient_name": g.get("ing") or g.get("prod") or code,
                    "qty": g.get("qty") or 0.0, "unit": norm(g.get("unit")) or None,
                    "prep_note": g.get("detail") or None, "loss_pct": loss,
                    # ฿0 ON PURPOSE only when the sheet really says 0 (salmon trimmings); a
                    # #N/A / #DIV/0! cell also reads 0 but means "broken", and is priced
                    # from the master instead
                    "zero_by_design": 1 if (not g.get("full_price") and not g.get("cost")
                                            and not str(g.get("raw_cost") or "").startswith("#")) else 0,
                    "_product": g.get("prod"), "_full_price": g.get("full_price"), "_full_qty": g.get("full_qty"),
                    "_sheet_cost": g.get("cost"), "_raw_cost": g.get("raw_cost")})
    return out


def write_recipe(con, release_id: int, sku_code: str, lines: list, *, source: str, source_cost=None,
                 source_ref=None, who: str = "-", register_codes: bool = True):
    """Replace one SKU's recipe in a draft release."""
    con.execute("DELETE FROM bc_recipe_line WHERE release_id=? AND sku_code=?", (release_id, sku_code))
    con.execute("INSERT OR REPLACE INTO bc_recipe(release_id, sku_code, kind, source, source_cost, source_ref)"
                " VALUES(?,?, 'recipe', ?, ?, ?)", (release_id, sku_code, source, source_cost, source_ref))
    known = {r[0] for r in con.execute("SELECT code FROM bc_ingredient")}
    for n, ln in enumerate(lines, start=1):
        code = (ln.get("ingredient_code") or "").strip() or None
        if code and register_codes and code not in known:
            from .aliases import resolve
            res = resolve(con, code, known)          # an ACTIVE alias of a register item is that item: use it, register nothing
            if res in known:
                code = res
        if code and register_codes and code not in known:
            add_from_purchase(con, code, ln.get("_product") or ln.get("ingredient_name"), ln.get("unit"), who)
            con.execute("UPDATE bc_ingredient SET kind=?, family=?, note=? WHERE code=? AND updated_by=?",
                        (kind_of(code, ln.get("ingredient_name") or ""), family_of(code),
                         "เพิ่มจากสูตร Cost-RW · ตรวจหน่วยและราคา", code, who))
            known.add(code)
        con.execute("INSERT INTO bc_recipe_line(release_id, sku_code, seq, ingredient_code, ingredient_name, qty, unit,"
                    " prep_note, zero_by_design, loss_pct) VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (release_id, sku_code, n, code, ln.get("ingredient_name"), float(ln.get("qty") or 0),
                     norm(ln.get("unit")) or None, ln.get("prep_note"), 1 if ln.get("zero_by_design") else 0,
                     ln.get("loss_pct")))


def import_cost_rw(con, parsed: dict, release_id: int, who: str) -> dict:
    """Write every matched block into the draft; park the rest as unmatched checks."""
    blocks = match_blocks(con, parsed["blocks"])
    con.execute("DELETE FROM bc_release_check WHERE release_id=? AND check_code='unmatched'", (release_id,))
    stats = {"blocks": len(blocks), "matched": 0, "by_method": {}, "unmatched": [], "lines": 0,
             "zero_lines": 0, "no_code_lines": 0, "error_cells": 0}
    stats["skipped_not_on_sale"] = []
    for b in blocks:
        lines = _block_lines(b)
        stats["error_cells"] += sum(1 for ln in lines if str(ln.get("_raw_cost") or "").startswith("#"))
        if b["method"] == "skip":
            stats["skipped_not_on_sale"].append({"tier": b["tier"], "th": b.get("th"), "en": b.get("en")})
            continue
        if not b["sku_code"]:
            payload = {k: b.get(k) for k in ("tier", "en", "th", "serve", "total", "sheet", "name_key", "why")}
            payload["lines"] = [{k: v for k, v in ln.items()} for ln in lines]
            con.execute("INSERT OR REPLACE INTO bc_release_check(release_id, check_code, ref, severity, message)"
                        " VALUES(?, 'unmatched', ?, 'block', ?)",
                        (release_id, f"{b['tier']}|{b['name_key'] or b.get('en')}", json.dumps(payload, ensure_ascii=False)))
            stats["unmatched"].append({"tier": b["tier"], "en": b.get("en"), "th": b.get("th"),
                                       "total": b.get("total"), "why": b.get("why")})
            continue
        _ensure_dish(con, b["sku_code"], b)
        write_recipe(con, release_id, b["sku_code"], lines, source="cost-rw", source_cost=b.get("total"),
                     source_ref=f"{b['sheet']} · {b.get('th') or b.get('en')}", who=who)
        stats["matched"] += 1
        stats["by_method"][b["method"]] = stats["by_method"].get(b["method"], 0) + 1
        stats["lines"] += len(lines)
        stats["zero_lines"] += sum(1 for ln in lines if ln["zero_by_design"])
        stats["no_code_lines"] += sum(1 for ln in lines if not ln["ingredient_code"])
    stats["fuzzy"] = [{"tier": b["tier"], "th": b.get("th"), "en": b.get("en"), "sku_code": b["sku_code"],
                       "score": b["score"]} for b in blocks if b["method"] == "fuzzy"]
    audit(con, who, "import_cost_rw", "bc_release", str(release_id), None,
          json.dumps({k: v for k, v in stats.items() if k in ("blocks", "matched", "lines")}))
    return stats


def apply_mapping(con, release_id: int, ref: str, sku_code: str, who: str, modules=None) -> dict:
    """The analyst answered an unmatched block: remember it, apply the block now."""
    rel = release_row(con, release_id)
    if rel and rel["status"] == "review" and modules is not None and "bom.center" not in modules:
        raise ValueError(REVIEW_LOCKED)      # a release in review is the analyst's to change, not the kitchen's
    r = con.execute("SELECT message FROM bc_release_check WHERE release_id=? AND check_code='unmatched' AND ref=?",
                    (release_id, ref)).fetchone()
    if not r:
        raise ValueError("ไม่พบสูตรนี้ในรายการที่ยังไม่จับคู่")
    if sku_code != SKIP and not con.execute("SELECT 1 FROM skus WHERE sku_code=?", (sku_code,)).fetchone():
        raise ValueError(f"ไม่มีรหัส {sku_code} ใน POS")
    b = json.loads(r[0])
    con.execute("INSERT OR REPLACE INTO bc_dish_name_map(tier, name_key, sku_code, mapped_by, mapped_at) VALUES(?,?,?,?,?)",
                (str(b["tier"]), b.get("name_key") or "", sku_code, who, now()))
    if sku_code == SKIP:
        con.execute("DELETE FROM bc_release_check WHERE release_id=? AND check_code='unmatched' AND ref=?",
                    (release_id, ref))
        audit(con, who, "skip_block", "bc_dish_name_map", ref, None, "not on sale")
        return {"sku_code": None, "skipped": True, "lines": 0}
    _ensure_dish(con, sku_code, b)
    write_recipe(con, release_id, sku_code, b["lines"], source="cost-rw", source_cost=b.get("total"),
                 source_ref=f"{b.get('sheet')} · {b.get('th') or b.get('en')} (จับคู่โดย {who})", who=who)
    con.execute("DELETE FROM bc_release_check WHERE release_id=? AND check_code='unmatched' AND ref=?", (release_id, ref))
    audit(con, who, "map_block", "bc_dish_name_map", ref, None, sku_code)
    return {"sku_code": sku_code, "lines": len(b["lines"])}


# ── costing ──────────────────────────────────────────────────────────────────
class Coster:
    """Prices recipe lines for one month. Built once per request."""

    def __init__(self, con, price_month: str, prep_book: dict = None):
        self.con, self.month = con, price_month
        self.conv = Converter(con)
        self.loss = dict(con.execute("SELECT code, loss_pct FROM bc_ingredient"))
        self._p = {}
        self.prep_book = prep_book or {}
        self._prep = None

    def price(self, code):
        """A prepared item with a complete sub-recipe is costed from its raw ingredients
        (Phase 1.5); otherwise — and for everything raw — the month's price."""
        if code not in self._p:
            p = None
            if code in self.prep_book:
                if self._prep is None:
                    from .prep import PrepCoster
                    self._prep = PrepCoster(self.con, self.month, self.prep_book, self.conv)
                c = self._prep.cost(code)
                if c["per_base"] is not None:
                    p = (c["per_base"], "sub_recipe", self.month, 0)
            self._p[code] = p or price_for(self.con, code, self.month)
        return self._p[code]

    def line(self, ln: dict) -> dict:
        """-> {cost, qty_base, base_unit, price_per_base, loss_pct, problem}"""
        out = {"cost": 0.0, "qty_base": None, "base_unit": None, "price_per_base": None,
               "loss_pct": None, "problem": None, "price_source": None}
        code = ln.get("ingredient_code")
        if ln.get("zero_by_design"):
            out["problem"] = None
            out["zero"] = True
            return out
        if not code:
            out["problem"] = "ไม่มีรหัสวัตถุดิบ"
            return out
        qb, how = self.conv.to_base(code, ln.get("qty") or 0, ln.get("unit"))
        out["base_unit"] = self.conv.base_unit(code)
        if qb is None:
            out["problem"] = how
            return out
        out["qty_base"] = qb
        p = self.price(code)
        if not p:
            out["problem"] = f"ไม่มีราคา {code} เดือน {self.month}"
            return out
        loss = ln.get("loss_pct")
        if loss is None:
            loss = self.loss.get(code) or 0.0
        if loss >= 100:
            out["problem"] = "สูญเสีย ≥ 100%"
            return out
        out.update(price_per_base=p[0], price_source=p[1], loss_pct=loss,
                   cost=qb * p[0] / (1.0 - loss / 100.0))
        return out


def coster_for(con, release_id: int, rel: dict = None) -> Coster:
    """The Coster a release is priced with: its price month and its sub-recipes
    (frozen into it once released — prep.book)."""
    from .prep import book
    rel = rel or release_row(con, release_id)
    return Coster(con, rel["price_month"] or rel["effective_month"], book(con, release_id))


def recipe_lines(con, release_id: int, sku_code: str) -> list:
    cols = ("seq", "ingredient_code", "ingredient_name", "qty", "unit", "prep_note", "zero_by_design",
            "loss_pct", "line_cost")
    return [dict(zip(cols, r)) for r in con.execute(
        f"SELECT {', '.join(cols)} FROM bc_recipe_line WHERE release_id=? AND sku_code=? ORDER BY seq",
        (release_id, sku_code))]


def release_row(con, release_id: int) -> dict:
    cols = ("release_id", "label", "kind", "effective_month", "status", "scope", "price_month", "created_by",
            "created_at", "released_by", "released_at", "acknowledged_by", "acknowledged_at", "parent_release_id", "notes",
            "submitted_by", "submitted_at")
    r = con.execute(f"SELECT {', '.join(cols)} FROM bc_release WHERE release_id=?", (release_id,)).fetchone()
    return dict(zip(cols, r)) if r else None


REVIEW_LOCKED = "ร่างนี้อยู่ระหว่างตรวจของฝ่ายวิเคราะห์ แก้ไขสูตรไม่ได้จนกว่าจะส่งกลับ"


def assert_editable(rel: dict, modules=None):
    """A recipe may change only while its release is a draft; once the kitchen has
    submitted it for review only the analyst (bom.center) may still edit.
    `modules=None` = no permission context (internal callers, tests)."""
    if not rel or rel["status"] not in ("draft", "review"):
        raise ValueError("แก้ไขได้เฉพาะเวอร์ชันที่ยังเป็นร่าง")
    if rel["status"] == "review" and modules is not None and "bom.center" not in modules:
        raise ValueError(REVIEW_LOCKED)


def cost_recipe(con, release_id: int, sku_code: str, coster: "Coster" = None) -> dict:
    rel = release_row(con, release_id)
    rec = con.execute("SELECT kind, flat_cost, alias_sku_code, source, source_cost, source_ref FROM bc_recipe"
                      " WHERE release_id=? AND sku_code=?", (release_id, sku_code)).fetchone()
    if not rec:
        return {"sku_code": sku_code, "in_release": False, "total": None, "lines": [], "problems": 0}
    kind, flat, alias_to, source, source_cost, source_ref = rec
    lines = recipe_lines(con, release_id, alias_to if kind == "alias" else sku_code)
    if rel["kind"] == "baseline":                        # v0 carries its exact legacy numbers
        for ln in lines:
            ln["cost"] = float(ln["line_cost"] or 0.0)
        total = float(flat or 0.0) if kind == "flat" else (0.0 if kind == "zero" else sum(ln["cost"] for ln in lines))
        problems = 0
    else:
        coster = coster or coster_for(con, release_id, rel)
        for ln in lines:
            ln.update(coster.line(ln))
        total = float(flat or 0.0) if kind == "flat" else sum(ln["cost"] for ln in lines)
        problems = sum(1 for ln in lines if ln.get("problem"))
    return {"sku_code": sku_code, "in_release": True, "kind": kind, "source": source, "source_ref": source_ref,
            "source_cost": source_cost, "total": round(total, 6), "lines": lines, "problems": problems}


def release_costs(con, release_id: int) -> dict:
    """{sku_code: (total, problems)} for every recipe in a release."""
    rel = release_row(con, release_id)
    coster = None if rel["kind"] == "baseline" else coster_for(con, release_id, rel)
    out = {}
    for (code,) in con.execute("SELECT sku_code FROM bc_recipe WHERE release_id=?", (release_id,)).fetchall():
        c = cost_recipe(con, release_id, code, coster)
        out[code] = (c["total"], c["problems"])
    return out


def tier_view(con, release_id: int, tier: int = None) -> dict:
    """Every buffet dish of the release (min_tier <= tier when given), with its
    cost now, what the sheet said, and what v0 (today's engine) says."""
    ensure(con)
    costs = release_costs(con, release_id)
    base = con.execute("SELECT release_id FROM bc_release WHERE kind='baseline' AND status='released'"
                       " ORDER BY release_id DESC LIMIT 1").fetchone()
    v0 = dict(con.execute("SELECT sku_code, unit_cost FROM bcd_cost_cache WHERE release_id=?", (base[0],))) if base else {}
    if not v0:
        v0 = {c: v for c, v in con.execute("SELECT s.sku_code, e.recipe_cost FROM sku_effective_cost e JOIN skus s USING(sku_id)")} \
            if con.execute("SELECT 1 FROM sqlite_master WHERE name='sku_effective_cost'").fetchone() else {}
    rows = []
    for code, desc, dish_tier, name_th, name_en, serve, src_cost, src_ref, n in con.execute(
            "SELECT r.sku_code, s.description, d.min_tier, d.name_th, d.name_en, ds.serve_unit, r.source_cost,"
            " r.source_ref, (SELECT COUNT(*) FROM bc_recipe_line l WHERE l.release_id=r.release_id AND l.sku_code=r.sku_code)"
            " FROM bc_recipe r LEFT JOIN skus s ON s.sku_code=r.sku_code"
            " LEFT JOIN bc_dish_sku ds ON ds.sku_code=r.sku_code LEFT JOIN bc_dish d ON d.dish_id=ds.dish_id"
            " WHERE r.release_id=?", (release_id,)):
        if tier and dish_tier and dish_tier > tier:
            continue
        total, problems = costs.get(code, (None, 0))
        old = v0.get(code)
        rows.append({"sku_code": code, "name": name_th or desc, "name_en": name_en, "pos_name": desc,
                     "min_tier": dish_tier, "serve": serve, "cost": total, "problems": problems, "lines": n,
                     "sheet_cost": src_cost, "source_ref": src_ref,
                     "sheet_diff": (round(total - src_cost, 4) if (total is not None and src_cost is not None) else None),
                     "v0_cost": old, "v0_diff_pct": (round(100.0 * (total - old) / old, 1) if (old and total is not None) else None)})
    rows.sort(key=lambda r: (r["min_tier"] or 0, r["name"] or ""))
    unmatched = [json.loads(m) | {"ref": ref} for ref, m in con.execute(
        "SELECT ref, message FROM bc_release_check WHERE release_id=? AND check_code='unmatched' ORDER BY ref", (release_id,))]
    for u in unmatched:
        u["n_lines"] = len(u.pop("lines", []))
    counts = {t: sum(1 for r in rows if r["min_tier"] == t) for t in TIERS}
    return {"release": release_row(con, release_id), "tier": tier, "rows": rows, "unmatched": unmatched,
            "counts": counts, "problems": sum(1 for r in rows if r["problems"]),
            "sheet_mismatch": sum(1 for r in rows if r["sheet_diff"] is not None and abs(r["sheet_diff"]) > 0.01)}


def save_lines(con, release_id: int, sku_code: str, lines: list, who: str, modules=None):
    assert_editable(release_row(con, release_id), modules)
    old = recipe_lines(con, release_id, sku_code)
    keep = con.execute("SELECT source_cost, source_ref FROM bc_recipe WHERE release_id=? AND sku_code=?",
                       (release_id, sku_code)).fetchone()
    clean = []
    for ln in lines:
        if not (ln.get("ingredient_code") or ln.get("ingredient_name")):
            continue
        q = float(ln.get("qty") or 0)
        if q < 0:
            raise ValueError("ปริมาณต้องไม่ติดลบ")
        lp = ln.get("loss_pct")
        if lp not in (None, "") and not (0 <= float(lp) < 100):
            raise ValueError("สูญเสีย % ต้องอยู่ระหว่าง 0–99")
        clean.append({**ln, "qty": q, "loss_pct": None if lp in (None, "") else float(lp)})
    write_recipe(con, release_id, sku_code, clean, source="edited", source_cost=keep[0] if keep else None,
                 source_ref=keep[1] if keep else None, who=who, register_codes=False)
    audit(con, who, "edit_recipe", "bc_recipe_line", f"{release_id}|{sku_code}",
          json.dumps([[o["ingredient_code"], o["qty"], o["unit"]] for o in old], ensure_ascii=False),
          json.dumps([[c.get("ingredient_code"), c["qty"], c.get("unit")] for c in clean], ensure_ascii=False))


# ── flat template (one sheet per tier) ───────────────────────────────────────
FLAT_HEADERS = ["รหัส POS", "ชื่อเมนู", "Tier", "ลำดับ", "รหัสวัตถุดิบ", "วัตถุดิบ", "ปริมาณ", "หน่วย",
                "สูญเสีย % (เว้นว่าง = ตามทะเบียน)", "ตั้งใจเป็น 0", "หมายเหตุการเตรียม"]


def _style_header(ws, widths=None):
    """Bold white-on-navy header row (row 1) and optional column widths."""
    from openpyxl.styles import Font, PatternFill
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A6B")
    for col, w in zip("ABCDEFGHIJKLMNOP", widths or ()):
        ws.column_dimensions[col].width = w


def _style_flat_sheet(ws):
    _style_header(ws, (16, 34, 7, 6, 13, 30, 9, 10, 12, 10, 30))


def export_flat(con, release_id: int):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    view = tier_view(con, release_id)
    by_tier = {}
    for r in view["rows"]:
        by_tier.setdefault(r["min_tier"] or 0, []).append(r)
    for t in sorted(by_tier):
        ws = wb.create_sheet(f"BF_{t}" if t else "BF_ไม่มีTier")
        ws.append(FLAT_HEADERS)
        for r in by_tier[t]:
            for ln in recipe_lines(con, release_id, r["sku_code"]):
                ws.append([r["sku_code"], r["name"], t, ln["seq"], ln["ingredient_code"], ln["ingredient_name"],
                           ln["qty"], ln["unit"], ln["loss_pct"], "ใช่" if ln["zero_by_design"] else "", ln["prep_note"]])
        _style_flat_sheet(ws)
    if not wb.sheetnames:
        ws = wb.create_sheet("BF")
        ws.append(FLAT_HEADERS)
        _style_flat_sheet(ws)
    return wb


# ── blank template: many menus in one file (PR-C, owner 2026-10-09) ───────────
_SAMPLE = "(ตัวอย่าง ลบก่อนนำเข้า)"
TEMPLATE_SAMPLE = [
    ["ตัวอย่าง-001", f"เมนูตัวอย่าง A {_SAMPLE}", 459, 1, "ตัวอย่าง-01", f"วัตถุดิบตัวอย่าง 1 {_SAMPLE}", 80, "กรัม",
     None, "", "หั่นชิ้นพอดีคำ"],
    ["ตัวอย่าง-001", f"เมนูตัวอย่าง A {_SAMPLE}", 459, 2, "ตัวอย่าง-02", f"วัตถุดิบตัวอย่าง 2 {_SAMPLE}", 15, "มิลลิลิตร",
     None, "", None],
    ["ตัวอย่าง-002", f"เมนูตัวอย่าง B {_SAMPLE}", 629, 1, "ตัวอย่าง-01", f"วัตถุดิบตัวอย่าง 1 {_SAMPLE}", 1, "ชิ้น",
     2.5, "", None],
]
TEMPLATE_HOWTO = [
    "กรอกสูตรที่แผ่นงาน “สูตร” บรรทัดละ 1 วัตถุดิบ เมนูเดียวกันใช้รหัส POS เดียวกันทุกบรรทัด",
    "รหัส POS ต้องตรงกับแผ่นงาน “รหัส POS” และรหัสวัตถุดิบต้องตรงกับแผ่นงาน “รหัสวัตถุดิบ” (คัดลอกได้)",
    "หน่วย: กรัม / มิลลิลิตร / ชิ้น ตามหน่วยฐาน หรือกิโลกรัม ลิตร แพ็ค (แปลงเป็นหน่วยฐานตอนนำเข้า)",
    "สูญเสีย % เว้นว่าง = ใช้ค่าตามทะเบียนวัตถุดิบ · “ตั้งใจเป็น 0” ใส่ ใช่ เมื่อบรรทัดนั้นไม่มีต้นทุนโดยเจตนา",
    "นำเข้าแล้ว สูตรของเมนูที่อยู่ในไฟล์ถูกแทนที่ทั้งสูตร เมนูที่ไม่อยู่ในไฟล์ไม่เปลี่ยน · ลบแถวตัวอย่างก่อนนำเข้า",
    "ดูตัวอย่างก่อนบันทึกทุกครั้ง · เมื่อครบแล้วกด “ส่งตรวจ” ให้ฝ่ายวิเคราะห์",
]
_KIND_TH = {"raw": "วัตถุดิบ", "prepared": "ของที่ผลิตเอง", "consumable": "วัสดุสิ้นเปลือง", "zero_by_design": "ตั้งใจเป็น 0"}


def pos_sku_rows(con, release_id: int = None) -> list:
    """[(sku_code, description, tier, has_recipe)] — the POS menus a recipe may be written for."""
    have = {c for (c,) in con.execute("SELECT sku_code FROM bc_recipe WHERE release_id=?", (release_id or -1,))}
    has = lambda t: bool(con.execute("SELECT 1 FROM sqlite_master WHERE name=?", (t,)).fetchone())
    active = ""
    if "active" in {r[1] for r in con.execute("PRAGMA table_info(skus)")}:
        active = " AND COALESCE(active,1)=1"
    if has("sku_tier"):
        tier = dict(con.execute("SELECT s.sku_code, t.tier FROM sku_tier t JOIN skus s USING(sku_id)"))
    else:
        tier = dict(con.execute("SELECT ds.sku_code, d.min_tier FROM bc_dish_sku ds JOIN bc_dish d USING(dish_id)"))
    return [(c, d, tier.get(c), c in have) for c, d in con.execute(
        "SELECT sku_code, description FROM skus WHERE (sku_code LIKE 'OTHBF%' OR sku_code LIKE 'OTHAL%')"
        f"{active} ORDER BY sku_code")]


def ingredient_rows(con) -> list:
    """[(code, name_th, base_unit, kind_th)] — the ingredient register a recipe line may point at."""
    return [(c, n, u, _KIND_TH.get(k, k)) for c, n, u, k in con.execute(
        "SELECT code, name_th, base_unit, kind FROM bc_ingredient WHERE COALESCE(active,1)=1 ORDER BY code")]


def template_flat(con, release_id: int = None):
    """A blank workbook for the kitchen: sheet “สูตร” (the flat layout, two sample menus), the POS menus,
    the ingredient codes and how to fill it. Only “สูตร” has A1 = “รหัส POS”, so import_flat reads it alone."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "สูตร"
    ws.append(FLAT_HEADERS)
    for row in TEMPLATE_SAMPLE:
        ws.append(list(row))
    _style_flat_sheet(ws)
    ws = wb.create_sheet("รหัส POS")
    # A1 must NOT read "รหัส POS": import_flat takes every sheet whose A1 is that text for recipe rows
    ws.append(["รหัสเมนู POS", "ชื่อเมนู", "Tier", "มีสูตรในร่างแล้ว"])
    for code, desc, tier, has_rec in pos_sku_rows(con, release_id):
        ws.append([code, desc, tier, "มี" if has_rec else ""])
    _style_header(ws, (16, 40, 8, 20))
    ws = wb.create_sheet("รหัสวัตถุดิบ")
    ws.append(["รหัสวัตถุดิบ", "ชื่อวัตถุดิบ", "หน่วยฐาน", "ชนิด"])
    for row in ingredient_rows(con):
        ws.append(list(row))
    _style_header(ws, (16, 40, 12, 18))
    ws = wb.create_sheet("วิธีใช้")
    ws.append(["วิธีใช้แม่แบบสูตร"])
    for i, t in enumerate(TEMPLATE_HOWTO, start=1):
        ws.append([f"{i}. {t}"])
    _style_header(ws, (120,))
    return wb


def import_flat(con, data: bytes, release_id: int, who: str, modules=None) -> dict:
    """The flat template back into a draft: every SKU it mentions is replaced."""
    import openpyxl
    assert_editable(release_row(con, release_id), modules)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    recipes, errors = {}, []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows or str(rows[0][0] or "").strip() != FLAT_HEADERS[0]:
            continue
        for i, r in enumerate(rows[1:], start=2):
            r = list(r) + [None] * 11
            sku = str(r[0] or "").strip()
            if not sku:
                continue
            if "(ตัวอย่าง" in " ".join(str(c) for c in r if c):
                continue                        # QA L4d: the template's example rows ("(ตัวอย่าง ลบก่อนนำเข้า)", and the older wording) are never recipes
            try:
                qty = float(r[6] or 0)
            except (TypeError, ValueError):
                errors.append({"sheet": ws.title, "row": i, "why": "ปริมาณไม่ใช่ตัวเลข"})
                continue
            loss = r[8]
            try:
                loss = None if loss in (None, "") else float(loss)
            except (TypeError, ValueError):
                errors.append({"sheet": ws.title, "row": i, "why": "สูญเสีย % ไม่ใช่ตัวเลข"})
                continue
            recipes.setdefault(sku, []).append({
                "ingredient_code": (str(r[4]).strip() if r[4] else None), "ingredient_name": r[5],
                "qty": qty, "unit": r[7], "loss_pct": loss,
                "zero_by_design": str(r[9] or "").strip() in ("ใช่", "1", "y", "yes", "true", "✓"),
                "prep_note": r[10]})
    known = {c for (c,) in con.execute("SELECT sku_code FROM skus")}
    unknown = sorted(s for s in recipes if s not in known)
    for s in unknown:
        errors.append({"sku": s, "why": "ไม่มีรหัส POS นี้"})
    for sku, lines in recipes.items():
        if sku in known:
            save_lines(con, release_id, sku, lines, who, modules=modules)
    return {"recipes": len(recipes) - len(unknown), "lines": sum(len(v) for k, v in recipes.items() if k in known),
            "errors": errors[:200]}
