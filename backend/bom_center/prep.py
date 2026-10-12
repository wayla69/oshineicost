# -*- coding: utf-8 -*-
"""Phase 1.5 — sub-recipes of prepared items (CK / HK / DK …) — docs/BOM_CENTER_PLAN.md §5.8.

A prepared item is made in the restaurant from raw ingredients: "น้ำซุปใส — 1 สูตร =
15,225 ml" uses 30 g dried fish, 175 ml mirin … This module

  * imports the kitchen's sheet (BoM_Semi + its "Cost per unit" lookup) or a flat file,
  * matches each sheet row to the master's prepared code by NAME and BATCH SIZE —
    never by code: the sheet's CK numbers are an older scheme (its CK-001 น้ำซุปใส
    is the master's HK-001; the master's CK-001 is a rice item),
  * costs a prepared item from its raw ingredients × actual prices (recursive: a
    sub-recipe may use another one, e.g. แป้งเทมปุระเปียก), and
  * explodes a prepared item down to raw ingredients for ordering.

Sub-recipes are master data the kitchen edits (bc_prep_recipe / bc_prep_line).
A release freezes the set it was costed with (bc_rel_prep_*), so editing a
sub-recipe later never changes a released month's cost.
"""
import io
import re
from difflib import SequenceMatcher

from .codes import _core
from .codes import key as _key
from .codes import suggest as _suggest
from .master import audit, now
from .prices import price_for
from .units import Converter, norm

SKIP = "-"
TOLERANCE_PCT = 5.0            # computed vs purchase price: more = the kitchen checks the yield (C14)

_YIELD = re.compile(r"=\s*([0-9][0-9,]*(?:\.[0-9]+)?)\s*([^\s0-9]+)")


def parse_yield(text):
    """'1 สูตร = 15,225 ml' -> (15225.0, 'มิลลิลิตร'); '1 สูตร = 5 PCS' -> (5.0, 'ชิ้น')."""
    m = _YIELD.search(str(text or ""))
    if not m:
        return None, None
    try:
        return float(m.group(1).replace(",", "")), norm(m.group(2))
    except ValueError:
        return None, None


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


# ── reading the kitchen's sheet ──────────────────────────────────────────────
def parse_bom_semi(data: bytes) -> dict:
    """The BoM_Semi sheet (+ 'Cost per unit' when the workbook has it).

    -> {"sheet", "blocks": [{row, src_code, name, kitchen, yield_text, yield_qty, yield_unit,
        net_cost, lines: [{seq, name, qty, src_price}]}], "cpu": {name: (code, unit, price)}}"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks, sheet_used = [], None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr_i, cols = None, {}
        for i, r in enumerate(rows[:15]):
            vals = [_clean(v) for v in r]
            if "Mat1" in vals and "Usage1" in vals and "Menu" in vals:
                hdr_i, cols = i, {v: j for j, v in enumerate(vals) if v}
                break
        if hdr_i is None:
            continue
        sheet_used = ws.title
        n_mat = max(int(k[3:]) for k in cols if re.fullmatch(r"Mat\d+", k))
        for i, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
            def g(k):
                j = cols.get(k)
                return r[j] if j is not None and j < len(r) else None
            name = _clean(g("Menu"))
            if not name:
                continue
            lines = []
            for n in range(1, n_mat + 1):
                mat, use = _clean(g(f"Mat{n}")), _num(g(f"Usage{n}"))
                if mat and use:
                    lines.append({"seq": n, "name": mat, "qty": use, "src_price": _num(g(f"Price{n}"))})
            if not lines:
                continue
            ytext = _clean(g("Description"))
            yq, yu = parse_yield(ytext)
            blocks.append({"row": i, "src_code": _clean(g("Code")), "name": name, "kitchen": _clean(g("Kitchen")),
                           "yield_text": ytext, "yield_qty": yq, "yield_unit": yu,
                           "net_cost": _num(g("Net Cost")), "lines": lines})
        break
    cpu = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True, max_row=4000))
        hdr_i = next((i for i, r in enumerate(rows[:10]) if "Mat-BoM" in [_clean(v) for v in r]), None)
        if hdr_i is None:
            continue
        h = {_clean(v): j for j, v in enumerate(rows[hdr_i]) if v}
        jc, jn, ju, jp = h.get("Items Code"), h.get("Mat-BoM"), h.get("Unit : BoM"), h.get("@BoM")
        for r in rows[hdr_i + 1:]:
            nm = _clean(r[jn]) if jn is not None and jn < len(r) else ""
            if not nm:
                continue
            code = _clean(r[jc]) if jc is not None else ""
            unit = _clean(r[ju]) if ju is not None else ""
            cpu.setdefault(nm, (code, norm(unit) if unit and unit != "รอถาม" else "", _num(r[jp]) if jp is not None else None))
        break
    if sheet_used is None:
        raise ValueError("ไม่พบแผ่นงานสูตรกึ่งสำเร็จ (BoM_Semi) ต้องมีหัวคอลัมน์ Menu / Mat1 / Usage1")
    return {"sheet": sheet_used, "blocks": blocks, "cpu": cpu}


# ── which master code is this sheet row? ─────────────────────────────────────
def _sim(a: str, b: str) -> float:
    ka, kb = _key(a), _key(b)
    if not ka or not kb:
        return 0.0
    if ka == kb:
        return 1.0
    r = 0.0
    for x, y in ((ka, kb), (_core(ka), _core(kb))):    # also without น้ำ/ซอส/…: น้ำจิ้มซารุ ≈ ซอสซารุ
        if not x or not y:
            continue
        r = max(r, SequenceMatcher(None, x, y, autojunk=False).ratio())
        if x in y or y in x:
            r = max(r, 0.6 + 0.4 * min(len(x), len(y)) / max(len(x), len(y)))
    return r


def _same_yield(b: dict, base: str, per) -> bool:
    if not b.get("yield_qty") or not per:
        return False
    if b.get("yield_unit") and base and norm(b["yield_unit"]) != norm(base):
        return False
    # the same batch (pack sizes are rounded: 7,500 for a 7,520 ml batch) — not a near one (3,122 vs 3,150)
    return abs(b["yield_qty"] - per) <= 0.005 * max(b["yield_qty"], per)


def prepared_master(con) -> dict:
    return {c: {"name": n, "base_unit": norm(u), "qty_per_recv": p} for c, n, u, p in con.execute(
        "SELECT code, name_th, base_unit, qty_per_recv FROM bc_ingredient WHERE kind='prepared' AND active=1")}


def match_blocks(con, blocks: list) -> dict:
    """Assign every sheet row to the master's prepared codes it describes.

    A master code takes the row whose name fits best — the same batch size
    (= the master's pack size) lets a weaker name through ("แป้งเทมปุระเปียก"
    is "แป้งทอดเทมปุระ HK-012", both 2,600 g). One row may serve several codes
    (the master carries old and new numbers for some items). A person's mapping
    (bc_prep_name_map) always wins; '-' marks a row as not ours."""
    prep = prepared_master(con)
    mapped = dict(con.execute("SELECT name_key, prepared_code FROM bc_prep_name_map"))
    for b in blocks:
        b["codes"], b["key"] = [], _key(b["name"])
        m = mapped.get(b["key"])
        if m:
            b["skip"] = m == SKIP
            if m != SKIP:
                b["codes"].append({"code": m, "name": (prep.get(m) or {}).get("name"), "method": "map", "score": 1.0})
    taken = {c["code"] for b in blocks for c in b["codes"]}
    for code, p in prep.items():
        if code in taken:
            continue
        best = None
        for b in blocks:
            if b.get("skip") or any(c["method"] == "map" for c in b["codes"]):
                continue
            if b.get("yield_unit") and p["base_unit"] and norm(b["yield_unit"]) != p["base_unit"]:
                continue                                    # 110 g of fried salmon is not 1 ลูก of larb
            ns = _sim(b["name"], p["name"])
            y = _same_yield(b, p["base_unit"], p["qty_per_recv"])
            if ns >= 0.85 or (y and ns >= 0.35):
                sc = ns + (0.3 if y else 0.0)
                if best is None or sc > best[0]:
                    best = (sc, b, "name+yield" if y else "name")
        if best:
            best[1]["codes"].append({"code": code, "name": p["name"], "method": best[2], "score": round(best[0], 2)})
    for b in blocks:
        b["codes"].sort(key=lambda c: (-c["score"], c["code"]))
    with_recipe = {c["code"] for b in blocks for c in b["codes"]}
    return {"blocks": blocks, "prepared_without": sorted(set(prep) - with_recipe),
            "unmatched": [b for b in blocks if not b["codes"] and not b.get("skip")]}


def _resolver(con, blocks: list, cpu: dict):
    """material name -> (code, unit, how) — sub-recipe → lookup sheet → master name."""
    master = {c: (n, k, norm(u)) for c, n, k, u in con.execute("SELECT code, name_th, kind, base_unit FROM bc_ingredient")}
    by_key = {}
    for c, (n, k, _u) in master.items():
        if k != "prepared":                                 # a prepared item is reached through its sheet row only
            by_key.setdefault(_key(n), []).append((0 if k == "raw" else 1, c))
    raw_master = {c: (n, k) for c, (n, k, _u) in master.items() if k != "prepared"}
    sub = {}
    for b in blocks:
        if b["codes"]:
            sub[b["key"]] = (b["codes"][0]["code"], b.get("yield_unit"))

    def resolve(name: str):
        k = _key(name)
        if k in sub:                                        # another prepared item (nested sub-recipe)
            code, _ = sub[k]
            return code, master.get(code, (None, None, ""))[2], "สูตรกึ่งสำเร็จ"
        hit = cpu.get(_clean(name))
        if hit and hit[0] in master and master[hit[0]][1] != "prepared":
            return hit[0], hit[1] or master[hit[0]][2], "ตาราง Cost per unit"
        if not k:                                           # an English-only name has no Thai key: never guess
            return None, (hit[1] if hit else ""), None
        if by_key.get(k):
            code = sorted(by_key[k])[0][1]
            return code, (hit[1] if hit and hit[1] else master[code][2]), "ชื่อตรงทะเบียน"
        if hit and hit[0] in master:
            return hit[0], hit[1] or master[hit[0]][2], "ตาราง Cost per unit"
        near = [s for s in _suggest(raw_master, name) if s["score"] >= 1.0]   # the name sits inside one master name
        if near:
            code = near[0]["code"]
            return code, (hit[1] if hit and hit[1] else master[code][2]), "ชื่อใกล้เคียง (ควรตรวจสอบ)"
        return None, (hit[1] if hit else ""), None
    return resolve


_MASS_VOL = {"กรัม", "มิลลิลิตร"}
_COUNT = {"ชิ้น", "ฟอง", "ลูก", "ตัว", "แผ่น", "ใบ", "ดอก"}


def _sheet_unit(unit: str, base: str) -> str:
    """The kitchen sheet counts liquids in g and eggs in ชิ้น: its price per "G" of mirin
    is our price per ml. g↔ml and count↔count therefore mean the master's base unit."""
    u = norm(unit)
    if base and u and u != base and ({u, base} <= _MASS_VOL or {u, base} <= _COUNT):
        return base
    return u or base or ""


def plan_import(con, parsed: dict) -> dict:
    """What an import WOULD write (nothing is written): per master code, its rows and lines."""
    m = match_blocks(con, parsed["blocks"])
    resolve = _resolver(con, m["blocks"], parsed.get("cpu") or {})
    have = {r[0] for r in con.execute("SELECT prepared_code FROM bc_prep_recipe")}
    portion = portion_codes(con)
    base_of = {c: norm(u) for c, u in con.execute("SELECT code, base_unit FROM bc_ingredient")}
    writes, unresolved = [], 0
    for b in m["blocks"]:
        lines = []
        own = {c["code"] for c in b["codes"]}
        for ln in b["lines"]:
            code, unit, how = resolve(ln["name"])
            if code in own:                                 # never a recipe of itself
                code, how = None, None
            unit = _sheet_unit(unit, base_of.get(code))
            unresolved += code is None
            lines.append({"seq": ln["seq"], "ingredient_code": code or "", "qty": ln["qty"], "unit": unit or "",
                          "note": ln["name"], "how": how, "src_price": ln.get("src_price")})
        b["resolved"] = lines
        for c in b["codes"]:
            writes.append({"prepared_code": c["code"], "prepared_name": c["name"], "method": c["method"],
                           "score": c["score"], "src_name": b["name"], "src_code": b["src_code"], "row": b["row"],
                           "yield_qty": b["yield_qty"], "yield_unit": b["yield_unit"], "yield_text": b["yield_text"],
                           "net_cost": b["net_cost"], "lines": lines, "replaces": c["code"] in have and c["code"] not in portion,
                           "unresolved": sum(1 for ln in lines if not ln["ingredient_code"]),
                           "portion_skip": c["code"] in portion, "note_th": PORTION_SKIP_NOTE if c["code"] in portion else None})
    no_yield = [w for w in writes if not w["yield_qty"]]
    return {"sheet": parsed.get("sheet"), "rows": len(m["blocks"]), "writes": writes,
            "unmatched": [{"name": b["name"], "src_code": b["src_code"], "row": b["row"], "yield_text": b["yield_text"],
                           "key": b["key"]} for b in m["unmatched"]],
            "prepared_without": m["prepared_without"], "unresolved_lines": unresolved, "no_yield": len(no_yield)}


MODE_PORTION = "portion"          # สูตรแบ่งชิ้น: one whole piece (a salmon) -> N pieces (PR-U); anything else is a batch
MODE_TH = {"batch": "สูตรผลิตเป็นชุด", "portion": "แบ่งชิ้น"}


def save(con, code: str, yield_qty, yield_unit, lines: list, who: str, source_ref: str = None,
         source_cost=None, mode: str = None, note: str = None):
    """Replace one prepared item's sub-recipe. `mode='portion'` marks a portion recipe: ONE line whose trim loss is not
    applied again (the yield already counts only the usable pieces); any other value stores a plain batch recipe."""
    mode = MODE_PORTION if mode == MODE_PORTION else None
    yq = _num(yield_qty)
    if not yq or yq <= 0:
        raise ValueError(f"{code}: ต้องระบุปริมาณที่ได้ต่อ 1 สูตร (> 0)")
    if not norm(yield_unit):
        raise ValueError(f"{code}: ต้องระบุหน่วยของปริมาณที่ได้")
    if not con.execute("SELECT 1 FROM bc_ingredient WHERE code=?", (code,)).fetchone():
        raise ValueError(f"ไม่พบรหัส {code} ในทะเบียน")
    clean = []
    for i, ln in enumerate(lines or [], start=1):
        q = _num(ln.get("qty"))
        ic = _clean(ln.get("ingredient_code"))
        if not q and not ic:
            continue
        if not q or q <= 0:
            raise ValueError(f"{code} บรรทัด {i}: ปริมาณต้องมากกว่า 0")
        if ic == code:
            raise ValueError(f"{code} บรรทัด {i}: ไม่สามารถใช้รายการนี้เป็นส่วนผสมของตัวเองได้")
        clean.append((len(clean) + 1, ic, q, norm(ln.get("unit")) or None, _clean(ln.get("note")) or None,
                      1 if mode else (1 if ln.get("no_loss") else 0)))
    if mode and len(clean) != 1:
        raise ValueError(f"{code}: สูตรแบ่งชิ้นต้องมีวัตถุดิบ 1 บรรทัด (ของทั้งตัวที่นำมาแบ่ง)")
    old = con.execute("SELECT batch_yield_qty, batch_unit FROM bc_prep_recipe WHERE prepared_code=?", (code,)).fetchone()
    con.execute("INSERT INTO bc_prep_recipe(prepared_code, batch_yield_qty, batch_unit, note, updated_at, source_cost,"
                " source_ref, updated_by, mode) VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(prepared_code) DO UPDATE SET"
                " batch_yield_qty=excluded.batch_yield_qty, batch_unit=excluded.batch_unit, updated_at=excluded.updated_at,"
                " source_cost=excluded.source_cost, source_ref=excluded.source_ref, updated_by=excluded.updated_by,"
                " mode=excluded.mode, note=CASE WHEN excluded.mode IS NOT NULL THEN excluded.note"
                " ELSE COALESCE(excluded.note, bc_prep_recipe.note) END",
                (code, yq, norm(yield_unit), _clean(note) or None, now(), _num(source_cost), source_ref, who, mode))
    con.execute("DELETE FROM bc_prep_line WHERE prepared_code=?", (code,))
    con.executemany("INSERT INTO bc_prep_line(prepared_code, seq, ingredient_code, qty, unit, note, no_loss) VALUES(?,?,?,?,?,?,?)",
                    [(code,) + c for c in clean])
    audit(con, who, "prep_save", "bc_prep_recipe", code, old, {"yield": [yq, norm(yield_unit)], "lines": len(clean),
                                                                 "source": source_ref, "mode": mode or "batch"})
    return len(clean)


PORTION_SKIP_NOTE = "สูตรแบ่งชิ้น ไม่ถูกแทนที่โดยไฟล์นำเข้า แก้ไขที่การ์ดสูตรแบ่งชิ้น"


def portion_codes(con) -> set:
    """The prepared codes whose recipe is a portion recipe (PR-U2). A file import (BoM_Semi or the flat template) never replaces
    one: it would turn "1 ตัว → 225 ชิ้น" back into a batch recipe with the trim loss added; the card is where it is edited."""
    return {c for (c,) in con.execute("SELECT prepared_code FROM bc_prep_recipe WHERE mode=?", (MODE_PORTION,))}


def apply_import(con, plan: dict, who: str) -> dict:
    n, skipped = 0, []
    for w in plan["writes"]:
        if w.get("portion_skip"):
            skipped.append({"code": w["prepared_code"], "name": w.get("prepared_name"), "note": PORTION_SKIP_NOTE})
            continue
        if not w["yield_qty"]:
            continue
        save(con, w["prepared_code"], w["yield_qty"], w["yield_unit"], w["lines"], who,
             source_ref=f"BoM_Semi แถว {w['row']}: {w['src_code']} {w['src_name']}", source_cost=w["net_cost"])
        n += 1
    return {"saved": n, "skipped": skipped}


def set_map(con, src_name: str, code: str, who: str):
    k = _key(src_name)
    if not k:
        raise ValueError("ไม่ได้ระบุชื่อ")
    code = (code or "").strip()
    if code != SKIP and not con.execute("SELECT 1 FROM bc_ingredient WHERE code=?", (code,)).fetchone():
        raise ValueError(f"ไม่พบรหัส {code} ในทะเบียน")
    con.execute("INSERT INTO bc_prep_name_map(name_key, prepared_code, mapped_by, mapped_at) VALUES(?,?,?,?)"
                " ON CONFLICT(name_key) DO UPDATE SET prepared_code=excluded.prepared_code,"
                " mapped_by=excluded.mapped_by, mapped_at=excluded.mapped_at", (k, code, who, now()))
    audit(con, who, "prep_map", "bc_prep_name_map", k, None, code)


# ── the set of sub-recipes a cost uses ───────────────────────────────────────
def _read(con, where: str, args: tuple, rtable: str, ltable: str) -> dict:
    book = {}
    mcol = "mode" if _has_col(con, rtable, "mode") else "NULL"
    ncol = "no_loss" if _has_col(con, ltable, "no_loss") else "0"
    fcol = "factor_base" if _has_col(con, ltable, "factor_base") else "NULL"      # PR-U2: frozen with a released version
    for code, yq, yu, mode in con.execute(f"SELECT prepared_code, batch_yield_qty, batch_unit, {mcol} FROM {rtable}{where}", args):
        book[code] = {"yield_qty": yq, "yield_unit": yu, "mode": mode or "batch", "lines": []}
    for code, seq, ic, q, u, note, nl, fb in con.execute(
            f"SELECT prepared_code, seq, ingredient_code, qty, unit, note, {ncol}, {fcol} FROM {ltable}{where}"
            f" ORDER BY prepared_code, seq", args):
        if code in book:
            book[code]["lines"].append({"seq": seq, "ingredient_code": ic, "qty": q, "unit": u, "note": note,
                                        "no_loss": 1 if nl else 0, "factor_base": fb})
    return book


def _has_col(con, table: str, col: str) -> bool:
    return any(r[1] == col for r in con.execute(f"PRAGMA table_info({table})"))


def book(con, release_id: int = None) -> dict:
    """{prepared_code: {yield_qty, yield_unit, lines}}. A released (or retired) release
    reads the set frozen into it — none frozen (released before Phase 1.5) = none,
    so its costs never move; a draft, or no release, reads the live sub-recipes."""
    if not con.execute("SELECT 1 FROM sqlite_master WHERE name='bc_rel_prep_recipe'").fetchone():
        return {}
    if release_id is not None:
        st = con.execute("SELECT status FROM bc_release WHERE release_id=?", (release_id,)).fetchone()
        if st and st[0] in ("released", "retired"):
            return _read(con, " WHERE release_id=?", (release_id,), "bc_rel_prep_recipe", "bc_rel_prep_line")
    return _read(con, "", (), "bc_prep_recipe", "bc_prep_line")


def freeze(con, release_id: int) -> int:
    """Copy the live sub-recipes into a release (at release time)."""
    con.execute("DELETE FROM bc_rel_prep_line WHERE release_id=?", (release_id,))
    con.execute("DELETE FROM bc_rel_prep_recipe WHERE release_id=?", (release_id,))
    con.execute("INSERT INTO bc_rel_prep_recipe(release_id, prepared_code, batch_yield_qty, batch_unit, mode)"
                " SELECT ?, prepared_code, batch_yield_qty, batch_unit, mode FROM bc_prep_recipe", (release_id,))
    con.execute("INSERT INTO bc_rel_prep_line(release_id, prepared_code, seq, ingredient_code, qty, unit, note, no_loss)"
                " SELECT ?, prepared_code, seq, ingredient_code, qty, unit, note, COALESCE(no_loss, 0) FROM bc_prep_line",
                (release_id,))
    # L6 (PR-U2): a line that passes an average step (or is a portion line) keeps the factor it is costed with, so a later
    # change of the average (learned from receipts, approved by the analyst) never moves a released month
    conv = Converter(con)
    for seq, code, ic, unit, nl in con.execute(
            "SELECT seq, prepared_code, ingredient_code, unit, COALESCE(no_loss, 0) FROM bc_rel_prep_line WHERE release_id=?",
            (release_id,)).fetchall():
        f, _how = conv.factor(ic, unit) if ic and unit else (None, None)
        if f is not None and (nl or conv.last_approx):
            con.execute("UPDATE bc_rel_prep_line SET factor_base=? WHERE release_id=? AND prepared_code=? AND seq=?",
                        (f, release_id, code, seq))
    return con.execute("SELECT COUNT(*) FROM bc_rel_prep_recipe WHERE release_id=?", (release_id,)).fetchone()[0]


# ── costing ──────────────────────────────────────────────────────────────────
class PrepCoster:
    """Cost per base unit of a prepared item from its sub-recipe, for one price month."""

    def __init__(self, con, month: str, book_: dict, conv: Converter = None):
        self.con, self.month, self.book = con, month, book_ or {}
        self.conv = conv or Converter(con)
        self.loss = dict(con.execute("SELECT code, loss_pct FROM bc_ingredient"))
        self._memo, self._p = {}, {}

    def _price(self, code):
        if code not in self._p:
            self._p[code] = price_for(self.con, code, self.month)
        return self._p[code]

    def cost(self, code: str, _stack: tuple = ()) -> dict:
        """-> {per_base, batch_cost, yield_base, problems: [..], lines: [..]}; per_base None when incomplete."""
        if code in self._memo:
            return self._memo[code]
        r = self.book.get(code)
        out = {"code": code, "per_base": None, "batch_cost": 0.0, "yield_base": None, "problems": [], "lines": [],
               "mode": (r or {}).get("mode"), "approx": False}
        if r is None:
            out["problems"].append("ไม่มีสูตรย่อย")
            return out
        if code in _stack:
            out["problems"].append("สูตรย่อยอ้างอิงวนกลับมาที่ตัวเอง: " + " → ".join(_stack + (code,)))
            return out
        yb, how = self.conv.to_base(code, r["yield_qty"], r["yield_unit"])
        if not yb:
            out["problems"].append(f"แปลงปริมาณที่ได้ต่อสูตร {r['yield_qty']:g} {r['yield_unit']} เป็นหน่วยฐานไม่ได้ ({how})")
        out["yield_base"] = yb
        for ln in r["lines"]:
            ic = ln["ingredient_code"]
            row = dict(ln, cost=None, price_per_base=None, qty_base=None, source=None, problem=None, approx=False,
                       no_loss=1 if ln.get("no_loss") else 0)
            if not ic:
                row["problem"] = f"ไม่พบรหัสวัตถุดิบของ “{ln.get('note') or ''}”"
            else:
                qb, how = self.conv.to_base(ic, ln["qty"], ln["unit"])
                if ln.get("factor_base"):                          # frozen with a released version (L6): the live average does not count
                    qb, how = float(ln["qty"]) * float(ln["factor_base"]), "ค่าที่ตรึงไว้ตามเวอร์ชันที่ประกาศใช้"
                if qb is None:
                    row["problem"] = f"{ic}: {how}"
                else:
                    row["qty_base"] = qb
                    row["approx"], row["how"] = self.conv.last_approx, how    # an average step is part of the chain
                    if ic in self.book:
                        sub = self.cost(ic, _stack + (code,))
                        p, src = sub["per_base"], "sub_recipe"
                        row["approx"] = row["approx"] or bool(sub.get("approx"))
                        if p is None:
                            row["problem"] = f"{ic}: คำนวณต้นทุนสูตรย่อยไม่ครบ: " + "; ".join(sub["problems"][:2])
                    else:
                        pr = self._price(ic)
                        p, src = (pr[0], pr[1]) if pr else (None, None)
                        if p is None:
                            row["problem"] = f"{ic}: ไม่มีราคาเดือน {self.month}"
                    if p is not None:
                        # a portion line (PR-U) is the whole piece and its yield already counts only usable pieces: no trim loss on top
                        loss = 0.0 if row["no_loss"] else min(self.loss.get(ic) or 0.0, 99.0)
                        row.update(price_per_base=p, source=src, cost=qb * p / (1.0 - loss / 100.0))
                        out["batch_cost"] += row["cost"]
            if row["problem"]:
                out["problems"].append(row["problem"])
            out["approx"] = out["approx"] or row["approx"]
            out["lines"].append(row)
        if not out["problems"] and yb:
            out["per_base"] = out["batch_cost"] / yb
        self._memo[code] = out
        return out


def explode_parts(book_: dict, conv: Converter, code: str, _stack: tuple = (), _memo: dict = None) -> tuple:
    """Raw ingredients in ONE base unit of a prepared item, split by whether the raw ingredient's trim loss still has to be
    added on top: ({raw_code: qty that takes the loss}, {raw_code: qty that does not}, problems). The second dict holds
    the quantities that come through a portion line (`no_loss`): a salmon cut into 225 pieces already counts only the
    usable meat, so ordering must not gross it up by the salmon's 40 % trim loss a second time (PR-U)."""
    _memo = {} if _memo is None else _memo
    if code in _memo:
        return _memo[code]
    r = book_.get(code)
    if r is None or code in _stack:
        return {}, {}, ["ไม่มีสูตรย่อย" if r is None else "สูตรย่อยอ้างอิงวนกลับมาที่ตัวเอง"]
    yb, _ = conv.to_base(code, r["yield_qty"], r["yield_unit"])
    if not yb:
        return {}, {}, ["แปลงปริมาณที่ได้ต่อสูตรไม่ได้"]
    lossy, flat, probs = {}, {}, []
    for ln in r["lines"]:
        ic = ln["ingredient_code"]
        qb, _ = conv.to_base(ic, ln["qty"], ln["unit"]) if ic else (None, None)
        if ic and ln.get("factor_base"):
            qb = float(ln["qty"]) * float(ln["factor_base"])      # frozen with a released version (L6)
        if qb is None:
            probs.append(ln.get("note") or ic or "?")
            continue
        if ic in book_:
            sl, sf, p2 = explode_parts(book_, conv, ic, _stack + (code,), _memo)
            if not sl and not sf:
                probs.append(ic)
                continue
            probs += p2
            for k, v in sl.items():
                lossy[k] = lossy.get(k, 0.0) + v * qb / yb
            for k, v in sf.items():
                flat[k] = flat.get(k, 0.0) + v * qb / yb
        elif ln.get("no_loss"):
            flat[ic] = flat.get(ic, 0.0) + qb / yb
        else:
            lossy[ic] = lossy.get(ic, 0.0) + qb / yb
    _memo[code] = (lossy, flat, probs)
    return lossy, flat, probs


def explode(book_: dict, conv: Converter, code: str, _stack: tuple = (), _memo: dict = None) -> tuple:
    """Raw ingredients in ONE base unit of a prepared item: ({raw_code: qty_base}, problems). Trim loss is not in this
    number (see explode_parts for the split the order plan needs)."""
    lossy, flat, probs = explode_parts(book_, conv, code, _stack, _memo)
    out = dict(lossy)
    for k, v in flat.items():
        out[k] = out.get(k, 0.0) + v
    return out, probs


def overview(con, month: str) -> list:
    """Every prepared code: its sub-recipe, the computed cost and the purchase price."""
    bk = book(con)
    pc = PrepCoster(con, month, bk)
    rows = []
    for code, name, base, per, recv in con.execute(
            "SELECT code, name_th, base_unit, qty_per_recv, recv_unit FROM bc_ingredient WHERE kind='prepared'"
            " AND active=1 ORDER BY code"):
        meta = con.execute("SELECT source_ref, source_cost, updated_at, updated_by FROM bc_prep_recipe"
                           " WHERE prepared_code=?", (code,)).fetchone()
        r = bk.get(code)
        own = price_for(con, code, month)
        row = {"code": code, "name": name, "base_unit": base, "qty_per_recv": per, "recv_unit": recv,
               "has_recipe": r is not None, "own_price": own[0] if own else None,
               "own_source": own[1] if own else None}
        if r:
            c = pc.cost(code)
            row.update(yield_qty=r["yield_qty"], yield_unit=r["yield_unit"], lines=len(r["lines"]),
                       per_base=c["per_base"], batch_cost=round(c["batch_cost"], 4), problems=c["problems"][:5],
                       n_problems=len(c["problems"]), source_ref=meta[0] if meta else None,
                       source_cost=meta[1] if meta else None, updated_at=meta[2] if meta else None,
                       updated_by=meta[3] if meta else None, mode=r.get("mode") or "batch",
                       mode_th=MODE_TH.get(r.get("mode") or "batch"), approx=bool(c["approx"]))
            if c["per_base"] and row["own_price"]:
                row["diff_pct"] = round(100.0 * (c["per_base"] - row["own_price"]) / row["own_price"], 1)
        row["status"] = ("none" if not r else "incomplete" if row.get("n_problems") else
                         "check" if abs(row.get("diff_pct") or 0) > TOLERANCE_PCT else "ok")
        rows.append(row)
    order = {"incomplete": 0, "check": 1, "none": 2, "ok": 3}
    rows.sort(key=lambda r: (order[r["status"]], r["code"]))
    return rows


def detail(con, code: str, month: str) -> dict:
    bk = book(con)
    r = bk.get(code)
    m = con.execute("SELECT name_th, base_unit, qty_per_recv, recv_unit FROM bc_ingredient WHERE code=?", (code,)).fetchone()
    if not m:
        raise ValueError(f"ไม่พบรหัส {code}")
    names = dict(con.execute("SELECT code, name_th FROM bc_ingredient"))
    out = {"code": code, "name": m[0], "base_unit": m[1], "qty_per_recv": m[2], "recv_unit": m[3],
           "has_recipe": r is not None, "month": month}
    if r:
        c = PrepCoster(con, month, bk).cost(code)
        for ln in c["lines"]:
            ln["ingredient_name"] = names.get(ln["ingredient_code"])
        raw, probs = explode(bk, Converter(con), code)
        out.update(yield_qty=r["yield_qty"], yield_unit=r["yield_unit"], lines=c["lines"], per_base=c["per_base"],
                   batch_cost=round(c["batch_cost"], 4), yield_base=c["yield_base"], problems=c["problems"],
                   raw_per_base=[{"code": k, "name": names.get(k), "qty_per_base": round(v, 6)}
                                 for k, v in sorted(raw.items(), key=lambda t: -t[1])], explode_problems=probs)
        meta = con.execute("SELECT source_ref, source_cost, updated_at, updated_by FROM bc_prep_recipe"
                           " WHERE prepared_code=?", (code,)).fetchone()
        out.update(dict(zip(("source_ref", "source_cost", "updated_at", "updated_by"), meta)))
        out.update(mode=r.get("mode") or "batch", mode_th=MODE_TH.get(r.get("mode") or "batch"), approx=bool(c["approx"]),
                   recipe_note=con.execute("SELECT note FROM bc_prep_recipe WHERE prepared_code=?", (code,)).fetchone()[0])
    own = price_for(con, code, month)
    out["own_price"] = own[0] if own else None
    return out


FLAT_COLS = ("prepared_code", "prepared_name", "batch_yield_qty", "batch_unit", "seq", "ingredient_code",
             "ingredient_name", "qty", "unit", "note", "mode", "no_loss")     # mode / no_loss (PR-U): a portion recipe survives the round trip


def export_flat(con) -> list:
    names = dict(con.execute("SELECT code, name_th FROM bc_ingredient"))
    rows = []
    for code, r in sorted(book(con).items()):
        for ln in r["lines"] or [{"seq": None, "ingredient_code": None, "qty": None, "unit": None, "note": None}]:
            rows.append([code, names.get(code), r["yield_qty"], r["yield_unit"], ln["seq"], ln["ingredient_code"],
                         names.get(ln["ingredient_code"]), ln["qty"], ln["unit"], ln["note"],
                         MODE_PORTION if r.get("mode") == MODE_PORTION else None, 1 if ln.get("no_loss") else None])
    return rows


def import_flat(con, data: bytes, who: str, dry_run: bool = True) -> dict:
    """The flat template back in: one row per line, grouped by prepared_code (replaces those codes)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(rows[:5]) if "prepared_code" in [_clean(v) for v in r]), None)
    if hdr is None:
        raise ValueError("ไม่พบหัวคอลัมน์ prepared_code · ใช้ไฟล์จากปุ่ม “ดาวน์โหลดแม่แบบ”")
    h = {_clean(v): j for j, v in enumerate(rows[hdr]) if v}
    grp, errors = {}, []
    for i, r in enumerate(rows[hdr + 1:], start=hdr + 2):
        def g(k):
            j = h.get(k)
            return r[j] if j is not None and j < len(r) else None
        code = _clean(g("prepared_code"))
        if not code:
            continue
        e = grp.setdefault(code, {"yield_qty": g("batch_yield_qty"), "yield_unit": g("batch_unit"), "lines": [],
                                  "mode": MODE_PORTION if _clean(g("mode")).lower() in (MODE_PORTION, "แบ่งชิ้น") else None})
        if _num(g("qty")) or _clean(g("ingredient_code")):
            e["lines"].append({"ingredient_code": g("ingredient_code"), "qty": g("qty"), "unit": g("unit"),
                               "note": g("note") or g("ingredient_name"), "no_loss": 1 if _num(g("no_loss")) else 0})
    saved, skipped = 0, []
    portion = portion_codes(con)
    for code, e in grp.items():
        if code in portion:                                  # a portion recipe is edited on its own card, never by a file
            skipped.append({"code": code, "name": None, "note": PORTION_SKIP_NOTE})
            continue
        try:
            save(con, code, e["yield_qty"], e["yield_unit"], e["lines"], who, source_ref="นำเข้าจากแม่แบบ (Excel)",
                 mode=e["mode"])
            saved += 1
        except ValueError as ex:
            errors.append(str(ex))
    if dry_run or errors:
        con.rollback()
    return {"codes": len(grp), "saved": 0 if (dry_run or errors) else saved, "errors": errors, "dry_run": dry_run,
            "skipped": skipped}


# ── สูตรแบ่งชิ้น: one whole piece (a salmon) cut into N pieces (PR-U) ────────────────────────────────────
def latest_price_month(con) -> str:
    r = con.execute("SELECT MAX(month) FROM bc_ingredient_price").fetchone()
    return r[0] if r and r[0] else now()[:7]


def _portion_inputs(con, raw_code, qty, unit, pieces):
    """Validate one portion input and convert it: -> (Converter, qty, pieces, factor, approx, how, raw row)."""
    q, n = _num(qty), _num(pieces)
    raw_code = _clean(raw_code)
    if not raw_code:
        raise ValueError("เลือกวัตถุดิบที่นำมาแบ่ง")
    if not q or q <= 0:
        raise ValueError("จำนวนที่ใช้ต้องมากกว่า 0")
    if not n or n <= 0:
        raise ValueError("จำนวนชิ้นที่ได้ต้องมากกว่า 0")
    row = con.execute("SELECT name_th, kind, base_unit FROM bc_ingredient WHERE code=?", (raw_code,)).fetchone()
    if not row:
        raise ValueError(f"ไม่พบรหัส {raw_code} ในทะเบียน")
    if row[1] == "prepared":
        raise ValueError(f"{raw_code} เป็นสินค้ากึ่งสำเร็จ · สูตรแบ่งชิ้นใช้วัตถุดิบดิบเท่านั้น")
    if not _clean(unit):
        raise ValueError("ระบุหน่วยที่ใช้")
    conv = Converter(con)
    f, how = conv.factor(raw_code, unit)
    if f is None:
        raise ValueError(f"{raw_code}: แปลงหน่วย {_clean(unit)} เป็นหน่วยฐานไม่ได้ ({how}) · เพิ่มหน่วยในหน้าหน่วยและการแปลงก่อน")
    return conv, q, n, f, conv.last_approx, how, row


def portion_preview(con, raw_code, qty, unit, pieces, month: str = None) -> dict:
    """What a portion recipe means, before it is saved: {factor_base, approx, how, range, base_per_piece, cost_per_piece, ...}.
    Cost per piece = price per base × factor(unit) × qty ÷ pieces — the raw item's trim loss is NOT added (the piece
    count already counts only the usable pieces)."""
    conv, q, n, f, approx, how, row = _portion_inputs(con, raw_code, qty, unit, pieces)
    raw_code = _clean(raw_code)
    month = month or latest_price_month(con)
    qb = q * f
    out = {"raw_code": raw_code, "raw_name": row[0], "qty": q, "unit": norm(unit), "pieces": n, "base_unit": conv.base_unit(raw_code),
           "factor_base": f, "qty_base": qb, "approx": bool(approx), "how": how, "base_per_piece": qb / n,
           "range": None, "cost_per_piece": None, "batch_cost": None, "price_per_base": None, "price_month": None,
           "price_source": None, "month": month}
    sp = conv.span(raw_code, unit)
    if sp:
        out["range"] = {"qty_base_min": q * sp[0], "qty_base_max": q * sp[1],
                        "per_piece_min": q * sp[0] / n, "per_piece_max": q * sp[1] / n}
    pr = price_for(con, raw_code, month)
    if pr:
        out.update(price_per_base=pr[0], price_source=pr[1], price_month=pr[2], batch_cost=qb * pr[0],
                   cost_per_piece=qb * pr[0] / n)
        if sp:
            out["range"].update(cost_per_piece_min=q * sp[0] * pr[0] / n, cost_per_piece_max=q * sp[1] * pr[0] / n)
    return out


def ensure_prepared(con, code: str, name: str, base_unit: str, who: str) -> bool:
    """A prepared item the master does not have yet is created for a portion recipe (kind prepared, family CK, base unit =
    the unit the pieces are counted in, bought locally). -> True when it was created."""
    from .master import _upsert
    code = _clean(code)
    if not code:
        raise ValueError("ระบุรหัสสินค้ากึ่งสำเร็จ")
    row = con.execute("SELECT kind FROM bc_ingredient WHERE code=?", (code,)).fetchone()
    if row:
        if row[0] != "prepared":
            raise ValueError(f"{code} ไม่ใช่สินค้ากึ่งสำเร็จ · เลือกรหัสกึ่งสำเร็จ หรือสร้างรหัสใหม่")
        return False
    nm = _clean(name)
    if not nm:
        raise ValueError("ใส่ชื่อของสินค้ากึ่งสำเร็จที่สร้างใหม่")
    stats = {"inserted": 0, "updated": 0, "unchanged": 0}
    _upsert(con, code, {"name_th": nm, "kind": "prepared", "family": "CK", "base_unit": norm(base_unit) or "ชิ้น",
                        "sourcing_default": "local", "note": "สร้างจากสูตรแบ่งชิ้น"}, who, False, stats)
    audit(con, who, "prep_portion_new_item", "bc_ingredient", code, None, nm)
    return True


def save_portion(con, code: str, raw_code: str, qty, unit, pieces, out_unit, who: str, note: str = None, name: str = None,
                 month: str = None) -> dict:
    """Write a portion recipe: ONE line (the whole piece) -> N pieces of the prepared item. Replaces any recipe the code had.
    Returns the preview (see portion_preview) plus {code, created, replaced}. Caller commits."""
    code, out_unit = _clean(code), norm(out_unit) or "ชิ้น"
    if code == _clean(raw_code):
        raise ValueError("สินค้ากึ่งสำเร็จต้องเป็นคนละรหัสกับวัตถุดิบที่นำมาแบ่ง")
    _portion_inputs(con, raw_code, qty, unit, pieces)
    created = ensure_prepared(con, code, name, out_unit, who)
    if Converter(con).factor(code, out_unit)[0] is None:
        raise ValueError(f"{code}: หน่วยที่ได้ “{out_unit}” ไม่ตรงกับหน่วยฐานของสินค้ากึ่งสำเร็จ")
    replaced = con.execute("SELECT 1 FROM bc_prep_recipe WHERE prepared_code=?", (code,)).fetchone() is not None
    save(con, code, pieces, out_unit, [{"ingredient_code": _clean(raw_code), "qty": qty, "unit": unit}], who,
         source_ref="สูตรแบ่งชิ้น", mode=MODE_PORTION, note=note)
    audit(con, who, "prep_portion", "bc_prep_recipe", code, None,
          {"from": _clean(raw_code), "qty": _num(qty), "unit": norm(unit), "pieces": _num(pieces), "created": created})
    out = portion_preview(con, raw_code, qty, unit, pieces, month)
    out.update(code=code, created=created, replaced=replaced)
    return out
