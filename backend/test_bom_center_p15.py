"""BoM Center — Phase 1.5: sub-recipes of prepared items (CK / HK / DK).

The promises:
  * the kitchen's BoM_Semi sheet is read as it is (batch yield "1 สูตร = 15,225 ml",
    Mat/Price/Usage columns, the "Cost per unit" lookup), and each row lands on the
    master's prepared code by NAME + BATCH SIZE — never by the sheet's own code, which
    is an older numbering (its CK-001 น้ำซุปใส is the master's HK-001);
  * the sheet counts liquids in g: g↔ml (and ชิ้น↔ฟอง) mean the master's base unit;
  * a prepared item with a complete sub-recipe is costed from its raw ingredients ×
    the month's prices — recursively — and a cycle or a gap is reported, not guessed;
  * checks: no sub-recipe (C10) and an incomplete one (C13) block but can be accepted
    with a reason; a computed cost > 5 % away from the purchase price warns (C14);
  * a release freezes the sub-recipes it was costed with: editing one later never
    moves a released month, while drafts see the edit;
  * the order plan orders the raw ingredients inside a prepared item, saying which
    prepared item each came through.
"""
import io
from datetime import date

import pytest

from bom_center import forecast as F
from bom_center import prep as P
from bom_center import recipes as R
from bom_center import release as RL
from bom_center import snapshot_v0 as V0
from test_bom_center_s2 import _db

DISH = "OTHBFGNT00010"           # sku 11


@pytest.fixture
def con(tmp_path):
    c = _db(tmp_path)
    c.executemany("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, qty_per_recv, loss_pct) VALUES(?,?,?,?,?,?)", [
        ("HK-001", "น้ำซุปใส HK-001", "prepared", "มิลลิลิตร", 1000, 0),
        ("HK-012", "แป้งทอดเทมปุระ HK-012", "prepared", "กรัม", 2600, 0),
        ("CK-001", "ข้าวสวยญี่ปุ่นปรุงรส CK-001", "prepared", "กรัม", 2200, 0),      # the master's own CK-001
        ("11-01-053", "มิริน 20 ลิตร", "raw", "มิลลิลิตร", 20000, 0),
        ("S-011", "เกลือป่น 1 กก.", "raw", "กรัม", 1000, 0),
        ("11-01-072", "แป้งเทมปุระ Tempuramix 1 กก.", "raw", "กรัม", 1000, 0),
        ("F-008", "ไข่ไก่ 30 ฟอง", "raw", "ฟอง", 30, 0)])
    c.executemany("INSERT INTO bc_ingredient_price(ingredient_code, month, price_per_base, source) VALUES(?,?,?,?)", [
        ("11-01-053", "2026-07", 0.08, "purchase_avg"), ("S-011", "2026-07", 0.01, "purchase_avg"),
        ("11-01-072", "2026-07", 0.03, "purchase_avg"), ("F-008", "2026-07", 4.0, "purchase_avg"),
        ("HK-001", "2026-07", 0.0081, "list_ref"),              # the supplier/reference price of the soup
        ("HK-012", "2026-07", 0.05, "list_ref")])               # way off what the sub-recipe costs → C14
    c.execute("INSERT INTO sku_cost(sku_id, unit_cost) VALUES(11, 1.0)")
    c.commit()
    V0.build_v0(c)
    c.commit()
    yield c
    c.close()


def _sheet() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BoM_Semi"
    ws.append(["Oshinei_BoM_Semi"])
    hdr = ["", "Groups", "Kitchen-Easy", "No", "Kitchen", "Code", "Menu", "Description", "Net Cost"]
    hdr += [f"Mat{i}" for i in range(1, 4)] + [f"Price{i}" for i in range(1, 4)] + [f"Usage{i}" for i in range(1, 4)]
    ws.append(hdr)

    def row(code, menu, desc, mats):
        m = [x[0] for x in mats] + [None] * (3 - len(mats))
        p = [x[1] for x in mats] + [None] * (3 - len(mats))
        u = [x[2] for x in mats] + [None] * (3 - len(mats))
        ws.append(["", "HK-0001", "Hot Kitchen", 1, "ร้อน", code, menu, desc, 8.1] + m + p + u)
    row("CK-001", "น้ำซุปใส", "1 สูตร = 1,000 ml", [("มิริน", 0.08, 100), ("เกลือป่น", 0.01, 10)])
    row("CK-010", "แป้งเทมปุระเปียก", "1 สูตร = 2,600 g", [("ผงแป้งเทมปุระ", 0.03, 1300), ("ไข่ไก่", 4, 2)])
    row("CK-099", "ซอสที่ไม่มีในทะเบียน", "1 สูตร = 500 ml", [("มิริน", 0.08, 50)])
    cpu = wb.create_sheet("Cost per unit")
    cpu.append([])
    cpu.append(["Items Code", "Mat-BoM", "ชื่อสินค้า", None, "หมายเหตุ", "BoM Unit", "Unit : BoM", "Yield", "ราคา", "@BoM"])
    cpu.append(["11-01-053", "มิริน", "มิริน 20L", "Gal", None, 20000, "G", 1, 1600, 0.08])   # liquid counted in G
    cpu.append(["S-011", "เกลือป่น", "เกลือ", "KG", None, 1000, "G", 1, 10, 0.01])
    cpu.append(["11-01-072", "ผงแป้งเทมปุระ", "Tempuramix", "KG", None, 1000, "G", 1, 30, 0.03])
    cpu.append(["F-008", "ไข่ไก่", "ไข่", "แผง", None, 30, "PCS", 1, 120, 4])              # eggs counted in ชิ้น
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _imported(con):
    plan = P.plan_import(con, P.parse_bom_semi(_sheet()))
    P.apply_import(con, plan, "kitchen")
    con.commit()
    return plan


def test_parse_yield():
    assert P.parse_yield("1 สูตร = 15,225 ml") == (15225.0, "มิลลิลิตร")
    assert P.parse_yield("1สูตร = 2,120 g") == (2120.0, "กรัม")
    assert P.parse_yield("1 สูตร = 5 PCS") == (5.0, "ชิ้น")
    assert P.parse_yield("ไม่มีตัวเลข") == (None, None)


def test_rows_land_on_master_codes_by_name_and_batch_not_by_sheet_code(con):
    parsed = P.parse_bom_semi(_sheet())
    assert parsed["sheet"] == "BoM_Semi" and len(parsed["blocks"]) == 3 and parsed["cpu"]["มิริน"][0] == "11-01-053"
    plan = P.plan_import(con, parsed)
    by = {w["prepared_code"]: w for w in plan["writes"]}
    assert set(by) == {"HK-001", "HK-012"}                    # never the master's CK-001 (a rice item)
    assert by["HK-001"]["src_code"] == "CK-001" and by["HK-001"]["method"] == "name+yield"
    assert by["HK-012"]["src_name"] == "แป้งเทมปุระเปียก"      # a different name, the same 2,600 g batch
    mirin = by["HK-001"]["lines"][0]
    assert mirin["ingredient_code"] == "11-01-053" and mirin["unit"] == "มิลลิลิตร"   # G in the sheet = ml
    egg = by["HK-012"]["lines"][1]
    assert egg["ingredient_code"] == "F-008" and egg["unit"] == "ฟอง"                 # PCS = ฟอง
    assert [u["name"] for u in plan["unmatched"]] == ["ซอสที่ไม่มีในทะเบียน"]
    assert "CK-001" in plan["prepared_without"]
    assert con.execute("SELECT COUNT(*) FROM bc_prep_recipe").fetchone()[0] == 0       # a plan writes nothing


def test_an_english_only_name_is_never_guessed(con):
    # "F-Salmon" has no Thai letters; it once matched an English-named master item (brisket)
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('03-01-319', 'Brisket Sliced 1.5 mm', 'raw', 'กรัม')")
    resolve = P._resolver(con, [], {})
    assert resolve("F-Salmon")[0] is None
    assert resolve("เกลือป่น")[0] == "S-011"


def test_a_person_can_map_or_skip_a_row(con):
    P.set_map(con, "ซอสที่ไม่มีในทะเบียน", "CK-001", "kitchen")
    plan = P.plan_import(con, P.parse_bom_semi(_sheet()))
    assert {w["prepared_code"]: w["method"] for w in plan["writes"]}["CK-001"] == "map"
    P.set_map(con, "ซอสที่ไม่มีในทะเบียน", P.SKIP, "kitchen")
    plan = P.plan_import(con, P.parse_bom_semi(_sheet()))
    assert not plan["unmatched"] and "CK-001" not in {w["prepared_code"] for w in plan["writes"]}
    with pytest.raises(ValueError):
        P.set_map(con, "ซอส", "NOPE-1", "kitchen")


def test_cost_from_raw_ingredients_and_nesting(con):
    _imported(con)
    pc = P.PrepCoster(con, "2026-07", P.book(con))
    assert pc.cost("HK-001")["per_base"] == pytest.approx((100 * 0.08 + 10 * 0.01) / 1000)
    assert pc.cost("HK-012")["per_base"] == pytest.approx((1300 * 0.03 + 2 * 4.0) / 2600)
    # a prepared item made from another one: 1 batch = 500 g from 250 g of batter + 100 g salt
    P.save(con, "CK-001", 500, "กรัม", [{"ingredient_code": "HK-012", "qty": 250, "unit": "กรัม"},
                                        {"ingredient_code": "S-011", "qty": 100, "unit": "กรัม"}], "kitchen")
    pc = P.PrepCoster(con, "2026-07", P.book(con))
    batter = (1300 * 0.03 + 2 * 4.0) / 2600
    assert pc.cost("CK-001")["per_base"] == pytest.approx((250 * batter + 100 * 0.01) / 500)
    raw, gaps = P.explode(P.book(con), pc.conv, "CK-001")
    assert not gaps
    assert raw["11-01-072"] == pytest.approx(250 / 500 * 1300 / 2600) and raw["S-011"] == pytest.approx(100 / 500)
    # a loop is reported, not followed
    P.save(con, "HK-012", 2600, "กรัม", [{"ingredient_code": "CK-001", "qty": 10, "unit": "กรัม"}], "kitchen")
    c = P.PrepCoster(con, "2026-07", P.book(con)).cost("CK-001")
    assert c["per_base"] is None and any("วน" in p for p in c["problems"])


def test_save_is_validated(con):
    with pytest.raises(ValueError):
        P.save(con, "HK-001", 0, "มิลลิลิตร", [], "k")
    with pytest.raises(ValueError):
        P.save(con, "HK-001", 1000, "", [], "k")
    with pytest.raises(ValueError):
        P.save(con, "HK-001", 1000, "มล.", [{"ingredient_code": "HK-001", "qty": 1, "unit": "มล."}], "k")
    with pytest.raises(ValueError):
        P.save(con, "NOPE", 1000, "มล.", [], "k")


def _draft_with_soup(con):
    rid = R.get_or_create_draft(con, "2026-11 draft", "2026-11", "2026-07", "analyst")
    R.write_recipe(con, rid, DISH, [{"ingredient_code": "HK-001", "ingredient_name": "น้ำซุปใส", "qty": 100,
                                     "unit": "มิลลิลิตร"},
                                    {"ingredient_code": "HK-012", "ingredient_name": "แป้ง", "qty": 20, "unit": "กรัม"}],
                   source="test")
    return rid


def _codes(s):
    return {r["check_code"]: r for r in s["rows"]}


def test_checks_c10_c13_c14(con):
    rid = _draft_with_soup(con)
    s = RL.run_checks(con, rid)
    c = _codes(s)
    assert c["C10"]["severity"] == "block" and c["C10"]["acceptable"] and not s["can_release"]
    _imported(con)
    c = _codes(RL.run_checks(con, rid))
    assert "C10" not in c and "C13" not in c
    assert "C14" in c and c["C14"]["ref"] == "HK-012"           # ฿0.018 from the recipe vs ฿0.05 bought
    # the dish is costed from the sub-recipes
    cost = R.cost_recipe(con, rid, DISH)
    soup = (100 * 0.08 + 10 * 0.01) / 1000
    batter = (1300 * 0.03 + 2 * 4.0) / 2600
    assert cost["total"] == pytest.approx(100 * soup + 20 * batter)
    assert cost["lines"][0]["price_source"] == "sub_recipe"
    # an incomplete sub-recipe falls back to the purchase price and blocks until accepted
    con.execute("DELETE FROM bc_ingredient_price WHERE ingredient_code='S-011'")
    s = RL.run_checks(con, rid)
    c = _codes(s)
    assert c["C13"]["ref"] == "HK-001" and not s["can_release"]
    assert R.cost_recipe(con, rid, DISH)["lines"][0]["price_source"] == "list_ref"
    RL.accept(con, rid, "C13", "HK-001", "ครัวยังไม่ส่งราคาเกลือ ใช้ราคาอ้างอิงไปก่อน", "analyst")
    assert RL.run_checks(con, rid)["can_release"]


def test_a_release_freezes_its_sub_recipes(con):
    rid = _draft_with_soup(con)
    _imported(con)
    out = RL.do_release(con, rid, "analyst", today=date(2026, 10, 6))
    assert out["released"] and out["sub_recipes"] == 2
    before = R.release_costs(con, rid)[DISH][0]
    # the kitchen doubles the mirin afterwards
    P.save(con, "HK-001", 1000, "มิลลิลิตร", [{"ingredient_code": "11-01-053", "qty": 200, "unit": "มิลลิลิตร"},
                                              {"ingredient_code": "S-011", "qty": 10, "unit": "กรัม"}], "kitchen")
    assert R.release_costs(con, rid)[DISH][0] == pytest.approx(before)               # the released month stays
    rid2 = RL.new_draft_from(con, rid, "2026-12 draft", "2026-12", "2026-07", "analyst")
    assert R.release_costs(con, rid2)[DISH][0] > before                             # a new draft sees the edit


def test_a_release_before_phase_15_keeps_purchase_prices(con):
    rid = _draft_with_soup(con)
    con.execute("UPDATE bc_release SET status='released' WHERE release_id=?", (rid,))   # released, nothing frozen
    _imported(con)
    assert P.book(con, rid) == {}
    assert R.cost_recipe(con, rid, DISH)["lines"][0]["price_source"] == "list_ref"


def test_the_order_plan_orders_what_is_inside(con):
    _imported(con)
    con.execute("INSERT INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name, qty_per_serving, unit,"
                " price_per_unit, ingredient_cost, ingredient_type) VALUES(11, 1, 'HK-001', 'น้ำซุปใส', 100, 'มิลลิลิตร',"
                " 0.0081, 0.81, 'sauce')")
    per, stats = F._per_dish(con, "2026-10")
    got = dict(per[11])
    assert "HK-001" not in got
    assert got["11-01-053"] == pytest.approx(100 * 100 / 1000) and got["S-011"] == pytest.approx(100 * 10 / 1000)
    assert stats["exploded_lines"] == 1 and stats["via"]["11-01-053"] == ["HK-001"]


def test_flat_template_round_trip(con):
    _imported(con)
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(P.FLAT_COLS))
    for r in P.export_flat(con):
        ws.append(r)
    ws.append(["CK-001", None, 500, "กรัม", 1, "S-011", None, 5, "กรัม", None])
    buf = io.BytesIO()
    wb.save(buf)
    out = P.import_flat(con, buf.getvalue(), "kitchen", dry_run=True)
    assert out["codes"] == 3 and not out["errors"] and out["saved"] == 0
    assert not con.execute("SELECT 1 FROM bc_prep_recipe WHERE prepared_code='CK-001'").fetchone()
    out = P.import_flat(con, buf.getvalue(), "kitchen", dry_run=False)
    con.commit()
    assert out["saved"] == 3
    assert P.book(con)["CK-001"]["lines"][0]["ingredient_code"] == "S-011"
    assert len(P.book(con)["HK-001"]["lines"]) == 2


def test_overview_statuses(con):
    _imported(con)
    rows = {r["code"]: r for r in P.overview(con, "2026-07")}
    assert rows["HK-001"]["status"] == "ok" and rows["HK-001"]["diff_pct"] == pytest.approx(0.0, abs=0.1)
    assert rows["HK-012"]["status"] == "check"
    assert rows["CK-001"]["status"] == "none"
