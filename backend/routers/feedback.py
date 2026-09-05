# -*- coding: utf-8 -*-
"""💬 เสียงลูกค้า — customer feedback (pillar 3 of the menu programme, 2026-09).

Reviews reach `customer_feedback` three ways — pasted or uploaded in the tab
(CSV/TSV/TXT, XLSX, a Google Takeout zip or reviews.json), or posted by the
nightly fetcher (backend/fetch_feedback.py: Google Business Profile, Facebook
Pages) — and each is analysed for sentiment, themes and the menus it mentions
(feedback_nlp: Claude when ANTHROPIC_API_KEY is set, the Thai lexicon
otherwise; the row records which). The summary turns them into what the menu
tabs cannot see from POS data: what guests say about the dishes and the branch.

  POST   /api/feedback/upload      file / pasted text / JSON rows → analysed, inserted or updated
  GET    /api/feedback/summary     counts, sentiment, themes, menus, branches, recent negatives
  GET    /api/feedback/items       the rows behind a summary figure
  GET    /api/feedback/sources     per source: rows, latest review, the fetcher's last run
  POST   /api/feedback/sync-log    the fetcher's run record
  GET    /api/feedback/{id}/reply-draft   a suggested reply (Claude when a key is set, else a Thai template)
  POST   /api/feedback/{id}/reply         queue a reply — the fetcher posts it to Google / a Facebook comment
  DELETE /api/feedback/{id}/reply         cancel a queued reply
  GET    /api/feedback/reply-queue        what the fetcher must send (status=pending|failed|manual|sent)
  POST   /api/feedback/reply-result       the fetcher reports sent / failed
  DELETE /api/feedback/{id}
RBAC module customer.feedback (auth.py); writes need the write token like every
upload. The tables are carried across the daily DB push (core._PRESERVED_EXACT /
_PRESERVED_PREFIXES): the live site is the only place these rows exist.
"""
import csv
import datetime
import hashlib
import hmac
import io
import json
import os
import zipfile
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse

import auth as _auth
from core import get_db, require_write_token
from feedback_nlp import analyse, clean_menu_name, draft_reply
from feedback_sources.google_gbp import parse_reviews_json, parse_takeout_zip

router = APIRouter()

DDL = """
CREATE TABLE IF NOT EXISTS customer_feedback(
  id          INTEGER PRIMARY KEY AUTOINCREMENT,
  branch_id   INTEGER,
  branch_name TEXT,
  source      TEXT,
  rated_at    TEXT,
  rating      REAL,
  text        TEXT NOT NULL,
  sentiment   TEXT,
  score       REAL,
  themes      TEXT,
  menus       TEXT,
  analysed_by TEXT,
  uploaded_at TEXT DEFAULT (datetime('now','+7 hours')),
  uploaded_by TEXT,
  fingerprint TEXT UNIQUE
);
CREATE INDEX IF NOT EXISTS ix_cf_branch_date ON customer_feedback(branch_id, rated_at);
CREATE TABLE IF NOT EXISTS feedback_sync_log(
  id        INTEGER PRIMARY KEY AUTOINCREMENT,
  source    TEXT NOT NULL,
  ran_at    TEXT DEFAULT (datetime('now','+7 hours')),
  since     TEXT,
  fetched   INTEGER DEFAULT 0,
  inserted  INTEGER DEFAULT 0,
  updated   INTEGER DEFAULT 0,
  unchanged INTEGER DEFAULT 0,
  status    TEXT,
  message   TEXT,
  by        TEXT
);
CREATE TABLE IF NOT EXISTS feedback_webhook_events(
  id           INTEGER PRIMARY KEY AUTOINCREMENT,
  received_at  TEXT DEFAULT (datetime('now','+7 hours')),
  platform     TEXT,
  object_id    TEXT,
  field        TEXT,
  payload      TEXT,
  processed_at TEXT,
  result       TEXT
);
"""
# columns added after the first release — ALTERed in when missing (the live DB
# keeps its rows across deploys, so CREATE IF NOT EXISTS alone never adds them)
_EXTRA_COLS = (("external_id", "TEXT"), ("url", "TEXT"), ("author_hash", "TEXT"), ("reply_text", "TEXT"),
               ("replied_at", "TEXT"), ("location", "TEXT"), ("updated_at", "TEXT"),
               # the reply queue: external_ref = where a reply is posted (Google review resource
               # name / Facebook comment id); reply_status = pending → sent | failed, or manual
               # (Facebook recommendations cannot be answered through the API)
               ("external_ref", "TEXT"), ("reply_draft", "TEXT"), ("reply_status", "TEXT"),
               ("reply_requested_at", "TEXT"), ("reply_requested_by", "TEXT"), ("reply_error", "TEXT"),
               # review = on the chain's own listing/page; comment = under the chain's own post/video;
               # mention = anywhere else (X, YouTube, news, blogs) — no rating, no reply channel
               ("kind", "TEXT"),
               # flag = a serious case pinned by a manager (food safety, fraud…) with the reason
               ("flag", "INTEGER DEFAULT 0"), ("flag_note", "TEXT"),
               # contact = a QR-survey guest's phone / LINE id for the call back (blanked after survey.CONTACT_DAYS)
               ("contact", "TEXT"),
               # E5: severity high|med|low (+ the reason when high) and the staff named in the voice (JSON list)
               ("severity", "TEXT"), ("severity_note", "TEXT"), ("staff", "TEXT"))
SLA_DAYS = 1                # a negative voice waits at most this long for the branch's answer
REPLY_STATUSES = ("pending", "sent", "failed", "manual")
KINDS = ("review", "comment", "mention", "message", "survey")   # message = a private inbox conversation (Messenger / IG DM), per day; survey = the QR form (routers/survey.py)
REPLY_MAX = {"google": 4096, "facebook": 8000, "tiktok": 150, "instagram": 2200}   # characters the platform accepts; others 4000

# header names the CSV/TSV/XLSX/JSON import understands (lower-cased, Thai or English)
_COLS = {
    "date": ("date", "วันที่", "rated_at", "time", "datetime", "created"),
    "branch": ("branch", "สาขา", "store", "location", "ร้าน"),
    "source": ("source", "ช่องทาง", "แหล่ง", "platform", "channel"),
    "rating": ("rating", "คะแนน", "ดาว", "stars", "score"),
    "text": ("text", "review", "ข้อความ", "รีวิว", "comment", "ความคิดเห็น", "content", "body"),
    "external_id": ("external_id", "review_id", "reviewid", "id"),
    "url": ("url", "link", "ลิงก์"),
    "reply": ("reply", "response", "ตอบกลับ", "คำตอบ"),
}
MAX_ROWS = 5000


def _thai_today() -> str:
    return (datetime.datetime.utcnow() + datetime.timedelta(hours=7)).strftime("%Y-%m-%d")


def _thai_now() -> str:
    return (datetime.datetime.utcnow() + datetime.timedelta(hours=7)).strftime("%Y-%m-%d %H:%M:%S")


def _ensure(con):
    con.executescript(DDL)
    have = {r[1] for r in con.execute("PRAGMA table_info(customer_feedback)")}
    missing = [(c, t) for c, t in _EXTRA_COLS if c not in have]
    for c, t in missing:
        con.execute(f"ALTER TABLE customer_feedback ADD COLUMN {c} {t}")
    con.execute("CREATE INDEX IF NOT EXISTS ix_cf_ext ON customer_feedback(source, external_id)")
    if missing:
        con.commit()


# ── parsing ──────────────────────────────────────────────────────────────────
def _col_map(header):
    out = {}
    for i, h in enumerate(header):
        hl = str(h if h is not None else "").strip().lower()
        for key, names in _COLS.items():
            if key not in out and any(hl == n or hl.startswith(n) for n in names):
                out[key] = i
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _rows_from_table(table: list, default_branch, default_source, default_date):
    """A header + rows grid → rows, or None when no header names a text column."""
    cmap = _col_map(table[0]) if table else {}
    if "text" not in cmap:
        return None
    rows = []
    for r in table[1:]:
        if not r or len(r) <= cmap["text"]:
            continue
        t = _cell(r[cmap["text"]])
        g = lambda k: (_cell(r[cmap[k]]) if k in cmap and len(r) > cmap[k] else "")  # noqa: E731
        if not t and not (g("external_id") and g("rating")):
            continue
        rows.append({"date": g("date") or default_date, "branch": g("branch") or default_branch,
                     "source": g("source") or default_source, "rating": g("rating") or None, "text": t,
                     "external_id": g("external_id"), "url": g("url"), "reply": g("reply")})
    return rows


def _xlsx_table(data: bytes) -> list:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    table = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return table


def _rows_from_json(data: bytes, default_branch, default_source, default_date) -> list:
    obj = json.loads(data.decode("utf-8-sig"))
    items = obj.get("reviews") if isinstance(obj, dict) else obj
    if not isinstance(items, list):
        return []
    if any(isinstance(x, dict) and ("starRating" in x or "reviewId" in x or "reviewer" in x) for x in items):
        return parse_reviews_json(items, default_branch)          # Google review shape
    header = []
    for x in items:
        if isinstance(x, dict):
            header.extend(k for k in x if k not in header)
    table = [header] + [[x.get(k) for k in header] for x in items if isinstance(x, dict)]
    return _rows_from_table(table, default_branch, default_source, default_date) or []


def _kind(filename: str, data: bytes) -> str:
    fn = (filename or "").lower()
    if fn.endswith(".xlsx"):
        return "xlsx"
    if fn.endswith(".zip"):
        return "zip"
    if fn.endswith(".json"):
        return "json"
    if data[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                return "xlsx" if any(n.startswith("xl/") for n in z.namelist()) else "zip"
        except zipfile.BadZipFile:
            return "text"
    if data.lstrip()[:1] in (b"{", b"["):
        return "json"
    return "text"


def parse_rows(text: str = "", data: bytes = b"", filename: str = "",
               default_branch: str = "", default_source: str = "", default_date: str = "") -> list:
    """Pasted lines and/or a file (CSV/TSV/TXT, XLSX, JSON, Google Takeout zip)
    → normalised rows {date, branch, source, rating, text, external_id, …}."""
    rows = []
    if data:
        kind = _kind(filename, data)
        if kind == "zip":
            got, _locations = parse_takeout_zip(data, default_branch)
            rows.extend(got)
        elif kind == "json":
            rows.extend(_rows_from_json(data, default_branch, default_source, default_date))
        elif kind == "xlsx":
            rows.extend(_rows_from_table(_xlsx_table(data), default_branch, default_source, default_date) or [])
        else:
            body = data.decode("utf-8-sig", errors="replace")
            delim = "\t" if body.count("\t") > body.count(",") else ","
            got = _rows_from_table(list(csv.reader(io.StringIO(body), delimiter=delim)),
                                   default_branch, default_source, default_date)
            if got is None:   # no recognisable header → every non-empty line is one review
                text = (text or "") + "\n" + body
            else:
                rows.extend(got)
    for line in (text or "").splitlines():
        t = line.strip()
        if t and not t.startswith("#"):
            rows.append({"date": default_date, "branch": default_branch, "source": default_source,
                         "rating": None, "text": t})
    return [r for r in (normalise_row(x) for x in rows) if r]


def normalise_row(r: dict):
    """One raw row (any source) → the stored shape, or None when it carries
    neither text nor (an external id with a rating: a star-only review)."""
    if not isinstance(r, dict):
        return None
    text = str(r.get("text") if r.get("text") is not None else "").strip()
    rating = r.get("rating")
    try:
        rating = float(str(rating).replace("★", "").strip()) if rating not in (None, "") else None
    except ValueError:
        rating = None
    ext = str(r.get("external_id") or "").strip() or None
    if not text and not (ext and rating is not None):
        return None
    d = str(r.get("date") or "").strip()
    hint = r.get("sentiment_hint")
    s = lambda k, n=None: str(r.get(k) or "").strip()[:n] if n else str(r.get(k) or "").strip()  # noqa: E731
    return {"date": d[:10] if len(d) >= 10 else d, "branch": s("branch"), "source": s("source"),
            "rating": rating, "text": text, "external_id": ext, "url": s("url"), "author": s("author"),
            "reply": s("reply"), "replied_at": s("replied_at", 10), "location": s("location"),
            "ref": s("ref"), "sentiment_hint": hint if hint in ("pos", "neg") else None,
            "kind": r.get("kind") if r.get("kind") in KINDS else "review",
            "detect_branch": bool(r.get("detect_branch"))}


def _detect_branch(text: str, table: dict):
    """The branch a mention talks about: the longest branch name the text
    contains ('โอชิเน ขอนแก่น คิวยาว' → ขอนแก่น); none when no name appears."""
    hits = [bn for bn in table if bn and bn in (text or "")]
    if not hits:
        return None, ""
    best = max(hits, key=len)
    return table[best], best


# ── storage ──────────────────────────────────────────────────────────────────
def _branches(con) -> dict:
    return {row["name"]: row["branch_id"] for row in con.execute("SELECT branch_id, name FROM branches")}


def _match_branch(name: str, table: dict):
    n = (name or "").strip()
    if not n:
        return None, ""
    if n in table:
        return table[n], n
    for bn, bid in table.items():
        if bn in n or n in bn:
            return bid, bn
    return None, n


def _menu_names(con) -> list:
    return [r[0] for r in con.execute("SELECT description FROM skus WHERE description IS NOT NULL")]


def _fingerprint(r) -> str:
    """Identity of a review: source + external id when the platform gave one
    (stable when the reviewer edits the text), else branch + date + text."""
    if r.get("external_id"):
        key = f"{r.get('source') or ''}#{r['external_id']}"
    else:
        key = f"{r.get('branch') or ''}|{r.get('date') or ''}|{(r.get('text') or '').strip()}"
    return hashlib.sha1(key.encode("utf-8")).hexdigest()


def _num(v):
    return None if v is None else float(v)


def insert_reviews(con, rows: list, by: str = "", today: str = None) -> dict:
    """rows (normalise_row shape) → customer_feedback. New rows are analysed
    and inserted; a row whose fingerprint exists is updated when its text or
    rating changed (re-analysed) or its owner reply changed (not re-analysed);
    an unchanged row costs nothing (no Claude call). Rows without a date are
    stamped with the Thai calendar day. Returns the counts, the location names
    no branch matched, and the inserted rows (for alerts)."""
    today = today or _thai_today()
    now = _thai_now()
    for r in rows:
        if not (r.get("date") or "").strip():
            r["date"] = today
    table = _branches(con)
    fps = [_fingerprint(r) for r in rows]
    existing = {}
    for i in range(0, len(fps), 400):
        chunk = fps[i:i + 400]
        for e in con.execute(f"SELECT fingerprint, id, text, rating, reply_text FROM customer_feedback "
                             f"WHERE fingerprint IN ({','.join('?' * len(chunk))})", chunk):
            existing[e[0]] = {"id": e[1], "text": e[2], "rating": e[3], "reply": e[4] or ""}
    plan, seen = [], set()
    for i, (r, fp) in enumerate(zip(rows, fps)):
        if fp in seen:
            plan.append(("dup", i)); continue
        seen.add(fp)
        e = existing.get(fp)
        if e is None:
            plan.append(("new", i))
        elif e["text"] != r["text"] or _num(e["rating"]) != _num(r.get("rating")):
            plan.append(("changed", i))
        elif (r.get("reply") or "") != e["reply"]:
            plan.append(("reply", i))
        else:
            plan.append(("same", i))
    to_analyse = [i for k, i in plan if k in ("new", "changed")]
    engine = None
    analysed = {}
    if to_analyse:
        results, engine = analyse([rows[i] for i in to_analyse], _menu_names(con))
        analysed = dict(zip(to_analyse, results))
    inserted = updated = unchanged = dups = 0
    new_items, unmatched = [], set()
    for kind, i in plan:
        r, fp = rows[i], fps[i]
        if kind == "dup":
            dups += 1; continue
        if kind == "same":
            unchanged += 1; continue
        if kind == "reply":
            # a reply that arrived from the platform closes a queued one (ours or a manual answer)
            con.execute("UPDATE customer_feedback SET reply_text=?, replied_at=?, updated_at=?, "
                        "reply_status=CASE WHEN ?!='' AND reply_status='pending' THEN 'sent' ELSE reply_status END WHERE id=?",
                        (r.get("reply") or None, r.get("replied_at") or None, now, r.get("reply") or "", existing[fp]["id"]))
            updated += 1; continue
        bid, bname = _match_branch(r.get("branch") or "", table)
        if bid is None and not r.get("branch") and r.get("detect_branch"):
            bid, bname = _detect_branch(r["text"], table)         # a mention names the branch in its text, if at all
        if r.get("branch") and bid is None:
            unmatched.add(r.get("location") or r["branch"])
        a = analysed[i]
        sent = a["sentiment"]
        if r.get("sentiment_hint") and sent == "neu":
            sent = r["sentiment_hint"]
        themes = list(a["themes"]) + [t for t in (r.get("themes_extra") or []) if t not in a["themes"]]   # a survey's tapped aspects
        sev, sev_note = a.get("severity") or "low", a.get("severity_note") or ""
        vals = (bid, bname, r.get("source") or "", r.get("date") or "", r.get("rating"), r["text"], sent, a["score"],
                json.dumps(themes, ensure_ascii=False), json.dumps(a["menus"], ensure_ascii=False), engine,
                r.get("external_id"), r.get("url") or None, r.get("author") or None, r.get("reply") or None,
                r.get("replied_at") or None, r.get("location") or None, r.get("ref") or None, r.get("kind") or "review",
                r.get("contact") or None, sev, sev_note or None, json.dumps(a.get("staff") or [], ensure_ascii=False))
        if kind == "new":
            cur = con.execute(
                "INSERT OR IGNORE INTO customer_feedback(branch_id, branch_name, source, rated_at, rating, text, "
                "sentiment, score, themes, menus, analysed_by, external_id, url, author_hash, reply_text, replied_at, "
                "location, external_ref, kind, contact, severity, severity_note, staff, uploaded_by, fingerprint) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                vals + (by or "", fp))
            if cur.rowcount and cur.rowcount > 0:
                inserted += 1
                if sev == "high":              # a safety / integrity matter pins itself for management
                    con.execute("UPDATE customer_feedback SET flag=1, flag_note=? WHERE id=?",
                                (f"อัตโนมัติ: ร้ายแรง ({sev_note or 'ตรวจพบคำสำคัญ'})", cur.lastrowid))
                new_items.append({"id": cur.lastrowid, "date": r.get("date"), "branch": bname or r.get("branch") or "",
                                  "source": r.get("source") or "", "rating": r.get("rating"), "text": r["text"][:300],
                                  "sentiment": sent, "url": r.get("url") or "", "severity": sev, "severity_note": sev_note})
            else:
                dups += 1
        else:
            con.execute(
                "UPDATE customer_feedback SET branch_id=?, branch_name=?, source=?, rated_at=?, rating=?, text=?, "
                "sentiment=?, score=?, themes=?, menus=?, analysed_by=?, external_id=?, url=?, author_hash=?, "
                "reply_text=?, replied_at=?, location=?, external_ref=COALESCE(?, external_ref), kind=?, contact=COALESCE(?, contact), severity=?, severity_note=?, staff=?, updated_at=?, "
                "reply_status=CASE WHEN ?!='' AND reply_status='pending' THEN 'sent' ELSE reply_status END WHERE id=?",
                vals + (now, r.get("reply") or "", existing[fp]["id"]))
            updated += 1
    return {"rows": len(rows), "inserted": inserted, "updated": updated, "unchanged": unchanged,
            "duplicates": dups + unchanged, "analysed_by": engine or "none",
            "unmatched": sorted(unmatched), "new_items": new_items[:50]}


# ── endpoints ────────────────────────────────────────────────────────────────
@router.post("/api/feedback/upload", dependencies=[Depends(require_write_token)])
async def feedback_upload(request: Request, file: Optional[UploadFile] = File(None),
                          text: Optional[str] = Form(None), branch: str = Form(""), source: str = Form(""),
                          date: str = Form(""), by: str = Form("")):
    """Accepts a multipart form (file and/or pasted `text`, with defaults) or a
    JSON body {rows:[{date, branch, source, rating, text, external_id, url,
    author, reply, replied_at, location, sentiment_hint}], by}."""
    rows = []
    ctype = (request.headers.get("content-type") or "").lower()
    if "application/json" in ctype:
        body = await request.json()
        by = body.get("by") or by
        rows = [r for r in (normalise_row(x) for x in (body.get("rows") or [])) if r]
    else:
        data = await file.read() if file is not None else b""
        try:
            rows = parse_rows(text or "", data, file.filename if file is not None else "", branch, source, date)
        except (ValueError, zipfile.BadZipFile) as e:
            raise HTTPException(400, f"อ่านไฟล์ไม่ได้: {e}")
    if not rows:
        raise HTTPException(400, "ไม่พบข้อความรีวิว — วางข้อความบรรทัดละ 1 รีวิว หรือแนบไฟล์ CSV/XLSX/JSON/Takeout zip ที่มีรีวิว")
    if len(rows) > MAX_ROWS:
        raise HTTPException(400, f"ครั้งละไม่เกิน {MAX_ROWS:,} รีวิว")
    scope = _scope(request)
    if scope:                                  # a branch login can only feed its own branch
        for r in rows:
            r["branch"] = scope
            r["detect_branch"] = False
    with get_db() as con:
        _ensure(con)
        out = insert_reviews(con, rows, by)
        con.commit()
    return out


def _where(branch, month_from, month_to, sentiment=None, theme=None, menu=None, source=None, kind=None):
    w, p = ["1=1"], []
    if kind:
        w.append("COALESCE(kind,'review') = ?"); p.append(kind)
    if branch:
        w.append("branch_name = ?"); p.append(branch)
    if month_from:
        w.append("substr(rated_at,1,7) >= ?"); p.append(month_from[:7])
    if month_to:
        w.append("substr(rated_at,1,7) <= ?"); p.append(month_to[:7])
    if sentiment:
        w.append("sentiment = ?"); p.append(sentiment)
    if theme:
        w.append("themes LIKE ?"); p.append(f'%"{theme}"%')
    if menu:
        w.append("menus LIKE ?"); p.append(f'%"{menu}"%')
    if source:
        w.append("source = ?"); p.append(source)
    return " AND ".join(w), p


def _needs_reply(r) -> bool:
    """A fetched review or comment (has an external id) that is negative or
    ≤3★, the branch has not answered and no reply is queued — the list to work
    through. Mentions elsewhere on the web are read, not answered."""
    kind = r.get("kind") or "review"
    answerable = bool(r.get("contact")) if kind == "survey" else kind in ("review", "comment", "message")
    return bool(r.get("external_id")) and answerable and \
        not (r.get("reply_text") or "") and r.get("reply_status") not in ("pending", "sent") and \
        (r.get("sentiment") == "neg" or (r.get("rating") is not None and float(r["rating"]) <= 3))


def _reply_channel(row) -> str:
    """Where a queued reply can go: google (review resource name), facebook (a
    comment id), tiktok (video/comment ids), manual (a Facebook recommendation,
    a mention, or a row without a ref — the draft is kept for pasting by hand),
    none (no external id at all)."""
    if not row.get("external_id"):
        return "none"
    src = (row.get("source") or "").strip().lower()
    ext, ref = str(row.get("external_id")), row.get("external_ref") or ""
    if src == "google" and ref.count("/reviews/"):
        return "google"
    if src == "facebook" and ext.startswith(("comment:", "post:")) and ref:     # a comment, or a visitor's post on the page
        return "facebook"
    if src == "tiktok" and ext.startswith("comment:") and "/" in ref:
        return "tiktok"
    if src == "instagram" and ext.startswith("comment:") and "/" in ref:
        return "instagram"
    return "manual"


_ROW_COLS = ("id, branch_name, source, rated_at, rating, text, sentiment, score, themes, menus, analysed_by, "
             "external_id, url, reply_text, replied_at, reply_status, reply_error, kind, flag, flag_note, location, contact, "
             "severity, severity_note, staff")


def _age_days(rated_at: str, today: str):
    try:
        return (datetime.date.fromisoformat(today) - datetime.date.fromisoformat((rated_at or "")[:10])).days
    except ValueError:
        return None


def _sla_breached(r, today: str) -> bool:
    age = _age_days(r.get("rated_at"), today)
    return _needs_reply(r) and age is not None and age > SLA_DAYS


def _prev_range(month_from, month_to):
    """The window of the same length right before [month_from, month_to]; (None, None) when unbounded."""
    if not month_from or not month_to:
        return None, None
    try:
        ya, ma, yb, mb = int(month_from[:4]), int(month_from[5:7]), int(month_to[:4]), int(month_to[5:7])
    except ValueError:
        return None, None
    span = (yb * 12 + mb) - (ya * 12 + ma) + 1
    if span <= 0:
        return None, None
    end = ya * 12 + (ma - 1) - 1              # 0-based month index of the month before month_from
    start = end - span + 1
    if start < 0:
        return None, None
    return f"{start // 12:04d}-{start % 12 + 1:02d}", f"{end // 12:04d}-{end % 12 + 1:02d}"


def _brief(r, today: str) -> dict:
    """One voice as the tab lists it (queue, negatives, positives)."""
    return {"id": r["id"], "date": r["rated_at"], "branch": r["branch_name"], "source": r["source"],
            "rating": r["rating"], "text": (r["text"] or "")[:300], "themes": json.loads(r["themes"] or "[]"),
            "menus": json.loads(r["menus"] or "[]"), "url": r["url"] or "", "sentiment": r["sentiment"],
            "replied": bool(r["reply_text"]), "needs_reply": _needs_reply(r), "kind": r["kind"] or "review",
            "external_id": r["external_id"], "reply_status": r["reply_status"], "reply_error": r["reply_error"],
            "flag": bool(r.get("flag")), "flag_note": r.get("flag_note") or "", "age_days": _age_days(r["rated_at"], today),
            "sla_breached": _sla_breached(r, today), "location": r.get("location") or "", "contact": r.get("contact") or "",
            "severity": r.get("severity") or "low", "severity_note": r.get("severity_note") or "",
            "staff": json.loads(r.get("staff") or "[]") if isinstance(r.get("staff"), str) else (r.get("staff") or [])}


def _aggregate(rows: list, today: str) -> dict:
    """Everything the tab shows for one window of rows."""
    total = len(rows)
    sent = {"pos": 0, "neg": 0, "neu": 0}
    themes, menus, branches, by_month, by_day, sources, engines, staff = {}, {}, {}, {}, {}, {}, {}, {}
    reply_queue = {s: 0 for s in REPLY_STATUSES}
    ratings, unreplied, sla, flagged, severe = [], 0, 0, 0, 0
    answerable = answered = 0
    reply_days = []
    for r in rows:
        s = r["sentiment"] or "neu"
        sent[s] = sent.get(s, 0) + 1
        if r["rating"] is not None:
            ratings.append(float(r["rating"]))
        nr = _needs_reply(r)
        unreplied += nr
        sla += _sla_breached(r, today)
        flagged += 1 if r.get("flag") else 0
        severe += 1 if r.get("severity") == "high" else 0
        for nm in (json.loads(r.get("staff") or "[]") if isinstance(r.get("staff"), str) else (r.get("staff") or [])):
            st = staff.setdefault(nm, {"name": nm, "n": 0, "pos": 0, "neg": 0, "branches": {}, "examples": []})
            st["n"] += 1; st["pos"] += (s == "pos"); st["neg"] += (s == "neg")
            st["branches"][r["branch_name"] or "-"] = st["branches"].get(r["branch_name"] or "-", 0) + 1
            if len(st["examples"]) < 2:
                st["examples"].append({"id": r["id"], "text": (r["text"] or "")[:120], "sentiment": s})
        if r["reply_status"] in reply_queue:
            reply_queue[r["reply_status"]] += 1
        # the answer rate: fetched negative / ≤3★ voices the branch could answer, and did
        if r.get("external_id") and (r["kind"] or "review") in ("review", "comment", "message") and \
                (s == "neg" or (r["rating"] is not None and float(r["rating"]) <= 3)):
            answerable += 1
            if r["reply_text"] or r["reply_status"] == "sent":
                answered += 1
                try:
                    reply_days.append((datetime.date.fromisoformat((r["replied_at"] or "")[:10])
                                       - datetime.date.fromisoformat((r["rated_at"] or "")[:10])).days)
                except ValueError:
                    pass
        bn = r["branch_name"] or "(ไม่ระบุสาขา)"
        b = branches.setdefault(bn, {"branch": bn, "n": 0, "neg": 0, "pos": 0, "ratings": [], "unreplied": 0, "flagged": 0, "themes": {}})
        b["n"] += 1; b["neg"] += (s == "neg"); b["pos"] += (s == "pos"); b["unreplied"] += nr; b["flagged"] += 1 if r.get("flag") else 0
        if r["rating"] is not None:
            b["ratings"].append(float(r["rating"]))
        for t in json.loads(r["themes"] or "[]"):
            d = themes.setdefault(t, {"theme": t, "n": 0, "neg": 0, "pos": 0, "branches": {}, "examples": []})
            d["n"] += 1; d["neg"] += (s == "neg"); d["pos"] += (s == "pos")
            d["branches"][bn] = d["branches"].get(bn, 0) + 1
            if s == "neg" and len(d["examples"]) < 3:
                d["examples"].append({"id": r["id"], "date": r["rated_at"], "branch": r["branch_name"], "text": (r["text"] or "")[:160]})
            if s == "neg":
                b["themes"][t] = b["themes"].get(t, 0) + 1
        for m in json.loads(r["menus"] or "[]"):
            d = menus.setdefault(m, {"menu": m, "n": 0, "pos": 0, "neg": 0}); d["n"] += 1; d["pos"] += (s == "pos"); d["neg"] += (s == "neg")
        ym = (r["rated_at"] or "")[:7] or "ไม่ระบุ"
        d = by_month.setdefault(ym, {"month": ym, "n": 0, "pos": 0, "neg": 0}); d["n"] += 1; d["pos"] += (s == "pos"); d["neg"] += (s == "neg")
        day = (r["rated_at"] or "")[:10]
        if len(day) == 10:
            d = by_day.setdefault(day, {"day": day, "n": 0, "neg": 0}); d["n"] += 1; d["neg"] += (s == "neg")
        sources[r["source"] or "(ไม่ระบุ)"] = sources.get(r["source"] or "(ไม่ระบุ)", 0) + 1
        engines[r["analysed_by"] or "?"] = engines.get(r["analysed_by"] or "?", 0) + 1
    for b in branches.values():
        b["avg_rating"] = round(sum(b["ratings"]) / len(b["ratings"]), 2) if b["ratings"] else None
        b["neg_pct"] = round(b["neg"] / b["n"] * 100, 1) if b["n"] else 0.0
        b["top_theme"] = max(b["themes"].items(), key=lambda kv: kv[1])[0] if b["themes"] else None
        del b["ratings"], b["themes"]
    for d in themes.values():
        d["neg_pct"] = round(d["neg"] / d["n"] * 100, 1) if d["n"] else 0.0
        d["branches"] = [k for k, _ in sorted(d["branches"].items(), key=lambda kv: -kv[1])[:3]]
    reply_days.sort()
    for st in staff.values():
        st["branch"] = max(st["branches"].items(), key=lambda kv: kv[1])[0] if st["branches"] else ""
        del st["branches"]
    return {
        "total": total, "avg_rating": round(sum(ratings) / len(ratings), 2) if ratings else None,
        "sentiment": sent, "neg_pct": round(sent["neg"] / total * 100, 1) if total else 0.0,
        "unreplied": unreplied, "sla_breached": sla, "flagged": flagged, "severe": severe, "reply_queue": reply_queue,
        "staff": sorted(staff.values(), key=lambda d: (-d["n"], d["name"]))[:12],
        "response": {"answerable": answerable, "answered": answered,
                     "rate": round(answered / answerable * 100) if answerable else None,
                     "median_days": reply_days[len(reply_days) // 2] if reply_days else None},
        "themes": sorted(themes.values(), key=lambda d: -d["n"]),
        "menus": sorted(menus.values(), key=lambda d: -d["n"])[:30],
        "branches": sorted(branches.values(), key=lambda d: (-d["neg_pct"], -d["n"])),
        "by_month": sorted(by_month.values(), key=lambda d: d["month"]),
        "by_day": sorted(by_day.values(), key=lambda d: d["day"]),
        "sources": sources, "analysed_by": engines,
        "recent_negative": [_brief(r, today) for r in rows if r["sentiment"] == "neg"][:12],
        "recent_positive": [_brief(r, today) for r in rows if r["sentiment"] == "pos" and (r["kind"] or "review") != "mention" and (r["text"] or "").strip()][:6],
    }


@router.get("/api/feedback/summary")
def feedback_summary(request: Request, branch: Optional[str] = Query(None), month_from: Optional[str] = Query(None),
                     month_to: Optional[str] = Query(None), source: Optional[str] = Query(None),
                     kind: Optional[str] = Query(None)):
    """The window's numbers plus the same window right before it (`prev`), so
    the tab can say which way every figure moved."""
    if kind and kind not in KINDS:
        raise HTTPException(400, f"kind must be one of {', '.join(KINDS)}")
    today = _thai_today()
    branch = _scope(request) or branch
    pf, pt = _prev_range(month_from, month_to)
    with get_db() as con:
        _ensure(con)
        where, params = _where(branch, month_from, month_to, source=source, kind=kind)
        rows = [dict(r) for r in con.execute(f"SELECT {_ROW_COLS} FROM customer_feedback WHERE {where} "
                                             f"ORDER BY rated_at DESC, id DESC", params)]
        prev_rows = []
        if pf:
            wp, pp = _where(branch, pf, pt, source=source, kind=kind)
            prev_rows = [dict(r) for r in con.execute(f"SELECT {_ROW_COLS} FROM customer_feedback WHERE {wp}", pp)]
        # the kind chips count every kind under the OTHER filters, so a selected kind never hides the rest
        where_k, params_k = _where(branch, month_from, month_to, source=source)
        kinds = {k: n for k, n in con.execute(
            f"SELECT COALESCE(kind,'review'), COUNT(*) FROM customer_feedback WHERE {where_k} GROUP BY 1", params_k)}
    out = _aggregate(rows, today)
    prev = None
    if pf:
        p = _aggregate(prev_rows, today)
        prev = {"range": [pf, pt], "total": p["total"], "avg_rating": p["avg_rating"], "neg_pct": p["neg_pct"],
                "sentiment": p["sentiment"], "response_rate": p["response"]["rate"], "sla_breached": p["sla_breached"],
                "themes": {t["theme"]: t["n"] for t in p["themes"]},
                "branches": {b["branch"]: {"avg_rating": b["avg_rating"], "neg_pct": b["neg_pct"], "n": b["n"]} for b in p["branches"]}}
    out.update({"kinds": kinds, "prev": prev, "today": today, "sla_days": SLA_DAYS,
                "filters": {"branch": branch, "month_from": month_from, "month_to": month_to, "source": source, "kind": kind}})
    return out


@router.get("/api/feedback/items")
def feedback_items(request: Request, branch: Optional[str] = Query(None), month_from: Optional[str] = Query(None),
                   month_to: Optional[str] = Query(None), sentiment: Optional[str] = Query(None),
                   theme: Optional[str] = Query(None), menu: Optional[str] = Query(None),
                   source: Optional[str] = Query(None), kind: Optional[str] = Query(None),
                   needs_reply: int = Query(0), flagged: int = Query(0), sla: int = Query(0),
                   sort: str = Query("newest"), limit: int = Query(200, ge=1, le=2000)):
    """sort: newest (default) · oldest · worst (lowest stars first) · priority
    (unanswered first, then pinned / severe, then SLA breaches, then the worst, then the oldest)."""
    if kind and kind not in KINDS:
        raise HTTPException(400, f"kind must be one of {', '.join(KINDS)}")
    if sort not in ("newest", "oldest", "worst", "priority"):
        raise HTTPException(400, "sort must be newest, oldest, worst or priority")
    today = _thai_today()
    branch = _scope(request) or branch
    with get_db() as con:
        _ensure(con)
        where, params = _where(branch, month_from, month_to, sentiment, theme, menu, source, kind)
        if flagged:
            where += " AND flag=1"
        rows = con.execute(
            f"SELECT {_ROW_COLS}, uploaded_at, reply_draft "
            f"FROM customer_feedback WHERE {where} ORDER BY rated_at DESC, id DESC LIMIT ?", params + [limit]).fetchall()
    items = [{**dict(r), "themes": json.loads(r["themes"] or "[]"), "menus": json.loads(r["menus"] or "[]"),
              "staff": json.loads(r["staff"] or "[]") if "staff" in r.keys() else [], "severity": (r["severity"] if "severity" in r.keys() else None) or "low"}
             for r in rows]
    for it in items:
        it["kind"] = it.get("kind") or "review"
        it["flag"] = bool(it.get("flag"))
        it["needs_reply"] = _needs_reply(it)
        it["age_days"] = _age_days(it["rated_at"], today)
        it["sla_breached"] = _sla_breached(it, today)
    if needs_reply:
        items = [it for it in items if it["needs_reply"]]
    if sla:
        items = [it for it in items if it["sla_breached"]]
    if sort == "oldest":
        items.sort(key=lambda it: (it["rated_at"] or "", it["id"]))
    elif sort == "worst":
        items.sort(key=lambda it: (it["rating"] if it["rating"] is not None else 6, -(it["age_days"] or 0)))
    elif sort == "priority":
        items.sort(key=lambda it: (not it["needs_reply"], not (it["flag"] or it.get("severity") == "high"), not it["sla_breached"],
                                   it["rating"] if it["rating"] is not None else 6, -(it["age_days"] or 0)))
    return {"items": items, "today": today}


@router.post("/api/feedback/{fid}/handled", dependencies=[Depends(require_write_token)])
async def feedback_handled(fid: int, request: Request):
    """The branch answered outside the system (in the app, by phone): record it
    so the voice leaves the queue and counts as answered."""
    b = await request.json()
    with get_db() as con:
        _ensure(con)
        _get_row(con, fid, request)
        note = str(b.get("note") or "").strip()[:2000] or "(ตอบแล้วนอกระบบ)"
        con.execute("UPDATE customer_feedback SET reply_status='sent', reply_text=COALESCE(NULLIF(reply_text,''), ?), "
                    "replied_at=COALESCE(replied_at, ?), reply_requested_by=?, reply_error=NULL, updated_at=? WHERE id=?",
                    (note, _thai_today(), str(b.get("by") or "")[:60], _thai_now(), fid))
        con.commit()
    return {"id": fid, "reply_status": "sent"}


@router.post("/api/feedback/{fid}/flag", dependencies=[Depends(require_write_token)])
async def feedback_flag(fid: int, request: Request):
    """Pin a serious case (food safety, fraud, harassment) with the reason — it
    stays on top of the queue and in the branch league until unpinned."""
    b = await request.json()
    on = bool(b.get("flag", True))
    with get_db() as con:
        _ensure(con)
        _get_row(con, fid, request)
        con.execute("UPDATE customer_feedback SET flag=?, flag_note=?, updated_at=? WHERE id=?",
                    (1 if on else 0, (str(b.get("note") or "").strip()[:500] or None) if on else None, _thai_now(), fid))
        con.commit()
    return {"id": fid, "flag": on}


@router.get("/api/feedback/sources")
def feedback_sources():
    """Per source: how many reviews, the latest review date, how many were
    fetched by a connector, the reply queue, and the fetcher's last run."""
    with get_db() as con:
        _ensure(con)
        per = [dict(r) for r in con.execute(
            "SELECT source, COUNT(*) AS n, MAX(rated_at) AS latest, "
            "SUM(CASE WHEN external_id IS NOT NULL THEN 1 ELSE 0 END) AS fetched, "
            "SUM(CASE WHEN external_id IS NOT NULL AND (reply_text IS NULL OR reply_text='') "
            "         AND COALESCE(kind,'review') IN ('review','comment','message') "
            "         AND COALESCE(reply_status,'') NOT IN ('pending','sent') "
            "         AND (sentiment='neg' OR rating<=3) THEN 1 ELSE 0 END) AS unreplied "
            "FROM customer_feedback GROUP BY source ORDER BY n DESC")]
        logs = [dict(r) for r in con.execute(
            "SELECT source, ran_at, since, fetched, inserted, updated, unchanged, status, message "
            "FROM feedback_sync_log ORDER BY id DESC LIMIT 300")]
        reply_queue = {s: 0 for s in REPLY_STATUSES}
        for st, n in con.execute("SELECT reply_status, COUNT(*) FROM customer_feedback "
                                 "WHERE reply_status IS NOT NULL GROUP BY reply_status"):
            if st in reply_queue:
                reply_queue[st] = n
    last, last_ok = {}, {}
    for entry in logs:
        last.setdefault(entry["source"], entry)
        if entry["status"] == "ok" and entry["source"] not in last_ok:
            last_ok[entry["source"]] = entry["ran_at"]
    seen = set()
    out = []
    for p in per:
        name = p["source"] or "(ไม่ระบุ)"
        seen.add(p["source"])
        out.append({**p, "source": name, "last_run": last.get(p["source"])})
    for s, entry in last.items():
        if s not in seen:
            out.append({"source": s, "n": 0, "latest": None, "fetched": 0, "unreplied": 0, "last_run": entry})
    return {"sources": out, "last_ok": last_ok, "runs": logs[:20], "reply_queue": reply_queue}


# ── the reply queue ──────────────────────────────────────────────────────────
def _get_row(con, fid: int, request: Request = None) -> dict:
    r = con.execute("SELECT * FROM customer_feedback WHERE id=?", (fid,)).fetchone()
    if not r:
        raise HTTPException(404, "not found")
    row = dict(r)
    if request is not None:
        _guard(request, row)
    return row


def _scope(request: Request):
    """The branch this login is confined to (E4), or None."""
    return _auth.scope_branch(request) if request is not None else None


def _guard(request: Request, row: dict):
    b = _scope(request)
    if b and (row.get("branch_name") or "") != b:
        raise HTTPException(403, "รายการนี้อยู่นอกสาขาของคุณ")


@router.get("/api/feedback/reply-queue")
def feedback_reply_queue(status: str = Query("pending")):
    """What the fetcher (backend/fetch_feedback.py) must post, plus the counts
    per status for the tab."""
    if status not in REPLY_STATUSES:
        raise HTTPException(400, f"status must be one of {', '.join(REPLY_STATUSES)}")
    with get_db() as con:
        _ensure(con)
        items = [dict(r) for r in con.execute(
            "SELECT id, source, external_id, external_ref, branch_name, location, reply_draft, reply_status, "
            "reply_requested_at, reply_requested_by, reply_error FROM customer_feedback WHERE reply_status=? "
            "ORDER BY reply_requested_at, id", (status,))]
        counts = {s: 0 for s in REPLY_STATUSES}
        for st, n in con.execute("SELECT reply_status, COUNT(*) FROM customer_feedback "
                                 "WHERE reply_status IS NOT NULL GROUP BY reply_status"):
            if st in counts:
                counts[st] = n
    return {"items": items, "counts": counts}


@router.post("/api/feedback/reply-result", dependencies=[Depends(require_write_token)])
async def feedback_reply_result(request: Request):
    """The fetcher reports one queued reply: {id, status: sent|failed, error, reply_text}."""
    b = await request.json()
    try:
        fid = int(b.get("id"))
    except (TypeError, ValueError):
        raise HTTPException(400, "id required")
    status = b.get("status")
    if status not in ("sent", "failed"):
        raise HTTPException(400, "status must be sent or failed")
    with get_db() as con:
        _ensure(con)
        row = _get_row(con, fid)
        if status == "sent":
            con.execute("UPDATE customer_feedback SET reply_status='sent', reply_error=NULL, "
                        "reply_text=COALESCE(?, reply_draft), replied_at=?, updated_at=? WHERE id=?",
                        ((b.get("reply_text") or "").strip() or None, (b.get("replied_at") or _thai_today())[:10], _thai_now(), fid))
        else:
            con.execute("UPDATE customer_feedback SET reply_status='failed', reply_error=?, updated_at=? WHERE id=?",
                        (str(b.get("error") or "unknown error")[:300], _thai_now(), fid))
        con.commit()
    return {"id": fid, "reply_status": status, "previous": row.get("reply_status")}


@router.get("/api/feedback/{fid}/reply-draft")
def feedback_reply_draft(fid: int, request: Request):
    """A suggested reply for the branch to edit: Claude when ANTHROPIC_API_KEY is
    set, the Thai template otherwise. Also says where the reply can be sent."""
    with get_db() as con:
        _ensure(con)
        row = _get_row(con, fid, request)
    if row.get("reply_status") in ("pending", "manual", "failed") and row.get("reply_draft"):
        draft, engine = row["reply_draft"], "saved"
    else:
        draft, engine = draft_reply(row)
    ch = _reply_channel(row)
    return {"id": fid, "draft": draft, "engine": engine, "channel": ch, "note": _reply_note(ch, row),
            "max_len": REPLY_MAX.get(ch, 4000), "kind": row.get("kind") or "review",
            "reply_status": row.get("reply_status"), "reply_error": row.get("reply_error"),
            "text": row.get("text") or "", "branch": row.get("branch_name") or "", "source": row.get("source") or "",
            "rating": row.get("rating"), "date": row.get("rated_at"), "url": row.get("url") or ""}


@router.post("/api/feedback/{fid}/reply", dependencies=[Depends(require_write_token)])
async def feedback_reply(fid: int, request: Request):
    """Queue a reply. Google reviews and Facebook comments are posted by the
    next fetcher run; a Facebook recommendation (or a row without a resource
    name) is kept as `manual` — the text to paste by hand."""
    b = await request.json()
    text = str(b.get("text") or "").strip()
    if not text:
        raise HTTPException(400, "ใส่ข้อความตอบกลับก่อน")
    with get_db() as con:
        _ensure(con)
        row = _get_row(con, fid, request)
        ch = _reply_channel(row)
        if ch == "none":
            raise HTTPException(400, "รีวิวนี้ไม่ได้มาจากช่องทางที่ระบบตอบกลับได้ (ไม่มี review id)")
        limit = REPLY_MAX.get(ch, 4000)
        if len(text) > limit:
            raise HTTPException(400, f"ข้อความตอบกลับยาวเกิน {limit:,} ตัวอักษร ({ch} รับได้เท่านั้น) — ตอนนี้ {len(text):,}")
        status = "pending" if ch in ("google", "facebook", "tiktok", "instagram") else "manual"
        con.execute("UPDATE customer_feedback SET reply_draft=?, reply_status=?, reply_requested_at=?, "
                    "reply_requested_by=?, reply_error=NULL, updated_at=? WHERE id=?",
                    (text, status, _thai_now(), str(b.get("by") or "")[:60], _thai_now(), fid))
        con.commit()
    return {"id": fid, "reply_status": status, "channel": ch, "note": _reply_note(ch, row)}


def _reply_note(ch: str, row: dict) -> str:
    if ch == "google":
        return "ส่งถึง Google ในรอบดึงถัดไป (09:00)"
    if ch == "facebook":
        return "โพสต์ใต้คอมเมนต์ในรอบดึงถัดไป (09:00)"
    if ch == "tiktok":
        return "โพสต์ตอบใต้คอมเมนต์ TikTok ในรอบดึงถัดไป (09:00) — TikTok รับไม่เกิน 150 ตัวอักษร"
    if ch == "instagram":
        return "โพสต์ตอบใต้คอมเมนต์ Instagram ในรอบดึงถัดไป (09:00)"
    if ch == "manual":
        if (row.get("kind") or "review") == "survey":
            return (f"แบบสอบถามหน้าร้าน — ลูกค้าฝากช่องทางติดต่อไว้: {row.get('contact') or '-'} "
                    "โทรหรือทัก LINE กลับด้วยข้อความนี้ แล้วกด ✓ ตอบแล้ว ในกล่องงาน (ระบบไม่ส่งข้อความแทน)")
        if (row.get("kind") or "review") == "message":
            return "ข้อความส่วนตัวจาก Inbox — ระบบไม่ส่งแทน บันทึกร่างไว้แล้วไปตอบในกล่องข้อความของ Meta Business Suite (เปิดลิงก์ ↗)"
        if (row.get("kind") or "review") == "mention":
            return "เป็นการพูดถึงบนเว็บ/โซเชียลของคนอื่น — ระบบไม่โพสต์แทน บันทึกข้อความไว้แล้วไปตอบเองในแพลตฟอร์มนั้น"
        if (row.get("source") or "").strip().lower() == "facebook":
            return "Facebook ไม่ให้ตอบ Recommendation ผ่าน API — คัดลอกข้อความไปตอบในแอป Facebook"
        return "รีวิวนี้ไม่มี resource name (นำเข้าจากไฟล์ที่ไม่ใช่ Takeout/API?) — คัดลอกข้อความไปตอบในแอป"
    return "รีวิวนี้ไม่มี review id จึงตอบผ่านระบบไม่ได้"


# ── the Meta webhook (Instagram @mentions) ───────────────────────────────────
# Meta pushes events here; there is no session (main.py exempts the path), so the
# endpoint carries its own two checks: the verify token on the subscription
# handshake and the HMAC-SHA256 of every body with the app secret. Both come
# from env (Fly secrets META_APP_SECRET / META_WEBHOOK_VERIFY_TOKEN); unset =
# fail closed. Events are only STORED — the nightly fetcher resolves them with
# the Page token it holds and marks them processed.
@router.get("/api/feedback/webhook/meta", include_in_schema=False)
def meta_webhook_verify(hub_mode: str = Query("", alias="hub.mode"), hub_challenge: str = Query("", alias="hub.challenge"),
                        hub_verify_token: str = Query("", alias="hub.verify_token")):
    expected = os.environ.get("META_WEBHOOK_VERIFY_TOKEN", "")
    if hub_mode == "subscribe" and expected and hmac.compare_digest(hub_verify_token, expected):
        return PlainTextResponse(hub_challenge)
    raise HTTPException(403, "verify token mismatch")


@router.post("/api/feedback/webhook/meta", include_in_schema=False)
async def meta_webhook(request: Request):
    secret = os.environ.get("META_APP_SECRET", "")
    if not secret:
        raise HTTPException(503, "webhook not configured (META_APP_SECRET)")
    raw = await request.body()
    if len(raw) > 262144:
        raise HTTPException(413, "payload too large")
    expected = "sha256=" + hmac.new(secret.encode("utf-8"), raw, hashlib.sha256).hexdigest()
    if not hmac.compare_digest(request.headers.get("x-hub-signature-256", ""), expected):
        raise HTTPException(403, "bad signature")
    try:
        payload = json.loads(raw.decode("utf-8") or "{}")
    except ValueError:
        raise HTTPException(400, "not JSON")
    platform = str(payload.get("object") or "")
    n = 0
    with get_db() as con:
        _ensure(con)
        for entry in payload.get("entry") or []:
            for ch in (entry.get("changes") or []) if isinstance(entry, dict) else []:
                con.execute("INSERT INTO feedback_webhook_events(platform, object_id, field, payload) VALUES (?,?,?,?)",
                            (platform, str(entry.get("id") or ""), str(ch.get("field") or ""),
                             json.dumps(ch.get("value") or {}, ensure_ascii=False)[:4000]))
                n += 1
        con.execute("DELETE FROM feedback_webhook_events WHERE id NOT IN "
                    "(SELECT id FROM feedback_webhook_events ORDER BY id DESC LIMIT 5000)")
        con.commit()
    return {"stored": n}


@router.get("/api/feedback/webhook-events")
def feedback_webhook_events(platform: str = Query("instagram"), unprocessed: int = Query(1),
                            limit: int = Query(500, ge=1, le=2000)):
    """What the webhook kept — the fetcher reads the unprocessed ones."""
    with get_db() as con:
        _ensure(con)
        where = "platform=?" + (" AND processed_at IS NULL" if unprocessed else "")
        rows = con.execute(f"SELECT id, received_at, platform, object_id, field, payload, processed_at, result "
                           f"FROM feedback_webhook_events WHERE {where} ORDER BY id LIMIT ?", (platform, limit)).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        try:
            d["payload"] = json.loads(d.get("payload") or "{}")
        except ValueError:
            d["payload"] = {}
        out.append(d)
    return {"events": out}


@router.post("/api/feedback/webhook-events/ack", dependencies=[Depends(require_write_token)])
async def feedback_webhook_ack(request: Request):
    b = await request.json()
    ids = [int(i) for i in (b.get("ids") or []) if str(i).isdigit()]
    if not ids:
        raise HTTPException(400, "ids required")
    with get_db() as con:
        _ensure(con)
        con.executemany("UPDATE feedback_webhook_events SET processed_at=?, result=? WHERE id=?",
                        [(_thai_now(), str(b.get("result") or "ok")[:200], i) for i in ids])
        con.commit()
    return {"acked": len(ids)}


@router.get("/api/feedback/oauth-callback", include_in_schema=False)
def feedback_oauth_callback(code: str = Query(""), state: str = Query(""), error: str = Query(""),
                            error_description: str = Query("")):
    """Where a platform's OAuth consent lands (TikTok insists on a registered
    https redirect, so it cannot be the fetcher's localhost): shows the code
    for pasting into `fetch_feedback.py auth …`."""
    from html import escape
    body = (f"<p>คัดลอก code นี้ไปวางในหน้าต่างคำสั่ง (<code>fetch_feedback.py auth …</code>):</p>"
            f"<pre style='font-size:1.1em;background:#f3f4f6;padding:.6rem;word-break:break-all'>{escape(code)}</pre>"
            if code else f"<p style='color:#b91c1c'>ไม่ได้รับ code: {escape(error)} {escape(error_description)}</p>")
    return HTMLResponse(f"<!doctype html><meta charset='utf-8'><title>Oshinei OAuth callback</title>"
                        f"<body style='font-family:sans-serif;max-width:640px;margin:3rem auto'>"
                        f"<h3>Oshinei — OAuth callback</h3>{body}"
                        f"<p style='color:#6b7280;font-size:.9em'>state: {escape(state)}</p></body>")


@router.delete("/api/feedback/{fid}/reply", dependencies=[Depends(require_write_token)])
def feedback_reply_cancel(fid: int, request: Request):
    with get_db() as con:
        _ensure(con)
        row = _get_row(con, fid, request)
        if row.get("reply_status") == "sent":
            raise HTTPException(400, "ส่งไปแล้ว — ยกเลิกไม่ได้")
        con.execute("UPDATE customer_feedback SET reply_status=NULL, reply_draft=NULL, reply_error=NULL, "
                    "reply_requested_at=NULL, reply_requested_by=NULL, updated_at=? WHERE id=?", (_thai_now(), fid))
        con.commit()
    return {"id": fid, "reply_status": None}


@router.post("/api/feedback/sync-log", dependencies=[Depends(require_write_token)])
async def feedback_sync_log(request: Request):
    """The fetcher records each run here so the tab can show when a source was
    last collected and whether it worked."""
    b = await request.json()
    source = str(b.get("source") or "").strip()
    if not source:
        raise HTTPException(400, "source required")
    n = lambda k: int(b.get(k) or 0)  # noqa: E731
    with get_db() as con:
        _ensure(con)
        cur = con.execute(
            "INSERT INTO feedback_sync_log(source, since, fetched, inserted, updated, unchanged, status, message, by) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (source, str(b.get("since") or "")[:10], n("fetched"), n("inserted"), n("updated"), n("unchanged"),
             "ok" if (b.get("status") or "ok") == "ok" else "error", str(b.get("message") or "")[:500],
             str(b.get("by") or "")[:60]))
        con.execute("DELETE FROM feedback_sync_log WHERE id NOT IN (SELECT id FROM feedback_sync_log ORDER BY id DESC LIMIT 1000)")
        con.commit()
    return {"logged": cur.lastrowid}


@router.delete("/api/feedback/{fid}", dependencies=[Depends(require_write_token)])
def feedback_delete(fid: int, request: Request):
    with get_db() as con:
        _ensure(con)
        _get_row(con, fid, request)
        n = con.execute("DELETE FROM customer_feedback WHERE id=?", (fid,)).rowcount
        con.commit()
    if not n:
        raise HTTPException(404, "not found")
    return {"deleted": fid}


__all__ = ["router", "parse_rows", "normalise_row", "insert_reviews", "clean_menu_name", "REPLY_STATUSES", "KINDS"]
