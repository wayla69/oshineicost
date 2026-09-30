# -*- coding: utf-8 -*-
"""รายงานประจำวัน/สัปดาห์/เดือน — เสียงลูกค้า, ออกแบบให้ทีมการตลาดใช้ได้ทันที:
เลือกช่วง (วันเดียว, สัปดาห์, เดือน หรือกำหนดเอง — และสาขาถ้าต้องการ) แล้วได้
สรุปพร้อมใช้ — คะแนน/สัดส่วนความรู้สึกของช่วงนั้น กี่รายการตอบแล้ว/รอตอบ/เกิน
SLA เคสร้ายแรงที่ต้องรู้ หัวข้อ/เมนูที่ถูกพูดถึง เสียงชมที่หยิบไปใช้ต่อได้ และ
รายชื่อที่ต้องติดตาม (รวมเบอร์โทรกลับจาก QR — ดู js/11-feedback.js:_fbContactHtml
สำหรับ tel: link ฝั่งแอป).

_aggregate() (routers/feedback.py) คำนวณให้ทุกอย่างอยู่แล้วสำหรับ /api/feedback/
summary — ที่นี่แค่เป็นหน้าต่าง [date_from, date_to] ของฟังก์ชันเดียวกัน ไม่ใช่
โมเดลใหม่ จึงรับประกันว่าตัวเลขตรงกับที่แท็บโชว์เป๊ะ. วันเดียว (date_to ไม่ใส่
หรือเท่ากับ date) คือ "รายงานประจำวัน" เดิม — ใส่ date_to ให้ต่างจาก date จะได้
รายงานตามช่วงที่เลือก (สัปดาห์/เดือน/กำหนดเอง แล้วแต่หน้าเว็บส่งมา).

ต่อกับ routers/feedback_actions.py (โมดูล ปิดวงจร) เพื่อสี่ส่วนนี้:
  - new_issues:      (สาขา,หัวข้อ) ที่มีเสียงลบในช่วงนี้ แต่ยังไม่มี action ใดรับผิดชอบ — "ปัญหามาใหม่"
  - resolved_issues: action ที่ปิด (done_at) อยู่ในช่วงนี้ — พร้อมระยะเวลาที่ใช้แก้และผลตรวจ (better/same/worse)
  - open_issues:     action ที่ยังเปิดอยู่ (open/doing) ณ วันนี้จริง — ไม่ผูกกับช่วงรายงาน เพราะเป็นสถานะสด
  - chronic_branches:(สาขา,หัวข้อ) ที่มีเสียงลบใน >=3 จาก CHRONIC_LOOKBACK ช่วงล่าสุด (ช่วงนี้+ก่อนหน้า ยาวเท่ากัน) — "ผิดซ้ำสม่ำเสมอ"

  GET /api/feedback/daily-report?date=&date_to=&branch=          JSON (ใช้ render ในแอป)
  GET /api/feedback/daily-report.print?date=&date_to=&branch=      ฉบับพิมพ์ A4 (พิมพ์/บันทึก PDF)
"""
import datetime
import html
import io
import json
from typing import Optional

from fastapi import APIRouter, Query, Request, Response

from feedback_host import branch_rows, get_db, scope_branch, thai_today
from routers.feedback import _ROW_COLS, _aggregate, _ensure
from routers.feedback_actions import _ensure as _ensure_actions
from routers.feedback_actions import _list as _list_actions

router = APIRouter()

CHRONIC_LOOKBACK = 4          # this period + 3 before it, same length each
CHRONIC_MIN_HITS = 3          # a (branch, theme) showing up in >=3 of those 4 counts as "chronic"


def _esc(s) -> str:
    return html.escape(str(s if s is not None else ""))


def _d(s: str) -> datetime.date:
    return datetime.date.fromisoformat(s)


def _period_label(date_from: str, date_to: str) -> str:
    return date_from if date_from == date_to else f"{date_from} – {date_to}"


def _period_word(date_from: str, date_to: str) -> str:
    """วันนี้ for a single day, ช่วงนี้ for any range — used everywhere a section
    heading would otherwise hard-code "today" for what may be a week or month."""
    return "วันนี้" if date_from == date_to else "ช่วงนี้"


def _prev_period(date_from: str, date_to: str) -> tuple:
    """The comparison window: for one day, the SAME WEEKDAY 7 days back (a
    Friday-vs-Thursday comparison is meaningless for restaurant traffic — see
    mtd_weekday_fair_baseline); for a range, the immediately preceding period
    of equal length (a week/month already balances weekdays on its own)."""
    d_from, d_to = _d(date_from), _d(date_to)
    span = (d_to - d_from).days + 1
    if span == 1:
        p = d_from - datetime.timedelta(days=7)
        return p.isoformat(), p.isoformat()
    pe = d_from - datetime.timedelta(days=1)
    ps = pe - datetime.timedelta(days=span - 1)
    return ps.isoformat(), pe.isoformat()


def _headline(agg: dict, date_from: str, date_to: str, branch: Optional[str]) -> str:
    label = _period_label(date_from, date_to)
    if not agg["total"]:
        return f"ไม่มีเสียงลูกค้าเข้ามาในช่วง {label}" + (f" ที่สาขา {branch}" if branch else "")
    bits = [f"{agg['total']} เสียง", f"บวก {agg['sentiment']['pos']} · ลบ {agg['sentiment']['neg']} · กลาง {agg['sentiment']['neu']}"]
    if agg["avg_rating"] is not None:
        bits.append(f"เฉลี่ย {agg['avg_rating']}★")
    if agg["unreplied"]:
        bits.append(f"รอตอบ {agg['unreplied']}")
    if agg["sla_breached"]:
        bits.append(f"เกิน SLA {agg['sla_breached']}")
    if agg["severe"]:
        bits.append(f"⚠ ร้ายแรง {agg['severe']}")
    return " · ".join(bits)


def compute_daily_report(con, date: str, branch: Optional[str] = None, date_to: Optional[str] = None) -> dict:
    """[date, date_to] of customer_feedback (date_to defaults to date — one
    calendar day), through the same _aggregate() the tab's /api/feedback/summary
    uses — so the numbers never drift apart."""
    _ensure(con)
    _ensure_actions(con)
    date_from = date
    date_to = date_to or date_from
    if date_to < date_from:
        date_from, date_to = date_to, date_from
    w, p = ["rated_at >= ?", "rated_at < ?"], [date_from, (_d(date_to) + datetime.timedelta(days=1)).isoformat()]
    if branch:
        w.append("branch_name = ?"); p.append(branch)
    rows = [dict(r) for r in con.execute(f"SELECT {_ROW_COLS} FROM customer_feedback WHERE {' AND '.join(w)}", p)]
    agg = _aggregate(rows, today=date_to)
    p_from, p_to = _prev_period(date_from, date_to)
    wp, pp = ["rated_at >= ?", "rated_at < ?"], [p_from, (_d(p_to) + datetime.timedelta(days=1)).isoformat()]
    if branch:
        wp.append("branch_name = ?"); pp.append(branch)
    prev_total = con.execute(f"SELECT COUNT(*) FROM customer_feedback WHERE {' AND '.join(wp)}", pp).fetchone()[0]
    all_branches = [b["name"] for b in branch_rows(con) if b["active"]]
    today_real = thai_today()
    new_issues = _new_issues(con, rows, date_from, date_to)
    resolved_issues = _resolved_issues(con, today_real, date_from, date_to, branch)
    open_issues = _open_issues(con, today_real, branch)
    chronic_branches = _chronic_branches(con, date_from, date_to, branch)
    con.commit()
    return {"date_from": date_from, "date_to": date_to, "branch": branch, "branch_count": 1 if branch else len(all_branches),
            "prev_period": {"date_from": p_from, "date_to": p_to, "total": prev_total},
            "headline": _headline(agg, date_from, date_to, branch),
            "new_issues": new_issues, "resolved_issues": resolved_issues,
            "open_issues": open_issues, "chronic_branches": chronic_branches, **agg}


@router.get("/api/feedback/daily-report")
def feedback_daily_report(request: Request, date: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
                          branch: Optional[str] = Query(None)):
    scope = scope_branch(request)
    if scope:
        branch = scope
    d = (date or "")[:10] or thai_today()
    dt = (date_to or "")[:10] or None
    with get_db() as con:
        return compute_daily_report(con, d, branch, date_to=dt)


# ── print edition: one A4 page, teal/indigo to match the tab, testimonial-ready quotes ──
_PRINT_CSS = """
@page { size: A4; margin: 12mm 12mm 14mm 12mm; }
* { box-sizing: border-box; }
body { margin: 0; font: 10pt/1.45 "Sarabun", "Segoe UI", Tahoma, sans-serif; color: #1b232a; background: #fff; }
h1, h2, h3 { font-family: "Bai Jamjuree", "Sarabun", sans-serif; margin: 0; }
.head { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2px solid #0f766e; padding-bottom: 6pt; margin-bottom: 8pt; }
.brand { font: 700 8pt "Bai Jamjuree"; letter-spacing: .1em; text-transform: uppercase; color: #0f766e; }
.head h1 { font-size: 18pt; margin-top: 2pt; }
.head .meta { text-align: right; font-size: 8.5pt; color: #64748b; }
.headline { background: #f0fdfa; border: .6pt solid #99f6e4; border-radius: 4pt; padding: 6pt 9pt; font-size: 10.5pt; font-weight: 600; margin-bottom: 9pt; }
.kpis { display: grid; grid-template-columns: repeat(6, 1fr); gap: 6pt; margin-bottom: 9pt; }
.kpi { border: .6pt solid #d9e2df; border-radius: 4pt; padding: 5pt 7pt; }
.kpi .l { font-size: 7pt; text-transform: uppercase; letter-spacing: .04em; color: #64748b; }
.kpi .v { font-size: 15pt; font-weight: 700; color: #0f172a; }
.kpi .v.warn { color: #b45309; } .kpi .v.bad { color: #b91c1c; }
h2.sec { font-size: 11pt; border-left: 3pt solid #0f766e; padding-left: 6pt; margin: 10pt 0 5pt; }
table { width: 100%; border-collapse: collapse; font-size: 8.8pt; margin-bottom: 4pt; }
th { text-align: left; font-size: 7.6pt; color: #64748b; text-transform: uppercase; letter-spacing: .03em; border-bottom: .8pt solid #cbd5e1; padding: 2pt 5pt; }
td { padding: 3pt 5pt; border-bottom: .5pt solid #e2e8f0; vertical-align: top; }
.quote { border-left: 3pt solid #a7f3d0; background: #f0fdf4; padding: 4pt 8pt; margin: 4pt 0; border-radius: 0 4pt 4pt 0; font-size: 9.2pt; }
.quote .m { font-size: 7.6pt; color: #64748b; margin-top: 1pt; }
.item { border-left: 3pt solid #fca5a5; background: #fef2f2; padding: 4pt 8pt; margin: 4pt 0; border-radius: 0 4pt 4pt 0; font-size: 9pt; }
.item .m { font-size: 7.6pt; color: #64748b; margin-top: 1pt; }
.item .m a { color: inherit; text-decoration: underline; }
.tag { display: inline-block; font-size: 7.4pt; background: #fef3c7; color: #92400e; border-radius: 3pt; padding: 0 4pt; margin-right: 2pt; }
.two { display: grid; grid-template-columns: 1fr 1fr; gap: 0 14pt; }
.three { display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 0 12pt; break-inside: avoid-page; }
.empty { color: #94a3b8; font-size: 9pt; font-style: italic; }
.branch-h { font: 700 8.6pt "Bai Jamjuree"; color: #0f766e; margin: 7pt 0 2pt; padding-top: 3pt; border-top: .5pt dashed #cbd5e1; }
.branch-h:first-child { border-top: 0; padding-top: 0; }
.ex { font-size: 7.6pt; color: #64748b; font-style: italic; }
.foot { margin-top: 10pt; font-size: 7.6pt; color: #94a3b8; }
"""


def _tel_href(contact: str):
    digits = "".join(c for c in (contact or "") if c.isdigit() or c == "+")
    return "tel:" + digits if digits and 8 <= len(digits.lstrip("+")) <= 15 else None


def _group_by_branch(items: list) -> list:
    """[(branch, [items…]), …] in descending count order — used to break the
    negative/positive lists out per branch on the ทุกสาขา view (a flat blended
    list across many branches is hard to scan for a branch manager)."""
    groups: dict = {}
    for it in items:
        groups.setdefault(it.get("branch") or "(ไม่ระบุสาขา)", []).append(it)
    return sorted(groups.items(), key=lambda kv: -len(kv[1]))


def _new_issues(con, rows: list, date_from: str, date_to: str) -> list:
    """(สาขา, หัวข้อ) ที่มีเสียงลบในช่วงนี้ แต่ไม่มี action (open/doing) ใดรับผิดชอบ
    อยู่เลย — คือปัญหาที่ "มาใหม่" จริงๆ ยังไม่มีใครเริ่มตาม (ต่างจาก "ต้องติดตาม"
    ด้านบนซึ่งรวมทุกเสียงลบ ไม่ว่าจะมีคนตามอยู่แล้วหรือไม่)."""
    claimed = {(r["branch_name"] or "", r["theme"] or "") for r in
               con.execute("SELECT branch_name, theme FROM feedback_action WHERE status IN ('open','doing')")}
    tally: dict = {}
    for r in rows:
        if r.get("sentiment") != "neg":
            continue
        b = r.get("branch_name") or ""
        for t in json.loads(r.get("themes") or "[]"):
            if (b, t) in claimed:
                continue
            d = tally.setdefault((b, t), {"branch": b, "theme": t, "n": 0, "examples": []})
            d["n"] += 1
            if len(d["examples"]) < 2 and r.get("text"):
                d["examples"].append(r["text"][:140])
    return sorted(tally.values(), key=lambda x: -x["n"])


def _resolved_issues(con, today: str, date_from: str, date_to: str, branch: Optional[str]) -> list:
    """Action ที่ถูกปิด (done_at) ในช่วงรายงานนี้ — ระยะเวลาที่ใช้แก้ (จากวันสร้างถึง
    วันปิด) และผลตรวจ (better/same/worse/no_data, วัดจากเสียงลบ/100 ก่อน-หลัง)."""
    out = []
    for a in _list_actions(con, today, status="done", branch=branch):
        done_at = (a.get("done_at") or "")[:10]
        if not (done_at and date_from <= done_at <= date_to):
            continue
        created = (a.get("created_at") or "")[:10]
        days = (_d(done_at) - _d(created)).days if created else None
        out.append({"id": a["id"], "title": a["title"], "branch": a.get("branch_name") or "", "theme": a.get("theme") or "",
                    "owner": a.get("owner") or "", "created_at": created, "done_at": done_at, "days_to_fix": days,
                    "result": a.get("result"), "after_days_left": a.get("after_days_left"),
                    "baseline_rate": a.get("baseline_rate"), "after_rate": a.get("after_rate"), "note": a.get("note") or ""})
    out.sort(key=lambda x: x["done_at"], reverse=True)
    return out


def _open_issues(con, today: str, branch: Optional[str]) -> list:
    """Action ที่ยังไม่ปิด (open/doing) ณ วันนี้จริง — เป็นสถานะสด ไม่ผูกกับช่วง
    ของรายงาน (ปัญหาที่ยังไม่ได้แก้ไม่ใช่สิ่งที่ "เกิดเฉพาะวันนั้น")."""
    out = []
    for a in _list_actions(con, today, branch=branch):
        if a["status"] not in ("open", "doing"):
            continue
        created = (a.get("created_at") or "")[:10]
        days_open = (_d(today) - _d(created)).days if created else None
        out.append({"id": a["id"], "title": a["title"], "branch": a.get("branch_name") or "", "theme": a.get("theme") or "",
                    "owner": a.get("owner") or "", "due": a.get("due"), "overdue": a.get("overdue"),
                    "status": a["status"], "created_at": created, "days_open": days_open})
    out.sort(key=lambda x: (not x["overdue"], x["days_open"] is None, -(x["days_open"] or 0)))
    return out


def _chronic_branches(con, date_from: str, date_to: str, branch: Optional[str]) -> list:
    """(สาขา, หัวข้อ) ที่มีเสียงลบใน >=CHRONIC_MIN_HITS จาก CHRONIC_LOOKBACK ช่วง
    ล่าสุด (ช่วงนี้ + ช่วงก่อนหน้าที่ยาวเท่ากันติดต่อกัน) — สาขาที่ปัญหาเดิมวนกลับมา
    ซ้ำๆ ไม่ใช่แค่เกิดครั้งเดียวแล้วหาย."""
    span = (_d(date_to) - _d(date_from)).days + 1
    windows, f, t = [], date_from, date_to
    for _ in range(CHRONIC_LOOKBACK):
        windows.append((f, t))
        t = (_d(f) - datetime.timedelta(days=1)).isoformat()
        f = (_d(t) - datetime.timedelta(days=span - 1)).isoformat()
    seen: dict = {}
    for i, (wf, wt) in enumerate(windows):
        w, p = ["rated_at >= ?", "rated_at < ?", "sentiment = 'neg'"], [wf, (_d(wt) + datetime.timedelta(days=1)).isoformat()]
        if branch:
            w.append("branch_name = ?"); p.append(branch)
        for r in con.execute(f"SELECT branch_name, themes FROM customer_feedback WHERE {' AND '.join(w)}", p):
            b = r["branch_name"] or ""
            for th in json.loads(r["themes"] or "[]"):
                seen.setdefault((b, th), set()).add(i)
    out = [{"branch": b, "theme": th, "periods": len(idxs), "of_periods": CHRONIC_LOOKBACK}
           for (b, th), idxs in seen.items() if len(idxs) >= CHRONIC_MIN_HITS]
    out.sort(key=lambda x: (-x["periods"], x["branch"], x["theme"]))
    return out


def _build_print_html(data: dict) -> str:
    date_from, date_to, branch = data["date_from"], data["date_to"], data["branch"]
    label, pw = _period_label(date_from, date_to), _period_word(date_from, date_to)
    kind = "รายงานประจำวัน" if date_from == date_to else "รายงานตามช่วง"
    title = f"{kind} {label}" + (f" · {branch}" if branch else " · ทุกสาขา")
    kpis = [
        ("เสียงทั้งหมด", data["total"], ""),
        ("คะแนนเฉลี่ย", f"{data['avg_rating']:.2f}★" if data["avg_rating"] is not None else "—", ""),
        ("เชิงลบ", f"{data['neg_pct']}%", "bad" if data["neg_pct"] >= 30 else ""),
        ("รอตอบ", data["unreplied"], "warn" if data["unreplied"] else ""),
        ("เกิน SLA", data["sla_breached"], "bad" if data["sla_breached"] else ""),
        ("ร้ายแรง", data["severe"], "bad" if data["severe"] else ""),
    ]
    kpi_html = "".join(f'<div class="kpi"><div class="l">{_esc(l)}</div><div class="v {c}">{_esc(v)}</div></div>' for l, v, c in kpis)
    pp_ = data["prev_period"]
    pp_label = _period_label(pp_["date_from"], pp_["date_to"])
    cmp_line = f"เทียบช่วงก่อนหน้าที่ยาวเท่ากัน ({pp_label}): {pp_['total']} เสียง" + \
               (f" ({'มากขึ้น' if data['total'] > pp_['total'] else 'น้อยลง' if data['total'] < pp_['total'] else 'เท่าเดิม'})" if pp_["total"] else "")

    def theme_rows():
        if not data["themes"]:
            return f'<p class="empty">ยังไม่มีหัวข้อที่ถูกพูดถึง{pw}</p>'
        rows = []
        for t in data["themes"][:12]:
            ex = t["examples"][0]["text"] if t.get("examples") else ""
            rows.append(f'<tr><td>{_esc(t["theme"])}</td><td>{t["n"]}</td><td>{t["neg"]}</td><td>{", ".join(_esc(b) for b in t["branches"])}</td>'
                       f'<td class="ex">{f"“{_esc(ex)}”" if ex else ""}</td></tr>')
        return f'<table><thead><tr><th>หัวข้อ</th><th>พูดถึง</th><th>ลบ</th><th>สาขา</th><th>ตัวอย่าง</th></tr></thead><tbody>{"".join(rows)}</tbody></table>'

    def menu_rows():
        if not data["menus"]:
            return f'<p class="empty">ยังไม่มีเมนูที่ถูกพูดถึง{pw}</p>'
        rows = "".join(f'<tr><td>{_esc(m["menu"])}</td><td>{m["n"]}</td><td>{m["pos"]}</td><td>{m["neg"]}</td></tr>' for m in data["menus"][:12])
        return f'<table><thead><tr><th>เมนู</th><th>พูดถึง</th><th>บวก</th><th>ลบ</th></tr></thead><tbody>{rows}</tbody></table>'

    def _star(v):
        return f" · {v}★" if v is not None else ""

    def _neg_card(it):
        tel = _tel_href(it.get("contact") or "")
        contact_html = f' · <a href="{tel}">{_esc(it["contact"])}</a>' if tel else (f' · {_esc(it["contact"])}' if it.get("contact") else "")
        contacted_html = ' · <span style="color:#166534">โทรกลับแล้ว</span>' if it.get("contacted") else ""
        sev = ' <span class="tag">ร้ายแรง</span>' if it.get("severity") == "high" else ""
        status = "ตอบแล้ว" if it["replied"] else ("เกิน SLA" if it["sla_breached"] else "รอตอบ")
        date_bit = f'{_esc((it.get("date") or "")[:10])} · ' if date_from != date_to else ""
        branch_bit = "" if data["branch"] else f'{_esc(it["branch"] or "-")} · '
        return (f'<div class="item">“{_esc(it["text"])}”{sev}<div class="m">{date_bit}{branch_bit}{_esc(it["source"] or "")}'
               f'{_star(it.get("rating"))} · {status}{contact_html}{contacted_html}</div></div>')

    def _pos_card(q):
        date_bit = f'{_esc((q.get("date") or "")[:10])} · ' if date_from != date_to else ""
        branch_bit = "" if data["branch"] else f'{_esc(q["branch"] or "-")} · '
        return f'<div class="quote">“{_esc(q["text"])}”<div class="m">— {date_bit}{branch_bit}{_esc(q["source"] or "")}{_star(q.get("rating"))}</div></div>'

    def negative_items():
        if not data["recent_negative"]:
            return f'<p class="empty">ไม่มีรายการเชิงลบที่ต้องติดตาม{pw}</p>'
        if data["branch"]:
            return "".join(_neg_card(it) for it in data["recent_negative"])
        return "".join(f'<div class="branch-h">{_esc(b)} ({len(items)})</div>' + "".join(_neg_card(it) for it in items)
                       for b, items in _group_by_branch(data["recent_negative"]))

    def positive_quotes():
        if not data["recent_positive"]:
            return f'<p class="empty">ยังไม่มีเสียงชมที่หยิบไปใช้ได้{pw}</p>'
        if data["branch"]:
            return "".join(_pos_card(q) for q in data["recent_positive"][:5])
        return "".join(f'<div class="branch-h">{_esc(b)} ({len(items)})</div>' + "".join(_pos_card(q) for q in items)
                       for b, items in _group_by_branch(data["recent_positive"]))

    def _star_line(rate):
        return f"{rate}/100" if rate is not None else "—"

    def new_issues_block():
        items = data["new_issues"]
        if not items:
            return f'<p class="empty">ไม่มีปัญหาใหม่ที่ยังไม่มีคนตาม{pw}</p>'
        rows = []
        for it in items[:12]:
            ex = f' — “{_esc(it["examples"][0])}”' if it.get("examples") else ""
            branch_bit = "" if data["branch"] else f'{_esc(it["branch"] or "-")} · '
            rows.append(f'<div class="item">{branch_bit}<b>{_esc(it["theme"])}</b> ({it["n"]} เสียงลบ){ex}</div>')
        return "".join(rows)

    def open_issues_block():
        items = data["open_issues"]
        if not items:
            return '<p class="empty">ไม่มีปัญหาที่ค้างอยู่</p>'
        rows = []
        for it in items[:15]:
            branch_bit = "" if data["branch"] else f'{_esc(it["branch"] or "-")} · '
            due = f' · กำหนด {_esc(it["due"])}' if it.get("due") else ""
            overdue = ' <span class="tag">เกินกำหนด</span>' if it.get("overdue") else ""
            status = "กำลังทำ" if it["status"] == "doing" else "ยังไม่เริ่ม"
            days = f' · เปิดมา {it["days_open"]} วัน' if it.get("days_open") is not None else ""
            rows.append(f'<div class="item">{branch_bit}<b>{_esc(it["title"])}</b>{overdue}'
                       f'<div class="m">{_esc(it.get("theme") or "-")} · {status} · {_esc(it.get("owner") or "ยังไม่มอบหมาย")}{due}{days}</div></div>')
        return "".join(rows)

    _RESULT_TH = {"better": "ดีขึ้น", "same": "เท่าเดิม", "worse": "แย่ลง", "no_data": "รอผล"}

    def resolved_issues_block():
        items = data["resolved_issues"]
        if not items:
            return '<p class="empty">ไม่มีปัญหาที่ปิดในช่วงนี้</p>'
        rows = []
        for it in items[:15]:
            branch_bit = "" if data["branch"] else f'{_esc(it["branch"] or "-")} · '
            days = f'{it["days_to_fix"]} วัน' if it.get("days_to_fix") is not None else "—"
            result = _RESULT_TH.get(it.get("result"), "รอผล" if it.get("after_days_left") is not None else "—")
            rate_line = f' (ก่อน {_star_line(it.get("baseline_rate"))} → หลัง {_star_line(it.get("after_rate"))})' if it.get("after_rate") is not None else ""
            rows.append(f'<div class="quote">{branch_bit}<b>{_esc(it["title"])}</b><div class="m">{_esc(it.get("theme") or "-")} · '
                       f'ปิดเมื่อ {_esc(it["done_at"])} · ใช้เวลาแก้ {days} · ผล {result}{rate_line}</div></div>')
        return "".join(rows)

    def chronic_block():
        items = data["chronic_branches"]
        if not items:
            return f'<p class="empty">ไม่มีสาขาที่ปัญหาเดิมวนซ้ำใน {CHRONIC_LOOKBACK} ช่วงล่าสุด</p>'
        claimed = {(it["branch"] or "", it["theme"] or "") for it in data["open_issues"]}
        rows = "".join(f'<tr><td>{_esc(it["branch"] or "-")}</td><td>{_esc(it["theme"])}</td><td>{it["periods"]}/{it["of_periods"]} ช่วง</td>'
                       f'<td>{"มีงานอยู่แล้ว" if (it["branch"] or "", it["theme"] or "") in claimed else "ยังไม่มีงาน"}</td></tr>'
                       for it in items[:15])
        return f'<table><thead><tr><th>สาขา</th><th>หัวข้อที่ซ้ำ</th><th>ความถี่</th><th>สถานะ</th></tr></thead><tbody>{rows}</tbody></table>'

    def branch_table():
        if data["branch"] or not data["branches"]:
            return ""
        rows = "".join(f'<tr><td>{_esc(b["branch"])}</td><td>{b["n"]}</td><td>{b["neg_pct"]}%</td><td>{b["avg_rating"] if b["avg_rating"] is not None else "—"}</td>'
                       f'<td>{b["unreplied"]}</td><td>{b["flagged"]}</td><td>{_esc(b["top_theme"] or "—")}</td></tr>' for b in data["branches"][:15])
        return (f'<h2 class="sec">รายสาขาที่มีเสียง{pw} ({len(data["branches"])})</h2>'
               f'<table><thead><tr><th>สาขา</th><th>เสียง</th><th>ลบ%</th><th>★</th><th>รอตอบ</th><th>ปักหมุด</th><th>หัวข้อเด่น</th></tr></thead><tbody>{rows}</tbody></table>')

    def sources_block():
        if not data["sources"]:
            return '<p class="empty">—</p>'
        items = sorted(data["sources"].items(), key=lambda kv: -kv[1])
        return "".join(f'<div style="display:flex;justify-content:space-between"><span>{_esc(s)}</span><b>{n}</b></div>' for s, n in items)

    def response_block():
        r = data["response"]
        rate = f'{r["rate"]}%' if r["rate"] is not None else "—"
        median = f'{r["median_days"]} วัน' if r["median_days"] is not None else "—"
        return (f'<div style="display:flex;justify-content:space-between"><span>ตอบได้ทั้งหมด</span><b>{r["answerable"]}</b></div>'
               f'<div style="display:flex;justify-content:space-between"><span>ตอบแล้ว</span><b>{r["answered"]}</b></div>'
               f'<div style="display:flex;justify-content:space-between"><span>อัตราตอบกลับ</span><b>{rate}</b></div>'
               f'<div style="display:flex;justify-content:space-between"><span>ตอบไวเฉลี่ย</span><b>{median}</b></div>')

    def staff_block():
        if not data["staff"]:
            return f'<p class="empty">ยังไม่มีพนักงานถูกระบุชื่อ{pw}</p>'
        return "".join(f'<div style="display:flex;justify-content:space-between"><span>{_esc(s["name"])} <span class="ex">({_esc(s["branch"])})</span></span>'
                       f'<b>{s["pos"]} {s["neg"]}</b></div>' for s in data["staff"][:8])

    return f"""<!doctype html>
<html lang="th"><head><meta charset="utf-8"><title>{_esc(title)}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Bai+Jamjuree:wght@600;700&family=Sarabun:wght@400;500;600;700&display=swap">
<style>{_PRINT_CSS}</style></head><body>
<div class="head"><div><div class="brand">Oshinei Dashboard · โมดูล เสียงลูกค้า</div><h1>{_esc(title)}</h1></div>
<div class="meta">พิมพ์เมื่อ {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}<br>{data['branch_count']} สาขาในรายงานนี้</div></div>
<div class="headline">{_esc(data['headline'])}<br><span style="font-weight:400;font-size:9pt;color:#3a4a45">{_esc(cmp_line)}</span></div>
<div class="kpis">{kpi_html}</div>
<div class="two">
<div><h2 class="sec">หัวข้อที่ถูกพูดถึง{pw}</h2>{theme_rows()}</div>
<div><h2 class="sec">เมนูที่ถูกพูดถึง{pw}</h2>{menu_rows()}</div>
</div>
<h2 class="sec">ต้องติดตาม{pw} ({len(data['recent_negative'])})</h2>
{negative_items()}
<h2 class="sec">ปัญหาใหม่{pw} ยังไม่มีคนตาม ({len(data['new_issues'])})</h2>
{new_issues_block()}
<h2 class="sec">ปัญหาที่ยังไม่ได้แก้ ({len(data['open_issues'])})</h2>
{open_issues_block()}
<h2 class="sec">ปัญหาที่แก้แล้วช่วงนี้ ({len(data['resolved_issues'])})</h2>
{resolved_issues_block()}
<h2 class="sec">สาขาที่มีความผิดซ้ำสม่ำเสมอ ({len(data['chronic_branches'])})</h2>
{chronic_block()}
<h2 class="sec">เสียงชมที่หยิบไปใช้ได้ ({len(data['recent_positive'])})</h2>
{positive_quotes()}
<div class="three">
<div><h2 class="sec">ช่องทางที่มา{pw}</h2>{sources_block()}</div>
<div><h2 class="sec">อัตราตอบกลับ</h2>{response_block()}</div>
<div><h2 class="sec">พนักงานที่ถูกพูดถึง</h2>{staff_block()}</div>
</div>
{branch_table()}
<div class="foot">Oshinei · {_esc(kind)}เสียงลูกค้า · คำนวณจากข้อมูลชุดเดียวกับแท็บ เสียงลูกค้าในแอป</div>
</body></html>"""


def _build_xlsx(data: dict):
    """The same report as a workbook — one sheet per section, for marketing to
    paste into a deck/sheet themselves (the print edition is for reading/PDF,
    this is for reuse). Mirrors reports.py's header()/autofit() styling."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    HDR = Font(bold=True, color="FFFFFF")
    BOLD = Font(bold=True)
    FILL_H = PatternFill("solid", fgColor="0F766E")
    CENTER = Alignment(horizontal="center")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = HDR; c.fill = FILL_H; c.alignment = CENTER

    def autofit(ws):
        for col in ws.columns:
            mx = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(mx + 2, 60)

    label = _period_label(data["date_from"], data["date_to"])
    ws = wb.active; ws.title = "Summary"
    ws.append(["Oshinei — รายงานเสียงลูกค้า"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append(["ช่วง", label]); ws.append(["สาขา", data["branch"] or "ทุกสาขา"])
    ws.append([]); ws.append([data["headline"]]); ws["A5"].font = BOLD
    pp = data["prev_period"]
    ws.append(["เทียบช่วงก่อนหน้า", f'{pp["date_from"]} – {pp["date_to"]}', pp["total"]])
    ws.append([])
    header(ws, ["ตัวชี้วัด", "ค่า"])
    for l, v in (("เสียงทั้งหมด", data["total"]), ("คะแนนเฉลี่ย", data["avg_rating"]), ("บวก", data["sentiment"]["pos"]),
                ("ลบ", data["sentiment"]["neg"]), ("กลาง", data["sentiment"]["neu"]), ("เชิงลบ %", data["neg_pct"]),
                ("รอตอบ", data["unreplied"]), ("เกิน SLA", data["sla_breached"]), ("ร้ายแรง", data["severe"])):
        ws.append([l, v])
    autofit(ws)

    ws = wb.create_sheet("หัวข้อ")
    header(ws, ["หัวข้อ", "พูดถึง", "ลบ", "บวก", "สาขา", "ตัวอย่าง"])
    for t in data["themes"]:
        ws.append([t["theme"], t["n"], t["neg"], t["pos"], ", ".join(t["branches"]), (t["examples"][0]["text"] if t.get("examples") else "")])
    autofit(ws)

    ws = wb.create_sheet("เมนู")
    header(ws, ["เมนู", "พูดถึง", "บวก", "ลบ"])
    for m in data["menus"]:
        ws.append([m["menu"], m["n"], m["pos"], m["neg"]])
    autofit(ws)

    ws = wb.create_sheet("ต้องติดตาม")
    header(ws, ["วันที่", "สาขา", "แหล่ง", "ดาว", "ข้อความ", "สถานะ", "ติดต่อกลับ", "โทรกลับแล้ว"])
    for it in data["recent_negative"]:
        ws.append([(it.get("date") or "")[:10], it.get("branch") or "-", it.get("source") or "", it.get("rating"), it.get("text") or "",
                  "ตอบแล้ว" if it["replied"] else ("เกิน SLA" if it["sla_breached"] else "รอตอบ"), it.get("contact") or "",
                  "ใช่" if it.get("contacted") else ""])
    autofit(ws)

    ws = wb.create_sheet("เสียงชม")
    header(ws, ["วันที่", "สาขา", "แหล่ง", "ดาว", "ข้อความ"])
    for q in data["recent_positive"]:
        ws.append([(q.get("date") or "")[:10], q.get("branch") or "-", q.get("source") or "", q.get("rating"), q.get("text") or ""])
    autofit(ws)

    ws = wb.create_sheet("งานแก้ปัญหา")
    header(ws, ["สถานะ", "สาขา", "หัวข้อ/ชื่องาน", "รายละเอียด", "วันที่/ผล"])
    for it in data["new_issues"]:
        ws.append(["ใหม่ ยังไม่มีคนตาม", it["branch"] or "-", it["theme"], f'{it["n"]} เสียงลบ', ""])
    for it in data["open_issues"]:
        ws.append(["ยังไม่ได้แก้" + (" (เกินกำหนด)" if it.get("overdue") else ""), it["branch"] or "-", it["title"],
                  it.get("theme") or "", f'{it.get("owner") or "ยังไม่มอบหมาย"} · กำหนด {it.get("due") or "-"}'])
    for it in data["resolved_issues"]:
        result = {"better": "ดีขึ้น", "same": "เท่าเดิม", "worse": "แย่ลง"}.get(it.get("result"), "รอผล")
        ws.append(["แก้แล้ว", it["branch"] or "-", it["title"], it.get("theme") or "",
                  f'ปิดเมื่อ {it["done_at"]} · ใช้เวลา {it.get("days_to_fix")} วัน · ผล {result}'])
    autofit(ws)

    ws = wb.create_sheet("ปัญหาซ้ำซาก")
    header(ws, ["สาขา", "หัวข้อ", "ความถี่", "สถานะ"])
    claimed = {(it["branch"] or "", it["theme"] or "") for it in data["open_issues"]}
    for it in data["chronic_branches"]:
        status = "มีงานอยู่แล้ว" if (it["branch"] or "", it["theme"] or "") in claimed else "ยังไม่มีงาน"
        ws.append([it["branch"] or "-", it["theme"], f'{it["periods"]}/{it["of_periods"]} ช่วง', status])
    autofit(ws)

    if not data["branch"] and data["branches"]:
        ws = wb.create_sheet("รายสาขา")
        header(ws, ["สาขา", "เสียง", "ลบ %", "★", "รอตอบ", "ปักหมุด", "หัวข้อเด่น"])
        for b in data["branches"]:
            ws.append([b["branch"], b["n"], b["neg_pct"], b["avg_rating"], b["unreplied"], b["flagged"], b["top_theme"] or "—"])
        autofit(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/api/feedback/daily-report.xlsx")
def feedback_daily_report_xlsx(request: Request, date: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
                               branch: Optional[str] = Query(None)):
    scope = scope_branch(request)
    if scope:
        branch = scope
    d = (date or "")[:10] or thai_today()
    dt = (date_to or "")[:10] or None
    with get_db() as con:
        data = compute_daily_report(con, d, branch, date_to=dt)
    buf = _build_xlsx(data)
    fname = f"FEEDBACK_REPORT_{data['date_from']}_{data['date_to']}" + (f"_{branch}" if branch else "") + ".xlsx"
    return Response(buf.read(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={fname}"})


@router.get("/api/feedback/daily-report.print")
def feedback_daily_report_print(request: Request, date: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
                                branch: Optional[str] = Query(None)):
    scope = scope_branch(request)
    if scope:
        branch = scope
    d = (date or "")[:10] or thai_today()
    dt = (date_to or "")[:10] or None
    with get_db() as con:
        data = compute_daily_report(con, d, branch, date_to=dt)
    return Response(_build_print_html(data), media_type="text/html")


__all__ = ["router", "compute_daily_report"]
