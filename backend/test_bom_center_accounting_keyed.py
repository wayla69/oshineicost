# -*- coding: utf-8 -*-
"""Sales Support marks the ใบสั่งซื้อ it has keyed into the external accounting system (PR-E).

The promises:
  * only an order the warehouse has processed (or the branch has received) can be marked keyed;
    the mark is audited, can be taken back, and an empty list or an unknown order is refused;
  * the queue carries the mark and can be filtered by it;
  * หน้างาน shows Sales Support how many processed orders are still to be keyed;
  * the consolidated workbook's "ต่อสาขา" sheet ends with a คีย์บัญชี column;
  * the mark is Sales Support's (management too): the warehouse, the analyst and the branch cannot set it,
    and the longer `/hub/keyed` rule beats the warehouse-only `/hub` rule while `/hub/process` stays the warehouse's.
"""
import io
import sqlite3
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import auth
from bom_center import hub as HB
from bom_center import todo as TD
from test_bom_center_s4b import con  # noqa: F401  (schema ensured on a scratch DB)

TOKEN = {"X-Upload-Token": "owner-token-xyz"}


def _order(con, status, cycle="2026-10-05", rnd="1"):  # noqa: F811
    con.execute("INSERT INTO bc_order(branch_id, cycle_date, round, status, value) VALUES(1, ?, ?, ?, 0)",
                (cycle, rnd, status))
    con.commit()                                                                    # a later rollback keeps the order
    return con.execute("SELECT MAX(order_id) FROM bc_order").fetchone()[0]


def _audit_actions(con):  # noqa: F811
    return [r[0] for r in con.execute("SELECT action FROM bc_audit WHERE action LIKE 'order_accounting_%' ORDER BY id")]


# ── marking ──────────────────────────────────────────────────────────────────
def test_only_processed_or_received_orders_can_be_keyed(con):  # noqa: F811
    sub, proc, rec = _order(con, "submitted"), _order(con, "processed", "2026-10-06"), _order(con, "received", "2026-10-07")
    with pytest.raises(ValueError) as e:
        HB.set_keyed(con, [proc, sub], True, "ss")
    assert HB.order_label_th(con, sub) in str(e.value) and f"#{sub}" not in str(e.value) and "คลังยังไม่ยืนยันยอดส่ง จึงคีย์เข้าระบบบัญชีไม่ได้" in str(e.value)
    con.rollback()
    assert con.execute("SELECT COUNT(*) FROM bc_order WHERE accounting_keyed_at IS NOT NULL").fetchone()[0] == 0
    out = HB.set_keyed(con, [proc, rec], True, "ss1")
    assert out["updated"] == 2 and out["keyed"] is True and out["at"]
    rows = con.execute("SELECT accounting_keyed_at, accounting_keyed_by FROM bc_order WHERE order_id IN (?,?)",
                       (proc, rec)).fetchall()
    assert all(r[0] == out["at"] and r[1] == "ss1" for r in rows)
    assert _audit_actions(con) == ["order_accounting_keyed"] * 2


def test_unkeying_clears_the_mark_and_is_audited(con):  # noqa: F811
    proc = _order(con, "processed")
    HB.set_keyed(con, [proc], True, "ss")
    out = HB.set_keyed(con, [proc], False, "ss")
    assert out == {"updated": 1, "keyed": False, "at": None}
    assert con.execute("SELECT accounting_keyed_at, accounting_keyed_by FROM bc_order WHERE order_id=?",
                       (proc,)).fetchone() == (None, None)
    assert _audit_actions(con) == ["order_accounting_keyed", "order_accounting_unkeyed"]


def test_empty_and_unknown_orders_are_refused(con):  # noqa: F811
    with pytest.raises(ValueError, match="เลือกใบสั่งซื้ออย่างน้อย 1 ใบ"):
        HB.set_keyed(con, [], True, "ss")
    proc = _order(con, "processed")
    with pytest.raises(ValueError, match="ไม่พบใบสั่งซื้อที่เลือก"):
        HB.set_keyed(con, [proc, 99999], True, "ss")
    con.rollback()


# ── the queue ────────────────────────────────────────────────────────────────
def test_the_queue_carries_and_filters_the_mark(con):  # noqa: F811
    a, b, c = _order(con, "processed"), _order(con, "received", "2026-10-06"), _order(con, "submitted", "2026-10-07")
    HB.set_keyed(con, [a], True, "ss")
    allq = {r["order_id"]: r for r in HB.queue(con, "all")}
    assert allq[a]["accounting_keyed"] is True and allq[a]["accounting_keyed_by"] == "ss"
    assert allq[a]["accounting_keyed_at"] and len(allq[a]["accounting_keyed_date"]) == 10
    assert allq[b]["accounting_keyed"] is False and allq[b]["accounting_keyed_date"] is None
    assert [r["order_id"] for r in HB.queue(con, "all", keyed=1)] == [a]
    assert [r["order_id"] for r in HB.queue(con, "all", keyed=0)] == [b, c]
    assert [r["order_id"] for r in HB.queue(con, "processed", keyed=0)] == []
    assert {r["order_id"] for r in HB.queue(con, "all")} == {a, b, c}               # no filter = everything, as before


# ── หน้างาน ──────────────────────────────────────────────────────────────────
def test_the_todo_counts_processed_orders_not_yet_keyed(con):  # noqa: F811
    today = date(2026, 10, 7)
    it = {i["key"]: i for i in TD.build(con, today=today)["items"]}["accounting_keyed"]
    assert it["status"] == "done" and it["detail"] == "คีย์ครบทุกใบ" and it["count"] is None
    a, b = _order(con, "processed"), _order(con, "received", "2026-10-06")
    _order(con, "submitted", "2026-10-07")                                          # not processed: not counted
    it = {i["key"]: i for i in TD.build(con, today=today)["items"]}["accounting_keyed"]
    assert it["status"] == "todo" and it["count"] == 2 and it["seq"] == 227
    assert it["title"] == "คีย์ใบสั่งซื้อที่คลังจัดส่งแล้วเข้าระบบบัญชี"
    assert it["detail"] == "ยังไม่คีย์ 2 ใบ · ดาวน์โหลดสรุปรวมตามวันส่ง (Excel) แล้วกด “คีย์บัญชีแล้ว”"
    assert it["owner_modules"] == ["supply.ordering"] and it["owner_th"] == "ประสานงานขาย"
    assert it["go"] == {"tab": "orderhub", "pane": "queue"}
    HB.set_keyed(con, [a, b], True, "ss")
    it = {i["key"]: i for i in TD.build(con, today=today)["items"]}["accounting_keyed"]
    assert it["status"] == "done" and it["count"] is None


# ── the permission matrix ────────────────────────────────────────────────────
def test_the_longer_keyed_rule_beats_the_warehouse_only_hub_rule():
    pa = auth.path_allowed
    sales, wh = auth.modules_for_role("salessupport"), auth.modules_for_role("warehouse")
    analyst, mgmt = auth.modules_for_role("analyst"), auth.modules_for_role("management")
    path = "/api/bom-center/hub/keyed"
    assert pa("POST", path, sales) and pa("POST", path, mgmt)
    assert not pa("POST", path, wh) and not pa("POST", path, analyst)
    assert not pa("POST", path, auth.modules_for_role("branch"))
    # the warehouse's own write is untouched, and Sales Support still cannot process
    assert pa("POST", "/api/bom-center/hub/process/1", wh) and pa("POST", "/api/bom-center/hub/process/1", mgmt)
    assert not pa("POST", "/api/bom-center/hub/process/1", sales)
    assert pa("GET", "/api/bom-center/hub/queue", sales) and pa("GET", "/api/bom-center/hub/queue", wh)


# ── over HTTP ────────────────────────────────────────────────────────────────
@pytest.fixture
def auth_on(monkeypatch):
    monkeypatch.setenv("OSHINEI_AUTH", "1")
    monkeypatch.setenv("OSHINEI_AUTH_SECRET", "test-secret-stable")
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "owner-token-xyz")
    yield


def _login(user, pw=auth.DEFAULT_PASSWORD):
    import main
    c = TestClient(main.app)
    assert c.post("/api/auth/login", json={"username": user, "password": pw}).status_code == 200
    return c


def _seed_orders(db_path):
    c = sqlite3.connect(str(db_path))
    bid = c.execute("SELECT branch_id FROM branches WHERE name<>'ALL' ORDER BY branch_id LIMIT 1").fetchone()[0]
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv, supplier, active)"
              " VALUES('T-KEY1','ของทดสอบ','raw','กรัม','แพ็ค',1000,'ผู้จำหน่าย ก',1)")
    ids = []
    for status, cyc in (("processed", "2098-03-02"), ("submitted", "2098-03-04")):
        c.execute("INSERT INTO bc_order(branch_id, cycle_date, round, status, value) VALUES(?,?, '1', ?, 0)", (bid, cyc, status))
        ids.append(c.execute("SELECT MAX(order_id) FROM bc_order").fetchone()[0])
        c.execute("INSERT INTO bc_order_line(order_id, ingredient_code, suggested_units, confirmed_units, processed_units, order_unit)"
                  " VALUES(?, 'T-KEY1', 2, 2, 2, 'pack')", (ids[-1],))
    c.commit()
    c.close()
    return ids


def _cleanup(db_path, ids):
    c = sqlite3.connect(str(db_path))
    for i in ids:
        c.execute("DELETE FROM bc_order_line WHERE order_id=?", (i,))
        c.execute("DELETE FROM bc_order WHERE order_id=?", (i,))
        c.execute("DELETE FROM bc_audit WHERE row_key=? AND action LIKE 'order_accounting_%'", (str(i),))
    c.execute("DELETE FROM bc_ingredient WHERE code='T-KEY1'")
    c.commit()
    c.close()


def _warehouse_client(mg):
    r = mg.post("/api/auth/users", headers=TOKEN,
                json={"username": "wh_keyed_test", "password": "Temp@12345", "role": "warehouse"})
    assert r.status_code in (200, 409), r.text
    return _login("wh_keyed_test", "Temp@12345")


def test_http_flow_roles_filter_and_excel(client, app_db, auth_on):
    import main
    client.get("/api/bom-center/hub/queue")                                         # creates the bc_* tables
    ids = _seed_orders(main.DB_PATH)
    proc, sub = ids
    try:
        ss, mg = _login("salessupport1"), _login("management1")
        wh = _warehouse_client(mg)
        body = {"order_ids": [proc], "keyed": True}
        # the warehouse is refused by role (not by token); Sales Support needs the owner token like every central write
        r = wh.post("/api/bom-center/hub/keyed", json=body, headers=TOKEN)
        assert r.status_code == 403 and r.headers.get("X-Denied") == "role"
        r = ss.post("/api/bom-center/hub/keyed", json=body)
        assert r.status_code == 403 and r.headers.get("X-Denied") == "token"
        # ... while the warehouse may still process, and Sales Support may not
        assert wh.post("/api/bom-center/hub/process/0", json={}, headers=TOKEN).headers.get("X-Denied") != "role"
        assert ss.post("/api/bom-center/hub/process/0", json={}, headers=TOKEN).headers.get("X-Denied") == "role"
        # refusals are 400 with the Thai message
        r = ss.post("/api/bom-center/hub/keyed", json={"order_ids": [], "keyed": True}, headers=TOKEN)
        assert r.status_code == 400 and "เลือกใบสั่งซื้ออย่างน้อย 1 ใบ" in r.json()["detail"]
        r = ss.post("/api/bom-center/hub/keyed", json={"order_ids": [sub], "keyed": True}, headers=TOKEN)
        assert r.status_code == 400 and "คีย์เข้าระบบบัญชีไม่ได้" in r.json()["detail"]
        # the queue says "not keyed" before, and the Excel ends with the column
        q = "date_from=2098-03-01&date_to=2098-03-31"
        orders = ss.get(f"/api/bom-center/hub/queue?status=all&{q}").json()["orders"]
        assert {o["order_id"]: o["accounting_keyed"] for o in orders} == {proc: False, sub: False}
        wb = load_workbook(io.BytesIO(ss.get(f"/api/bom-center/hub/consolidated.xlsx?{q}&status=all").content), read_only=True)
        rows = list(wb["ต่อสาขา"].iter_rows(min_row=1, values_only=True))
        header = next(r for r in rows if r and r[0] == "สาขา")
        assert header[-1] == "คีย์บัญชี"
        data = [r for r in rows if r and r[9] == proc]
        assert data and data[0][-1] == "ยังไม่คีย์"
        # mark it
        r = ss.post("/api/bom-center/hub/keyed", json=body, headers=TOKEN)
        assert r.status_code == 200 and r.json()["updated"] == 1 and r.json()["keyed"] is True
        assert [o["order_id"] for o in ss.get(f"/api/bom-center/hub/queue?status=all&keyed=1&{q}").json()["orders"]] == [proc]
        assert [o["order_id"] for o in ss.get(f"/api/bom-center/hub/queue?status=all&keyed=0&{q}").json()["orders"]] == [sub]
        assert ss.get("/api/bom-center/hub/queue?keyed=2").status_code == 400
        wb = load_workbook(io.BytesIO(ss.get(f"/api/bom-center/hub/consolidated.xlsx?{q}&status=all").content), read_only=True)
        rows = list(wb["ต่อสาขา"].iter_rows(values_only=True))
        done = next(r for r in rows if r and r[9] == proc)
        assert done[-1].startswith("คีย์แล้ว 20") and len(done[-1]) == len("คีย์แล้ว 2026-10-09")
        # take it back
        r = ss.post("/api/bom-center/hub/keyed", json={"order_ids": [proc], "keyed": False}, headers=TOKEN)
        assert r.status_code == 200 and r.json()["keyed"] is False
        assert [o["order_id"] for o in ss.get(f"/api/bom-center/hub/queue?status=all&keyed=1&{q}").json()["orders"]] == []
    finally:
        _cleanup(main.DB_PATH, ids)
