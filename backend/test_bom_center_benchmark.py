# -*- coding: utf-8 -*-
"""BoM Center — PR-X: the benchmark ("ค่ากลาง") from the company-owned branches.

The promises:
  * the branches that count are a setting (default: อุบลราชธานี, อุดรธานี, โคราช1, สุรินทร์ by name); the benchmark of a figure is the
    MEDIAN of the benchmark branches that have it (at least 2, else none); the buffer is 0-50 and defaults to 15;
  * a branch-month with a submitted begin AND end count reads its purchases with the stock change and is judged by the normal
    variance threshold; every other branch by the buffer;
  * the per-฿100,000 report carries the benchmark, "เทียบค่ากลาง %" on every branch and item, and the chip "ต่ำกว่าค่ากลาง ·
    อาจซื้อจากภายนอก" on a branch OUTSIDE the benchmark only; a branch login sees none of it;
  * a new branch's default reference is the benchmark (one chosen branch, or none, still possible);
  * the recipe check reads the price file's per-branch purchases against the recipes, per ฿100,000 of sales, with the statuses
    ok / more / less / x10 / no_recipe / no_buy, scales purchases by the days that have daily sales, and is the analyst's alone;
  * หน้างาน lists the ten-times rows for the analyst.
"""
import io
import sqlite3
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

import auth
from bom_center import benchmark as BM
from bom_center import forecast as FC
from bom_center import orders as OR
from bom_center import per_sales as PS
from bom_center import recipe_check as RC
from bom_center import stock as ST
from bom_center import todo as TD
from bom_center.clock import BKK
from test_bom_center_issue_summary import TOKEN, _login, _restore, _snap, auth_on  # noqa: F401
from test_bom_center_s4b import PATTERN, con  # noqa: F401  (branch 1: crab 45 g on sku 10, salmon 20 g on sku 11, 1 Sep - 5 Oct 2026)

CRAB, SALMON, VEG, SEAWEED = "13-01-009", "02-01-001", "V-007", "11-02-073"
WED = date(2026, 10, 7)
NOW = datetime(2026, 10, 5, 8, 0, tzinfo=BKK)
AS_OF = date(2026, 10, 5)
TODAY = date(2026, 10, 12)


@pytest.fixture(autouse=True)
def _flat_season(monkeypatch):
    monkeypatch.setattr(FC, "season_factor", lambda *a, **k: {"factor": 1.0, "mode": "off", "reason": ""})
    monkeypatch.setattr(PS, "bkk_today", lambda: TODAY)


def _days(a, b):
    return [a + timedelta(days=i) for i in range((b - a).days + 1)]


def _branch(c, bid, name, typ):
    c.execute("INSERT OR REPLACE INTO branches(branch_id, name, branch_type, active) VALUES(?,?,?,1)", (bid, name, typ))


def _sell(c, bid, a, b, q10=10, q11=10, p10=100, p11=200):
    """Branch `bid` sells q10 plates of sku 10 (฿p10) and q11 of sku 11 (฿p11) every day from a to b (a zero quantity: not sold)."""
    for d in _days(a, b):
        if q10:
            c.execute("INSERT INTO sales_sku_daily VALUES(?,?,10,?,?)", (bid, d.isoformat(), q10, q10 * p10))
        if q11:
            c.execute("INSERT INTO sales_sku_daily VALUES(?,?,11,?,?)", (bid, d.isoformat(), q11, q11 * p11))


@pytest.fixture
def sold(con):  # noqa: F811
    """The forecast fixture with prices on the dishes (sku 10 = ฿100 a plate, sku 11 = ฿200) and the company branches' names."""
    con.execute("UPDATE sales_sku_daily SET subtotal_ex_vat = qty * 100 WHERE sku_id=10")
    con.execute("UPDATE sales_sku_daily SET subtotal_ex_vat = qty * 200 WHERE sku_id=11")
    con.commit()
    return con


# ── 1 · which branches count ─────────────────────────────────────────────────────────────────
def test_default_set_resolves_the_four_names(sold):
    assert BM.default_ids(sold) == [] and BM.branch_ids(sold) == []                    # the fixture has none of the four names
    for i, (n, t) in enumerate([("อุบลราชธานี", "master"), ("อุดรธานี", "master"), ("โคราช 1", "master"), ("สุรินทร์", "master"),
                                ("เชียงใหม่", "master"), ("หาดใหญ่", "franchise")], start=20):
        _branch(sold, i, n, t)
    sold.execute("INSERT INTO branches(branch_id, name, branch_type, active) VALUES(30, 'โคราช2', 'master', 0)")
    sold.commit()
    assert BM.default_ids(sold) == [20, 21, 22, 23]                                    # the owner's four, in his order; เชียงใหม่ and a closed branch are out
    assert BM.branch_ids(sold) == [20, 21, 22, 23] and BM.settings(sold)["is_default"] is True
    s = BM.settings(sold)
    assert s["buffer_pct"] == 15.0 and s["buffer_default"] == 15.0 and s["min_branches"] == 2
    assert set(s["selected"]) == {20, 21, 22, 23}


def test_the_set_is_saved_audited_and_validated(sold):
    for i in (3, 4):
        _branch(sold, i, f"Co{i}", "master")
    out = BM.set_benchmark(sold, "sales", [1, 3, 4])
    assert out["selected"] == [1, 3, 4] and out["is_default"] is False and BM.branch_ids(sold) == [1, 3, 4]
    assert sold.execute("SELECT value FROM app_settings WHERE key='bc_benchmark_branches'").fetchone()[0] == "[1, 3, 4]"
    a = sold.execute("SELECT action, who, new_value FROM bc_audit WHERE action='benchmark_branches'").fetchone()
    assert a == ("benchmark_branches", "sales", "[1, 3, 4]")
    with pytest.raises(ValueError, match="อย่างน้อย 2 สาขา"):
        BM.set_benchmark(sold, "sales", [1])
    with pytest.raises(ValueError, match="ไม่พบสาขา"):
        BM.set_benchmark(sold, "sales", [1, 99])
    assert BM.branch_ids(sold) == [1, 3, 4]                                            # a refused save changes nothing
    sold.execute("DELETE FROM branches WHERE branch_id=4")                             # a branch that no longer exists drops out
    assert BM.branch_ids(sold) == [1, 3]


def test_buffer_bounds_default_and_audit(sold):
    assert BM.buffer_pct(sold) == 15.0
    assert BM.set_benchmark(sold, "sales", None, 25)["buffer_pct"] == 25.0 and BM.buffer_pct(sold) == 25.0
    assert BM.set_benchmark(sold, "sales", None, 0)["buffer_pct"] == 0.0
    assert BM.set_benchmark(sold, "sales", None, 50)["buffer_pct"] == 50.0
    for bad in (-1, 50.5, "abc"):
        with pytest.raises(ValueError):
            BM.set_benchmark(sold, "sales", None, bad)
    assert BM.buffer_pct(sold) == 50.0
    assert sold.execute("SELECT old_value, new_value FROM bc_audit WHERE action='benchmark_buffer' ORDER BY id LIMIT 1").fetchone() == ("15", "25")
    sold.execute("UPDATE app_settings SET value='abc' WHERE key='bc_stock_buffer_pct'")
    assert BM.buffer_pct(sold) == 15.0                                                 # a bad stored value reads as the default


# ── 2 · the median ────────────────────────────────────────────────────────────────────────────
def test_median_needs_two_branches_and_ignores_one_odd_branch():
    assert BM.median_of([5.0]) is None and BM.median_of([]) is None and BM.median_of([None, 7.0]) is None
    assert BM.median_of([10.0, 20.0]) == 15.0                                          # two branches: the middle
    assert BM.median_of([10.0, 12.0, 400.0]) == 12.0                                   # three: the odd one does not move it
    assert BM.median_of([10.0, 12.0, 14.0, 400.0]) == 13.0                             # four: the mean of the two in the middle
    assert BM.median_of([1.0, None, 3.0]) == 2.0
    assert BM.vs_pct(115, 100) == 15.0 and BM.vs_pct(80, 100) == -20.0 and BM.vs_pct(5, 0) is None and BM.vs_pct(None, 5) is None


# ── 3 · purchases read with the stock change ────────────────────────────────────────────────────
def _count(db, day, lines, bid):
    return ST.save(db, bid, [{"code": c, "qty": q, "unit": "กรัม"} for c, q in lines.items()], "udon",
                   submit=True, count_date=day, require_all=False)["count_id"]


def test_begin_and_end_counts_switch_a_branch_month_to_stock_adjusted_use(sold):
    bought = {CRAB: 20000.0, SALMON: 5000.0, VEG: 700.0}
    use, adj, per = BM.adjusted_use(sold, 1, "2026-09", bought)
    assert use == bought and adj == set() and per is None                              # no counts: the plain purchase
    _count(sold, "2026-09-01", {CRAB: 4000, SALMON: 0}, 1)
    _count(sold, "2026-10-01", {CRAB: 8500, SALMON: 6000}, 1)
    use, adj, per = BM.adjusted_use(sold, 1, "2026-09", bought)
    assert per["begin"]["date"] == "2026-09-01" and per["end"]["date"] == "2026-10-01"
    assert use[CRAB] == 20000 + 4000 - 8500 and CRAB in adj                           # bought + begin - end
    assert use[SALMON] == 5000 and SALMON not in adj                                  # 5,000 + 0 - 6,000 < 0: a miscount or a missing receipt, the purchase stays
    assert use[VEG] == 700 and VEG not in adj                                          # not counted: the purchase
    use2, adj2, _ = BM.adjusted_use(sold, 1, "2026-09", {})                            # counted at both ends, nothing bought, the shelf filled up: nothing used
    assert use2 == {} and adj2 == set()
    assert BM.adjusted_use(sold, 1, "2026-08", bought)[2] is None                      # the pair is September's, not August's
    assert BM.adjusted_use(sold, 2, "2026-09", bought)[2] is None                      # and branch 1's


def test_a_branch_with_both_counts_uses_the_normal_threshold_the_others_the_buffer():
    assert BM.tolerance("A", True, 15.0) == 5.0 and BM.tolerance("B", True, 15.0) == 10.0 and BM.tolerance("C", True, 15.0) == 20.0
    assert BM.tolerance("A", False, 15.0) == 15.0 and BM.tolerance(None, False, 22.0) == 22.0
    assert BM.tolerance("zzz", True, 15.0) == 20.0                                    # an unknown class reads as C


# ── 4 · the per-฿100,000 report against the benchmark ───────────────────────────────────────────
@pytest.fixture
def ps(sold):
    """Company branches 1 (the fixture's pattern), 3 (both dishes) and 4 (the crab dish alone); franchise 5 (sells the crab dish at
    ฿200) and 6 (same as 3). Crab per ฿100,000: 15,000 g · 15,000 · 45,000 · 11,250 · 15,000; salmon: 6,667 · 6,667 · - · 5,000 · 6,667."""
    for i, n, t in ((3, "Co3", "master"), (4, "Co4", "master"), (5, "Fr5", "franchise"), (6, "Fr6", "franchise")):
        _branch(sold, i, n, t)
    a = date(2026, 9, 1)
    _sell(sold, 3, a, AS_OF)
    _sell(sold, 4, a, AS_OF, q11=0)
    _sell(sold, 5, a, AS_OF, p10=200)
    _sell(sold, 6, a, AS_OF)
    BM.set_benchmark(sold, "t", [1, 3, 4])
    sold.commit()
    return sold


def test_report_carries_the_median_the_comparison_and_the_chip(ps):
    r = PS.report(ps, 4, today=date(2026, 10, 6))
    bm = r["benchmark"]
    assert bm["label"] == "ค่ากลางสาขาของบริษัท" and bm["n"] == 3 and [b["branch_id"] for b in bm["branches"]] == [1, 3, 4]
    bi = {x["code"]: x for x in bm["items"]}
    assert bi[CRAB]["per_100k"] == pytest.approx(15.0)                                 # median of 15,000 · 15,000 · 45,000 g, in packs of 1 kg
    assert bi[SALMON]["per_100k"] == pytest.approx(6.667, abs=1e-3)                    # only the branches that use it: 1 and 3
    assert bm["value_per_100k"] == pytest.approx(15000 * 0.15 + 6666.667 * 0.3714, abs=0.05)
    assert bm["buffer_pct"] == 15.0 and r["buffer_pct"] == 15.0 and "ใช้ชั่วคราว" in r["buffer_note"]
    b = {x["branch_id"]: x for x in r["branches"]}
    assert b[1]["is_benchmark"] and b[3]["is_benchmark"] and b[4]["is_benchmark"] and not b[5]["is_benchmark"]
    assert b[3]["vs_benchmark_pct"] == 0.0 and b[6]["vs_benchmark_pct"] == 0.0
    # the odd company branch is far above the median and is NOT chipped (it is part of it)
    assert b[4]["vs_benchmark_pct"] > 15 and b[4]["below_benchmark"] is False
    # franchise 5 uses 25 % less value per ฿100,000: below the benchmark by more than the buffer
    assert b[5]["vs_benchmark_pct"] == pytest.approx(-25.0, abs=0.1) and b[5]["below_benchmark"] is True
    assert b[5]["below_benchmark_th"] == "ต่ำกว่าค่ากลาง · อาจซื้อจากภายนอก" and b[6]["below_benchmark"] is False
    c5 = {x["code"]: x for x in b[5]["items"]}[CRAB]
    assert c5["bench_per_100k"] == pytest.approx(15.0) and c5["vs_benchmark_pct"] == pytest.approx(-25.0, abs=0.1)
    # a company branch's own item is compared as well
    assert {x["code"]: x for x in b[4]["items"]}[CRAB]["vs_benchmark_pct"] == pytest.approx(200.0)
    # the all-branch total is still there, as its own line
    assert r["chain"] is not None and r["chain"]["branches"] == len(r["branches"])
    # the chip follows the buffer: with 30 % the 25 % gap is inside it
    BM.set_benchmark(ps, "t", None, 30)
    assert {x["branch_id"]: x for x in PS.report(ps, 4, today=date(2026, 10, 6))["branches"]}[5]["below_benchmark"] is False


def test_one_branch_asked_for_still_reads_the_whole_benchmark(ps):
    r = PS.report(ps, 4, [5], today=date(2026, 10, 6))
    assert [x["branch_id"] for x in r["branches"]] == [5] and r["benchmark"]["n"] == 3
    assert r["branches"][0]["below_benchmark"] is True and r["chain"]["branches"] == 1


def test_fewer_than_two_benchmark_branches_with_sales_gives_no_benchmark(ps):
    BM.set_benchmark(ps, "t", [1, 2])                                                  # branch 2 sold nothing
    r = PS.report(ps, 4, today=date(2026, 10, 6))
    assert r["benchmark"] is None
    assert all(b["vs_benchmark_pct"] is None and b["below_benchmark"] is False for b in r["branches"])
    assert all(x["vs_benchmark_pct"] is None for b in r["branches"] for x in b["items"])


def test_report_cache_follows_the_benchmark_setting(ps, monkeypatch):
    monkeypatch.setattr(FC, "USAGE_CACHE_TTL", 600)
    a = PS.report(ps, 4, today=date(2026, 10, 6))
    assert PS.report(ps, 4, today=date(2026, 10, 6)) is a
    BM.set_benchmark(ps, "t", [1, 3])
    b = PS.report(ps, 4, today=date(2026, 10, 6))
    assert b is not a and b["benchmark"]["n"] == 2
    assert {x["code"]: x for x in b["benchmark"]["items"]}[CRAB]["per_100k"] == pytest.approx(15.0)


# ── 5 · a new branch: the benchmark is the default reference ─────────────────────────────────────
@pytest.fixture
def newb(ps):
    """Branch 7 opened five days ago (the crab dish and the salmon dish); the benchmark is branches 1, 3 and 4."""
    _branch(ps, 7, "New7", "franchise")
    _sell(ps, 7, date(2026, 10, 1), AS_OF)
    ps.commit()
    return ps


def test_the_benchmark_is_the_default_reference(newb):
    assert PS.reference_spec(newb, 7) == {"mode": "benchmark", "ref_branch_id": None}
    assert PS.reference_for(newb, 7) is None and PS.resolved_reference(newb, 7) == PS.BENCH
    p = OR.plan(newb, 7, "1", WED, now=NOW)
    rm = p["ref_mode"]
    assert rm["active"] and rm["benchmark"] is True and rm["needs_sales"] and rm["branch_id"] is None
    assert rm["name"] == "ค่ากลางสาขาของบริษัท" and rm["label"] == "คำนวณจากค่ากลางสาขาของบริษัท"
    assert rm["message"] == PS.NEED_SALES_TH and p["rows"] == []                      # it must still say what it expects to sell
    assert "BranchA" not in str(rm) and "Co3" not in str(rm)                           # the branch sees no other branch's name
    with pytest.raises(ValueError) as e:
        OR.save_flow(newb, 7, "1", WED, [], "u", now=NOW)
    assert str(e.value) == PS.NEED_SALES_TH
    newb.rollback()
    q = OR.plan(newb, 7, "1", WED, now=NOW, expected_sales=100000)
    crab = {r["code"]: r for r in q["rows"]}[CRAB]
    assert crab["forecast"] == pytest.approx(15000.0)                                  # the MEDIAN of 15,000 · 15,000 · 45,000 g per ฿100,000 x ฿100,000
    assert crab["ref_label"] == "คำนวณจากค่ากลางสาขาของบริษัท" and q["ref_mode"]["ref_ok"] is True
    assert q["ref_mode"]["ref_sales_normal"] is not None


def test_a_chosen_branch_or_none_still_works_and_a_branch_never_borrows_from_itself(newb):
    assert PS.set_reference(newb, "t", 7, 1) == {"branch_id": 7, "ref_branch_id": 1}
    assert PS.reference_spec(newb, 7) == {"mode": "branch", "ref_branch_id": 1} and PS.resolved_reference(newb, 7) == 1
    p = OR.plan(newb, 7, "1", WED, now=NOW, expected_sales=100000)
    assert p["ref_mode"]["benchmark"] is False and p["ref_mode"]["label"] == "คำนวณจากสาขาอ้างอิง BranchA"
    assert PS.set_reference(newb, "t", 7, None, mode="none") == {"branch_id": 7, "ref_branch_id": None}
    assert PS.reference_spec(newb, 7)["mode"] == "none" and PS.resolved_reference(newb, 7) is None
    assert OR.plan(newb, 7, "1", WED, now=NOW)["ref_mode"] is None                  # opted out: today's behaviour
    assert PS.reference_map(newb) == {} and newb.execute("SELECT value FROM app_settings WHERE key='bc_ref_branch'").fetchone()[0] == '{"7": "none"}'
    PS.set_reference(newb, "t", 7, None)                                              # cleared: back to the default
    assert PS.reference_spec(newb, 7)["mode"] == "benchmark"
    # the benchmark needs two OTHER company branches: branch 3 among {1, 3} has one
    BM.set_benchmark(newb, "t", [1, 3])
    assert PS.resolved_reference(newb, 3) is None and PS.resolved_reference(newb, 7) == PS.BENCH


def test_no_benchmark_configured_leaves_every_branch_as_before(sold):
    _branch(sold, 7, "New7", "franchise")
    _sell(sold, 7, date(2026, 10, 1), AS_OF)
    assert PS.resolved_reference(sold, 7) is None
    p = OR.plan(sold, 7, "1", WED, now=NOW)
    assert p["ref_mode"] is None and {r["code"]: r for r in p["rows"]}[CRAB]["ref_label"] is None


# ── 6 · the recipe check ─────────────────────────────────────────────────────────────────────────
@pytest.fixture
def rc(sold):
    """September 2026 for the benchmark branches 1, 3 and 4 (4 sold on the first 15 days only).
    Bought, against the recipes: crab x1.05 · x1.00 · x1.10 (ok) · salmon x1.4 on 1 and 3 (more) · V-007 x20 (x10) · T-LS x0.5 (less)
    · seaweed bought, in no recipe · T-NB in a recipe, never bought. Branch 4's purchases are those of the whole month."""
    sold.execute("UPDATE bc_ingredient SET loss_pct=0 WHERE code='V-007'")
    for code, nm in (("T-LS", "ของซื้อน้อย"), ("T-NB", "ของไม่ซื้อ")):
        sold.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, loss_pct, active) VALUES(?,?, 'raw', 'กรัม', 0, 1)", (code, nm))
        sold.execute("INSERT INTO bc_ingredient_price(ingredient_code, month, price_per_base, source) VALUES(?, '2026-07', 0.05, 'list_ref')", (code,))
    sold.executemany("INSERT INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name, qty_per_serving, unit, price_per_unit,"
                     " ingredient_cost, ingredient_type) VALUES(?,?,?,?,?,?,?,?,?)", [
                         (10, 2, VEG, "หัวไชเท้า", 10, "กรัม", 0.04, 0.4, "veg"), (10, 3, "T-LS", "ของซื้อน้อย", 10, "กรัม", 0.05, 0.5, "x"),
                         (11, 2, "T-NB", "ของไม่ซื้อ", 5, "กรัม", 0.05, 0.25, "x")])
    for i, n, t in ((3, "Co3", "master"), (4, "Co4", "master"), (5, "Fr5", "franchise")):
        _branch(sold, i, n, t)
    _sell(sold, 3, date(2026, 9, 1), AS_OF)
    _sell(sold, 4, date(2026, 9, 1), date(2026, 9, 15), q11=0)
    BM.set_benchmark(sold, "t", [1, 3, 4])
    sold.commit()
    use = {b: FC.usage_history(sold, b, date(2026, 9, 1), date(2026, 9, 30), "2026-09")[0] for b in (1, 3, 4)}
    tot = lambda b, c: sum(use[b].get(c, {}).values())          # noqa: E731
    rows = []
    for b, f in ((1, 1.05), (3, 1.0), (4, 1.10 * 2)):          # branch 4 sold 15 of 30 days: its month's purchases are twice what its 15 days used
        rows.append((CRAB, b, tot(b, CRAB) * f))
    for b in (1, 3):
        rows.append((SALMON, b, tot(b, SALMON) * 1.4))
    for b, f in ((1, 20), (3, 20), (4, 20 * 2)):
        rows.append((VEG, b, tot(b, VEG) * f))
    for b, f in ((1, 0.5), (3, 0.5), (4, 0.5 * 2)):
        rows.append(("T-LS", b, tot(b, "T-LS") * f))
    rows += [(SEAWEED, 1, 100.0), (SEAWEED, 3, 100.0)]
    sold.executemany("INSERT INTO bc_ingredient_price_branch(ingredient_code, month, branch_id, price_per_base, qty_bought, value_bought)"
                     " VALUES(?, '2026-09', ?, 1.0, ?, ?)", [(c, b, q, q) for c, b, q in rows])
    sold.commit()
    return sold


def _by(r):
    return {x["code"]: x for x in r["rows"]}


def test_statuses_on_a_fixture_with_one_of_each(rc):
    r = RC.check(rc, "2026-09", TODAY)
    assert r["month"] == "2026-09" and r["enough"] and [b["branch_id"] for b in r["branches"]] == [1, 3, 4]
    by = _by(r)
    assert by[CRAB]["status"] == "ok" and by[CRAB]["diff_pct"] == pytest.approx(5.0, abs=0.1)         # median of x1.05 · x1.00 · x1.10
    assert by[SALMON]["status"] == "more" and by[SALMON]["diff_pct"] == pytest.approx(40.0, abs=0.1) and by[SALMON]["n_branches"] == 2
    assert by["T-LS"]["status"] == "less" and by["T-LS"]["diff_pct"] == pytest.approx(-50.0, abs=0.1)
    assert by[VEG]["status"] == "x10" and by[VEG]["bought"] / by[VEG]["recipe"] == pytest.approx(20.0, rel=1e-3)
    assert by[SEAWEED]["status"] == "no_recipe" and by[SEAWEED]["recipe"] == 0
    assert by["T-NB"]["status"] == "no_buy" and by["T-NB"]["bought"] == 0
    assert r["counts"] == {"x10": 1, "no_recipe": 1, "no_buy": 1, "more": 1, "less": 1, "ok": 1}
    assert by[CRAB]["status_th"] == "ในเกณฑ์ (±15%)" and by[VEG]["status_th"] == "ต่างกันเกิน 10 เท่า · ตรวจหน่วย"
    assert by[SALMON]["status_th"] == "ซื้อมากกว่าสูตร" and by["T-LS"]["status_th"] == "ซื้อน้อยกว่าสูตร"
    assert by[SEAWEED]["status_th"] == "มีการซื้อ แต่ไม่มีในสูตร" and by["T-NB"]["status_th"] == "มีในสูตร แต่ไม่มีการซื้อ · ตรวจรหัส"
    # per ฿100,000 of sales: crab 15,000 g in the recipe (45 g per ฿300 of two dishes)
    assert by[CRAB]["recipe"] == pytest.approx(15000.0, rel=1e-3) and by[CRAB]["bought"] == pytest.approx(15750.0, rel=1e-3)
    # sorted by value impact, the biggest first; every row has its branch columns
    imp = [x["impact_value"] or 0 for x in r["rows"]]
    assert imp == sorted(imp, reverse=True) and r["rows"][0]["code"] == VEG
    assert [b["branch_id"] for b in by[CRAB]["branches"]] == [1, 3, 4] and by[CRAB]["branches"][2]["recipe"] == pytest.approx(45000.0, rel=1e-3)
    assert by[SALMON]["branches"][2]["bought"] is None                                             # branch 4 has no salmon at all


def test_purchases_are_scaled_by_the_days_that_have_daily_sales(rc):
    r = RC.check(rc, "2026-09", TODAY)
    b = {x["branch_id"]: x for x in r["branches"]}
    assert b[1]["scale"] == 1.0 and b[1]["scaled"] is False and b[3]["sales_days"] == 30
    assert b[4]["sales_days"] == 15 and b[4]["month_days"] == 30 and b[4]["scale"] == 0.5 and b[4]["scaled"] is True
    assert r["scaled"] is True and r["partial_month"] is False
    # branch 4 bought twice what its fifteen days used (x1.10 of it): scaled it reads +10 %, not +120 %
    c4 = _by(r)[CRAB]["branches"][2]
    assert c4["diff_pct"] == pytest.approx(10.0, abs=0.1)
    assert r["from"] == "2026-09-01" and r["to"] == "2026-09-30"


def test_sales_are_ex_vat(rc):
    base = RC.check(rc, "2026-09", TODAY)
    rc.execute("ALTER TABLE sales_sku_daily ADD COLUMN total_inc_vat REAL")
    rc.execute("UPDATE sales_sku_daily SET total_inc_vat = subtotal_ex_vat * 1.07")
    rc.commit()
    RC.clear_cache()
    again = RC.check(rc, "2026-09", TODAY)
    assert again["branches"][0]["sales"] == base["branches"][0]["sales"]
    assert _by(again)[CRAB]["recipe"] == pytest.approx(_by(base)[CRAB]["recipe"])


def test_a_branch_without_purchases_or_sales_is_left_out_and_two_are_needed(rc):
    rc.execute("DELETE FROM bc_ingredient_price_branch WHERE branch_id=4")
    RC.clear_cache()
    r = RC.check(rc, "2026-09", TODAY)
    assert [b["branch_id"] for b in r["branches"]] == [1, 3] and r["skipped_branches"][0]["branch_id"] == 4 and r["enough"]
    assert "ไม่มียอดซื้อรายสาขา" in r["skipped_branches"][0]["why"]
    rc.execute("DELETE FROM bc_ingredient_price_branch WHERE branch_id=3")
    RC.clear_cache()
    one = RC.check(rc, "2026-09", TODAY)
    assert one["enough"] is False and "อย่างน้อย 2 สาขา" in one["message"]
    BM.set_benchmark(rc, "t", [1, 5])                                                              # branch 5 has no purchases or sales
    assert RC.check(rc, "2026-09", TODAY)["enough"] is False


def test_months_and_the_default_month(rc):
    rc.execute("INSERT INTO bc_ingredient_price_branch VALUES('x', '2026-10', 1, 1, 1, 1)")
    rc.execute("INSERT INTO bc_ingredient_price_branch VALUES('y', '2026-08', 1, 1, 1, 1)")
    rc.execute("INSERT INTO bc_ingredient_price_branch VALUES('y', '2026-08', 3, 1, 1, 1)")
    rc.execute("INSERT INTO bc_ingredient_price_branch VALUES('z', '2026-07', 5, 1, 1, 1)")      # branch 5 is not a benchmark branch
    assert RC.months(rc) == [{"month": "2026-10", "branches": 1, "complete": False}, {"month": "2026-09", "branches": 3, "complete": True},
                             {"month": "2026-08", "branches": 2, "complete": True}]
    assert RC.default_month(rc) == "2026-09"                                                       # the newest finished month with two branches
    assert RC.check(rc, None, TODAY)["month"] == "2026-09"


def test_stock_counts_move_a_branch_to_the_normal_threshold(rc):
    use = FC.usage_history(rc, 3, date(2026, 9, 1), date(2026, 9, 30), "2026-09")[0]
    crab = sum(use[CRAB].values())
    rc.execute("UPDATE bc_ingredient_price_branch SET qty_bought=?, value_bought=? WHERE branch_id=3 AND ingredient_code=?",
               (crab + 4500, crab + 4500, CRAB))
    rc.commit()
    RC.clear_cache()
    before = _by(RC.check(rc, "2026-09", TODAY))[CRAB]["branches"][1]
    assert before["adjusted"] is False and before["diff_pct"] > 15
    _count(rc, "2026-09-01", {CRAB: 4000}, 3)
    _count(rc, "2026-10-01", {CRAB: 8500}, 3)
    RC.clear_cache()
    r = RC.check(rc, "2026-09", TODAY)
    after = _by(r)[CRAB]["branches"][1]
    assert after["adjusted"] is True and after["diff_pct"] == pytest.approx(0.0, abs=0.1)          # bought + 4,000 - 8,500 = what the recipes used
    b3 = {x["branch_id"]: x for x in r["branches"]}[3]
    assert b3["stock_adjusted"] is True and b3["adjusted_items"] == 1 and b3["begin_date"] == "2026-09-01" and b3["end_date"] == "2026-10-01"
    assert r["stock_adjusted_branches"] == 1
    assert _by(r)[CRAB]["stock_adjusted"] is False and _by(r)[CRAB]["tolerance_pct"] == 15.0         # the other two are still read with the buffer


def test_the_buffer_decides_the_band(rc):
    assert _by(RC.check(rc, "2026-09", TODAY))[CRAB]["status"] == "ok"
    BM.set_benchmark(rc, "t", None, 3)                                                             # +5 % is outside a 3 % band
    assert _by(RC.check(rc, "2026-09", TODAY))[CRAB]["status"] == "more"


def test_the_recipe_check_cache_is_dropped_with_the_usage_cache(rc, monkeypatch):
    monkeypatch.setattr(FC, "USAGE_CACHE_TTL", 600)
    a = RC.check(rc, "2026-09", TODAY)
    assert RC.check(rc, "2026-09", TODAY) is a
    FC.clear_usage_cache()
    assert RC.check(rc, "2026-09", TODAY) is not a


# ── 7 · หน้างาน ─────────────────────────────────────────────────────────────────────────────────────
def test_todo_lists_the_ten_times_rows_for_the_analyst_only(rc):
    out = TD.build(rc, today=TODAY)
    it = next(i for i in out["items"] if i["key"] == "recipe_check_units")
    assert it["title"] == "ตรวจหน่วยของวัตถุดิบที่ยอดซื้อต่างจากสูตรเกิน 10 เท่า (1 รายการ)" and it["count"] == 1 and it["status"] == "todo"
    assert "หัวไชเท้า" in it["detail"] and "กันยายน" in it["detail"] and "กรุณา" not in it["detail"]
    assert it["owner_modules"] == ["bom.center"] and it["owner_th"] == "ฝ่ายวิเคราะห์"
    assert it["go"] == {"tab": "bomcenter", "pane": "rcheck", "month": "2026-09"}
    mine = TD.build(rc, today=TODAY, scoped_branch="BranchA", branch_id=1)
    assert "recipe_check_units" not in {i["key"] for i in mine["items"]}                            # a branch login gets its own list only
    rc.execute("UPDATE bc_ingredient_price_branch SET qty_bought = qty_bought / 20.0 WHERE ingredient_code=?", (VEG,))
    rc.execute("UPDATE bc_ingredient_price_branch SET qty_bought = qty_bought / 2.0 WHERE ingredient_code=? AND branch_id=4", (VEG,))
    rc.commit()
    RC.clear_cache()
    assert "recipe_check_units" not in {i["key"] for i in TD.build(rc, today=TODAY)["items"]}       # fixed: the item goes away


# ── 8 · HTTP: who reaches what ─────────────────────────────────────────────────────────────────────
def _seed_http(db_path):
    """Branches 1 and 2 as in the app database; 3 and 4 company branches. ฿1,000 and 10 plates a day (branch 1, 3) or 20 (2, 4), 1 Sep - 4 Oct;
    the dish holds 0.5 g of T-BM1; September's purchases are exactly what the recipes used."""
    c = sqlite3.connect(str(db_path))
    c.execute("INSERT OR REPLACE INTO branches(branch_id, name, branch_type, active) VALUES(3, 'BranchC', 'master', 1)")
    c.execute("INSERT OR REPLACE INTO branches(branch_id, name, branch_type, active) VALUES(4, 'BranchD', 'master', 1)")
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, family, active) VALUES('T-BM1', 'ของทดสอบ ค่ากลาง', 'raw', 'กรัม', 'F', 1)")
    c.execute("INSERT OR REPLACE INTO bc_ingredient_price(ingredient_code, month, price_per_base, source) VALUES('T-BM1', '2026-09', 2.0, 'purchase_avg')")
    c.execute("INSERT INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name, qty_per_serving, unit, price_per_unit,"
              " ingredient_cost, ingredient_type) VALUES(2, 1, 'T-BM1', 'ของทดสอบ ค่ากลาง', 0.5, 'กรัม', 0, 0, 'x')")
    for d in _days(date(2026, 9, 1), date(2026, 10, 4)):
        for b, q in ((1, 10), (2, 20), (3, 10), (4, 20)):
            c.execute("INSERT INTO sales_sku_daily VALUES(?, ?, 2, ?, 1000)", (b, d.isoformat(), q))
    for b, q in ((1, 150.0), (3, 150.0), (4, 300.0)):
        c.execute("INSERT INTO bc_ingredient_price_branch VALUES('T-BM1', '2026-09', ?, 2.0, ?, ?)", (b, q, q * 2))
    c.commit()
    c.close()


def test_http_roles_pages_workbook_and_pinning(client, app_db, auth_on, monkeypatch):  # noqa: F811
    import main
    from routers.bom_center import RECIPE_CHECK_HEADS
    monkeypatch.setattr(PS, "bkk_today", lambda: date(2026, 10, 12))
    client.get("/api/bom-center/hub/queue")
    snap = _snap(main.DB_PATH)
    _seed_http(main.DB_PATH)
    try:
        mg, an, ss, acc = _login("management1"), _login("analyst1"), _login("salessupport1"), _login("accounting1")
        mg.post("/api/auth/users", headers=TOKEN, json={"username": "br_bm_test", "password": "Temp@12345", "role": "branch", "branch": "BranchA"})
        br = _login("br_bm_test", "Temp@12345")
        # the settings card: the sales team and the analyst read and write; accounting reads nothing; a branch login reaches neither
        url = "/api/bom-center/benchmark"
        s = ss.get(url).json()
        assert s["buffer_pct"] == 15.0 and s["is_default"] is True and s["min_branches"] == 2 and {b["name"] for b in s["branches"]} >= {"BranchA", "BranchC"}
        assert an.get(url).status_code == 200 and mg.get(url).status_code == 200
        assert acc.get(url).status_code == 403 and br.get(url).status_code == 403
        assert br.post(url, headers=TOKEN, json={"branch_ids": [1, 3]}).status_code == 403
        assert acc.post(url, headers=TOKEN, json={"branch_ids": [1, 3]}).status_code == 403
        assert ss.post(url, headers=TOKEN, json={"branch_ids": [1]}).status_code == 400
        assert ss.post(url, headers=TOKEN, json={"branch_ids": [1, 99]}).status_code == 400
        assert ss.post(url, headers=TOKEN, json={"buffer_pct": 60}).status_code == 400
        out = ss.post(url, headers=TOKEN, json={"branch_ids": [1, 3, 4], "buffer_pct": 20}).json()
        assert out["selected"] == [1, 3, 4] and out["buffer_pct"] == 20.0
        cal = ss.get("/api/bom-center/order-calendar").json()
        assert cal["benchmark"]["names"] == ["BranchA", "BranchC", "BranchD"] and all(r["ref_mode"] == "benchmark" for r in cal["rows"])
        ref = "/api/bom-center/order-calendar/ref-branch"
        assert ss.post(ref, headers=TOKEN, json={"branch_id": 2, "ref_mode": "none"}).status_code == 200
        assert {r["branch_id"]: r["ref_mode"] for r in ss.get("/api/bom-center/order-calendar").json()["rows"]}[2] == "none"
        assert ss.post(ref, headers=TOKEN, json={"branch_id": 2, "ref_mode": "sideways"}).status_code == 400
        assert ss.post(ref, headers=TOKEN, json={"branch_id": 2, "ref_branch_id": 1}).json() == {"branch_id": 2, "ref_branch_id": 1}
        assert ss.post(ref, headers=TOKEN, json={"branch_id": 2}).status_code == 200
        # the recipe check: the analyst and management; nobody else
        for who in (an, mg):
            j = who.get("/api/bom-center/recipe-check").json()
            assert j["month"] == "2026-09" and j["enough"] and j["buffer_pct"] == 20.0
            row = {x["code"]: x for x in j["rows"]}["T-BM1"]
            assert row["status"] == "ok" and row["bought"] == pytest.approx(500.0) and row["recipe"] == pytest.approx(500.0)
        assert an.get("/api/bom-center/recipe-check?month=2026-09").status_code == 200
        assert an.get("/api/bom-center/recipe-check?month=garbage").status_code == 400
        for who in (ss, acc, br):
            assert who.get("/api/bom-center/recipe-check").status_code == 403, who
            assert who.get("/api/bom-center/recipe-check.xlsx").status_code == 403
        r = an.get("/api/bom-center/recipe-check.xlsx?month=2026-09")
        assert r.status_code == 200 and "bom_center_recipe_check_2026-09.xlsx" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["ตรวจสูตรด้วยยอดซื้อ"] == [RC.SHEET]
        rows = list(wb[RC.SHEET].iter_rows(values_only=True))
        heads = list(rows[1])
        assert heads[:9] == RECIPE_CHECK_HEADS == ["รหัส", "วัตถุดิบ", "หน่วยฐาน", "ซื้อจริงต่อยอดขาย ฿100,000 (ค่ากลาง)", "ตามสูตรต่อยอดขาย ฿100,000 (ค่ากลาง)",
                                                   "ส่วนต่าง %", "ช่วงยอมรับ ±%", "ผลตรวจ", "มูลค่าส่วนต่างต่อยอดขาย ฿100,000"]
        assert heads[9:12] == ["BranchA · ซื้อจริง", "BranchA · ตามสูตร", "BranchA · ส่วนต่าง %"] and heads[-1] == "หมายเหตุ"
        mine = next(x for x in rows[2:] if x[0] == "T-BM1")
        assert mine[3] == pytest.approx(500.0) and mine[4] == pytest.approx(500.0) and mine[7] == "ในเกณฑ์ (±20%)"
        # the per-฿100,000 report: the benchmark for the centre, nothing of it for a branch login
        j = mg.get("/api/bom-center/per-sales?weeks=2").json()
        assert j["benchmark"]["n"] == 3 and {x["code"]: x for x in j["benchmark"]["items"]}["T-BM1"]["per_100k"] == pytest.approx(500.0)
        assert {b["branch"]: b for b in j["branches"]}["BranchB"]["is_benchmark"] is False
        mine_ps = br.get("/api/bom-center/per-sales?weeks=2").json()
        assert mine_ps["benchmark"] is None and mine_ps["chain"] is None and [b["branch"] for b in mine_ps["branches"]] == ["BranchA"]
        assert mine_ps["branches"][0]["vs_benchmark_pct"] is None and mine_ps["branches"][0]["is_benchmark"] is None
        assert all(i["vs_benchmark_pct"] is None and i["bench_per_100k"] is None for i in mine_ps["branches"][0]["items"])
        assert "benchmark" in str(ss.get("/api/bom-center/per-sales?weeks=2").json().keys())
        x = load_workbook(io.BytesIO(mg.get("/api/bom-center/per-sales.xlsx").content))
        names = {r_[0] for r_ in x["ต่อยอดขาย 100000"].iter_rows(min_row=3, values_only=True)}
        assert "ค่ากลางสาขาของบริษัท" in names and "ทั้งเครือ" in names
        xb = load_workbook(io.BytesIO(br.get("/api/bom-center/per-sales.xlsx").content))
        assert {r_[0] for r_ in xb["ต่อยอดขาย 100000"].iter_rows(min_row=3, values_only=True)} == {"BranchA"}
        # หน้างาน: nothing for a branch login
        assert "recipe_check_units" not in {i["key"] for i in br.get("/api/bom-center/todo").json()["items"]}
    finally:
        c = sqlite3.connect(str(main.DB_PATH))
        c.execute("DELETE FROM sales_sku_daily WHERE day >= '2026-09-01'")
        c.execute("DELETE FROM bom_ingredients WHERE ingredient_code='T-BM1'")
        c.execute("DELETE FROM bc_ingredient_price_branch WHERE ingredient_code='T-BM1'")
        c.execute("DELETE FROM branches WHERE branch_id IN (3, 4)")                                # the session database is shared: leave it as it was
        c.execute("DELETE FROM app_user WHERE username='br_bm_test'")
        c.execute("DELETE FROM app_settings WHERE key IN ('bc_ref_branch', 'bc_benchmark_branches', 'bc_stock_buffer_pct')")
        c.commit()
        c.close()
        _restore(main.DB_PATH, snap)


# ── 9 · who may read and set ──────────────────────────────────────────────────────────────────────────
def test_auth_matrix():
    R = auth.ROLE_MODULES
    check = "/api/bom-center/recipe-check"
    for role, mods in R.items():
        allowed = auth.path_allowed("GET", check, mods)
        assert allowed == (role in ("analyst", "management")), role
        assert auth.path_allowed("GET", check + ".xlsx", mods) == allowed
    cfg = "/api/bom-center/benchmark"
    for role in ("salessupport", "analyst", "management"):
        assert auth.path_allowed("GET", cfg, R[role]) and auth.path_allowed("POST", cfg, R[role]), role
    for role in ("branch", "franchise", "warehouse", "accounting", "innovation"):
        assert not auth.path_allowed("POST", cfg, R[role]), role
        assert not auth.path_allowed("GET", cfg, R[role]), role
    assert auth.path_allowed("GET", cfg, R["operation"])                                           # Operations reads the settings it is measured by
