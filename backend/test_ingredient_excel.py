# -*- coding: utf-8 -*-
"""The ordering tab's Excel round-trips and the weekly ingredient-spend tab.

Three pairs, one shape: a GET that hands out a sheet already carrying the keys,
a POST that reads the filled sheet back. What is pinned here is the CONTRACT —
that a template can be filled and re-imported, that a file from another system
(renamed headers, a title row, text-formatted numbers) still lands, and that a
row which cannot be read is reported rather than silently dropped.
"""
import io
import sqlite3

import openpyxl
import pytest

import main

# The shared fixture carries no recipes, so the pack-size sheet would be empty.
# Same three lines test_ingredients.py uses — that module wipes the ordering
# tables in its own setup, and it sorts after this one, so nothing leaks either way.
BOM = [
    (2, 1, "ING-SAL",  "แซลมอน",      0.12, "กก.", 500.0, 60.0, "protein"),
    (2, 2, "ING-RICE", "ข้าวญี่ปุ่น", 0.05, "กก.",  40.0,  2.0, "carb"),
    (3, 1, "ING-BEEF", "เนื้อ",        0.20, "กก.", 600.0, 120.0, "protein"),
]


@pytest.fixture(scope="module", autouse=True)
def _seed_bom(client):
    con = sqlite3.connect(str(main.DB_PATH))
    con.executemany(
        "INSERT OR REPLACE INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name,"
        " qty_per_serving, unit, price_per_unit, ingredient_cost, ingredient_type, source_file)"
        " VALUES(?,?,?,?,?,?,?,?,?,'test')", BOM)
    con.commit()
    con.close()
    try:
        main.invalidate_cache()
    except Exception:
        pass
    yield
    con = sqlite3.connect(str(main.DB_PATH))
    for t in ("ingredient_order_unit", "ingredient_purchase_week"):
        try:
            con.execute(f"DELETE FROM {t}")
        except sqlite3.OperationalError:
            pass
    con.commit()
    con.close()
    try:
        main.invalidate_cache()
    except Exception:
        pass


def _xlsx(rows):
    """rows = list of lists → an .xlsx in memory, ready to POST."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _post(client, url, buf, name="upload.xlsx", **params):
    return client.post(url, params=params,
                       files={"file": (name, buf.getvalue(),
                                       "application/vnd.openxmlformats-officedocument."
                                       "spreadsheetml.sheet")})


def _sheet(resp):
    assert resp.status_code == 200, resp.text[:300]
    assert "spreadsheetml" in resp.headers["content-type"]
    return openpyxl.load_workbook(io.BytesIO(resp.content)).active


# ── window: an explicit range instead of "N days back from the last sale" ────
def test_order_plan_accepts_an_explicit_date_range(client):
    j = client.get("/api/ingredients/order-plan",
                   params={"date_from": "2026-06-01", "date_to": "2026-06-03"}).json()
    if j.get("available") is False:
        pytest.skip("fixture has no BoM-mapped daily sales")
    m = j["meta"]
    assert m["date_from"] == "2026-06-01" and m["date_to"] == "2026-06-03"
    assert m["days"] == 3, "days must be the measured span of the range, not the default 28"
    assert m["window_mode"] == "range"


def test_order_plan_rolling_window_still_default(client):
    j = client.get("/api/ingredients/order-plan", params={"days": 7}).json()
    if j.get("available") is False:
        pytest.skip("fixture has no BoM-mapped daily sales")
    assert j["meta"]["days"] == 7 and j["meta"]["window_mode"] == "rolling"


def test_order_plan_flags_a_window_that_straddles_the_menu_change(client):
    """A window spanning 2026-09-01 averages two different recipe books. The
    plan still computes — the caller is told, and told where to start instead."""
    j = client.get("/api/ingredients/order-plan",
                   params={"date_from": "2026-08-20", "date_to": "2026-09-05"}).json()
    if j.get("available") is False:
        pytest.skip("fixture has no BoM-mapped daily sales")
    era = j["meta"]["era"]
    if j["meta"]["date_from"] >= "2026-09-01" or j["meta"]["date_to"] < "2026-09-01":
        pytest.skip("fixture coverage does not reach across the changeover")
    assert era and era["straddles"] is True
    assert era["suggest_from"] == "2026-09-01"


# ── ตั้งค่าการสั่งซื้อ: template → fill → import ────────────────────────────
def test_order_units_template_round_trips(client):
    ws = _sheet(client.get("/api/ingredients/order-units.template.xlsx"))
    head = [c.value for c in ws[3]]
    assert head[:5] == ["รหัส", "วัตถุดิบ", "หน่วยนับ", "หน่วยสั่ง", "ขนาดบรรจุ"]
    body = [[c.value for c in r] for r in ws.iter_rows(min_row=4)]
    if not body:
        pytest.skip("fixture has no bom_ingredients rows")
    code = body[0][0]
    r = _post(client, "/api/ingredients/order-units/import",
              _xlsx([head, [code, "x", body[0][2], "ลัง", 12, 240, 1, 1, "จากไฟล์", ""]]))
    assert r.status_code == 200, r.text[:300]
    d = r.json()
    assert d["saved"] == 1 and not d["errors"]
    back = {x["code"]: x for x in client.get("/api/ingredients/order-units").json()["rows"]}
    assert back[code]["order_unit"] == "ลัง" and back[code]["qty_per_order"] == 12


def test_order_units_import_reports_bad_rows_instead_of_dropping_them(client):
    ws = _sheet(client.get("/api/ingredients/order-units.template.xlsx"))
    head = [c.value for c in ws[3]]
    body = [[c.value for c in r] for r in ws.iter_rows(min_row=4)]
    if not body:
        pytest.skip("fixture has no bom_ingredients rows")
    code = body[0][0]
    d = _post(client, "/api/ingredients/order-units/import", _xlsx([
        head,
        [code, "", "กก.", "ลัง", 0, None, 1, 1, "", ""],            # pack size 0 → divides by zero
        ["NO_SUCH_CODE", "", "กก.", "ลัง", 5, None, 1, 1, "", ""],   # unknown ingredient
    ]), dry_run=True).json()
    assert d["would_save"] == 0
    why = {e["code"]: e["why"] for e in d["errors"]}
    assert "0" in why[code] or "มากกว่า" in why[code]
    assert "NO_SUCH_CODE" in why


def test_order_units_import_reads_a_file_from_another_system(client):
    """Renamed headers, a title row above them, and numbers stored as text."""
    rows = [c.value for r in _sheet(client.get("/api/ingredients/order-units.template.xlsx"))
            .iter_rows(min_row=4, max_row=4) for c in r]
    if not rows or not rows[0]:
        pytest.skip("fixture has no bom_ingredients rows")
    code = rows[0]
    d = _post(client, "/api/ingredients/order-units/import", _xlsx([
        ["รายงานขนาดบรรจุ ส่งออกเมื่อ 2026-09-10"],       # a title row, not the header
        [],
        ["Code", "Ingredient Name", "Unit", "Order Unit", "Pack Size", "Pack Price"],
        [code, "ชื่อจากอีกระบบ", "กก.", "ถุง", " 2,500.5 ", "฿1,200"],
    ]), name="from_erp.xlsx").json()
    assert d["saved"] == 1, d
    back = {x["code"]: x for x in client.get("/api/ingredients/order-units").json()["rows"]}
    assert back[code]["qty_per_order"] == 2500.5
    assert back[code]["price_per_order"] == 1200


def test_import_refuses_a_file_with_no_recognisable_header(client):
    r = _post(client, "/api/ingredients/order-units/import",
              _xlsx([["ก", "ข", "ค"], [1, 2, 3]]))
    assert r.status_code == 400
    assert "หัวตาราง" in r.json()["detail"]


# ── ยอดซื้อวัตถุดิบรายสัปดาห์ vs ยอดขาย ─────────────────────────────────────
def test_purchases_template_lists_every_branch_and_week(client):
    ws = _sheet(client.get("/api/ingredients/purchases.template.xlsx", params={"weeks": 3}))
    head = [c.value for c in ws[3]]
    # column 2 is the sales team's branch code (2026-09-21); the week columns follow it
    assert head[:6] == ["สาขา", "รหัสสาขา POS", "สัปดาห์เริ่ม", "ถึงวันที่", "สัปดาห์ที่",
                        "ยอดซื้อวัตถุดิบ (บาท)"]
    W, SUN, LABEL = head.index("สัปดาห์เริ่ม"), head.index("ถึงวันที่"), head.index("สัปดาห์ที่")
    body = [[c.value for c in r] for r in ws.iter_rows(min_row=4)]
    assert {r[0] for r in body} == {"BranchA", "BranchB"}
    assert len({r[W] for r in body}) == 3                     # three weeks offered
    for r in body:
        import datetime as _dt
        d = _dt.date.fromisoformat(str(r[W])[:10])
        assert d.weekday() == 0, "every week must be keyed on its Monday"
        # the other system is queried by range — each row spells out its Mon→Sun + ISO week
        assert _dt.date.fromisoformat(r[SUN]) == d + _dt.timedelta(days=6)
        assert r[LABEL].startswith(f"W{d.isocalendar()[1]:02d}/")


def test_purchases_template_round_trips_with_its_extra_columns(client):
    """The added 'ถึงวันที่'/'สัปดาห์ที่' columns must not steal the week or amount."""
    from routers.ingredients import _week_span
    assert _week_span("2026-09-07") == ("2026-09-13", "W37/2026 · 07/09–13/09/2026")
    wk = "2026-06-01"
    try:
        r = _post(client, "/api/ingredients/purchases/import", _xlsx([
            ["สาขา", "สัปดาห์เริ่ม", "ถึงวันที่", "สัปดาห์ที่", "ยอดซื้อวัตถุดิบ (บาท)", "ยอดขาย (ระบบ)"],
            ["BranchA", wk, "2026-06-07", "W23/2026 · 01/06–07/06/2026", 4321, 99999],
        ]))
        assert r.status_code == 200 and r.json()["saved"] == 1, r.text[:300]
        j = client.get("/api/ingredients/purchases",
                       params={"date_from": wk, "date_to": "2026-06-07"}).json()
        got = [x for x in j["rows"] if x["branch"] == "BranchA" and x["amount"] is not None]
        assert got and got[0]["week_start"] == wk and got[0]["amount"] == 4321
    finally:
        client.request("DELETE", "/api/ingredients/purchases",
                       params={"branch": "BranchA", "week_start": wk})


def test_purchases_import_then_compare_against_sales(client):
    d = client.get("/api/ingredients/purchases").json()
    weeks = d["meta"]["weeks"]
    if not weeks:
        pytest.skip("fixture has no weekly sales")
    wk = weeks[-1]
    try:
        r = _post(client, "/api/ingredients/purchases/import", _xlsx([
            ["สาขา", "สัปดาห์เริ่ม", "ยอดซื้อวัตถุดิบ (บาท)"],
            ["BranchA", wk, 12000],
            ["BranchB", wk, 6000],
        ]))
        assert r.status_code == 200, r.text[:300]
        assert r.json()["saved"] == 2 and not r.json()["errors"]
        j = client.get("/api/ingredients/purchases").json()
        cell = {(x["branch"], x["week_start"]): x for x in j["rows"]}
        a = cell[("BranchA", wk)]
        assert a["amount"] == 12000
        assert a["revenue"] > 0
        assert a["pct"] == round(12000 / a["revenue"] * 100, 1)
        assert j["totals"]["amount"] == 18000
    finally:
        for b in ("BranchA", "BranchB"):
            client.request("DELETE", "/api/ingredients/purchases",
                           params={"branch": b, "week_start": wk})


def test_purchases_import_snaps_any_day_to_its_monday(client):
    """The same week filed as a Wednesday and as a Sunday must be ONE row."""
    try:
        r = _post(client, "/api/ingredients/purchases/import", _xlsx([
            ["สาขา", "วันที่", "ยอดซื้อ"],
            ["BranchA", "2026-06-03", 1000],      # Wednesday
            ["BranchA", "2026-06-07", 2500],      # Sunday of the same week — overwrites
        ]))
        assert r.status_code == 200, r.text[:300]
        assert r.json()["saved"] == 2
        j = client.get("/api/ingredients/purchases",
                       params={"date_from": "2026-06-01", "date_to": "2026-06-07"}).json()
        got = [x for x in j["rows"] if x["branch"] == "BranchA" and x["amount"] is not None]
        assert len(got) == 1 and got[0]["week_start"] == "2026-06-01"
        assert got[0]["amount"] == 2500
    finally:
        client.request("DELETE", "/api/ingredients/purchases",
                       params={"branch": "BranchA", "week_start": "2026-06-01"})


def test_purchases_import_reports_an_unknown_branch(client):
    d = _post(client, "/api/ingredients/purchases/import", _xlsx([
        ["สาขา", "วันที่", "ยอดซื้อ"],
        ["ไม่มีสาขานี้", "2026-06-03", 1000],
        ["BranchA", "ไม่ใช่วันที่", 500],
    ]), dry_run=True).json()
    assert d["would_save"] == 0
    whys = " ".join(e["why"] for e in d["errors"])
    assert "ไม่พบสาขานี้" in whys and "วันที่" in whys


def test_purchases_shows_a_sold_week_with_no_import_as_a_gap(client):
    """A week that was sold but never imported is the finding — it must not
    quietly disappear from the comparison."""
    j = client.get("/api/ingredients/purchases").json()
    if not j["rows"]:
        pytest.skip("fixture has no weekly sales")
    assert any(r["amount"] is None for r in j["rows"])
    assert j["meta"]["missing_weeks"] >= 1
    for r in j["rows"]:
        if r["amount"] is None:
            assert r["pct"] is None, "no import means no ratio, not a zero"


# ── ปริมาณสั่งจริง: template needs an issued cycle ──────────────────────────
def test_warning_letter_draft_is_visibly_a_draft(client):
    """A warning notice names a branch and asks it to answer for a shortfall, so
    the preview must never be mistakable for the issued letter."""
    import conftest as _cf
    ps = f"{_cf.LATEST_MONTH}-01"
    r = client.post("/api/ingredients/order-plan/issue",
                    json={"period_start": ps, "horizon": "week",
                          "date_to": f"{_cf.LATEST_MONTH}-03", "days": 28})
    if r.status_code != 200 or not r.json().get("issued"):
        pytest.skip("no plan rows to issue on this fixture")
    p = {"period_start": ps, "horizon": "week"}
    sent = client.get("/api/ingredients/compliance-letter.print", params=p)
    draft = client.get("/api/ingredients/compliance-letter.print",
                       params={**p, "preview": "true"})
    if sent.status_code == 404:
        pytest.skip("every branch is compliant — no notice is produced")
    assert sent.status_code == 200 and draft.status_code == 200
    assert "ร่าง" in draft.text and "DRAFT" in draft.text
    assert "ร่าง" not in sent.text, "the real letter must carry no draft stamp"


def test_order_actual_template_needs_an_issued_cycle(client):
    r = client.get("/api/ingredients/order-actual.template.xlsx",
                   params={"period_start": "1999-01-04"})
    assert r.status_code == 404
    assert "ออกรอบสั่ง" in r.json()["detail"]
