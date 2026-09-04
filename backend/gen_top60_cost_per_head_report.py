#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Top-60 menu COST-PER-HEAD report for the MASTER branches, one month.

Deliverable: a standalone, printable Thai HTML report (house style, A4) that
opens with a combined all-master overview, then breaks down each master branch
with its FULL Top-60 menus by cost-per-head, flags the material deviations of
each branch vs the overview, and ends with an inspect / develop / follow-up
action plan.

Methodology (mirrors the app's single source of truth, _unified_cogs_expr):
  • Per menu line COGS = 0 for Easy buffet TIER rows (BUFFET/350/500/590/799/
    CHILD) — consumption is logged on the item-level B0xxx rows, so counting the
    tier too would double-count — else qty × effective unit cost, where the unit
    cost prefers the owner-maintained sku_cost over the BoM recipe (2026-06-11
    priority flip) via the pre-built sku_effective_cost lookup.
  • cost-per-head of a menu = (that menu's monthly COGS in the scope)
                              ÷ (the scope's monthly covers = Σ customer_count).
  • Top-60 = order by COGS desc (≡ cost-per-head desc within a fixed scope).
  • Overview scope = all 5 active master branches combined.

Run:  python backend/gen_top60_cost_per_head_report.py [--period 2026-05] [--db PATH]
Out:  docs/TOP60_COST_PER_HEAD_<period>.html
"""
from __future__ import annotations
import argparse
import html
import os
import sqlite3
import sys

import rules   # the app's own revenue / cost definitions

# ── locate things ────────────────────────────────────────────────────────────
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)


def _find_db(explicit: str | None) -> str:
    cands = []
    if explicit:
        cands.append(explicit)
    cands += [
        os.path.join(HERE, "oshinei.db"),
        os.path.join(REPO, "backend", "oshinei.db"),
        # worktrees share the same OneDrive folder; the live DB lives in the
        # canonical (non-worktree) backend dir.
        r"C:\Users\ASUS\OneDrive\Daiki\2026\Cost Analysis\backend\oshinei.db",
    ]
    best = None
    for c in cands:
        try:
            if os.path.exists(c) and os.path.getsize(c) > 1_000_000:
                if best is None or os.path.getsize(c) > os.path.getsize(best):
                    best = c
        except OSError:
            pass
    if not best:
        sys.exit("Could not find a populated oshinei.db (looked in %s)" % cands)
    return best


# ── COGS: the app's unified rule, replicated verbatim ───────────────────────
TIER = rules.EASY_TIER_CODES_SQL
COGS = rules.cogs_expr('ssm')      # identical to every report in the app

# premium / high-value ingredient keywords → drives the "what to check" hint
PREMIUM_KW = [
    "แซลมอน", "ทุเรียน", "หอยเชลล์", "กุ้ง", "ปู", "ทูน่า", "อูนิ", "ไข่ปลา",
    "ไข่หอยเม่น", "คาเวียร์", "ปลาดิบ", "ซาชิมิ", "ดอง", "เนื้อ", "วากิว",
    "โอโทโร่", "ชูโทโร่", "อิคุระ", "หอยนางรม", "สเต็ก", "ออส", "ปลาไหล", "อูนางิ",
]


def is_premium(desc: str) -> bool:
    d = desc or ""
    return any(k in d for k in PREMIUM_KW)


# ── data layer ───────────────────────────────────────────────────────────────
def covers_of(con, period, branch_ids) -> int:
    ph = ",".join("?" * len(branch_ids))
    r = con.execute(
        f"SELECT COALESCE(SUM(customer_count),0) FROM sales_month_summary "
        f"WHERE period_month=? AND branch_id IN ({ph})",
        [period, *branch_ids],
    ).fetchone()
    return int(r[0] or 0)


def scope_revenue(con, period, branch_ids) -> float:
    """Total food revenue ex-VAT over ALL rows in the scope — INCLUDING the
    buffet package (tier) rows that carry the buffet revenue. Used as the
    food-cost-% denominator. (The Top-60 menu list excludes tier rows because
    their COGS is 0, so their revenue must be added back here.)"""
    ph = ",".join("?" * len(branch_ids))
    r = con.execute(
        f"SELECT COALESCE(SUM(subtotal_ex_vat),0) FROM sales_sku_monthly "
        f"WHERE period_month=? AND branch_id IN ({ph})",
        [period, *branch_ids],
    ).fetchone()
    return float(r[0] or 0)


def scope_menus(con, period, branch_ids):
    """All costed menus (cogs>0) for a scope, rolled up by sku, COGS desc."""
    ph = ",".join("?" * len(branch_ids))
    rows = con.execute(
        f"""
        SELECT ssm.sku_id                         AS sku_id,
               sk.sku_code                        AS code,
               sk.description                     AS desc,
               ROUND(SUM(ssm.qty),0)              AS qty,
               ROUND(SUM({COGS}),2)               AS cogs,
               ROUND(SUM(ssm.subtotal_ex_vat),2)  AS rev
        FROM sales_sku_monthly ssm
        JOIN skus sk ON sk.sku_id=ssm.sku_id
        WHERE ssm.period_month=? AND ssm.branch_id IN ({ph})
        GROUP BY ssm.sku_id
        HAVING SUM({COGS}) > 0
        ORDER BY cogs DESC
        """,
        [period, *branch_ids],
    ).fetchall()
    return [dict(r) for r in rows]


def build_scope(con, period, name, branch_ids, overview_covers=None):
    covers = covers_of(con, period, branch_ids)
    menus = scope_menus(con, period, branch_ids)
    total_cogs = sum(m["cogs"] for m in menus)
    total_rev = scope_revenue(con, period, branch_ids)   # ALL rows incl. buffet tiers
    for m in menus:
        m["cph"] = (m["cogs"] / covers) if covers else 0.0       # cost / head
        m["qph"] = (m["qty"] / covers) if covers else 0.0        # qty  / head
        m["share"] = (m["cogs"] / total_cogs) if total_cogs else 0.0
    return {
        "name": name,
        "branch_ids": branch_ids,
        "covers": covers,
        "menus": menus,
        "total_cogs": total_cogs,
        "total_rev": total_rev,
        "cph_total": (total_cogs / covers) if covers else 0.0,
        "food_pct": (total_cogs / total_rev * 100) if total_rev else 0.0,
        "by_sku": {m["sku_id"]: m for m in menus},
    }


# ── number / html helpers ────────────────────────────────────────────────────
def b0(n):   return f"{n:,.0f}"
def b2(n):   return f"{n:,.2f}"
def pct(n, d=1): return f"{n:.{d}f}%"
def esc(s):  return html.escape(str(s or ""))


def signed(n, d=2):
    s = f"{n:+,.{d}f}"
    return s


# ── report builder ───────────────────────────────────────────────────────────
TOPN = 60

CSS = """
:root{
  --ink:#1B2A30; --ink2:#465056; --muted:#8B979D; --faint:#B7C0C4;
  --teal:#13595A; --teal-d:#0E4445; --accent:#B0472D; --accent-d:#933A23;
  --green:#1F7A4D; --red:#B0472D; --gold:#8A6D1F;
  --line:#E3E6E7; --line2:#EEF0F1; --soft:#F5F7F7; --soft2:#FAFBFB; --bg:#FFFFFF;
}
*{box-sizing:border-box}
html{-webkit-print-color-adjust:exact;print-color-adjust:exact}
body{font-family:'IBM Plex Sans Thai','Sarabun',-apple-system,'Segoe UI',Tahoma,sans-serif;
  color:var(--ink);background:#EDEFF0;margin:0;line-height:1.62;font-size:14px}
.page{width:210mm;min-height:296mm;margin:0 auto 9mm;padding:19mm 18mm 16mm;background:var(--bg);
  position:relative;box-shadow:0 1px 8px rgba(0,0,0,.07)}
@page{size:A4;margin:0}
@media print{
  *{-webkit-print-color-adjust:exact;print-color-adjust:exact}
  body{background:#fff;font-size:10.2px}
  .page{box-shadow:none;margin:0;width:auto;min-height:auto;padding:12mm 12mm;page-break-after:always;overflow:hidden}
  .page:last-child{page-break-after:auto}
  .noprint{display:none!important}
  h2,h3,h4{page-break-after:avoid}
  table{page-break-inside:auto} tr{page-break-inside:avoid}
  thead{display:table-header-group}
  .callout,.figrow{page-break-inside:avoid}
}
h1,h2,h3,h4{line-height:1.3;margin:0 0 .5em;font-weight:700;color:var(--ink)}
h2{font-size:21px;letter-spacing:-.2px;padding-bottom:.3em;border-bottom:2px solid var(--ink);margin-bottom:.7em}
h3{font-size:16px;color:var(--teal-d);margin-top:1.4em}
h4{font-size:14px;color:var(--ink);margin-top:1.1em;font-weight:700}
p{margin:0 0 .75em} a{color:var(--teal-d)}
.lead{font-size:15px;color:var(--ink2);line-height:1.68}
.small{font-size:11.5px;color:var(--muted)} .tiny{font-size:10.5px;color:var(--muted)}
.kicker{font-size:11px;font-weight:700;letter-spacing:2.5px;text-transform:uppercase;color:var(--accent);margin-bottom:.5em}
.sec-no{font-size:11px;font-weight:700;letter-spacing:2px;text-transform:uppercase;color:var(--muted);margin-bottom:.3em}
ul,ol{margin:.2em 0 1em;padding-left:1.2em} li{margin:.28em 0}
strong,b{color:var(--ink);font-weight:700} em{color:var(--ink2)}
hr{border:none;border-top:1px solid var(--line);margin:1.3em 0}
.accent{color:var(--accent)} .tealc{color:var(--teal-d)} .greenc{color:var(--green)} .redc{color:var(--red)}
.pfoot{position:absolute;left:18mm;right:18mm;bottom:9mm;display:flex;justify-content:space-between;
  font-size:9.5px;color:var(--muted);border-top:1px solid var(--line);padding-top:6px;letter-spacing:.3px}
@media print{.pfoot{left:12mm;right:12mm;bottom:7mm}}
.tag{display:inline-block;font-size:10px;font-weight:700;padding:1px 7px;border-radius:3px;
  color:var(--teal-d);border:1px solid var(--line);letter-spacing:.3px;vertical-align:middle}
.tag.s{color:var(--accent);border-color:#e8cfc7} .tag.green{color:var(--green);border-color:#cfe4d8}
.tag.red{color:var(--red);border-color:#e8cfc7} .tag.g{color:var(--gold);border-color:#e6dcc0}
.pdfbtn{position:fixed;top:14px;right:14px;z-index:9999;background:var(--ink);color:#fff;border:none;
  padding:9px 16px;border-radius:5px;font-size:12.5px;font-weight:700;cursor:pointer;font-family:inherit;
  box-shadow:0 2px 10px rgba(0,0,0,.22)}
.pdfbtn:hover{background:#000}
@media print{.pdfbtn{display:none!important}}
/* cover */
.cover .cbody{padding:30mm 18mm 14mm;border-top:3px solid var(--ink)}
.cover .wm{font-size:12px;font-weight:700;letter-spacing:3px;text-transform:uppercase;color:var(--ink)}
.cover h1{font-size:36px;line-height:1.14;letter-spacing:-.5px;margin:.4em 0 .25em;color:var(--ink)}
.cover .sub{font-size:17px;color:var(--ink2);font-weight:500;line-height:1.5;max-width:92%}
.cover .rule{height:1px;background:var(--line);margin:22px 0}
.cover .meta{font-size:12.5px;color:var(--ink2);line-height:1.9}
.cover .kpis{display:flex;margin-top:26px;border-top:2px solid var(--ink)}
.cover .kpis .k{flex:1;padding:13px 14px 0 0}
.cover .kpis .k+ .k{border-left:1px solid var(--line);padding-left:16px}
.cover .kpis .k .n{font-size:25px;font-weight:800;color:var(--ink);line-height:1.05;letter-spacing:-.5px}
.cover .kpis .k .n.a{color:var(--accent)}
.cover .kpis .k .l{font-size:10.5px;color:var(--ink2);margin-top:5px;line-height:1.4}
/* figrow */
.figrow{display:flex;gap:0;margin:1.1em 0;border-top:2px solid var(--ink);border-bottom:1px solid var(--line)}
.figrow .fig{flex:1;padding:12px 14px 12px 0}
.figrow .fig+ .fig{border-left:1px solid var(--line);padding-left:16px}
.fig .n{font-size:22px;font-weight:800;color:var(--ink);line-height:1.05;letter-spacing:-.4px}
.fig .n.a{color:var(--accent)} .fig .n.green{color:var(--green)} .fig .n.teal{color:var(--teal-d)}
.fig .l{font-size:11px;color:var(--ink2);margin-top:5px;font-weight:600;line-height:1.35}
.fig .h{font-size:10px;color:var(--muted);margin-top:3px;line-height:1.35}
/* callouts */
.callout{border-left:3px solid var(--teal);background:var(--soft2);padding:10px 14px;margin:1em 0}
.callout.s{border-left-color:var(--accent)} .callout.green{border-left-color:var(--green)}
.callout.red{border-left-color:var(--red)} .callout.g{border-left-color:var(--gold)}
.callout .ct{font-weight:700;font-size:13px;margin-bottom:.25em;color:var(--ink)}
.callout.s .ct{color:var(--accent-d)} .callout.red .ct{color:var(--red)} .callout.green .ct{color:var(--green)}
.callout p:last-child{margin-bottom:0}
/* tables */
table{width:100%;border-collapse:collapse;margin:.8em 0;font-size:11.6px}
th,td{padding:5px 7px;text-align:left;border-bottom:1px solid var(--line2);vertical-align:top}
th{color:var(--ink);font-weight:700;font-size:10.8px;letter-spacing:.2px;border-bottom:2px solid var(--ink);
  background:var(--soft)}
th.r,td.r{text-align:right;font-variant-numeric:tabular-nums} th.c,td.c{text-align:center}
tbody tr:last-child td{border-bottom:1px solid var(--line)}
tr.tot td{border-top:2px solid var(--ink);font-weight:700;background:var(--soft)}
td.pos{color:var(--red);font-weight:700} td.neg{color:var(--green);font-weight:700}
td.muted{color:var(--muted)}
.rk{color:var(--muted);font-variant-numeric:tabular-nums;font-size:10.5px}
.code{font-variant-numeric:tabular-nums;color:var(--ink2);font-size:10.6px;white-space:nowrap}
.prem{display:inline-block;width:7px;height:7px;border-radius:50%;background:var(--accent);margin-right:5px;vertical-align:middle}
.tbl-cap{font-size:10.6px;color:var(--muted);margin:-.3em 0 1em}
.legend{font-size:10.5px;color:var(--muted);margin:.4em 0 0}
.legend b{color:var(--ink)}
.bar-mini{display:inline-block;height:8px;background:var(--teal);border-radius:2px;vertical-align:middle;margin-left:6px}
.toc{columns:2;column-gap:30px;font-size:12.5px}
.toc a{display:block;color:var(--ink2);text-decoration:none;padding:3px 0;border-bottom:1px dotted var(--line)}
.toc a b{color:var(--accent);margin-right:8px}
"""


def fig(n, label, hint="", cls=""):
    return (f'<div class="fig"><div class="n {cls}">{n}</div>'
            f'<div class="l">{label}</div>'
            + (f'<div class="h">{hint}</div>' if hint else "") + "</div>")


def menu_table(scope, overview, with_variance):
    """Render a Top-60 menu table. with_variance adds overview comparison cols."""
    ov = overview["by_sku"]
    ov_cph_total = overview["cph_total"]
    rows = scope["menus"][:TOPN]
    cum = 0.0
    out = ['<table>']
    if with_variance:
        out.append(
            '<thead><tr>'
            '<th class="c">#</th><th>เมนู</th>'
            '<th class="r">จำนวน</th><th class="r">/หัว</th>'
            '<th class="r">ต้นทุนรวม</th><th class="r">ต้นทุน/หัว</th>'
            '<th class="r">ภาพรวม/หัว</th><th class="r">Δ/หัว</th><th class="r">Δ%</th>'
            '</tr></thead><tbody>')
    else:
        out.append(
            '<thead><tr>'
            '<th class="c">#</th><th>เมนู</th>'
            '<th class="r">จำนวน</th><th class="r">/หัว</th>'
            '<th class="r">ต้นทุนรวม</th><th class="r">ต้นทุน/หัว</th>'
            '<th class="r">% ต้นทุน</th><th class="r">สะสม%</th>'
            '</tr></thead><tbody>')
    for i, m in enumerate(rows, 1):
        cum += m["share"]
        dot = '<span class="prem"></span>' if is_premium(m["desc"]) else ''
        name = f'{dot}{esc(m["desc"])} <span class="code">[{esc(m["code"])}]</span>'
        if with_variance:
            o = ov.get(m["sku_id"])
            o_cph = o["cph"] if o else 0.0
            d = m["cph"] - o_cph
            dp = (d / o_cph * 100) if o_cph > 0 else (float("inf") if d > 0 else 0.0)
            dcls = "pos" if d > 0 else ("neg" if d < 0 else "muted")
            dp_txt = "ใหม่" if dp == float("inf") else (signed(dp, 0) + "%")
            o_txt = b2(o_cph) if o else '<span class="muted">—</span>'
            out.append(
                f'<tr><td class="rk c">{i}</td><td>{name}</td>'
                f'<td class="r">{b0(m["qty"])}</td><td class="r">{b2(m["qph"])}</td>'
                f'<td class="r">{b0(m["cogs"])}</td><td class="r"><b>{b2(m["cph"])}</b></td>'
                f'<td class="r muted">{o_txt}</td>'
                f'<td class="r {dcls}">{signed(d)}</td>'
                f'<td class="r {dcls}">{dp_txt}</td></tr>')
        else:
            out.append(
                f'<tr><td class="rk c">{i}</td><td>{name}</td>'
                f'<td class="r">{b0(m["qty"])}</td><td class="r">{b2(m["qph"])}</td>'
                f'<td class="r">{b0(m["cogs"])}</td><td class="r"><b>{b2(m["cph"])}</b></td>'
                f'<td class="r">{pct(m["share"]*100)}</td>'
                f'<td class="r muted">{pct(cum*100)}</td></tr>')
    # totals row for Top-60
    top_cogs = sum(m["cogs"] for m in rows)
    top_cph = sum(m["cph"] for m in rows)
    ncol_pre = 4
    if with_variance:
        out.append(
            f'<tr class="tot"><td></td><td>รวม Top {len(rows)}</td>'
            f'<td></td><td></td>'
            f'<td class="r">{b0(top_cogs)}</td><td class="r">{b2(top_cph)}</td>'
            f'<td></td><td></td><td></td></tr>')
    else:
        out.append(
            f'<tr class="tot"><td></td><td>รวม Top {len(rows)} '
            f'(จากทั้งหมด {len(scope["menus"])} เมนู)</td>'
            f'<td></td><td></td>'
            f'<td class="r">{b0(top_cogs)}</td><td class="r">{b2(top_cph)}</td>'
            f'<td class="r">{pct(top_cogs/scope["total_cogs"]*100)}</td><td></td></tr>')
    out.append('</tbody></table>')
    return "".join(out)


def deviation_decomp(scope, overview):
    """Per-menu contribution to (branch cph - overview cph), across ALL menus
    (union of branch + overview), sorted by signed contribution."""
    ov = overview["by_sku"]
    union = set(scope["by_sku"]) | set(ov)
    items = []
    for sid in union:
        b = scope["by_sku"].get(sid)
        o = ov.get(sid)
        b_cph = b["cph"] if b else 0.0
        o_cph = o["cph"] if o else 0.0
        d = b_cph - o_cph
        desc = (b or o)["desc"]
        code = (b or o)["code"]
        items.append({"sku": sid, "desc": desc, "code": code, "d": d,
                      "b_cph": b_cph, "o_cph": o_cph, "prem": is_premium(desc)})
    items.sort(key=lambda x: x["d"])
    lows = [x for x in items if x["d"] < 0][:8]            # most negative
    highs = [x for x in items if x["d"] > 0][-8:][::-1]    # most positive
    return highs, lows


def page(inner, foot_left, pno):
    return (f'<div class="page"><div class="pfoot"><span>{foot_left}</span>'
            f'<span>หน้า {pno}</span></div>{inner}</div>')


def build_html(period, overview, branches, gen_dt, db_path, scope_label="Master"):
    th_period = thai_month(period)
    nbr = len(branches)
    pages = []
    pno = [0]

    def addp(inner, foot):
        pno[0] += 1
        pages.append(page(inner, foot, pno[0]))

    FOOT = "Oshinei Enterprise — รายงาน Top 60 ต้นทุน/หัว · " + th_period

    # ── COVER ────────────────────────────────────────────────────────────
    cover = f"""
    <div class="cover"><div class="cbody">
      <div class="wm">OSHINEI ENTERPRISE · INTERNAL</div>
      <div class="kicker">รายงานวิเคราะห์ต้นทุนอาหารต่อหัว</div>
      <h1>Top 60 เมนู — ต้นทุนต่อหัวต่อเดือน<br>สาขา {scope_label}</h1>
      <div class="sub">ภาพรวมทุกสาขารวมกัน แล้วแยกราย {nbr} สาขา {scope_label} พร้อมวิเคราะห์ผลต่าง
        ที่มีสาระสำคัญเทียบกับภาพรวม และแนวทางตรวจสอบ–พัฒนา–ติดตามแก้ไข</div>
      <div class="rule"></div>
      <div class="meta">
        <b>รอบข้อมูล:</b> {th_period} &nbsp;·&nbsp;
        <b>ขอบเขต:</b> {nbr} สาขา {scope_label} ({", ".join(b["name"] for b in branches)})<br>
        <b>หน่วยวัดหลัก:</b> ต้นทุนวัตถุดิบ (COGS) ต่อหัวลูกค้า — บาท/หัว<br>
        <b>จัดทำเมื่อ:</b> {gen_dt}
      </div>
      <div class="kpis">
        <div class="k"><div class="n">{b0(overview['covers'])}</div><div class="l">จำนวนหัวลูกค้ารวม (covers)</div></div>
        <div class="k"><div class="n a">฿{b2(overview['cph_total'])}</div><div class="l">ต้นทุนอาหาร/หัว เฉลี่ยทุกสาขา</div></div>
        <div class="k"><div class="n">{pct(overview['food_pct'])}</div><div class="l">Food cost (COGS ÷ ยอดขาย ex-VAT)</div></div>
        <div class="k"><div class="n">฿{b0(overview['total_cogs'])}</div><div class="l">ต้นทุนอาหารรวม (COGS)</div></div>
      </div>
    </div></div>"""

    # methodology + TOC on the cover page tail
    method = f"""
    <h3 style="margin-top:1.6em">วิธีการคำนวณ (อ้างอิงสูตรมาตรฐานของระบบ)</h3>
    <ul class="small" style="font-size:12px;color:var(--ink2)">
      <li><b>ต้นทุน/หัวของเมนู</b> = (ต้นทุนวัตถุดิบของเมนูนั้นทั้งเดือนในขอบเขตนั้น) ÷ (จำนวนหัวลูกค้าของขอบเขต = ผลรวม customer_count). บอกว่า “เมนูนี้กินต้นทุนกี่บาทต่อลูกค้า 1 คน”.</li>
      <li><b>ต้นทุนวัตถุดิบ (COGS)</b> ต่อบรรทัด = จำนวน × ต้นทุนต่อหน่วยที่มีผล โดยใช้ <b>ต้นทุนจริงจากตารางต้นทุน (sku_cost) ก่อน BoM</b> (นโยบาย 2026-06-11). <b>ตัดแพ็กเกจบุฟเฟต์ (BUFFET350/500/590/799/CHILD) = 0</b> เพราะต้นทุนถูกบันทึกที่รายการ B0xxx แล้ว — ถ้านับแพ็กเกจซ้ำจะนับซ้ำสองครั้ง.</li>
      <li><b>Top 60</b> = เรียงเมนูตามต้นทุนรวม (มาก→น้อย) เอา 60 อันดับแรก (ในขอบเขตเดียวกัน การเรียงด้วยต้นทุนรวม = การเรียงด้วยต้นทุน/หัว).</li>
      <li><b>ภาพรวม</b> = 5 สาขา Master รวมกัน. <b>Δ/หัว</b> ในตารางรายสาขา = (ต้นทุน/หัวของสาขา) − (ต้นทุน/หัวของภาพรวมเมนูเดียวกัน); <span class="redc">แดง = สูงกว่าภาพรวม</span>, <span class="greenc">เขียว = ต่ำกว่า</span>. จุด <span class="prem"></span> = วัตถุดิบราคาแพง (แซลมอน/ทุเรียน/หอยเชลล์/กุ้ง/ปลาดิบ ฯลฯ).</li>
    </ul>
    <div class="callout"><div class="ct">อ่านรายงานนี้อย่างไร</div>
    <p class="small" style="margin:0">ต้นทุน/หัว คือมุมที่ “แฟร์” ที่สุดในการเทียบสาขาที่มีขนาดต่างกัน เพราะหารด้วยจำนวนลูกค้าแล้ว
    ถ้าสาขาหนึ่งมีต้นทุน/หัวของแซลมอนสูงกว่าภาพรวมมาก แปลว่าลูกค้าสาขานั้น “กินแซลมอนต่อหัวมากกว่า” หรือ
    มีปัญหาพอร์ชั่น/ของเสีย/การแจกฟรี/การ map รหัสสินค้า — ส่วนที่ต้องตรวจสอบจะชี้ไว้ให้แล้วในแต่ละสาขาและบทวิเคราะห์รวม.</p></div>
    """

    toc_items = [
        ("1", "ภาพรวมทุกสาขารวมกัน — Top 60 ต้นทุน/หัว", "#overview"),
        ("2", "บทวิเคราะห์ผลต่างที่มีสาระสำคัญ (รายสาขาเทียบภาพรวม)", "#analysis"),
    ]
    for i, b in enumerate(branches, 1):
        toc_items.append((f"3.{i}", f"สาขา {b['name']} — Top 60 + ผลต่าง", f"#br{b['branch_ids'][0]}"))
    toc_items.append(("4", "แนวทางตรวจสอบ · พัฒนา · ติดตามแก้ไข", "#action"))
    toc = '<h3>สารบัญ</h3><div class="toc">' + "".join(
        f'<a href="{href}"><b>{no}</b>{esc(t)}</a>' for no, t, href in toc_items) + '</div>'

    addp(cover + method + toc, FOOT)

    # ── §1 OVERVIEW ──────────────────────────────────────────────────────
    ov_top = overview["menus"][:TOPN]
    top10_share = sum(m["share"] for m in ov_top[:10]) * 100
    prem_cph = sum(m["cph"] for m in overview["menus"] if is_premium(m["desc"]))
    prem_share = prem_cph / overview["cph_total"] * 100 if overview["cph_total"] else 0
    ov_inner = f"""
    <a id="overview"></a>
    <div class="sec-no">ส่วนที่ 1</div>
    <h2>ภาพรวมทุกสาขารวมกัน — Top 60 เมนูตามต้นทุน/หัว</h2>
    <p class="lead">รวม {nbr} สาขา {scope_label} เดือน {th_period}: ลูกค้า <b>{b0(overview['covers'])}</b> หัว,
    ต้นทุนอาหารรวม <b>฿{b0(overview['total_cogs'])}</b>, คิดเป็น
    <b>฿{b2(overview['cph_total'])}/หัว</b> (Food cost {pct(overview['food_pct'])} ของยอดขาย ex-VAT).
    เมนู 10 อันดับแรกกินไปแล้ว <b>{pct(top10_share)}</b> ของต้นทุนทั้งหมด และวัตถุดิบราคาแพง (จุดแดง)
    รวมกัน <b>{pct(prem_share)}</b> ของต้นทุน/หัว — นี่คือจุดที่การควบคุมให้ผลตอบแทนสูงสุด.</p>
    <div class="figrow">
      {fig('฿'+b2(overview['cph_total']),'ต้นทุน/หัว เฉลี่ย','ฐานเทียบของทุกสาขา','a')}
      {fig(b0(overview['covers']),'จำนวนหัวลูกค้า',f'{nbr} สาขารวมกัน')}
      {fig(pct(top10_share),'Top 10 = สัดส่วนต้นทุน','กระจุกตัวสูง')}
      {fig(pct(prem_share),'วัตถุดิบแพง = สัดส่วน','แซลมอน/ทุเรียน/ทะเล','teal')}
    </div>
    {menu_table(overview, overview, with_variance=False)}
    <div class="legend">หมายเหตุ: “/หัว” = ปริมาณหรือบาท หารด้วยจำนวนหัวลูกค้า {b0(overview['covers'])}.
    “สะสม%” ช่วยดูการกระจุกตัวของต้นทุน (กฎ 80/20).</div>
    """
    # overview is long — split table across pages naturally; one page is fine in print
    addp(ov_inner, FOOT)

    # ── §2 CROSS-BRANCH ANALYSIS (built from data) ───────────────────────
    # rank branches by total cph
    ranked = sorted(branches, key=lambda b: b["cph_total"], reverse=True)
    rows = []
    for b in ranked:
        gap = b["cph_total"] - overview["cph_total"]
        highs, lows = deviation_decomp(b, overview)
        drv = " · ".join(f'{esc(h["desc"][:18])} {signed(h["d"],1)}' for h in highs[:3])
        sav = " · ".join(f'{esc(l["desc"][:18])} {signed(l["d"],1)}' for l in lows[:2])
        rows.append((b, gap, drv, sav))
    arr = '<table><thead><tr><th>สาขา</th><th class="r">หัวลูกค้า</th>' \
          '<th class="r">ต้นทุน/หัว</th><th class="r">Δ vs ภาพรวม</th>' \
          '<th class="r">Food%</th><th>เมนูที่ดันต้นทุนขึ้น (บาท/หัว)</th></tr></thead><tbody>'
    for b, gap, drv, sav in rows:
        gcls = "pos" if gap > 0 else "neg"
        arr += (f'<tr><td><b>{esc(b["name"])}</b></td><td class="r">{b0(b["covers"])}</td>'
                f'<td class="r"><b>฿{b2(b["cph_total"])}</b></td>'
                f'<td class="r {gcls}">{signed(gap)}</td>'
                f'<td class="r">{pct(b["food_pct"])}</td>'
                f'<td class="tiny">{drv}</td></tr>')
    arr += (f'<tr class="tot"><td>ภาพรวม (ฐานเทียบ)</td><td class="r">{b0(overview["covers"])}</td>'
            f'<td class="r">฿{b2(overview["cph_total"])}</td><td class="r">—</td>'
            f'<td class="r">{pct(overview["food_pct"])}</td><td></td></tr>')
    arr += '</tbody></table>'

    hottest = ranked[0]
    coolest = ranked[-1]
    h_gap = hottest["cph_total"] - overview["cph_total"]
    h_highs, _ = deviation_decomp(hottest, overview)
    # excess baht for hottest branch on its covers
    h_excess_month = h_gap * hottest["covers"]
    drv_txt = ", ".join(f'{esc(h["desc"])} (+฿{b2(h["d"])}/หัว)' for h in h_highs[:3])

    analysis_inner = f"""
    <a id="analysis"></a>
    <div class="sec-no">ส่วนที่ 2</div>
    <h2>บทวิเคราะห์ผลต่างที่มีสาระสำคัญ (รายสาขาเทียบภาพรวม)</h2>
    <p class="lead">เรียงสาขาจาก “ต้นทุน/หัวสูงสุด” ลงมา และแยกว่าเมนูใดเป็นตัวดันต้นทุนของแต่ละสาขา
    ให้สูง/ต่ำกว่าภาพรวม (Δ/หัว เป็นบาทต่อหัว). ช่องนี้คือจุดตั้งต้นของการตรวจสอบ.</p>
    {arr}
    <div class="legend"><b>Δ vs ภาพรวม</b> = ต้นทุน/หัวรวมของสาขา − ฿{b2(overview['cph_total'])} (ภาพรวม).
    <span class="redc">แดง = สาขาแพงกว่าค่าเฉลี่ยกลุ่ม</span>, <span class="greenc">เขียว = ถูกกว่า</span>.</div>

    <div class="callout {'red' if h_gap>0 else 'green'}">
      <div class="ct">สาขาที่ต้นทุน/หัวสูงสุด: {esc(hottest['name'])} (฿{b2(hottest['cph_total'])}/หัว, {signed(h_gap)} จากภาพรวม)</div>
      <p>ตัวขับหลักคือ {drv_txt}. ถ้านำส่วนต่าง {signed(h_gap)} บาท/หัว คูณจำนวนลูกค้า {b0(hottest['covers'])} หัวของเดือนนี้
      ≈ <b>฿{b0(abs(h_excess_month))}</b> ที่สาขานี้ใช้ต้นทุน{'มากกว่า' if h_gap>0 else 'น้อยกว่า'}ค่าเฉลี่ยกลุ่ม — เป็นเป้าตรวจสอบพอร์ชั่น/ของเสีย/การ map รหัสเป็นลำดับแรก.</p>
    </div>
    <p class="small">รายละเอียดเมนูครบ Top 60 และคอลัมน์ Δ/หัวของแต่ละสาขา อยู่ในส่วนที่ 3 ถัดไป (1 สาขา ต่อ 1 ชุดตาราง).</p>
    """
    addp(analysis_inner, FOOT)

    # ── §3 PER-BRANCH ────────────────────────────────────────────────────
    for i, b in enumerate(branches, 1):
        gap = b["cph_total"] - overview["cph_total"]
        gcls = "pos" if gap > 0 else "neg"
        highs, lows = deviation_decomp(b, overview)
        hi_html = "".join(
            f'<li>{"<span class=prem></span>" if h["prem"] else ""}<b>{esc(h["desc"])}</b> '
            f'<span class="redc">{signed(h["d"])}/หัว</span> '
            f'(สาขา ฿{b2(h["b_cph"])} vs ภาพรวม ฿{b2(h["o_cph"])})</li>'
            for h in highs[:5])
        lo_html = "".join(
            f'<li><b>{esc(l["desc"])}</b> <span class="greenc">{signed(l["d"])}/หัว</span> '
            f'(สาขา ฿{b2(l["b_cph"])} vs ภาพรวม ฿{b2(l["o_cph"])})</li>'
            for l in lows[:5])
        inner = f"""
        <a id="br{b['branch_ids'][0]}"></a>
        <div class="sec-no">ส่วนที่ 3.{i}</div>
        <h2>สาขา {esc(b['name'])} — Top 60 ต้นทุน/หัว และผลต่างจากภาพรวม</h2>
        <div class="figrow">
          {fig(b0(b['covers']),'จำนวนหัวลูกค้า','เดือน '+th_period)}
          {fig('฿'+b2(b['cph_total']),'ต้นทุน/หัว','vs ภาพรวม ฿'+b2(overview['cph_total']),'a' if gap>0 else 'green')}
          {fig(signed(gap),'Δ ต้นทุน/หัว','แดง=สูงกว่า / เขียว=ต่ำกว่า', 'a' if gap>0 else 'green')}
          {fig(pct(b['food_pct']),'Food cost %','COGS ÷ ยอดขาย ex-VAT')}
          {fig(b0(b['total_cogs']),'ต้นทุนรวม (฿)','{} เมนูมีต้นทุน'.format(len(b['menus'])))}
        </div>
        <div style="display:flex;gap:18px;margin:.6em 0 0">
          <div style="flex:1"><h4 style="margin-top:0;color:var(--red)">▲ เมนูที่ดันต้นทุนสูงกว่าภาพรวม</h4>
            <ul class="small">{hi_html or '<li class="muted">— ไม่มีรายการเด่น</li>'}</ul></div>
          <div style="flex:1"><h4 style="margin-top:0;color:var(--green)">▼ เมนูที่ต่ำกว่าภาพรวม</h4>
            <ul class="small">{lo_html or '<li class="muted">— ไม่มีรายการเด่น</li>'}</ul></div>
        </div>
        {menu_table(b, overview, with_variance=True)}
        <div class="legend">Δ/หัว = ต้นทุน/หัวของสาขา − ภาพรวมเมนูเดียวกัน · “ใหม่” = เมนูที่ภาพรวมแทบไม่มีแต่สาขานี้มี (เมนูเฉพาะถิ่น/โปรโมชัน/ตรวจการ map รหัส).</div>
        """
        addp(inner, FOOT)

    # ── §4 ACTION PLAN ───────────────────────────────────────────────────
    # collect premium hot spots across branches for concrete targeting
    hot_targets = {}
    for b in branches:
        highs, _ = deviation_decomp(b, overview)
        for h in highs[:4]:
            if h["prem"] and h["d"] > 1:
                hot_targets.setdefault(h["desc"], []).append((b["name"], h["d"]))
    ht_lines = ""
    for desc, lst in sorted(hot_targets.items(), key=lambda kv: -sum(d for _, d in kv[1]))[:8]:
        brs = ", ".join(f'{esc(n)} {signed(d,1)}' for n, d in sorted(lst, key=lambda x: -x[1]))
        ht_lines += f'<li><b>{esc(desc)}</b> — สาขาที่สูงกว่าภาพรวม: {brs} (บาท/หัว)</li>'

    action_inner = f"""
    <a id="action"></a>
    <div class="sec-no">ส่วนที่ 4</div>
    <h2>แนวทางตรวจสอบ · พัฒนา · ติดตามแก้ไข</h2>
    <p class="lead">กรอบลงมือ 3 จังหวะ — ตรวจสอบหาสาเหตุก่อน, แก้ที่กระบวนการ, แล้ววัดซ้ำให้แน่ใจว่าลดได้จริงและไม่กลับมา.</p>

    <h3>① ตรวจสอบ (หาสาเหตุของผลต่าง)</h3>
    <ul>
      <li><b>โฟกัสที่จุดแดงราคาแพงก่อน</b> — เมนูที่ Δ/หัวบวกมากในตารางรายสาขา ให้ดึง “บิลขาดทุน/บิลส่วนลด” และ Void/ยกเลิก ของเมนูนั้นในเดือนนี้มาดูว่าเป็นการ <i>บริโภคจริง, ของเสีย, แจกฟรี, หรือการ map รหัสผิด</i>.</li>
      <li><b>ตรวจพอร์ชั่นจริงเทียบมาตรฐาน (yield test)</b> — ชั่งน้ำหนัก/นับชิ้นต่อจานจริงของแซลมอน/หอยเชลล์/กุ้ง/ทุเรียน เทียบสูตร BoM; ส่วนต่างพอร์ชั่นแปลงเป็นบาท/หัวได้ทันที.</li>
      <li><b>ตรวจ Mapping รหัสสินค้า/ต้นทุน (sku_cost)</b> — เมนูที่ขึ้น “ใหม่” ในบางสาขาหรือต้นทุน/หน่วยผิดปกติ มักเกิดจากรหัสซ้ำ/ต้นทุนต่อหน่วยยังไม่อัปเดต; ตรวจ sku_cost เทียบราคาซื้อจริงล่าสุด.</li>
      <li><b>กระทบยอดของเสีย (waste log)</b> — ให้แต่ละสาขาลงบันทึกของเสีย/ของหมดอายุรายวันของ Top-10 วัตถุดิบแพง 2 สัปดาห์ แล้วเทียบกับส่วนต่างต้นทุน/หัว.</li>
    </ul>
    <div class="callout s"><div class="ct">เป้าตรวจสอบเฉพาะจุด (วัตถุดิบแพงที่บางสาขาสูงกว่าภาพรวม)</div>
      <ul class="small" style="margin:.2em 0 0">{ht_lines or '<li class="muted">ไม่มีจุดเด่นเกินเกณฑ์</li>'}</ul></div>

    <h3>② พัฒนา (แก้ที่กระบวนการ ไม่ใช่แค่ตักเตือน)</h3>
    <ul>
      <li><b>มาตรฐานพอร์ชั่นเป็นภาพ (visual portion guide)</b> ติดที่ครัวทุกสาขา + เครื่องชั่งประจำสเตชั่นแซลมอน/ทะเล เพื่อให้ทุกสาขาเสิร์ฟต่อหัวเท่ากัน.</li>
      <li><b>คุมการเติม (refill) ของวัตถุดิบแพงในบุฟเฟต์</b> — เสิร์ฟเป็นรอบ/จำกัดต่อโต๊ะสำหรับแซลมอน/หอยเชลล์/ทุเรียน เพื่อลดต้นทุน/หัวในสาขาที่สูงผิดปกติ โดยไม่กระทบประสบการณ์.</li>
      <li><b>Menu engineering</b> — เมนูต้นทุน/หัวสูงแต่คนสั่งน้อย พิจารณาปรับสูตร/ลดความถี่/ทดแทนวัตถุดิบ; เมนูที่ทำกำไรดีดันให้ขายมากขึ้น.</li>
      <li><b>จัดซื้อ/สเปกวัตถุดิบรวมศูนย์</b> — ล็อกสเปกและราคาแซลมอน/กุ้ง/หอยเชลล์ให้เท่ากันทุกสาขา ลดผลต่างต้นทุน/หน่วยที่ทำให้ต้นทุน/หัวเพี้ยน.</li>
    </ul>

    <h3>③ ติดตามแก้ไข (วัดซ้ำ ปิดลูป)</h3>
    <ul>
      <li><b>ตั้งเป้าต้นทุน/หัวรายสาขา</b> = ภาพรวม ฿{b2(overview['cph_total'])}/หัว (หรือดีกว่า). สาขาที่เกินเป้าให้มีแผนลดเป็นรายเมนูที่ชี้ไว้ในส่วนที่ 3.</li>
      <li><b>ทบทวนรายเดือนด้วยรายงานชุดเดิม</b> — รันรายงานนี้ทุกเดือน (สคริปต์เดียวกัน เปลี่ยน --period) ดูว่า Δ/หัวของเมนูเป้าหมายลดลงหรือไม่.</li>
      <li><b>มอบหมายเจ้าของงาน (owner) ต่อสาขา</b> + กำหนดวันวัดซ้ำ (เช่น 30 วัน) และเกณฑ์ผ่าน (Δ/หัวของเมนูเป้าหมายลด ≥ 50%).</li>
      <li><b>ยกระดับเมื่อไม่ดีขึ้น</b> — ถ้าผ่านการอบรม+คุมพอร์ชั่นแล้ว Δ/หัวยังสูง ให้สอบทานเชิงลึก (สต๊อก, การรับเข้า, การเข้าถึงระบบ) ตามแนวทางควบคุมภายใน.</li>
    </ul>
    <div class="callout green"><div class="ct">เกณฑ์ “มีสาระสำคัญ” ที่ใช้ในรายงานนี้</div>
      <p class="small" style="margin:0">เมนูถูกชี้เป็นผลต่างสำคัญเมื่อ Δ/หัว ของสาขาต่างจากภาพรวมอย่างเห็นได้ชัดและเป็นบาท/หัวที่มีนัย
      (ให้ความสำคัญกับวัตถุดิบราคาแพงเป็นพิเศษ เพราะให้ผลตอบแทนต่อการควบคุมสูงสุด).
      ตัวเลขทั้งหมดคำนวณจากฐานข้อมูลจริง — ส่วน “สาเหตุ” เป็นสมมุติฐานเพื่อการตรวจสอบ ไม่ใช่ข้อสรุปความผิด.</p></div>
    """
    addp(action_inner, FOOT)

    pdfbtn = ('<button class="pdfbtn noprint" onclick="window.print()">'
              '⬇ บันทึก / พิมพ์ PDF</button>')
    head = (f'<!DOCTYPE html><html lang="th"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>Top 60 ต้นทุน/หัว สาขา {scope_label} — {th_period}</title>'
            f'<link rel="preconnect" href="https://fonts.googleapis.com">'
            f'<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans+Thai:wght@400;500;600;700'
            f'&family=Sarabun:wght@400;500;600;700;800&display=swap" rel="stylesheet">'
            f'<style>{CSS}</style></head><body>{pdfbtn}')
    return head + "".join(pages) + "</body></html>"


def thai_month(period):
    y, m = period.split("-")
    th = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
          "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
    return f"{th[int(m)]} {int(y)+543} ({period})"


# ── Excel export (same data pipeline as the HTML — single source of truth) ────
def write_xlsx(path, period, overview, branches, gen_dt, scope_label="Master"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Side, Border
    from openpyxl.utils import get_column_letter

    INK = "FF1B2A30"; RED = "FFB0472D"; GREEN = "FF1F7A4D"
    SOFT = "FFF1F4F4"; TOTF = "FFE7EBEC"; LINE = "FFD9DEDF"
    f_title = Font(name="Tahoma", size=13, bold=True, color=INK)
    f_sub   = Font(name="Tahoma", size=9,  color="FF8B979D")
    f_kpi   = Font(name="Tahoma", size=10, bold=True, color=INK)
    f_hdr   = Font(name="Tahoma", size=9.5, bold=True, color="FFFFFFFF")
    f_cell  = Font(name="Tahoma", size=10, color=INK)
    f_bold  = Font(name="Tahoma", size=10, bold=True, color=INK)
    f_red   = Font(name="Tahoma", size=10, bold=True, color=RED)
    f_green = Font(name="Tahoma", size=10, bold=True, color=GREEN)
    fill_hdr = PatternFill("solid", fgColor=INK)
    fill_tot = PatternFill("solid", fgColor=TOTF)
    AL_R = Alignment(horizontal="right"); AL_C = Alignment(horizontal="center")
    AL_L = Alignment(horizontal="left"); AL_W = Alignment(horizontal="left", wrap_text=True, vertical="top")
    thin = Border(bottom=Side(style="thin", color=LINE))
    I = "#,##0"; D2 = "#,##0.00"; P = "0.0%"

    def style_header(ws, row, ncols, first_data_row, autofilter=True):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = f_hdr; cell.fill = fill_hdr
            cell.alignment = AL_C if c >= 4 else AL_L
        ws.freeze_panes = ws.cell(row=first_data_row, column=1)
        if autofilter:
            ws.auto_filter.ref = f"A{row}:{get_column_letter(ncols)}{ws.max_row}"

    th_period = thai_month(period)

    # column plans: (header, width, numfmt, align)  numfmt None = text
    COLS_OV = [
        ("อันดับ", 7, I, AL_C), ("รหัส", 11, None, AL_L), ("เมนู", 46, None, AL_L),
        ("วัตถุดิบแพง", 11, None, AL_C), ("จำนวน", 11, I, AL_R), ("จำนวน/หัว", 11, D2, AL_R),
        ("ต้นทุนรวม (฿)", 14, I, AL_R), ("ต้นทุน/หัว (฿)", 13, D2, AL_R),
        ("% ต้นทุน", 10, P, AL_R), ("สะสม %", 10, P, AL_R),
    ]
    COLS_BR = [
        ("อันดับ", 7, I, AL_C), ("รหัส", 11, None, AL_L), ("เมนู", 46, None, AL_L),
        ("วัตถุดิบแพง", 11, None, AL_C), ("จำนวน", 11, I, AL_R), ("จำนวน/หัว", 11, D2, AL_R),
        ("ต้นทุนรวม (฿)", 14, I, AL_R), ("ต้นทุน/หัว (฿)", 13, D2, AL_R),
        ("ภาพรวม/หัว (฿)", 13, D2, AL_R), ("Δ/หัว (฿)", 12, D2, AL_R), ("Δ %", 10, P, AL_R),
    ]

    def write_menu_sheet(ws, scope, variant):
        cols = COLS_BR if variant else COLS_OV
        gap = scope["cph_total"] - overview["cph_total"]
        ws.cell(1, 1, f"{scope['name']} — Top {TOPN} ต้นทุน/หัว · {th_period}").font = f_title
        kpi = (f"ลูกค้า {scope['covers']:,} หัว   ·   ต้นทุน/หัว ฿{scope['cph_total']:,.2f}"
               + (f"   ·   Δ vs ภาพรวม {gap:+,.2f}" if variant else "")
               + f"   ·   Food cost {scope['food_pct']:.1f}%   ·   ต้นทุนรวม ฿{scope['total_cogs']:,.0f}"
               f"   ·   เมนูมีต้นทุน {len(scope['menus'])} รายการ")
        ws.cell(2, 1, kpi).font = f_kpi
        HR = 4
        for ci, (name, _, _, _) in enumerate(cols, 1):
            ws.cell(HR, ci, name)
        rows = scope["menus"][:TOPN]
        cum = 0.0
        r = HR + 1
        ov = overview["by_sku"]
        for i, m in enumerate(rows, 1):
            cum += m["share"]
            prem = "●" if is_premium(m["desc"]) else ""
            if variant:
                o = ov.get(m["sku_id"]); o_cph = o["cph"] if o else None
                d = m["cph"] - (o_cph or 0.0)
                dp = (d / o_cph) if (o_cph and o_cph > 0) else None
                vals = [i, m["code"], m["desc"], prem, m["qty"], m["qph"],
                        m["cogs"], m["cph"], o_cph, d, dp]
            else:
                vals = [i, m["code"], m["desc"], prem, m["qty"], m["qph"],
                        m["cogs"], m["cph"], m["share"], cum]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(r, ci, v)
                _, _, fmt, al = cols[ci - 1]
                cell.font = f_cell; cell.alignment = al
                if fmt:
                    cell.number_format = fmt
            if variant:
                dcell = ws.cell(r, 10)            # Δ/หัว
                dcell.font = f_red if vals[9] > 0 else (f_green if vals[9] < 0 else f_cell)
                if vals[3] == "●":
                    ws.cell(r, 4).font = f_red
            r += 1
        # totals
        top_cogs = sum(m["cogs"] for m in rows); top_cph = sum(m["cph"] for m in rows)
        ws.cell(r, 3, f"รวม Top {len(rows)} (จากทั้งหมด {len(scope['menus'])} เมนู)")
        ws.cell(r, 7, top_cogs).number_format = I
        ws.cell(r, 8, top_cph).number_format = D2
        if not variant:
            ws.cell(r, 9, (top_cogs / scope["total_cogs"]) if scope["total_cogs"] else 0).number_format = P
        for ci in range(1, len(cols) + 1):
            c = ws.cell(r, ci); c.font = f_bold; c.fill = fill_tot
        for ci, (_, w, _, _) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        style_header(ws, HR, len(cols), HR + 1)

    wb = Workbook()

    # ── สรุป ─────────────────────────────────────────────────────────────
    ws = wb.active; ws.title = "สรุป"; ws.sheet_properties.tabColor = "1B2A30"
    ws.cell(1, 1, f"รายงาน Top 60 เมนู — ต้นทุนต่อหัว · สาขา {scope_label}").font = f_title
    ws.cell(2, 1, f"รอบข้อมูล {th_period}   ·   จัดทำ {gen_dt}   ·   "
                  f"ขอบเขต {len(branches)} สาขา {scope_label}   ·   ตัวเลขจากฐานข้อมูลจริง").font = f_sub
    # KPI block
    ws.cell(4, 1, "ภาพรวมทุกสาขารวมกัน").font = f_bold
    kpis = [("จำนวนหัวลูกค้า (covers)", overview["covers"], I),
            ("ต้นทุนอาหาร/หัว เฉลี่ย (฿)", overview["cph_total"], D2),
            ("Food cost (COGS ÷ ยอดขาย ex-VAT)", overview["food_pct"] / 100, P),
            ("ต้นทุนอาหารรวม (฿)", overview["total_cogs"], I),
            ("ยอดขาย ex-VAT รวม (฿)", overview["total_rev"], I),
            ("จำนวนเมนูที่มีต้นทุน", len(overview["menus"]), I)]
    rr = 5
    for label, val, fmt in kpis:
        ws.cell(rr, 1, label).font = f_cell
        c = ws.cell(rr, 2, val); c.font = f_bold; c.number_format = fmt; c.alignment = AL_R
        rr += 1
    # branch ranking
    hr = rr + 1
    ws.cell(hr, 1, "การจัดอันดับสาขา (เรียงต้นทุน/หัว มาก→น้อย)").font = f_bold
    hr += 1
    headers = ["สาขา", "หัวลูกค้า", "ต้นทุน/หัว (฿)", "Δ vs ภาพรวม (฿)", "Food cost %",
               "ต้นทุนรวม (฿)", "เมนูที่ดันต้นทุนขึ้น (บาท/หัว)"]
    for ci, h in enumerate(headers, 1):
        ws.cell(hr, ci, h)
    ranked = sorted(branches, key=lambda b: b["cph_total"], reverse=True)
    dr = hr + 1
    for b in ranked:
        gap = b["cph_total"] - overview["cph_total"]
        highs, _ = deviation_decomp(b, overview)
        drv = " · ".join(f'{h["desc"][:20]} {h["d"]:+.1f}' for h in highs[:3])
        vals = [b["name"], b["covers"], b["cph_total"], gap, b["food_pct"] / 100, b["total_cogs"], drv]
        fmts = [None, I, D2, D2, P, I, None]
        for ci, (v, fmt) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(dr, ci, v); c.font = f_cell
            if fmt:
                c.number_format = fmt; c.alignment = AL_R
            if ci == 4:
                c.font = f_red if gap > 0 else f_green
            if ci == 7:
                c.alignment = AL_W
        dr += 1
    # overview baseline row
    base = [f"ภาพรวม (ฐานเทียบ)", overview["covers"], overview["cph_total"], 0.0,
            overview["food_pct"] / 100, overview["total_cogs"], ""]
    for ci, v in enumerate(base, 1):
        c = ws.cell(dr, ci, v); c.font = f_bold; c.fill = fill_tot
        if ci in (2, 6):
            c.number_format = I; c.alignment = AL_R
        if ci in (3,):
            c.number_format = D2; c.alignment = AL_R
        if ci == 5:
            c.number_format = P; c.alignment = AL_R
        if ci == 4:
            c.value = None
    style_header(ws, hr, len(headers), hr + 1, autofilter=False)
    # methodology
    mr = dr + 2
    ws.cell(mr, 1, "วิธีคิด").font = f_bold; mr += 1
    notes = [
        "• ต้นทุน/หัวของเมนู = (ต้นทุนวัตถุดิบของเมนูนั้นทั้งเดือน) ÷ (จำนวนหัวลูกค้าของขอบเขต = Σ customer_count).",
        "• COGS ใช้สูตรมาตรฐานระบบ: แพ็กเกจบุฟเฟต์ (BUFFET350/500/590/799/CHILD) = 0 (กันนับซ้ำ), อื่นๆ = จำนวน × ต้นทุน/หน่วย (ใช้ sku_cost จริงก่อน BoM).",
        "• Top 60 = เรียงตามต้นทุนรวมมาก→น้อย.  Δ/หัว = ต้นทุน/หัวสาขา − ภาพรวมเมนูเดียวกัน (แดง=สูงกว่า, เขียว=ต่ำกว่า).",
        "• Food cost % = COGS ÷ ยอดขาย ex-VAT (รวมแพ็กเกจบุฟเฟต์).  ● = วัตถุดิบราคาแพง.",
        "• Δ % เว้นว่าง = เมนูที่ภาพรวมแทบไม่มีแต่สาขานี้มี (เมนูเฉพาะถิ่น/โปรโมชัน/ตรวจการ map รหัส).",
        "• ชีต ‘ข้อมูลดิบ’ = ทุกเมนูทุกสาขา (long format) สำหรับทำ PivotTable.",
    ]
    for n in notes:
        c = ws.cell(mr, 1, n); c.font = f_sub; c.alignment = AL_L; mr += 1
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["G"].width = 52

    # ── ภาพรวม + per-branch ──────────────────────────────────────────────
    ov_ws = wb.create_sheet("ภาพรวม Top60"); ov_ws.sheet_properties.tabColor = "B0472D"
    write_menu_sheet(ov_ws, overview, variant=False)
    used = {"สรุป", "ภาพรวม Top60"}
    for b in branches:
        title = b["name"][:28]
        t = title; k = 2
        while t in used:
            t = f"{title[:26]} {k}"; k += 1
        used.add(t)
        bws = wb.create_sheet(t); bws.sheet_properties.tabColor = "13595A"
        write_menu_sheet(bws, b, variant=True)

    # ── ข้อมูลดิบ (long format, ALL menus) ───────────────────────────────
    dws = wb.create_sheet("ข้อมูลดิบ"); dws.sheet_properties.tabColor = "8B979D"
    dcols = [("ขอบเขต", 16, None, AL_L), ("สาขา_id", 9, None, AL_C), ("รหัส", 11, None, AL_L),
             ("เมนู", 44, None, AL_L), ("วัตถุดิบแพง", 11, None, AL_C), ("อันดับ", 8, I, AL_C),
             ("จำนวน", 11, I, AL_R), ("จำนวน/หัว", 11, D2, AL_R), ("ต้นทุนรวม (฿)", 14, I, AL_R),
             ("ต้นทุน/หัว (฿)", 13, D2, AL_R), ("% ต้นทุน", 10, P, AL_R),
             ("ภาพรวม/หัว (฿)", 13, D2, AL_R), ("Δ/หัว (฿)", 12, D2, AL_R), ("Δ %", 10, P, AL_R)]
    for ci, (name, _, _, _) in enumerate(dcols, 1):
        dws.cell(1, ci, name)
    ov = overview["by_sku"]
    r = 2

    def emit(scope, scope_label, bid, variant):
        nonlocal r
        for i, m in enumerate(scope["menus"], 1):  # ALL menus, not just 60
            prem = "●" if is_premium(m["desc"]) else ""
            if variant:
                o = ov.get(m["sku_id"]); o_cph = o["cph"] if o else None
                d = m["cph"] - (o_cph or 0.0)
                dp = (d / o_cph) if (o_cph and o_cph > 0) else None
            else:
                o_cph = d = dp = None
            vals = [scope_label, bid, m["code"], m["desc"], prem, i, m["qty"], m["qph"],
                    m["cogs"], m["cph"], m["share"], o_cph, d, dp]
            for ci, v in enumerate(vals, 1):
                c = dws.cell(r, ci, v); c.font = f_cell
                _, _, fmt, al = dcols[ci - 1]
                c.alignment = al
                if fmt:
                    c.number_format = fmt
            r += 1

    emit(overview, "ภาพรวม(รวม)", "", variant=False)
    for b in branches:
        emit(b, b["name"], b["branch_ids"][0], variant=True)
    for ci, (_, w, _, _) in enumerate(dcols, 1):
        dws.column_dimensions[get_column_letter(ci)].width = w
    style_header(dws, 1, len(dcols), 2)

    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--period", default="2026-05")
    ap.add_argument("--db", default=None)
    ap.add_argument("--out", default=None)
    ap.add_argument("--gen-dt", default="")
    ap.add_argument("--no-html", action="store_true", help="skip the HTML output")
    ap.add_argument("--no-xlsx", action="store_true", help="skip the Excel output")
    args = ap.parse_args()

    db = _find_db(args.db)
    con = sqlite3.connect(db)
    con.row_factory = sqlite3.Row

    # active master branches with data for the period (exclude inactive/no-data)
    ph_master = con.execute(
        "SELECT branch_id, name FROM branches WHERE branch_type='master' AND active=1 "
        "ORDER BY branch_id"
    ).fetchall()
    master = []
    for r in ph_master:
        if covers_of(con, args.period, [r["branch_id"]]) > 0:
            master.append((r["branch_id"], r["name"]))
    if not master:
        sys.exit(f"No master-branch data for {args.period}")

    all_ids = [bid for bid, _ in master]
    overview = build_scope(con, args.period, "ภาพรวมทุกสาขา", all_ids)
    branches = [build_scope(con, args.period, name, [bid]) for bid, name in master]

    gen_dt = args.gen_dt or "(ดูชื่อไฟล์)"
    base = args.out or os.path.join(REPO, "docs", f"TOP60_COST_PER_HEAD_{args.period}")
    base = base[:-5] if base.lower().endswith(".html") else base
    os.makedirs(os.path.dirname(base), exist_ok=True)

    out = base + ".html"
    xlsx = base + ".xlsx"
    if not args.no_html:
        htmlout = build_html(args.period, overview, branches, gen_dt, db)
        with open(out, "w", encoding="utf-8") as f:
            f.write(htmlout)
    if not args.no_xlsx:
        write_xlsx(xlsx, args.period, overview, branches, gen_dt)

    # console summary
    print(f"DB:   {db}")
    if not args.no_html:
        print(f"HTML: {out}  ({os.path.getsize(out):,} bytes)")
    if not args.no_xlsx:
        print(f"XLSX: {xlsx}  ({os.path.getsize(xlsx):,} bytes)")
    print(f"Period {args.period} · masters: {', '.join(n for _, n in master)}")
    print(f"Overview: covers={overview['covers']:,}  COGS={overview['total_cogs']:,.0f}  "
          f"cph={overview['cph_total']:.2f}  food%={overview['food_pct']:.1f}  "
          f"menus={len(overview['menus'])}")
    for b in branches:
        print(f"  {b['name']:<12} covers={b['covers']:>7,}  cph={b['cph_total']:>7.2f}  "
              f"Δ={b['cph_total']-overview['cph_total']:+7.2f}  food%={b['food_pct']:>5.1f}  "
              f"menus={len(b['menus'])}  top60_ok={len(b['menus'])>=60}")


if __name__ == "__main__":
    main()
