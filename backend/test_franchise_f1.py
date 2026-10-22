"""PR-F1 - ร้านของฉัน, more to see: the daily table, the 12-month trend, the target, the cost-coverage chip and the download.

  * payload shape of every new section (sales.daily, trend_12m, target, gp.coverage);
  * the daily table and the 30-day bars hold FINISHED days only (rules.closed_day_clause), and the day a week earlier is
    the comparison;
  * the 12-month trend is month by month, never a sum over months, and its months read as the management report's;
  * the target: a typed target, last year's actual, none (neutral), pro-rated by the report's weekday weights;
  * GET /api/franchise/summary.xlsx: four sheets (F3: a fifth, วัตถุดิบ), a Thai file name, the same 403 rules as the summary, and none of the
    forbidden words or another branch's name in any cell.
"""
import io
import json
import sqlite3
from datetime import date, timedelta

import openpyxl
import pytest

import auth
from test_franchise_role import TOKEN, _login, _walk, auth_on, world      # noqa: F401  (fixtures used by name)

FORBIDDEN = ("recipe", "price", "cashier", "staff", "supplier", "receipt_no", "bill_no", "username", "password", "token", "bom")
CUMULATIVE = ("total", "cumulative", "ytd", "sum", "running", "accum")


def _summary(c, **q):
    r = c.get("/api/franchise/summary", params={"month": "2026-06", **q})
    assert r.status_code == 200, r.text
    return r.json()


# ── the rule (no server) ─────────────────────────────────────────────────────
def test_the_download_is_under_the_one_get_rule_and_nothing_writes():
    mods = auth.ROLE_MODULES["franchise"]
    assert auth.path_allowed("GET", "/api/franchise/summary.xlsx", mods)
    for m in ("POST", "PUT", "PATCH", "DELETE"):
        assert not auth.path_allowed(m, "/api/franchise/summary.xlsx", mods), m


# ── the payload ──────────────────────────────────────────────────────────────
def test_the_new_sections_have_the_agreed_shape(world):
    d = _summary(world["b"])
    s = d["sales"]
    assert [x["date"] for x in s["daily"]] == ["2026-06-01", "2026-06-02", "2026-06-03"]          # the month up to its last day
    assert s["daily_through"] == "2026-06-03" == s["through"] == s["last_day"]["date"]
    for x in s["daily"]:
        assert set(x) == {"date", "weekday", "revenue", "bills", "customers", "per_bill", "vs_last_week_pct"}
    first = s["daily"][0]
    assert first["weekday"] == "จันทร์"                                                         # 1 June 2026 is a Monday
    # the dashboard's revenue rule: 1070 / 1.07 - 50 service charge = 950, one bill of 4 customers
    assert first["revenue"] == 950.0 and first["bills"] == 1 and first["customers"] == 4 and first["per_bill"] == 950
    assert first["vs_last_week_pct"] is None                                                     # 25 May has no sales
    t = d["trend_12m"]
    assert {"months", "latest"} == set(t) and t["months"] and t["latest"] == t["months"][-1]
    for m in t["months"]:
        assert set(m) == {"month", "label", "partial", "through", "days", "revenue", "customers", "bills", "per_bill",
                          "gp_pct", "mom_pct", "mom_basis", "mom_customers_pct", "mom_gp_pp", "yoy_pct"}
    assert t["latest"]["month"] == "2026-06" and t["latest"]["partial"] is True and t["latest"]["through"] == "2026-06-03"
    assert {"source", "basis_month", "month_target", "to_date", "mtd", "vs_to_date_pct", "of_month_pct",
            "elapsed_share_pct", "weighting", "month_complete"} == set(d["target"])
    cov = d["gp"]["coverage"]
    assert set(cov) == {"uncosted_plates_pct", "uncosted_sales_pct", "floor"}
    assert all(v is None or 0 <= v <= 100 for v in (cov["uncosted_plates_pct"], cov["uncosted_sales_pct"]))


def test_the_trend_is_month_by_month_and_never_a_sum(world):
    d = _summary(world["b"])
    bad = [p for p, kind, k in _walk(d) if kind == "key" and any(w in k.lower() for w in CUMULATIVE)]
    assert not bad, bad
    months = d["trend_12m"]["months"]
    assert len(months) <= 13 and [m["month"] for m in months] == sorted(m["month"] for m in months)
    total = round(sum(m["revenue"] or 0 for m in months), 2)
    flat = [v for v in _flat_numbers(d)]
    assert len(months) < 2 or total not in flat                                                    # no 12-month total anywhere
    # each month reads as the management report's month-by-month series does (same engine, same branch)
    from routers.reports import report_management_months
    ref = json.loads(report_management_months(branch="BranchB", group=None, lang="th", closed=12).body)
    by = {m["month"]: m for m in ref["months"]}
    for m in months:
        assert m["revenue"] == round(by[m["month"]]["revenue"], 2) and m["gp_pct"] == by[m["month"]]["gp_pct"]
    # the month in progress is the page's month: same sales, weekday-fair MoM as the headline
    assert months[-1]["revenue"] == d["sales"]["mtd"]


def _flat_numbers(o):
    if isinstance(o, dict):
        for v in o.values():
            yield from _flat_numbers(v)
    elif isinstance(o, list):
        for v in o:
            yield from _flat_numbers(v)
    elif isinstance(o, (int, float)) and not isinstance(o, bool):
        yield round(float(o), 2)


def test_the_daily_table_holds_finished_days_only():
    """rules.closed_day_clause: the day still trading (Thai today) never reaches the table or the bars, even when the
    caller asks for it, and the change is on the same weekday a week earlier."""
    from routers import franchise as F
    con = sqlite3.connect(":memory:")
    con.execute("CREATE TABLE sales_transaction(branch_id INT, sold_at TEXT, total_inc_vat REAL, service_charge REAL, "
                "customer_count INT, source TEXT)")
    today = con.execute("SELECT date('now','+7 hours')").fetchone()[0]
    td = date.fromisoformat(today)
    y, wk = (td - timedelta(days=1)).isoformat(), (td - timedelta(days=8)).isoformat()
    for d, amt in ((today, 10700.0), (y, 2140.0), (wk, 1070.0)):
        con.execute("INSERT INTO sales_transaction VALUES(2, ?, ?, 0, 2, 'Easy')", (d + "T12:00:00", amt))
    sales = {}
    F._daily(con, sales, 2, today)                                     # asked for today: today stays out
    assert all(x["revenue"] is None for x in sales["daily"] if x["date"] == today)
    assert {x["day"]: x["revenue"] for x in sales["series_30d"]}[today] is None
    assert sales["last_day"]["revenue"] == 0.0
    sales = {}
    F._daily(con, sales, 2, y)
    assert sales["daily_through"] == y and sales["last_day"]["revenue"] == 2000.0
    assert {x["day"]: x["revenue"] for x in sales["series_30d"]}[y] == 2000.0
    row = [x for x in sales["daily"] if x["date"] == y][0]
    assert row["vs_last_week_pct"] == 100.0 and row["per_bill"] == 2000 and row["customers"] == 2
    assert all(x["date"] < today for x in sales["daily"])


# ── the target ───────────────────────────────────────────────────────────────
def _fake_targets(monkeypatch, row):
    import targets
    monkeypatch.setattr(targets, "targets_for_month", lambda con, month, branch=None, group=None: [row] if row else [])


def test_target_absent_is_neutral(monkeypatch):
    from routers import franchise as F
    _fake_targets(monkeypatch, {"source": "none", "revenue_target": 0.0, "basis_month": "2025-06"})
    t = F._target(None, {"pace": {}}, {"mtd": 1000.0, "month_complete": False}, "BranchB", "2026-06")
    assert t["source"] == "none" and t["month_target"] is None and t["vs_to_date_pct"] is None and t["of_month_pct"] is None
    _fake_targets(monkeypatch, None)
    assert F._target(None, {"pace": {}}, {"mtd": 1000.0}, "BranchB", "2026-06")["source"] == "none"


def test_target_present_is_pro_rated_by_the_reports_weekday_weights(monkeypatch):
    from routers import franchise as F
    _fake_targets(monkeypatch, {"source": "last_year", "revenue_target": 31000.0, "basis_month": "2025-06"})
    # open month: the report projects 3100 sold so far to 31000 on its weekday weights, so 10 % of the month has gone
    rep = {"pace": {"proj_method": "dow", "projected": 31000.0, "mtd_revenue": 3100.0, "days_elapsed": 3, "days_total": 30}}
    t = F._target(None, rep, {"mtd": 3100.0, "month_complete": False}, "BranchB", "2026-06")
    assert t["weighting"] == "dow" and t["to_date"] == 3100 and t["vs_to_date_pct"] == 0.0 and t["of_month_pct"] == 10.0
    # no weekday weights: elapsed days over days in the month
    rep = {"pace": {"days_elapsed": 3, "days_total": 30}}
    t = F._target(None, rep, {"mtd": 3100.0, "month_complete": False}, "BranchB", "2026-06")
    assert t["weighting"] == "days" and t["to_date"] == 3100 and t["vs_to_date_pct"] == 0.0
    # a complete month: the full target
    t = F._target(None, {"pace": {}}, {"mtd": 34100.0, "month_complete": True}, "BranchB", "2026-06")
    assert t["weighting"] == "month" and t["to_date"] == 31000 and t["vs_to_date_pct"] == 10.0 and t["of_month_pct"] == 110.0


def test_a_typed_target_reaches_the_page(world):
    import main
    c = sqlite3.connect(str(main.DB_PATH))
    assert _summary(world["b"])["target"]["source"] == "none"            # the test data has no year-ago month
    c.execute("INSERT OR REPLACE INTO branch_targets(branch_id, period_month, revenue_target, customer_target) VALUES(2, '2026-06', 30000, 0)")
    c.commit()
    try:
        t = _summary(world["b"])["target"]
        assert t["source"] == "manual" and t["month_target"] == 30000 and t["to_date"] is not None and 0 < t["to_date"] <= 30000
        assert t["of_month_pct"] == round(100.0 * t["mtd"] / 30000, 1)
        assert _summary(world["a"])["target"]["source"] == "none"        # the other branch has no typed target
    finally:
        c.execute("DELETE FROM branch_targets WHERE branch_id=2 AND period_month='2026-06'")
        c.commit()
        c.close()


# ── the download ─────────────────────────────────────────────────────────────
def _cells(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    yield ws.title, v


def test_the_download_is_the_month_in_four_sheets(world):
    r = world["b"].get("/api/franchise/summary.xlsx", params={"month": "2026-06"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    cd = r.headers["content-disposition"]
    from urllib.parse import quote
    assert "filename*=UTF-8''" in cd and quote("ร้านของฉัน_BranchB_2026-06.xlsx") in cd
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["สรุป", "รายวัน", "12 เดือน", "เมนูขายดี", "วัตถุดิบ"]          # F3: + วัตถุดิบ (test_franchise_f3)
    days = [row[0] for row in wb["รายวัน"].iter_rows(min_row=3, values_only=True)]
    assert days == ["2026-06-01", "2026-06-02", "2026-06-03"]
    d = _summary(world["b"])
    assert len(list(wb["12 เดือน"].iter_rows(min_row=3, values_only=True))) == len(d["trend_12m"]["months"])
    assert len(list(wb["เมนูขายดี"].iter_rows(min_row=3, values_only=True))) == len(d["top_menus"])
    vals = {row[0]: row[1] for row in wb["สรุป"].iter_rows(min_row=3, values_only=True)}
    assert vals["ยอดขาย (ไม่รวม VAT)"] == d["sales"]["mtd"]
    text = [str(v).lower() for _, v in _cells(wb)]
    hits = [t for t in text if any(w in t for w in FORBIDDEN)]
    assert not hits, hits
    assert not any("brancha" in t for t in text)                         # no other branch, no chain figure


def test_the_download_keeps_the_summarys_403_rules(world):
    r = world["b"].get("/api/franchise/summary.xlsx", params={"branch_id": 1, "month": "2026-06"})
    assert r.status_code == 403 and r.headers.get("X-Denied") == "role"
    for q in ({}, {"branch_id": 2}):
        r = world["nobranch"].get("/api/franchise/summary.xlsx", params=q)
        assert r.status_code == 403 and r.headers.get("X-Denied") == "role", q
    assert world["mg"].get("/api/franchise/summary.xlsx").status_code == 400                    # the centre names a branch
    assert world["mg"].get("/api/franchise/summary.xlsx", params={"branch_id": 1, "month": "2026-06"}).status_code == 200
    for m in ("POST", "PUT", "DELETE"):
        r = world["b"].request(m, "/api/franchise/summary.xlsx", headers={"X-Upload-Token": TOKEN})
        assert r.status_code == 403 and r.headers.get("X-Denied") == "role", m


def test_the_new_payload_holds_no_forbidden_word_or_other_branch(world):
    for who, other in (("b", "BranchA"), ("a", "BranchB")):
        d = _summary(world[who])
        part = {"daily": d["sales"]["daily"], "trend": d["trend_12m"], "target": d["target"], "coverage": d["gp"]["coverage"]}
        hits = [(p, kind, s) for p, kind, s in _walk(part) if any(w in s.lower() for w in FORBIDDEN)]
        assert not hits, hits
        assert other not in json.dumps(d, ensure_ascii=False)


@pytest.mark.parametrize("path", ["tabs/56-franchise.html", "js/17-franchise.js"])
def test_the_page_files_hold_no_forbidden_endpoint(path):
    """The page reads /api/franchise/* only: no other API path is named in its files."""
    import re
    from pathlib import Path
    src = (Path(__file__).resolve().parent.parent / path).read_text(encoding="utf-8")
    apis = set(re.findall(r"/api/[a-z0-9_\-/.]+", src))
    assert apis and all(a.startswith("/api/franchise/") for a in apis), apis
