# -*- coding: utf-8 -*-
"""PL1 - polish after the RV1 review (docs: the findings table of 2026-10-23).

The promises:
  #1  a dish costed only by the analyst's imported list is COSTED in the coverage (rules.costed_sql / coverage_sums_sql) and in the costed
      revenue bases (core._bom_costed_rev_expr / _costed_rev_case): coverage and food cost agree with COGS;
  #2  stock mode: the branch's month-close task "ปิดเดือน…: บอกเหตุผลส่วนต่าง" reads the month-end reconcile (the สิ้นเดือน tab), never the
      recipe period, and opens that tab; full mode keeps variance.period;
  #3  the pilot scorecard's reason metric reads the reconcile in stock mode;
  #5  a dish-cost write does not rebuild sku_effective_cost;
  #6  an import / edit says that the fact-table reports follow at the next sync;
  #7  ONE set behind "เฉลี่ยทุกสาขา" (forecast.median_branch_ids);
  #9  a free-text cell that starts with = + - @ tab CR is written with a leading apostrophe;
  #11 the count sheet classes items from whatever slips exist (not only from 14 days);
  #16 a dish-cost month beyond next month is refused, a 0 baht cost needs an explicit confirm;
  #19 a swap carries the triggers of the preserved tables; #21 an empty live answer never wipes the local cost list;
  the stock-mode month close is ONE list of five steps.
"""
import io
import sqlite3
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

import core
import rules
import xl_safe
from bom_center import close_settings as CS
from bom_center import dish_cost as DC
from bom_center import forecast as FC
from bom_center import issue_reports as IR
from bom_center import pilot as PL
from bom_center import todo as TD
from bom_center import variance as VA
from bom_center.clock import BKK
from test_bom_center_s4b import con  # noqa: F401
from test_bom_center_s7 import db  # noqa: F401  (branch 1: counts 8 and 15 Sep 2026, the crab stick 30 % over the recipes)
from test_dish_cost import _file, con as dc_con, web  # noqa: F401
from test_bom_center_s11 import udon  # noqa: F401
from test_slip_variance_nobom import world  # noqa: F401  (stock mode, branches 1, 3, 4 with 28 confirmed slip days)


# ── #1 coverage sees the imported dish cost ──────────────────────────────────
def _sell_month(c, sid, month, qty, rev):
    c.execute("INSERT INTO sales_sku_monthly(branch_id, period_month, sku_id, qty, subtotal_ex_vat, tax_total, total_inc_vat, source)"
              " VALUES(1,?,?,?,?,0,0,'Easy')", (month, sid, qty, rev))


def _cov(c, month):
    r = c.execute(f"SELECT {rules.coverage_sums_sql('ssm')}, SUM({core._bom_costed_rev_expr('ssm')}) AS crev,"
                  f" SUM({core._costed_rev_case('ssm', 'sk', 'sc')}) AS crev2 "
                  "FROM sales_sku_monthly ssm JOIN skus sk ON sk.sku_id = ssm.sku_id WHERE ssm.period_month=?", (month,))
    r.row_factory = None
    row = r.fetchone()
    names = [d[0] for d in r.description]
    return dict(zip(names, row))


def test_a_dish_costed_only_by_the_list_is_costed_in_the_coverage(dc_con):
    c = dc_con
    _sell_month(c, 9002, "2026-09", 40, 400.0)          # DCT-B: no cost anywhere
    _sell_month(c, 9001, "2026-09", 100, 1000.0)        # DCT-A: the old chain (฿10)
    before = _cov(c, "2026-09")
    assert before["cov_plates"] == 140 and before["cov_plates_costed"] == 100
    assert before["cov_rev_costed"] == 1000.0 and before["crev"] == 1000.0 and before["crev2"] == 1000.0
    DC.commit(c, "2026-09", [{"code": "DCT-B", "name": "", "cost": 12, "note": ""}], "analyst1")
    after = _cov(c, "2026-09")
    assert after["cov_plates_costed"] == 140 and after["cov_rev_costed"] == 1400.0         # the list counts, as it does in COGS
    assert after["crev"] == 1400.0 and after["crev2"] == 1400.0                            # the costed-revenue bases agree
    cov = rules.coverage_from_sums(after)
    assert cov["plates_pct"] == 100.0 and cov["uncosted_plates"] == 0
    cogs = c.execute(f"SELECT SUM({rules.cogs_expr('ssm')}) FROM sales_sku_monthly ssm WHERE ssm.sku_id=9002").fetchone()[0]
    assert cogs == 40 * 12                                                                 # COGS and coverage see the same number


def test_a_cost_list_month_after_the_row_month_does_not_cost_it(dc_con):
    c = dc_con
    _sell_month(c, 9002, "2026-09", 40, 400.0)
    DC.commit(c, "2026-10", [{"code": "DCT-B", "name": "", "cost": 12, "note": ""}], "analyst1")      # October's list is not September's cost
    assert _cov(c, "2026-09")["cov_plates_costed"] == 0


# ── #2 the month-close task, #13 the five steps ──────────────────────────────
def _item(c, today):
    d = TD.build(c, "2026-09", branch_id=1, today=today, now=datetime(2026, 9, 25, 8, 0, tzinfo=BKK))
    return next(i for i in d["items"] if TD._base_key(i["key"]) == "variance_open")


def test_stock_mode_month_end_task_reads_the_reconcile_and_opens_its_tab(db, monkeypatch):  # noqa: F811
    c, end = db
    c.execute("INSERT OR REPLACE INTO app_settings(key, value) VALUES('bc_reason_min_baht', '0')")
    c.commit()
    assert CS.mode(c) == "stock"
    # the recipe period is never asked in stock mode
    monkeypatch.setattr(VA, "period", lambda *a, **k: (_ for _ in ()).throw(AssertionError("recipe period read in stock mode")))
    it = _item(c, date(2026, 9, 25))
    rc = IR.reconcile(c, 1, end)
    assert it["status"] == "todo" and it["count"] == rc["totals"]["needs_reason"] + 1          # + the receipts not yet confirmed
    assert it["go"]["pane"] == "var" and it["go"]["extra"] == {"end_count_id": end, "var_tab": "recon"}
    VA.confirm_receipts(c, 1, end, "udon")
    c.commit()
    it = _item(c, date(2026, 9, 25))
    assert it["count"] == (rc["totals"]["needs_reason"] or None) and it["status"] == ("todo" if rc["totals"]["needs_reason"] else "done")


def test_full_mode_month_end_task_keeps_the_recipe_period(db):  # noqa: F811
    c, end = db
    c.execute("INSERT OR REPLACE INTO app_settings(key, value) VALUES('bc_reason_min_baht', '0')")
    CS.set_mode(c, "t", "full")
    c.commit()
    p = VA.period(c, 1, summary_only=True)
    it = _item(c, date(2026, 9, 25))
    assert it["count"] == p["needs_reason"] + 1 and "var_tab" not in it["go"]["extra"]


def test_stock_close_is_one_list_of_five_steps(world):  # noqa: F811
    c = world
    now = datetime(2026, 9, 28, 8, 0, tzinfo=BKK)
    d = TD.build(c, "2026-09", None, None, today=date(2026, 9, 28), now=now)
    steps = d["close"]["steps"]
    assert [s["key"] for s in steps] == ["central_report", "code_links", "counts", "receipts_confirm", "issue_posted"]
    assert [s["seq"] for s in steps] == [1, 2, 3, 4, 5] and d["close"]["total"] == 5
    assert [p["key"] for p in steps[0]["parts"]] == ["central_report", "deliveries"]
    assert [p["key"] for p in steps[2]["parts"]] == ["begin_count", "end_count"]
    assert d["close"]["done"] == sum(1 for s in steps if s["status"] == "done")
    # the full mode still lists nine
    CS.set_mode(c, "t", "full")
    full = TD.build(c, "2026-09", None, None, today=date(2026, 9, 28), now=now)["close"]
    assert full["total"] == 9 and [s["seq"] for s in full["steps"]] == list(range(1, 10))


def test_merged_status_is_the_worst_of_its_parts():
    assert TD._worst(["done", "done"]) == "done" and TD._worst(["done", "na"]) == "done"
    assert TD._worst(["na", "na"]) == "na" and TD._worst(["done", "todo"]) == "todo"
    assert TD._worst(["blocked", "done"]) == "blocked" and TD._worst(["blocked", "todo"]) == "todo"


# ── #3 the pilot scorecard ───────────────────────────────────────────────────
def test_pilot_reason_metric_reads_the_reconcile_in_stock_mode(db, monkeypatch):  # noqa: F811
    c, end = db
    monkeypatch.setattr(VA, "period", lambda *a, **k: (_ for _ in ()).throw(AssertionError("recipe period read in stock mode")))
    rows = [{"ask_reason": True, "reason": "waste"}, {"ask_reason": True, "reason": "waste"}, {"ask_reason": True},
            {"ask_reason": False}]
    monkeypatch.setattr(IR, "reconcile", lambda con, bid, eid: {"rows": rows})
    s = PL.scorecard(c, 1, weeks=4, as_of=date(2026, 9, 20))
    assert s["flagged"] == 3 and s["metrics"]["reason_coverage"] == pytest.approx(66.7)


# ── #5 #6 dish cost: no rebuild, the sync line ───────────────────────────────
def test_a_dish_cost_flush_does_not_rebuild_the_effective_cost(monkeypatch):
    calls = []
    monkeypatch.setattr(core, "_build_sku_effective_cost", lambda con, force=False: calls.append(force))
    core.invalidate_cache(rebuild_cost=False)
    assert calls == []
    core.invalidate_cache()
    assert calls == [True]


def test_the_note_names_the_next_sync():
    t = datetime(2026, 10, 23, 13, 0, tzinfo=BKK)
    assert DC.facts_note(t) == "รายงานเมนูและรายงานรายบิลใช้ต้นทุนนี้หลังการซิงก์รอบถัดไป เวลา 16:00 น. (วันนี้)"
    t = datetime(2026, 10, 23, 17, 0, tzinfo=BKK)
    assert "06:00 น. (พรุ่งนี้)" in DC.facts_note(t)


def test_http_dish_cost_writes_flush_without_rebuild_and_carry_the_sync_line(client, web, monkeypatch):
    import routers.bom_center as rb
    monkeypatch.delenv("OSHINEI_AUTH", raising=False)
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "admin:tok")   # the machine endpoints take the owner ("admin") token only
    monkeypatch.delenv("OSHINEI_WRITE_TOKEN", raising=False)
    calls = []
    monkeypatch.setattr(rb, "invalidate_cache", lambda *a, **k: calls.append((a, k)))
    h = {"X-Upload-Token": "tok"}
    r = client.post("/api/bom-center/dish-cost/import?month=2026-10&dry_run=false", files=_file([["DCW-1", "x", 15, ""]]), headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["facts_note"].startswith("รายงานเมนูและรายงานรายบิลใช้ต้นทุนนี้หลังการซิงก์รอบถัดไป")
    assert calls == [((), {"rebuild_cost": False})]
    r = client.post("/api/bom-center/dish-cost/edit", json={"month": "2026-10", "sku_code": "DCW-1", "unit_cost": 16}, headers=h)
    assert r.status_code == 200 and "facts_note" in r.json() and calls[-1] == ((), {"rebuild_cost": False})
    # 0 baht: refused without the confirm, written with it
    r = client.post("/api/bom-center/dish-cost/edit", json={"month": "2026-10", "sku_code": "DCW-1", "unit_cost": 0}, headers=h)
    assert r.status_code == 400 and "ยืนยันอีกครั้ง" in r.json()["detail"]
    r = client.post("/api/bom-center/dish-cost/edit", json={"month": "2026-10", "sku_code": "DCW-1", "unit_cost": 0, "confirm_zero": True}, headers=h)
    assert r.status_code == 200


# ── #16 bounds and zero ──────────────────────────────────────────────────────
def test_dish_cost_month_stops_at_next_month(dc_con):
    today = date(2026, 10, 23)
    assert DC.check_month("2026-11", today) == "2026-11"
    with pytest.raises(ValueError, match="ไม่เกินเดือน 2026-11"):
        DC.check_month("2062-10", today)
    assert DC.max_month(date(2026, 12, 5)) == "2027-01"


def test_zero_cost_needs_an_explicit_confirm(dc_con):
    c = dc_con
    rows = [{"code": "DCT-B", "name": "", "cost": 0, "note": ""}, {"code": "DCT-A", "name": "", "cost": 11, "note": ""}]
    pre = DC.preview(c, "2026-10", rows)
    # audit 2026-10: an import never writes a 0 cost - the row is a per-row error in the preview, the other rows go in
    assert pre["zero_costs"] == [] and [b["code"] for b in pre["bad_values"]] == ["DCT-B"] and pre["imported"] == 1
    DC.commit(c, "2026-10", rows, "analyst1", confirm_zero=True)
    assert c.execute("SELECT COUNT(*) FROM bc_sku_cost_list WHERE sku_id=9002").fetchone()[0] == 0
    with pytest.raises(ValueError, match="ยืนยันอีกครั้ง"):
        DC.set_one(c, "2026-10", "DCT-D", 0, "", "analyst1")
    DC.set_one(c, "2026-10", "DCT-D", 0, "", "analyst1", confirm_zero=True)


def test_the_import_day_is_the_bangkok_day():
    assert DC.bkk_day("2026-10-22T20:30:00") == "2026-10-23"          # 03:30 in Bangkok
    assert DC.bkk_day("2026-10-22T10:00:00") == "2026-10-22" and DC.bkk_day(None) is None


# ── #7 one median set ────────────────────────────────────────────────────────
def test_one_set_of_branches_behind_the_all_branch_figure(world):  # noqa: F811
    c = world
    from bom_center import benchmark as BM
    company = FC.company_branch_ids(c)
    assert set(company) >= {1, 3, 4}
    BM.set_benchmark(c, "t", ids=[1, 3])                         # the ค่ากลาง page chooses two branches
    c.commit()
    assert FC.median_branch_ids(c) == company                    # stock mode: every company branch, whatever the page says
    assert FC.median_branch_ids(c, exclude=1) == [b for b in company if b != 1]
    CS.set_mode(c, "t", "full")
    assert FC.median_branch_ids(c) == [1, 3] and FC.median_branch_ids(c, exclude=1) == [3]


def test_the_per_100k_report_benchmark_follows_the_company_set_in_stock_mode(world):  # noqa: F811
    from bom_center import benchmark as BM
    from bom_center import per_sales as PS
    c = world
    BM.set_benchmark(c, "t", ids=[1, 3])
    c.commit()
    PS.clear_cache()
    r = PS.report(c, weeks=2, today=date(2026, 10, 6))
    bench = {b["branch_id"] for b in (r["benchmark"] or {}).get("branches", [])}
    assert bench == {1, 3, 4}                                    # branch 4 is outside the settings page and inside the one definition
    assert all(b["is_benchmark"] for b in r["branches"] if b["branch_id"] in (1, 3, 4))


# ── #9 Excel formula injection ───────────────────────────────────────────────
@pytest.mark.parametrize("v", ["=HYPERLINK(\"http://x\")", "+1+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1"])
def test_safe_cell_prefixes_an_apostrophe(v):
    assert xl_safe.safe_cell(v) == "'" + v
    assert xl_safe.safe_cell("ธรรมดา") == "ธรรมดา" and xl_safe.safe_cell(5) == 5 and xl_safe.safe_cell("-") == "-"
    assert xl_safe.safe_cell("") == "" and xl_safe.safe_cell(None) is None


def test_the_workbook_response_guards_every_text_cell():
    import routers.bom_center as rb
    wb = Workbook()
    ws = wb.active
    ws.append(["หมายเหตุ", "=HYPERLINK(\"http://evil\",\"x\")", "+cmd", "-cmd", "@x", 12, "ปกติ"])
    ws2 = wb.create_sheet("อีกแผ่น")
    ws2["B2"] = "=1+1"
    resp = rb._xlsx(wb, "ทดสอบ.xlsx")
    data = b"".join(resp.body_iterator) if hasattr(resp, "body_iterator") and not hasattr(resp.body_iterator, "__aiter__") else None
    if data is None:
        import asyncio

        async def _read():
            return b"".join([chunk async for chunk in resp.body_iterator])
        data = asyncio.run(_read())
    out = load_workbook(io.BytesIO(data))
    row = [c.value for c in out.worksheets[0][1]]
    assert row == ["หมายเหตุ", "'=HYPERLINK(\"http://evil\",\"x\")", "'+cmd", "'-cmd", "'@x", 12, "ปกติ"]
    assert out.worksheets[1]["B2"].value == "'=1+1" and out.worksheets[1]["B2"].data_type == "s"


# ── #11 the count sheet reads whatever slips exist ───────────────────────────
def test_the_count_sheet_classes_items_from_fewer_than_14_slip_days(world):  # noqa: F811
    from bom_center import stock as ST
    c = world
    c.execute("DELETE FROM bc_issue_line WHERE slip_id IN (SELECT slip_id FROM bc_issue_slip WHERE branch_id=1 AND slip_date > '2026-09-05')")
    c.execute("DELETE FROM bc_issue_slip WHERE branch_id=1 AND slip_date > '2026-09-05'")
    c.commit()
    assert FC.demand_source(c, 1, date(2026, 9, 1), date(2026, 9, 30), "2026-09")[2]["source"] == "none"      # under 14 days: no suggestion
    per_day, recent = ST._usage(c, 1, date(2026, 9, 30), with_recent=True)
    assert per_day and per_day.get("13-01-009", 0) > 0                  # ... but the sheet still knows what the branch uses


# ── #19 triggers survive a swap, #21 an empty live list never wipes ──────────
def test_a_swap_carries_the_triggers_of_a_preserved_table(tmp_path):
    live, inc = tmp_path / "live.db", tmp_path / "inc.db"
    c = sqlite3.connect(str(live))
    c.execute("CREATE TABLE franchise_msg_message(id INTEGER PRIMARY KEY, body TEXT)")
    c.execute("CREATE TRIGGER trg_fmsg_message_no_update BEFORE UPDATE ON franchise_msg_message BEGIN SELECT RAISE(ABORT, 'immutable'); END")
    c.execute("INSERT INTO franchise_msg_message(body) VALUES('hello')")
    c.commit()
    c.close()
    sqlite3.connect(str(inc)).close()
    core._carry_forward_preserved(str(live), str(inc))
    d = sqlite3.connect(str(inc))
    assert d.execute("SELECT body FROM franchise_msg_message").fetchone() == ("hello",)
    with pytest.raises(sqlite3.DatabaseError, match="immutable"):
        d.execute("UPDATE franchise_msg_message SET body='x'")
    d.close()


def test_an_empty_live_answer_keeps_the_local_cost_list(tmp_path, monkeypatch):
    import json
    from bom_center import dated as D
    sq = sqlite3.connect(str(tmp_path / "staging.db"))
    sq.executescript("CREATE TABLE skus(sku_id INTEGER PRIMARY KEY, sku_code TEXT);INSERT INTO skus VALUES (77, 'DCT-A');")
    D.ensure(sq)
    sq.execute("INSERT INTO bc_sku_cost_list(sku_id, month, unit_cost, note, imported_by, imported_at) VALUES(77,'2026-10',11,NULL,'a','x')")
    sq.commit()
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "admin:tok")   # the machine endpoints take the owner ("admin") token only
    monkeypatch.delenv("OSHINEI_SYNC_ON_LIVE", raising=False)

    class _Resp:
        def __init__(self, body): self.body = body
        def read(self): return self.body
        def __enter__(self): return self
        def __exit__(self, *a): return False

    monkeypatch.setattr("urllib.request.urlopen", lambda req, timeout=0: _Resp(json.dumps({"rows": []}).encode()))
    D.pull_from_live(sq, api="http://x")
    assert sq.execute("SELECT COUNT(*) FROM bc_sku_cost_list").fetchone()[0] == 1          # an empty answer is not "the list was emptied"


# ── #13 / #14 the words a branch reads ───────────────────────────────────────
def test_the_branch_fallback_message_uses_the_glossary_word():
    from pathlib import Path
    src = (Path(__file__).parent / "bom_center" / "forecast.py").read_text(encoding="utf-8")
    assert 'f"จำนวนที่ควรสั่งจากเฉลี่ยทุกสาขา {k} สาขา' in src and "จำนวนที่ควรสั่งจากค่ากลาง" not in src


def test_the_comparator_phrase_has_one_home():
    from pathlib import Path
    from bom_center import fmt
    assert fmt.NORMAL_OR_MEDIAN_TH == "ที่สาขาใช้ปกติหรือเฉลี่ยทุกสาขา" and fmt.comparator_th(False) == "สูตร"
    root = Path(__file__).parent
    for rel in ("bom_center/todo.py", "bom_center/reminders.py", "bom_center/issue_reports.py", "routers/franchise.py", "routers/bom_center.py"):
        assert "ที่สาขาใช้ปกติหรือเฉลี่ยทุกสาขา" not in (root / rel).read_text(encoding="utf-8"), rel
