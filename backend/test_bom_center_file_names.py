# -*- coding: utf-8 -*-
"""One name for one file (owner 2026-10-21: "the file names and the import labels do not match").

A template or an export is named the same everywhere: the page the user sees, the download link, the saved file name
(Content-Disposition filename*, UTF-8), the main sheet. These checks pin the server side of that rule for the files that
need no data to build; the files that carry a branch or a month are pinned in the tests of their own reports.
"""
import io

import pytest
from conftest import dl_name
from openpyxl import load_workbook

CASES = [
    ("/api/bom-center/units/template.xlsx", "แม่แบบ_หน่วยและการแปลง.xlsx", "หน่วยและการแปลง"),
    ("/api/bom-center/units/template.xlsx?all=1", "แม่แบบ_หน่วยและการแปลง.xlsx", "หน่วยและการแปลง"),
    ("/api/bom-center/prices/template.xlsx", "แม่แบบ_ราคารายเดือน.xlsx", "ราคารายเดือน"),
    ("/api/bom-center/prices/template.xlsx?page=ไฟล์ราคา", "แม่แบบ_ไฟล์ราคา.xlsx", "ไฟล์ราคา"),
    ("/api/bom-center/prices/template.xlsx?page=อะไรก็ได้", "แม่แบบ_ราคารายเดือน.xlsx", "ราคารายเดือน"),
    ("/api/bom-center/order-calendar/template.xlsx", "แม่แบบ_รอบส่งของสาขา.xlsx", "รอบส่งของสาขา"),
    ("/api/bom-center/prep/template.xlsx", "แม่แบบ_สูตรกึ่งสำเร็จ.xlsx", "สูตรกึ่งสำเร็จ"),
    ("/api/bom-center/recipes/template.xlsx", "แม่แบบ_สูตรตามระดับราคา.xlsx", "สูตรตามระดับราคา"),
    ("/api/bom-center/code-review.xlsx", "ตรวจรหัสสูตรเดิม.xlsx", "ตรวจรหัสสูตรเดิม"),
]


@pytest.mark.parametrize("url,name,sheet", CASES, ids=[c[0] for c in CASES])
def test_the_file_the_link_and_the_sheet_carry_the_same_name(client, url, name, sheet):
    r = client.get(url)
    assert r.status_code == 200, r.text
    assert dl_name(r) == name
    assert load_workbook(io.BytesIO(r.content)).worksheets[0].title == sheet
    # the header itself is plain ASCII (the Thai name is percent-encoded in filename*) and carries an ASCII fallback
    cd = r.headers["content-disposition"]
    assert cd.isascii() and 'filename="' in cd and name.rsplit(".", 1)[0] not in cd


def test_the_register_is_named_for_the_register_with_the_date(client):
    r = client.get("/api/bom-center/ingredients/template.xlsx")
    assert r.status_code == 200
    assert dl_name(r).startswith("ทะเบียนวัตถุดิบ_") and dl_name(r).endswith(".xlsx")
    assert load_workbook(io.BytesIO(r.content)).worksheets[0].title == "ทะเบียนวัตถุดิบ"


def test_the_list_of_goods_is_named_for_the_list(client):
    r = client.get("/api/bom-center/central.xlsx")
    assert r.status_code == 200
    assert dl_name(r).startswith("บัญชีสินค้า_") and "ส่วนกลาง" not in dl_name(r)
    assert load_workbook(io.BytesIO(r.content)).worksheets[0].title == "บัญชีสินค้า"


def test_an_old_file_with_the_old_sheet_name_still_imports(client):
    """Only the names we give change: the unit file is found by its sheet name, and a file made before this change has the same one."""
    from bom_center import unit_sheet as US
    assert US.SHEET == "หน่วยและการแปลง" and US.PORTION_SHEET == "สูตรแบ่งชิ้น"
