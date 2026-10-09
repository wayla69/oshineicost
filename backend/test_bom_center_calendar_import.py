# -*- coding: utf-8 -*-
"""Bulk delivery days from a file (PR-I, 30-branch onboarding): the template lists every active branch with
its current days, a dry run previews and writes nothing, a save writes the changed rounds through the same
path as the page save and audits them. Sales Support's; no other role."""
import io
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

import auth
from bom_center import forecast as FC

HDR = ["สาขา", "รอบ 1 วันส่ง", "รอบ 2 วันส่ง", "แซลมอนสด วันส่ง"]
TOK = {"X-Upload-Token": "t"}
URL = "/api/bom-center/order-calendar/import"


@pytest.fixture(autouse=True)
def _clean(client, monkeypatch):
    import main
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "t")

    def wipe():
        c = sqlite3.connect(str(main.DB_PATH))
        try:
            c.execute("DELETE FROM bc_order_calendar")
            c.execute("DELETE FROM bc_audit WHERE action IN ('order_calendar','order_calendar_import')")
            c.commit()
        except sqlite3.OperationalError:               # the BoM Center tables appear with the first request
            pass
        finally:
            c.close()
    wipe()
    yield
    wipe()


def _book(rows, header=HDR, note="ตั้งวันส่ง (ตัวอย่างหมายเหตุ)"):
    wb = Workbook()
    ws = wb.active
    ws.append([note])
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, data, dry=True, headers=None, name="cal.xlsx"):
    return client.post(f"{URL}?dry_run={'true' if dry else 'false'}", headers=headers or {},
                       files={"file": (name, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def _rows(client):
    return {r["branch"]: r for r in client.get("/api/bom-center/order-calendar").json()["rows"]}


def _audit(action):
    import main
    c = sqlite3.connect(str(main.DB_PATH))
    try:
        return c.execute("SELECT who, row_key, old_value, new_value FROM bc_audit WHERE action=? ORDER BY id", (action,)).fetchall()
    finally:
        c.close()


def test_template_lists_every_active_branch_with_the_current_days(client):
    names = list(_rows(client))
    r = client.post("/api/bom-center/order-calendar", headers=TOK,
                    json=[{"branch_id": _rows(client)[names[0]]["branch_id"], "round": "2", "delivery_weekday": 5}])
    assert r.status_code == 200, r.text
    t = client.get("/api/bom-center/order-calendar/template.xlsx")
    assert t.status_code == 200 and "spreadsheetml" in t.headers["content-type"]
    ws = load_workbook(io.BytesIO(t.content)).worksheets[0]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    assert "13:15" in grid[0][0] and not grid[0][0].startswith("สาขา")        # the note row explains the cutoff rule
    assert grid[1] == HDR
    body = {r[0]: r[1:] for r in grid[2:] if r[0]}
    assert set(body) == set(names)                                              # every active branch, nothing else
    assert body[names[0]] == ["พุธ", "ศุกร์", "จันทร์"]                           # round 2 was saved as Friday; the rest are defaults
    assert body[names[1]] == [FC.WD_TH[3], FC.WD_TH[6], FC.WD_TH[1]]


def test_dry_run_previews_and_writes_nothing(client):
    names = list(_rows(client))
    a, b = names[0], names[1]
    data = _book([[a, "อังคาร", "", "จันทร์"], [b, "พุธ", "เสาร์", None]])        # a: round 1 changes; b: all as the defaults
    r = _post(client, data, dry=True)                                            # no write token on a preview
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["dry_run"] is True and j["errors"] == [] and j["saved"] == 0
    assert j["branches"] == 2 and j["changed_branches"] == 1 and j["changed_rounds"] == 1
    ra = next(x for x in j["rows"] if x["branch"] == a)
    assert ra["change"] is True
    assert ra["r1"]["changed"] and ra["r1"]["was_th"] == "พุธ" and ra["r1"]["now_th"] == "อังคาร"
    assert ra["r1"]["cutoff_th"] == "วันเสาร์ 13:15"                              # Tuesday delivery → Saturday cutoff
    assert not ra["r2"]["changed"] and ra["r2"]["now"] == FC.DEFAULT_CAL["2"]["delivery_weekday"]
    rb = next(x for x in j["rows"] if x["branch"] == b)
    assert rb["change"] is False
    assert not any(v["r1"]["filled"] for v in _rows(client).values())            # nothing was written
    assert _audit("order_calendar") == [] and _audit("order_calendar_import") == []


def test_save_writes_the_changed_rounds_and_audits(client):
    rows = _rows(client)
    a, b = list(rows)[:2]
    data = _book([[a, "อังคาร", "พฤหัสบดี", "ศุกร์"], [b, "พุธ", "เสาร์", "จันทร์"]])
    assert _post(client, data, dry=False).status_code == 403                    # a real save needs the write token
    r = _post(client, data, dry=False, headers=TOK)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["saved"] == 3 and j["changed_branches"] == 1                       # a: three rounds; b: nothing differs
    cal = _rows(client)
    assert cal[a]["r1"]["delivery_weekday"] == 2 and cal[a]["r1"]["cutoff_th"] == "วันเสาร์ 13:15"
    assert cal[a]["r2"]["delivery_weekday"] == 4 and cal[a]["rsalmon"]["delivery_weekday"] == 5
    assert cal[a]["r1"]["filled"] and not cal[b]["r1"]["filled"]                # an unchanged branch stays a default
    log = _audit("order_calendar")
    assert len(log) == 3 and {k for _, k, _, _ in log} == {f"{rows[a]['branch_id']}|{x}" for x in ("1", "2", "salmon")}
    assert any(old == "delivery 3" and new.startswith("delivery 2 / cutoff 6") for _, _, old, new in log)
    imp = _audit("order_calendar_import")
    assert len(imp) == 1 and imp[0][1] == "cal.xlsx" and "1 branches / 3 rounds" in imp[0][3]
    again = _post(client, data, dry=True).json()                                # idempotent: nothing left to change
    assert again["changed_branches"] == 0


@pytest.mark.parametrize("cell,want", [
    ("อาทิตย์", 0), ("วันจันทร์", 1), ("อังคาร", 2), ("พุธ", 3), ("พฤหัส", 4), ("พฤหัสบดี", 4), ("ศุกร์", 5), ("เสาร์", 6),
    ("อา", 0), ("พฤ", 4), ("Fri", 5), (" sat ", 6), ("3", 3), (3, 3), (3.0, 3), (0, 0), ("0", 0), (6, 6),
])
def test_weekday_parsing(cell, want):
    assert FC.parse_weekday(cell) == want


@pytest.mark.parametrize("cell", ["วันที่ 7", 7, "7", -1, 2.5, "พรุ่งนี้", True])
def test_weekday_parsing_refuses_what_it_cannot_read(cell):
    with pytest.raises(ValueError):
        FC.parse_weekday(cell)


def test_blank_weekday_cells_mean_no_change():
    assert FC.parse_weekday(None) is None and FC.parse_weekday("  ") is None


def test_numbers_and_thai_names_in_one_file(client):
    a = list(_rows(client))[0]
    j = _post(client, _book([[a, 2, "6", "วันเสาร์"]]), dry=True).json()
    ra = j["rows"][0]
    assert [ra[k]["now"] for k in ("r1", "r2", "rsalmon")] == [2, 6, 6]
    assert j["errors"] == [] and ra["r2"]["changed"] is False and ra["rsalmon"]["changed"] is True


def test_unknown_branch_and_bad_cells_are_errors_and_block_the_save(client):
    a = list(_rows(client))[0]
    data = _book([["สาขาที่ไม่มีอยู่จริง", "พุธ", "เสาร์", "จันทร์"], [a, "พรุ่งนี้", "เสาร์", "จันทร์"], [a, "พุธ", "", ""]])
    j = _post(client, data, dry=True).json()
    why = {e["line"]: e["why"] for e in j["errors"]}
    assert "ไม่พบสาขา" in why[1] and "ไม่ถูกต้อง" in why[2]
    assert any(e["line"] == 3 and "ซ้ำ" in e["why"] for e in j["errors"])        # the same branch twice
    assert j["changed_branches"] == 0
    r = _post(client, data, dry=False, headers=TOK)
    assert r.status_code == 400 and "ไม่ถูกต้อง" in r.json()["detail"]
    assert not any(v["r1"]["filled"] for v in _rows(client).values())            # the whole file was refused


def test_an_inactive_branch_is_reported_not_written(client, app_db):
    import main
    a = list(_rows(client))[0]
    c = sqlite3.connect(str(main.DB_PATH))
    bid = _rows(client)[a]["branch_id"]
    c.execute("UPDATE branches SET active=0 WHERE branch_id=?", (bid,))
    c.commit()
    try:
        j = _post(client, _book([[a, "อังคาร", "", ""]]), dry=True).json()
        assert j["rows"] == [] and "ปิดใช้งาน" in j["errors"][0]["why"]
    finally:
        c.execute("UPDATE branches SET active=1 WHERE branch_id=?", (bid,))
        c.commit()
        c.close()


def test_a_branch_prefix_and_csv_are_accepted(client):
    a = list(_rows(client))[0]
    csv = ("สาขา,รอบ 1 วันส่ง,รอบ 2 วันส่ง,แซลมอนสด วันส่ง\nสาขา" + a + ",ศุกร์,,\n").encode("utf-8-sig")
    j = _post(client, csv, dry=True, name="cal.csv").json()
    assert j["errors"] == [] and j["rows"][0]["r1"]["now"] == 5 and j["changed_rounds"] == 1


def test_a_file_without_the_columns_is_refused(client):
    r = _post(client, _book([["ก", "ข"]], header=["สินค้า", "จำนวน"]), dry=True)
    assert r.status_code == 400
    r = _post(client, _book([["x"]], header=["สาขา"]), dry=True)
    assert r.status_code == 400 and "วันส่ง" in r.json()["detail"]


def test_path_rules():
    ss = auth.ROLE_MODULES["salessupport"]
    assert auth.path_allowed("POST", URL, ss) and auth.path_allowed("GET", "/api/bom-center/order-calendar/template.xlsx", ss)
    for role in ("branch", "warehouse", "innovation"):
        mods = auth.ROLE_MODULES[role]
        assert not auth.path_allowed("POST", URL, mods), role
    for role, mods in auth.ROLE_MODULES.items():              # the template is exactly as readable as the calendar it lists
        assert auth.path_allowed("GET", "/api/bom-center/order-calendar/template.xlsx", mods) == \
            auth.path_allowed("GET", "/api/bom-center/order-calendar", mods), role
