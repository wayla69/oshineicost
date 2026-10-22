"""BoM Center — PR-Y: a new code for an existing item is an ALIAS; the recipes keep the register code.

The promises:
  * the code the recipes use (bc_ingredient.code) never changes: a code from the accounting system, the warehouse or a supplier
    becomes an alias of the item, so a recipe needs no edit and the usage the forecast counts does not move;
  * Sales Support reports the new code (a request); the analyst applies it (the alias is active and becomes the code to use for
    that system, the one in use before stops being current) or rejects it with a reason; nothing in the register changes without
    the analyst - an import run by someone who may not decide files a request instead;
  * an alias is retired (active=0), never deleted, and a retired alias no longer resolves;
  * an import never adds an unknown code on its own when the screen sends its answers: it lists the code with the register items
    it could be a new code of (price file, central report, register file, unit file) and the person chooses per code;
  * the files that go to the accounting system show the code that system uses now, with the register code in a last column.
"""
import io
import json
import sqlite3
from datetime import date

import pytest
from openpyxl import Workbook

import auth
from bom_center import aliases as AL
from bom_center import central as CE
from bom_center import forecast as F
from bom_center import master as M
from bom_center import prices as P
from bom_center import todo as TD
from test_bom_center_s11 import ITEMS, _report, udon  # noqa: F401
from test_bom_center_s4b import AS_OF, con  # noqa: F401

CRAB, SALMON = "13-01-009", "02-01-001"
TODAY = date(2026, 10, 12)


@pytest.fixture
def ac(con):  # noqa: F811
    """The S4b fixture (crab 45 g and salmon 20 g per plate, five weeks of sales) with a pack size on the crab."""
    con.execute("UPDATE bc_ingredient SET recv_unit='แพ็ค 1 กก.', qty_per_recv=1000, moq=1 WHERE code=?", (CRAB,))
    con.commit()
    return con


def _alias_row(c, code):
    return c.execute("SELECT ingredient_code, system, valid_from, COALESCE(active,1), COALESCE(is_current,0), created_by"
                     " FROM bc_ingredient_alias WHERE alias_code=?", (code,)).fetchone()


# ── 1 · schema ─────────────────────────────────────────────────────────────────────────────────────
def test_the_new_columns_and_table_are_owned_and_the_migrations_are_idempotent(tmp_path):
    from bom_center import schema as S
    assert "bc_code_request" in S.OWNED_TABLES
    c = sqlite3.connect(str(tmp_path / "m.db"))
    S.ensure(c, commit=True)
    # a row written before PR-Y has none of the new columns: it must read as active, not current
    c.execute("CREATE TABLE old_al AS SELECT * FROM bc_ingredient_alias WHERE 0")
    S.ensure(c, commit=True)
    S.ensure(c, commit=True)
    cols = lambda t: {r[1] for r in c.execute(f"PRAGMA table_info({t})")}  # noqa: E731
    assert {"valid_from", "active", "is_current", "created_by", "created_at"} <= cols("bc_ingredient_alias")
    assert {"request_id", "old_code", "new_code", "system", "ingredient_code", "status", "decision_note"} <= cols("bc_code_request")
    c.execute("INSERT INTO bc_ingredient_alias(alias_code, ingredient_code, system) VALUES('LEGACY-1','X-1','warehouse')")
    assert _alias_row(c, "LEGACY-1")[3:5] == (1, 0)
    assert AL.active_map(c) == {"LEGACY-1": "X-1"}
    c.close()


# ── 2 · asking for a new code ──────────────────────────────────────────────────────────────────────
def test_a_request_is_checked_and_changes_nothing_in_the_register(ac):
    out = AL.request(ac, "NEW-CRAB", "acc", CRAB, "ss", name="ปูอัดแท่งใหม่", valid_from="2026-11-01", note="บัญชีเปลี่ยนรหัส")
    ac.commit()
    assert out["request_id"] and not out["duplicate"]
    r = ac.execute("SELECT old_code, new_code, system, ingredient_code, valid_from, status, requested_by FROM bc_code_request").fetchone()
    assert r == (CRAB, "NEW-CRAB", "acc", CRAB, "2026-11-01", "pending", "ss")           # the old code defaults to the one in use
    assert ac.execute("SELECT COUNT(*) FROM bc_ingredient_alias").fetchone()[0] == 0     # the register is untouched
    assert ac.execute("SELECT COUNT(*) FROM bc_ingredient").fetchone()[0] == 4
    for kw, msg in (({"new_code": ""}, "ระบุรหัสใหม่"),
                    ({"new_code": SALMON}, "เป็นรหัสของ แซลมอน ในทะเบียนอยู่แล้ว"),         # a register code of another item
                    ({"new_code": CRAB}, "ต้องไม่ซ้ำกับรหัสในทะเบียนของรายการนี้"),
                    ({"system": "x"}, "ระบบต้องเป็น"),
                    ({"valid_from": "soon"}, "วันที่เริ่มใช้ไม่ถูกต้อง"),
                    ({"old_code": "WRONG-9"}, "ไม่ใช่รหัสของ"),
                    ({"ingredient_code": "NOPE"}, "ไม่พบรหัส NOPE ในทะเบียน"),
                    ({"ingredient_code": ""}, "เลือกวัตถุดิบ")):
        args = {"new_code": "ANOTHER-1", "system": "acc", "ingredient_code": CRAB, **kw}
        with pytest.raises(ValueError) as e:
            AL.request(ac, args.pop("new_code"), args.pop("system"), args.pop("ingredient_code"), "ss", **args)
        assert msg in str(e.value), (kw, str(e.value))


def test_the_same_request_twice_is_one_request_and_a_clash_with_another_item_is_refused(ac):
    a = AL.request(ac, "NEW-CRAB", "acc", CRAB, "ss")
    b = AL.request(ac, "NEW-CRAB", "acc", CRAB, "ss2")
    assert b["duplicate"] and b["request_id"] == a["request_id"]
    with pytest.raises(ValueError, match="รออยู่แล้ว"):
        AL.request(ac, "NEW-CRAB", "acc", SALMON, "ss")
    # a code that is already an alias of another item is refused too
    AL.add_alias(ac, SALMON, "SAL-ACC", "acc", None, None, "ana")
    with pytest.raises(ValueError, match="เป็นรหัสในระบบอื่นของ แซลมอน อยู่แล้ว"):
        AL.request(ac, "SAL-ACC", "acc", CRAB, "ss")
    with pytest.raises(ValueError, match="เป็นรหัสในระบบอื่นของรายการนี้อยู่แล้ว"):
        AL.request(ac, "SAL-ACC", "acc", SALMON, "ss")


# ── 3 · the analyst decides ────────────────────────────────────────────────────────────────────────
def test_apply_records_the_alias_makes_it_current_and_switches_the_one_before(ac):
    r1 = AL.request(ac, "CRAB-A1", "acc", CRAB, "ss", valid_from="2026-10-01")["request_id"]
    out = AL.decide(ac, r1, "applied", None, "ana")
    ac.commit()
    assert out["alias"] == "CRAB-A1" and "สูตรไม่ต้องแก้" in out["message"]
    assert _alias_row(ac, "CRAB-A1") == (CRAB, "acc", "2026-10-01", 1, 1, "ss")           # created by the requester, current
    assert AL.current_code(ac, CRAB, "acc", on="2026-10-12") == "CRAB-A1"
    assert AL.current_code(ac, CRAB, "warehouse") == CRAB                                 # another system keeps the register code
    assert ac.execute("SELECT status, decided_by FROM bc_code_request WHERE request_id=?", (r1,)).fetchone() == ("applied", "ana")
    assert ac.execute("SELECT 1 FROM bc_audit WHERE action='code_request_apply' AND row_key=?", (str(r1),)).fetchone()
    assert ac.execute("SELECT 1 FROM bc_audit WHERE action='code_alias_add' AND row_key='CRAB-A1'").fetchone()
    with pytest.raises(ValueError, match="ตัดสินใจไปแล้ว"):
        AL.decide(ac, r1, "applied", None, "ana")
    # a second change in the same system: the first stops being current, but stays an alias
    r2 = AL.request(ac, "CRAB-A2", "acc", CRAB, "ss", old_code="CRAB-A1")["request_id"]
    AL.decide(ac, r2, "applied", None, "ana")
    ac.commit()
    assert _alias_row(ac, "CRAB-A1")[3:5] == (1, 0) and _alias_row(ac, "CRAB-A2")[3:5] == (1, 1)
    assert AL.current_code(ac, CRAB, "acc") == "CRAB-A2" and AL.current_map(ac, "acc") == {CRAB: "CRAB-A2"}
    assert AL.active_map(ac) == {"CRAB-A1": CRAB, "CRAB-A2": CRAB}                        # both still resolve to the item
    # the warehouse's new code is current for the warehouse only
    r3 = AL.request(ac, "CRAB-W1", "warehouse", CRAB, "ss")["request_id"]
    AL.decide(ac, r3, "applied", None, "ana")
    assert AL.current_code(ac, CRAB, "acc") == "CRAB-A2" and AL.current_code(ac, CRAB, "warehouse") == "CRAB-W1"
    d = AL.aliases_of(ac, CRAB)
    assert {x["alias"]: (x["system"], x["in_use"]) for x in d["aliases"]} == {
        "CRAB-A1": ("acc", False), "CRAB-A2": ("acc", True), "CRAB-W1": ("warehouse", True)}
    assert {c["system"]: c["code"] for c in d["current"]} == {"acc": "CRAB-A2", "warehouse": "CRAB-W1", "supplier": CRAB, "other": CRAB}


def test_a_rejection_needs_a_reason_and_leaves_the_register_alone(ac):
    rid = AL.request(ac, "CRAB-A1", "acc", CRAB, "ss")["request_id"]
    with pytest.raises(ValueError, match="ต้องระบุเหตุผล"):
        AL.decide(ac, rid, "rejected", " ", "ana")
    with pytest.raises(ValueError):
        AL.decide(ac, rid, "maybe", "x", "ana")
    out = AL.decide(ac, rid, "rejected", "รหัสนี้เป็นของสินค้าคนละรายการ", "ana")
    ac.commit()
    assert "ไม่ใช้" in out["message"]
    assert ac.execute("SELECT status, decision_note FROM bc_code_request").fetchone() == ("rejected", "รหัสนี้เป็นของสินค้าคนละรายการ")
    assert ac.execute("SELECT COUNT(*) FROM bc_ingredient_alias").fetchone()[0] == 0
    assert AL.pending_count(ac) == 0
    assert [r["status_th"] for r in AL.listing(ac, "all")] == ["ไม่ใช้"]


def test_the_code_is_checked_again_when_the_analyst_decides(ac):
    rid = AL.request(ac, "LATE-1", "acc", CRAB, "ss")["request_id"]
    ac.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('LATE-1','ของใหม่','raw','กรัม')")      # meanwhile
    with pytest.raises(ValueError, match="เป็นรหัสของ ของใหม่ ในทะเบียนอยู่แล้ว"):
        AL.decide(ac, rid, "applied", None, "ana")
    assert AL.pending_count(ac) == 1                                                       # still waiting, nothing half-written


def test_a_new_code_that_starts_later_does_not_replace_the_one_in_use_until_that_day(ac):
    AL.add_alias(ac, CRAB, "CRAB-A1", "acc", "2026-09-01", None, "ana")
    AL.add_alias(ac, CRAB, "CRAB-A2", "acc", "2026-11-01", None, "ana")
    assert AL.current_code(ac, CRAB, "acc", on="2026-10-12") == "CRAB-A1"                 # A2 has not started
    assert AL.current_code(ac, CRAB, "acc", on="2026-11-01") == "CRAB-A2"
    assert AL.current_map(ac, "acc", on="2026-10-12") == {CRAB: "CRAB-A1"}
    assert AL.in_use(ac, on="2026-11-02") == {"CRAB-A2"}


def test_retiring_never_deletes_and_a_retired_alias_no_longer_resolves(ac):
    AL.add_alias(ac, CRAB, "CRAB-A1", "acc", None, None, "ana")
    lines = [{"code": "CRAB-A1", "qty": 2, "unit": "กิโลกรัม", "value": 600}]
    assert "13-01-009" in P.aggregate(ac, lines)["codes"]
    AL.retire(ac, "CRAB-A1", "ana")
    assert _alias_row(ac, "CRAB-A1")[3:5] == (0, 0)                                        # the row stays, retired and not current
    assert AL.current_code(ac, CRAB, "acc") == CRAB and AL.active_map(ac) == {}
    agg = P.aggregate(ac, lines)
    assert agg["codes"] == {} and agg["issues"][0]["why"] == "ไม่มีในทะเบียนวัตถุดิบ"
    with pytest.raises(ValueError, match="เลิกใช้แล้ว"):
        AL.retire(ac, "CRAB-A1", "ana")
    with pytest.raises(ValueError, match="เลิกใช้แล้ว"):
        AL.set_current(ac, "CRAB-A1", "ana")
    # the same code for ANOTHER item is refused while it is only retired; for the same item it is brought back
    with pytest.raises(ValueError, match="เคยเป็นรหัสในระบบอื่นของ"):
        AL.add_alias(ac, SALMON, "CRAB-A1", "acc", None, None, "ana")
    AL.add_alias(ac, CRAB, "CRAB-A1", "acc", None, None, "ana")
    assert _alias_row(ac, "CRAB-A1")[3:5] == (1, 1) and "13-01-009" in P.aggregate(ac, lines)["codes"]
    # back to the register code for the system
    AL.use_register_code(ac, CRAB, "acc", "ana")
    assert AL.current_code(ac, CRAB, "acc") == CRAB and _alias_row(ac, "CRAB-A1")[3] == 1
    assert ac.execute("SELECT COUNT(*) FROM bc_ingredient_alias").fetchone()[0] == 1


# ── 4 · matching a name to a register item ─────────────────────────────────────────────────────────
def test_suggest_items_ranks_by_name_and_prefers_the_same_kind_of_unit(ac):
    ac.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('13-01-010','ปูอัดแท่งสลัด','raw','ชิ้น')")
    got = AL.suggest_items(ac, "ปูอัดซาชิมิ Kani Kamoboko 0.5Kg/Pack(20Pk)", code="NEW-1", unit="กิโลกรัม")
    assert got[0]["code"] == CRAB and got[0]["same_unit"] is True and got[0]["base_unit"] == "กรัม"
    assert all(h["code"] != "NEW-1" for h in got) and len(got) <= 3
    assert {h["code"] for h in AL.suggest_items(ac, "ปูอัด", unit="ชิ้น")} >= {CRAB, "13-01-010"}
    assert AL.suggest_items(ac, "") == [] and AL.suggest_items(ac, "รหัส 123") == []
    ac.execute("UPDATE bc_ingredient SET active=0 WHERE code=?", (CRAB,))                    # a retired item is not offered
    assert all(h["code"] != CRAB for h in AL.suggest_items(ac, "ปูอัดซาชิมิ"))


# ── 5 · the price file ─────────────────────────────────────────────────────────────────────────────
LINES = [{"code": "ACC-CRAB-2", "name": "ปูอัดซาชิมิ Kani 0.5Kg", "qty": 10, "unit": "กิโลกรัม", "value": 1500, "branch_id": 1},
         {"code": "ACC-NEW-9", "name": "ของที่ไม่มีที่ไหนเลย", "qty": 1, "unit": "ชิ้น", "value": 50, "branch_id": 1},
         {"code": SALMON, "name": "แซลมอน", "qty": 5, "unit": "กิโลกรัม", "value": 2000, "branch_id": 1}]


def test_the_dry_run_lists_unknown_codes_with_suggestions_and_adds_nothing(ac):
    before = ac.execute("SELECT COUNT(*) FROM bc_ingredient").fetchone()[0]
    agg = P.aggregate(ac, LINES, month="2026-09")
    pv = P.preview(ac, "2026-09", agg)
    assert pv["unknown_codes"] == ["ACC-CRAB-2", "ACC-NEW-9"] and pv["added_codes"] == []
    first = pv["unknown_items"][0]
    assert first["code"] == "ACC-CRAB-2" and first["value"] == 1500 and first["suggestions"][0]["code"] == CRAB
    assert pv["unknown_items"][1]["suggestions"] == []
    assert ac.execute("SELECT COUNT(*) FROM bc_ingredient").fetchone()[0] == before          # nothing was added
    assert "ACC-CRAB-2" not in agg["codes"] and set(agg["codes"]) == {SALMON}


def test_once_aliased_the_price_import_prices_the_item_under_its_register_code(ac):
    AL.add_alias(ac, CRAB, "ACC-CRAB-2", "acc", None, None, "ana")
    agg = P.aggregate(ac, LINES, month="2026-09")
    assert agg["codes"][CRAB]["qty_base"] == 10000 and agg["codes"][CRAB]["avg"] == pytest.approx(0.15)
    assert agg["branches"][(CRAB, 1)]["value"] == 1500
    assert P.preview(ac, "2026-09", agg)["unknown_codes"] == ["ACC-NEW-9"]


def test_choices_alias_a_code_for_the_analyst_and_file_a_request_for_everyone_else(ac):
    unknown = AL.unknown_in(ac, LINES)
    assert [u["code"] for u in unknown] == ["ACC-CRAB-2", "ACC-NEW-9"] and unknown[0]["lines"] == 1
    ch = AL.choices_from(json.dumps({"ACC-CRAB-2": {"as": "alias", "item": CRAB}, "ACC-NEW-9": {"as": "new"}}))
    added = []
    res = AL.apply_choices(ac, unknown, ch, "ana", True, add_new=lambda c, u: added.append(c))
    assert res == {"aliased": ["ACC-CRAB-2"], "requested": [], "new": ["ACC-NEW-9"], "open": []} and added == ["ACC-NEW-9"]
    assert _alias_row(ac, "ACC-CRAB-2")[3:5] == (1, 1)
    # someone who may not decide: the alias is a pending request, the register is untouched
    ac.execute("DELETE FROM bc_ingredient_alias")
    res = AL.apply_choices(ac, unknown, ch, "ss", False, add_new=lambda c, u: None)
    assert res["requested"] == ["ACC-CRAB-2"] and res["aliased"] == []
    assert AL.pending_count(ac) == 1 and ac.execute("SELECT COUNT(*) FROM bc_ingredient_alias").fetchone()[0] == 0
    # a code with no answer stays open
    assert AL.apply_choices(ac, unknown, {}, "ana", True)["open"] == ["ACC-CRAB-2", "ACC-NEW-9"]
    for bad in ('{"X": {"as": "maybe"}}', '{"X": {"as": "alias"}}', "[1]", "{"):
        with pytest.raises(ValueError):
            AL.choices_from(bad)
    assert AL.choices_from(None) is None and AL.choices_from("") is None and AL.choices_from("{}") == {}


# ── 6 · recipes and the forecast ───────────────────────────────────────────────────────────────────
def _usage(c):
    use, days, stats = F.usage_history(c, 1, date(2026, 9, 1), AS_OF, "2026-10")
    return {k: round(sum(d.values()), 3) for k, d in use.items()}, stats


def test_the_usage_the_forecast_counts_does_not_move_when_an_item_gets_a_new_code(ac, monkeypatch):
    monkeypatch.setattr(F, "USAGE_CACHE_TTL", 0)
    before, _ = _usage(ac)
    assert before[CRAB] > 0 and before[SALMON] > 0
    rid = AL.request(ac, "ACC-CRAB-2", "acc", CRAB, "ss")["request_id"]
    AL.decide(ac, rid, "applied", None, "ana")
    after, _ = _usage(ac)
    assert after == before                                                                   # the recipes keep the register code
    # a recipe line that already carries the NEW code is counted on the item (and was not before the alias)
    ac.execute("INSERT INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name, qty_per_serving, unit, price_per_unit,"
               " ingredient_cost, ingredient_type) VALUES(10, 2, 'ACC-CRAB-3', 'ปูอัด', 5, 'กรัม', 0, 0, 'protein')")
    lost, st = _usage(ac)
    assert lost[CRAB] == before[CRAB] and st["unconverted_lines"] >= 0
    AL.add_alias(ac, CRAB, "ACC-CRAB-3", "acc", None, None, "ana")
    more, _ = _usage(ac)
    assert more[CRAB] == pytest.approx(before[CRAB] * (50.0 / 45.0), rel=1e-3)               # 45 g + 5 g per plate of dish 10
    AL.retire(ac, "ACC-CRAB-3", "ana")
    assert _usage(ac)[0][CRAB] == lost[CRAB]


# ── 7 · the files that go to the accounting system ─────────────────────────────────────────────────
def _sheet_book():
    wb = Workbook()
    ws = wb.active
    ws.title = "รายการ"
    ws.append(["note"])
    ws.append(["วันที่", "รหัส", "วัตถุดิบ", "จำนวน"])
    ws.append(["2026-09-01", CRAB, "ปูอัด", 3])
    ws.append(["2026-09-01", SALMON, "แซลมอน", 4])
    ws.append(["", "", "รวม", 7])
    ws2 = wb.create_sheet("สรุป")
    ws2.append(["note"])
    ws2.append(["สาขา", "มูลค่า"])
    ws2.append(["อุดร", 100])
    return wb


def test_the_code_column_shows_the_accounting_code_and_keeps_the_register_code_in_the_last_column(ac):
    wb = _sheet_book()
    assert AL.recode_workbook(wb, {}) == 0
    assert [c.value for c in wb["รายการ"][2]] == ["วันที่", "รหัส", "วัตถุดิบ", "จำนวน"]       # no new code anywhere: left exactly as it was
    AL.add_alias(ac, CRAB, "CRAB-A2", "acc", None, None, "ana")
    AL.add_alias(ac, SALMON, "SAL-W", "warehouse", None, None, "ana")                         # a warehouse code is not the accounting code
    cur = AL.current_map(ac, "acc")
    assert cur == {CRAB: "CRAB-A2"}
    assert AL.recode_workbook(wb, cur) == 1
    ws = wb["รายการ"]
    assert [c.value for c in ws[2]] == ["วันที่", "รหัส", "วัตถุดิบ", "จำนวน", "รหัสในทะเบียน"]
    assert [c.value for c in ws[3]] == ["2026-09-01", "CRAB-A2", "ปูอัด", 3, CRAB]
    assert [c.value for c in ws[4]][1::3] == [SALMON, SALMON]                                  # unchanged items show the register code twice
    assert ws[5][4].value is None                                                              # the totals row has no code
    assert "รหัสในทะเบียน" in ws["A1"].value
    assert [c.value for c in wb["สรุป"][2]] == ["สาขา", "มูลค่า"]                              # a sheet with no code column is not touched


def test_exports_for_accounting_are_recoded_and_only_those(client, app_db):
    import inspect
    from routers import bom_center as RB
    for fn in (RB.bc_issues_export, RB.bc_issue_export_one, RB.bc_issue_summary_xlsx, RB.bc_issue_summary_all_xlsx,
               RB.bc_hub_consolidated_xlsx):
        assert "_for_accounting(wb)" in inspect.getsource(fn), fn.__name__
    for fn in (RB.bc_ir_reconcile_xlsx, RB.bc_code_review_xlsx, RB.bc_per_sales_xlsx):
        assert "_for_accounting" not in inspect.getsource(fn), fn.__name__


# ── 8 · the central report and the register file ───────────────────────────────────────────────────
def test_a_central_report_with_the_new_code_lands_on_the_item(udon):
    AL.add_alias(udon, CRAB, "03-NEW-CRAB", "acc", None, None, "ana")
    p = CE.parse_report(udon, _report([("03-NEW-CRAB", *ITEMS[0][1:])]))
    assert p["items"][0]["code"] == CRAB and p["items"][0]["file_code"] == "03-NEW-CRAB"
    row = CE.plan_catalogue(udon, p)["rows"][0]
    assert not row["is_new"] and row["aliased_from"] == "03-NEW-CRAB"
    out = CE.apply_catalogue(udon, p, "ss")
    assert out["inserted"] == 0
    assert udon.execute("SELECT COUNT(*) FROM bc_ingredient WHERE code='03-NEW-CRAB'").fetchone()[0] == 0   # no second item
    assert CE.items(udon)["items"][0]["code"] == CRAB


def test_a_new_central_code_is_listed_with_suggestions_and_added_only_when_chosen(udon):
    src = [("03-NEW-CRAB", ITEMS[0][1], ITEMS[0][2], ITEMS[0][3], ITEMS[0][4], ITEMS[0][5]),
           ("16-01-021", ITEMS[2][1], ITEMS[2][2], ITEMS[2][3], ITEMS[2][4], ITEMS[2][5])]
    p = CE.parse_report(udon, _report(src))
    rows = {r["code"]: r for r in CE.plan_catalogue(udon, p)["rows"]}
    assert rows["03-NEW-CRAB"]["is_new"] and rows["03-NEW-CRAB"]["suggestions"][0]["code"] == CRAB
    assert rows["16-01-021"]["suggestions"] == []
    # a strict run with no answers adds nothing and says which codes wait
    out = CE.apply_catalogue(udon, p, "ss", choices={})
    assert out["open_codes"] == ["03-NEW-CRAB", "16-01-021"] and out["saved"] == 0 and out["inserted"] == 0
    assert udon.execute("SELECT COUNT(*) FROM bc_ingredient WHERE code IN ('03-NEW-CRAB','16-01-021')").fetchone()[0] == 0
    # the analyst answers: the crab code is an alias, the apron is a new item
    ch = AL.choices_from(json.dumps({"03-NEW-CRAB": {"as": "alias", "item": CRAB}, "16-01-021": {"as": "new"}}))
    p = CE.parse_report(udon, _report(src))
    out = CE.apply_catalogue(udon, p, "ana", choices=ch, can_alias=True)
    assert out["aliased"] == 1 and out["inserted"] == 1 and out["open_codes"] == []
    assert udon.execute("SELECT COUNT(*) FROM bc_ingredient WHERE code='03-NEW-CRAB'").fetchone()[0] == 0
    assert udon.execute("SELECT 1 FROM bc_ingredient WHERE code='16-01-021'").fetchone()
    assert CRAB in {i["code"] for i in CE.items(udon)["items"]}
    assert _alias_row(udon, "03-NEW-CRAB")[:2] == (CRAB, "acc")


def test_a_central_run_by_someone_who_cannot_decide_files_a_request_and_skips_the_code(udon):
    src = [("03-NEW-CRAB", ITEMS[0][1], ITEMS[0][2], ITEMS[0][3], ITEMS[0][4], ITEMS[0][5])]
    ch = AL.choices_from(json.dumps({"03-NEW-CRAB": {"as": "alias", "item": CRAB}}))
    p = CE.parse_report(udon, _report(src))
    out = CE.apply_catalogue(udon, p, "ss", choices=ch, can_alias=False)
    assert out["requested"] == 1 and out["saved"] == 0 and AL.pending_count(udon) == 1
    assert udon.execute("SELECT COUNT(*) FROM bc_ingredient_alias").fetchone()[0] == 0
    dl = CE.deliveries(udon, p, "ss", None, dry_run=True)                                    # the month's quantity of that code waits too
    assert dl["branches"][0]["lines"] == 0 and "ยังไม่ได้เลือก" in dl["branches"][0]["skipped"][0]["why"]


def test_the_register_file_resolves_alias_codes_and_asks_before_adding(ac):
    AL.add_alias(ac, CRAB, "ACC-CRAB-2", "acc", None, None, "ana")
    ac.commit()
    rows = [{"code": "ACC-CRAB-2", "name": "ปูอัดซาชิมิ", "loss_pct": 3},
            {"code": "ZZ-NEW-1", "name": "ปูอัดแท่งใหม่ล่าสุด", "base_unit": "กรัม"},
            {"code": "ZZ-NEW-2", "name": "ของใหม่แท้ๆ", "base_unit": "ชิ้น"}]
    st = M.import_rows(ac, rows, "ana")                                                       # the old behaviour: no choices sent
    ac.rollback()
    assert {n["code"] for n in st["new_codes"]} == {"ZZ-NEW-1", "ZZ-NEW-2"}
    assert st["new_codes"][0]["suggestions"][0]["code"] == CRAB
    assert st["inserted"] == 2 and st["updated"] == 1                                         # the alias row updated the item, no 3rd item
    # strict: nothing is added without an answer
    st = M.import_rows(ac, rows, "ana", choices={})
    ac.rollback()
    assert st["inserted"] == 0 and {s["code"] for s in st["skipped"]} == {"ZZ-NEW-1", "ZZ-NEW-2"} and st["new_codes"] == []
    ch = AL.choices_from(json.dumps({"ZZ-NEW-1": {"as": "alias", "item": CRAB}, "ZZ-NEW-2": {"as": "new"}}))
    st = M.import_rows(ac, rows, "ana", choices=ch)
    ac.commit()
    assert st["aliased"] == 1 and st["inserted"] == 1
    assert ac.execute("SELECT COUNT(*) FROM bc_ingredient WHERE code='ZZ-NEW-1'").fetchone()[0] == 0
    assert AL.active_map(ac)["ZZ-NEW-1"] == CRAB and ac.execute("SELECT 1 FROM bc_ingredient WHERE code='ZZ-NEW-2'").fetchone()


def test_the_unit_file_never_adds_an_item_and_points_at_the_likely_one(ac):
    from bom_center import unit_sheet as US
    AL.add_alias(ac, CRAB, "ACC-CRAB-2", "acc", None, None, "ana")
    wb = Workbook()
    ws = wb.active
    ws.title = US.SHEET
    ws.append([US.NOTE])
    ws.append(US.HEADERS)
    ws.append(["ACC-CRAB-2", "ปูอัดซาชิมิ", "กรัม", "แท่ง", 15, None, "แน่นอน", None, None, None])
    ws.append(["ZZ-UNKNOWN", "ปูอัดแท่งสลัดใหม่", "กรัม", "ชิ้น", 5, None, "แน่นอน", None, None, None])
    ws.append(["ZZ-NONE", "ไม่มีชื่อคล้ายอะไรเลย", "กรัม", "ชิ้น", 5, None, "แน่นอน", None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    res = US.apply(ac, US.parse(buf.getvalue()), "ana")
    ac.rollback()
    by = {r["file_code"] or r["code"]: r for r in res["ladders"]}
    assert by["ACC-CRAB-2"]["code"] == CRAB and by["ACC-CRAB-2"]["problems"] == [] and by["ACC-CRAB-2"]["ladder"][0]["unit"] == "แท่ง"
    assert "ปูอัด" in by["ZZ-UNKNOWN"]["problems"][0] and "ไม่พบรหัส ZZ-UNKNOWN ในทะเบียน" in by["ZZ-UNKNOWN"]["problems"][0]
    assert "อาจเป็นรหัสใหม่ของ 13-01-009" in by["ZZ-UNKNOWN"]["problems"][0] or "รหัสในระบบอื่น" in by["ZZ-UNKNOWN"]["problems"][0]
    assert "รหัสในระบบอื่น" in by["ZZ-NONE"]["problems"][0] and res["problems"] == 2
    assert ac.execute("SELECT COUNT(*) FROM bc_ingredient WHERE code LIKE 'ZZ-%'").fetchone()[0] == 0


# ── 9 · หน้างาน ─────────────────────────────────────────────────────────────────────────────────────
def test_todo_lists_the_waiting_requests_for_the_analyst_only(ac):
    assert all(i["key"] != "code_request_review" for i in TD.build(ac, today=TODAY)["items"])
    AL.request(ac, "CRAB-A1", "acc", CRAB, "ss", old_code=CRAB)
    AL.request(ac, "SAL-A1", "acc", SALMON, "ss")
    out = TD.build(ac, today=TODAY)
    it = next(i for i in out["items"] if i["key"] == "code_request_review")
    assert it["title"] == "ตรวจรหัสใหม่ที่ประสานงานขายแจ้ง (2 เรื่อง)" and it["status"] == "todo" and it["count"] == 2
    assert "CRAB-A1" in it["detail"] and "ระบบบัญชี" in it["detail"] and "กรุณา" not in it["detail"]
    assert it["owner_modules"] == ["bom.center", "bom.prices"] and it["owner_th"] == "ฝ่ายวิเคราะห์"
    assert it["go"] == {"tab": "bomcenter", "pane": "ing", "anchor": "bmcCodeReqCard"}
    for r in AL.listing(ac, "pending"):
        AL.decide(ac, r["request_id"], "rejected", "ตรวจแล้ว", "ana")
    assert all(i["key"] != "code_request_review" for i in TD.build(ac, today=TODAY)["items"])
    AL.request(ac, "CRAB-A1", "acc", CRAB, "ss")
    assert all(i["key"] != "code_request_review" for i in TD.build(ac, scoped_branch="BranchA", today=TODAY)["items"])   # not a branch's task


# ── 10 · permissions ───────────────────────────────────────────────────────────────────────────────
def test_the_routes_follow_the_units_writers_and_sales_support_may_only_ask():
    ana, acc, br = auth.ROLE_MODULES["analyst"], auth.ROLE_MODULES["accounting"], auth.ROLE_MODULES["branch"]
    kit, ss, wh = auth.ROLE_MODULES["innovation"], auth.ROLE_MODULES["salessupport"], auth.ROLE_MODULES["warehouse"]
    op = auth.ROLE_MODULES["operation"]
    ask, decide = "/api/bom-center/codes/request", "/api/bom-center/codes/request/decide"
    assert auth.path_allowed("POST", ask, ss) and auth.path_allowed("POST", ask, ana)
    for other in (acc, br, kit, wh, op):
        assert not auth.path_allowed("POST", ask, other), other
    assert auth.path_allowed("POST", decide, ana) and auth.path_allowed("POST", decide, {"bom.prices"})
    for other in (ss, acc, br, kit, wh, op):                                                  # Sales Support asks, the analyst decides
        assert not auth.path_allowed("POST", decide, other), other
    reg = "/api/bom-center/ingredients/13-01-009/aliases"
    assert auth.path_allowed("POST", reg, ana) and not auth.path_allowed("POST", reg, ss)     # the register row is the analyst's
    assert not auth.path_allowed("POST", reg, kit) and not auth.path_allowed("POST", reg, br)
    for path in ("/api/bom-center/codes/requests", reg, "/api/bom-center/codes/suggest"):
        assert auth.path_allowed("GET", path, ana) and auth.path_allowed("GET", path, ss), path
        assert not auth.path_allowed("GET", path, br) and not auth.path_allowed("GET", path, wh), path


def test_the_write_handlers_clear_the_usage_cache_and_a_request_does_not():
    import inspect
    from routers import bom_center as RB
    for fn in (RB.bc_code_request_decide, RB.bc_ingredient_aliases_save, RB.bc_central_import, RB.bc_central_month_import):
        src = inspect.getsource(fn)
        assert "clear_usage_cache()" in src and "_drop_bom_days()" in src, fn.__name__
    assert "clear_usage_cache" not in inspect.getsource(RB.bc_code_request)                  # a request changes no recipe-side input
    for fn in (RB.bc_prices_import, RB.bc_prices_from_requisition, RB.bc_price_submission_import, RB.bc_ingredients_import):
        assert "code_choices" in inspect.getsource(fn), fn.__name__


# ── 11 · over HTTP ─────────────────────────────────────────────────────────────────────────────────
@pytest.fixture
def auth_on(monkeypatch):
    monkeypatch.setenv("OSHINEI_AUTH", "1")
    monkeypatch.setenv("OSHINEI_AUTH_SECRET", "test-secret-stable")
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "owner-token-xyz")
    yield


TOK = {"X-Upload-Token": "owner-token-xyz"}


def _login(user, pw=auth.DEFAULT_PASSWORD):
    import main
    from fastapi.testclient import TestClient
    c = TestClient(main.app)
    assert c.post("/api/auth/login", json={"username": user, "password": pw}).status_code == 200
    return c


def _seed_items():
    import core
    from bom_center.schema import ensure
    with core.get_db() as c:
        ensure(c)
        c.executemany("INSERT OR IGNORE INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv) VALUES(?,?,?,?,?,?)",
                      [("T-AL1", "ปลาแซลมอนทดสอบ", "raw", "กรัม", "แพ็ค", 1000), ("T-AL2", "เห็ดทดสอบ", "raw", "กรัม", "แพ็ค", 500)])
        c.commit()


def _cleanup_http():
    import core
    with core.get_db() as c:
        c.execute("DELETE FROM bc_ingredient_alias WHERE ingredient_code LIKE 'T-AL%' OR alias_code LIKE 'T-AL%' OR alias_code LIKE 'TAL-%'")
        c.execute("DELETE FROM bc_code_request WHERE ingredient_code LIKE 'T-AL%'")
        c.execute("DELETE FROM bc_ingredient_price WHERE ingredient_code LIKE 'T-AL%'")
        c.execute("DELETE FROM bc_ingredient WHERE code LIKE 'T-AL%' OR code LIKE 'TAL-%'")
        c.execute("DELETE FROM bc_audit WHERE row_key LIKE 'T-AL%' OR row_key LIKE 'TAL-%'")
        c.commit()


def _price_csv(rows) -> bytes:
    out = ["รหัส,ชื่อสินค้า,จำนวน,หน่วย,มูลค่ารวม,สาขา,เดือน"]
    out += [",".join(str(x) for x in r) for r in rows]
    return ("\n".join(out) + "\n").encode("utf-8-sig")


def test_http_flow_ask_decide_import_and_export(client, app_db, auth_on):
    _seed_items()
    try:
        ss, ana, mg = _login("salessupport1"), _login("analyst1"), _login("management1")
        # Sales Support asks; the owner's token is needed like every central write
        body = {"new_code": "TAL-NEW1", "ingredient_code": "T-AL1", "system": "acc", "valid_from": "2026-09-01", "note": "บัญชีเปลี่ยนรหัส"}
        assert ss.post("/api/bom-center/codes/request", json=body).headers.get("X-Denied") == "token"
        r = ss.post("/api/bom-center/codes/request", json=body, headers=TOK)
        assert r.status_code == 200 and r.json()["request_id"], r.text
        rid = r.json()["request_id"]
        assert ss.post("/api/bom-center/codes/request", json={**body, "new_code": "T-AL2"}, headers=TOK).status_code == 400   # a register code
        lst = ss.get("/api/bom-center/codes/requests").json()
        assert lst["pending"] == 1 and lst["requests"][0]["new_code"] == "TAL-NEW1" and lst["requests"][0]["system_th"] == "ระบบบัญชี"
        assert ss.get("/api/bom-center/codes/requests?status=zzz").status_code == 400
        # she cannot decide, the analyst can
        dec = {"request_id": rid, "decision": "applied"}
        assert ss.post("/api/bom-center/codes/request/decide", json=dec, headers=TOK).headers.get("X-Denied") == "role"
        assert ana.post("/api/bom-center/codes/request/decide", json=dec).headers.get("X-Denied") == "token"
        assert ana.post("/api/bom-center/codes/request/decide", json={**dec, "decision": "rejected"}, headers=TOK).status_code == 400
        done = ana.post("/api/bom-center/codes/request/decide", json=dec, headers=TOK)
        assert done.status_code == 200 and "สูตรไม่ต้องแก้" in done.json()["message"], done.text
        # the register row now shows it as the code in use
        row = next(x for x in ana.get("/api/bom-center/ingredients?q=T-AL1").json()["rows"] if x["code"] == "T-AL1")
        assert row["aliases"] == [{"alias": "TAL-NEW1", "system": "acc", "system_th": "ระบบบัญชี", "valid_from": "2026-09-01", "in_use": True}]
        assert [x["code"] for x in ana.get("/api/bom-center/ingredients?q=TAL-NEW1").json()["rows"]] == ["T-AL1"]       # found by the new code
        det = ana.get("/api/bom-center/ingredients/T-AL1/aliases").json()
        assert det["current"][0] == {"system": "acc", "system_th": "ระบบบัญชี", "code": "TAL-NEW1", "is_alias": True}
        assert ana.get("/api/bom-center/ingredients/NOPE/aliases").status_code == 404
        # the analyst records a warehouse code by hand, switches and retires; Sales Support cannot
        add = {"action": "add", "alias": "TAL-W1", "system": "warehouse", "note": "รหัสคลัง"}
        assert ss.post("/api/bom-center/ingredients/T-AL1/aliases", json=add, headers=TOK).headers.get("X-Denied") == "role"
        a = ana.post("/api/bom-center/ingredients/T-AL1/aliases", json=add, headers=TOK)
        assert a.status_code == 200 and {x["alias"] for x in a.json()["state"]["aliases"]} == {"TAL-NEW1", "TAL-W1"}
        assert ana.post("/api/bom-center/ingredients/T-AL2/aliases", json={"action": "retire", "alias": "TAL-W1"}, headers=TOK).status_code == 400
        assert ana.post("/api/bom-center/ingredients/T-AL1/aliases", json={"action": "retire", "alias": "TAL-W1"}, headers=TOK).status_code == 200
        assert ana.post("/api/bom-center/ingredients/T-AL1/aliases", json={"action": "fly"}, headers=TOK).status_code == 400
        assert ana.post("/api/bom-center/ingredients/T-AL1/aliases", json={"action": "register", "system": "acc"}, headers=TOK).status_code == 200
        assert ana.get("/api/bom-center/ingredients/T-AL1/aliases").json()["current"][0]["code"] == "T-AL1"
        # the price file with the new code: the dry run lists it with suggestions and adds nothing
        csv = _price_csv([("TAL-NEW1", "ปลาแซลมอนทดสอบ", 2, "กิโลกรัม", 800, "", "2026-09"),
                          ("TAL-NEW2", "ปลาแซลมอนทดสอบ ชุดใหม่", 3, "กิโลกรัม", 1200, "", "2026-09"),
                          ("TAL-NONE", "สินค้าแปลกประหลาดที่ไม่มีใครรู้จัก", 1, "ชิ้น", 10, "", "2026-09")])
        pre = ana.post("/api/bom-center/prices/import?dry_run=true", files={"file": ("p.csv", csv, "text/csv")})
        assert pre.status_code == 200, pre.text
        j = pre.json()
        assert [u["code"] for u in j["unknown_items"]] == ["TAL-NEW2", "TAL-NONE"]            # TAL-NEW1 is already an alias
        assert j["unknown_items"][0]["suggestions"][0]["code"] == "T-AL1" and j["unknown_items"][1]["suggestions"] == []
        assert j["added_codes"] == [] and j["priced"] == 1
        _noadd = ana.get("/api/bom-center/ingredients?q=TAL-NONE").json()["rows"]
        assert _noadd == []
        # the same dry run with the answers: the alias is applied inside the dry run and rolled back with it
        ch = json.dumps({"TAL-NEW2": {"as": "alias", "item": "T-AL1"}, "TAL-NONE": {"as": "new"}})
        pre2 = ana.post("/api/bom-center/prices/import?dry_run=true", files={"file": ("p.csv", csv, "text/csv")},
                        data={"code_choices": ch})
        j2 = pre2.json()
        assert j2["codes_decided"]["aliased"] == ["TAL-NEW2"] and j2["codes_decided"]["new"] == ["TAL-NONE"] and j2["unknown_items"] == []
        assert j2["priced"] == 2                                                                # T-AL1 (both codes) and the new item
        assert ana.get("/api/bom-center/ingredients/T-AL1/aliases").json()["aliases"][0]["alias"] == "TAL-NEW1"
        assert not any(x["alias"] == "TAL-NEW2" for x in ana.get("/api/bom-center/ingredients/T-AL1/aliases").json()["aliases"])
        # committing records them
        ok = ana.post("/api/bom-center/prices/import?dry_run=false", files={"file": ("p.csv", csv, "text/csv")},
                      data={"code_choices": ch}, headers=TOK)
        assert ok.status_code == 200 and ok.json()["codes_decided"]["aliased"] == ["TAL-NEW2"], ok.text
        assert {x["alias"] for x in ana.get("/api/bom-center/ingredients/T-AL1/aliases").json()["aliases"]} >= {"TAL-NEW1", "TAL-NEW2"}
        assert ana.get("/api/bom-center/prices/history?code=T-AL1").json()["rows"][0]["qty_bought"] == 5000.0
        # a bad answer is a plain 400
        bad = ana.post("/api/bom-center/prices/import?dry_run=true", files={"file": ("p.csv", csv, "text/csv")}, data={"code_choices": "{"})
        assert bad.status_code == 400
        # the branch cannot read any of it
        br = mg.post("/api/auth/users", headers=TOK, json={"username": "brn_al", "password": "Temp@12345", "role": "branch", "branch": "BranchA"})
        assert br.status_code in (200, 409)
        b = _login("brn_al", "Temp@12345")
        assert b.get("/api/bom-center/codes/requests").status_code == 403
        assert b.post("/api/bom-center/codes/request", json=body, headers=TOK).status_code == 403
    finally:
        _cleanup_http()


def test_http_central_choices_are_filed_as_a_request_for_sales_support(client, app_db, auth_on):
    _seed_items()
    try:
        ss = _login("salessupport1")
        raw = _report([("03-AL-9", "ปลาแซลมอนทดสอบ Fresh", "KG1.00/KG", 3, 0, "KG")])
        ch = json.dumps({"03-AL-9": {"as": "alias", "item": "T-AL1"}})
        pre = ss.post("/api/bom-center/central/import?dry_run=true", files={"file": ("r.xlsx", raw)})
        assert pre.status_code == 200 and pre.json()["rows"][0]["suggestions"][0]["code"] == "T-AL1"
        done = ss.post("/api/bom-center/central/import?dry_run=false&set_local=false", files={"file": ("r.xlsx", raw)},
                       data={"code_choices": ch}, headers=TOK)
        assert done.status_code == 200, done.text
        assert done.json()["requested"] == 1 and done.json()["aliased"] == 0 and done.json()["saved"] == 0
        reqs = ss.get("/api/bom-center/codes/requests").json()["requests"]
        assert [(r["new_code"], r["ingredient_code"], r["requested_by"]) for r in reqs if r["new_code"] == "03-AL-9"] == [("03-AL-9", "T-AL1", "salessupport1")]
    finally:
        import core
        with core.get_db() as c:
            c.execute("DELETE FROM bc_code_request WHERE new_code='03-AL-9'")
            c.execute("DELETE FROM bc_central_item WHERE code LIKE '03-AL%'")
            c.execute("DELETE FROM bc_ingredient WHERE code LIKE '03-AL%'")
            c.commit()
        _cleanup_http()


def test_http_the_consolidated_file_shows_the_accounting_code_and_keeps_the_register_code(client, app_db, auth_on):
    import main
    from test_bom_center_accounting_keyed import _cleanup, _seed_orders
    client.get("/api/bom-center/hub/queue")                                                  # creates the bc_* tables
    ids = _seed_orders(main.DB_PATH)
    try:
        ss, ana = _login("salessupport1"), _login("analyst1")
        url = "/api/bom-center/hub/consolidated.xlsx?date_from=2098-03-01&date_to=2098-03-31&status=all"

        def sheet(name):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(ss.get(url).content), read_only=True)
            return [r for r in wb[name].iter_rows(values_only=True) if r]
        head = next(r for r in sheet("ต่อสาขา") if r[0] == "สาขา")
        assert head[-1] == "คีย์บัญชี" and "รหัสในทะเบียน" not in head                       # no new code anywhere: the file is as it was
        r = ana.post("/api/bom-center/ingredients/T-KEY1/aliases", json={"action": "add", "alias": "TKEY-ACC", "system": "acc"}, headers=TOK)
        assert r.status_code == 200, r.text
        for name in ("ต่อสาขา", "สรุปรวมตามวันส่ง"):
            rows = sheet(name)
            head = next(r for r in rows if r[0] in ("สาขา", "วันส่ง"))
            assert head[-1] == "รหัสในทะเบียน" and head[head.index("รหัส")] == "รหัส", name
            data = [r for r in rows if "T-KEY1" in r]
            assert data and all(r[head.index("รหัส")] == "TKEY-ACC" and r[-1] == "T-KEY1" for r in data), name
        # a warehouse code is not the accounting code: the file does not change for it
        ana.post("/api/bom-center/ingredients/T-KEY1/aliases", json={"action": "retire", "alias": "TKEY-ACC"}, headers=TOK)
        ana.post("/api/bom-center/ingredients/T-KEY1/aliases", json={"action": "add", "alias": "TKEY-WH", "system": "warehouse"}, headers=TOK)
        assert next(r for r in sheet("ต่อสาขา") if r[0] == "สาขา")[-1] == "คีย์บัญชี"
    finally:
        import core
        with core.get_db() as c:
            c.execute("DELETE FROM bc_ingredient_alias WHERE ingredient_code='T-KEY1'")
            c.execute("DELETE FROM bc_audit WHERE row_key='T-KEY1' OR row_key LIKE 'TKEY-%'")
            c.commit()
        _cleanup(main.DB_PATH, ids)


import pytest as _pytest_s3  # noqa: E402
pytestmark = _pytest_s3.mark.usefixtures("recipe_mode")   # PR-S3: this module builds its forecast, plans and count sheets from recipes (mode full)
