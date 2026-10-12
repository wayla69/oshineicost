"""BoM Center — PR-U: the unit ladder, average conversions and portion recipes.

The promises:
  * an item has a LADDER of up to 3 units above its base unit, each stated against the one below
    ("1 กล่อง = 10 แพ็ค", "1 แพ็ค = 1,000 กรัม"); the metric step (กิโลกรัม → กรัม) is universal and not a step;
  * a bad ladder (a 4th step, a loop, a parent that is not in the ladder, กิโลกรัม redefined for a กรัม item) is refused
    as a whole with a reason, and nothing is written;
  * a step can be an AVERAGE (1 ตัว ≈ 4.5 กก., range 4–5): every conversion through it says so (approx, "≈");
  * rules written before the ladder keep converting; qty_per_recv follows the ladder when the receiving unit is one of its units;
  * the ladder travels through an Excel file (dry run writes nothing, a bad code blocks the whole file);
  * a portion recipe turns ONE whole piece into N pieces: cost per piece, grams per piece, no trim loss on top — in the
    cost AND in the order plan — and a release freezes it.
"""
import io
import sqlite3

import pytest

import auth
import conftest as _cf
from bom_center import forecast as F
from bom_center import master as M
from bom_center import prep as P
from bom_center import schema as S
from bom_center import unit_sheet as US
from bom_center import units as U
from bom_center.units import Converter, check_ladder
from test_bom_center_s2 import _db

SALMON = "02-01-001"            # base กรัม, trim loss 40 %, ฿0.5 per กรัม (฿500 per กก.) in these tests
DISH = "OTHBFGNT00010"          # sku 11


def _row(unit, qty, parent=None, kind="exact", rmin=None, rmax=None, label=None):
    return {"unit": unit, "qty_in_parent": qty, "parent_unit": parent, "conv_kind": kind, "range_min": rmin, "range_max": rmax,
            "label": label}


@pytest.fixture
def con(tmp_path):
    c = _db(tmp_path)
    c.execute("UPDATE bc_ingredient SET recv_unit='กิโลกรัม', qty_per_recv=1000 WHERE code=?", (SALMON,))
    c.execute("DELETE FROM bc_ingredient_price WHERE ingredient_code=?", (SALMON,))
    c.execute("INSERT INTO bc_ingredient_price(ingredient_code, month, price_per_base, source) VALUES(?,?,?,?)",
              (SALMON, "2026-09", 0.5, "purchase_avg"))
    c.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit, loss_pct) VALUES('X-001','ซอส','raw','มิลลิลิตร',0)")
    c.commit()
    yield c
    c.close()


# ── the converter and the validator ──────────────────────────────────────────
def test_a_chain_box_pack_gram(con):
    M.set_ladder(con, "X-001", [_row("กล่อง", 10, "แพ็ค"), _row("แพ็ค", 1000)], "t")
    cv = Converter(con)
    assert cv.to_base("X-001", 2, "กล่อง")[0] == 20000                      # 2 × 10 × 1000
    f, how = cv.factor("X-001", "กล่อง")
    assert f == 10000 and how == "1 กล่อง = 10 แพ็ค · 1 แพ็ค = 1,000 มิลลิลิตร" and not cv.last_approx
    assert [r["unit"] for r in cv.ladder("X-001")] == ["แพ็ค", "กล่อง"]            # lowest step first
    assert cv.ladder("X-001")[1]["level"] == 2 and cv.ladder("X-001")[1]["parent_unit"] == "แพ็ค"
    assert cv.from_base("X-001", 5000, "กล่อง")[0] == 0.5
    stored = con.execute("SELECT factor_to_base, parent_unit, qty_in_parent FROM bc_unit_rule WHERE unit='กล่อง'").fetchone()
    assert stored == (10000.0, "แพ็ค", 10.0)                                 # factor_to_base is recomputed for old readers


def test_three_steps_are_fine_and_a_fourth_is_refused(con):
    three = [_row("ลัง", 4, "กล่อง"), _row("กล่อง", 10, "แพ็ค"), _row("แพ็ค", 500)]
    assert check_ladder("มิลลิลิตร", three) == []
    four = three + [_row("พาเลท", 20, "ลัง")]
    probs = check_ladder("มิลลิลิตร", four)
    assert any("ไม่เกิน 3" in p for p in probs)
    with pytest.raises(ValueError, match="ไม่เกิน 3"):
        M.set_ladder(con, "X-001", four, "t")
    assert con.execute("SELECT COUNT(*) FROM bc_unit_rule WHERE ingredient_code='X-001'").fetchone()[0] == 0   # nothing written
    # three units side by side on the base is also the limit of "3 units above the base"
    side = [_row(u, 5) for u in ("ถุง", "กล่อง", "ถัง")]
    assert check_ladder("กรัม", side) == []
    assert any("ไม่เกิน 3" in p for p in check_ladder("กรัม", side + [_row("ลัง", 5)]))


def test_a_loop_a_missing_parent_and_a_redefined_metric_unit_are_refused():
    assert any("วนกลับ" in p for p in check_ladder("กรัม", [_row("ก", 2, "ข"), _row("ข", 2, "ก")]))
    assert any("ไม่อยู่ในบันไดหน่วย" in p for p in check_ladder("กรัม", [_row("ถุง", 2, "ลัง")]))
    assert any("หน่วยฐานหรือหน่วยมาตรฐาน" in p for p in check_ladder("กรัม", [_row("กิโลกรัม", 900)]))
    assert any("หน่วยฐานหรือหน่วยมาตรฐาน" in p for p in check_ladder("กรัม", [_row("g", 2)]))
    assert any("ซ้ำ" in p for p in check_ladder("กรัม", [_row("ถุง", 2), _row("ถุง", 3)]))
    assert any("มากกว่า 0" in p for p in check_ladder("กรัม", [_row("ถุง", 0)]))
    assert any("ยังไม่ระบุหน่วยฐาน" in p for p in check_ladder("", [_row("ถุง", 2)]))
    assert check_ladder("มิลลิลิตร", [_row("ขวด", 1, "ลิตร")]) == []                      # a metric parent is fine


def test_a_metric_parent_is_not_a_step(con):
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, 5)], "t")
    cv = Converter(con)
    assert cv.to_base(SALMON, 1, "ตัว")[0] == pytest.approx(4500)
    assert cv.ladder(SALMON)[0]["level"] == 1                                    # กิโลกรัม → กรัม does not count
    # a box of those pieces: ตัว (level 1) under กล่อง (level 2) — the metric step still does not count, so 3 fit
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, 5), _row("กล่อง", 6, "ตัว"), _row("ลัง", 2, "กล่อง")], "t")
    assert Converter(con).to_base(SALMON, 1, "ลัง")[0] == pytest.approx(2 * 6 * 4500)


def test_an_average_step_is_approximate_and_says_so(con):
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, 5), _row("ถุง", 2000)], "t")
    cv = Converter(con)
    f, how = cv.factor(SALMON, "ตัว")
    assert f == 4500 and cv.last_approx
    assert how.startswith("≈ 1 ตัว = 4.5 กิโลกรัม") and how.endswith("(เฉลี่ย ช่วง 4–5 กิโลกรัม)")
    cv.factor(SALMON, "ถุง")
    assert not cv.last_approx                                                    # the flag follows the LAST call
    cv.to_base(SALMON, 3, "ตัว")
    assert cv.last_approx
    assert cv.approx_units(SALMON) == {"ตัว"}
    assert cv.approx_unit(SALMON) == ("ตัว", 4500)
    assert cv.span(SALMON, "ตัว") == (4000, 5000) and cv.span(SALMON, "ถุง") is None
    # an average outside its range is refused; a range needs both ends; exact has no range
    assert any("อยู่ในช่วง" in p for p in check_ladder("กรัม", [_row("ตัว", 6, "กิโลกรัม", "avg", 4, 5)]))
    assert any("ทั้งค่าต่ำสุดและค่าสูงสุด" in p for p in check_ladder("กรัม", [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, None)]))
    assert any("ไม่มีช่วง" in p for p in check_ladder("กรัม", [_row("ตัว", 4.5, "กิโลกรัม", "exact", 4, 5)]))
    assert any("ชนิดการแปลง" in p for p in check_ladder("กรัม", [_row("ตัว", 4.5, "กิโลกรัม", "แปลก")]))
    # the average a person left empty is the midpoint of the range (the screen does the same while typing)
    r = M.set_ladder(con, SALMON, [_row("ตัว", None, "กิโลกรัม", "เฉลี่ย", 4, 5)], "t")
    assert r["ladder"][0]["qty_in_parent"] == 4.5 and r["ladder"][0]["approx"]


def test_a_chain_through_an_average_is_approximate_too(con):
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, 5), _row("กล่อง", 6, "ตัว")], "t")
    cv = Converter(con)
    f, how = cv.factor(SALMON, "กล่อง")
    assert f == pytest.approx(27000) and cv.last_approx and how.startswith("≈ 1 กล่อง = 6 ตัว · 1 ตัว = 4.5 กิโลกรัม")
    assert cv.approx_units(SALMON) == {"ตัว", "กล่อง"}


def test_legacy_rules_without_the_ladder_columns_still_convert(tmp_path):
    p = tmp_path / "old.db"
    c = sqlite3.connect(str(p))
    c.execute("CREATE TABLE bc_ingredient(code TEXT PRIMARY KEY, base_unit TEXT, recv_unit TEXT, qty_per_recv REAL)")
    c.execute("CREATE TABLE bc_unit_rule(ingredient_code TEXT NOT NULL, unit TEXT NOT NULL, factor_to_base REAL NOT NULL,"
              " label TEXT, PRIMARY KEY (ingredient_code, unit))")
    c.execute("INSERT INTO bc_ingredient VALUES('13-01-009','กรัม',NULL,NULL)")
    c.execute("INSERT INTO bc_unit_rule VALUES('13-01-009','แท่ง',15,NULL)")
    assert Converter(c).to_base("13-01-009", 3, "แท่ง") == (45, "1 แท่ง = 15 กรัม")
    c.close()


def test_ensure_backfills_old_rules_as_level_one_exact_steps(tmp_path):
    p = tmp_path / "mig.db"
    _cf._build_db(str(p))
    c = sqlite3.connect(str(p))
    c.execute("CREATE TABLE bc_unit_rule(ingredient_code TEXT NOT NULL, unit TEXT NOT NULL, factor_to_base REAL NOT NULL,"
              " label TEXT, PRIMARY KEY (ingredient_code, unit))")
    c.execute("INSERT INTO bc_unit_rule VALUES('A','แท่ง',15,NULL)")
    S.ensure(c, commit=True)
    S.ensure(c, commit=True)                                                    # idempotent
    assert c.execute("SELECT qty_in_parent, parent_unit, conv_kind FROM bc_unit_rule").fetchall() == [(15.0, None, None)]
    c.close()


def test_the_pack_size_follows_the_ladder_when_the_receiving_unit_is_a_ladder_unit(con):
    con.execute("UPDATE bc_ingredient SET recv_unit='กล่อง', qty_per_recv=999 WHERE code='X-001'")
    r = M.set_ladder(con, "X-001", [_row("กล่อง", 10, "แพ็ค"), _row("แพ็ค", 1000)], "t")
    assert r["qty_per_recv_change"]["to"] == 10000 and r["qty_per_recv_change"]["from"] == 999
    assert con.execute("SELECT qty_per_recv FROM bc_ingredient WHERE code='X-001'").fetchone()[0] == 10000
    # a receiving unit that is not in the ladder is left alone
    r = M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg")], "t")
    assert r["qty_per_recv_change"] is None
    assert con.execute("SELECT qty_per_recv FROM bc_ingredient WHERE code=?", (SALMON,)).fetchone()[0] == 1000


def test_set_ladder_replaces_the_whole_ladder_and_is_audited(con):
    M.set_ladder(con, "X-001", [_row("ถุง", 500), _row("กล่อง", 10, "ถุง")], "ana")
    M.set_ladder(con, "X-001", [_row("ถุง", 450)], "ana")
    assert [r[0] for r in con.execute("SELECT unit FROM bc_unit_rule WHERE ingredient_code='X-001'")] == ["ถุง"]
    aud = con.execute("SELECT who, old_value, new_value FROM bc_audit WHERE action='unit_ladder' AND row_key='X-001' ORDER BY id").fetchall()
    assert len(aud) == 2 and aud[1][0] == "ana" and "กล่อง" in aud[1][1] and "ถุง=450" in aud[1][2]
    # the single-rule writer is the same ladder: a rule on an unknown parent is refused, a delete behind a child too
    with pytest.raises(ValueError, match="ไม่อยู่ในบันไดหน่วย"):
        M.set_unit_rule(con, "X-001", "ลัง", None, None, "t", parent_unit="หีบ", qty_in_parent=2)
    M.set_unit_rule(con, "X-001", "กล่อง", None, None, "t", parent_unit="ถุง", qty_in_parent=10)
    with pytest.raises(ValueError, match="อ้างอิงหน่วยนี้อยู่"):
        M.set_unit_rule(con, "X-001", "ถุง", None, None, "t")
    M.set_unit_rule(con, "X-001", "กล่อง", None, None, "t")
    M.set_unit_rule(con, "X-001", "ถุง", 450, None, "t")                          # legacy call shape
    assert Converter(con).to_base("X-001", 2, "ถุง")[0] == 900


# ── Excel ────────────────────────────────────────────────────────────────────
def _file(ladder_rows, portion_rows=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = US.SHEET
    ws.append([US.NOTE])
    ws.append(US.HEADERS)
    for r in ladder_rows:
        ws.append(r)
    if portion_rows is not None:
        w2 = wb.create_sheet(US.PORTION_SHEET)
        w2.append([US.PORTION_NOTE])
        w2.append(US.PORTION_HEADERS)
        for r in portion_rows:
            w2.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_the_headers_are_the_fixed_layout():
    assert US.HEADERS == ["รหัส", "ชื่อวัตถุดิบ", "หน่วยฐาน", "หน่วย", "1 หน่วยนี้ =", "หน่วยที่อ้างอิง", "ชนิดการแปลง", "ต่ำสุด", "สูงสุด", "หมายเหตุ"]
    assert US.PORTION_HEADERS == ["รหัสสินค้ากึ่งสำเร็จ", "ชื่อสินค้ากึ่งสำเร็จ", "รหัสวัตถุดิบ", "จำนวนที่ใช้", "หน่วยที่ใช้", "ได้ (ชิ้น)",
                                  "หน่วยที่ได้", "หมายเหตุ"]
    assert US.SHEET == "หน่วยและการแปลง" and US.PORTION_SHEET == "สูตรแบ่งชิ้น"


def test_template_import_round_trip_dry_run_commit_and_a_bad_code(con):
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, 5)], "t")
    con.commit()
    rows = US.ladder_rows(con)
    assert rows == [[SALMON, "แซลมอน", "กรัม", "ตัว", 4.5, "กิโลกรัม", "เฉลี่ย", 4.0, 5.0, None]]
    assert len(US.ladder_rows(con, all_items=True)) > 1                         # every active item with no ladder: an empty row
    data = _file([[SALMON, "x", "กรัม", "ตัว", 4.8, "กิโลกรัม", "เฉลี่ย", 4, 5, "ไซส์ 4–5 กก."],
                  ["X-001", "ซอส", "มิลลิลิตร", "ขวด", 500, None, "แน่นอน", None, None, None],
                  ["X-001", "ซอส", "มิลลิลิตร", "ลัง", 12, "ขวด", "แน่นอน", None, None, None]])
    parsed = US.parse(data)
    before = con.execute("SELECT COUNT(*) FROM bc_unit_rule").fetchone()[0]
    res = US.apply(con, parsed, "ana")                                          # the router rolls a dry run back
    assert res["problems"] == 0 and {r["code"] for r in res["ladders"]} == {SALMON, "X-001"}
    x = next(r for r in res["ladders"] if r["code"] == "X-001")
    assert [(r["unit"], r["factor_to_base"]) for r in x["ladder"]] == [("ขวด", 500.0), ("ลัง", 6000.0)]
    con.rollback()
    assert con.execute("SELECT COUNT(*) FROM bc_unit_rule").fetchone()[0] == before          # dry run wrote nothing
    assert Converter(con).to_base(SALMON, 1, "ตัว")[0] == pytest.approx(4500)
    res = US.apply(con, parsed, "ana")                                          # commit
    con.commit()
    assert Converter(con).to_base(SALMON, 1, "ตัว")[0] == pytest.approx(4800)
    assert Converter(con).to_base("X-001", 1, "ลัง")[0] == 6000
    # a bad code blocks the whole file: the good code in it is not written either (the router refuses the commit)
    bad = US.parse(_file([["X-001", "ซอส", "มิลลิลิตร", "ขวด", 999, None, "แน่นอน", None, None, None],
                          ["NOPE-1", "ไม่มี", "กรัม", "ถุง", 5, None, "แน่นอน", None, None, None]]))
    res = US.apply(con, bad, "ana")
    assert res["problems"] == 1
    assert "ไม่พบรหัส" in next(r for r in res["ladders"] if r["code"] == "NOPE-1")["problems"][0]


def test_a_different_base_unit_in_the_file_is_reported_not_applied_and_blank_units_keep_the_ladder(con):
    M.set_ladder(con, "X-001", [_row("ขวด", 500)], "t")
    parsed = US.parse(_file([["X-001", "ซอส", "กรัม", "ถุง", 100, None, "แน่นอน", None, None, None],
                             [SALMON, "แซลมอน", "กรัม", None, None, None, None, None, None, None]]))
    assert list(parsed["ladders"]) == ["X-001"]                                  # a row with no unit never reaches a code
    res = US.apply(con, parsed, "ana")
    assert res["problems"] == 0 and any("ต่างจากทะเบียน" in w for w in res["warnings"])
    assert Converter(con).base_unit("X-001") == "มิลลิลิตร"


def test_the_template_and_import_endpoints(client, monkeypatch):
    import main
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "pru-test-token")
    c = sqlite3.connect(str(main.DB_PATH))
    tables = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'bc%'")]
    snap = {t: c.execute(f"SELECT * FROM {t}").fetchall() for t in tables}
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('PRU-1','ทดสอบ','raw','กรัม')")
    c.commit()
    c.close()
    try:
        r = client.get("/api/bom-center/units/template.xlsx?all=1")
        assert r.status_code == 200
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == [US.SHEET, US.PORTION_SHEET]
        assert [c.value for c in wb[US.SHEET][2]] == US.HEADERS and wb[US.SHEET]["A1"].value == US.NOTE
        data = _file([["PRU-1", "ทดสอบ", "กรัม", "ถุง", 250, None, "แน่นอน", None, None, None]])
        pre = client.post("/api/bom-center/units/import?dry_run=true", files={"file": ("u.xlsx", data)})      # no token needed
        assert pre.status_code == 200 and pre.json()["codes_ok"] == 1 and pre.json()["ladders"][0]["ladder"][0]["factor_to_base"] == 250
        assert client.get("/api/bom-center/units/ladder?code=PRU-1").json()["ladder"] == []                  # nothing written
        no = client.post("/api/bom-center/units/import?dry_run=false", files={"file": ("u.xlsx", data)})
        assert no.status_code == 403 and no.headers.get("X-Denied") == "token"
        hdr = {"X-Upload-Token": "pru-test-token"}
        bad = _file([["PRU-1", "ทดสอบ", "กรัม", "ถุง", 250, None, "แน่นอน", None, None, None],
                     ["NOPE-9", "ไม่มี", "กรัม", "ถุง", 5, None, "แน่นอน", None, None, None]])
        r = client.post("/api/bom-center/units/import?dry_run=false", files={"file": ("u.xlsx", bad)}, headers=hdr)
        assert r.status_code == 400 and "มีรายการที่ต้องแก้ไข 1 รายการ จึงยังไม่บันทึก" in r.text
        assert client.get("/api/bom-center/units/ladder?code=PRU-1").json()["ladder"] == []
        r = client.post("/api/bom-center/units/import?dry_run=false", files={"file": ("u.xlsx", data)}, headers=hdr)
        assert r.status_code == 200, r.text
        lad = client.get("/api/bom-center/units/ladder?code=PRU-1").json()
        assert lad["base_unit"] == "กรัม" and lad["ladder"][0]["unit"] == "ถุง" and lad["problems"] == []
        # the ladder endpoints
        r = client.post("/api/bom-center/units/ladder", headers=hdr, json={"code": "PRU-1", "rows": [
            {"unit": "ตัว", "qty_in_parent": 4.5, "parent_unit": "กิโลกรัม", "conv_kind": "avg", "range_min": 4, "range_max": 5}]})
        assert r.status_code == 200 and r.json()["ladder"][0]["approx"] is True
        bad = client.post("/api/bom-center/units/ladder", headers=hdr, json={"code": "PRU-1", "rows": [
            {"unit": "กิโลกรัม", "qty_in_parent": 900}]})
        assert bad.status_code == 400 and "หน่วยฐานหรือหน่วยมาตรฐาน" in bad.text
        conv = client.get("/api/bom-center/convert", params={"code": "PRU-1", "qty": 2, "unit": "ตัว"}).json()
        assert conv["qty_base"] == 9000 and conv["approx"] is True and conv["how"].startswith("≈ ") and conv["range_base"] == [8000, 10000]
        ing = client.get("/api/bom-center/ingredients?q=PRU-1").json()["rows"][0]
        assert ing["unit_rules"][0]["approx"] is True and ing["unit_rules"][0]["qty_in_parent"] == 4.5
        old = client.post("/api/bom-center/units", headers=hdr, json=[{"code": "PRU-1", "unit": "ถุง", "factor_to_base": 250}])
        assert old.status_code == 200
    finally:
        c = sqlite3.connect(str(main.DB_PATH))
        for t, rows in snap.items():
            c.execute(f"DELETE FROM {t}")
            if rows:
                c.executemany(f"INSERT INTO {t} VALUES({','.join('?' * len(rows[0]))})", rows)
        c.commit()
        c.close()


# ── portion recipes ──────────────────────────────────────────────────────────
def _portion(con, pieces=225):
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, 5)], "t")
    return P.save_portion(con, "CK-801", SALMON, 1, "ตัว", pieces, "ชิ้น", "kitchen", note="หัว หนัง", name="แซลมอนแล่ชิ้น",
                          month="2026-09")


def test_a_portion_recipe_costs_a_piece_without_the_trim_loss(con):
    r = _portion(con)
    assert r["created"] and r["approx"] and r["how"].startswith("≈ ")
    assert r["factor_base"] == 4500 and r["base_per_piece"] == 20 and r["cost_per_piece"] == pytest.approx(10.0)
    assert r["price_month"] == "2026-09" and r["range"]["per_piece_min"] == pytest.approx(4000 / 225)
    # the master row the recipe made
    row = con.execute("SELECT kind, family, base_unit, sourcing_default, note FROM bc_ingredient WHERE code='CK-801'").fetchone()
    assert row == ("prepared", "CK", "ชิ้น", "local", "สร้างจากสูตรแบ่งชิ้น")
    c = P.PrepCoster(con, "2026-09", P.book(con)).cost("CK-801")
    assert c["per_base"] == pytest.approx(10.0) and c["approx"] and c["mode"] == "portion"
    assert c["lines"][0]["no_loss"] == 1 and c["lines"][0]["cost"] == pytest.approx(2250)         # NOT 2250 / 0.6
    # the same line as a plain batch recipe WOULD add the 40 % loss — that is the difference no_loss makes
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('CK-802','x','prepared','ชิ้น')")
    P.save(con, "CK-802", 225, "ชิ้น", [{"ingredient_code": SALMON, "qty": 1, "unit": "ตัว"}], "kitchen")
    assert P.PrepCoster(con, "2026-09", P.book(con)).cost("CK-802")["per_base"] == pytest.approx(10.0 / 0.6)
    ov = {x["code"]: x for x in P.overview(con, "2026-09")}["CK-801"]
    assert ov["mode"] == "portion" and ov["mode_th"] == "แบ่งชิ้น" and ov["approx"] and ov["per_base"] == pytest.approx(10.0)
    d = P.detail(con, "CK-801", "2026-09")
    assert d["mode"] == "portion" and d["recipe_note"] == "หัว หนัง" and d["raw_per_base"][0]["qty_per_base"] == pytest.approx(20)


def test_a_dish_of_five_pieces_costs_50_and_orders_100_g_with_no_loss(con):
    _portion(con)
    con.execute("INSERT INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name, qty_per_serving, unit, price_per_unit,"
                " ingredient_cost, ingredient_type) VALUES(11, 1, 'CK-801', 'แซลมอนแล่ชิ้น', 5, 'ชิ้น', 10, 50, 'protein')")
    con.commit()
    # cost: 5 pieces × ฿10.00
    from bom_center.recipes import Coster
    out = Coster(con, "2026-09", P.book(con)).line({"ingredient_code": "CK-801", "qty": 5, "unit": "ชิ้น"})
    assert out["cost"] == pytest.approx(50.0)
    # explode: 20 g per piece, the loss is not in it, and the order plan does not gross it up by 40 % either
    assert P.explode(P.book(con), Converter(con), "CK-801")[0][SALMON] == pytest.approx(20.0)
    per, stats = F._per_dish(con, "2026-10")
    got = dict(per[11])
    assert got[SALMON] == pytest.approx(100.0) and stats["via"][SALMON] == ["CK-801"]
    # contrast: the same salmon inside an ordinary batch recipe IS grossed up (÷ 0.6)
    con.execute("DELETE FROM bom_ingredients WHERE sku_id=11")
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('CK-802','x','prepared','ชิ้น')")
    P.save(con, "CK-802", 225, "ชิ้น", [{"ingredient_code": SALMON, "qty": 1, "unit": "ตัว"}], "kitchen")
    con.execute("INSERT INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name, qty_per_serving, unit, price_per_unit,"
                " ingredient_cost, ingredient_type) VALUES(11, 1, 'CK-802', 'x', 5, 'ชิ้น', 10, 50, 'protein')")
    per, _ = F._per_dish(con, "2026-10")
    assert dict(per[11])[SALMON] == pytest.approx(100.0 / 0.6)


def test_a_nested_portion_keeps_its_no_loss_through_another_prepared_item(con):
    _portion(con)
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('CK-803','ซูชิเซ็ต','prepared','ชิ้น')")
    P.save(con, "CK-803", 1, "ชิ้น", [{"ingredient_code": "CK-801", "qty": 2, "unit": "ชิ้น"},
                                      {"ingredient_code": SALMON, "qty": 100, "unit": "กรัม"}], "kitchen")
    lossy, flat, _ = P.explode_parts(P.book(con), Converter(con), "CK-803")
    assert flat[SALMON] == pytest.approx(40) and lossy[SALMON] == pytest.approx(100)


def test_portion_input_is_validated(con):
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg")], "t")
    for args, why in [((SALMON, 0, "ตัว", 225), "มากกว่า 0"), ((SALMON, 1, "ตัว", 0), "มากกว่า 0"),
                      (("NOPE", 1, "ตัว", 225), "ไม่พบรหัส"), ((SALMON, 1, "ก้อน", 225), "แปลงหน่วย"),
                      (("CK-001", 1, "ชิ้น", 225), "ไม่พบรหัส")]:
        with pytest.raises(ValueError, match=why):
            P.portion_preview(con, *args)
    with pytest.raises(ValueError, match="คนละรหัส"):
        P.save_portion(con, SALMON, SALMON, 1, "ตัว", 225, "ชิ้น", "k")
    with pytest.raises(ValueError, match="ใส่ชื่อ"):
        P.save_portion(con, "CK-900", SALMON, 1, "ตัว", 225, "ชิ้น", "k")                      # a new code needs a name
    con.execute("INSERT INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('CK-901','x','prepared','ชิ้น')")
    with pytest.raises(ValueError, match="วัตถุดิบดิบ"):
        P.portion_preview(con, "CK-901", 1, "ชิ้น", 5)
    with pytest.raises(ValueError, match="ไม่ใช่สินค้ากึ่งสำเร็จ"):
        P.save_portion(con, "X-001", SALMON, 1, "ตัว", 225, "ชิ้น", "k")
    # a portion recipe has ONE line
    with pytest.raises(ValueError, match="1 บรรทัด"):
        P.save(con, "CK-901", 5, "ชิ้น", [{"ingredient_code": SALMON, "qty": 1, "unit": "ตัว"},
                                          {"ingredient_code": "X-001", "qty": 1, "unit": "มิลลิลิตร"}], "k", mode="portion")


def test_a_release_freezes_mode_and_no_loss(con):
    from bom_center import release as RL
    from bom_center import snapshot_v0 as V0
    _portion(con)
    con.commit()
    V0.build_v0(con)
    con.execute("INSERT INTO bc_release(label, kind, effective_month, status, scope, price_month) VALUES('t','monthly','2026-10','draft','all','2026-09')")
    rid = con.execute("SELECT MAX(release_id) FROM bc_release").fetchone()[0]
    assert P.freeze(con, rid) >= 1
    assert con.execute("SELECT mode FROM bc_rel_prep_recipe WHERE release_id=? AND prepared_code='CK-801'", (rid,)).fetchone()[0] == "portion"
    assert con.execute("SELECT no_loss FROM bc_rel_prep_line WHERE release_id=? AND prepared_code='CK-801'", (rid,)).fetchone()[0] == 1
    con.execute("UPDATE bc_release SET status='released' WHERE release_id=?", (rid,))
    frozen = P.book(con, rid)["CK-801"]
    assert frozen["mode"] == "portion" and frozen["lines"][0]["no_loss"] == 1
    # the kitchen edits the live recipe afterwards: the released month does not move
    P.save_portion(con, "CK-801", SALMON, 1, "ตัว", 100, "ชิ้น", "k")
    assert P.book(con, rid)["CK-801"]["yield_qty"] == 225
    assert RL is not None


def test_saving_a_portion_recipe_through_the_generic_editor_or_the_flat_file_keeps_it_a_portion_recipe(con):
    _portion(con)
    flat = P.export_flat(con)
    row = next(r for r in flat if r[0] == "CK-801")
    assert row[-2:] == ["portion", 1] and P.FLAT_COLS[-2:] == ("mode", "no_loss")
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(list(P.FLAT_COLS))
    wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    assert P.import_flat(con, buf.getvalue(), "k", dry_run=False)["saved"] == 1
    assert P.book(con)["CK-801"]["mode"] == "portion" and P.book(con)["CK-801"]["lines"][0]["no_loss"] == 1
    # a plain batch save (the BoM_Semi import) turns it back into a batch recipe
    P.save(con, "CK-801", 225, "ชิ้น", [{"ingredient_code": SALMON, "qty": 1, "unit": "ตัว"}], "k")
    assert P.book(con)["CK-801"]["mode"] == "batch" and P.book(con)["CK-801"]["lines"][0]["no_loss"] == 0


def test_the_portion_sheet_of_the_excel_file(con):
    M.set_ladder(con, SALMON, [_row("ตัว", 4.5, "กิโลกรัม", "avg", 4, 5)], "t")
    con.commit()
    data = _file([], [["CK-810", "แซลมอนแล่ชิ้น", SALMON, 1, "ตัว", 225, "ชิ้น", "หัว หนัง"],
                      ["CK-811", "ผิดหน่วย", SALMON, 1, "ก้อน", 100, "ชิ้น", None]])
    res = US.apply(con, US.parse(data), "ana", month="2026-09")
    ok, bad = res["portions"]
    assert res["problems"] == 1 and "แปลงหน่วย" in bad["problems"][0]
    assert ok["factor_base"] == 4500 and ok["approx"] and ok["base_per_piece"] == 20 and ok["cost_per_piece"] == pytest.approx(10.0)
    con.rollback()
    assert not con.execute("SELECT 1 FROM bc_ingredient WHERE code='CK-810'").fetchone()
    rows = US.portion_rows(con)
    assert rows == []                                                             # nothing saved yet
    P.save_portion(con, "CK-810", SALMON, 1, "ตัว", 225, "ชิ้น", "k", note="หัว หนัง", name="แซลมอนแล่ชิ้น")
    assert US.portion_rows(con) == [["CK-810", "แซลมอนแล่ชิ้น", SALMON, 1.0, "ตัว", 225.0, "ชิ้น", "หัว หนัง"]]


def test_portion_endpoints_and_order_plan_fields(client, monkeypatch):
    import main
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "pru-test-token")
    c = sqlite3.connect(str(main.DB_PATH))
    tables = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'bc%'")]
    snap = {t: c.execute(f"SELECT * FROM {t}").fetchall() for t in tables}
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, loss_pct) VALUES('PRU-S','ปลาทดสอบ','raw','กรัม',40)")
    c.execute("INSERT OR REPLACE INTO bc_ingredient_price(ingredient_code, month, price_per_base, source) VALUES('PRU-S','2026-09',0.5,'purchase_avg')")
    c.commit()
    c.close()
    hdr = {"X-Upload-Token": "pru-test-token"}
    try:
        assert client.post("/api/bom-center/units/ladder", headers=hdr, json={"code": "PRU-S", "rows": [
            {"unit": "ตัว", "qty_in_parent": 4.5, "parent_unit": "กิโลกรัม", "conv_kind": "avg", "range_min": 4, "range_max": 5}]}).status_code == 200
        q = {"raw_code": "PRU-S", "qty": 1, "unit": "ตัว", "pieces": 225, "month": "2026-09"}
        pv = client.get("/api/bom-center/prep/portion/preview", params=q).json()
        assert pv["approx"] and pv["base_per_piece"] == 20 and pv["cost_per_piece"] == pytest.approx(10.0)
        assert client.get("/api/bom-center/prep/portion/preview", params={**q, "pieces": 0}).status_code == 400
        body = {"code": "CK-PRU", "name": "ปลาแล่ชิ้น", "raw_code": "PRU-S", "qty": 1, "unit": "ตัว", "pieces": 225, "out_unit": "ชิ้น",
                "note": "หัว หนัง"}
        no = client.post("/api/bom-center/prep/portion", json=body)
        assert no.status_code in (200, 403)                                   # token env is set: the gate refuses a write without it
        assert no.status_code == 403 and no.headers.get("X-Denied") == "token"
        r = client.post("/api/bom-center/prep/portion", json=body, headers=hdr)
        assert r.status_code == 200 and r.json()["created"] and r.json()["cost_per_piece"] == pytest.approx(10.0)
        d = client.get("/api/bom-center/prep/CK-PRU").json()
        assert d["mode"] == "portion" and d["approx"] and d["per_base"] == pytest.approx(10.0)
        # the generic editor round trip keeps the portion recipe a portion recipe
        put = client.put("/api/bom-center/prep/CK-PRU", headers=hdr, json={"yield_qty": 225, "yield_unit": "ชิ้น", "mode": "portion",
                         "lines": [{"ingredient_code": "PRU-S", "qty": 1, "unit": "ตัว", "no_loss": 1}]})
        assert put.status_code == 200
        assert client.get("/api/bom-center/prep/CK-PRU").json()["lines"][0]["no_loss"] == 1
        bad = client.post("/api/bom-center/prep/portion", headers=hdr, json={**body, "pieces": -1})
        assert bad.status_code == 400
    finally:
        c = sqlite3.connect(str(main.DB_PATH))
        for t, rows in snap.items():
            c.execute(f"DELETE FROM {t}")
            if rows:
                c.executemany(f"INSERT INTO {t} VALUES({','.join('?' * len(rows[0]))})", rows)
        c.commit()
        c.close()


# ── permissions: the new routes follow the existing /units and sub-recipe rules ──
def test_new_routes_follow_the_unit_and_sub_recipe_write_rules():
    ana, kit = auth.ROLE_MODULES["analyst"], auth.ROLE_MODULES["innovation"]
    acc, br = auth.ROLE_MODULES["accounting"], auth.ROLE_MODULES["branch"]
    for path in ("/api/bom-center/units/ladder", "/api/bom-center/units/import"):
        assert auth.path_allowed("POST", path, ana), path
        assert auth.path_allowed("POST", path, {"bom.prices"}), path
        assert not auth.path_allowed("POST", path, acc), path
        assert not auth.path_allowed("POST", path, br), path
        assert not auth.path_allowed("POST", path, kit), path                      # same as POST /units: the kitchen does not set units
    for path in ("/api/bom-center/prep/portion",):
        assert auth.path_allowed("POST", path, kit), path                          # same as PUT /prep/{code}: bom.recipes
        assert auth.path_allowed("POST", path, ana), path
        assert not auth.path_allowed("POST", path, acc), path
        assert not auth.path_allowed("POST", path, br), path
    for path in ("/api/bom-center/units/ladder?code=X", "/api/bom-center/units/template.xlsx", "/api/bom-center/prep/portion/preview"):
        assert auth.path_allowed("GET", path.split("?")[0], ana), path
        assert auth.path_allowed("GET", path.split("?")[0], kit), path
