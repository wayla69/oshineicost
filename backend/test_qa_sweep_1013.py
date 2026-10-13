# -*- coding: utf-8 -*-
"""The fixes of the QA sweep of 2026-10-13 (PR "fix/qa-sweep"), one block per item.

H1 หน้างาน close steps use the period rule of the ใบสรุปตัดเบิก (test_bom_center_todo.py has the end-count cases)
H2 warehouse phone: ส่งตามที่สาขายืนยันทั้งหมด + the reason shown only when ส่งจริง differs
M1 stored UTC stamps are shown on the Bangkok clock
M2 the units export has no "แพ็ค = 1 แพ็ค" step and the import skips one
M3 the central list round-trips without losing "ที่มาของอัตราแปลง" and "งวดราคา"
M4 a supply role lands on its supply page on a phone, with a link back from the summary
M6 the rounds of หน้างาน are sorted by cutoff and named like the order page
L2 a carried-forward price is a note, not a problem
L4 the prep template has Thai headers (the old English ones still import)
"""
import contextlib
import io
import json
import re
from datetime import date

import openpyxl
import pytest

from bom_center import central as CE
from bom_center import fmt
from bom_center import prep as P
from bom_center import todo as TD
from bom_center import unit_sheet as US
from routers import bom_center as RB
from test_bom_center_batch2 import REPO, _fn, _node, _src, needs_node
from test_bom_center_s11 import _report, udon  # noqa: F401
from test_bom_center_s4b import con  # noqa: F401

CRAB = "13-01-009"
CORE = _src("js/01-core.js")
BRN = _src("js/14-branch.js")
OHB = _src("js/15-orderhub.js")


# ── M1 · times on the Bangkok clock ─────────────────────────────────────────────────────────────────────────
STAMPS = [
    ("2026-10-09T06:15:42", "9/10/2026 13:15"),            # naive = UTC, as master.now() writes it
    ("2026-10-09T06:15:42+00:00", "9/10/2026 13:15"),      # aware UTC
    ("2026-10-09T06:15:42Z", "9/10/2026 13:15"),
    ("2026-10-09T13:15:42+07:00", "9/10/2026 13:15"),      # already Bangkok
    ("2026-10-09T20:30:00", "10/10/2026 03:30"),           # the day rolls over
    ("2026-10-09 13:15", "9/10/2026 13:15"),               # a space-separated stamp is the reports' own Bangkok clock
    ("2026-10-09", "9/10/2026"),
    ("", ""),
]


@needs_node
def test_fmtat_converts_stored_utc_to_bangkok():
    pre = "\n".join([_fn(CORE, "_isoParts"), _fn(CORE, "fmtDay"), _fn(CORE, "_bkkClock"), _fn(CORE, "fmtAt"), _fn(CORE, "fmtDayAt")])
    got = _node(pre, json.dumps([s for s, _ in STAMPS]) + ".map(fmtAt)")
    assert got == [w for _, w in STAMPS]
    assert _node(pre, "[fmtDayAt('2026-10-09T18:30:00+00:00'), fmtDayAt('2026-10-09T10:30:00'), fmtDayAt('2026-10-09')]") == ["10/10/2026", "9/10/2026", "9/10/2026"]


@needs_node
def test_the_other_time_formatters_convert_too():
    pre = "\n".join([_fn(CORE, "_bkkClock"), _fn(BRN, "_brnIssTime"), _fn(OHB, "_ohbIsDt")])
    assert _node(pre, "[_brnIssTime('2026-10-09T06:15:42'), _brnIssTime('2026-10-09T06:15:42+00:00'), _brnIssTime(''), _ohbIsDt('2026-10-09T06:15:42'), _ohbIsDt('')]") \
        == ["13:15", "13:15", "", "9/10 13:15", ""]


def test_the_server_side_formatter_converts_too():
    for s, want in STAMPS[:7]:
        assert fmt.at_th(s) == want, s
    assert fmt.at_th(None) == "-" and fmt.day_at_th("2026-10-09T18:30:00+00:00") == "10/10/2026"


# ── H2 · the warehouse on a phone ─────────────────────────────────────────────────────────────────────────
@needs_node
def test_fill_all_sets_the_untouched_lines_to_what_the_branch_confirmed():
    fns = [_fn(OHB, n) for n in ("_ohbPRead", "_ohbPLine", "ohbPLineFull", "ohbPLineCheck", "_ohbDiffers", "_ohbPXToggle", "ohbFillAll")]
    pre = """
const els = {};
const mk = (id, value) => (els[id] = { id, value: String(value ?? ''), dataset: {}, style: {}, hidden: false, textContent: '' });
const document = { getElementById: id => els[id] || null };
function bmcStatus(id, ok, t) { els._status = t; }
let _ohbOrder = { lines: [{ code: 'A', confirmed_units: 5, processed_units: null }, { code: 'B', confirmed_units: 3, processed_units: null }, { code: 'C', confirmed_units: 2, processed_units: 1 }] };
mk('ohbP_A', 4); mk('ohbR_A', 'short'); mk('ohbN_A', 'x'); mk('ohbX_A', ''); els.ohbX_A.hidden = false; mk('ohbPE_A');
mk('ohbP_B', 9); els.ohbP_B.dataset.touched = '1'; mk('ohbR_B', ''); mk('ohbN_B', ''); mk('ohbX_B', ''); mk('ohbPE_B');
mk('ohbP_C', 1); mk('ohbR_C', 'x'); mk('ohbN_C', ''); mk('ohbX_C', ''); mk('ohbPE_C');
""" + "\n".join(fns)
    out = _node(pre, "(() => { ohbFillAll(); return [els.ohbP_A.value, els.ohbR_A.value, els.ohbN_A.value, els.ohbX_A.hidden, els.ohbP_B.value, els.ohbX_B.hidden, els.ohbP_C.value, els._status]; })()")
    assert out[:7] == ["5", "", "", True, "9", False, "2"]            # A and C reset to the confirmed quantity, B (typed by hand) kept and its reason field shown
    assert "ตั้งส่งจริงตามที่สาขายืนยัน 2 รายการ" in out[7] and "1 รายการที่แก้เอง" in out[7]
    # typing a different quantity shows the reason fields; typing the confirmed quantity hides them again
    typ = pre + "\n" + "\n".join([])
    out = _node(typ, "(() => { els.ohbP_C.value = '5'; ohbPLineCheck('C', true); const shown = !els.ohbX_C.hidden; els.ohbP_C.value = '2'; ohbPLineCheck('C', true); return [shown, els.ohbX_C.hidden, els.ohbP_C.dataset.touched]; })()")
    assert out == [True, True, "1"]


def test_the_warehouse_bar_is_sticky_and_the_button_exists():
    html = (REPO / "index.html").read_text(encoding="utf-8")
    assert re.search(r"\.ohb-actbar\s*\{\s*position:\s*sticky", html) and "body.supply-phone .ohb-actbar" in html
    assert "ส่งตามที่สาขายืนยันทั้งหมด" in OHB and "ohbFillAll()" in OHB
    assert "reason: r && !hidden" in OHB                                   # a reason hidden on the phone is not sent


# ── M2 · the units file ───────────────────────────────────────────────────────────────────────────────────
def _units_workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = US.SHEET
    ws.append([US.NOTE])
    ws.append(US.HEADERS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_units_export_has_no_identity_step_and_the_import_skips_one(udon):  # noqa: F811
    c = udon
    c.execute("UPDATE bc_ingredient SET base_unit='แพ็ค' WHERE code=?", (CRAB,))
    c.execute("DELETE FROM bc_unit_rule WHERE ingredient_code=?", (CRAB,))
    c.execute("INSERT INTO bc_unit_rule(ingredient_code, unit, factor_to_base, label, qty_in_parent, updated_at, updated_by)"
              " VALUES(?, 'แพ็ค', 1, '1/20 ของ กล่อง (20 แพ็ค) (รายการส่วนกลาง)', 1, 'n', 't')", (CRAB,))
    rows = US.ladder_rows(c, all_items=True)
    mine = [r for r in rows if r[0] == CRAB]
    assert len(mine) == 1 and mine[0][3] is None                           # an empty row to fill in, not "แพ็ค = 1 แพ็ค"
    # a file written by the old export still carries the step: it is skipped, nothing is reported, and the stored rule stays
    data = _units_workbook([[CRAB, "ปูอัด", "แพ็ค", "แพ็ค", 1, None, "แน่นอน", None, None, "x"]])
    res = US.apply(c, US.parse(data), "tester")
    assert res["problems"] == 0 and res["ladders"] == []
    assert c.execute("SELECT COUNT(*) FROM bc_unit_rule WHERE ingredient_code=? AND unit='แพ็ค'", (CRAB,)).fetchone()[0] == 1
    # a real step next to it is written, the identity rule is not touched, and a "แพ็ค = 2 แพ็ค" is still an error
    data = _units_workbook([[CRAB, "ปูอัด", "แพ็ค", "แพ็ค", 1, None, "แน่นอน", None, None, None],
                            [CRAB, "ปูอัด", "แพ็ค", "กล่อง", 20, None, "แน่นอน", None, None, None]])
    res = US.apply(c, US.parse(data), "tester")
    assert res["problems"] == 0 and [r["code"] for r in res["ladders"]] == [CRAB]
    assert {u for (u,) in c.execute("SELECT unit FROM bc_unit_rule WHERE ingredient_code=?", (CRAB,))} == {"แพ็ค", "กล่อง"}
    bad = US.apply(c, US.parse(_units_workbook([[CRAB, "ปูอัด", "แพ็ค", "แพ็ค", 2, None, "แน่นอน", None, None, None]])), "tester")
    assert bad["problems"] == 1


def test_the_units_file_round_trips_with_no_problem(udon):  # noqa: F811
    c = udon
    CE.apply_catalogue(c, CE.parse_report(c, _report()), "ss")
    before = US.ladder_rows(c, all_items=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = US.SHEET
    ws.append([US.NOTE])
    ws.append(US.HEADERS)
    for r in before:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    res = US.apply(c, US.parse(buf.getvalue()), "tester")
    assert res["problems"] == 0, res
    assert US.ladder_rows(c, all_items=True) == before


# ── M3 · the central list ───────────────────────────────────────────────────────────────────────────────
def _export_rows(c):
    wb = RB._central_workbook(CE.items(c))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), [list(r) for r in wb.active.iter_rows(min_row=3, values_only=True)]


def test_the_central_list_round_trips_its_own_columns(udon):  # noqa: F811
    c = udon
    CE.apply_catalogue(c, CE.parse_report(c, _report()), "ss")
    data, rows0 = _export_rows(c)
    assert rows0 and any(r[9] and r[9] != "ระบุในไฟล์" for r in rows0) and all(r[10] == "2026-09" for r in rows0 if r[4])
    p = CE.parse_report(c, data)
    assert p["month"] is None                                              # the list's own export carries no month
    plan = CE.plan_catalogue(c, p)["rows"]
    stored = {r[1]: r[9] for r in rows0}
    assert {r["code"]: r["factor_how"] for r in plan if r["status"] != "skip"} == {k: v for k, v in stored.items() if k in {x["code"] for x in plan}}
    CE.apply_catalogue(c, p, "ss")
    _, rows1 = _export_rows(c)
    assert rows1 == rows0                                                  # export -> import -> export: identical
    # a factor a person typed that differs IS "ระบุในไฟล์"
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    hdr = [x.value for x in ws[2]]
    for r in ws.iter_rows(min_row=3):
        if r[1].value == CRAB:
            r[hdr.index("1 หน่วยขาย = หน่วยฐาน")].value = 12345
    buf = io.BytesIO()
    wb.save(buf)
    plan = {r["code"]: r for r in CE.plan_catalogue(c, CE.parse_report(c, buf.getvalue()))["rows"]}
    assert plan[CRAB]["factor_how"] == "ระบุในไฟล์" and plan[CRAB]["base_per_sale"] == 12345


# ── M4 · a supply role on a phone ─────────────────────────────────────────────────────────────────────────
def test_a_supply_role_lands_on_its_supply_page_on_a_phone():
    assert re.search(r"else if \(_supplyPhoneRole\(\) && _phoneModeSaved\(\) !== 'summary'\) document\.body\.classList\.add\('force-desktop'\)", CORE)
    main = _src("js/03-tabs-main.js")
    assert "_supplyPhoneRole() && _phoneModeSaved() !== 'summary') supplyOpen()" in main
    assert 'id="msSupply"' in (REPO / "index.html").read_text(encoding="utf-8") and "async function msOpenSupply" in _src("js/09c-mobile.js")


@needs_node
def test_the_supply_phone_role_follows_the_landing_table():
    pre = ("const window = { AUTH: { enabled: true, role: 'salessupport' } }; const tabAllowed = t => t === 'supplyhome';\n"
           + re.search(r"const _SUPPLY_LANDING = \{.*?\n\};", CORE, re.S).group(0) + "\n" + _fn(CORE, "_supplyPhoneRole"))
    assert _node(pre, "_supplyPhoneRole()") is True
    for role, want in (("warehouse", True), ("accounting", True), ("management", False), ("marketing", False)):
        assert _node(pre.replace("'salessupport'", repr(role)), "_supplyPhoneRole()") is want, role


# ── M6 · the rounds of หน้างาน ────────────────────────────────────────────────────────────────────────────
def test_the_rounds_are_sorted_by_cutoff_and_named_like_the_order_page(udon):  # noqa: F811
    from test_bom_center_issues import NOW
    d = TD.build(udon, "2026-09", None, None, today=date(2026, 10, 5), now=NOW)
    due = [i for i in d["items"] if i["key"].startswith("order_due_")]
    assert len(due) == 3 and [i["cutoff_at"] for i in due] == sorted(i["cutoff_at"] for i in due)
    for i in due:
        rnd = i["key"].split(":")[0].replace("order_due_", "")
        assert i["title"].endswith(" · " + fmt.round_day_th(rnd, i["delivery"]))
    assert any(re.search(r" · รอบ [12] \(ส่ง \d+/\d+\)$", i["title"]) for i in due) and any(i["title"].endswith("แซลมอนสด (ส่ง " + i["delivery"][8:].lstrip("0") + "/10)") for i in due)


# ── L2 · carried-forward prices ──────────────────────────────────────────────────────────────────────────
def test_a_carried_forward_price_is_a_note_not_a_problem(udon, monkeypatch):  # noqa: F811
    @contextlib.contextmanager
    def _db():
        yield udon
    monkeypatch.setattr(RB, "get_db", _db)
    monkeypatch.setattr(RB, "_ensure", lambda con: None)
    rows = RB.bc_ingredients(month="2026-10", q=None, kind=None, issues=False)["rows"]
    assert rows
    assert not any(p.startswith("ราคาเก่า") for r in rows for p in r["problems"])
    assert all("notes" in r for r in rows)
    only = RB.bc_ingredients(month="2026-10", q=None, kind=None, issues=True)["rows"]
    assert all(r["problems"] for r in only)


# ── L4 · small items ─────────────────────────────────────────────────────────────────────────────────────
def test_the_prep_template_has_thai_headers_and_the_old_english_ones_still_import(con):  # noqa: F811
    assert len(P.FLAT_HEADERS_TH) == len(P.FLAT_COLS) and not any(re.search(r"[A-Za-z_]{4,}", h) for h in P.FLAT_HEADERS_TH)

    def book(header):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(list(header))
        ws.append(["CK-901", None, 500, "กรัม", 1, "V-007", None, 5, "กรัม", None, None, None])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    for header in (P.FLAT_HEADERS_TH, P.FLAT_COLS):
        con.execute("INSERT OR IGNORE INTO bc_ingredient(code, name_th, kind, base_unit) VALUES('CK-901', 'ซอสทดสอบ', 'prepared', 'กรัม')")      # a dry run rolls back
        out = P.import_flat(con, book(header), "kitchen", dry_run=True)
        assert out["codes"] == 1 and not out["errors"], header
    with pytest.raises(ValueError, match="รหัสสินค้ากึ่งสำเร็จ"):
        P.import_flat(con, book(["ก", "ข", "ค"]), "kitchen")


def test_the_round_columns_of_the_two_exports_use_the_round_names():
    src = (REPO / "backend" / "routers" / "bom_center.py").read_text(encoding="utf-8")
    assert src.count("_fc.ROUND_TH.get(str(") >= 2                           # consolidated.xlsx and compliance.xlsx


def test_the_register_and_dialogs_wording_is_formal():
    for rel in ("backend/bom_center/recipes.py", "backend/routers/bom_center.py", "js/13-bom-center.js", "js/15-orderhub.js"):
        assert "กรุณา" not in (REPO / rel).read_text(encoding="utf-8"), rel
    assert "(ตัวอย่าง ลบก่อนนำเข้า)" in (REPO / "backend/bom_center/recipes.py").read_text(encoding="utf-8")


def test_no_console_403_from_roles_without_the_module():
    assert "if (!hasModule('cost.analysis')) return;" in _src("js/05-pnl-settings.js")
    assert "['branch.stock', 'supply.ordering', 'bom.center', 'ops.report'].some(hasModule)" in CORE


@needs_node
def test_every_js_slice_parses():
    import subprocess
    from test_bom_center_batch2 import NODE
    bad = []
    for p in sorted((REPO / "js").glob("*.js")):
        r = subprocess.run([NODE, "--check", str(p)], capture_output=True, text=True, encoding="utf-8")
        if r.returncode:
            bad.append(f"{p.name}: {r.stderr.strip().splitlines()[-1][:160] if r.stderr.strip() else ''}")
    assert not bad, bad
