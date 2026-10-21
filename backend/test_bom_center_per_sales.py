# -*- coding: utf-8 -*-
"""BoM Center — PR-V: ingredient usage per ฿100,000 of sales, a second way to read the same history.

The promises:
  * the ratio is the forecast's own usage (purchase quantities, base unit) over the sales of the same days, ex VAT, per 100,000;
    a week is a whole block of 7 days; the swing of the weekly ratios decides "steady" (<= 15 %); fewer than two weeks = no verdict;
  * "ปกติ" is the sales the forecast assumes for the cover days (weekday-fair mean x trend x season), so it moves with them;
  * a branch types the sales it expects in baht: the % is (expected / normal - 1), held to -50..+100, and the SAME buffer field drives
    the suggestion; the baht is kept on the order;
  * every plan row carries its ratio; the report is the last N weeks against the N before, per branch and for the chain
    (sum of usage / sum of sales); a branch login sees its own branch only;
  * a new branch (fewer than 14 trading days) with a reference branch is planned from that branch's ratio x the sales it expects, and
    cannot order without saying what it expects; without a reference branch nothing changes.
"""
import io
import sqlite3
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook
from conftest import dl_name

import auth
from bom_center import forecast as FC
from bom_center import orders as OR
from bom_center import per_sales as PS
from bom_center.clock import BKK
from test_bom_center_issue_summary import TOKEN, _login, _restore, _snap, auth_on  # noqa: F401
from test_bom_center_s4b import PATTERN, con  # noqa: F401  (branch 1: crab 45 g on sku 10, salmon 20 g on sku 11, 1 Sep - 5 Oct 2026)

CRAB, SALMON = "13-01-009", "02-01-001"
WED = date(2026, 10, 7)                                   # round 1's delivery (cutoff Monday 5 Oct 13:15)
NOW = datetime(2026, 10, 5, 8, 0, tzinfo=BKK)
AS_OF = date(2026, 10, 5)
WEEK_SALES = 132 * 100 + 132 * 200                        # 132 plates of each dish a week at ฿100 and ฿200


@pytest.fixture(autouse=True)
def _flat_season(monkeypatch):
    """The season / festival factor is the calendar's business (and may be ≠ 1 in October): these tests fix it at 1.0,
    except the one that proves "ปกติ" follows it."""
    monkeypatch.setattr(FC, "season_factor", lambda *a, **k: {"factor": 1.0, "mode": "off", "reason": ""})


@pytest.fixture
def sold(con):  # noqa: F811
    """The forecast fixture with prices on the dishes: sku 10 = ฿100 a plate, sku 11 = ฿200 (subtotal_ex_vat = qty x price)."""
    con.execute("UPDATE sales_sku_daily SET subtotal_ex_vat = qty * 100 WHERE sku_id=10")
    con.execute("UPDATE sales_sku_daily SET subtotal_ex_vat = qty * 200 WHERE sku_id=11")
    con.commit()
    return con


def _days(a, b):
    return [a + timedelta(days=i) for i in range((b - a).days + 1)]


# ── 1 · the ratio ───────────────────────────────────────────────────────────────────────────
def test_ratio_is_the_forecasts_usage_over_the_sales_per_100k(sold):
    r = PS.ratio(sold, 1, date(2026, 9, 1), AS_OF, "2026-10")
    assert r["sales"] == 5 * WEEK_SALES and r["days"] == 35 and r["weeks"] == 5
    crab, sal = r["items"][CRAB], r["items"][SALMON]
    assert crab["usage"] == pytest.approx(5 * 132 * 45)                       # the forecast's own purchase quantity, in grams
    assert crab["per_100k"] == pytest.approx(132 * 45 / WEEK_SALES * 1e5)     # 15,000 g per ฿100,000
    assert crab["per_100k"] == pytest.approx(15000.0) and crab["unit"] == "กรัม"
    assert sal["per_100k"] == pytest.approx(132 * 20 / WEEK_SALES * 1e5)
    assert crab["weekly"] == [pytest.approx(15000.0)] * 5 and crab["swing_pct"] == 0.0 and crab["steady"] is True
    # it is literally usage_history over sales: no second definition of usage
    use, _d, _s = FC.usage_history(sold, 1, date(2026, 9, 1), AS_OF, "2026-10")
    assert crab["usage"] == pytest.approx(sum(use[CRAB].values()))


def test_vat_is_never_in_the_sales(sold):
    base = PS.ratio(sold, 1, date(2026, 9, 1), AS_OF, "2026-10")
    sold.execute("ALTER TABLE sales_sku_daily ADD COLUMN total_inc_vat REAL")
    sold.execute("UPDATE sales_sku_daily SET total_inc_vat = subtotal_ex_vat * 1.07")
    sold.commit()
    again = PS.ratio(sold, 1, date(2026, 9, 1), AS_OF, "2026-10")
    assert again["sales"] == base["sales"] == 5 * WEEK_SALES                  # ex VAT, whatever the other columns hold
    assert again["items"][CRAB]["per_100k"] == pytest.approx(base["items"][CRAB]["per_100k"])
    assert PS.sales_by_day(sold, 1, date(2026, 9, 1), date(2026, 9, 1)) == {"2026-09-01": pytest.approx(10 * 100 + 10 * 200)}


def test_weekly_swing_and_the_steady_flag(sold):
    # one week of crab sold twice as much: 22,500 g against 15,000 g in the other four weeks
    sold.execute("UPDATE sales_sku_daily SET qty = qty * 2, subtotal_ex_vat = subtotal_ex_vat * 2"
                 " WHERE sku_id=10 AND day BETWEEN '2026-09-15' AND '2026-09-21'")
    sold.commit()
    it = PS.ratio(sold, 1, date(2026, 9, 1), AS_OF, "2026-10")["items"][CRAB]
    assert it["weekly"] == [pytest.approx(15000.0), pytest.approx(15000.0), pytest.approx(22500.0),
                            pytest.approx(15000.0), pytest.approx(15000.0)]
    mean = (4 * 15000 + 22500) / 5
    sd = (((mean - 15000) ** 2 * 4 + (22500 - mean) ** 2) / 5) ** 0.5
    assert it["swing_pct"] == pytest.approx(sd / mean * 100, abs=0.06) and it["swing_pct"] > PS.STEADY_PCT
    assert it["steady"] is False
    assert PS.STEADY_PCT == 15.0
    # the salmon did not move: its ratio changes only through the bigger sales that week, still inside 15 %
    sal = PS.ratio(sold, 1, date(2026, 9, 1), AS_OF, "2026-10")["items"][SALMON]
    assert sal["swing_pct"] == pytest.approx(10.5, abs=0.1) and sal["steady"] is True


def test_fewer_than_two_whole_weeks_gives_no_verdict(sold):
    one = PS.ratio(sold, 1, date(2026, 9, 26), AS_OF, "2026-10")               # 10 days = one whole week
    assert one["weeks"] == 1 and one["items"][CRAB]["steady"] is None and one["items"][CRAB]["swing_pct"] is None
    assert one["items"][CRAB]["per_100k"] == pytest.approx(15000.0)            # the ratio itself is still there
    five = PS.ratio(sold, 1, date(2026, 10, 1), AS_OF, "2026-10")
    assert five["weeks"] == 0 and five["items"][CRAB]["steady"] is None
    assert PS.week_blocks(date(2026, 9, 1), AS_OF)[-1][-1] == "2026-10-05" and len(PS.week_blocks(date(2026, 9, 1), AS_OF)) == 5


# ── 2 · "ปกติ" ─────────────────────────────────────────────────────────────────────────────
def test_normal_sales_is_the_weekday_fair_mean(sold):
    n = PS.normal_sales(sold, 1, date(2026, 10, 7), date(2026, 10, 9), factor=1.0, as_of=AS_OF)
    # Wed 10 plates, Thu 12, Fri 20 of both dishes at ฿300 a pair
    assert n["total"] == pytest.approx(10 * 300 + 12 * 300 + 20 * 300) and n["trend"] == 1.0
    assert [d["normal"] for d in n["days"]] == [3000.0, 3600.0, 6000.0]
    assert PS.normal_sales(sold, 1, date(2026, 10, 7), date(2026, 10, 9), factor=1.5, as_of=AS_OF)["total"] == pytest.approx(12600 * 1.5)


def test_normal_sales_carries_the_forecasts_trend(sold):
    sold.execute("UPDATE sales_sku_daily SET subtotal_ex_vat = subtotal_ex_vat * 1.2 WHERE day > '2026-09-21'")
    sold.commit()
    n = PS.normal_sales(sold, 1, date(2026, 10, 7), date(2026, 10, 9), factor=1.0, as_of=AS_OF)
    # each weekday has 5 samples, 2 of them in the last 14 days (x1.2): mean x 1.08; the last 14 days against the 14 before: x 1.2
    assert n["trend"] == pytest.approx(1.2) and n["total"] == pytest.approx(12600 * 1.08 * 1.2)


def test_suggest_and_plan_carry_the_same_normal_and_the_season(sold, monkeypatch):
    out = FC.suggest(sold, 1, "1", date(2026, 10, 6), season="off", as_of=AS_OF)
    assert out["sales_normal"] == pytest.approx(12600.0)                      # the cover of round 1 (Wed 7 - Fri 9)
    crab = {r["code"]: r for r in out["rows"]}[CRAB]
    assert crab["per_sales"]["per_100k"] == pytest.approx(15000.0) and crab["per_sales"]["steady"] is True
    # the season factor the forecast uses scales "ปกติ" with it
    monkeypatch.setattr(FC, "season_factor", lambda *a, **k: {"factor": 1.25, "mode": "test", "reason": ""})
    s2 = FC.suggest(sold, 1, "1", date(2026, 10, 6), season="auto", as_of=AS_OF)
    assert s2["sales_normal"] == pytest.approx(12600 * 1.25)
    assert {r["code"]: r for r in s2["rows"]}[CRAB]["forecast"] == pytest.approx(42 * 45 * 1.25)


# ── 3 · baht -> % ────────────────────────────────────────────────────────────────────────────
def test_baht_to_percent_and_the_clamp():
    assert PS.pct_from_sales(15120, 12600) == {"pct": 20.0, "clamped": False, "message": ""}
    assert PS.pct_from_sales(12600, 12600)["pct"] == 0.0
    hi = PS.pct_from_sales(12600 * 3, 12600)
    assert hi["pct"] == 100.0 and hi["clamped"] and "+100%" in hi["message"] and "ระบบใช้" in hi["message"]
    lo = PS.pct_from_sales(12600 * 0.2, 12600)
    assert lo["pct"] == -50.0 and lo["clamped"]
    assert PS.pct_from_sales(100, 0)["pct"] is None and PS.pct_from_sales("x", 5)["pct"] is None and PS.pct_from_sales(-1, 5)["pct"] is None
    assert PS.sales_clean("") is None and PS.sales_clean(0) is None and PS.sales_clean("15,120") == 15120.0
    for bad in ("abc", -5):
        with pytest.raises(ValueError):
            PS.sales_clean(bad)


def test_the_baht_is_kept_on_the_order_and_the_percent_drives_the_suggestion(sold):
    p = OR.plan(sold, 1, "1", WED, now=NOW)
    assert p["sales_normal"] == pytest.approx(12600.0) and p["expected_sales"] is None and p["ref_mode"] is None
    pct = PS.pct_from_sales(15120, p["sales_normal"])["pct"]
    out = OR.save_flow(sold, 1, "1", WED, [], "u", buffer={"pct": pct, "reason": "event", "sales": 15120}, now=NOW)
    sold.commit()
    o = OR.detail(sold, out["order_id"])
    assert o["buffer_pct"] == 20.0 and o["buffer_sales"] == 15120.0 and o["buffer_reason"] == "event"
    again = OR.plan(sold, 1, "1", WED, now=NOW)                               # the page reopens on the order's own figures
    assert again["buffer"]["pct"] == 20.0 and again["expected_sales"] == 15120.0
    crab0, crab1 = ({r["code"]: r for r in x["rows"]}[CRAB] for x in (p, again))
    assert crab1["forecast"] == pytest.approx(crab0["forecast"] * 1.2)        # the % still drives the suggestion
    # a bad figure is refused, and clearing the buffer clears the baht
    with pytest.raises(ValueError) as e:
        OR.save_flow(sold, 1, "1", WED, [], "u", buffer={"pct": 0, "sales": "abc"}, now=NOW)
    assert "ตัวเลข" in str(e.value)
    sold.rollback()
    OR.save_flow(sold, 1, "1", WED, [], "u", buffer={"pct": 0}, now=NOW)
    assert OR.detail(sold, out["order_id"])["buffer_sales"] is None


def test_plan_payload_keys(sold):
    p = OR.plan(sold, 1, "1", WED, now=NOW)
    assert {"sales_normal", "expected_sales", "ref_mode", "buffer"} <= set(p)
    r = {x["code"]: x for x in p["rows"]}[CRAB]
    assert set(r["per_sales"]) == {"per_100k", "steady", "swing_pct"} and r["ref_label"] is None


# ── 4 · a new branch with a reference branch ──────────────────────────────────────────────
@pytest.fixture
def newb(sold):
    """Branch 2 opened five days ago (5 trading days, the same crab dish); branch 1 is its reference."""
    sold.executemany("INSERT INTO sales_sku_daily VALUES(2, ?, 10, 5, 500)", [(d.isoformat(),) for d in _days(date(2026, 10, 1), AS_OF)])
    PS.set_reference(sold, "t", 2, 1)
    sold.commit()
    return sold


def test_reference_branch_setting(sold):
    assert PS.reference_for(sold, 2) is None
    out = PS.set_reference(sold, "sales", 2, 1)
    assert out == {"branch_id": 2, "ref_branch_id": 1} and PS.reference_for(sold, 2) == 1 and PS.reference_map(sold) == {2: 1}
    raw = sold.execute("SELECT value FROM app_settings WHERE key='bc_ref_branch'").fetchone()[0]
    assert raw == '{"2": 1}'
    assert sold.execute("SELECT action, old_value, new_value FROM bc_audit WHERE action='ref_branch_set'").fetchone() == ("ref_branch_set", None, "1")
    for args in ((2, 2), (2, 99), (99, 1)):
        with pytest.raises(ValueError):
            PS.set_reference(sold, "sales", *args)
    assert PS.set_reference(sold, "sales", 2, None)["ref_branch_id"] is None and PS.reference_map(sold) == {}


def test_a_new_branch_must_say_what_it_expects_to_sell(newb):
    p = OR.plan(newb, 2, "1", WED, now=NOW)
    rm = p["ref_mode"]
    assert rm["active"] and rm["needs_sales"] and rm["branch_id"] == 1 and rm["name"] == "BranchA" and rm["own_days"] == 5
    assert rm["message"] == "ยังไม่มีประวัติการขายพอ กรอกยอดขายที่คาดของรอบนี้" and "กรุณา" not in rm["message"]
    assert rm["ref_sales_normal"] == pytest.approx(12600.0)                    # the reference branch's "ปกติ" for the same cover
    assert p["rows"] == [] and p["sales_normal"] is None
    with pytest.raises(ValueError) as e:                                       # nothing is saved without it
        OR.save_flow(newb, 2, "1", WED, [], "u", now=NOW)
    assert str(e.value) == rm["message"]
    newb.rollback()
    assert newb.execute("SELECT COUNT(*) FROM bc_order WHERE branch_id=2").fetchone()[0] == 0


def test_a_new_branch_is_planned_from_the_reference_ratio(newb):
    p = OR.plan(newb, 2, "1", WED, now=NOW, expected_sales=100000)
    assert p["ref_mode"]["active"] and not p["ref_mode"]["needs_sales"] and p["ref_mode"]["expected_sales"] == 100000
    crab = {r["code"]: r for r in p["rows"]}[CRAB]
    assert crab["forecast"] == pytest.approx(15000.0)                          # 15,000 g per ฿100,000 x ฿100,000
    assert crab["safety"] == 0 and crab["order_units"] == 15.0 and crab["order_unit"] == "pack"
    assert crab["ref_label"] == "คำนวณจากสาขาอ้างอิง BranchA" and crab["per_sales"]["per_100k"] == pytest.approx(15000.0)
    assert SALMON not in {r["code"] for r in p["rows"]}                        # the salmon round is its own
    sal = {r["code"]: r for r in OR.plan(newb, 2, "salmon", date(2026, 10, 12), now=NOW, expected_sales=100000)["rows"]}[SALMON]
    assert sal["forecast"] == pytest.approx(132 * 20 / WEEK_SALES * 1e5)
    # half the expected sales = half the order
    half = {r["code"]: r for r in OR.plan(newb, 2, "1", WED, now=NOW, expected_sales=50000)["rows"]}[CRAB]
    assert half["forecast"] == pytest.approx(7500.0)
    # a % on top does not apply: the expected sales are the expectation
    up = OR.plan(newb, 2, "1", WED, now=NOW, expected_sales=100000, buffer_pct=50)
    assert up["buffer"]["pct"] == 0 and {r["code"]: r for r in up["rows"]}[CRAB]["forecast"] == pytest.approx(15000.0)


def test_the_new_branch_order_keeps_the_expected_sales(newb):
    out = OR.save_flow(newb, 2, "1", WED, [], "u", buffer={"pct": 0, "sales": 100000}, now=NOW)
    o = OR.detail(newb, out["order_id"])
    assert o["buffer_sales"] == 100000.0 and o["buffer_pct"] is None
    assert {l["code"]: l for l in o["lines"]}[CRAB]["confirmed_units"] == 15.0
    assert OR.plan(newb, 2, "1", WED, now=NOW)["expected_sales"] == 100000.0   # reopened on the order's own figure


def test_nothing_changes_without_a_reference_or_with_enough_history(newb):
    PS.set_reference(newb, "t", 2, None)
    p = OR.plan(newb, 2, "1", WED, now=NOW)
    assert p["ref_mode"] is None and {r["code"]: r for r in p["rows"]}[CRAB]["ref_label"] is None     # today's behaviour
    PS.set_reference(newb, "t", 1, 2)                                          # branch 1 has 35 trading days: the reference is ignored
    p1 = OR.plan(newb, 1, "1", WED, now=NOW)
    assert p1["ref_mode"] is None and {r["code"]: r for r in p1["rows"]}[CRAB]["forecast"] == pytest.approx(42 * 45)


# ── 5 · the report, the chain total and branch pinning ──────────────────────────────────────
def test_report_chain_total_is_sum_of_usage_over_sum_of_sales(sold):
    sold.executemany("INSERT INTO sales_sku_daily VALUES(2, ?, 10, 5, 500)", [(d.isoformat(),) for d in _days(date(2026, 9, 1), AS_OF)])
    sold.commit()
    r = PS.report(sold, 2, today=date(2026, 10, 6))
    assert r["end"] == "2026-10-05" and r["current"] == {"from": "2026-09-22", "to": "2026-10-05"}
    assert r["previous"] == {"from": "2026-09-08", "to": "2026-09-21"} and [b["branch_id"] for b in r["branches"]] == [1, 2]
    b1, b2 = r["branches"]
    c1 = {x["code"]: x for x in b1["items"]}[CRAB]
    # in packs of 1 kg (the unit the branch counts in): 15 per ฿100,000, unchanged against the weeks before
    assert c1["unit"] == "แพ็ค 1 กก." and c1["per_100k"] == pytest.approx(15.0) and c1["prev_per_100k"] == pytest.approx(15.0)
    assert c1["change_pct"] == 0.0 and c1["steady"] is True and c1["swing_pct"] == 0.0
    assert c1["value_per_100k"] == pytest.approx(15000 * 0.15)                 # the item's price per gram (฿0.15) x grams
    assert {x["code"]: x for x in b2["items"]}[CRAB]["per_100k"] == pytest.approx(45.0)
    # items are listed by value, highest first
    assert [x["value_per_100k"] for x in b1["items"]] == sorted((x["value_per_100k"] or 0) for x in b1["items"])[::-1]
    # chain = sum of usage / sum of sales (not an average of the two branch ratios)
    use = 2 * (132 * 45 + 35 * 45)
    sales = 2 * (WEEK_SALES + 3500)
    chain = {x["code"]: x for x in r["chain"]["items"]}[CRAB]
    assert r["chain"]["sales"] == sales and chain["per_100k"] == pytest.approx(use / sales * 1e5 / 1000)
    assert chain["per_100k"] != pytest.approx((15.0 + 45.0) / 2)
    # a single branch asked for: no chain; the cache key never mixes the two answers
    one = PS.report(sold, 2, [2], today=date(2026, 10, 6))
    assert [b["branch_id"] for b in one["branches"]] == [2]


def test_report_without_a_previous_window_says_so(sold):
    r = PS.report(sold, 8, [1], today=date(2026, 10, 6))                       # 8 weeks back reaches before the Grand Menu
    assert r["previous"] is None and r["current"]["from"] == "2026-09-01"       # the old menu never counts (the 999 plates of 20 Aug)
    it = {x["code"]: x for x in r["branches"][0]["items"]}[CRAB]
    assert it["per_100k"] == pytest.approx(15.0)


def test_report_cache_is_dropped_with_the_usage_cache(sold, monkeypatch):
    monkeypatch.setattr(FC, "USAGE_CACHE_TTL", 600)
    a = PS.report(sold, 2, [1], today=date(2026, 10, 6))
    assert PS.report(sold, 2, [1], today=date(2026, 10, 6)) is a                # served from the cache
    FC.clear_usage_cache()                                                      # what every recipe-side write calls
    assert PS.report(sold, 2, [1], today=date(2026, 10, 6)) is not a


def _seed_http(db_path):
    """Branch 1: ฿1,000 and 10 plates a day, branch 2: ฿1,000 and 20 plates a day, 1 Sep - 4 Oct; the dish holds 0.5 g of T-PS1."""
    c = sqlite3.connect(str(db_path))
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, family, active) VALUES('T-PS1', 'ของทดสอบ ต่อยอดขาย', 'raw', 'กรัม', 'F', 1)")
    c.execute("INSERT OR REPLACE INTO bc_ingredient_price(ingredient_code, month, price_per_base, source) VALUES('T-PS1', '2026-09', 2.0, 'purchase_avg')")
    c.execute("INSERT INTO bom_ingredients(sku_id, seq, ingredient_code, ingredient_name, qty_per_serving, unit, price_per_unit,"
              " ingredient_cost, ingredient_type) VALUES(2, 1, 'T-PS1', 'ของทดสอบ ต่อยอดขาย', 0.5, 'กรัม', 0, 0, 'x')")
    for d in _days(date(2026, 9, 1), date(2026, 10, 4)):
        c.execute("INSERT INTO sales_sku_daily VALUES(1, ?, 2, 10, 1000)", (d.isoformat(),))
        c.execute("INSERT INTO sales_sku_daily VALUES(2, ?, 2, 20, 1000)", (d.isoformat(),))
    c.commit()
    c.close()


def test_http_report_workbook_pinning_and_the_reference_setting(client, app_db, auth_on, monkeypatch):  # noqa: F811
    import main
    monkeypatch.setattr(PS, "bkk_today", lambda: date(2026, 10, 12))
    from routers.bom_center import PER_SALES_COMPARE, PER_SALES_SHEET, _per_sales_headers
    client.get("/api/bom-center/hub/queue")                                                           # creates the bc_* tables
    snap = _snap(main.DB_PATH)
    _seed_http(main.DB_PATH)
    try:
        mg, ss, acc = _login("management1"), _login("salessupport1"), _login("accounting1")
        mg.post("/api/auth/users", headers=TOKEN, json={"username": "br_ps_test", "password": "Temp@12345",
                                                         "role": "branch", "branch": "BranchA"})
        br = _login("br_ps_test", "Temp@12345")
        j = mg.get("/api/bom-center/per-sales?weeks=2").json()
        assert j["end"] == "2026-10-04" and j["current"] == {"from": "2026-09-21", "to": "2026-10-04"}
        bb = {b["branch"]: b for b in j["branches"]}
        a, b = ({x["code"]: x for x in bb[n]["items"]}["T-PS1"] for n in ("BranchA", "BranchB"))
        assert a["per_100k"] == pytest.approx(500.0) and b["per_100k"] == pytest.approx(1000.0)
        assert a["value_per_100k"] == pytest.approx(1000.0) and a["steady"] is True
        assert {x["code"]: x for x in j["chain"]["items"]}["T-PS1"]["per_100k"] == pytest.approx(750.0)   # (5 + 10) / 2,000 x 100,000
        assert ss.get("/api/bom-center/per-sales?weeks=2").status_code == 200 and acc.get("/api/bom-center/per-sales?weeks=2").status_code == 200
        assert mg.get("/api/bom-center/per-sales?weeks=0").status_code == 422
        # a branch login: its own branch, no chain, never another branch
        mine = br.get("/api/bom-center/per-sales?weeks=2").json()
        assert [x["branch"] for x in mine["branches"]] == ["BranchA"] and mine["chain"] is None and mine["scoped"] is True
        assert br.get("/api/bom-center/per-sales?weeks=2&branch_id=1").status_code == 200
        assert br.get("/api/bom-center/per-sales?weeks=2&branch_id=2").status_code == 403
        assert br.get("/api/bom-center/per-sales.xlsx?branch_id=2").status_code == 403
        assert [x["branch"] for x in mg.get("/api/bom-center/per-sales?weeks=2&branch_id=2").json()["branches"]] == ["BranchB"]
        # the workbook
        r = mg.get("/api/bom-center/per-sales.xlsx")
        assert r.status_code == 200 and dl_name(r) == "ใช้วัตถุดิบต่อยอดขาย_ทุกสาขา_2026-10-04.xlsx"
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["ใช้วัตถุดิบต่อยอดขาย", "เทียบสาขา"] == [PER_SALES_SHEET, PER_SALES_COMPARE]
        rows = list(wb[PER_SALES_SHEET].iter_rows(values_only=True))
        assert list(rows[1]) == ["สาขา", "รหัส", "วัตถุดิบ", "หน่วย", "ต่อยอดขาย ฿100,000 (4 สัปดาห์ล่าสุด)", "4 สัปดาห์ก่อนหน้า",
                                 "เปลี่ยนแปลง %", "แกว่งต่อสัปดาห์ %", "มูลค่าต่อยอดขาย ฿100,000", "การใช้ตามสูตรเทียบค่ากลาง %", "ยอดซื้อเทียบค่ากลาง %", "หมายเหตุ"] == _per_sales_headers(4)
        assert [x[0] for x in rows[2:] if x[1] == "T-PS1"] == ["BranchA", "BranchB", "ทั้งเครือ"]
        chain = next(x for x in rows[2:] if x[0] == "ทั้งเครือ" and x[1] == "T-PS1")
        assert chain[4] == pytest.approx(750.0) and chain[3] == "กรัม"
        cmp_ = list(wb[PER_SALES_COMPARE].iter_rows(values_only=True))
        assert list(cmp_[1]) == ["รหัส", "วัตถุดิบ", "หน่วย", "BranchA", "BranchB", "ทั้งเครือ"]
        assert next(x for x in cmp_[2:] if x[0] == "T-PS1")[3:] == pytest.approx((500.0, 1000.0, 750.0))
        # a branch login: its own column only, no chain row
        mine_x = load_workbook(io.BytesIO(br.get("/api/bom-center/per-sales.xlsx").content))
        assert {x[0] for x in mine_x[PER_SALES_SHEET].iter_rows(min_row=3, values_only=True)} == {"BranchA"}
        assert list(next(mine_x[PER_SALES_COMPARE].iter_rows(min_row=2, max_row=2, values_only=True))) == ["รหัส", "วัตถุดิบ", "หน่วย", "BranchA"]
        # the reference branch setting: the sales team and the analyst; not a branch, not accounting
        url = "/api/bom-center/order-calendar/ref-branch"
        assert br.post(url, headers=TOKEN, json={"branch_id": 2, "ref_branch_id": 1}).status_code == 403
        assert acc.post(url, headers=TOKEN, json={"branch_id": 2, "ref_branch_id": 1}).status_code == 403
        assert ss.post(url, headers=TOKEN, json={"branch_id": 2, "ref_branch_id": 2}).status_code == 400
        assert ss.post(url, headers=TOKEN, json={"branch_id": 2, "ref_branch_id": 1}).json() == {"branch_id": 2, "ref_branch_id": 1}
        cal = ss.get("/api/bom-center/order-calendar").json()
        assert {x["branch_id"]: x["ref_branch_id"] for x in cal["rows"]} == {1: None, 2: 1} and cal["new_branch_days"] == 14
        assert ss.post(url, headers=TOKEN, json={"branch_id": 2, "ref_branch_id": None}).status_code == 200
    finally:
        c = sqlite3.connect(str(main.DB_PATH))
        c.execute("DELETE FROM sales_sku_daily WHERE day >= '2026-09-01'")
        c.execute("DELETE FROM bom_ingredients WHERE ingredient_code='T-PS1'")
        c.execute("DELETE FROM app_settings WHERE key='bc_ref_branch'")
        c.commit()
        c.close()
        _restore(main.DB_PATH, snap)


def test_the_plan_route_takes_the_expected_sales(client, app_db, auth_on):  # noqa: F811
    import main
    client.get("/api/bom-center/hub/queue")
    snap = _snap(main.DB_PATH)
    try:
        mg = _login("management1")
        p = mg.get("/api/bom-center/order/plan?branch_id=1&round=1&expected_sales=50000")
        assert p.status_code == 200 and {"sales_normal", "expected_sales", "ref_mode"} <= set(p.json())
        assert mg.get("/api/bom-center/order/plan?branch_id=1&round=1&expected_sales=abc").status_code == 422
    finally:
        _restore(main.DB_PATH, snap)


# ── 6 · who may read and set ─────────────────────────────────────────────────────────────────
def test_auth_matrix_follows_the_issue_reports_rule():
    for role, mods in auth.ROLE_MODULES.items():
        for path in ("/api/bom-center/per-sales", "/api/bom-center/per-sales.xlsx"):
            assert auth.path_allowed("GET", path, mods) == auth.path_allowed("GET", "/api/bom-center/issue-reports", mods), role
    br, ss, an = (auth.ROLE_MODULES[r] for r in ("branch", "salessupport", "analyst"))
    assert auth.path_allowed("GET", "/api/bom-center/per-sales", br) and auth.path_allowed("GET", "/api/bom-center/per-sales", ss)
    assert not auth.path_allowed("GET", "/api/bom-center/per-sales", auth.ROLE_MODULES["franchise"])
    assert not auth.path_allowed("GET", "/api/bom-center/per-sales", auth.ROLE_MODULES["warehouse"])
    assert not auth.path_allowed("POST", "/api/bom-center/per-sales", br)                       # nothing to write there (no POST route exists)
    ref = "/api/bom-center/order-calendar/ref-branch"
    assert auth.path_allowed("POST", ref, ss) and auth.path_allowed("POST", ref, an)
    assert not auth.path_allowed("POST", ref, br) and not auth.path_allowed("POST", ref, auth.ROLE_MODULES["accounting"])


def test_a_week_the_branch_barely_traded_is_not_a_week(newb):
    five = PS.ratio(newb, 2, date(2026, 9, 29), AS_OF, "2026-10")            # sold on 1-5 Oct: five of the seven days
    assert five["weeks"] == 1 and five["items"][CRAB]["per_100k"] == pytest.approx(45000.0)
    newb.execute("DELETE FROM sales_sku_daily WHERE branch_id=2 AND day='2026-10-01'")
    newb.commit()
    four = PS.ratio(newb, 2, date(2026, 9, 29), AS_OF, "2026-10")            # four days: the ratio is there, the week is not
    assert four["weeks"] == 0 and four["items"][CRAB]["swing_pct"] is None and four["items"][CRAB]["steady"] is None
