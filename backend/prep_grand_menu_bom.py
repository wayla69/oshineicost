# -*- coding: utf-8 -*-
"""
prep_grand_menu_bom.py — turn the new Grand Menu workbooks (effective 1 Sep 2026)
into clean, import-ready CSVs + a validation report.

Inputs (defaults point at ~/Downloads):
  Final BOM Grand Menu.xlsx    — recipe sheets BF_*/AC_* (ingredient lines),
                                 SUM (dish master: Item Code BF/AL, tier, cost),
                                 DATAP (ingredient master: code, unit, price, loss%)
  Master Item Grand Menu.xlsx  — flat item list (code, tier, Thai name, cost/serve)

Outputs (BoM/GrandMenu_2026-09/):
  menu_master.csv        one row per POS item code (BF + AL)   -> skus / sku_cost
  bom_lines.csv          one row per recipe ingredient line     -> bom_ingredients
  ingredient_master.csv  DATAP, deduped                         -> reference
  issues.csv             every exception that needs a human     -> send back to ops

NOTHING is written to the database. This is a prepare-and-validate step only.

Usage:
  python backend/prep_grand_menu_bom.py
  python backend/prep_grand_menu_bom.py --bom "<path>" --master "<path>" --out "<dir>"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent.resolve()
ROOT = HERE.parent
DL = Path.home() / "Downloads"

try:
    import ingest_bom
    classify_ingredient = ingest_bom.classify_ingredient
except Exception:                                    # standalone use
    ingest_bom = None

    def classify_ingredient(code, name):
        return "other"


def classify(code: str, name: str) -> str:
    """Keyword-FIRST classification.

    ingest_bom.classify_ingredient matches on code PREFIX first, which was tuned
    for the old BoM code space ('13-' => sauce). The new workbook uses DATAP's
    ITEM CODE(ACC), where 13-01-009 is ปูอัด (protein) — so prefixes lie here.
    Match the name first, fall back to the legacy classifier.
    """
    if ingest_bom is None:
        return "other"
    n = (name or "").lower()

    def has(kws):
        return any(k.lower() in n for k in kws)

    for kws, label in ((ingest_bom._PROTEIN_KW, "protein"),
                       (ingest_bom._CARB_KW, "carb"),
                       (ingest_bom._VEG_KW, "veg"),
                       (ingest_bom._OIL_KW, "oil"),
                       (ingest_bom._SAUCE_KW, "sauce"),
                       (ingest_bom._SEASON_KW, "seasoning")):
        if has(kws):
            return label
    return classify_ingredient(code, name)

# recipe-sheet column indices (0-based)
C_SEQ, C_CODE, C_PROD, C_ING, C_DETAIL = 1, 2, 3, 4, 5
C_FULLPRICE, C_FULLQTY, C_LOSS, C_NETQTY = 6, 7, 8, 9
C_PPU, C_QTY, C_UNIT, C_COST = 10, 11, 12, 13


def norm(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("​", " ")).strip().lower()


def f(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip() or 0)
    except ValueError:
        return 0.0


def s(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


# --- parsers ----------------------------------------------------------------
def parse_sum(ws) -> list:
    """SUM sheet = dish master, one row per dish carrying BOTH item codes."""
    out = []
    for r in ws.iter_rows(values_only=True):
        if r[3] and str(r[3]).startswith("OTH"):
            out.append(dict(no=r[1], en=s(r[2]), bf=s(r[3]), al=s(r[4]), th=s(r[5]),
                            main=s(r[6]), serve_bf=s(r[9]), cost_bf=f(r[10]),
                            serve_al=s(r[11]), cost_al=f(r[12]),
                            tier=s(r[13]).replace(",", "")))
    return out


def parse_recipe_sheet(ws) -> list:
    """A recipe sheet is stacked blocks: header / name row / table / total row."""
    rows = list(ws.iter_rows(values_only=True))
    blocks, i, n = [], 0, len(rows)
    while i < n:
        if rows[i][C_PROD] != "Product Item Code":
            i += 1
            continue
        head = rows[i + 1] if i + 1 < n else (None,) * 20
        code, en, th = s(head[C_PROD]), s(head[C_ING]), s(head[8])
        j = i + 2
        if j < n and rows[j][C_SEQ] == "ลำดับ":   # 'ลำดับ'
            j += 1
        ings, serve, total = [], "", None
        while j < n:
            rr = rows[j]
            if s(rr[C_SEQ]).startswith("ราคา"):        # 'ราคา (บาท)'
                serve, total = s(rr[C_DETAIL]), rr[C_COST]
                j += 1
                break
            if rr[C_PROD] == "Product Item Code":
                break
            seq = s(rr[C_SEQ]).replace(".0", "")
            if seq.isdigit():
                ings.append(dict(
                    seq=int(seq), code=s(rr[C_CODE]), prod=s(rr[C_PROD]),
                    ing=s(rr[C_ING]), detail=s(rr[C_DETAIL]),
                    full_price=f(rr[C_FULLPRICE]), full_qty=f(rr[C_FULLQTY]),
                    loss=f(rr[C_LOSS]), net_qty=f(rr[C_NETQTY]),
                    ppu=f(rr[C_PPU]), qty=f(rr[C_QTY]), unit=s(rr[C_UNIT]),
                    cost=f(rr[C_COST]), raw_cost=s(rr[C_COST])))
            j += 1
        blocks.append(dict(code=code, en=en, th=th, serve=serve,
                           total=None if total is None else f(total), ings=ings))
        i = j
    return blocks


def parse_datap(ws) -> dict:
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        code = s(r[6])
        if not code:
            continue
        out[code] = dict(code=code, category=s(r[3]), supplier=s(r[2]), name=s(r[5]),
                         unit=s(r[7]), full_price=f(r[8]), full_qty=f(r[9]),
                         loss_pct=f(r[10]), remarks=s(r[12]))
    return out


def parse_decisions(path: Path) -> dict:
    """Ops rulings on recipes that carry no Item Code, keyed (tier, normalised name).

    Hand-maintained CSV, never overwritten by this script:
        tier,name_en,name_th,decision,cost_override,note
    decision: 'exclude'      -> not on the new menu, stop reporting it as blocking
              'pending_code' -> real menu item, still waiting for its POS code
    cost_override: cost/serve to trust when the workbook's own formula returns 0.
    """
    out = {}
    if not path.exists():
        return out
    with path.open(encoding="utf-8-sig", newline="") as fh:
        for r in csv.DictReader(fh):
            tier = s(r.get("tier")).replace(",", "")
            rec = dict(decision=s(r.get("decision")).lower(),
                       cost_override=s(r.get("cost_override")), note=s(r.get("note")))
            for key in (r.get("name_en"), r.get("name_th")):
                if s(key):
                    out[(tier, norm(key))] = rec
    return out


def parse_master_item(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = {}
    for r in ws.iter_rows(values_only=True):
        code = s(r[4]) if len(r) > 4 else ""
        if not code.startswith("OTH"):
            continue
        # à-la-carte rows carry the literal text "None" in the tier column
        tier = s(r[5]).replace(",", "")
        out[code] = dict(code=code, tier="" if tier.lower() == "none" else tier,
                         name=s(r[6]), cost_per_serve=f(r[7]))
    wb.close()
    return out


# --- main -------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--bom", default=str(DL / "Final BOM Grand Menu.xlsx"))
    ap.add_argument("--master", default=str(DL / "Master Item Grand Menu.xlsx"))
    ap.add_argument("--out", default=str(ROOT / "BoM" / "GrandMenu_2026-09"))
    ap.add_argument("--decisions", default="",
                    help="ops rulings CSV (default: <out>/menu_decisions.csv)")
    a = ap.parse_args()

    out_dir = Path(a.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(a.bom, read_only=True, data_only=True)
    dishes = parse_sum(wb["SUM"])
    datap = parse_datap(wb["DATAP"])
    master = parse_master_item(Path(a.master))
    decisions = parse_decisions(Path(a.decisions) if a.decisions
                                else out_dir / "menu_decisions.csv")

    by_name = defaultdict(list)
    for d in dishes:
        by_name[(d["tier"], norm(d["en"]))].append(d)
        by_name[(d["tier"], norm(d["th"]))].append(d)

    issues = []
    lines = []
    seen_sku = {}

    for sheet in wb.sheetnames:
        m = re.match(r"^(BF|AC)_([\d,]+)NET", sheet)
        if not m:
            continue
        kind, tier = m.group(1), m.group(2).replace(",", "")
        for b in parse_recipe_sheet(wb[sheet]):
            if b["en"] == "MENU NAME":          # stray template row
                continue
            cand = by_name.get((tier, norm(b["en"]))) or by_name.get((tier, norm(b["th"])))
            if not cand:
                rule = decisions.get((tier, norm(b["en"]))) or decisions.get((tier, norm(b["th"])))
                if rule and rule["decision"] == "exclude":
                    issues.append(dict(kind="EXCLUDED_BY_OPS", sheet=sheet, sku_code="",
                                       menu=b["en"] + " / " + b["th"], tier=tier,
                                       detail=rule["note"] or "excluded in menu_decisions.csv"))
                    continue
                issues.append(dict(kind="NO_ITEM_CODE", sheet=sheet, sku_code="",
                                   menu=b["en"] + " / " + b["th"], tier=tier,
                                   detail="recipe has %d ingredient lines but no Item Code "
                                          "in SUM/Master Item%s"
                                          % (len(b["ings"]),
                                             " — " + rule["note"] if rule else "")))
                continue
            d = cand[0]
            sku = d["bf"] if kind == "BF" else d["al"]
            variant = "B" if kind == "BF" else "A"
            if not sku:
                issues.append(dict(kind="NO_ITEM_CODE", sheet=sheet, sku_code="",
                                   menu=d["en"], tier=tier,
                                   detail="%s code blank in SUM" % kind))
                continue

            stated = d["cost_bf"] if kind == "BF" else d["cost_al"]
            calc = sum(g["cost"] for g in b["ings"])
            if b["total"] is not None and abs(b["total"] - stated) > 0.01:
                issues.append(dict(kind="COST_MISMATCH_SUM", sheet=sheet, sku_code=sku,
                                   menu=d["en"], tier=tier,
                                   detail="recipe total %.4f vs SUM %.4f" % (b["total"], stated)))
            if abs(calc - (b["total"] or 0)) > 0.01:
                issues.append(dict(kind="COST_MISMATCH_LINES", sheet=sheet, sku_code=sku,
                                   menu=d["en"], tier=tier,
                                   detail="sum(lines) %.4f vs block total %s" % (calc, b["total"])))
            mrow = master.get(sku)
            if not mrow:
                issues.append(dict(kind="NOT_IN_MASTER_ITEM", sheet=sheet, sku_code=sku,
                                   menu=d["en"], tier=tier,
                                   detail="code missing from Master Item file"))
            elif abs(mrow["cost_per_serve"] - stated) > 0.01:
                issues.append(dict(kind="COST_MISMATCH_MASTER", sheet=sheet, sku_code=sku,
                                   menu=d["en"], tier=tier,
                                   detail="Master Item %.4f vs SUM %.4f"
                                          % (mrow["cost_per_serve"], stated)))

            if sku in seen_sku:
                issues.append(dict(kind="DUPLICATE_SKU", sheet=sheet, sku_code=sku,
                                   menu=d["en"], tier=tier,
                                   detail="already defined on sheet %s" % seen_sku[sku]["sheet"]))
                continue

            rule_c = decisions.get((tier, norm(d["en"]))) or decisions.get((tier, norm(d["th"])))
            th_name = mrow["name"] if mrow else "(%s) %s" % (kind, d["th"])
            desc = re.sub(r"^\((BF|AL)\)\s*", "(%s) " % variant, th_name)
            if not desc.startswith("("):
                desc = "(%s) %s" % (variant, desc)
            seen_sku[sku] = dict(
                sku_code=sku, variant=variant, tier=tier if kind == "BF" else "",
                description=desc, raw_name=th_name, name_en=d["en"], name_th=d["th"],
                serving=d["serve_bf"] if kind == "BF" else d["serve_al"],
                cost_per_serve=round(stated, 4), recipe_lines=len(b["ings"]),
                cost_override=(rule_c["cost_override"] if rule_c else ""),
                sheet=sheet, main_ingredient=d["main"])

            for g in b["ings"]:
                if g["code"] == "#N/A" or g["unit"] == "#N/A" or g["raw_cost"] == "#N/A":
                    issues.append(dict(kind="NA_CELL", sheet=sheet, sku_code=sku, menu=d["en"],
                                       tier=tier,
                                       detail='seq %d "%s" has #N/A' % (g["seq"], g["ing"])))
                elif g["code"] not in datap:
                    issues.append(dict(kind="INGREDIENT_NOT_IN_DATAP", sheet=sheet, sku_code=sku,
                                       menu=d["en"], tier=tier,
                                       detail="%s %s" % (g["code"], g["ing"])))
                if g["cost"] == 0:
                    issues.append(dict(kind="ZERO_COST_LINE", sheet=sheet, sku_code=sku,
                                       menu=d["en"], tier=tier,
                                       detail="seq %d %s %s" % (g["seq"], g["code"], g["ing"])))
                lines.append(dict(
                    sku_code=sku, variant=variant, tier=tier, seq=g["seq"],
                    ingredient_code=g["code"], ingredient_name=g["ing"] or g["prod"],
                    ingredient_detail=g["detail"], product_name=g["prod"],
                    qty_per_serving=g["qty"], unit=g["unit"],
                    price_per_unit=g["ppu"], ingredient_cost=round(g["cost"], 6),
                    full_price=g["full_price"], full_qty=g["full_qty"], loss_qty=g["loss"],
                    net_qty=g["net_qty"],
                    ingredient_type=classify(g["code"], g["ing"] or g["prod"]),
                    source_sheet=sheet))

    for d in dishes:
        for code, kind in ((d["bf"], "BF"), (d["al"], "AL")):
            if code and code not in seen_sku:
                issues.append(dict(kind="NO_RECIPE", sheet="SUM", sku_code=code, menu=d["en"],
                                   tier=d["tier"], detail="%s code has no recipe block" % kind))
    # Codes that exist in Master Item but have no recipe block. The rule is "go by the
    # items that have a POS code", so these are still menu items and must exist as skus
    # — otherwise the first September sync stubs them and their COGS reads 0. They carry
    # no BoM, so the loader seeds sku_cost from Master Item's cost/serve instead.
    for code, mrow in master.items():
        if code in seen_sku:
            continue
        variant = "B" if code[3:5] == "BF" else "A"
        rule_m = decisions.get((mrow["tier"], norm(mrow["name"])))
        if rule_m and rule_m["decision"] == "exclude":
            issues.append(dict(kind="EXCLUDED_BY_OPS", sheet="MasterItem", sku_code=code,
                               menu=mrow["name"], tier=mrow["tier"],
                               detail=rule_m["note"] or "excluded in menu_decisions.csv"))
            continue
        issues.append(dict(kind="MASTER_WITHOUT_RECIPE", sheet="MasterItem", sku_code=code,
                           menu=mrow["name"], tier=mrow["tier"],
                           detail="no recipe in the BoM workbook — loaded with cost/serve "
                                  "%.2f from Master Item, no BoM lines"
                                  % mrow["cost_per_serve"]))
        desc = re.sub(r"^\((BF|AL)\)\s*", "(%s) " % variant, mrow["name"])
        if not desc.startswith("("):
            desc = "(%s) %s" % (variant, desc)
        seen_sku[code] = dict(
            sku_code=code, variant=variant, tier=mrow["tier"] if variant == "B" else "",
            description=desc, raw_name=mrow["name"], name_en="", name_th="",
            serving="", cost_per_serve=round(mrow["cost_per_serve"], 4), recipe_lines=0,
            cost_override=(rule_m["cost_override"] if rule_m else ""),
            sheet="MasterItem", main_ingredient="")

    def dump(name, rows, cols):
        p = out_dir / name
        with p.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        print("  %-24s %5d rows" % (p.name, len(rows)))

    print("\nParsed %d dishes from SUM, %d ingredients from DATAP, %d codes from Master Item"
          % (len(dishes), len(datap), len(master)))
    print("Output -> %s" % out_dir)
    dump("menu_master.csv", sorted(seen_sku.values(), key=lambda r: r["sku_code"]),
         ["sku_code", "variant", "tier", "description", "raw_name", "name_en", "name_th",
          "serving", "cost_per_serve", "cost_override", "recipe_lines", "main_ingredient",
          "sheet"])
    dump("bom_lines.csv", lines,
         ["sku_code", "variant", "tier", "seq", "ingredient_code", "ingredient_name",
          "ingredient_detail", "product_name", "qty_per_serving", "unit", "price_per_unit",
          "ingredient_cost", "full_price", "full_qty", "loss_qty", "net_qty",
          "ingredient_type", "source_sheet"])
    dump("ingredient_master.csv", sorted(datap.values(), key=lambda r: r["code"]),
         ["code", "category", "supplier", "name", "unit", "full_price", "full_qty",
          "loss_pct", "remarks"])
    dump("issues.csv", issues, ["kind", "sheet", "sku_code", "menu", "tier", "detail"])

    print("\nIssue summary:")
    tally = defaultdict(int)
    for i in issues:
        tally[i["kind"]] += 1
    for k in sorted(tally, key=lambda k: -tally[k]):
        print("  %-26s %d" % (k, tally[k]))
    print("\nReady to import: %d menu codes, %d ingredient lines"
          % (len(seen_sku), len(lines)))
    wb.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
