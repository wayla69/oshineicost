"""The POS's own branch id on sheets that go to other teams.

Owner, 2026-09-21, on the weekly purchases template: "Please add the branch ID using for
sales team too". The app knew a branch only by its own 1-29 id and a Thai name; the sales
team and every Easy POS report know it as Branch.BranchCode ("0001" = ขอนแก่น). The live app
has no access to the POS, so the sync writes the mapping (branch_pos) and it travels with
the data.
"""
import io
import sqlite3

import openpyxl
import pytest

import main
import routers.ingredients as ing
import sync_mssql


@pytest.fixture
def pos_map(client):
    """Give the fixture DB's two branches the ids the POS would have for them."""
    con = sqlite3.connect(str(main.DB_PATH))
    ids = dict(con.execute("SELECT name, branch_id FROM branches WHERE name IN ('BranchA','BranchB')"))
    sync_mssql.write_branch_pos(con, [(ids["BranchA"], 5, "0001", "OSHINEI BranchA (1)"),
                                      (ids["BranchB"], 26, " 00021 ", "OSHINEI BranchB")])
    con.commit(); con.close()
    main.invalidate_cache()
    yield ids
    con = sqlite3.connect(str(main.DB_PATH))
    con.execute("DROP TABLE IF EXISTS branch_pos")
    try:
        con.execute("DELETE FROM ingredient_purchase_week")
    except sqlite3.OperationalError:
        pass
    con.commit(); con.close()
    main.invalidate_cache()


def _template(client, **params):
    r = client.get("/api/ingredients/purchases.template.xlsx", params={"weeks": 2, **params})
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(io.BytesIO(r.content))


def test_the_sync_writes_the_pos_ids_and_rewrites_them_whole():
    con = sqlite3.connect(":memory:")
    assert sync_mssql.write_branch_pos(con, [(3, 5, "0001", "OSHINEI ขอนแก่น (1)"), (16, 9, None, "OSHINEI มุกดาหาร1")]) == 2
    assert con.execute("SELECT * FROM branch_pos ORDER BY 1").fetchall() == [
        (3, 5, "0001", "OSHINEI ขอนแก่น (1)"), (16, 9, None, "OSHINEI มุกดาหาร1")]       # "0001" stays TEXT
    sync_mssql.write_branch_pos(con, [(3, 5, "0001", "renamed")])                          # a branch closed in the POS
    assert con.execute("SELECT branch_id, pos_branch_name FROM branch_pos").fetchall() == [(3, "renamed")]


def test_the_template_carries_the_pos_code_beside_the_name(client, pos_map):
    wb = _template(client)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[3]]
    assert head[:3] == ["สาขา", "รหัสสาขา POS", "สัปดาห์เริ่ม"] and "ยอดซื้อวัตถุดิบ (บาท)" in head
    rows = {(r[0].value, r[1].value) for r in ws.iter_rows(min_row=4)}
    assert rows == {("BranchA", "0001"), ("BranchB", "00021")}                  # text with its zeros, trimmed
    assert ws.cell(row=4, column=2).number_format == "@"                         # Excel must not turn it into 1
    assert wb.active.title == ws.title                                           # the import reads the active sheet
    ref = wb["รหัสสาขา"]
    assert [c.value for c in ref[1]] == ["สาขา (ชื่อในระบบนี้)", "รหัสสาขา POS", "POS BranchId", "ชื่อสาขาใน POS", "branch_id (ระบบนี้)"]
    got = {r[0]: r[1:] for r in ref.iter_rows(min_row=2, values_only=True)}
    assert got["BranchA"] == ("0001", 5, "OSHINEI BranchA (1)", pos_map["BranchA"])


def test_without_a_mapping_the_column_is_simply_blank(client):
    con = sqlite3.connect(str(main.DB_PATH)); con.execute("DROP TABLE IF EXISTS branch_pos"); con.commit(); con.close()
    ws = _template(client).worksheets[0]
    assert [c.value for c in ws[3]][1] == ing.POS_CODE_HEADER and ws.cell(row=4, column=2).value is None


def _sheet(rows):
    wb = openpyxl.Workbook(); ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    return {"file": ("purchases.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_the_filled_template_imports_and_a_retyped_name_is_rescued_by_the_code(client, pos_map, monkeypatch):
    monkeypatch.delenv("OSHINEI_SEED_TOKEN", raising=False)
    week = _template(client).worksheets[0].cell(row=4, column=3).value
    up = _sheet([["สาขา", "รหัสสาขา POS", "สัปดาห์เริ่ม", "ยอดซื้อวัตถุดิบ (บาท)"],
                 ["BranchA", "0001", week, 1000],
                 ["Branch-B (retyped)", 21, week, 2000],            # name wrong, and Excel ate the zeros: 00021 -> 21
                 ["Nowhere", "9999", week, 3000]])
    r = client.post("/api/ingredients/purchases/import", params={"dry_run": True}, files=up).json()
    assert r["would_save"] == 2 and [e["branch"] for e in r["errors"]] == ["Nowhere"]


def test_a_sheet_without_the_new_column_still_imports_as_before(client, pos_map, monkeypatch):
    monkeypatch.delenv("OSHINEI_SEED_TOKEN", raising=False)
    week = _template(client).worksheets[0].cell(row=4, column=3).value
    r = client.post("/api/ingredients/purchases/import", params={"dry_run": True},
                    files=_sheet([["สาขา", "สัปดาห์เริ่ม", "ยอดซื้อวัตถุดิบ (บาท)"], ["BranchA", week, 500]])).json()
    assert r["would_save"] == 1 and r["errors"] == []
    # two recognised columns were enough to find the header before; the optional one must not raise that bar
    r = client.post("/api/ingredients/purchases/import", params={"dry_run": True},
                    files=_sheet([["สาขา", "ยอดซื้อ"], ["BranchA", 500]]))
    assert r.status_code == 200 and r.json()["would_save"] == 0 and "วันที่" in r.json()["errors"][0]["why"]


def test_the_code_header_can_never_be_read_as_the_branch_column():
    assert not ing._norm_head(ing.POS_CODE_HEADER).startswith(ing._norm_head("สาขา"))
