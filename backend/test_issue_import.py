# -*- coding: utf-8 -*-
"""PR-SI — the branch forms as Excel: the issue slip, the order and the receipt can be downloaded as a template, filled in and imported.

The promises:
  * the slip template lists the items the app's sheet lists (same order), prefilled with the saved draft, with the 8-week normal and the
    system stock as grey locked columns, a reasons sheet and a drop-down of the reason labels; `days` repeats the block per date;
  * an import goes through issues.save: the same errors as the app, a dry run writes nothing, a bad date does not stop a good one,
    a confirmed date is refused (never appended), `submit` confirms with the late flag, the log carries the file name, the cache is cleared;
  * the order and the receipt imports go through orders.save_flow / orders.receive: the app's own rules (reason, cut-off, today's slip,
    acknowledgement, shortage reason), a dry run writes nothing, the branch is pinned.
"""
import io
import sqlite3
from datetime import date, datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook

import auth
from bom_center import branch_forms as BF
from bom_center import claims as CL
from bom_center import forecast as F
from bom_center import issues as IS
from bom_center import orders as OR
from bom_center import stock as ST
from bom_center import variance as VA
from bom_center.clock import BKK
from test_bom_center_claims import CODE  # noqa: F401
from test_bom_center_issues import AS_OF, CRAB, CRAB_NM, NOW, PACK, at, build_stock, used  # noqa: F401
from test_bom_center_s4b import con  # noqa: F401  (sales to 5 Oct 2026 at branch 1; crab stick 1 kg packs)


pytestmark = pytest.mark.usefixtures("recipe_mode")        # the fixtures build their sheets from recipes (mode full)
SPILL, OTHER = IS.WASTE_REASON_TH["spill"], IS.WASTE_REASON_TH["other"]


def _book(wb) -> bytes:
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _tmpl(con, day="2026-10-05", days=1, now=NOW):  # noqa: F811
    wb, name, d0, d1 = BF.slip_template(con, 1, date.fromisoformat(day), days, now=now)
    return wb, name


def _fill(wb, edits):
    """edits: {(date, code): {column name: value}} -> the same workbook with those cells typed."""
    ws = wb[BF.SLIP_SHEET]
    head = [c.value for c in ws[2]]
    for r in range(3, ws.max_row + 1):
        key = (ws.cell(r, 1).value, ws.cell(r, 2).value)
        for col, v in (edits.get(key) or {}).items():
            ws.cell(r, head.index(col) + 1).value = v
    return wb


def _parse(wb):
    return BF.read_slip_file(_book(wb))


def _imp(con, wb, **kw):  # noqa: F811
    kw.setdefault("dry_run", True)
    kw.setdefault("now", NOW)
    return BF.import_slips(con, 1, _parse(wb), "udon", "ใบตัดเบิก.xlsx", **kw)


def _state(con):  # noqa: F811
    return {t: sorted(repr(tuple(r)) for r in con.execute(f"SELECT * FROM {t}"))
            for t in ("bc_issue_slip", "bc_issue_line", "bc_issue_log", "bc_audit")}


def _row_of(wb, day, code):
    ws = wb[BF.SLIP_SHEET]
    return next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == day and ws.cell(r, 2).value == code)


# ── 1 · the template ─────────────────────────────────────────────────────────
def test_template_rows_are_the_sheet_items_prefilled_with_the_draft(con):  # noqa: F811
    build_stock(con)
    IS.save(con, 1, "2026-10-05", [{"code": CRAB, "qty_used": 3, "qty_waste": 2, "waste_reason": "spill", "note": "หก", "unit": PACK}],
            "udon", now=NOW)
    con.commit()
    wb, name = _tmpl(con)
    ws = wb[BF.SLIP_SHEET]
    assert name and "รหัสสาขา 1" in ws["A1"].value and "2026-10-05" in ws["A1"].value
    assert "กรอกเฉพาะช่องเบิกใช้ ของเสีย และสาเหตุ · ห้ามแก้รหัส" in ws["A1"].value
    assert [c.value for c in ws[2]] == ["วันที่", "รหัส", "วัตถุดิบ", "หน่วย", "ปกติ (8 สัปดาห์)", "ของในระบบ", "เบิกใช้", "ของเสีย",
                                         "สาเหตุของเสีย", "หมายเหตุ"]
    sh = IS.sheet(con, 1, "2026-10-05", now=NOW)
    want = [it["code"] for z in sh["zones"] for it in z["items"]]
    got = [ws.cell(r, 2).value for r in range(3, ws.max_row + 1)]
    assert got == want and len(got) >= 1                                          # same items, same order
    r = _row_of(wb, "2026-10-05", CRAB)
    row = {ws.cell(2, c).value: ws.cell(r, c).value for c in range(1, 11)}
    assert row["หน่วย"] == PACK and row["เบิกใช้"] == 3 and row["ของเสีย"] == 2 and row["สาเหตุของเสีย"] == "ของเสียหายระหว่างเตรียม" or \
        row["สาเหตุของเสีย"] in dict(IS.WASTE_REASONS).values()
    assert row["สาเหตุของเสีย"] == IS.WASTE_REASON_TH["spill"] and row["หมายเหตุ"] == "หก"
    it = next(i for z in sh["zones"] for i in z["items"] if i["code"] == CRAB)
    assert row["ของในระบบ"] == BF.sys_text(it["system_qty"], PACK)


def test_template_has_locked_grey_columns_a_reasons_sheet_and_a_dropdown(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con)
    ws = wb[BF.SLIP_SHEET]
    assert ws.protection.sheet is True and not ws.protection.password
    r = _row_of(wb, "2026-10-05", CRAB)
    for c in range(1, 7):                                                             # วันที่ รหัส วัตถุดิบ หน่วย ปกติ คงเหลือ: locked + grey
        assert ws.cell(r, c).protection.locked is True and ws.cell(r, c).fill.fgColor.rgb.endswith("E5E7EB")
    for c in range(7, 11):
        assert ws.cell(r, c).protection.locked is False
    rs = wb[BF.SLIP_REASONS_SHEET]
    assert [(x[0], x[1]) for x in rs.iter_rows(min_row=2, values_only=True)] == [(t, c) for c, t in IS.WASTE_REASONS]
    dv = ws.data_validations.dataValidation
    assert len(dv) == 1 and dv[0].type == "list" and "สาเหตุของเสีย" in dv[0].formula1 and f"I3:I{ws.max_row}" in str(dv[0].sqref)


def test_the_normal_column_is_the_8_week_mean_and_blank_without_history(con):  # noqa: F811
    build_stock(con)
    for k in (1, 2):                                                                  # two confirmed Mondays: 1000 and 3000 g -> mean 2 packs
        d = date(2026, 10, 5) - timedelta(days=7 * k)
        IS.save(con, 1, d.isoformat(), [used(qty=1000 * (2 * k - 1))], "udon", submit=True, now=at(d.day, 21, month=d.month))
    con.commit()
    wb, _ = _tmpl(con)
    ws = wb[BF.SLIP_SHEET]
    assert ws.cell(_row_of(wb, "2026-10-05", CRAB), 5).value == pytest.approx(2.0)
    other = [r for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value != CRAB]
    assert all(ws.cell(r, 5).value is None for r in other)


def test_days_repeats_the_block_for_consecutive_dates_and_is_limited(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con, "2026-10-03", days=3)
    ws = wb[BF.SLIP_SHEET]
    days = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
    assert sorted(set(days)) == ["2026-10-03", "2026-10-04", "2026-10-05"]
    n = days.count("2026-10-03")
    assert days.count("2026-10-04") == n == days.count("2026-10-05") and "ถึง 2026-10-05" in ws["A1"].value
    for bad in (0, 8):
        with pytest.raises(ValueError, match="จำนวนวันต้องอยู่ระหว่าง 1 ถึง 7"):
            BF.slip_template(con, 1, date(2026, 10, 3), bad, now=NOW)
    with pytest.raises(ValueError, match="ใบตัดเบิกต้องไม่เป็นวันล่วงหน้า"):
        BF.slip_template(con, 1, date(2026, 10, 5), 2, now=NOW)


# ── 2 · the dry run ──────────────────────────────────────────────────────────
def test_a_dry_run_reports_per_date_and_touches_nothing(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con, "2026-10-04", days=2)
    _fill(wb, {("2026-10-04", CRAB): {"เบิกใช้": 2}, ("2026-10-05", CRAB): {"เบิกใช้": 1, "ของเสีย": 1, "สาเหตุของเสีย": SPILL}})
    con.commit()
    before = _state(con)
    out = _imp(con, wb)
    assert out["dry_run"] is True and out["ok"] is True and out["ignored"] > 0
    d4, d5 = out["dates"]
    assert (d4["date"], d4["state"], d4["lines"], d4["used_lines"], d4["waste_lines"], d4["errors"]) == ("2026-10-04", "draft", 1, 1, 0, [])
    assert (d5["date"], d5["state"], d5["used_lines"], d5["waste_lines"]) == ("2026-10-05", "new", 1, 1)
    assert _state(con) == before
    con.rollback()
    assert _state(con) == before


@pytest.mark.parametrize("edit, msg", [
    ({"เบิกใช้": "abc"}, f"{CRAB_NM}: จำนวนต้องเป็นตัวเลข"),
    ({"เบิกใช้": -1}, f"{CRAB_NM}: จำนวนต้องไม่ติดลบ"),
    ({"เบิกใช้": 1, "ของเสีย": 1}, f"{CRAB_NM}: เลือกสาเหตุของของเสีย"),
    ({"ของเสีย": 1, "สาเหตุของเสีย": OTHER}, f"{CRAB_NM}: สาเหตุ “อื่น ๆ” ต้องระบุคำอธิบาย"),
])
def test_the_errors_are_the_apps_errors_with_the_row(con, edit, msg):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con)
    _fill(wb, {("2026-10-05", CRAB): edit})
    out = _imp(con, wb)
    d = out["dates"][0]
    assert out["ok"] is False and d["result"] == "refused"
    assert [(e["row"], e["message"]) for e in d["errors"]] == [(_row_of(wb, "2026-10-05", CRAB), msg)]
    with pytest.raises(ValueError, match=msg):                                        # the very message of the app's save
        IS.save(con, 1, "2026-10-05", [{"code": CRAB, "qty_used": edit.get("เบิกใช้"), "qty_waste": edit.get("ของเสีย"),
                                        "waste_reason": "other" if edit.get("สาเหตุของเสีย") else None, "unit": PACK}], "udon", now=NOW)


def _extra(row, day, code=CRAB, **kw):
    """A row appended to a parsed file by hand (a date or a code the template would not offer)."""
    return (row, day, {"code": code, "qty_used": 1, "qty_waste": None, "waste_reason": None, "note": None, **kw})


def test_an_unconvertible_unit_an_unknown_code_and_a_repeated_code_are_errors(con, monkeypatch):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con)
    _fill(wb, {("2026-10-05", CRAB): {"เบิกใช้": 1}})
    parsed = _parse(wb)
    parsed["rows"] += [_extra(98, "2026-10-05", "NO-SUCH"), _extra(99, "2026-10-05")]
    out = BF.import_slips(con, 1, parsed, "udon", "f.xlsx", dry_run=True, now=NOW)
    errs = {(e["row"], e["code"]) for e in out["dates"][0]["errors"]}
    assert (98, "code") in errs and (99, "dup") in errs and any("NO-SUCH" in e["message"] for e in out["dates"][0]["errors"])
    parsed["rows"] = [r for r in parsed["rows"] if r[0] < 98]
    monkeypatch.setattr(IS, "Converter", type("NoUnits", (IS.Converter,), {"to_base": lambda self, code, q, unit: (None, None)}))
    out = BF.import_slips(con, 1, parsed, "udon", "f.xlsx", dry_run=True, now=NOW)
    e = out["dates"][0]["errors"]
    assert [(x["row"], x["code"]) for x in e] == [(_row_of(wb, "2026-10-05", CRAB), "unit")] and "ยังไม่ได้ตั้งหน่วย" in e[0]["message"]


def test_future_old_and_confirmed_dates_are_refused_with_their_state(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con, "2026-10-04", days=2)
    _fill(wb, {("2026-10-04", CRAB): {"เบิกใช้": 1}, ("2026-10-05", CRAB): {"เบิกใช้": 1}})
    parsed = _parse(wb)
    parsed["rows"] += [_extra(90, "2026-10-06"), _extra(91, "2026-07-01"), _extra(92, "2026-10-03")]       # 3 Oct was confirmed by build_stock
    out = BF.import_slips(con, 1, parsed, "udon", "f.xlsx", dry_run=True, now=NOW)
    by = {d["date"]: d for d in out["dates"]}
    assert by["2026-10-06"]["state"] == "future" and by["2026-10-06"]["errors"][0]["message"] == "ใบตัดเบิกต้องไม่เป็นวันล่วงหน้า"
    assert by["2026-07-01"]["state"] == "too_old" and "ย้อนหลังได้ไม่เกิน 60 วัน" in by["2026-07-01"]["errors"][0]["message"]
    assert by["2026-10-03"]["state"] == "confirmed" and by["2026-10-03"]["result"] == "refused"
    assert by["2026-10-03"]["errors"][0]["message"] == "ใบตัดเบิกวันที่ 3/10/2026 ส่งแล้ว หากต้องแก้ไข แจ้งประสานงานขาย"
    assert by["2026-10-04"]["state"] == "draft" and by["2026-10-05"]["state"] == "new"
    assert out["ok"] is False and [d["date"] for d in out["dates"]] == sorted(by)


def test_a_period_accounting_has_posted_refuses_the_date(con):  # noqa: F811
    ST.save(con, 1, [{"code": CRAB, "qty": 10, "unit": PACK}], "udon", submit=True, count_date="2026-09-01", require_all=False)
    end = ST.save(con, 1, [{"code": CRAB, "qty": 8, "unit": PACK}], "udon", submit=True, count_date="2026-10-02", require_all=False)["count_id"]
    con.commit()
    VA.confirm_receipts(con, 1, end, "udon")
    VA.set_posted(con, end, True, "accounting1")
    con.commit()
    wb, _ = _tmpl(con, "2026-09-30")
    _fill(wb, {("2026-09-30", CRAB): {"เบิกใช้": 1}})
    out = _imp(con, wb)
    d = out["dates"][0]
    assert d["state"] == "posted" and d["result"] == "refused" and "อยู่ในงวดที่บัญชีบันทึกใบสรุปตัดเบิกแล้ว" in d["errors"][0]["message"]
    before = _state(con)
    assert _imp(con, wb, dry_run=False)["refused"] == 1 and _state(con) == before


# ── 3 · writing ──────────────────────────────────────────────────────────────
def test_a_multi_day_file_saves_each_good_day_and_skips_a_bad_one(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con, "2026-10-04", days=2)
    _fill(wb, {("2026-10-04", CRAB): {"เบิกใช้": "x"}, ("2026-10-05", CRAB): {"เบิกใช้": 1}})
    parsed = _parse(wb)
    parsed["rows"].append(_extra(97, "2026-10-02"))                                    # a confirmed day in the same file
    con.commit()
    out = BF.import_slips(con, 1, parsed, "udon", "f.xlsx", dry_run=False, now=NOW)
    con.commit()
    assert {d["date"]: d["result"] for d in out["dates"]} == {"2026-10-02": "refused", "2026-10-04": "refused", "2026-10-05": "saved"}
    assert (out["saved"], out["refused"]) == (1, 2)
    # 4 Oct keeps build_stock's draft (5 packs), 5 Oct is a new draft with the file's line
    assert con.execute("SELECT status, (SELECT qty_used FROM bc_issue_line l WHERE l.slip_id=s.slip_id) FROM bc_issue_slip s"
                       " WHERE slip_date='2026-10-04'").fetchone()[:] == ("draft", 5)
    assert con.execute("SELECT status, (SELECT qty_used FROM bc_issue_line l WHERE l.slip_id=s.slip_id) FROM bc_issue_slip s"
                       " WHERE slip_date='2026-10-05'").fetchone()[:] == ("draft", 1)


def test_submit_confirms_with_the_late_flag_and_the_log_carries_the_file_name(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con, "2026-10-04", days=2)
    _fill(wb, {("2026-10-04", CRAB): {"เบิกใช้": 2}, ("2026-10-05", CRAB): {"เบิกใช้": 1, "ของเสีย": 1, "สาเหตุของเสีย": SPILL}})
    con.commit()
    out = _imp(con, wb, dry_run=False, submit=True)
    con.commit()
    assert [(d["date"], d["result"], d["late"]) for d in out["dates"]] == [("2026-10-04", "confirmed", 1), ("2026-10-05", "confirmed", 0)]
    assert out["confirmed"] == 2
    rows = con.execute("SELECT s.slip_date, g.action, g.note, g.who FROM bc_issue_log g JOIN bc_issue_slip s USING(slip_id)"
                       " WHERE s.slip_date IN ('2026-10-04','2026-10-05') ORDER BY g.log_id").fetchall()
    assert {r[0] for r in rows if r[1] == "confirm" and r[2] == "นำเข้าจากไฟล์ ใบตัดเบิก.xlsx" and r[3] == "udon"} == {"2026-10-04", "2026-10-05"}
    acts = [r[0] for r in con.execute("SELECT action FROM bc_audit WHERE action LIKE 'issue_%' ORDER BY id")]
    assert acts.count("issue_confirm") >= 2 and acts.count("issue_import") == 2
    # a confirmed date is refused afterwards - never appended to
    again = _imp(con, wb, dry_run=False)
    assert [d["result"] for d in again["dates"]] == ["refused", "refused"] and again["saved"] == 0
    assert con.execute("SELECT qty_used FROM bc_issue_line l JOIN bc_issue_slip s USING(slip_id) WHERE s.slip_date='2026-10-05'"
                       ).fetchone()[0] == 1


def test_a_draft_import_logs_the_file_when_it_creates_the_slip(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con)
    _fill(wb, {("2026-10-05", CRAB): {"เบิกใช้": 1}})
    _imp(con, wb, dry_run=False)
    con.commit()
    assert con.execute("SELECT g.action, g.note FROM bc_issue_log g JOIN bc_issue_slip s USING(slip_id) WHERE s.slip_date='2026-10-05'"
                       ).fetchall()[0][:] == ("create", "นำเข้าจากไฟล์ ใบตัดเบิก.xlsx")


def test_a_submit_with_a_bad_line_confirms_nothing_of_that_date(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con)
    _fill(wb, {("2026-10-05", CRAB): {"เบิกใช้": 1, "ของเสีย": 1}})
    before = _state(con)
    out = _imp(con, wb, dry_run=False, submit=True)
    assert out["refused"] == 1 and out["confirmed"] == 0 and _state(con) == before




def test_a_file_of_another_branch_or_another_kind_is_refused(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con)
    _fill(wb, {("2026-10-05", CRAB): {"เบิกใช้": 1}})
    parsed = _parse(wb)
    with pytest.raises(ValueError, match="ไฟล์นี้เป็นของสาขาอื่น"):
        BF.import_slips(con, 2, parsed, "udon", "f.xlsx", dry_run=True, now=NOW)
    ws = wb[BF.SLIP_SHEET]
    ws["A1"].value = "ใบสั่งซื้อ สาขาอื่น"
    with pytest.raises(ValueError, match="ไม่ใช่แบบฟอร์มใบตัดเบิก"):
        _parse(wb)
    junk = Workbook()
    junk.active.append(["x"])
    with pytest.raises(ValueError):
        BF.read_slip_file(_book(junk))
    with pytest.raises(ValueError, match="อ่านไฟล์ไม่ได้"):
        BF.read_slip_file(b"not a workbook")


def test_the_reason_may_be_the_code_and_blank_means_zero(con):  # noqa: F811
    build_stock(con)
    wb, _ = _tmpl(con)
    _fill(wb, {("2026-10-05", CRAB): {"ของเสีย": 1, "สาเหตุของเสีย": "spill"}})
    out = _imp(con, wb)
    assert out["ok"] is True and out["dates"][0]["used_lines"] == 0 and out["dates"][0]["waste_lines"] == 1
    out = _imp(con, wb, dry_run=False)
    con.commit()
    assert con.execute("SELECT qty_used, qty_waste, waste_reason FROM bc_issue_line l JOIN bc_issue_slip s USING(slip_id)"
                       " WHERE s.slip_date='2026-10-05'").fetchone()[:] == (0, 1, "spill")


# ── 4 · the order form ───────────────────────────────────────────────────────
WED = date(2026, 10, 7)
LATE = at(5, 15)
SLIP_LINE = [{"code": "V-007", "qty_used": 10, "unit": "กรัม"}]


def _plan(c, now=NOW):  # noqa: F811
    return OR.plan(c, 1, "1", WED, as_of=AS_OF, now=now)


def _ofile(c, fill=None, plan=None):  # noqa: F811
    """The order template of branch 1 / round 1 / 7 Oct with `fill` {code: {column: value}} typed into it -> (workbook, plan)."""
    p = plan or _plan(c)
    wb, _name = BF.order_template(c, p, 1, "1", WED)
    ws = wb[BF.ORDER_SHEET]
    head = [x.value for x in ws[2]]
    for r in range(3, ws.max_row + 1):
        for col, v in ((fill or {}).get(ws.cell(r, 1).value) or {}).items():
            ws.cell(r, head.index(col) + 1).value = v
    return wb, p


def _oimp(c, wb, p, **kw):  # noqa: F811
    kw.setdefault("dry_run", True)
    return BF.import_order(c, p, 1, "1", WED, BF.read_order_file(_book(wb)), "udon", "สั่ง.xlsx", now=kw.pop("now", NOW), as_of=AS_OF, **kw)


def _ostate(c):  # noqa: F811
    return {t: sorted(repr(tuple(r)) for r in c.execute(f"SELECT * FROM {t}")) for t in ("bc_order", "bc_order_line", "bc_audit")}


def test_the_order_template_carries_the_plans_numbers_read_only_and_the_reasons(con):  # noqa: F811
    p = _plan(con)
    wb, name = BF.order_template(con, p, 1, "1", WED)
    ws = wb[BF.ORDER_SHEET]
    assert name and "รหัสสาขา 1" in ws["A1"].value and "รอบ 1" in ws["A1"].value and "วันส่ง 2026-10-07" in ws["A1"].value
    assert [c.value for c in ws[2]] == ["รหัส", "วัตถุดิบ", "หน่วย", "ของในระบบ", "จำนวนที่ควรสั่ง", "สั่งครั้งก่อน", "สั่ง", "เหตุผล", "หมายเหตุ"]
    assert [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)] == [r["code"] for r in p["rows"]]
    r = next(x for x in range(3, ws.max_row + 1) if ws.cell(x, 1).value == CRAB)
    crab = next(x for x in p["rows"] if x["code"] == CRAB)
    assert ws.cell(r, 5).value == crab["order_units"] and ws.cell(r, 7).value is None and ws.cell(r, 9).value is None
    for c in (1, 2, 3, 4, 5, 6):
        assert ws.cell(r, c).protection.locked is True
    for c in (7, 8, 9):
        assert ws.cell(r, c).protection.locked is False
    assert [x[0] for x in wb[BF.ORDER_REASONS_SHEET].iter_rows(min_row=2, values_only=True)] == [t for _c, t in OR.REASONS]
    dv = ws.data_validations.dataValidation
    assert len(dv) == 1 and str(dv[0].sqref) in (f"H3:H{ws.max_row}", "H3")


def test_without_a_suggestion_the_template_has_no_suggested_column_and_says_why(con):  # noqa: F811
    import copy
    p = copy.deepcopy(_plan(con))
    p["fallback"] = {"mode": "none", "message": "ใบตัดเบิกยังไม่พอ (3/14 วัน) · กรอกยอดสั่งเอง"}
    for r in p["rows"]:
        r["suggested"], r["last_order_qty"] = None, 4000
    wb, _ = BF.order_template(con, p, 1, "1", WED)
    ws = wb[BF.ORDER_SHEET]
    assert "ใบตัดเบิกยังไม่พอ (3/14 วัน) · กรอกยอดสั่งเอง" in ws["A1"].value
    assert all(ws.cell(r, 5).value is None for r in range(3, ws.max_row + 1))
    assert all(ws.cell(r, 6).value is not None for r in range(3, ws.max_row + 1))


def test_an_order_dry_run_writes_nothing_and_a_real_import_is_the_apps_draft(con):  # noqa: F811
    p = _plan(con)
    sug = next(x for x in p["rows"] if x["code"] == CRAB)["order_units"]
    wb, p = _ofile(con, {CRAB: {"สั่ง": sug, "เหตุผล": OR.REASON_TH["event"], "หมายเหตุ": "งานเลี้ยง"}}, p)
    con.commit()
    before = _ostate(con)
    out = _oimp(con, wb, p)
    assert out["ok"] is True and out["errors"] == [] and out["lines"] == 1 and out["status"] == "suggested" and out["order_id"] is None
    con.rollback()
    assert _ostate(con) == before
    out = _oimp(con, wb, p, dry_run=False)
    con.commit()
    assert out["status"] == "suggested" and out["order_id"]
    d = OR.detail(con, out["order_id"])
    assert {ln["code"]: (ln["confirmed_units"], ln["reason_code"], ln["note"]) for ln in d["lines"] if ln["code"] == CRAB} == {
        CRAB: (sug, "event", "งานเลี้ยง")}
    a = con.execute("SELECT new_value FROM bc_audit WHERE action='order_import'").fetchone()[0]
    assert "สั่ง.xlsx" in a and "'submit': False" in a


def test_the_order_errors_are_the_apps_errors(con):  # noqa: F811
    wb, p = _ofile(con, {CRAB: {"สั่ง": "abc"}})
    out = _oimp(con, wb, p)
    assert out["ok"] is False and [(e["row"], e["message"]) for e in out["errors"]] == [(3, f"{CRAB_NM}: จำนวนต้องเป็นตัวเลข")]
    with pytest.raises(ValueError, match="ไฟล์มี 1 แถวที่ต้องแก้ไข"):
        _oimp(con, wb, p, dry_run=False)
    wb, p = _ofile(con, {CRAB: {"สั่ง": -1}})
    assert _oimp(con, wb, p)["errors"][0]["message"] == f"{CRAB_NM}: จำนวนต้องไม่ติดลบ"
    wb, p = _ofile(con, {CRAB: {"สั่ง": 1, "เหตุผล": "ไม่มีเหตุผลนี้"}})
    assert "ไม่อยู่ในรายการ" in _oimp(con, wb, p)["errors"][0]["message"]
    wb, p = _ofile(con, {})
    assert _oimp(con, wb, p)["errors"][0]["message"] == "ไม่พบจำนวนที่สั่งในไฟล์"
    parsed = BF.read_order_file(_book(wb))
    parsed["rows"].append((50, {"code": "NO-SUCH", "confirmed_units": 1, "reason_code": None, "note": None}))
    out = BF.import_order(con, p, 1, "1", WED, parsed, "udon", "f.xlsx", now=NOW, as_of=AS_OF)
    assert out["errors"][0]["row"] == 50 and "NO-SUCH" in out["errors"][0]["message"]


def test_the_reason_rule_the_slip_rule_and_the_cutoff_are_the_apps(con):  # noqa: F811
    p = _plan(con)
    sug = next(x for x in p["rows"] if x["code"] == CRAB)["order_units"]
    wb, p = _ofile(con, {CRAB: {"สั่ง": sug * 3 + 5}}, p)
    out = _oimp(con, wb, p)                                                          # a draft only lists what needs a reason
    assert out["ok"] is True and out["needs_reason_names"] == [CRAB_NM]
    assert "ต้องระบุเหตุผล" in _oimp(con, wb, p, submit=True)["errors"][0]["message"]          # sending does not
    wb, p = _ofile(con, {CRAB: {"สั่ง": sug * 3 + 5, "เหตุผล": OR.REASON_TH["other"]}}, p)
    assert "ต้องระบุเหตุผล" in _oimp(con, wb, p, submit=True)["errors"][0]["message"]         # "other" needs a note
    wb, p = _ofile(con, {CRAB: {"สั่ง": sug * 3 + 5, "เหตุผล": OR.REASON_TH["event"]}}, p)
    assert "ต้องส่งใบตัดเบิกของวันนี้ก่อนส่งใบสั่งซื้อ" in _oimp(con, wb, p, submit=True)["errors"][0]["message"]
    IS.save(con, 1, "2026-10-05", SLIP_LINE, "udon", submit=True, now=NOW)
    con.commit()
    ok = _oimp(con, wb, p, submit=True)
    assert ok["ok"] is True and ok["status"] == "submitted" and ok["late"] == 0 and ok["order_id"] is None
    done = _oimp(con, wb, p, submit=True, dry_run=False)
    con.commit()
    assert done["status"] == "submitted" and OR.detail(con, done["order_id"])["status"] == "submitted"
    # after the cut-off: a dry run reports the late fee, a real import needs the branch's own acknowledgement
    con.execute("DELETE FROM bc_order_line")
    con.execute("DELETE FROM bc_order")
    con.commit()
    assert _oimp(con, wb, p, submit=True, now=LATE)["late"] == 1
    with pytest.raises(OR.LateSubmit):
        _oimp(con, wb, p, submit=True, dry_run=False, now=LATE)
    con.rollback()
    late = _oimp(con, wb, p, submit=True, dry_run=False, ack_late=True, now=LATE)
    assert late["late"] == 1 and late["status"] == "submitted"


def test_an_order_file_of_another_round_branch_or_kind_is_refused(con):  # noqa: F811
    wb, p = _ofile(con, {CRAB: {"สั่ง": 1}})
    parsed = BF.read_order_file(_book(wb))
    with pytest.raises(ValueError, match="ไฟล์นี้เป็นของสาขาอื่น"):
        BF.import_order(con, p, 2, "1", WED, parsed, "udon", "f.xlsx", now=NOW, as_of=AS_OF)
    with pytest.raises(ValueError, match="ไฟล์นี้เป็นของรอบหรือวันส่งอื่น"):
        BF.import_order(con, p, 1, "2", WED, parsed, "udon", "f.xlsx", now=NOW, as_of=AS_OF)
    with pytest.raises(ValueError, match="ไฟล์นี้เป็นของรอบหรือวันส่งอื่น"):
        BF.import_order(con, p, 1, "1", WED + timedelta(days=7), parsed, "udon", "f.xlsx", now=NOW, as_of=AS_OF)
    slip_wb, _ = _tmpl(con)
    with pytest.raises(ValueError, match="ไม่ใช่แบบฟอร์มใบสั่งซื้อ"):
        BF.read_order_file(_book(slip_wb))


# ── 5 · the receipt form ─────────────────────────────────────────────────────
def _sent(c):  # noqa: F811
    from test_bom_center_receive_claims import _sent as sent
    return sent(c)


def _rfile(c, oid, fill=None):  # noqa: F811
    wb, _ = BF.receive_template(c, 1, oid)
    ws = wb[BF.RECEIVE_SHEET]
    head = [x.value for x in ws[2]]
    for r in range(3, ws.max_row + 1):
        for col, v in ((fill or {}).get(ws.cell(r, 1).value) or {}).items():
            ws.cell(r, head.index(col) + 1).value = v
    return wb


def _rimp(c, wb, oid, **kw):  # noqa: F811
    kw.setdefault("dry_run", True)
    return BF.import_receive(c, 1, oid, BF.read_receive_file(_book(wb)), "udon", "รับ.xlsx", receipt_date="2026-10-07", **kw)


def _rstate(c):  # noqa: F811
    return {t: sorted(repr(tuple(r)) for r in c.execute(f"SELECT * FROM {t}"))
            for t in ("bc_order", "bc_order_line", "bc_receipt", "bc_receipt_line", "bc_claim", "bc_audit")}


def test_the_receive_template_prefills_what_was_sent_and_locks_it(con):  # noqa: F811
    oid, q = _sent(con)
    wb = _rfile(con, oid)
    ws = wb[BF.RECEIVE_SHEET]
    assert f"ใบสั่งซื้อเลขที่ {oid}" in ws["A1"].value and "รหัสสาขา 1" in ws["A1"].value
    assert [c.value for c in ws[2]] == ["รหัส", "วัตถุดิบ", "หน่วย", "ส่งจริง", "รับจริง", "สาเหตุที่ขาด", "หมายเหตุ"]
    r = next(x for x in range(3, ws.max_row + 1) if ws.cell(x, 1).value == CODE)
    assert ws.cell(r, 4).value == q and ws.cell(r, 5).value == q and ws.cell(r, 6).value is None
    assert [ws.cell(r, c).protection.locked for c in range(1, 8)] == [True, True, True, True, False, False, False]
    assert [x[0] for x in wb[BF.RECEIVE_REASONS_SHEET].iter_rows(min_row=2, values_only=True)] == [t for _c, t in CL.REASONS]


def test_a_receive_dry_run_shows_the_claim_and_writes_nothing(con):  # noqa: F811
    oid, q = _sent(con)
    before = _rstate(con)
    wb = _rfile(con, oid, {CODE: {"รับจริง": q - 2}})
    bad = _rimp(con, wb, oid)                                                        # short, no reason: the app's message
    assert bad["ok"] is False and f"{CRAB_NM}: รับจริงน้อยกว่าที่คลังส่ง ({q:g}) · เลือกสาเหตุที่ขาด" == bad["errors"][0]["message"]
    wb = _rfile(con, oid, {CODE: {"รับจริง": q - 2, "สาเหตุที่ขาด": CL.REASON_TH["short"]}})
    out = _rimp(con, wb, oid)
    assert out["ok"] is True and out["claim_lines"] == 1 and out["received"] is False and out["claim_id"] is None
    con.rollback()
    assert _rstate(con) == before and OR.detail(con, oid)["status"] == "processed"
    over = _rimp(con, _rfile(con, oid, {CODE: {"รับจริง": q + 1}}), oid)             # more than sent is only reported
    assert over["ok"] is True and over["over"] and over["claim_lines"] == 0
    con.rollback()
    assert _rstate(con) == before


def test_a_real_receive_import_needs_the_acknowledgement_and_makes_the_claim(con):  # noqa: F811
    oid, q = _sent(con)
    wb = _rfile(con, oid, {CODE: {"รับจริง": q - 2, "สาเหตุที่ขาด": "short"}})
    before = _rstate(con)
    with pytest.raises(ValueError, match="ติ๊ก “ตรวจสภาพและจำนวนครบแล้ว” ก่อนบันทึก"):
        _rimp(con, wb, oid, dry_run=False)
    con.rollback()
    assert _rstate(con) == before
    out = _rimp(con, wb, oid, dry_run=False, ack=True, claim_note="ของไม่ครบ")
    con.commit()
    assert out["received"] is True and out["claim_id"] and out["claim_lines"] == 1
    assert OR.detail(con, oid)["status"] == "received" and OR._order_row(con, oid)["ack_by"] == "udon"
    assert con.execute("SELECT COUNT(*) FROM bc_audit WHERE action='receive_import'").fetchone()[0] == 1
    with pytest.raises(ValueError, match="รับของได้เมื่อส่งใบสั่งซื้อแล้วเท่านั้น"):
        BF.receive_template(con, 1, oid)


def test_a_receive_file_of_another_order_branch_or_kind_is_refused(con):  # noqa: F811
    oid, q = _sent(con)
    wb = _rfile(con, oid)
    parsed = BF.read_receive_file(_book(wb))
    with pytest.raises(ValueError, match="ไฟล์นี้เป็นของใบสั่งซื้อเลขที่"):
        BF.import_receive(con, 1, oid + 1, parsed, "udon", "f.xlsx")
    with pytest.raises(ValueError, match="ไฟล์นี้เป็นของสาขาอื่น"):
        BF.import_receive(con, 2, oid, parsed, "udon", "f.xlsx")
    with pytest.raises(ValueError, match="ไม่พบใบสั่งซื้อของสาขานี้"):
        BF.receive_template(con, 2, oid)
    slip_wb, _ = _tmpl(con)
    with pytest.raises(ValueError, match="ไม่ใช่แบบฟอร์มรับของ"):
        BF.read_receive_file(_book(slip_wb))
    parsed["rows"].append((60, {"code": "NO-SUCH", "received_units": 1, "reason": "", "note": None}))
    out = BF.import_receive(con, 1, oid, parsed, "udon", "f.xlsx")
    assert out["errors"][0]["row"] == 60 and "NO-SUCH" in out["errors"][0]["message"]


# ── 6 · over HTTP: gates, pinning, the cache, the routes ─────────────────────
def test_http_templates_dry_run_pinning_token_log_and_cache(client, app_db, monkeypatch):
    import main
    from conftest import dl_name
    from routers import bom_center as RB
    from test_bom_center_claims import _branch_user, _login
    from test_bom_center_export import _restore, _rows
    monkeypatch.setenv("OSHINEI_AUTH", "1")
    monkeypatch.setenv("OSHINEI_AUTH_SECRET", "test-secret-stable")
    monkeypatch.setenv("OSHINEI_SEED_TOKEN", "owner-token-xyz")
    monkeypatch.setattr(IS, "now_bkk", lambda: at(5, 20))
    monkeypatch.setattr(F, "now_bkk", lambda: at(5, 8))
    monkeypatch.setattr(RB, "bkk_today", lambda: date(2026, 10, 5))
    cleared = []
    real_clear = F.clear_slip_cache
    monkeypatch.setattr(F, "clear_slip_cache", lambda: (cleared.append(1), real_clear())[1])
    client.get("/api/bom-center/hub/queue")                                  # creates the bc_* tables
    snap = _rows(main.DB_PATH)
    c = sqlite3.connect(str(main.DB_PATH))
    c.execute("DELETE FROM bc_order_calendar")
    c.execute("INSERT OR REPLACE INTO bc_ingredient(code, name_th, kind, base_unit, recv_unit, qty_per_recv, count_unit, supplier, active, monitored)"
              " VALUES('X-SI', 'ของ ทดสอบ', 'raw', 'กรัม', 'แพ็ค', 1000, 'แพ็ค', 'ผู้จำหน่าย ก', 1, 1)")
    ids = [r[0] for r in c.execute("SELECT branch_id FROM branches WHERE name IN ('BranchA', 'BranchB') ORDER BY name")]
    c.commit()
    c.close()
    a_id, b_id = ids
    try:
        mg = _login("management1")
        br_a = _branch_user("br_si_a", "BranchA")
        br_b = _branch_user("br_si_b", "BranchB")
        # the template: a branch login is pinned, the centre names the branch, the days are limited
        r = br_a.get("/api/bom-center/issues/template.xlsx?date=2026-10-05")
        assert r.status_code == 200 and dl_name(r).startswith("แบบฟอร์ม_ใบตัดเบิก_BranchA_2026-10-05"), r.text[:200]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == [BF.SLIP_SHEET, BF.SLIP_REASONS_SHEET] and f"รหัสสาขา {a_id}" in wb[BF.SLIP_SHEET]["A1"].value
        ws = wb[BF.SLIP_SHEET]
        row = next(i for i in range(3, ws.max_row + 1) if ws.cell(i, 2).value == "X-SI")
        assert br_a.get(f"/api/bom-center/issues/template.xlsx?date=2026-10-05&branch_id={b_id}").status_code == 403
        assert mg.get("/api/bom-center/issues/template.xlsx?date=2026-10-05").status_code == 400            # the centre must name a branch
        assert br_a.get("/api/bom-center/issues/template.xlsx?date=2026-10-05&days=8").status_code == 400
        assert br_a.get("/api/bom-center/issues/template.xlsx?date=2026-10-06").status_code == 400
        ws.cell(row, 7).value = 2
        body = _book(wb)
        files = {"file": ("ใบตัดเบิก.xlsx", body)}
        # a dry run needs no owner token and writes nothing
        before = _rows(main.DB_PATH)
        r = br_a.post("/api/bom-center/issues/import", files=files)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["dry_run"] is True and d["ok"] is True and d["dates"][0]["state"] == "new" and d["dates"][0]["used_lines"] == 1
        assert _rows(main.DB_PATH) == before and cleared == []
        # a branch login cannot import for, or with the file of, another branch
        assert br_b.post(f"/api/bom-center/issues/import?branch_id={a_id}", files=files).status_code == 403
        r = br_b.post("/api/bom-center/issues/import?dry_run=false", files=files)
        assert r.status_code == 400 and "ไฟล์นี้เป็นของสาขาอื่น" in r.json()["detail"]
        assert _rows(main.DB_PATH) == before
        # the centre's real import needs the owner token and a role with the module; the branch's own passes without
        r = mg.post(f"/api/bom-center/issues/import?dry_run=false&branch_id={a_id}", files=files)
        assert r.status_code == 403 and r.headers.get("X-Denied") == "token"
        r = br_a.post("/api/bom-center/issues/import?dry_run=false", files=files)
        assert r.status_code == 200 and r.json()["saved"] == 1 and r.json()["dates"][0]["result"] == "saved", r.text
        assert cleared == [1]
        sid = br_a.get(f"/api/bom-center/issues?branch_id={a_id}&date_from=2026-10-05&date_to=2026-10-05").json()["slips"][0]["slip_id"]
        log = br_a.get(f"/api/bom-center/issues/{sid}").json()["log"]
        assert [(x["action"], x["note"]) for x in log] == [("create", "นำเข้าจากไฟล์ ใบตัดเบิก.xlsx")]
        # confirming through the file, then the date is locked
        r = br_a.post("/api/bom-center/issues/import?dry_run=false&submit=true", files=files)
        assert r.status_code == 200 and r.json()["confirmed"] == 1 and r.json()["dates"][0]["late"] == 0, r.text
        assert cleared == [1, 1]
        r = br_a.post("/api/bom-center/issues/import?dry_run=false", files=files)
        assert r.status_code == 200 and r.json()["refused"] == 1
        assert r.json()["dates"][0]["message"] == "ใบตัดเบิกวันที่ 5/10/2026 ส่งแล้ว หากต้องแก้ไข แจ้งประสานงานขาย"
        assert [x["action"] for x in br_a.get(f"/api/bom-center/issues/{sid}").json()["log"]] == ["create", "confirm"]
        # the order and receipt routes are not swallowed by /order/{id}
        q = f"round=1&delivery=2026-10-07"
        r = br_a.get(f"/api/bom-center/order/template.xlsx?{q}")
        assert r.status_code == 200 and dl_name(r).startswith("แบบฟอร์ม_ใบสั่งซื้อ_BranchA_รอบ1_2026-10-07"), r.text[:200]
        owb = load_workbook(io.BytesIO(r.content))
        assert owb.sheetnames == [BF.ORDER_SHEET, BF.ORDER_REASONS_SHEET]
        assert br_a.get(f"/api/bom-center/order/template.xlsx?{q}&branch_id={b_id}").status_code == 403
        r = br_a.post(f"/api/bom-center/order/import?{q}", files={"file": ("ใบสั่ง.xlsx", body)})        # a slip file is not an order form
        assert r.status_code == 400 and "ไม่ใช่แบบฟอร์มใบสั่งซื้อ" in r.json()["detail"]
        r = br_a.post(f"/api/bom-center/order/import?{q}", files={"file": ("ใบสั่ง.xlsx", _book(owb))})   # nothing typed
        assert r.status_code == 200 and r.json()["ok"] is False and r.json()["errors"][0]["message"] == "ไม่พบจำนวนที่สั่งในไฟล์"
        assert br_a.get("/api/bom-center/order/999999/receive-template.xlsx").status_code == 404
        r = br_a.post("/api/bom-center/order/999999/receive-import", files={"file": ("รับ.xlsx", body)})
        assert r.status_code == 400
    finally:
        _restore(main.DB_PATH, snap)
        c = sqlite3.connect(str(main.DB_PATH))
        c.execute("DELETE FROM bc_ingredient WHERE code='X-SI'")
        c.commit()
        c.close()


def test_a_line_received_by_weight_is_received_in_the_app_not_from_a_file(con, monkeypatch):  # noqa: F811
    oid, q = _sent(con)
    real = OR.detail

    def weighted(c, order_id, receive=False):
        o = real(c, order_id, receive=receive)
        for ln in o["lines"]:
            if ln["code"] == CODE:
                ln["pieces_unit"] = "ตัว"
        return o
    monkeypatch.setattr(OR, "detail", weighted)
    wb = _rfile(con, oid)
    ws = wb[BF.RECEIVE_SHEET]
    assert "รับในแอป" in ws["A1"].value
    r = next(x for x in range(3, ws.max_row + 1) if ws.cell(x, 1).value == CODE)
    assert all(ws.cell(r, c).protection.locked for c in range(1, 8))
    out = _rimp(con, wb, oid)
    assert out["ok"] is False and out["errors"][0]["code"] == "weight" and "รับรายการนี้ในแอป" in out["errors"][0]["message"]
    before = _rstate(con)
    with pytest.raises(ValueError, match="รับรายการนี้ในแอป"):
        _rimp(con, wb, oid, dry_run=False, ack=True)
    assert _rstate(con) == before


def test_system_stock_is_never_negative_in_a_file_for_the_slip_and_the_order(con, monkeypatch):  # noqa: F811
    import copy
    assert BF.sys_text(-0.5, "แพ็ค") == "น้อยกว่า 1 แพ็ค" and BF.sys_text(0, "แพ็ค") == "น้อยกว่า 1 แพ็ค"
    assert BF.sys_text(0.4, "แพ็ค") == "น้อยกว่า 1 แพ็ค" and BF.sys_text(12.7, "แพ็ค") == "12.7 แพ็ค"
    build_stock(con)
    real = IS.sheet

    def neg(c, bid, d, now=None):
        sh = real(c, bid, d, now)
        for z in sh["zones"]:
            for it in z["items"]:
                it["system_qty"] = -0.45
        return sh
    monkeypatch.setattr(IS, "sheet", neg)
    wb, _ = _tmpl(con)
    ws = wb[BF.SLIP_SHEET]
    vals = [ws.cell(r, 6).value for r in range(3, ws.max_row + 1)]
    assert vals and all(v.startswith("น้อยกว่า 1 ") for v in vals)
    p = copy.deepcopy(_plan(con))
    for r in p["rows"]:
        r["system_qty"] = -62.9
    owb, _ = BF.order_template(con, p, 1, "1", WED)
    ows = owb[BF.ORDER_SHEET]
    assert all(str(ows.cell(r, 4).value).startswith("น้อยกว่า 1 ") for r in range(3, ows.max_row + 1))
