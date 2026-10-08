# -*- coding: utf-8 -*-
"""Month-end inventory workbook (accounting / Express "Inventory Checklists") → our count template.

The accounting files list every stock item with the store count and the cutting-board count
under THEIR item codes. Our system counts under OUR codes, so each source row is matched to
an ingredient of ours by NAME first (the codes are only a hint), then converted into a file the
system's own import accepts: ปิดงวดเดือน → ผลตรวจนับจากแบบฟอร์ม (dry run first, then save).

Matching order for every source row
  1. the conversion template (inventory_conversion_template.xlsx — rows a person confirmed with Y, or filled in by hand; always wins)
  2. the same name (spaces, case and punctuation ignored)
  3. the same code AND a similar name (≥ 0.60)  — a code reused for another item is NOT accepted
  4. the closest name (≥ 0.85, one clear winner)
  Anything else is listed as "ไม่พบ" for a person to map; nothing is guessed.

Outputs (in --out)
  count_<branch>_<date>.xlsx    sheet "ตรวจนับ" (the system's template format) + sheet "ตรวจสอบการจับคู่"
  inventory_conversion_template.xlsx   every match of this run as a fill-in sheet (source code → our code); correct it,
                                       put Y in the last column, and the next run uses it. Rows you confirmed are never overwritten.

    py -3.12 scripts/inventory_to_count.py "D:/69.08.31_Inventory-UBN.xlsx" --branch อุบลราชธานี --date 2026-09-01 --db <db>

The registry comes from a database copy (--db; default backend/oshinei.db). It writes nothing to
that database and never calls the live system.
"""
import argparse
import difflib
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))


def norm(s) -> str:
    s = str(s or "").lower().replace("\xa0", " ")
    return re.sub(r"[\s\-_/.,()\[\]*]+", "", s)


def sim(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a, b).ratio()


def same_sizes(a: str, b: str) -> bool:
    """Pack sizes and clothing sizes must agree: "2KG (6PACK)" is not "1KG (10PACK)", "SS" is not "S"."""
    nums = lambda t: sorted(re.findall(r"\d+(?:\.\d+)?", str(t)))
    size = lambda t: re.findall(r"size\s*(\w+)", str(t).lower())
    return nums(a) == nums(b) and sorted(size(a)) == sorted(size(b))


def same_garment_size(a: str, b: str) -> bool:
    """Same code and a similar name is accepted, but never across a different clothing size."""
    size = lambda t: re.findall(r"size\s*(\w+)", str(t).lower())
    return sorted(size(a)) == sorted(size(b))


def read_source(path: str) -> list:
    """Rows of the first sheet ("... Inventory Checklists"): code, name, store/board, total + its unit."""
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    out = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        code = str(r[2] or "").strip()
        if not re.match(r"^[0-9A-Za-z]{1,3}-\d{2,3}", code) and not re.match(r"^[A-Z]-\d+", code):
            continue
        if not r[4]:
            continue
        total = r[13] if len(r) > 13 else None
        try:
            total = float(total) if total not in (None, "") else 0.0
        except (TypeError, ValueError):
            total = 0.0
        unit = str(r[7] or r[5] or "").strip()
        out.append({"code": code, "type": r[3], "name": str(r[4]).strip(), "unit": unit, "qty": round(total, 4),
                    "store": r[6], "board": r[8]})
    return out


def load_aliases(path: Path) -> dict:
    """{source_code: (our_code, note)} from the conversion template: a row counts when a person
    ticked "ยืนยัน" (Y) or typed our code against a row the run could not match."""
    if not path.exists():
        return {}
    import openpyxl
    out = {}
    for r in openpyxl.load_workbook(path, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
        src, ours, method, ok = (str(r[0] or "").strip(), str(r[2] or "").strip(), str(r[4] or ""), str(r[6] or "").strip().lower() if len(r) > 6 else "")
        if src and ours and (ok in ("y", "yes", "ใช่", "1") or method.startswith("ไม่พบ")):
            out[src] = (ours, str(r[5] or "").strip())
    return out


def build_matches(rows: list, registry: dict, aliases: dict) -> list:
    by_name, names = {}, []
    for code, (name_th, name_en) in registry.items():
        for n in (name_th, name_en):
            k = norm(n)
            if k:
                by_name.setdefault(k, set()).add(code)
                names.append((k, code))
    out = []
    for r in rows:
        key, m = norm(r["name"]), {"our_code": None, "method": "ไม่พบ", "score": 0.0}
        if r["code"] in aliases and aliases[r["code"]][0] in registry:
            m = {"our_code": aliases[r["code"]][0], "method": "ตารางจับคู่ที่ยืนยันแล้ว", "score": 1.0}
        elif key in by_name and len(by_name[key]) == 1:
            m = {"our_code": next(iter(by_name[key])), "method": "ชื่อตรงกัน", "score": 1.0}
        elif r["code"] in registry and sim(key, norm(registry[r["code"]][0])) >= 0.60 and same_garment_size(
                r["name"], registry[r["code"]][0]):
            m = {"our_code": r["code"], "method": "รหัสตรงและชื่อใกล้เคียง", "score": round(sim(key, norm(registry[r["code"]][0])), 2)}
        else:
            best = sorted(((sim(key, k), c) for k, c in names if same_sizes(r["name"], registry[c][0]) or same_sizes(r["name"], registry[c][1])),
                          reverse=True)[:2]
            if best and best[0][0] >= 0.85 and (len(best) == 1 or best[0][0] - best[1][0] >= 0.03 or best[0][1] == best[1][1]):
                m = {"our_code": best[0][1], "method": "ชื่อใกล้เคียงที่สุด (ควรตรวจ)", "score": round(best[0][0], 2)}
        out.append({**r, **m})
    return out


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", help="the accounting inventory workbook (.xlsx)")
    ap.add_argument("--branch", required=True, help="our branch name, e.g. อุดรธานี")
    ap.add_argument("--date", required=True, help="count date to import as, YYYY-MM-DD (the 1st of the month)")
    ap.add_argument("--db", default=str(ROOT / "backend" / "oshinei.db"))
    ap.add_argument("--out", default=str(Path.home() / "Oshinei_inventory_import"))
    ap.add_argument("--aliases", help="the conversion template to read (default: <out>/inventory_conversion_template.xlsx)")
    a = ap.parse_args()
    os.environ["OSHINEI_DB"] = a.db
    import core
    from bom_center import stock as st
    from bom_center.units import Converter
    from openpyxl import Workbook

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    alias_path = Path(a.aliases) if a.aliases else out / "inventory_conversion_template.xlsx"
    aliases = load_aliases(alias_path)
    rows = read_source(a.source)
    with core.get_db() as con:
        bid = con.execute("SELECT branch_id FROM branches WHERE name=?", (a.branch,)).fetchone()
        if not bid:
            sys.exit(f"ไม่พบสาขา {a.branch}")
        bid = bid[0]
        registry = {c: (n or "", e or "") for c, n, e in con.execute("SELECT code, name_th, name_en FROM bc_ingredient")}
        sheet_items = {r["code"]: r for r in st.template_rows(con, bid)}
        conv = Converter(con)
        matches = build_matches(rows, registry, aliases)
        taken, lines = {}, []
        for m in matches:
            code = m["our_code"]
            m["status"], m["note"] = "ไม่พบรหัสของเรา — ต้องจับคู่เอง", ""
            if not code:
                continue
            if code in taken:
                m["status"], m["note"] = "ซ้ำกับแถวก่อนหน้า (ไม่นำเข้า)", f"ซ้ำกับ {taken[code]}"
                continue
            unit = m["unit"] or conv.base_unit(code)
            base = conv.to_base(code, m["qty"], unit)[0]
            if base is None:
                m["status"], m["note"] = "แปลงหน่วยไม่ได้ — ตั้งกฎหน่วยก่อน", f"หน่วย {unit or '-'}"
                continue
            taken[code] = m["code"]
            on_sheet = code in sheet_items
            m["status"] = "นำเข้า" if on_sheet else "นำเข้า (ไม่อยู่ในแบบฟอร์มของสาขา)"
            m["our_name"] = registry[code][0]
            lines.append({"zone": (sheet_items.get(code) or {}).get("zone", ""), "code": code, "name": registry[code][0],
                          "unit": unit, "qty": m["qty"], "per": (sheet_items.get(code) or {}).get("per_count_unit"),
                          "base": (sheet_items.get(code) or {}).get("base_unit"), "cls": (sheet_items.get(code) or {}).get("class"),
                          "req": (sheet_items.get(code) or {}).get("required"), "src": m["code"]})
        not_counted = [c for c in sheet_items if c not in taken]

    wb = Workbook()
    ws = wb.active
    ws.title = "ตรวจนับ"
    ws.append([f"ผลตรวจนับต้นงวด สาขา{a.branch} · วันที่ {a.date} · แปลงจากไฟล์บัญชี {Path(a.source).name} (จับคู่ด้วยชื่อ) · นำเข้าที่ ปิดงวดเดือน → ผลตรวจนับจากแบบฟอร์ม"])
    ws.append(["โซน", "รหัส", "ชื่อ", "หน่วยนับ", "จำนวนที่นับ", "1 หน่วยนับ = หน่วยฐาน", "หน่วยฐาน", "ชั้น", "รายการบังคับนับ", "หมายเหตุ"])
    for ln in lines:
        ws.append([ln["zone"], ln["code"], ln["name"], ln["unit"], ln["qty"], ln["per"], ln["base"], ln["cls"],
                   "บังคับ" if ln["req"] else "", f"จากรหัสบัญชี {ln['src']}"])
    ws.column_dimensions["C"].width = 44
    ws.freeze_panes = "A3"
    rv = wb.create_sheet("ตรวจสอบการจับคู่")
    rv.append(["รหัสในไฟล์บัญชี", "ชื่อในไฟล์บัญชี", "จำนวน", "หน่วย", "รหัสของเรา", "ชื่อของเรา", "วิธีจับคู่", "คะแนน", "ผล", "หมายเหตุ"])
    for m in matches:
        rv.append([m["code"], m["name"], m["qty"], m["unit"], m["our_code"], registry.get(m["our_code"], ("", ""))[0] if m["our_code"] else "",
                   m["method"], m["score"], m["status"], m["note"]])
    for col, w in zip("ABCDEFGHIJ", (16, 46, 10, 10, 14, 46, 28, 8, 34, 24)):
        rv.column_dimensions[col].width = w
    xlsx = out / f"count_{a.branch}_{a.date}.xlsx"
    wb.save(xlsx)

    tpl = out / "inventory_conversion_template.xlsx"
    old = {}
    if tpl.exists():
        import openpyxl
        for r in openpyxl.load_workbook(tpl, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
            if r[0]:
                old[str(r[0]).strip()] = list(r)
    tw = Workbook()
    t = tw.active
    t.title = "แปลงรหัส"
    t.append(["รหัสในไฟล์บัญชี", "ชื่อในไฟล์บัญชี", "รหัสของเรา (กรอก/แก้)", "ชื่อของเรา", "วิธีจับคู่", "หมายเหตุ", "ยืนยัน (ใส่ Y เมื่อตรวจแล้ว)"])
    fresh = {m["code"]: m for m in matches}
    for code, m in fresh.items():
        prev = old.pop(code, None)
        if prev and (str(prev[6] or "").strip().lower() in ("y", "yes", "ใช่", "1") or str(prev[4] or "").startswith("ไม่พบ") and prev[2]):
            t.append([prev[0], prev[1], prev[2], registry.get(str(prev[2]), ("", ""))[0], prev[4], prev[5], prev[6]])   # keep a person's work
        else:
            t.append([code, m["name"], m["our_code"] or "", registry.get(m["our_code"], ("", ""))[0] if m["our_code"] else "", m["method"], m["note"], ""])
    for prev in old.values():                                                         # rows from other months' files
        t.append(prev[:7] + [None] * (7 - len(prev[:7])))
    for col, w in zip("ABCDEFG", (16, 50, 20, 50, 30, 24, 16)):
        t.column_dimensions[col].width = w
    tw.save(tpl)

    cnt = {}
    for m in matches:
        cnt[m["method"]] = cnt.get(m["method"], 0) + 1
    print(f"สาขา{a.branch} · {len(rows)} แถวในไฟล์บัญชี → นำเข้า {len(lines)} รายการ")
    for k, v in sorted(cnt.items(), key=lambda x: -x[1]):
        print(f"  {k}: {v}")
    print(f"  ไม่พบ/แปลงไม่ได้/ซ้ำ: {sum(1 for m in matches if not m['status'].startswith('นำเข้า'))} (ดูแผ่น 'ตรวจสอบการจับคู่')")
    print(f"  รายการในแบบฟอร์มสาขาที่ไม่มีในไฟล์บัญชี: {len(not_counted)}")
    print("→", xlsx)


if __name__ == "__main__":
    main()
