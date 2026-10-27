# -*- coding: utf-8 -*-
"""The branch list the onboarding scripts share (PR-I, 30-branch rollout).

One row per branch, in an Excel or CSV file:

    สาขา        the branch name exactly as the system shows it (required)
    ตัวย่อ       the branch abbreviation, a-z 0-9 only: the company's 3-letter code (cnx, ubn, kr1 ...); the login becomes wh_<abbreviation>
    ชื่อผู้ใช้     the login name; blank = wh_<ตัวย่อ> (owner 2026-10-10); only for an exception
    LINE/ผู้รับ   who receives the handout (a LINE group, a person) - optional

With `--role franchise` the same file gives the franchise owners' logins own_<ตัวย่อ> (PR-L; give it the franchise
branches only). `branch_accounts.py` creates the logins from it and writes accounts_result.csv (same columns plus the
outcome); `branch_handout.py` reads the same file, or that result file, to write the messages and cards.
No password and no token is ever part of a list.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

H_BRANCH, H_ABBR, H_USER, H_RECIP, H_RESULT = "สาขา", "ตัวย่อ", "ชื่อผู้ใช้", "LINE/ผู้รับ", "ผล"
HEADERS = [H_BRANCH, H_ABBR, H_USER, H_RECIP]
RESULT_HEADERS = HEADERS + [H_RESULT, "หมายเหตุ"]
USER_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{2,39}$")
ABBR_RE = re.compile(r"^[a-z0-9]{2,12}$")
PREFIX = "wh_"                                   # every branch login is wh_<ตัวย่อ> (owner 2026-10-10)
OWNER_PREFIX = "own_"                            # the franchise owner's login is own_<ตัวย่อ> (PR-L, owner 2026-10-10)
ROLE_PREFIX = {"branch": PREFIX, "franchise": OWNER_PREFIX}

# the five pilot branches (owner 2026-10-22, opening count 1 Nov 2026) and their abbreviations: the company's 3-letter
# branch codes (owner 2026-10-27; the same codes as SHEET2BR in backend/routers/requisition.py and bulk_recon_upload.py)
PILOT_ABBRS = {
    "เชียงใหม่": "cnx",
    "อุบลราชธานี": "ubn",
    "อุดรธานี": "udn",
    "โคราช1": "kr1",
    "สุรินทร์": "srn",
}
PILOT_USERNAMES = {name: PREFIX + abbr for name, abbr in PILOT_ABBRS.items()}
# the four logins created on 2026-10-07 under the old naming (the old pilot: ubon udon khonkaen rayong);
# branch_accounts.py --replace-pilot retires exactly these (--replace-logins retires every branch login that is not the
# wanted wh_<ตัวย่อ>: these four and the long wh_* names of 2026-10-22 alike). ขอนแก่น and ระยอง are no longer pilot branches.
LEGACY_PILOT_USERNAMES = {"อุบลราชธานี": "branch_ubon", "อุดรธานี": "branch_udon",
                          "ขอนแก่น": "branch_khonkaen", "ระยอง": "branch_rayong"}

# result labels written by branch_accounts.py (column ผล)
R_CREATED, R_EXISTS, R_FAILED, R_PLAN = "สร้างแล้ว", "มีอยู่แล้ว", "ไม่สำเร็จ", "จะสร้าง"

_ALIASES = {
    "branch": ("สาขา", "branch", "ชื่อสาขา"),
    "abbr": ("ตัวย่อ", "ตัวย่อสาขา", "abbr", "abbreviation", "short"),
    "user": ("ชื่อผู้ใช้", "username", "user", "login"),
    "recip": ("line/ผู้รับ", "line", "ผู้รับ", "recipient", "ผู้รับ/line"),
    "result": ("ผล", "result", "status"),
}


def _norm(s) -> str:
    return re.sub(r"[\s_\-.·:()\[\]*]+", "", str(s or "").strip().lower())


def _grid(path: Path) -> list:
    p = Path(path)
    if not p.exists():
        raise SystemExit(f"ไม่พบไฟล์รายชื่อสาขา: {p}")
    if p.suffix.lower() in (".csv", ".txt"):
        raw = p.read_bytes()
        for enc in ("utf-8-sig", "utf-8", "cp874"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise SystemExit("อ่านไฟล์ CSV ไม่ได้ (ต้องเป็น UTF-8)")
        return [list(r) for r in csv.reader(text.splitlines())]
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("ต้องติดตั้ง openpyxl เพื่ออ่านไฟล์ Excel (py -3.12 -m pip install openpyxl) หรือใช้ไฟล์ .csv")
    wb = load_workbook(p, read_only=True, data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def read_list(path) -> list:
    """-> [{branch, abbr, username, recipient, result, line}] in file order. Raises SystemExit with a plain Thai message."""
    grid = _grid(path)
    head_at, cols = -1, {}
    for i, row in enumerate(grid[:15]):
        cells = [_norm(c) for c in row]
        m = {}
        for k, al in _ALIASES.items():
            for j, c in enumerate(cells):
                if c and c in [_norm(a) for a in al] and j not in m.values():
                    m[k] = j
                    break
        if "branch" in m:
            head_at, cols = i, m
            break
    if head_at < 0:
        raise SystemExit(f"ไม่พบหัวตารางในไฟล์ (ต้องมีคอลัมน์ “{H_BRANCH}”) ใช้แม่แบบจาก branch_handout.py --template")
    out, seen, abbr_at, problems = [], {}, {}, []
    for i, row in enumerate(grid[head_at + 1:], start=head_at + 2):
        def cell(k):
            j = cols.get(k)
            return str(row[j]).strip() if j is not None and j < len(row) and row[j] is not None else ""
        branch = cell("branch")
        if not branch:
            continue
        user = cell("user")
        abbr = cell("abbr").lower()
        if abbr and not ABBR_RE.match(abbr):
            problems.append(f"แถวที่ {i}: ตัวย่อ “{abbr}” ใช้ได้เฉพาะ a-z 0-9 ยาว 2–12 ตัวอักษร")
        if user and not USER_RE.match(user):
            problems.append(f"แถวที่ {i}: ชื่อผู้ใช้ “{user}” ใช้ได้เฉพาะ a-z 0-9 _ . - ยาว 3–40 ตัวอักษร")
        if branch in seen:
            problems.append(f"แถวที่ {i}: สาขา “{branch}” ซ้ำกับแถวที่ {seen[branch]}")
        seen[branch] = i
        if abbr and abbr in abbr_at:
            problems.append(f"แถวที่ {i}: ตัวย่อ “{abbr}” ซ้ำกับแถวที่ {abbr_at[abbr]} (หนึ่งตัวย่อต่อหนึ่งสาขา)")
        abbr_at.setdefault(abbr, i)
        out.append({"branch": branch, "abbr": abbr, "username": user, "recipient": cell("recip"),
                    "result": cell("result"), "line": i})
    names = [login_of(r) for r in out if login_of(r)]
    for u in sorted({u for u in names if names.count(u) > 1}):
        problems.append(f"ชื่อผู้ใช้ “{u}” ซ้ำในหลายสาขา (ตรวจคอลัมน์ ตัวย่อ และ ชื่อผู้ใช้)")
    if problems:
        raise SystemExit("ไฟล์รายชื่อสาขามีข้อผิดพลาด:\n  " + "\n  ".join(problems))
    if not out:
        raise SystemExit("ไฟล์รายชื่อสาขาไม่มีรายการ")
    return out


def login_of(row: dict, role: str = "branch") -> str:
    """The login a list row names without any database: its own username, else wh_<ตัวย่อ>, else the pilot name; '' if none.
    For role franchise only an own_* username counts (the column normally holds a branch's wh_* exception), else own_<ตัวย่อ>."""
    user = row.get("username") or ""
    if user and (role != "franchise" or user.startswith(OWNER_PREFIX)):
        return user
    abbr = (row.get("abbr") or PILOT_ABBRS.get(row.get("branch"), "")).lower()
    return ROLE_PREFIX.get(role, PREFIX) + abbr if abbr else ""


def username_for(branch: str, branch_id=None, abbr: str = "", role: str = "branch"):
    """The login name for a branch whose list row leaves it blank: wh_<ตัวย่อ> (own_<ตัวย่อ> for the franchise owner);
    the five pilot branches have their abbreviation built in; a branch with no abbreviation gets wh_<branch id>. None
    when none of these is known (the caller must say so: branch names are Thai, so there is no ASCII name to fall back
    on, and two scripts guessing differently would hand a branch a login that does not exist)."""
    prefix = ROLE_PREFIX.get(role, PREFIX)
    abbr = (abbr or PILOT_ABBRS.get(branch, "")).strip().lower()
    if abbr:
        return prefix + abbr
    if branch_id is not None:
        return f"{prefix}{int(branch_id)}"
    return None


def write_template(path) -> Path:
    """An example list: the five pilot branches with their abbreviations, a note row on top."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "รายชื่อสาขา"
    ws.append(["กรอกหนึ่งสาขาต่อหนึ่งแถว · ชื่อสาขาต้องตรงกับที่ระบบแสดง · ตัวย่อใช้ a-z 0-9 ยาว 2–12 ตัวอักษร ระบบตั้งชื่อผู้ใช้เป็น wh_<ตัวย่อ> · "
               "เว้นชื่อผู้ใช้ว่างไว้ (กรอกเฉพาะกรณีพิเศษ) · ช่อง LINE/ผู้รับใช้จดว่าจะส่งข้อความให้ใคร (ไม่บังคับ) · ห้ามใส่รหัสผ่านในไฟล์นี้"])
    ws["A1"].font = Font(italic=True, color="5F6C80")
    ws.append(HEADERS)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A6B")
    for name, abbr in PILOT_ABBRS.items():
        ws.append([name, abbr, "", ""])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 30
    ws.freeze_panes = "A3"
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def write_result(path, rows) -> Path:
    """accounts_result.csv: the list columns plus the outcome (UTF-8 with BOM so Excel opens Thai text)."""
    p = Path(path)
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(RESULT_HEADERS)
        for r in rows:
            w.writerow([r.get("branch", ""), r.get("abbr", ""), r.get("username", ""), r.get("recipient", ""),
                        r.get("result", ""), r.get("note", "")])
    return p
